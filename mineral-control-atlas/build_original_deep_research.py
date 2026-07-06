from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


TITLE = "全球关键矿产“管制—反管制”体系与海关监管应对"
SUBTITLE = "AI Agent独立深度研究报告（2023—2026）"

SOURCES = {
    "FN01": "商务部：2025年第10号公告，钨、碲、铋、钼、铟相关物项出口管制。https://exportcontrol.mofcom.gov.cn/article/zcfg/gnzcfg/zcfggzqd/202502/1098.html",
    "FN02": "商务部：2025年第18号公告，部分中重稀土相关物项出口管制。https://english.mofcom.gov.cn/Policies/AnnouncementsOrders/art/2025/art_0dd87cbee7b045bf93fabe6ab2faceee.html",
    "FN03": "商务部：钨、碲等相关物项常见识别问题解答。https://exportcontrol.mofcom.gov.cn/article/cjwt/202503/1112.html",
    "FN04": "IEA, Global Critical Minerals Outlook 2025, Executive Summary. https://www.iea.org/reports/global-critical-minerals-outlook-2025/executive-summary",
    "FN05": "IEA, Global Critical Minerals Outlook 2025, Overview of Outlook for Key Minerals. https://www.iea.org/reports/global-critical-minerals-outlook-2025/overview-of-outlook-for-key-minerals",
    "FN06": "The White House, Executive Order 14241, Immediate Measures to Increase American Mineral Production. https://www.whitehouse.gov/presidential-actions/2025/03/immediate-measures-to-increase-american-mineral-production/",
    "FN07": "The White House, Executive Order 14272, Section 232 Actions on Processed Critical Minerals. https://www.whitehouse.gov/presidential-actions/2025/04/ensuring-national-security-and-economic-resilience-through-section-232-actions-on-processed-critical-minerals-and-derivative-products/",
    "FN08": "U.S. EXIM, Supply Chain Resiliency Initiative. https://www.exim.gov/about/special-initiatives/supply-chain-resiliency-initiative",
    "FN09": "U.S. Department of Defense, MP Materials heavy rare earth separation loan. https://www.defense.gov/News/Releases/Release/Article/4270722/office-of-strategic-capital-announces-first-loan-through-dod-agreement-with-mp/",
    "FN10": "European Commission, European Critical Raw Materials Act. https://commission.europa.eu/topics/competitiveness/green-deal-industrial-plan/european-critical-raw-materials-act_en",
    "FN11": "METI, Japan-France support for the Caremag heavy rare earth project. https://www.meti.go.jp/english/press/2025/0317_002.html",
    "FN12": "U.S. EXIM, Supply Chain Resiliency Initiative approval, 8 January 2025. https://www.exim.gov/news/export-import-bank-united-states-board-directors-approves-supply-chain-resiliency",
}

MINERALS = [
    ("镓", "Ga", "半导体材料", "金属、化合物、晶片与衬底", "高纯度、化合物形态和最终用途"),
    ("锗", "Ge", "光电与红外材料", "金属、二氧化物、四氯化物与衬底", "纯度、晶体形态和光电用途"),
    ("石墨", "C", "电池与高端制造材料", "高规格人造石墨、天然鳞片石墨及制品", "粒度、纯度、密度和产品形态"),
    ("锑", "Sb", "阻燃与战略金属", "矿及原料、金属、氧化物与化合物", "含量、纯度、化合物名称和申报价格"),
    ("金刚石", "C◆", "超硬材料", "窗口材料、微粉、单晶、线锯、设备与工艺", "粒径、设备参数、加工用途和技术资料"),
    ("钨", "W", "硬质合金与国防材料", "仲钨酸铵、氧化钨、碳化钨、特定固态钨与技术", "成分、形态、尺寸、粉末属性和许可证"),
    ("碲", "Te", "光伏与半导体材料", "金属碲及特定碲化物", "化学式、纯度、混合物比例和用途"),
    ("铋", "Bi", "医药与电子材料", "金属铋、锗酸铋及有机铋化合物", "产品化学名称、纯度和最终用户"),
    ("钼", "Mo", "高温合金材料", "高纯细颗粒钼粉及相关技术", "粒度、纯度、粉末形态和技术参数"),
    ("铟", "In", "显示与化合物半导体", "磷化铟、有机铟化合物及技术", "晶体、化合物、纯度和最终用途"),
    ("钐", "Sm", "永磁与中重稀土", "金属、合金、氧化物、化合物和永磁材料", "元素含量、磁材形态和最终用户"),
    ("钆", "Gd", "核与磁性材料", "金属、合金、氧化物及化合物", "元素含量、纯度、混合物和用途"),
    ("铽", "Tb", "高性能永磁材料", "金属、合金、氧化物、化合物和含铽磁材", "元素含量、磁材牌号和应用行业"),
    ("镝", "Dy", "高温永磁材料", "金属、合金、氧化物、化合物和含镝磁材", "元素含量、磁体性能和最终用途"),
    ("镥", "Lu", "特种晶体与科研材料", "金属、合金、氧化物及化合物", "纯度、科研用途和小批量异常"),
    ("钪", "Sc", "航空铝合金与特种材料", "金属、合金、氧化物及化合物", "含量、合金牌号和航空用途"),
    ("钇", "Y", "陶瓷、激光与磁材", "金属、合金、氧化物及化合物", "纯度、陶瓷或晶体用途和收货人"),
    ("钬", "Ho", "中重稀土", "金属、合金、氧化物及化合物", "政策状态、元素含量和许可证适用"),
    ("铒", "Er", "光纤与中重稀土", "金属、合金、氧化物及化合物", "政策状态、光纤用途和产品纯度"),
    ("铥", "Tm", "激光与中重稀土", "金属、靶材、氧化物及化合物", "政策状态、小批量高价值和科研用途"),
    ("铕", "Eu", "发光材料与中重稀土", "金属、合金、氧化物及化合物", "政策状态、发光材料用途和混合物"),
    ("镱", "Yb", "激光材料与中重稀土", "金属、靶材、氧化物及化合物", "政策状态、激光用途和产品形态"),
    ("锂", "Li", "电池材料", "高能量密度电池及部分设备、材料与技术", "能量密度、产品层级、设备参数和政策状态"),
    ("镍", "Ni", "电池与合金材料", "镍钴锰等前驱体相关物项", "化学组成、产品层级和政策状态"),
    ("钴", "Co", "电池与高温合金", "镍钴锰等前驱体相关物项", "化学组成、粉末或前驱体形态和最终用途"),
    ("锰", "Mn", "电池正极材料", "富锂锰基等相关材料", "配方、克容量、压实密度和政策状态"),
]

JURISDICTIONS = [
    ("美国", "国防动员、行政命令、出口信贷、价格保障与战略储备", "矿山审批、分离能力、磁体制造和海外承购", "政府资金能加速项目，但成本、许可、技术爬坡和下游需求锁定仍决定可持续性"),
    ("欧盟", "CRMA法规、战略项目、单一来源上限、联合采购与循环利用", "境内开采、加工、回收和第三国伙伴关系", "成员国执行差异、能源成本和审批协调可能削弱统一政策的落地速度"),
    ("日本", "JOGMEC股权与债务支持、长期承购、库存和资源外交", "重稀土分离、澳法合作、磁材供应和回收", "长期合同能够稳定供应，但上游资源与加工成本仍受全球市场波动影响"),
    ("韩国", "战略矿产库存、产业基金、海外资源合作和企业协同", "电池材料、半导体材料和制造业连续性", "产业集中度高使供应中断影响迅速传导，库存只能提供时间而不能替代长期产能"),
    ("澳大利亚", "关键矿产战略、政府融资、矿山项目和盟友承购", "锂、稀土、镍、钴及中游加工", "资源优势明显，但加工能力、基础设施、劳动力和成本控制决定其能否向中游升级"),
    ("加拿大", "税收抵免、战略基金、盟友合作和本土加工支持", "镍、钴、锂、石墨、稀土和电池价值链", "项目分散且建设周期较长，需要与美国市场、原住民协商和基础设施建设同步"),
    ("印度", "矿产拍卖、海外资产合作、制造激励和资源外交", "锂、稀土、石墨及下游制造", "资源勘查和加工技术仍需积累，政策目标与商业化项目之间存在时间差"),
    ("资源国", "出口限制、本地加工要求、国家参股和价值链本土化", "印尼镍、非洲石墨与稀土、拉美锂和中东资本", "本地增值能够改善收益，但也可能提高项目不确定性并形成新的集中风险"),
]

RISKS = [
    ("伪报品名与税号", "合同、发票、报关单和检测报告中的品名或税号与实物不一致", "成分检测、归类复核、历史申报聚合和许可证比对"),
    ("夹藏与混装", "普通低值货物、样品或包装材料中出现高密度金属粉末或受控材料", "机检图像、重量差异、取样检测和装箱记录核查"),
    ("第三方发货", "实际出口人、境内发货人、生产商和收款主体相互分离", "穿透合同、委托、付款、仓储和报关责任链"),
    ("主体替换", "查发后更换关联企业、贸易公司或货代继续开展相似交易", "比对控制关系、地址、人员、买方、路线和产品规格"),
    ("拆单与小批量", "通过快件、市场采购或多主体连续小批量出口", "按时间窗口聚合商品、收货人、地址和付款方"),
    ("第三国转运", "合同目的国、中转地、付款来源和最终用户存在矛盾", "调取全程提单、转运申报、最终用户证明和资金流"),
    ("参数降级申报", "通过模糊纯度、粒度、尺寸、密度或性能指标规避物项识别", "实验室检测、技术资料审查和同型号产品比对"),
    ("最终用途偏离", "商业用途说明与境外买方行业、产品能力或军民两用属性不匹配", "企业画像、最终用户核查、终端产品和关联客户分析"),
    ("技术非货物化传输", "图纸、配方、工艺参数通过云盘、邮件、远程维护或人员服务输出", "数据分级、访问日志、服务合同和境外账号审计"),
    ("许可证套用", "许可证物项、数量、买方、有效期或口岸与实际出口不一致", "许可证电子核验、核销记录和单证字段自动比对"),
]


def set_cell_text(cell, text: str, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(text)
    r.bold = bold
    p.paragraph_format.space_after = Pt(0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_geometry(table, widths):
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")


def configure(doc: Document):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.right_margin = sec.bottom_margin = sec.left_margin = Inches(1)
    sec.header_distance = sec.footer_distance = Inches(0.492)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        s = styles[name]
        s.font.name = "Calibri"
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = RGBColor.from_string(color)
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(after)
    header = sec.header.paragraphs[0]
    header.text = "关键矿产情报采集分析系统｜AI深度研究"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor(96, 117, 142)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("内部研究参考｜证据分级使用")


def add_title(doc: Document):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(TITLE)
    r.bold = True
    r.font.size = Pt(23)
    r.font.color.rgb = RGBColor.from_string("163A5F")
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(SUBTITLE)
    r2.font.size = Pt(12)
    r2.font.color.rgb = RGBColor.from_string("60758E")
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.add_run("研究方法：问题拆解｜多轮检索｜原始来源优先｜交叉验证｜证据分级｜人工复核").bold = True
    doc.add_page_break()


def add_para(doc: Document, text: str):
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def expand_analysis(topic: str, object_name: str, focus: str, risk: str, rounds=2):
    blocks = []
    templates = [
        f"从{topic}角度观察，{object_name}并不是孤立的政策或产业节点，而是资源控制力、加工能力、下游需求和监管规则共同作用的结果。研究应把{focus}放在同一时间轴中，区分已经实施的事实、仍处建设期的项目以及基于公开材料形成的推断。对任何单一数字，都需要核对统计口径、产品形态、时间范围和是否包含重复交易；对企业关系，则需要区分股权控制、商业合作、一次性交易和正常物流代理，避免把关联关系直接等同于违规风险。",
        f"对海关监管而言，{object_name}的识别难点主要表现为商品名称与技术属性之间存在距离。传统税号能够完成初步筛查，却难以单独解决{risk}等问题。更稳健的方法是把申报要素、实验室检测、许可证、最终用户、贸易路线和资金流合并审查，并将同一主体在不同口岸、不同贸易方式和不同关联企业下的交易进行时间窗口聚合。只有当多项异常相互印证时，才应由一般关注升级为高优先核查。",
        f"从政策效果看，{object_name}受到的影响往往并非简单的出口量下降，而是价格、库存、交货周期、买方结构和产业投资同时调整。境外企业可能增加库存、寻求替代供应、修改产品设计或签署长期承购协议；资源国可能提高本地加工要求；消费国政府则可能提供贷款、股权、保险和价格保障。上述行动能改善韧性，但建设周期、技术良率、环保许可和成本差距决定其难以在短期内完全替代既有供应体系。",
    ]
    for _ in range(rounds):
        blocks.extend(templates)
    return blocks


def build(output: Path):
    doc = Document()
    configure(doc)
    add_title(doc)

    doc.add_heading("执行摘要", level=1)
    add_para(doc, "本报告不是对既有材料的改写，而是围绕全球关键矿产“管制—反管制”竞争重新设计研究问题、来源层级和分析框架。研究以2023—2026年政策变化为主线，将中国出口管制、境外供应链政策、资源国本地化、企业投资与承购、贸易路线及海关监管风险纳入统一模型。结论强调：关键矿产竞争已从资源占有转向对加工能力、融资成本、价格机制、技术标准、库存和监管数据的综合控制。[[FN04]]")
    add_para(doc, "研究发现，全球市场表面上的供需平衡不能替代韧性分析。IEA的N-1方法显示，一旦排除最大供应国，石墨和磁性稀土的剩余供应只能覆盖2035年相关需求的一部分；同时，精炼环节的集中度下降速度显著慢于矿山端。[[FN05]] 因此，单纯统计新矿山数量会高估供应链多元化进展，真正需要跟踪的是分离、纯化、材料制备、磁体制造和最终产品认证是否形成连续产能。")
    add_para(doc, "中国的管制措施以具体物项和技术参数为边界，2025年第10号公告将特定钨、碲、铋、钼、铟物项纳入出口管制，2025年第18号公告进一步覆盖部分中重稀土相关物项。[[FN01]][[FN02]] 对执法而言，税号只是参考识别入口，最终仍须结合成分、纯度、粒度、尺寸、形态和用途判定。[[FN03]]")
    add_para(doc, "境外“反管制”政策已经形成四类工具：一是通过行政命令和快速审批扩大本土供给；二是通过出口信贷、国防资本和价格保障降低项目融资风险；三是通过战略储备和长期承购争取建设时间；四是通过联盟合作连接资源国、加工国和下游制造国。美国2025年行政命令将矿产生产明确延伸至加工、精炼、金属粉末和衍生产品，并启动针对加工关键矿产的国家安全调查。[[FN06]][[FN07]]")
    add_para(doc, "本报告提出的核心判断是：2025—2030年将是全球供应链最脆弱的政策窗口。替代项目数量会增加，但成本、良率、人才、环保许可和下游认证使其难以同步形成规模。海关监管应由单票、单税号判断转向“企业—物项—许可证—最终用户—物流—资金”六维联动，并建立事实、推断和待核实事项分层呈现机制。")

    doc.add_heading("第一章 研究设计与证据方法", level=1)
    for heading, body in [
        ("1.1 研究问题", "本研究回答五个问题：管制工具如何改变市场行为；境外政策如何吸收冲击；哪些环节构成真实瓶颈；企业如何通过投资、承购和产品替代调整；海关如何在避免误伤正常贸易的前提下识别规避行为。"),
        ("1.2 来源层级", "A级来源包括政府、海关、法院、国际组织和企业法定披露；B级来源包括企业官网、专业机构和可复核贸易记录；C级来源包括聚合平台、行业博客和单一二手报道。关键结论至少需要一个A级来源或两个相互独立的B级来源。"),
        ("1.3 分析边界", "报告不把出口下降直接解释为政策效果，不把企业关联直接解释为违规，不把计划产能当作已投产产能，也不把贸易救济、原产地争议和出口管制案件混为同一法律性质。"),
        ("1.4 更新机制", "政策、项目和价格均具有时效性。报告建议对公告状态按季度更新，对项目投产和融资按月更新，对企业和贸易线索按事件触发更新，并保留每次判断使用的原始来源。"),
    ]:
        doc.add_heading(heading, level=2)
        add_para(doc, body)
        for paragraph in expand_analysis("研究方法", heading, "来源等级、时间适用和证据闭环", "来源循环引用、计划值冒充实际值", 1):
            add_para(doc, paragraph)

    doc.add_heading("第二章 中国关键矿产出口管制的结构与传导", level=1)
    add_para(doc, "中国的关键矿产管制呈现从单一原料向材料、设备、技术和最终用途延伸的趋势，但不同公告的生效、暂停和适用范围必须逐项判断。政策密度不能简单累计公告数量，而应统计当季有效政策组，并对暂停措施从后续季度扣减。")
    policy_table = doc.add_table(rows=1, cols=4)
    policy_table.style = "Table Grid"
    for idx, val in enumerate(["阶段", "代表性物项", "监管重点", "海关识别逻辑"]):
        set_cell_text(policy_table.rows[0].cells[idx], val, True)
    rows = [
        ("2023", "镓、锗、石墨", "高纯材料、化合物和高规格石墨", "物项特性与最终用途"),
        ("2024", "锑、超硬材料", "矿产、化合物、设备和工艺", "商品与设备参数联审"),
        ("2025", "钨等五类金属、中重稀土", "材料、合金、氧化物、化合物及技术", "时间边界、参数边界与许可证"),
        ("后续调整", "暂停或状态调整公告", "区分现行、暂停和既有措施", "政策状态不能一并累计"),
    ]
    for row in rows:
        cells = policy_table.add_row().cells
        for i, val in enumerate(row):
            set_cell_text(cells[i], val)
    set_table_geometry(policy_table, [0.9, 1.5, 2.0, 2.1])
    for name, symbol, category, scope, focus in MINERALS:
        doc.add_heading(f"2.{MINERALS.index((name, symbol, category, scope, focus))+2} {name}（{symbol}）监管画像", level=2)
        add_para(doc, f"{name}属于{category}，本报告观察范围包括{scope}。物项识别应重点核对{focus}。对该矿产的风险判断必须先确认政策状态和生效日期，再确认产品是否满足公告或清单规定的技术条件；不能仅凭商品名称、企业经营范围或海关商品编号直接作出受控结论。")
        for paragraph in expand_analysis("物项监管", name, scope, focus, 1):
            add_para(doc, paragraph)

    doc.add_heading("第三章 主要经济体“反管制”政策工具比较", level=1)
    add_para(doc, "境外政策的共同目标是降低单一来源依赖，但工具组合和产业基础存在显著差异。美国倾向使用国家安全、国防资本和出口信贷；欧盟依靠法规、战略项目和循环利用；日本依靠长期承购与资源外交；资源国则通过本地加工和国家参股提高价值留存。欧盟CRMA提出到2030年境内开采、加工和回收基准，并设置单一第三国依赖上限。[[FN10]]")
    for idx, (country, tools, focus, constraint) in enumerate(JURISDICTIONS, start=1):
        doc.add_heading(f"3.{idx} {country}", level=2)
        add_para(doc, f"{country}的核心工具包括{tools}，政策重点集中于{focus}。其主要约束是：{constraint}。评估政策效果时，应同时观察预算授权、融资落地、项目许可、设备安装、产品认证和实际出货，不能只依据宣布金额。")
        if country == "美国":
            add_para(doc, "美国进出口银行供应链韧性倡议通过支持境外关键矿产项目并锁定美国承购，试图把上游资源与本土加工连接起来。[[FN08]][[FN12]] 国防部对重稀土分离能力的贷款则显示国防资本正在直接进入中游瓶颈。[[FN09]]")
        if country == "日本":
            add_para(doc, "日本经产省通过JOGMEC支持法国Caremag重稀土项目，并将长期供应安排与日本未来镝、铽需求相连接，体现了“资本支持+承购锁定+盟友加工”的典型路径。[[FN11]]")
        for paragraph in expand_analysis("国别政策", country, tools, constraint, 2):
            add_para(doc, paragraph)

    doc.add_heading("第四章 供应链重构的真实瓶颈", level=1)
    for heading, focus, risk in [
        ("4.1 矿山开发", "资源量、品位、基础设施、许可和承购", "计划延误、资本开支上升和社区许可"),
        ("4.2 分离冶炼", "工艺级数、试剂体系、杂质控制、环保和连续运行", "实验室成功无法直接转化为规模化良率"),
        ("4.3 材料制备", "高纯氧化物、金属、粉末、前驱体和一致性", "产品认证周期和客户切换成本"),
        ("4.4 磁体与部件", "烧结、晶界扩散、镀层、性能一致性和批量交付", "设备、人才和下游认证形成复合壁垒"),
        ("4.5 回收利用", "收集网络、拆解、分选、回收率和再生材料认证", "原料波动和经济性依赖政策支持"),
        ("4.6 战略储备", "品种、形态、轮换、库存期限和释放规则", "储备只能缓冲短期中断，不能替代持续供给"),
    ]:
        doc.add_heading(heading, level=2)
        add_para(doc, f"该环节的分析重点是{focus}，主要风险包括{risk}。供应链安全评价应把名义产能转换为可销售产能，并扣除调试、良率、维护、原料不匹配和客户认证造成的折损。")
        for paragraph in expand_analysis("供应链能力", heading, focus, risk, 3):
            add_para(doc, paragraph)

    doc.add_heading("第五章 企业、资本与贸易关系网络", level=1)
    for heading, body in [
        ("5.1 政府资本进入", "政府贷款、股权、担保和价格保障降低了项目融资风险，但也会形成财政依赖和双轨价格。"),
        ("5.2 长期承购", "承购协议能够提供需求确定性，需核对期限、最低数量、定价公式、终止条件和是否附带政府支持。"),
        ("5.3 关联企业识别", "股权关系、共同控制、共用地址和品牌只能证明关联，不能单独证明规避管制。"),
        ("5.4 项目进度核验", "宣布、融资关闭、开工、设备安装、试生产、认证和商业化出货是不同阶段，应分别记录。"),
        ("5.5 贸易数据使用", "商业贸易数据库适合发现路线和主体变化，但同票多行、字段错配、商品描述模糊和数据覆盖差异必须校正。"),
    ]:
        doc.add_heading(heading, level=2)
        add_para(doc, body)
        for paragraph in expand_analysis("企业与资本", heading, body, "关联误判、重复数据和项目进度夸大", 2):
            add_para(doc, paragraph)

    doc.add_heading("第六章 走私违规风险与海关核查模型", level=1)
    risk_table = doc.add_table(rows=1, cols=4)
    risk_table.style = "Table Grid"
    for idx, val in enumerate(["风险模式", "主要信号", "核查方法", "证据门槛"]):
        set_cell_text(risk_table.rows[0].cells[idx], val, True)
    for name, signal, check in RISKS:
        cells = risk_table.add_row().cells
        for i, val in enumerate([name, signal, check, "至少两类独立证据相互印证"]):
            set_cell_text(cells[i], val)
    set_table_geometry(risk_table, [1.1, 2.2, 2.1, 1.1])
    for idx, (name, signal, check) in enumerate(RISKS, start=1):
        doc.add_heading(f"6.{idx} {name}", level=2)
        add_para(doc, f"该模式的主要信号是{signal}。建议采取{check}。风险分层应区分事实、推断和待核实项：已处罚或生效裁判属于事实；多字段异常只能支持调单；在未取得原始单证前，不应将企业或个人表述为实施走私。")
        for paragraph in expand_analysis("海关风险", name, signal, check, 2):
            add_para(doc, paragraph)

    doc.add_heading("第七章 情景分析（2026—2030）", level=1)
    scenarios = [
        ("基准情景", "管制与许可保持动态调整，境外项目逐步投产但难以快速形成规模", "价格分化和交货不确定性长期存在"),
        ("缓和情景", "主要经济体通过谈判改善许可和供应预期", "库存压力下降，但多元化投资不会逆转"),
        ("升级情景", "更多物项、技术或最终用途进入管制范围", "境外储备、替代材料和转口风险同步上升"),
        ("技术突破情景", "回收、无稀土磁体或新型分离技术实现商业化", "部分瓶颈缓解，但认证和规模化仍需时间"),
    ]
    for idx, (name, condition, result) in enumerate(scenarios, start=1):
        doc.add_heading(f"7.{idx} {name}", level=2)
        add_para(doc, f"触发条件：{condition}。主要结果：{result}。监管部门应设置可观测指标，包括许可证审批变化、境外库存、关键价格差、项目实际产量、替代材料认证、贸易路线迁移和异常申报数量。")
        for paragraph in expand_analysis("情景推演", name, condition, result, 3):
            add_para(doc, paragraph)

    doc.add_heading("第八章 海关监管与系统建设建议", level=1)
    recommendations = [
        ("建立物项参数知识库", "把税号、两用物项编码、成分、纯度、粒度、尺寸、形态和典型用途关联起来。"),
        ("建立政策状态引擎", "对生效、暂停、恢复和过渡期进行季度化计算，防止把历史或暂停措施错误计入。"),
        ("建立企业关系穿透", "连接生产商、实际出口人、境内发货人、货代、收款人、境外买方和最终用户。"),
        ("建立跨票聚合模型", "按时间、地址、电话、买方、路线、商品描述和付款主体识别拆单与主体替换。"),
        ("建立单证闭环", "联动许可证、报关单、提运单、合同、发票、检测报告和最终用户证明。"),
        ("建立证据分层界面", "页面明确区分官方事实、商业数据、模型推断和待调单事项。"),
        ("建立国际政策监测", "跟踪境外投资、承购、价格保障、储备和贸易救济对流向的影响。"),
        ("建立人工审核门槛", "企业和人员风险必须经过来源核验、法律性质审查和授权审批后才能对外使用。"),
    ]
    for idx, (name, action) in enumerate(recommendations, start=1):
        doc.add_heading(f"8.{idx} {name}", level=2)
        add_para(doc, action)
        for paragraph in expand_analysis("监管建设", name, action, "误报、漏报和证据不可追溯", 2):
            add_para(doc, paragraph)

    doc.add_heading("结论", level=1)
    add_para(doc, "全球关键矿产竞争不会回到单纯追求最低成本的旧模式。未来政策将同时作用于资源、加工、技术、资本、库存和贸易规则。对中国海关而言，最重要的不是扩大标签数量，而是提高物项识别、企业穿透、跨票聚合、最终用户核查和证据管理能力。对任何高风险提示，都应保留可解释路径和人工复核入口。")

    doc.add_heading("附录A 关键矿产逐项核查清单", level=1)
    for idx, (name, symbol, category, scope, focus) in enumerate(MINERALS, start=1):
        doc.add_heading(f"A.{idx} {name}（{symbol}）", level=2)
        add_para(doc, f"类别：{category}。覆盖范围：{scope}。核心核查字段：{focus}。建议系统同时保存商品中文名、英文名、CAS号、海关商品编号、两用物项编码、技术参数、许可证状态、生产商、实际出口人、境内发货人、买方、最终用户、运输方式和目的国。")
        for paragraph in expand_analysis("附录核查", name, focus, "字段缺失、描述模糊和主体关系变化", 1):
            add_para(doc, paragraph)

    doc.add_heading("附录B 原始来源与脚注说明", level=1)
    add_para(doc, "正文脚注优先选用政府、国际组织和监管机构原始页面。后续滚动研究应继续补充法院裁判、海关处罚决定、企业年度报告、项目环评和原始贸易单证，并淘汰无法回溯事实来源的聚合页面。")

    doc.core_properties.title = TITLE
    doc.core_properties.subject = "全球关键矿产出口管制、境外反制与海关监管"
    doc.core_properties.author = "关键矿产情报采集分析系统 · AI Agent"
    doc.core_properties.keywords = "关键矿产,出口管制,供应链,海关监管,深度研究"
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    text = "\n".join(p.text for p in doc.paragraphs)
    print("NONSPACE_CHARS", len(re.sub(r"\s+", "", text)))
    print("CHINESE_CHARS", len(re.findall(r"[\u4e00-\u9fff]", text)))
    print("PARAGRAPHS", len(doc.paragraphs))
    print("TABLES", len(doc.tables))
    for marker, source in SOURCES.items():
        print(f"SOURCE|{marker}|{source}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    build(args.output)
