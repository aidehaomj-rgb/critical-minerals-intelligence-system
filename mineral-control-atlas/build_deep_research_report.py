from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Mm, Pt


TITLE = "全球关键矿产“管制—反管制”体系与海关监管应对"
SUBTITLE = "AI Agent深度研究专题报告｜研究周期：2023—2026年｜形成时间：2026年7月"


def add_point(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    label, separator, body = text.partition("：")
    if separator:
        paragraph.add_run(label + separator).bold = True
        paragraph.add_run(body)
    else:
        paragraph.add_run(text)


def build(source: Path, output: Path) -> None:
    doc = Document(str(source))

    doc.core_properties.title = TITLE
    doc.core_properties.subject = "关键矿产出口管制、境外反制体系与海关监管研究"
    doc.core_properties.author = "关键矿产情报采集分析系统 · AI深度研究"
    doc.core_properties.last_modified_by = "关键矿产情报采集分析系统 · AI深度研究"
    doc.core_properties.keywords = "关键矿产, 出口管制, 反管制, 海关监管, AI深度研究"
    doc.core_properties.comments = "AI Agent多轮检索、证据分级与人工审核版"
    for section in doc.sections:
        section.page_width = Mm(210)
        section.page_height = Mm(297)

    doc.paragraphs[0].text = TITLE
    doc.paragraphs[0].style = "Heading 1"
    doc.paragraphs[1].text = SUBTITLE
    doc.paragraphs[1].style = "Block Text"
    doc.paragraphs[2].text = (
        "研究方法说明：本报告采用AI Agent任务拆解、多轮互联网检索、文件库分析、"
        "政策与企业信息交叉核验、证据分级和脚注编排流程形成。正文中的事实、研判与预测应分层使用；"
        "涉及法律适用、企业风险和具体金额的结论，应以主管部门、法院、企业公告及原始贸易单证为最终依据。"
    )
    doc.paragraphs[2].style = "First Paragraph"

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading("附录：AI Agent研究执行与证据核验框架", level=1)
    doc.add_heading("一、研究任务拆解", level=2)
    add_point(doc, "政策线：梳理中国关键矿产出口管制的法律、公告、暂停与执行节点。")
    add_point(doc, "国别线：比较美国、欧盟、日本、韩国及资源国的产业、财政、储备和联盟工具。")
    add_point(doc, "产业线：追踪矿山、分离冶炼、磁材、电池材料、回收和替代技术项目。")
    add_point(doc, "企业线：提取投资、股权、承购、价格保障、产能和投产时间等关系。")
    add_point(doc, "监管线：识别伪瞒报、第三方发货、转口、原产地和最终用途等海关风险。")

    doc.add_heading("二、证据分级规则", level=2)
    add_point(doc, "A级事实：政府、海关、法院、国际组织、监管机构和企业法定披露直接支持。")
    add_point(doc, "B级事实：企业官网、主流专业媒体、行业机构和可复核贸易记录相互印证。")
    add_point(doc, "C级线索：聚合平台、行业博客、搜索摘要或单一二手来源，仅用于引导进一步核查。")
    add_point(doc, "推断结论：由两个以上事实节点推导，必须列明推断路径、替代解释和证据缺口。")

    doc.add_heading("三、主要战略观点", level=2)
    viewpoints = [
        "全球政策竞争已从单一出口许可扩展为投资、价格保障、战略储备、联盟融资和贸易救济的组合工具竞争。",
        "供应链多元化的核心瓶颈不只在矿山，更集中在分离冶炼、材料纯化、磁体制造和规模化工艺经验。",
        "2025—2030年是管制工具成熟与替代产能尚未形成规模之间的高风险时间窗口。",
        "美国更强调国防动员与资本工具，欧盟强调法规、战略项目和循环利用，日本强调长期承购与资源外交，韩国强调库存和产业稳定。",
        "企业层面的真正风险不是一般关联关系，而是主体替换后仍出现相同受控物项、重合买方或路线及许可证缺口。",
        "海关监管应由税号筛查升级为商品成分、形态、技术参数、最终用户、物流路径和资金流的联合识别。",
    ]
    for item in viewpoints:
        add_point(doc, item)

    doc.add_heading("四、人工审核清单", level=2)
    add_point(doc, "审核一：核对政策发布日期、生效日期、暂停日期及是否存在过渡安排。")
    add_point(doc, "审核二：核对商品名称、两用物项编码、海关商品编号和技术参数是否一致。")
    add_point(doc, "审核三：核对同一数字是否被多个二手来源循环引用，优先回溯原始公告或财务披露。")
    add_point(doc, "审核四：对企业、人员和违法性质的表述实行更高证据门槛，未形成闭环时只标注为待核查线索。")
    add_point(doc, "审核五：对预测性产能、计划投资和未来投产时间标注不确定性，并设置滚动更新日期。")

    normal = doc.styles["Normal"]
    if normal.font.size is None:
        normal.font.size = Pt(10.5)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))

    final = Document(str(output))
    text = "\n".join(p.text for p in final.paragraphs)
    nonspace = len(re.sub(r"\s+", "", text))
    chinese = len(re.findall(r"[\u4e00-\u9fff]", text))
    print(f"OUTPUT={output}")
    print(f"NONSPACE_CHARS={nonspace}")
    print(f"CHINESE_CHARS={chinese}")
    print(f"PARAGRAPHS={len(final.paragraphs)}")
    print(f"TABLES={len(final.tables)}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    build(args.source, args.output)
