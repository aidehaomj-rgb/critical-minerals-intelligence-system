from __future__ import annotations

import argparse
import re
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BASE = Path(__file__).resolve().parent
ASSET_DIR = BASE / "reports" / "_professional_report_assets"

NAVY = "17365D"
BLUE = "2F5597"
TEAL = "0F6B78"
GOLD = "B18432"
PALE_BLUE = "EAF1F8"
PALE_GOLD = "F7F1E5"
PALE_GRAY = "F3F5F7"
MID_GRAY = "667085"
DARK = "1F2937"
WHITE = "FFFFFF"
RED = "A63A3A"


SOURCES = [
    ("FN01", "中华人民共和国国务院：《中华人民共和国两用物项出口管制条例》（国务院令第792号），2024年。https://xkzj.mofcom.gov.cn/tzgg/art/2024/art_49503f2524484d8d9488dfd37395a731.html"),
    ("FN02", "商务部、海关总署等：《中华人民共和国两用物项出口管制清单》，自2024年12月1日起实施。https://app.www.gov.cn/govdata/gov/202411/19/521898/article.html"),
    ("FN03", "商务部、海关总署公告2023年第23号：关于对镓、锗相关物项实施出口管制。https://www.mofcom.gov.cn/zcfb/blgg/art/2023/art_ca2e9d349361441f847bdabac5d8331b.html"),
    ("FN04", "商务部、海关总署公告2023年第39号：关于优化调整石墨物项临时出口管制措施。https://interview.mofcom.gov.cn/mofcom_interview/front/opdata/downlodePdf?id=20231003447368"),
    ("FN05", "商务部、海关总署公告2024年第33号：关于对锑等物项实施出口管制。https://www.mofcom.gov.cn/zcfb/zc/art/2024/art_8d1d0c07124e4c01b69339c275c69cbe.html"),
    ("FN06", "商务部公告2024年第46号：关于加强相关两用物项对美国出口管制。https://www.mofcom.gov.cn/zcfb/zc/art/2024/art_a362d9e4d4944ff3854b76c572899e7e.html"),
    ("FN07", "商务部、海关总署公告2025年第10号：对钨、碲、铋、钼、铟相关物项实施出口管制。https://exportcontrol.mofcom.gov.cn/article/zcfg/gnzcfg/zcfggzqd/202502/1098.html"),
    ("FN08", "商务部、海关总署公告2025年第18号：对部分中重稀土相关物项实施出口管制。https://www.mofcom.gov.cn/zfxxgk/fdzdgknr/ztfl/blgg/art/2025/art_41ba05e0a9f24635a5fbc182b58aec00.html"),
    ("FN09", "商务部、海关总署公告2025年第70号：暂停实施第55、56、57、58、61、62号公告至2026年11月10日。https://www.mofcom.gov.cn/zfxxgk/fdzdgknr/ztfl/dwmygl/art/2025/art_00667414c0524b018985abd28b8847a5.html"),
    ("FN10", "International Energy Agency, Global Critical Minerals Outlook 2025, Executive Summary. https://www.iea.org/reports/global-critical-minerals-outlook-2025/executive-summary"),
    ("FN11", "International Energy Agency, With new export controls on critical minerals, supply concentration risks become reality, 23 October 2025. https://www.iea.org/commentaries/with-new-export-controls-on-critical-minerals-supply-concentration-risks-become-reality"),
    ("FN12", "International Energy Agency, Rare Earth Elements, Executive Summary, 2026. https://www.iea.org/reports/rare-earth-elements/executive-summary"),
    ("FN13", "International Energy Agency, Recycling of Critical Minerals, Executive Summary, 2024. https://www.iea.org/reports/recycling-of-critical-minerals/executive-summary"),
    ("FN14", "The White House, Executive Order 14241, Immediate Measures to Increase American Mineral Production, 20 March 2025. https://www.whitehouse.gov/presidential-actions/2025/03/immediate-measures-to-increase-american-mineral-production/"),
    ("FN15", "The White House, Executive Order 14272, Ensuring National Security and Economic Resilience Through Section 232 Actions on Processed Critical Minerals, 15 April 2025. https://www.whitehouse.gov/presidential-actions/2025/04/ensuring-national-security-and-economic-resilience-through-section-232-actions-on-processed-critical-minerals-and-derivative-products/"),
    ("FN16", "The White House, Adjusting Imports of Processed Critical Minerals and Their Derivative Products into the United States, January 2026. https://www.whitehouse.gov/presidential-actions/2026/01/adjusting-imports-of-processed-critical-minerals-and-their-derivative-products-into-the-united-states/"),
    ("FN17", "Export-Import Bank of the United States, Supply Chain Resiliency Initiative, 8 January 2025. https://www.exim.gov/news/export-import-bank-united-states-board-directors-approves-supply-chain-resiliency"),
    ("FN18", "U.S. Department of Defense, Office of Strategic Capital Announces First Loan Through DoD Agreement With MP Materials, 10 August 2025. https://www.defense.gov/News/Releases/Release/Article/4270722/office-of-strategic-capital-announces-first-loan-through-dod-agreement-with-mp/"),
    ("FN19", "MP Materials Corp., Form 10-Q for the quarter ended 31 March 2026, U.S. SEC; includes the NdPr price protection agreement. https://www.sec.gov/Archives/edgar/data/1801368/000180136826000029/mp-20260331.htm"),
    ("FN20", "Export-Import Bank of the United States, Project Vault loan approval, 2 February 2026. https://www.exim.gov/news/project-vault"),
    ("FN21", "U.S. Geological Survey, About the 2025 List of Critical Minerals, 6 November 2025. https://www.usgs.gov/programs/mineral-resources-program/science/about-2025-list-critical-minerals"),
    ("FN22", "U.S. Geological Survey, Mineral Commodity Summaries 2026. https://pubs.usgs.gov/periodicals/mcs2026/mcs2026.pdf"),
    ("FN23", "European Commission, Critical Raw Materials Act: objectives, benchmarks and implementation. https://single-market-economy.ec.europa.eu/sectors/raw-materials/areas-specific-interest/critical-raw-materials/critical-raw-materials-act_en"),
    ("FN24", "European Commission, Selected Strategic Projects under the Critical Raw Materials Act, 2025. https://single-market-economy.ec.europa.eu/sectors/raw-materials/areas-specific-interest/critical-raw-materials/strategic-projects-under-crma/selected-projects_en"),
    ("FN25", "European Commission, RESourceEU action plan: measures to secure raw materials and strengthen economic security, 3 December 2025. https://commission.europa.eu/news-and-media/news/new-measures-secure-raw-materials-and-strengthen-eus-economic-security-2025-12-03_en"),
    ("FN26", "Ministry of Economy, Trade and Industry of Japan, Japanese and French Governments Support a Heavy Rare Earth Project in France, 17 March 2025. https://www.meti.go.jp/english/press/2025/0317_002.html"),
    ("FN27", "JOGMEC, Securing Supply of Heavy Rare Earths to Japan with Additional Investment to Lynas, 7 March 2023. https://www.jogmec.go.jp/english/news/release/release_00228.html"),
    ("FN28", "Ministry of Economy, Trade and Industry of Japan, Japan-U.S. Action Plan and project cooperation on critical minerals, 20 March 2026. https://www.meti.go.jp/english/press/2026/0320_001.html"),
    ("FN29", "Australian Department of Industry, Science and Resources, Critical Minerals Strategic Reserve, 2026. https://www.industry.gov.au/mining-oil-and-gas/minerals/critical-minerals/critical-minerals-strategic-reserve"),
    ("FN30", "Government of India, National Critical Mineral Mission, Cabinet approval, 29 January 2025. https://www.pib.gov.in/Pressreleaseshare.aspx?PRID=2097309&lang=2&reg=48"),
    ("FN31", "Government of Canada, G7 Critical Minerals Action Plan, 17 June 2025. https://g7.canada.ca/en/news-and-media/news/g7-critical-minerals-action-plan/"),
    ("FN32", "Australian Minister for Foreign Affairs, Quad Critical Minerals Initiative Framework, 26 May 2026. https://www.foreignminister.gov.au/minister/penny-wong/media-release/quad-critical-minerals-initiative-framework-among-united-states-japan-australia-and-india"),
    ("FN33", "Natural Resources Canada, Critical Minerals Resilience and Production Alliance, 2026. https://www.canada.ca/en/campaign/critical-minerals-in-canada/our-critical-minerals-strategic-partnerships/critical-minerals-production-alliance.html"),
    ("FN34", "World Trade Organization, China — Measures Related to the Exportation of Rare Earths, Tungsten and Molybdenum, DS432 one-page summary. https://www.wto.org/english/tratop_e/dispu_e/cases_e/1pagesum_e/ds432sum_e.pdf"),
    ("FN35", "U.S. Geological Survey, Quantifying potential effects of China’s gallium and germanium export restrictions on the U.S. economy. https://www.usgs.gov/publications/quantifying-potential-effects-chinas-gallium-and-germanium-export-restrictions-us"),
    ("FN36", "Ministry of Economy, Trade and Industry of Japan, Interim Report on Strengthening the Manufacturing Base in Light of Geopolitical Risks, 15 April 2026. https://www.meti.go.jp/english/press/2026/0415_001.html"),
    ("FN37", "International Energy Agency Policies Database, Republic of Korea: Strategy for Securing Reliable Critical Minerals Supply. https://www.iea.org/policies/17942-the-strategy-for-securing-reliable-critical-minerals-supply"),
]


def marker(code: str) -> str:
    return f"[[{code}]]"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = tc_mar.find(qn(f"w:{side}"))
        if tag is None:
            tag = OxmlElement(f"w:{side}")
            tc_mar.append(tag)
        tag.set(qn("w:w"), str(value))
        tag.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int], indent=120) -> None:
    total = sum(widths)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")
    old_grid = table._tbl.tblGrid
    for child in list(old_grid):
        old_grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        old_grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths[min(idx, len(widths) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_run_font(run, name="宋体", size=10.5, color=DARK, bold=False, italic=False) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    pf = normal.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(5)
    pf.line_spacing = 1.45
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for style_name, size, color, before, after in [
        ("Heading 1", 16, NAVY, 18, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 11.5, TEAL, 8, 4),
        ("Heading 4", 10.5, DARK, 6, 3),
    ]:
        style = doc.styles[style_name]
        style.font.name = "微软雅黑"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Caption Pro" not in [s.name for s in doc.styles]:
        cap = doc.styles.add_style("Caption Pro", WD_STYLE_TYPE.PARAGRAPH)
    else:
        cap = doc.styles["Caption Pro"]
    cap.font.name = "微软雅黑"
    cap._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    cap.font.size = Pt(9)
    cap.font.color.rgb = RGBColor.from_string(MID_GRAY)
    cap.paragraph_format.space_before = Pt(4)
    cap.paragraph_format.space_after = Pt(5)
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(p.add_run("关键矿产“反管制”体系深度研究｜专业审阅稿"), "微软雅黑", 8.5, MID_GRAY)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(fp.add_run("独立研究 · 资料截至2026年7月5日  |  "), "微软雅黑", 8, MID_GRAY)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    fp._p.append(fld)


def add_body(doc: Document, text: str, bold_lead: str | None = None, after=5) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.45
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_lead and text.startswith(bold_lead):
        set_run_font(p.add_run(bold_lead), bold=True)
        set_run_font(p.add_run(text[len(bold_lead) :]))
    else:
        set_run_font(p.add_run(text))


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.32)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.3
    set_run_font(p.add_run(text))


def add_number(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.34)
    p.paragraph_format.first_line_indent = Inches(-0.2)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.3
    set_run_font(p.add_run(text))


def add_callout(doc: Document, title: str, text: str, fill=PALE_BLUE) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    set_run_font(p.add_run(title), "微软雅黑", 10.5, NAVY, bold=True)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.35
    set_run_font(p2.add_run(text), size=10)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int], caption: str | None = None) -> None:
    if caption:
        cp = doc.add_paragraph(caption, style="Caption Pro")
        cp.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    hdr = table.rows[0]
    tr_pr = hdr._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)
    for idx, text in enumerate(headers):
        set_cell_shading(hdr.cells[idx], NAVY)
        p = hdr.cells[idx].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(text), "微软雅黑", 9, WHITE, bold=True)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for cidx, text in enumerate(row):
            if ridx % 2:
                set_cell_shading(cells[cidx], PALE_GRAY)
            p = cells[cidx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if cidx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            set_run_font(p.add_run(str(text)), size=8.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_picture(doc: Document, path: Path, caption: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    shape = p.add_run().add_picture(str(path), width=Inches(6.55))
    shape._inline.docPr.set("descr", caption)
    shape._inline.docPr.set("title", caption.split("  ", 1)[0])
    doc.add_paragraph(caption, style="Caption Pro")


def make_charts() -> dict[str, Path]:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    regular_path = Path(r"C:\Windows\Fonts\msyh.ttc")
    bold_path = Path(r"C:\Windows\Fonts\msyhbd.ttc")
    def font(size: int, bold=False):
        path = bold_path if bold and bold_path.exists() else regular_path
        return ImageFont.truetype(str(path), size=size)
    def text_center(draw, xy, text, fnt, fill):
        box = draw.textbbox((0, 0), text, font=fnt)
        w, h = box[2] - box[0], box[3] - box[1]
        draw.text((xy[0] - w / 2, xy[1] - h / 2), text, font=fnt, fill=fill)
    paths = {}
    img = Image.new("RGB", (1800, 920), "white")
    draw = ImageDraw.Draw(img)
    draw.text((70, 45), "磁性稀土价值链集中度（2024年）", font=font(42, True), fill="#17365D")
    stages = ["磁性稀土采矿", "磁性稀土分离精炼", "烧结永磁体制造"]
    values = [60, 91, 94]
    colors = ["#6B8DB5", "#2F5597", "#17365D"]
    x0, x1 = 470, 1640
    for tick in range(0, 101, 20):
        x = x0 + (x1 - x0) * tick / 100
        draw.line((x, 145, x, 780), fill="#E2E8F0", width=2)
        text_center(draw, (x, 820), str(tick), font(25), "#667085")
    for idx, (stage, val, color) in enumerate(zip(stages, values, colors)):
        y = 225 + idx * 200
        draw.text((70, y + 18), stage, font=font(30, True), fill="#1F2937")
        draw.rounded_rectangle((x0, y, x1, y + 82), radius=16, fill="#EEF2F6")
        end = x0 + (x1 - x0) * val / 100
        draw.rounded_rectangle((x0, y, end, y + 82), radius=16, fill=color)
        draw.text((end + 18, y + 20), f"{val}%", font=font(30, True), fill="#17365D")
    text_center(draw, ((x0 + x1) / 2, 880), "中国占全球比重（%）", font(25), "#667085")
    p = ASSET_DIR / "rare-earth-concentration.png"
    img.save(p)
    paths["concentration"] = p

    img = Image.new("RGB", (1800, 1040), "white")
    draw = ImageDraw.Draw(img)
    draw.text((70, 40), "主要经济体关键矿产政策工具组合（研究编码）", font=font(40, True), fill="#17365D")
    economies = ["美国", "欧盟", "日本", "澳大利亚", "印度", "韩国"]
    tools = ["法规/清单", "政策融资", "长期承购", "战略储备", "价格工具", "国际联盟"]
    scores = [
        [3, 3, 3, 3, 3, 3],
        [3, 2, 2, 2, 1, 3],
        [2, 3, 3, 3, 1, 3],
        [2, 3, 3, 3, 3, 3],
        [3, 2, 2, 2, 1, 3],
        [2, 2, 3, 3, 1, 3],
    ]
    palette = ["#F8FAFC", "#DCE8F5", "#83A8D2", "#2F5597"]
    labels = ["—", "起步", "形成", "强化"]
    left, top, cell_w, cell_h = 300, 190, 235, 115
    for j, tool in enumerate(tools):
        text_center(draw, (left + j * cell_w + cell_w / 2, 135), tool, font(24, True), "#1F2937")
    for i in range(len(economies)):
        text_center(draw, (160, top + i * cell_h + cell_h / 2), economies[i], font(28, True), "#1F2937")
        for j in range(len(tools)):
            score = scores[i][j]
            x, y = left + j * cell_w, top + i * cell_h
            draw.rounded_rectangle((x + 5, y + 5, x + cell_w - 5, y + cell_h - 5), radius=12,
                                   fill=palette[score], outline="#D0D5DD")
            text_center(draw, (x + cell_w / 2, y + cell_h / 2), labels[score], font(24, score >= 2),
                        "white" if score >= 2 else "#1F2937")
    draw.text((70, 930), "说明：研究组依据法规、融资、承购、库存、价格机制和国际合作的公开政策进行定性编码。", font=font(22), fill="#667085")
    p = ASSET_DIR / "policy-toolkit-heatmap.png"
    img.save(p)
    paths["toolkit"] = p

    img = Image.new("RGB", (1900, 920), "white")
    draw = ImageDraw.Draw(img)
    draw.text((70, 45), "海关风险预警的证据链", font=font(42, True), fill="#17365D")
    boxes = [
        ("政策状态", "生效/暂停/恢复"),
        ("物项参数", "成分/纯度/形态"),
        ("主体关系", "企业/货代/买方"),
        ("交易聚合", "跨票/跨口岸/时间"),
        ("证据分层", "事实/推断/待核"),
        ("人工处置", "调单/查验/审核"),
    ]
    box_w, gap, y = 250, 55, 260
    x0 = 45
    for i, (title, sub) in enumerate(boxes):
        x = x0 + i * (box_w + gap)
        draw.rounded_rectangle((x, y, x + box_w, y + 190), radius=18, fill="#EAF1F8", outline="#2F5597", width=4)
        text_center(draw, (x + box_w / 2, y + 72), title, font(28, True), "#17365D")
        text_center(draw, (x + box_w / 2, y + 128), sub, font(21), "#667085")
        if i < len(boxes) - 1:
            sx, ex, cy = x + box_w + 8, x + box_w + gap - 8, y + 95
            draw.line((sx, cy, ex, cy), fill="#B18432", width=6)
            draw.polygon([(ex, cy), (ex - 18, cy - 12), (ex - 18, cy + 12)], fill="#B18432")
    text_center(draw, (950, 610), "任何高风险标签均须保留来源、适用时间和人工复核入口", font(30, True), "#A63A3A")
    text_center(draw, (950, 700), "异常信号 → 调取原始单证 → 物项与许可核验 → 人工审核 → 决定是否处置", font(26), "#667085")
    p = ASSET_DIR / "customs-evidence-pipeline.png"
    img.save(p)
    paths["pipeline"] = p
    return paths


def add_cover(doc: Document) -> None:
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("独立深度研究报告"), "微软雅黑", 11, GOLD, bold=True)
    p.paragraph_format.space_after = Pt(18)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("全球主要经济体应对中国关键矿产出口管制的"), "微软雅黑", 22, NAVY, bold=True)
    p.paragraph_format.space_after = Pt(4)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("政策与产业重构"), "微软雅黑", 28, NAVY, bold=True)
    p.paragraph_format.space_after = Pt(12)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("“反管制”体系、技术替代与海关监管影响（2023—2026）"), "微软雅黑", 14, TEAL)
    p.paragraph_format.space_after = Pt(46)
    add_callout(
        doc,
        "研究定位",
        "本报告借鉴前期材料的主题覆盖方式，但全部事实重新回溯原始来源，全部观点重新论证。所谓“反管制”仅为分析概念，指进口经济体为降低供应中断风险而形成的政策、资本、库存、技术和联盟工具组合，并非特定国家的正式法律术语。",
        PALE_GOLD,
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(45)
    set_run_font(p.add_run("专业审阅稿"), "微软雅黑", 12, NAVY, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("资料截至：2026年7月5日"), "微软雅黑", 10, MID_GRAY)
    doc.add_page_break()


def add_static_toc(doc: Document) -> None:
    doc.add_heading("目录", level=1)
    items = [
        "执行摘要",
        "1 研究范围、概念与证据方法",
        "2 中国出口管制政策体系：从物项许可到产业链规则",
        "3 美欧日韩澳印“反管制”政策工具比较",
        "4 全球替代供应源：项目数量增长不等于有效供给",
        "5 加工技术与中游瓶颈：真正的竞争发生在“转化”",
        "6 替代、减量与回收：三条路径对应不同时间尺度",
        "7 多边合作：从供应链联盟走向标准化市场",
        "8 企业供应链重组：从采购多元化到资本共担",
        "9 市场影响：从单一价格叙事转向多变量传导",
        "10 法律争端与企业合规：规则冲突不会自动消失",
        "11 新兴供应国：资源民族主义与本地增值并行",
        "12 国防与先进制造：需求小、规格高、替代慢",
        "13 2030情景分析：不预测单一路径，识别可观测触发器",
        "14 独立战略洞察与海关监管建议",
        "附录A 重点矿产与监管观察清单",
        "附录B 项目阶段与证据等级",
        "附录C 十二项专题深度分析",
        "附录D 重点矿种独立研判",
        "附录E 核心来源说明",
    ]
    for idx, item in enumerate(items):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.08 if idx == 0 else 0.18)
        p.paragraph_format.space_after = Pt(4)
        set_run_font(p.add_run(item), "微软雅黑", 10.5, NAVY if idx == 0 else DARK, bold=idx == 0)
    doc.add_page_break()


def build_report(output: Path) -> tuple[int, int, int, int]:
    charts = make_charts()
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_static_toc(doc)

    doc.add_heading("执行摘要", level=1)
    add_body(doc, "本报告围绕一个核心问题展开：当中国以物项清单、许可审查、最终用户管理、技术参数和特定目的地规则控制关键矿产及相关技术的跨境转移时，主要经济体如何把资源政策、产业政策、金融工具和联盟机制组合为可持续的供应链安全能力。研究不把“宣布投资”视为“形成产能”，也不把“贸易下降”直接归因于管制，而是把法律状态、项目阶段、商业可行性和实际交付能力分开评估。")
    add_body(doc, f"第一，关键矿产竞争的重心已从“谁拥有矿山”转向“谁能把矿石稳定转化为合格材料”。IEA数据显示，2024年中国在磁性稀土采矿、分离精炼和烧结永磁体制造环节的全球份额分别约为60%、91%和94%。这种沿价值链向下游递增的集中度意味着，新增矿权并不能自动消除供应风险；分离工艺、杂质控制、金属化、合金粉末、磁体良率和客户认证才是决定有效供给的约束。{marker('FN12')}")
    add_body(doc, f"第二，主要经济体正在把政府从“规则制定者”改造为“市场组织者”。美国通过行政命令、政策性贷款、价格保护、承购和战略储备同时介入项目收益和需求预期；欧盟以法规基准、战略项目、许可提速和联合采购构造制度市场；日本通过JOGMEC股权、长期承购和资源外交锁定特定材料；澳大利亚则进一步把远期承购、合同价差和选择性库存纳入战略储备工具箱。{marker('FN14')}{marker('FN29')}")
    add_body(doc, "第三，所谓“反管制”不是取消对中国依赖，而是购买选择权。短期选择权来自库存、许可协调和替代贸易路线；中期选择权来自回收、材料减量和在盟友国家形成可销售产能；长期选择权来自技术替代和产品重新设计。三类选择权的建设周期、成本和适用场景完全不同，不能用一个“去依赖率”概括。")
    add_body(doc, f"第四，2025—2030年构成显著的“政策—产能错配窗口”。政策工具可以在数日内调整，矿山、分离厂和磁体工厂却需要多年建设与认证。IEA同时指出，2024年关键矿产投资增速明显放缓，而精炼集中度继续上升；这意味着低价环境反而削弱了新进入者融资能力，强化了既有供应中心。{marker('FN10')}")
    add_body(doc, f"第五，2025年以后“价格支持”成为产业政策的重要分水岭。MP Materials披露的钕镨价格保护安排把政府支持从一次性补贴推进到长期收入稳定机制；Project Vault和澳大利亚战略储备又把库存、远期合同与下游需求聚合结合起来。由此可能形成“商业市场价格”和“安全溢价价格”并存的双层市场。{marker('FN19')}{marker('FN20')}")
    add_body(doc, "第六，多边合作正从原则声明转向项目和市场规则。G7强调标准化市场、资本动员和技术创新；Quad在2026年框架中提出动员最高200亿美元公共与私人支持，并列出贷款、股权、保险、承购和价格机制等工具。联盟的实质，不再是成员数量，而是能否把项目、融资、规格、承购和物流接成闭环。")
    add_body(doc, "第七，企业层面的最优策略不是简单“撤出中国”，而是分层管理：对不可替代材料建立库存与许可缓冲；对可替代材料推进双源认证；对中游瓶颈通过长期承购和共同投资锁定能力；对技术和最终用途建立可证明的合规链。供应链安全的目标不是完全本地化，而是在冲击发生时保留可切换的合格供应。")
    add_body(doc, "第八，对海关监管而言，真正有价值的不是扩大风险标签，而是把政策状态、物项参数、许可证、主体关系、跨票交易、物流和资金证据连接起来。任何企业关联、路线变化或价格异常都只能构成线索；只有与受控物项、许可缺口、单证矛盾和最终用途风险形成证据闭环，才适合升级为高优先核查。")
    add_callout(doc, "本报告的独立判断", "全球关键矿产体系正在从“全球统一价格下的效率竞争”转向“安全溢价、政策融资和受控市场准入并存”的双层结构。未来五年最稀缺的不是地质资源，而是能够同时满足技术、合规、融资和下游认证要求的可交付中游能力。")

    doc.add_heading("1 研究范围、概念与证据方法", level=1)
    doc.add_heading("1.1 “反管制”作为分析概念", level=2)
    add_body(doc, "本报告使用“反管制”一词描述进口经济体和跨国企业为降低出口许可不确定性而采取的系统性应对，包括增加本土或盟友供给、建立库存、提供政策融资、锁定承购、发展替代技术、调整产品设计和建立多边规则。它不同于法律意义上的贸易反制，也不预设这些政策一定针对中国。")
    add_body(doc, "分析对象限定为具有明显技术门槛、供应集中度或国家安全用途的矿产及其下游材料，重点覆盖镓、锗、石墨、锑、钨、碲、铋、钼、铟、中重稀土、永磁体和部分电池材料。研究时段以2023年镓锗公告至2026年上半年为主，历史材料仅用于解释制度演变。")
    doc.add_heading("1.2 证据等级与事实边界", level=2)
    add_body(doc, "A级证据为政府法规、海关或监管公告、国际组织数据库、法院裁判和企业法定披露；B级证据为企业官网、项目融资文件、专业机构可复核数据；C级证据为新闻聚合、行业评论和无法追溯原始口径的二手数字。关键结论原则上至少有一个A级来源，企业项目进度尽量使用政府文件与企业法定披露交叉验证。")
    add_body(doc, "报告严格区分四类状态：已经生效的政策、已经暂停或失效的政策、已宣布但尚未落地的项目、基于公开信息形成的研究判断。对外部数据库中的提单、价格和企业关系，不在缺少原始单证时作违法定性。")
    doc.add_heading("1.3 分析模型", level=2)
    add_body(doc, "本报告把供应链韧性拆解为五个相乘而非相加的条件：物理资源可得性、转化加工能力、商业可融资性、监管可进入性和物流可交付性。任何一项接近零，名义资源量都无法转换为有效供应。该模型解释了为什么部分资源丰富国家仍长期依赖外部精炼，也解释了价格保护和长期承购为何逐渐成为政策核心。")
    add_body(doc, "项目评价采用七阶段法：资源确认、预可研、可研与许可、融资关闭、建设、试生产与认证、商业化交付。只公布资源量或签署谅解备忘录的项目，不计入2030年有效产能；完成融资但尚未认证的项目，只计入条件性产能。")
    doc.add_heading("1.4 数据限制", level=2)
    add_body(doc, "关键矿产统计存在品位、纯度、氧化物当量、金属量、精矿量和制品量等多种口径。海关税号通常覆盖受控与非受控产品，公告中的参考税号不能替代技术参数判断。商业价格还受合同期限、交付地、含税状态和规格影响，因此本报告避免把单点报价直接解释为政策冲击。")

    doc.add_heading("2 中国出口管制政策体系：从物项许可到产业链规则", level=1)
    doc.add_heading("2.1 制度骨架", level=2)
    add_body(doc, f"2024年《两用物项出口管制条例》把清单、临时管制、许可、最终用户和最终用途、管控名单、海关查验以及境外转移等要素纳入统一框架。条例规定单项许可、通用许可和登记填报三类路径，并要求出口经营者了解物项性能指标和主要用途；无法确定时可以申请业务咨询。{marker('FN01')}")
    add_body(doc, f"2024年12月实施的统一两用物项出口管制清单，进一步把分散公告和既有清单整合为编码化体系。对监管和企业而言，这一变化的重要性不在于编码形式本身，而在于物项、许可证和报关数据可以形成机器可读的关联。{marker('FN02')}")
    doc.add_heading("2.2 2023—2026政策时间线与当前状态", level=2)
    timeline_rows = [
        ["2023.08", "镓、锗", "金属、化合物、晶片/衬底等", "许可实施", "第23号公告"],
        ["2023.12", "石墨", "高规格人造石墨、天然鳞片石墨及制品", "许可实施", "第39号公告"],
        ["2024.09", "锑、超硬材料", "锑矿、金属、氧化物、设备与技术等", "许可实施", "第33号公告"],
        ["2024.12", "特定目的地强化", "镓、锗、锑、超硬材料、石墨对美", "差异化审查", "第46号公告"],
        ["2025.02", "钨、碲、铋、钼、铟", "满足技术条件的相关物项", "许可实施", "第10号公告"],
        ["2025.04", "部分中重稀土", "钐、钆、铽、镝、镥、钪、钇", "许可实施", "第18号公告"],
        ["2025.10", "扩围措施", "超硬材料、稀土设备/元素、锂电材料、境外规则", "后续暂停", "第55—62号公告"],
        ["2025.11", "暂停决定", "第55、56、57、58、61、62号公告", "暂停至2026.11.10", "第70号公告"],
    ]
    add_table(doc, ["时间", "政策对象", "主要范围", "截至2026.7状态", "依据"], timeline_rows,
              [950, 1300, 3120, 1850, 2140], "表2-1  中国关键矿产相关出口管制政策时间线（状态口径）")
    add_body(doc, f"镓锗公告首次把多种金属、化合物和半导体材料形态纳入许可范围；石墨公告则通过纯度、强度、密度和天然鳞片石墨形态组合界定边界。两者都说明，海关商品编号只是参考，真正决定是否受控的是公告列明的物项属性。{marker('FN03')}{marker('FN04')}")
    add_body(doc, f"2024年锑及超硬材料公告把矿料、金属、氧化物、特殊化合物、设备和技术并置，显示管制对象从资源端向工艺端延伸。2024年第46号公告又对特定目的地设置原则上不予许可或更严格最终用途审查，说明同一物项的风险还取决于目的地和最终用户。{marker('FN05')}{marker('FN06')}")
    add_body(doc, f"2025年第10号和第18号公告分别覆盖钨等五类物项以及七种中重稀土。其共同特点是技术条件和产品形态非常具体，且要求在报关备注栏标明两用物项编码。监管系统因此必须能够处理“一税号多物项”和“一物项多税号”的多对多关系。{marker('FN07')}{marker('FN08')}")
    add_body(doc, f"2025年10月发布的若干扩围公告并非全部持续有效。第70号公告明确将第55、56、57、58、61、62号公告暂停至2026年11月10日。任何趋势统计、企业风险识别或政策密度计算都必须从暂停之日起扣除相应政策组，不能把“曾发布”误计为“当前有效”。{marker('FN09')}")
    doc.add_heading("2.3 政策机制的三个层次", level=2)
    add_body(doc, "第一层是物项识别：通过成分、纯度、粒度、尺寸、形态、性能和用途确定是否落入清单。第二层是交易许可：结合目的地、最终用户、最终用途和合同关系判断是否准予出口。第三层是持续控制：通过许可证核销、最终用途承诺、变更报告和海关质疑机制约束出口后的风险。")
    add_body(doc, "从政策效果看，许可制与全面禁运不同。它可以维持正常贸易，同时把不确定性集中于高风险用途和不完整信息交易。对市场主体而言，最具影响的往往不是公告发布本身，而是分类咨询、材料补正、最终用户证明和审批时长对交付计划的改变。")
    add_callout(doc, "独立判断：政策强度应按“有效物项组”而非公告数量衡量", "公告数量会被扩围、整合、暂停和恢复反复影响。更稳健的指标是：每季度实际生效的物项组数量、覆盖的技术环节、目的地差异、许可证审查条件和暂停状态。")

    doc.add_heading("3 美欧日韩澳印“反管制”政策工具比较", level=1)
    doc.add_heading("3.1 美国：国家安全、资本与价格机制叠加", level=2)
    add_body(doc, "美国政策的鲜明特征是把矿产生产定义延伸至加工、精炼、金属粉末、母合金和衍生产品，并通过联邦土地、国防生产法、出口信贷和国家安全审查协调推进。其目标不是单纯增加采矿量，而是把可被国防和先进制造直接使用的材料留在本土或盟友体系。")
    add_body(doc, f"2025年行政命令要求加快矿产生产，并明确“生产”包括采矿、加工、精炼和冶炼；随后针对加工关键矿产及衍生品启动国家安全调查。2026年后续措施进一步强调，仅有国内采矿而依赖外国加工仍不足以保障安全。{marker('FN15')}{marker('FN16')}")
    add_body(doc, f"美国2025年关键矿产清单扩大到60种，并采用包含经济影响和供应中断情景的新方法；2026年矿产品年度摘要则提供进口依赖和产量基线。两类文件共同表明，美国正在把“是否关键”从静态名单转向可更新的经济风险模型。{marker('FN21')}{marker('FN22')}")
    add_body(doc, f"政策融资开始从项目贷款走向收益结构重塑。美国国防部战略资本办公室向MP Materials提供1.5亿美元贷款用于重稀土分离能力；其价格保护协议以长期基准价降低低价周期中的现金流风险。EXIM的供应链韧性倡议则把海外矿产项目与美国制造业承购相连接。{marker('FN17')}{marker('FN18')}")
    doc.add_heading("3.2 欧盟：法规目标、战略项目与共同采购", level=2)
    add_body(doc, f"欧盟《关键原材料法案》以2030年基准组织政策：战略原材料消费量的至少10%由欧盟开采、40%由欧盟加工、25%来自回收，并把对单一第三国依赖控制在65%以内。该框架同时设置许可时限、压力测试、战略库存协调和大企业风险准备义务。{marker('FN23')}")
    add_body(doc, f"2025年首批战略项目把法规目标转化为项目清单，覆盖开采、加工、回收和替代。RESourceEU又增加需求聚合、联合采购、承购和协调储备。欧盟模式的优势是规则和市场规模，约束则在于成员国许可、能源成本和公共资金协调。{marker('FN24')}{marker('FN25')}")
    doc.add_heading("3.3 日本：少量资本撬动长期承购", level=2)
    add_body(doc, f"日本政策更强调“可获得份额”而非“项目所有权”。JOGMEC与企业共同投资Lynas并锁定部分重稀土供应，说明政策目标是把海外资源、加工能力和日本需求形成长期合同闭环。{marker('FN27')}")
    add_body(doc, f"日本和法国支持Caremag重稀土项目，通过约1亿欧元日方支持和长期供货安排，争取相当于日本未来镝铽需求一定比例的供应。该项目同时使用回收磁体和原矿，体现“海外加工+回收原料+承购锁定”的复合路径。2026年日美又把铜、稀土、第三国项目和深海资源合作纳入行动计划，显示日本正把双边项目组合化。{marker('FN26')}{marker('FN28')}")
    doc.add_heading("3.4 韩国、澳大利亚与印度：制造连续性、资源金融化和全链条任务", level=2)
    add_body(doc, f"韩国作为电池、半导体和汽车制造中心，政策重点是扩大关键矿产清单、提高储备、支持海外开发和强化供应链预警。其结构性难题是中游材料生产虽较强，但部分前驱体、石墨和原料仍依赖中国，库存只能换取调整时间。日本2026年制造基础报告同样把支持对象从单点项目扩展到材料、部件、循环资源和物流，反映东北亚制造经济体正在由“保原料”转向“保制造能力”。{marker('FN37')}{marker('FN36')}")
    add_body(doc, f"澳大利亚正在从资源出口国转向“资源+金融工具”提供者。2026年战略储备允许使用长期承购、远期合同、需求聚合、选择性库存和合同价差，首批聚焦锑、镓和轻重稀土。这类工具能够对冲低价周期，但也要求政府具备商品交易和项目尽调能力。")
    add_body(doc, f"印度国家关键矿产任务覆盖勘探、采矿、选矿、加工、海外资产、尾矿回收、加工园区、回收和研发，七年总投入由财政支出与国有企业投资共同构成。印度的比较优势在于市场和工程能力，短板则在资源勘查、工艺积累和基础设施。{marker('FN30')}")
    add_picture(doc, charts["toolkit"], "图3-1  主要经济体关键矿产政策工具组合（研究组基于官方政策的定性编码；3=强化，2=形成，1=起步）")
    policy_rows = [
        ["美国", "国防资本、出口信贷、价格保护、储备、国家安全调查", "收入稳定和需求锁定", "财政成本、项目成本和许可"],
        ["欧盟", "法规基准、战略项目、许可提速、联合采购、回收", "统一规则和大市场", "成员国执行与能源成本"],
        ["日本", "JOGMEC投资、长期承购、库存、资源外交", "合同精细、产业协同", "海外资源与价格风险"],
        ["韩国", "储备、产业基金、海外合作、供应预警", "制造端需求明确", "上游和部分中游依赖"],
        ["澳大利亚", "政策贷款、税收抵免、承购、合同价差、储备", "资源禀赋与盟友需求", "加工成本和项目融资"],
        ["印度", "国家任务、勘探拍卖、加工园区、海外资产、回收", "市场规模和政策集中", "技术、基础设施与建设周期"],
    ]
    add_table(doc, ["经济体", "核心工具", "主要优势", "主要约束"], policy_rows,
              [1050, 3700, 2300, 2310], "表3-1  主要经济体“反管制”政策模式比较")

    doc.add_heading("4 全球替代供应源：项目数量增长不等于有效供给", level=1)
    doc.add_heading("4.1 稀土：资源端多点出现，中游仍高度集中", level=2)
    add_body(doc, "美国、澳大利亚、巴西、加拿大、非洲和印度均有稀土项目推进，但矿体类型、放射性伴生、选冶流程和重稀土含量差异显著。项目宣传中的“稀土总量”不能直接换算为磁性稀土供给，更不能换算为镝铽或合格磁体产量。")
    add_body(doc, "IEA指出，稀土项目从发现到投产平均需要较长时间，分离精炼设施比矿山更少，永久磁体项目管线又弱于分离环节。由此形成三个断点：矿石能否形成适合分离的精矿、分离氧化物能否形成稳定金属与合金、磁体能否通过汽车和国防客户认证。")
    add_picture(doc, charts["concentration"], "图4-1  磁性稀土价值链集中度（数据来源：IEA《Rare Earth Elements 2026》）")
    doc.add_heading("4.2 钨、镓、锗、锑和石墨：副产属性与规格门槛", level=2)
    add_body(doc, "钨项目的瓶颈通常不是矿物存在，而是品位、回收率、选矿稳定性、冶炼成本和长期价格。现有矿山重启可以缩短地质勘探时间，但设备更新、尾矿处理和承购仍决定现金流。对中国以外供应而言，APT和碳化钨等中间品能力比精矿本身更具战略意义。")
    add_body(doc, "镓和锗多为铝、锌、煤等产业链的副产物，增产取决于主产品产量、回收设施和精制能力。单独观察储量会夸大可供性；更关键的指标是含镓/锗原料流量、回收率、精炼纯度和下游晶圆或光学材料认证。USGS针对镓锗限制的研究也将经济影响与进口替代能力结合评估，而非只看矿产资源量。")
    add_body(doc, f"锑和石墨的供应替代面临不同问题。锑受矿山集中、冶炼环保和军民需求共同影响；石墨则必须区分天然精矿、球化提纯、人造石墨和负极材料。低纯产品的新增供应不能直接替代高规格人造石墨或通过电池认证的负极材料。{marker('FN35')}")
    project_rows = [
        ["Mountain Pass/MP Materials", "美国", "稀土矿山—分离—磁体", "商业生产/扩建", "政府贷款、价格保护、承购"],
        ["Lynas/Mt Weld—Kalgoorlie—Malaysia", "澳/马", "矿山—裂解浸出—分离", "商业生产/重稀土扩展", "日本投资与承购"],
        ["Caremag Lacq", "法国", "回收磁体/原矿—重稀土分离", "建设推进", "法日政府支持、长期供货"],
        ["Eneabba", "澳大利亚", "矿砂副产物—稀土精炼", "建设/融资", "政府融资、战略项目"],
        ["Serra Verde", "巴西", "离子吸附型稀土", "商业爬坡", "重稀土潜力、产品稳定性"],
        ["Sangdong", "韩国", "钨矿—精矿", "重启建设", "承购、资本开支与爬坡"],
        ["北美石墨项目群", "美/加", "矿山—球化—提纯—负极", "多项目建设", "IRA/贷款、客户认证"],
        ["副产镓锗回收项目", "美欧日", "铝土矿/锌冶炼/煤灰—精制", "研发至扩产", "原料流和纯度经济性"],
    ]
    add_table(doc, ["项目/集群", "地区", "价值链位置", "阶段判断", "关键条件"], project_rows,
              [2050, 900, 2300, 1400, 2710], "表4-1  代表性替代供应项目及其有效供给条件")
    add_callout(doc, "项目评估原则", "只有完成融资关闭、建设、试生产、客户认证并形成商业化交付的产能，才应计入可用供应。资源量、设计产能、政府意向和谅解备忘录分别属于不同证据层级。")

    doc.add_heading("5 加工技术与中游瓶颈：真正的竞争发生在“转化”", level=1)
    doc.add_heading("5.1 稀土分离与磁体制造", level=2)
    add_body(doc, "稀土元素化学性质相近，分离过程需要多级萃取、精确控制杂质和稳定处理放射性副产物。实验室流程可以证明技术可行，却不能证明连续工业运行的回收率、试剂消耗、废水成本和产品一致性。")
    add_body(doc, "从氧化物到磁体还要经历金属化、合金、制粉、取向压制、烧结、热处理、晶界扩散、机械加工和镀层。每一步都会引入良率损失和性能波动。中国在磁体制造端的高份额反映的不只是设备数量，还包括供应商集群、工艺经验和下游客户共同改进。")
    doc.add_heading("5.2 电池材料与石墨", level=2)
    add_body(doc, "电池供应链的脆弱性并不只在电芯。负极材料的球化、提纯、包覆和石墨化，以及正极前驱体和磷酸铁锂材料的一致性，均是决定电池安全、寿命和快充性能的中游环节。IEA指出，许多电池中游细分领域的中国份额达到80%以上，部分前驱体和LFP材料接近高度集中。")
    add_body(doc, "新建海外电池工厂如果仍依赖中国负极和前驱体，只是把最终组装地迁移，并未改变上游风险。有效多元化必须同步建设材料、设备、工艺服务和质量体系，否则在技术变更或出口许可收紧时仍会暴露。")
    doc.add_heading("5.3 经济性约束：低价周期与新项目融资", level=2)
    add_body(doc, "关键矿产项目普遍具有高前期资本、长建设周期和价格波动大的特点。低价会改善下游成本，却压缩新进入者现金流并延迟投资；价格突然上涨又可能刺激过度扩产。政府价格保护、承购和合同价差的兴起，本质上是在弥补“安全价值无法被现货价格充分计价”的市场失灵。")
    add_body(doc, "但价格支持也会产生纪律问题：项目可能依赖财政补贴而缺乏成本竞争力，下游可能因保底价格承担更高成本，盟友之间还可能争夺同一有限需求。因此政策设计应包含成本下降路径、产量门槛、收益分享和退出机制。")
    bottleneck_rows = [
        ["资源与采矿", "品位、矿物学、许可、基础设施", "资源量、可采储量、回收率、许可状态"],
        ["选矿与分离", "流程适配、杂质、放射性、连续运行", "试车时长、回收率、单位试剂/能耗"],
        ["高纯材料", "纯度、一致性、批间波动", "合格率、客户退货率、认证周期"],
        ["磁体/部件", "制粉、烧结、镀层、性能稳定", "磁性能、良率、批量交付"],
        ["商业融资", "价格波动、承购、资本开支", "融资关闭、承购覆盖率、现金成本"],
        ["物流与合规", "许可证、危险品、路线、最终用途", "审批时长、路线冗余、单证完整率"],
    ]
    add_table(doc, ["环节", "主要瓶颈", "应观察的可验证指标"], bottleneck_rows,
              [1600, 3750, 4010], "表5-1  从资源到交付的瓶颈与验证指标")

    doc.add_heading("6 替代、减量与回收：三条路径对应不同时间尺度", level=1)
    doc.add_heading("6.1 产品替代不能只比较材料性能", level=2)
    add_body(doc, "无稀土电机、铁氧体磁体、铁氮磁体和其他新材料可以降低特定场景的稀土需求，但必须比较系统级指标：功率密度、体积、效率、温度稳定性、控制系统复杂度、制造成本和认证周期。材料层面的成功不必然转化为整车、风机或国防系统的可替代性。")
    add_body(doc, "短期最现实的技术策略通常是减量而非完全替代，例如通过晶界扩散减少镝铽用量、优化磁路设计、提高回收磁体占比和延长设备寿命。完全替代需要重新设计产品和供应链，适合新平台，不适合在存量产品上仓促切换。")
    doc.add_heading("6.2 回收是“近端矿山”，但受原料和规格约束", level=2)
    add_body(doc, f"IEA认为，扩大回收可以在满足各国气候承诺的情景下，到2050年减少25%—40%的新采矿需求；但回收并不会消除新增矿山和精炼投资。回收价值在于缩短供应链、降低进口暴露并形成应急缓冲。{marker('FN13')}")
    add_body(doc, "回收项目的约束包括废料收集、产品拆解、成分混杂、商业秘密、运输规则和再生材料认证。生产废料在短期最易回收，报废电机和消费电子需要更复杂的拆解体系，新能源汽车电池则受寿命周期影响，真正的大规模退役原料要滞后于销售高峰。")
    pathway_rows = [
        ["0—2年", "库存、许可协调、双源认证、生产废料回收", "快速缓冲", "库存成本、规格不匹配"],
        ["2—5年", "回收扩产、材料减量、盟友分离/精炼爬坡", "增加可切换供给", "原料量、良率、认证"],
        ["5年以上", "新矿山、全新材料体系、产品重设计", "结构性降低依赖", "资本、技术和市场接受"],
    ]
    add_table(doc, ["时间尺度", "主要工具", "作用", "关键限制"], pathway_rows,
              [1200, 3600, 2100, 2460], "表6-1  替代与韧性工具的时间尺度")

    doc.add_heading("7 多边合作：从供应链联盟走向标准化市场", level=1)
    doc.add_heading("7.1 G7：把环境和劳工标准转化为市场条件", level=2)
    add_body(doc, f"G7《关键矿产行动计划》提出标准化市场、资本动员和创新三条主线，强调负责任开采的真实成本、可追溯性、当地价值创造和开发金融机构协作。其政策含义是：盟友试图通过标准、融资和承购形成区别于纯现货市场的“安全供应市场”。{marker('FN31')}")
    doc.add_heading("7.2 Quad：从倡议转向可量化资本工具", level=2)
    add_body(doc, f"2026年Quad框架提出动员最高200亿美元政府与私人支持，覆盖采矿、加工和回收，并列举担保、贷款、股权、保险、补贴、承购和其他商业安排。框架同时关注许可、投资审查、地质调查、价格机制和非市场政策。{marker('FN32')}")
    doc.add_heading("7.3 联盟的四个执行门槛", level=2)
    add_number(doc, "规格互认：成员国必须对产品纯度、环境标准、原产地和可追溯数据形成兼容规则。")
    add_number(doc, "融资协同：开发金融机构需要明确风险分担，避免项目在多个审批体系中重复尽调。")
    add_number(doc, "承购闭环：矿山、加工厂和下游用户的数量、价格和期限必须匹配。")
    add_number(doc, "物流冗余：港口、危险品运输、保险和紧急调拨安排必须能在冲击时运行。")
    add_body(doc, f"加拿大主导的关键矿产韧性与生产联盟进一步强调以主权工具加速项目、应对市场集中和操纵。它反映了联盟合作正在从外交声明转向项目筛选与市场干预。{marker('FN33')}")
    add_callout(doc, "独立判断：联盟数量不是韧性指标", "如果多个联盟仍支持同一批矿山、依赖同一分离技术、争夺同一承购客户，名义合作网络会产生重复计算。应以新增合格产能、物流替代路径和可执行承购量衡量联盟成效。")

    doc.add_heading("8 企业供应链重组：从采购多元化到资本共担", level=1)
    doc.add_heading("8.1 企业行为的四个层级", level=2)
    add_body(doc, "第一层是库存和订单调整，速度快但成本高；第二层是双源认证和供应商替换，需要技术与质量验证；第三层是长期承购、预付款和股权投资，企业开始承担上游项目风险；第四层是产品重新设计和材料替代，改变需求结构但周期最长。")
    add_body(doc, "企业不应把“供应商国别”作为唯一风险字段。同一企业可能在不同国家拥有矿山、分离、仓储和销售实体；同一货物也可能经过多次转运。更有解释力的画像应包括实际控制、生产地、加工地、收款主体、合同承担者、最终用户和许可证责任人。")
    doc.add_heading("8.2 政府资本改变企业边界", level=2)
    add_body(doc, "MP Materials案例表明，政府贷款、价格保护和长期承购可以同时改变企业的融资成本、销售去向和扩产节奏。企业由普通商品生产商转变为承担国家安全任务的政策载体，政府则从补贴者转为价格和需求风险的共同承担者。")
    add_body(doc, "这种模式对竞争政策和贸易规则提出新问题：受支持产品可能长期高于国际现货价格，下游企业需要决定是否支付安全溢价；政府还需防止保底机制掩盖成本失控。评价项目时，应同时披露财政承诺、产量义务、承购期限和超额收益分享。")
    enterprise_rows = [
        ["采购", "双源/多源、库存、替代规格", "供应连续性", "库存天数、合格供应商数"],
        ["合同", "长期承购、价格公式、最低量", "需求可融资性", "承购覆盖率、终止条款"],
        ["资本", "股权、预付款、共同投资", "项目控制与优先供应", "资金到位、治理权、份额"],
        ["技术", "减量、替代、产品重设计", "结构性降低需求", "性能、认证、单位用量"],
        ["合规", "分类、许可、最终用户、追溯", "降低执法与交付风险", "单证完整率、异常闭环"],
    ]
    add_table(doc, ["企业层面", "典型动作", "目的", "关键绩效指标"], enterprise_rows,
              [1300, 3000, 2300, 2760], "表8-1  企业供应链重组工具与评价指标")

    doc.add_heading("9 市场影响：从单一价格叙事转向多变量传导", level=1)
    doc.add_heading("9.1 价格、数量和交期必须分开观察", level=2)
    add_body(doc, "出口管制可能同时改变出口量、审批时间、库存、区域价差和买方结构，但这些变量也受需求周期、汇率、企业去库存和新增供给影响。专业评估应建立公告前基线、对照产品和事件窗口，不能把公告后所有变化都归因于政策。")
    add_body(doc, f"IEA数据显示，2024年锂需求增长接近30%，镍、钴、石墨和稀土需求增长约6%—8%，但大规模新增供应使多种电池金属价格下降。这说明“战略重要性上升”可以与“现货价格下跌”同时发生，低价甚至可能增加中长期供应集中风险。")
    add_body(doc, f"在稀土领域，2025年出口许可冲击使部分进口市场价格显著高于中国市场，并造成汽车企业阶段性供应困难；但随后许可证发放和出口恢复说明，市场影响取决于许可节奏而非简单禁运。{marker('FN11')}")
    doc.add_heading("9.2 建议采用四组量化指标", level=2)
    metric_rows = [
        ["监管可得性", "有效政策组、许可证申请/批准、审批时长、补正次数", "判断行政摩擦"],
        ["物理可得性", "出口量、库存天数、交付周期、路线冗余", "判断短缺程度"],
        ["价格与成本", "区域价差、合同/现货差、保险与运输成本", "判断安全溢价"],
        ["结构调整", "买方集中度、第三国份额、替代材料认证、项目交付", "判断长期重构"],
    ]
    add_table(doc, ["指标组", "建议字段", "解释目标"], metric_rows,
              [1700, 4900, 2760], "表9-1  关键矿产政策影响量化框架")
    add_body(doc, "对贸易流变化，应先处理税号口径、转口重复、同票多行、重量与金额异常、自由贸易区再出口和企业名称标准化。只有在目的国、货描、价格、路线和主体同时变化时，才适合进一步研判转运或规避风险。")

    doc.add_heading("10 法律争端与企业合规：规则冲突不会自动消失", level=1)
    doc.add_heading("10.1 中国出口合规的操作重点", level=2)
    add_body(doc, "企业首先需要判断货物、软件、技术或服务是否属于清单或临时管制物项；其次核对目的地、最终用户和最终用途；再次检查许可证的物项、数量、有效期、口岸和买方是否与实际出口一致；最后保存技术说明、检测报告、合同、付款和运输证据。")
    add_body(doc, "对参数接近管制阈值的产品，企业不应仅依赖税号或商品名称，应建立产品主数据、实验室检测和版本控制。技术服务、云端访问和远程维护也可能构成跨境提供，需纳入内部授权和日志审计。")
    doc.add_heading("10.2 WTO与国家安全边界", level=2)
    add_body(doc, f"WTO在此前稀土、钨和钼出口限制争端中审查了出口税、配额及其抗辩，但当前以国家安全和两用物项许可为核心的制度具有不同法律结构。历史裁决仍能说明，措施设计、适用一致性和证据记录会影响国际争端风险。{marker('FN34')}")
    add_body(doc, "对跨国企业而言，真正困难的是多个法域同时要求最终用途调查、制裁筛查、数据保存和供应链尽调。合规体系需要把规则冲突升级到法律部门，而不能由采购或报关人员单独决定。")
    doc.add_heading("10.3 执法证据边界", level=2)
    add_body(doc, "公开贸易数据库、船舶轨迹、企业关联和图片识别适合发现线索，但不应直接用于违法定性。完整证据通常需要原始提运单、报关单、合同、发票、付款、许可证、装箱记录、实验室检测和最终用户证明相互印证。")

    doc.add_heading("11 新兴供应国：资源民族主义与本地增值并行", level=1)
    doc.add_heading("11.1 资源国追求价值留存", level=2)
    add_body(doc, "资源国越来越倾向通过出口限制、本地加工、国家参股、税收和基础设施条件提高国内价值留存。对进口经济体而言，这既提供了中国以外资源，也可能形成新的政策风险。供应链多元化不能只把依赖从一个国家转移到另一个国家。")
    add_body(doc, "澳大利亚和加拿大拥有较成熟的监管、资本市场和盟友网络，更容易获得政策融资；印度具有市场规模和产业政策优势；非洲和拉美项目则往往需要同步建设电力、道路、港口和社区关系。不同国家的风险应按项目基础设施和治理条件分解，而非使用统一“国家风险分”。")
    doc.add_heading("11.2 中游本地化是谈判核心", level=2)
    add_body(doc, "未来项目谈判的重点将从矿权和税率，转向本地加工比例、技术转移、就业、环境治理和产品承购。进口国希望获得稳定供应，资源国希望保留更高附加值，两者之间需要通过长期合同和基础设施融资形成交换。")
    add_body(doc, "这也意味着，资源外交若只承诺购买原矿，可能不再足以获得项目支持；相反，能够提供加工技术、人才、融资和市场准入的合作方更具吸引力。")

    doc.add_heading("12 国防与先进制造：需求小、规格高、替代慢", level=1)
    doc.add_heading("12.1 国防供应链的特殊性", level=2)
    add_body(doc, "国防应用往往占总矿产需求的比例不高，却要求高纯度、极端温度性能、长期可靠性和保密认证。商业市场中存在的材料不一定能直接用于雷达、导引、航空发动机、夜视、弹药和精密控制系统。")
    add_body(doc, "供应风险因此集中在少量特种规格和单一合格供应商，而不是总吨位。国防评估应建立“系统—部件—材料—规格—供应商”映射，并区分可替代、需再认证和不可替代物项。")
    doc.add_heading("12.2 储备与产能的边界", level=2)
    add_body(doc, "战略储备适合应对短期中断，但材料可能氧化、降级或因产品设计变更而失去适用性。储备政策必须明确形态、规格、轮换、保管、释放条件和优先用户，不能只公布金额。")
    add_body(doc, "对极小批量、高价值材料，政府承购和最低采购量可能比大规模库存更有效；对广泛工业使用的磁体和石墨，则需要商业产能与多行业需求共同支撑。")
    add_callout(doc, "独立判断：国防需求的关键不是“量大”，而是“无法在需要时快速替换”", "国防材料风险应采用最长认证周期、最少合格供应商和最窄技术容差来排序，而不是只看年度进口金额。")

    doc.add_heading("13 2030情景分析：不预测单一路径，识别可观测触发器", level=1)
    add_body(doc, "情景分析的目的不是给出看似精确的概率，而是检验不同外部条件下的准备方案。以下四种情景可以并存于不同矿种：同一时期稀土可能进入价格支持市场，锂则仍处低价竞争，镓锗可能因副产扩产缓慢而维持高集中。")
    scenario_rows = [
        ["A 管理型许可与渐进多元化", "许可保持、项目逐步投产、贸易谈判稳定", "区域价差下降但不消失；中游集中缓慢下降", "审批时长、出口恢复、项目认证"],
        ["B 集团化分链", "原产地、采购和投资规则进一步集团化", "盟友市场与全球现货市场分化，成本上升", "承购排他、关税/配额、投资审查"],
        ["C 价格支持型双层市场", "更多政府采用价格底线、合同价差和储备", "安全供应获得溢价，财政承担周期风险", "政府付款、保底产量、长期合同价差"],
        ["D 技术与循环突破", "回收、减量或替代材料通过规模化认证", "部分重稀土和电池材料需求弹性提高", "再生料占比、单位用量、客户认证"],
    ]
    add_table(doc, ["情景", "触发条件", "主要结果", "领先指标"], scenario_rows,
              [1650, 2850, 2920, 1940], "表13-1  2026—2030关键矿产四种情景")
    add_body(doc, "基准判断是：到2030年前，供应链会比2024年更分散，但并不会形成完全可替代的全球网络。矿山端多元化快于分离和材料端，政策融资增长快于商业成本下降，联盟项目增加快于合格交付能力。")
    add_body(doc, "最大的上行风险是价格保护和长期承购推动一批中游项目跨过融资门槛；最大的下行风险是低价、许可延误和客户认证失败使项目管线大量推迟。监管部门应按季度更新政策状态，按月跟踪项目里程碑，按事件触发更新企业和物流风险。")

    doc.add_heading("14 独立战略洞察与海关监管建议", level=1)
    doc.add_heading("14.1 八项独立战略洞察", level=2)
    insights = [
        "洞察一：加工能力正在成为比矿权更持久的结构性权力。资源可被发现，连续工业工艺和客户认证更难复制。",
        "洞察二：政府正在成为关键矿产市场的做市商。贷款、承购、价格底线和储备共同重塑价格发现。",
        "洞察三：未来市场可能长期存在“商业价格”和“安全价格”双轨。安全溢价反映冗余、标准和政策条件，而非单纯材料成本。",
        "洞察四：联盟化会降低单一国家风险，也可能制造集团内重复投资和集团间规则不兼容。真正指标是新增合格产能，不是签署文件数量。",
        "洞察五：2025—2030年最危险的不是资源绝对短缺，而是政策即时变化与产能慢速建设之间的时间错配。",
        "洞察六：回收、减量和替代不是同一策略。回收解决原料循环，减量改善单位需求，替代则改变产品体系，时间尺度完全不同。",
        "洞察七：贸易绕行与合法转口并存。路线变化只有与受控物项、许可证缺口和异常单证共同出现时，才具有高风险意义。",
        "洞察八：海关数据的战略价值将上升。许可、报关、检测、物流、企业关系和资金数据的连接能力，决定监管能否从事后查发转向事前预警。",
    ]
    for item in insights:
        add_number(doc, item)
    doc.add_heading("14.2 海关监管系统建设建议", level=2)
    recommendations = [
        ["政策状态引擎", "按生效、暂停、恢复、过渡期和目的地差异计算当期有效规则，保留历史快照。", "立即"],
        ["物项参数知识库", "连接管制编码、税号、成分、纯度、粒度、尺寸、形态、用途和检测方法。", "立即"],
        ["许可证字段核验", "自动比对物项、数量、买方、最终用户、有效期、口岸和核销记录。", "立即"],
        ["跨票聚合模型", "按主体、地址、电话、货描、路线、收款方和时间窗口识别拆单与主体替换。", "6个月"],
        ["企业关系穿透", "区分股权控制、共同人员、共同地址、货代关系、一次性交易和长期承购。", "6—12个月"],
        ["最终用户画像", "结合买方行业、产能、设备、终端产品和关联客户核验用途合理性。", "6—12个月"],
        ["境外政策与项目监测", "记录融资、开工、试产、认证和商业交付，避免把计划产能当作现实供给。", "持续"],
        ["证据分层界面", "将官方事实、企业披露、商业数据、模型推断和待调单事项分别展示。", "立即"],
        ["实验室与机检联动", "针对高风险参数建立抽样、检测、图像和重量差异规则。", "12个月"],
        ["人工审核门槛", "涉及企业、个人和执法结论的预警必须经过来源核验、法律审查和授权。", "立即"],
    ]
    add_table(doc, ["能力模块", "建设内容", "建议时序"], recommendations,
              [1750, 6000, 1610], "表14-1  海关监管能力建设路线图")
    add_picture(doc, charts["pipeline"], "图14-1  从政策状态到人工处置的证据链")
    add_body(doc, "在应用层面，系统不应直接输出“走私企业”或“违规人员”标签，而应输出可解释的风险理由、使用的数据、适用政策状态、证据等级和下一步核查动作。对于未取得原始单证的线索，应明确标注为待核实，避免把模型相关性转化为执法结论。")
    add_body(doc, "在组织层面，建议建立政策、归类、实验室、风险、稽查和技术部门的联合维护机制。物项参数由专业部门确认，政策状态由法规人员维护，模型规则由风险部门解释，企业和人员风险由执法授权链审核。")
    add_callout(doc, "最终结论", "关键矿产治理的竞争，不是简单比拼谁拥有更多矿山或发布更多政策，而是比拼谁能把资源、加工、资本、技术、库存、合同和监管数据组织成可持续的交付体系。对海关而言，最重要的能力不是标签数量，而是政策适用准确、物项识别精准、关系穿透可解释、证据闭环可复核。")

    doc.add_heading("附录A 重点矿产与监管观察清单", level=1)
    mineral_rows = [
        ["镓", "金属、氮化物、氧化物、砷化镓等", "纯度、晶片/粉末/碎料形态、用途", "副产回收和高纯精制能力"],
        ["锗", "金属、二氧化物、四氯化物、衬底", "纯度、光学/红外/半导体用途", "锌冶炼副产与光学材料认证"],
        ["石墨", "高规格人造石墨、天然鳞片石墨及制品", "纯度、强度、密度、球化/膨胀形态", "负极材料提纯、包覆和认证"],
        ["锑", "矿料、金属、氧化物、特殊化合物", "纯度、形态、是否夹藏混装", "矿山集中、冶炼和国防需求"],
        ["钨", "矿、APT、氧化物、碳化物、合金等", "物项编码、成分、粒度、最终用途", "精矿到APT/碳化钨的中游能力"],
        ["碲", "金属、化合物和相关材料", "纯度、化合物形态、光伏/热电用途", "铜冶炼副产回收"],
        ["铋", "金属及特定化合物", "纯度、合金/化合物、用途", "铅锌冶炼副产和高纯能力"],
        ["钼", "特定粉末、合金和相关材料", "粒度、纯度、材料性能", "矿山较分散但高规格材料集中"],
        ["铟", "金属、化合物、靶材等", "纯度、靶材和显示/半导体用途", "锌冶炼副产与靶材制造"],
        ["中重稀土", "钐、钆、铽、镝、镥、钪、钇", "元素、合金、氧化物、磁体、混合物", "分离、金属化和磁体认证"],
        ["永磁体", "钐钴、含镝/铽钕铁硼等", "元素含量、性能、尺寸、最终产品", "烧结、晶界扩散、镀层和客户认证"],
        ["电池材料", "高性能电芯、正极、负极、设备与技术", "能量密度、压实密度、容量、设备参数", "部分2025扩围措施当前暂停"],
    ]
    add_table(doc, ["矿产/材料", "重点产品形态", "海关识别要点", "境外替代瓶颈"], mineral_rows,
              [1200, 2850, 2850, 2460], "表A-1  重点矿产监管与供应链观察")

    doc.add_heading("附录B 项目阶段与证据等级", level=1)
    stage_rows = [
        ["S0 资源线索", "企业或政府公布资源前景", "不能计入产能"],
        ["S1 资源确认", "有合规资源量/储量报告", "仅证明地质基础"],
        ["S2 预可研", "形成工艺路线和初步经济性", "高不确定性"],
        ["S3 可研与许可", "完成可研、环境和关键许可", "具备融资基础"],
        ["S4 融资关闭", "资金和承购基本落实", "可计条件性产能"],
        ["S5 建设", "设备安装和工程进度可验证", "关注延期与超支"],
        ["S6 试产/认证", "连续运行并向客户送样", "关注良率和规格"],
        ["S7 商业交付", "稳定出货并形成收入", "计入有效供应"],
    ]
    add_table(doc, ["阶段", "判定依据", "分析处理"], stage_rows,
              [1500, 4900, 2960], "表B-1  关键矿产项目七阶段法")
    add_body(doc, "证据等级建议：A级为政府、国际组织、法院和企业法定披露；B级为企业官网、项目融资和可复核贸易记录；C级为新闻和行业评论。项目阶段至少使用一个A级或两个相互独立B级来源确认，计划值与实际值必须分栏存储。")

    doc.add_heading("附录C 十二项专题深度分析", level=1)
    doc.add_heading("C.1 许可制为何比全面禁运更具持续性", level=2)
    add_body(doc, "全面禁运的政策效果直观，但会迅速刺激替代、库存和外交反应，也会切断本国出口企业的正常市场。许可制保留正常贸易通道，把政策压力集中于高风险用途、敏感买方和信息不完整的交易，因此更适合长期运行。其影响也更难用单一出口量衡量：分类咨询、材料补正、最终用户证明和审批节奏都会改变企业交期。")
    add_body(doc, "许可制的战略弹性来自可调参数。主管部门可以调整物项技术条件、目的地、最终用途审查、证明材料、许可证期限和核销要求，而不必修改整个法律框架。对进口方而言，这种不确定性会被转化为库存成本、替代供应溢价和产品设计保守度；对出口方而言，则表现为合规人力、检测和客户尽调成本上升。")
    add_body(doc, "因此，判断政策强弱不能只看批准或拒绝数量。更完整的指标还包括申请完整率、平均补正轮次、不同用途的批准差异、许可证有效期、实际核销量和获批后运输时长。若只统计出口总量，可能把需求下降、季节性和企业预防性库存误判为政策效果。")

    doc.add_heading("C.2 政策状态引擎的必要性", level=2)
    add_body(doc, "关键矿产公告具有生效、扩围、暂停、恢复、替代和目的地差异等多种状态。传统法规库通常只保存发布日期和全文，无法回答“某一季度、某一目的地、某一技术参数是否有效”。这会直接污染趋势图、企业风险画像和自动预警。")
    add_body(doc, "建议把每项规则拆成“政策组—物项—参数—目的地—生效区间—法律动作”六个字段。暂停不是删除原公告，而是在时间轴上形成不适用区间；恢复则新增区间。统计政策密度时，应计算当期有效物项组，而不是累计公告篇数。")
    add_body(doc, "政策状态还需要版本化。系统应保存每次判断使用的规则版本，确保事后复核时能够还原当时的适用条件。对于报关日期、合同日期、许可证签发日期和实际离境日期不一致的交易，还要明确采用哪一时点判断法律适用。")

    doc.add_heading("C.3 价格底线的经济逻辑与财政风险", level=2)
    add_body(doc, "新建关键矿产项目往往面临“成本高于现货、现货低于长期安全价值”的困境。价格底线通过补足市场价与约定价之间的差额，为项目提供可融资收入；若价格上升，再通过收益分享回收部分财政支出。相比一次性补贴，这种安排更接近长期风险共担。")
    add_body(doc, "价格保护的关键不是底线高低，而是覆盖哪些产品和数量、如何定义基准价格、是否允许库存计入、项目何时达到满产、成本下降目标如何设置。若设计过宽，企业可能失去降本动力；若设计过窄，银行仍不认可其现金流。")
    add_body(doc, "从国际影响看，价格底线可能催生两个市场：一个按全球现货竞争，另一个以安全、原产地和政策合规为条件获得溢价。后者能够维持冗余产能，却也可能被贸易伙伴视为补贴。政策制定需要同时评估产业安全、下游成本、财政敞口和贸易规则。")

    doc.add_heading("C.4 战略储备不只是“买货入库”", level=2)
    add_body(doc, "关键矿产储备首先要解决形态问题。同一元素可以以矿石、精矿、氧化物、金属、合金、粉末、磁体或部件储存，不同形态对应不同保质期、转换时间和用户范围。储存上游原料可能仍受制于加工能力，储存下游部件则可能因规格迭代快速贬值。")
    add_body(doc, "现代储备制度正在把实物库存与承购权、远期合同、需求聚合和合同价差结合。这样可以减少仓储成本，并在市场中断时获得优先供货。但合同型储备依赖对手方履约、物流和产能真实存在，不能完全替代实物库存。")
    add_body(doc, "储备规模应以“关键用户可维持时间”而非采购金额衡量。释放规则还要明确触发阈值、优先用户、价格机制、补库节奏和跨国协调，防止市场一紧张各国同时抢购，反而放大价格波动。")

    doc.add_heading("C.5 长期承购如何把资源变成可融资项目", level=2)
    add_body(doc, "矿产项目融资的核心是未来现金流是否可预测。长期承购通过锁定最低数量、定价公式和买方信用，把资源量转化为银行可评估的收入。日本在稀土项目中常将政策资金与承购份额绑定，正是因为只有投资而没有销售安排，难以保证本国获得供应。")
    add_body(doc, "承购协议需要核对五个要素：期限、最低数量、价格公式、质量规格和终止条件。宣传中“已签承购”并不等于全部产能已有买家；无约束意向书、附条件承购和具有最低采购义务的合同，融资价值差异很大。")
    add_body(doc, "对下游企业而言，长期承购也是风险资产。若替代技术出现、需求预测下降或项目成本长期高企，买方可能承担高于市场的采购义务。因此安全承购最好包含分阶段数量、质量未达标退出、价格上下限和共同降本机制。")

    doc.add_heading("C.6 稀土分离为何是最难复制的中游能力", level=2)
    add_body(doc, "稀土矿不是同质原料。不同矿体的元素配分、杂质、放射性、浸出行为和试剂体系不同，一套在某矿成功的流程不能无成本迁移到另一矿。分离厂还必须在数百乃至上千级萃取中维持稳定界面和物料平衡，任何波动都会影响纯度和回收率。")
    add_body(doc, "工业能力的核心指标不是设计吨位，而是连续运行时间、产品合格率、单耗、废水与放射性副产处理以及设备腐蚀。试生产阶段能产出样品，并不等于能够以稳定成本连续交付。客户通常还要求长期批次数据，进一步拉长商业化周期。")
    add_body(doc, "因此，海外项目应优先建设与具体矿源绑定的工艺包、现场工程团队和客户联合认证，而不是只购买通用设备。政策融资也应按试车和认证里程碑拨付，避免工程完工即被统计为有效产能。")

    doc.add_heading("C.7 磁体制造的隐性壁垒", level=2)
    add_body(doc, "高性能永磁体的壁垒集中在合金成分、粉末粒度、氧含量、取向、烧结曲线、热处理、晶界扩散、机械加工和镀层的联动控制。单个工序达标不代表最终磁性能、耐腐蚀和批次一致性满足要求。")
    add_body(doc, "汽车、风电、机器人和国防客户的认证周期不同。汽车强调大批量一致性和十年以上可靠性，风电关注长期退磁和维护成本，国防则可能要求极端环境和供应保密。一个磁体工厂获得某一客户认证，不能自动替代其他市场供应。")
    add_body(doc, "评价非中国产能时，应区分实验线、商业线、设计能力、实际产量和已认证产量。只有最后一项最接近可替代供应。设备安装和首批样品适合视为进度里程碑，不应直接写入全球市场份额。")

    doc.add_heading("C.8 副产金属的“资源量幻觉”", level=2)
    add_body(doc, "镓、锗、铟和碲等金属大量来自铝、锌或铜产业链副产。它们的可供量取决于主产品产量、杂质进入哪一股物料、回收设施是否存在以及精制经济性。即使地质资源丰富，没有相应主产业和回收回路，也难形成商业供应。")
    add_body(doc, "副产金属的供给弹性通常低于主产品。价格上涨并不会立即刺激独立矿山扩产，而是先推动现有冶炼厂增加回收、改造流程或处理历史废物。新增产能受制于原料流量和回收率上限。")
    add_body(doc, "监管分析也应关注产品形态。高纯金属、化合物、晶片、靶材和光学材料的供应商并不相同。统计“金属产量”无法解释半导体或红外产业是否真正获得合格材料。")

    doc.add_heading("C.9 石墨供应链为何容易被低估", level=2)
    add_body(doc, "石墨同时存在天然和人造路线。天然石墨需要选矿、球化、提纯和包覆；人造石墨需要针状焦或其他碳源、长周期高温石墨化和表面改性。两条路线的成本、碳排放、倍率性能和客户认证不同。")
    add_body(doc, "在电池供应链中，矿山精矿只是起点。负极材料企业需要控制粒径分布、比表面积、振实密度、首次效率和循环性能。新增天然石墨矿如果没有球化提纯和包覆能力，不能直接替代电池级负极材料。")
    add_body(doc, "对海关而言，同一税号下可能存在普通石墨制品、高规格人造石墨和负极材料。风险筛查应结合纯度、强度、密度、粒径、用途和客户行业，而不能按“石墨”关键词一概拦截。")

    doc.add_heading("C.10 回收原料的时间与空间错配", level=2)
    add_body(doc, "回收被称为“城市矿山”，但废料产生地、处理地和材料需求地往往不同。生产废料集中、成分清楚，最适合早期商业化；报废产品分散、拆解成本高且含量波动，需要回收网络和标准化设计。")
    add_body(doc, "电池和永磁体的回收还存在时间错配。产品销售高峰多年后才形成报废高峰，当前快速扩张的需求不能完全依靠退役料满足。短期内，生产废料和工业设备更新是更现实的来源。")
    add_body(doc, "跨境运输规则会影响回收规模。若废料被认定为危险废物，运输、保险和许可成本上升；若规则过松，又可能形成环境风险。联盟国家需要统一废料分类、追溯和再生材料认证，才能把分散原料聚合为规模经济。")

    doc.add_heading("C.11 标准化市场与“安全溢价”", level=2)
    add_body(doc, "G7等机制提出标准化市场，试图让环境、劳工、反腐、可追溯和当地价值创造进入价格。其经济逻辑是：若低标准供应只按现金成本竞争，高标准新项目难以获得融资；安全和责任标准必须通过承购、采购规则或价格机制获得回报。")
    add_body(doc, "但标准也可能成为新的市场壁垒。不同联盟若采用不兼容的碳足迹、原产地和追溯规则，会增加企业合规成本，甚至排除发展中国家项目。标准设计需要透明、可验证、逐步实施，并为资源国提供能力建设。")
    add_body(doc, "对中国企业而言，标准化市场既是挑战也是机会。能够提供可追溯、低碳和稳定交付证明的企业，可能在安全溢价市场继续获得位置；缺乏产品级数据和供应链尽调的企业，则可能被排除于高标准采购。")

    doc.add_heading("C.12 贸易数据和企业关系的正确用法", level=2)
    add_body(doc, "商业提单和海关数据适合识别交易路线、主体变化和货物描述，但覆盖范围、字段质量和重复记录差异很大。分析前必须进行企业名称归一、同票多行合并、单位换算、贸易方式区分和转口校正。")
    add_body(doc, "企业关联也有多种强度：控股、共同实控、共同董事、同地址、品牌授权、货代代理和一次性交易不能等量齐观。主体替换风险应要求多项关系同时出现，并与查发时间、相似产品、相同买方或路线重合形成时间证据。")
    add_body(doc, "任何自动化模型都应输出“为什么预警”，例如：政策已生效、商品参数接近阈值、许可证字段不一致、同一买方在多个主体下连续小批量采购。模型不应输出未经人工审核的违法结论。")

    doc.add_heading("附录D 重点矿种独立研判", level=1)
    mineral_deep_dives = [
        ("D.1 镓：从原料优势到化合物和晶圆能力",
         "镓的战略价值集中在氮化镓、砷化镓等化合物半导体。其供应风险不能只用金属镓吨位判断，还要观察高纯精制、外延片、衬底和器件制造。海外扩产的现实路径多是从铝土矿或氧化铝流程回收粗镓，再建设高纯精制；这要求铝产业愿意改造副产回路。",
         "监管上，金属、化合物、晶片、粉末和碎料可能具有不同技术属性。企业若只在发票中写“半导体材料”或“样品”，会增加分类不确定性。建议以化学式、纯度、晶体形态、尺寸和最终用途构建产品主数据。"),
        ("D.2 锗：光纤、红外与半导体用途的规格分化",
         "锗供应主要受锌冶炼副产、煤相关资源和回收影响。红外光学、光纤和半导体对纯度和形态的要求差别较大，普通金属锗不能直接替代高规格四氯化锗、二氧化锗或单晶材料。",
         "海外替代项目应关注锗从哪一股冶炼物料回收、精炼纯度、客户认证和副产原料长期可得性。海关核查则应比对化合物形态、纯度、光学参数和买方行业，避免只按锗含量判断用途。"),
        ("D.3 锑：矿山、冶炼与军民需求共同驱动",
         "锑的供给风险来自矿山和冶炼集中、环保约束以及阻燃、电池、光学和国防需求叠加。新增矿山如果缺少稳定冶炼和有害元素处理，难以快速形成合格金属或氧化物供应。",
         "锑产品从矿料、金属到高纯化合物差异很大。高风险交易通常需要同时出现品名模糊、重量密度异常、第三方发货或最终用户不匹配。单纯的价格上涨或关联企业变更不应直接等同于走私。"),
        ("D.4 钨：中游产品比精矿更决定产业安全",
         "钨矿资源在多国存在，但从精矿到仲钨酸铵、氧化钨、钨粉、碳化钨和硬质合金的转化决定刀具、航空航天和国防供应。矿山重启可以增加原料，却不能立即补足高规格粉末和合金。",
         "钨产品识别需要成分、粒度、碳含量、牌号和用途。碳化钨粉末、混合料和制品可能分属不同税号或货描，监管系统应把两用物项编码与材料参数、生产工艺和最终用户连接。"),
        ("D.5 石墨：矿端多元化与负极端集中并存",
         "非洲、加拿大、澳大利亚和美国的天然石墨项目增加了矿端选择，但球化、提纯、包覆和客户认证仍是瓶颈。人造石墨又依赖高温石墨化设备、能源和工艺，不能用天然石墨产量替代。",
         "政策分析应分别统计精矿、球化石墨、提纯球形石墨、人造石墨和成品负极材料。若把不同层级合并，容易得出“供应已经多元化”的错误结论。"),
        ("D.6 稀土：必须按元素和价值链阶段拆分",
         "“稀土”是17种元素集合，轻稀土与重稀土在资源、分离和用途上差异显著。钕镨决定磁体主体性能，镝铽用于高温稳定，钇、钆、钐等又有不同光学、医疗和国防用途。总稀土氧化物产量不能代表关键元素可得性。",
         "研究和监管应至少区分矿石、混合精矿、单一氧化物、金属、合金、磁粉和磁体。对磁体还要观察元素含量、牌号、尺寸、性能和最终部件。"),
        ("D.7 铟与碲：显示、光伏和热电产业的副产约束",
         "铟多与锌产业相关，碲多与铜冶炼相关。两者需求受显示、薄膜光伏、热电和电子材料技术路线影响，供给却受主产品冶炼量制约。价格上升能推动回收和工艺改进，但难以像主矿产品一样快速扩产。",
         "项目评估应关注回收位置、精炼收率和高纯材料客户，而不是只看国家储量。监管则需区分金属、化合物、靶材和制品，结合纯度和用途判断。"),
        ("D.8 铋与钼：不能因市场更分散而忽视特定规格",
         "铋和钼在全球资源和生产上相对部分稀有金属更分散，但特定高纯粉末、化合物、合金或军工用途仍可能集中。总量供应充足不意味着所有受控规格都容易替代。",
         "风险评估应从“矿种级”下沉到“产品规格级”。若境外替代能力只覆盖普通商品，受控高规格材料仍可能形成单点故障。"),
        ("D.9 电池材料：扩围暂停不等于风险消失",
         "2025年部分锂电池、正负极材料、设备和技术扩围措施目前处于暂停状态。系统应在暂停区间停止把这些新增物项计为有效管制，但仍需保留公告内容和可能恢复的时间节点。",
         "产业风险仍然存在，因为电池中游集中、设备和工艺依赖并未因法律暂停而改变。政策风险与产业集中风险应在界面上分开显示，避免把产业脆弱性误写成当前法律限制。"),
        ("D.10 超硬材料：粒径、设备参数和用途边界",
         "人造金刚石从装饰用途到精密加工、半导体和先进制造跨度很大。管制边界通常依赖粒径、线径、破断拉力、工作速度或设备工艺等技术条件，而不是“金刚石”这一商品名称。",
         "暂停措施下，新增扩围物项的法律状态需要准确标识；既有锑及超硬材料公告中的相关管制仍应单独判断。海关核查应保留检测和设备技术资料，避免将首饰培育钻石与工业物项混同。"),
    ]
    for title, p1, p2 in mineral_deep_dives:
        doc.add_heading(title, level=2)
        add_body(doc, p1)
        add_body(doc, p2)

    doc.add_heading("附录E 核心来源说明", level=1)
    add_body(doc, "本报告脚注选取37项核心来源，优先使用中国商务部、国务院及海关相关文件，IEA和USGS统计与方法报告，美国、欧盟、日本、澳大利亚、印度、加拿大等政府政策文件，以及企业向美国证券监管机构提交的法定披露。二手报道仅作为检索线索，不作为核心判断的唯一依据。")
    for code, text in SOURCES:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.left_indent = Inches(0.34)
        p.paragraph_format.first_line_indent = Inches(-0.2)
        p.paragraph_format.space_after = Pt(2)
        set_run_font(p.add_run(text), size=8.5, color=MID_GRAY)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    full_text = "\n".join(p.text for p in doc.paragraphs)
    nonspace = len(re.sub(r"\s+", "", full_text))
    chinese = len(re.findall(r"[\u4e00-\u9fff]", full_text))
    return nonspace, chinese, len(doc.paragraphs), len(doc.tables)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("output")
    args = parser.parse_args()
    output = Path(args.output)
    counts = build_report(output)
    print("NONSPACE_CHARS", counts[0])
    print("CHINESE_CHARS", counts[1])
    print("PARAGRAPHS", counts[2])
    print("TABLES", counts[3])
    for code, text in SOURCES:
        print(f"SOURCE|{code}|{text}")


if __name__ == "__main__":
    main()
