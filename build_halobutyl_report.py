from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


BASE = Path(r"D:\易迅数据\反倾销税深度分析报告\01_卤化丁基橡胶")
SUMMARY = json.loads((BASE / "卤化丁基橡胶_易迅综合分析摘要.json").read_text(encoding="utf-8"))
OUT = BASE / "卤化丁基橡胶_反倾销税与第三国转运风险深度分析报告.docx"

BLUE = "17324D"
MID = "2A6F97"
LIGHT = "EAF2F8"
AMBER = "FFF3CD"
GREEN = "E2F0D9"
RED = "FDE2E2"
GRAY = "E5E7EB"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, size=10.5, bold=False, color="253444", latin="Aptos", east="Microsoft YaHei") -> None:
    run.font.name = latin
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east)


def add_text(p, text: str, **kwargs):
    r = p.add_run(text)
    set_run_font(r, **kwargs)
    return r


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_table(doc, headers, rows, widths=None, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, MID)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p, str(h), size=font_size, bold=True, color="FFFFFF")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_text(p, "" if value is None else str(value), size=font_size)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    return table


def add_callout(doc, title: str, text: str, fill=AMBER, title_color="674D00"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_repeat_table_header(table.rows[0])
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=160, bottom=160, start=180, end=180)
    p = cell.paragraphs[0]
    add_text(p, title + "\n", size=11.5, bold=True, color=title_color)
    add_text(p, text, size=10.5, color="253444")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(1.8)
sec.bottom_margin = Cm(1.6)
sec.left_margin = Cm(2.0)
sec.right_margin = Cm(2.0)
sec.header_distance = Cm(0.8)
sec.footer_distance = Cm(0.8)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(10.5)
normal.paragraph_format.line_spacing = 1.25
normal.paragraph_format.space_after = Pt(6)
for style_name, size, color in (("Heading 1", 15, BLUE), ("Heading 2", 12.5, MID), ("Heading 3", 11.5, MID)):
    st = styles[style_name]
    st.font.name = "Aptos Display"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(10)
    st.paragraph_format.space_after = Pt(5)
    st.paragraph_format.keep_with_next = True

header = sec.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
add_text(hp, "中国反倾销税商品深度排查｜卤化丁基橡胶", size=8.5, color="6B7280")
footer = sec.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_text(fp, "仅供风险研判使用｜易迅数据不等同海关全量统计", size=8, color="6B7280")

# Cover
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(70)
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
add_text(p, "卤化丁基橡胶", size=28, bold=True, color=BLUE, latin="Aptos Display")
p2 = doc.add_paragraph()
p2.paragraph_format.space_after = Pt(24)
add_text(p2, "反倾销税与第三国转运风险深度分析报告", size=20, bold=True, color=MID, latin="Aptos Display")
meta = add_table(doc, ["项目", "内容"], [
    ["查询平台", "易迅数据（页面逐页采集）"],
    ["数据期", "2025-08-06至2026-08-06"],
    ["查询口径", "HS 400239；关键词 BROMOBUTYL"],
    ["全量范围", "HS结果24页/4,738票；关键词结果14页/2,762票"],
    ["报告日期", "2026-08-12"],
    ["保密提示", "仅供内部风险研判和后续执法核查，不作为直接定案依据"],
], widths=[4.0, 12.0], font_size=9.5)
doc.add_paragraph()
add_callout(doc, "核心结论", "目前没有形成卤化丁基橡胶经第三国绕道进入中国的具体证据。已识别的沙特2222/2255链路存在当地真实产能这一强反证；9票英国/比利时原产直达中国记录属于高优先级核税线索，而非转口证据。", fill=AMBER)
doc.add_page_break()

doc.add_heading("一、执行摘要", level=1)
doc.add_paragraph("本次对两种查询口径的全部页面、全部记录进行采集和逐票判定。两口径原始合计7,500票，按页面可见字段精确去重后为5,582票；其中明确纳入3,206票、待核1,472票、排除904票。排除项主要为药用胶塞、轮胎或其他制成品，以及货描与被征税原料不符的记录。")
add_table(doc, ["风险事项", "数量/规模", "结论", "置信度"], [
    ["受税来源直达中国", "9票，335,972.8kg", "应核查反倾销税是否足额缴纳", "高优先级线索"],
    ["沙特原产进入中国", "31票，5,063,218.02kg", "当地有真实产能，暂不支持绕道判断", "低风险/需证件核验"],
    ["印度原产进入中国", "1票，10件，USD25样品", "印度已终止调查，不涉及该项反倾销税", "低风险"],
    ["受税来源→第三国", "3,257票（纳入/待核）", "仅表明全球贸易流，不足以证明再出口中国", "背景信息"],
    ["两段主体重合", "Exxon Mobil集团名重合", "集团内生产销售可解释，缺少同箱/提单/数量闭合", "线索，不是证据"],
], widths=[4.0, 4.0, 7.0, 3.0], font_size=9)

doc.add_heading("二、政策口径与涉及税种", level=1)
doc.add_paragraph("被征税产品为卤化丁基橡胶（Chlorobutyl Rubber/Bromobutyl Rubber），中国税则号列为40023910、40023990。")
add_table(doc, ["来源/企业", "反倾销税率", "生效情况"], [
    ["美国公司", "75.5%", "自2024-08-20起继续实施5年"],
    ["ARLANXEO Belgium NV", "27.4%", "自2024-08-20起继续实施5年"],
    ["其他欧盟公司", "71.9%", "自2024-08-20起继续实施5年"],
    ["英国公司", "71.9%", "自2024-08-20起继续实施5年"],
    ["ARLANXEO Singapore Pte. Ltd.", "23.1%", "自2024-08-20起继续实施5年"],
    ["其他新加坡公司", "45.2%", "自2024-08-20起继续实施5年"],
    ["日本丁基株式会社/其他日本", "15.0%/30.1%", "自2026-03-14起实施5年"],
    ["加拿大公司", "13.8%", "自2026-03-14起实施5年"],
], widths=[7.0, 4.0, 6.0], font_size=9)
doc.add_paragraph("税款计算口径：反倾销税额=海关确定的进口货物计税价格×适用反倾销税率。进口环节增值税计税基础还应计入关税和反倾销税；卤化丁基橡胶通常适用13%进口增值税率。因此，如反倾销税未征或少征，还会产生相应进口环节增值税差额。")
add_callout(doc, "为何本报告不直接给出9票的税额", "易迅记录未提供中国海关审定完税价格，且平台金额字段来源/币种不统一。以重量或第三方金额替代完税价格计算税款会造成误导。取得报关单和税款缴款书后，可按上述公式逐票核算：少缴反倾销税=完税价格×适用税率；由此增加的进口增值税差额≈少缴反倾销税×13%（最终以海关组成计税价格核定为准）。")

doc.add_heading("三、易迅数据完整性与方法", level=1)
doc.add_paragraph("查询一：HS前6位400239，结果4,738票，页面设为200条/页后共24页，已逐页采集，末页138票。查询二：产品英文关键词BROMOBUTYL，结果2,762票，共14页，已逐页采集，末页162票。两口径按数据源、方向、日期、HS、货描、买卖双方、重量、数量、金额、目的地、原产地等页面可见字段精确去重。")
doc.add_paragraph("判定分为“纳入、待核、排除”。明确出现BROMOBUTYL、CHLOROBUTYL、CIIR、BIIR或可识别牌号的原料记录纳入；仅命中宽税号或泛称橡胶的记录待核；4014/4015/4016项下胶塞、密封件等制成品及明确轮胎/部件货描排除。")

doc.add_heading("四、高优先级核税线索：受税来源直达中国", level=1)
direct = [r for r in SUMMARY["对华明细"] if r["链路类型"] == "受税来源直达中国"]
rows = []
for r in direct:
    rate = "71.9%" if r["原产国地区"] == "United Kingdom" else "27.4%或71.9%（视生产商）"
    rows.append([r["日期"], r["商品描述"], r["采购商"], r["供应商"], f'{r["重量数值"]:,.2f}', r["原产国地区"], rate])
add_table(doc, ["日期", "货描", "采购商", "供应商", "重量kg", "原产地", "适用税率线索"], rows, widths=[2.0, 5.0, 3.7, 3.7, 2.0, 2.0, 2.5], font_size=7.5)
doc.add_paragraph("上述9票中：英国原产7票、227,434.8kg；比利时原产2票、108,538kg。重点实体包括WEST PHARMACEUTICAL SERVICES、ARCHILINKS TRANSPAC LIMITED、SINOCHEM PLASTICS CO., LTD、RAVAGO DISTRIBUTION CENTER NV、JIUCHUAN INTERNATIONAL (HK) CO.，以及供应链中的EXXONMOBIL CHEMICAL ASIA PACIFIC、ELITE PALOUME EXXONMOBIL、RAVAGO HONG KONG LIMITED。")
doc.add_paragraph("风险点不在“是否绕道”，而在申报原产地、生产商身份与税率是否匹配。例如，比利时原产只有ARLANXEO Belgium NV适用27.4%，其他欧盟生产商适用71.9%；贸易商RAVAGO并不自动等同于列名生产商。英国原产适用71.9%。")

doc.add_heading("五、第三国链路专项分析", level=1)
doc.add_heading("（一）沙特阿拉伯→中国：大体量但有真实产能反证", level=2)
doc.add_paragraph("关键词查询识别31票沙特原产进入中国记录，合计5,063,218.02kg，货描集中为EXXON BROMOBUTYL 2222/2255，平台显示买卖双方均为Exxon Mobil Corporation。仅看主体和牌号会形成集团链路线索，但公开资料显示：沙特朱拜勒KEMYA项目建有年产110,000吨卤化丁基橡胶装置，采用ExxonMobil技术；公开牌号资料亦确认2222为低门尼溴化丁基橡胶、2255为高门尼溴化丁基橡胶。")
add_callout(doc, "判定", "沙特具备与票载品名和牌号一致的本地生产能力，5,063吨规模远低于11万吨/年产能，因此不存在明显供给能力矛盾。现有证据更支持真实生产后对华销售，第三国绕道风险评为低。仍建议核对KEMYA/生产厂声明、批次号和沙特原产证，以排除集团主体字段映射错误。", fill=GREEN, title_color="215E21")

doc.add_heading("（二）印度→中国：只有1票原料样品，不涉及现行反倾销税", level=2)
doc.add_paragraph("排除胶塞等制成品后，仅余1票原料样品：2025-10-15，CIIR IMPRAMER C 1139，10件，申报金额USD25，Reliance Sibur Elastomers Private Limited发往SIBUR INTERNATIONAL TRADING (SHANGH...)。商务部初裁因印度进口量可忽略而终止对印度的调查，终裁亦未对印度征税。因此，该票不涉及规避卤化丁基橡胶反倾销税。")

doc.add_heading("（三）其他中间国：未形成双端闭合证据", level=2)
doc.add_paragraph("受税来源向墨西哥、印度、越南、印度尼西亚等第三国流量较大，但在本次对华端记录中没有形成同一中间国、同一牌号、同一企业和相容时间/数量的闭环。跨段实体名仅发现Exxon Mobil集团层面重合；集团拥有多国产能和全球销售体系，主体重合本身不能证明换单或原产地规避。未发现同一集装箱、同一提单、相同批次或近似数量在短期内从受税来源进入第三国后再对华的证据。")

doc.add_heading("六、风险分级与建议核查对象", level=1)
add_table(doc, ["优先级", "对象", "核查事项", "拟解决问题"], [
    ["A", "9票英国/比利时原产直达中国记录", "进口报关单、原产证、生产商声明、合同发票、反倾销税缴款书、进口增值税缴款书", "是否按71.9%或27.4%/71.9%足额缴税"],
    ["A", "RAVAGO HONG KONG LIMITED相关2票", "确认实际生产商是否ARLANXEO Belgium NV；核对货物批次和原产证", "比利时原产适用列名或其他欧盟税率"],
    ["A", "EXXONMOBIL/ELITE PALOUME相关英国原产票", "确认生产商和装运地、原产规则、税款缴纳", "英国71.9%是否准确执行"],
    ["B", "沙特31票Exxon 2222/2255", "KEMYA生产厂、批次、沙特原产证、集团内合同和运输单证", "验证真实生产，排除平台原产字段错误"],
    ["C", "印度10件样品", "用途、无商业价值样品申报和后续流向", "确认非批量规避"],
], widths=[1.5, 4.5, 7.5, 5.0], font_size=8.5)

doc.add_heading("七、证据等级、反证与局限", level=1)
doc.add_paragraph("1. 已证实：两种查询口径的全部页面和全部行已完成采集；9票受税来源直达中国及其重量、主体和货描可由页面记录复核；沙特31票及其牌号和重量可由页面记录复核。")
doc.add_paragraph("2. 线索：直达中国记录是否少缴税，需以中国报关单和税款缴款书确认；平台“原产国”字段可能来自不同数据源，不能等同原产地法律结论。")
doc.add_paragraph("3. 未证实：没有第三国绕道的具体证据；没有同一集装箱、提单、批次、换单记录或数量闭合证据。")
doc.add_paragraph("4. 反证：沙特有与2222/2255匹配的本地卤化丁基产能；印度不在终裁征税范围；部分关键词记录是胶塞等制成品，已排除。")
doc.add_paragraph("5. 局限：易迅覆盖多国数据源但不等同海关全量统计；金额字段币种/口径不统一，不能直接作为中国完税价格；缺少装运港、卸货港、船名航次、箱号和报关单号等关键字段。")

doc.add_heading("八、下一步取证清单", level=1)
for text in [
    "调取9票对应中国进口报关单：商品编号、规格型号、原产国、启运国、境外发货人、境内收货人、申报价格和征免性质。",
    "调取反倾销税和进口环节增值税税款缴款书，逐票按适用生产商税率复算。",
    "核验原产地证、生产商声明、批次/牌号、包装唛头和工厂装运证明，特别是RAVAGO贸易链。",
    "取得历史提单字段：提单号、箱号、封志号、船名航次、装卸港、承运人、通知方，以开展同箱/短期换单匹配。",
    "对沙特31票核对KEMYA/朱拜勒工厂批次与产地证；如能对应，则可进一步排除转口疑点。",
]:
    p = doc.add_paragraph(style=None)
    p.style = doc.styles["List Bullet"]
    add_text(p, text, size=10.5)

doc.add_heading("九、查询日志与来源", level=1)
add_table(doc, ["来源", "查询/资料", "用途"], [
    ["易迅数据", "HS 400239，2025-08-06至2026-08-06，4,738票/24页", "全球基线、直达中国与第三国流量"],
    ["易迅数据", "关键词BROMOBUTYL，同期2,762票/14页", "补充错分/缺失HS记录和牌号链路"],
    ["商务部2024年第32号公告", "https://www.mofcom.gov.cn/zwgk/zcfb/art/2024/art_fb218a61ce0f449cb678ce91617cfe1e.html", "美国、欧盟、英国、新加坡税率及公式"],
    ["商务部2026年第15号公告", "https://cacs.mofcom.gov.cn/cacscms/article/jkdc?articleId=187465&type=1", "日本、加拿大终裁税率；印度终止调查"],
    ["ExxonMobil公开资料", "https://investor.exxonmobil.com/sec-filings/all-sec-filings/content/0001193125-14-112142/d687041dex991.htm", "沙特朱拜勒卤化丁基项目和产能背景"],
    ["TechnipFMC公告", "https://www.technipfmc.com/en/investors/archives/technip/press-releases/technip-awarded-contract-for-a-substantial-elastomer-project-in-saudi-arabia/", "KEMYA卤化丁基装置110,000吨/年"],
    ["ExxonMobil牌号表", "https://www.exxonmobilchemical.com/-/media/project/wep/shared/aprimo/2026/03/31/04/29/40666/butyl-gradeslate-85years-en.pdf", "2222/2255牌号属性"],
    ["国家税务总局", "https://tianjin.chinatax.gov.cn/11200000000/0300/030005/20260212160138427.shtml", "进口货物一般税率13%"],
], widths=[4.0, 10.0, 5.0], font_size=8)

doc.add_heading("十、随附成果", level=1)
doc.add_paragraph("1. 《卤化丁基橡胶_易迅逐票判定台账.xlsx》：含核查摘要、路线汇总、政策口径和5,582票逐票明细。")
doc.add_paragraph("2. 《卤化丁基橡胶_易迅逐票判定台账_合并去重.csv》：机器可读逐票底表。")
doc.add_paragraph("3. 《卤化丁基橡胶_易迅综合分析摘要.json》：查询完整性、路线汇总和对华明细。")

doc.core_properties.title = "卤化丁基橡胶反倾销税与第三国转运风险深度分析报告"
doc.core_properties.subject = "中国进口反倾销税风险核查"
doc.core_properties.author = ""
doc.core_properties.keywords = "卤化丁基橡胶,反倾销税,第三国转运,易迅数据"
doc.save(OUT)
print(OUT)
