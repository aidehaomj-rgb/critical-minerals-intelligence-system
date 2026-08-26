from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


DIR = Path(r"D:\易迅数据\反倾销税深度分析报告\05_氯氰菊酯")
OUT = DIR / "氯氰菊酯_反倾销税与第三国转运风险深度分析报告.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
MUTED = "666666"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
PALE_GOLD = "FFF4CC"
PALE_RED = "FCE8E6"
PALE_GREEN = "E6F4EA"
BORDER = "CBD5E1"
RISK_RED = "9B1C1C"
CAUTION = "7A5A00"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill: str):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    node = trpr.find(qn("w:tblHeader"))
    if node is None:
        node = OxmlElement("w:tblHeader")
        trpr.append(node)
    node.set(qn("w:val"), "true")


def set_table_geometry(table, widths_dxa):
    table.autofit = False
    tbl = table._tbl
    tblpr = tbl.tblPr
    for tag in ("w:tblW", "w:tblInd", "w:tblLayout", "w:tblCellMar"):
        old = tblpr.find(qn(tag))
        if old is not None:
            tblpr.remove(old)
    tblw = OxmlElement("w:tblW")
    tblw.set(qn("w:w"), "9360")
    tblw.set(qn("w:type"), "dxa")
    tblpr.append(tblw)
    tblind = OxmlElement("w:tblInd")
    tblind.set(qn("w:w"), "120")
    tblind.set(qn("w:type"), "dxa")
    tblpr.append(tblind)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr.append(layout)
    margins = OxmlElement("w:tblCellMar")
    for tag, val in (("top", 80), ("bottom", 80), ("start", 120), ("end", 120)):
        node = OxmlElement(f"w:{tag}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    tblpr.append(margins)
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths_dxa):
            tcpr = cell._tc.get_or_add_tcPr()
            tcw = tcpr.find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                tcpr.append(tcw)
            tcw.set(qn("w:w"), str(width))
            tcw.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def set_table_borders(table, color=BORDER, size="4"):
    tblpr = table._tbl.tblPr
    borders = tblpr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblpr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def format_cell(cell, bold=False, color="000000", size=9.3):
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.05
        for run in paragraph.runs:
            set_run_font(run, size=size, color=color, bold=bold)


def add_table(doc, headers, rows, widths_dxa, header_fill=LIGHT_GRAY, font_size=9.3):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for idx, text in enumerate(headers):
        table.cell(0, idx).text = str(text)
        set_cell_shading(table.cell(0, idx), header_fill)
        format_cell(table.cell(0, idx), bold=True, color=INK, size=font_size)
    set_repeat_header(table.rows[0])
    for row_values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_values):
            cells[idx].text = "" if value is None else str(value)
            format_cell(cells[idx], size=font_size)
    set_table_geometry(table, widths_dxa)
    set_table_borders(table)
    return table


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_run_font(r, size=9, color=MUTED, italic=True)
    return p


def add_callout(doc, title, text, fill=CALLOUT, title_color=INK):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, size=10.5, color=title_color, bold=True)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.1
    r2 = p2.add_run(text)
    set_run_font(r2, size=10.5, color="333333")
    set_table_geometry(table, [9360])
    set_table_borders(table, color=BORDER, size="6")
    return table


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    new_run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), "Calibri")
    rfonts.set(qn("w:hAnsi"), "Calibri")
    rfonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    rpr.extend([rfonts, color, underline])
    new_run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_page_field(paragraph):
    r = paragraph.add_run("第 ")
    set_run_font(r, size=9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    field_run = paragraph.add_run()._r
    field_run.extend([begin, instr, separate, text, end])
    r2 = paragraph.add_run(" 页")
    set_run_font(r2, size=9, color=MUTED)


def add_body_rule(paragraph, color=BLUE, size="14"):
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)
    ppr.append(pbdr)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    return p


def add_body(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run_font(r)
    return p


summary = json.loads((DIR / "氯氰菊酯_全量分析结果.json").read_text(encoding="utf-8"))

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

# standard_business_brief preset
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(11)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.10
normal.paragraph_format.widow_control = True
for level, size, color, before, after in (
    (1, 16, BLUE, 16, 8),
    (2, 13, BLUE, 12, 6),
    (3, 12, DARK_BLUE, 8, 4),
):
    style = doc.styles[f"Heading {level}"]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(size)
    style.font.color.rgb = rgb(color)
    style.font.bold = True
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

# memo_masthead running furniture
header_p = section.header.paragraphs[0]
header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
hr = header_p.add_run("反倾销税深度分析  |  氯氰菊酯")
set_run_font(hr, size=9, color=MUTED)
footer_p = section.footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
add_page_field(footer_p)

p = doc.add_paragraph()
r = p.add_run("贸易救济合规风险研究")
set_run_font(r, size=10.5, color=BLUE, bold=True)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(5)
r = p.add_run("氯氰菊酯")
set_run_font(r, size=23, color="000000", bold=True)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(14)
r = p.add_run("反倾销税、直接进口与第三国转运风险深度分析报告")
set_run_font(r, size=14, color="373737")
for label, value in (
    ("措施对象", "原产于印度的Cypermethrin / Cypermethrin technical / Cipermethrin及终裁列明3个CAS"),
    ("政策口径", "商务部公告2025年第24号；2025年5月7日起实施5年；中国申报商品编号2926909013"),
    ("易迅口径", "5个两年工作簿共1,145条；精确去重1,029条；每条完成范围与路线判定；另保留对华页面19条"),
    ("报告日期", "2026年8月13日（中国标准时间）"),
    ("证据原则", "商业中介/A腿、物理转运、伪报原产地、少缴税款分层；平台字段不等同中国海关原始申报"),
):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r1 = p.add_run(f"{label}：")
    set_run_font(r1, size=10.5, bold=True)
    r2 = p.add_run(value)
    set_run_font(r2, size=10.5)
rule = doc.add_paragraph()
rule.paragraph_format.space_before = Pt(8)
rule.paragraph_format.space_after = Pt(10)
add_body_rule(rule)

add_heading(doc, "一、执行摘要", 1)
add_callout(
    doc,
    "核心结论（直接进口核查风险高；已发生第三国绕道中国的证据低）",
    "两年下载数据和保存页面记录共同证明：措施后仍有列名印度生产商向中国发运技术级氯氰菊酯；最具体一票是2026年2月28日Tagros→IPO LTD SHNGHAI，16,000kg、USD95,200、92%、CAS52315-07-8。另一方面，5个下载表中没有第三国→中国的技术级B腿，公开来源也没有查缉、法院判决、行政处罚或官方反规避认定。UPL Mauritius、Sundat Singapore等记录可以证明第三国商业开票/关联销售与印度→越南A腿，但记录仍明确印度原产，不能改写为洗产地或逃税证据。",
    fill=PALE_GOLD,
    title_color=CAUTION,
)
add_caption(doc, "表1  主要结论、数量和证据等级")
add_table(
    doc,
    ["事项", "数量/金额", "证据等级", "结论"],
    [
        ["下载数据覆盖", "1,145原始行；1,029精确唯一；1,024格式镜像去重；1,009保守商业票", "已证实", "所有精确唯一行均在逐票台账保留范围分类、路线风险、理由与建议动作"],
        ["措施范围候选", "格式镜像去重531条；约4,699,472.685kg", "规则判定", "候选不等于最终进口措施范围；须以纯度、用途、CAS、COA及中国申报编号复核"],
        ["措施后直接印度→中国", "下载表1票/16,000kg/USD95,200；页面保守普通技术级3批/68,000kg/USD392,800", "A至B", "直接应税核查强；页面批次仍需中国报关单和提单去重"],
        ["第三国→中国B腿", "下载表0条", "未发现", "现有数据无法闭合印度→第三国→中国链"],
        ["越南进口侧A腿", "23票/162,325.005kg/VND29,053,847,387.5", "A腿已证实", "全部明确印度原产；无越南→中国B腿"],
        ["UPL Mauritius", "措施后1票/18,000kg/VND3,200,868,000至UPL Vietnam", "B+商业链", "第三国商业供应商明确；原产地仍为印度，不证明物理经过毛里求斯"],
    ],
    [1900, 2300, 1700, 3460],
)

add_heading(doc, "二、政策范围、税率与税款逻辑", 1)
add_body(doc, "商务部公告2025年第24号于2025年5月6日发布，自2025年5月7日起，对原产于印度的氯氰菊酯征收反倾销税，期限5年，预计至2030年5月6日。调查产品包括Cypermethrin、Cypermethrin technical和Cipermethrin，分子式C22H19Cl2NO3；有效成分CAS为52315-07-8、67375-30-8和1315501-18-8。8位税号29269090仅供参考，海关执行使用10位商品编号2926909013；同一税则号下其他产品不在措施范围。")
add_caption(doc, "表2  生产商税率与核查含义")
add_table(
    doc,
    ["印度生产商/出口商", "终裁税率", "本次数据线索", "核查重点"],
    [
        ["Tagros Chemicals India Pvt. Ltd", "48.4%", "措施后对华16吨；对越南多票", "产品范围、实际中国进口人、税款缴款书；出口HS38089135与中国执行税号需交叉核验"],
        ["Meghmani Organics Limited", "62.0%", "页面对华36吨；对UPL Vietnam去重4票/72吨", "中国进口人；CAS/纯度；批次；是否按62%缴纳"],
        ["Gharda Chemicals Limited", "75.7%", "包装指纹可检索", "93—94%、25/225kg桶、UN3352与批号"],
        ["UPL Limited", "166.2%", "官方关联UPL Mauritius与UPL Shanghai；越南集团链", "合同/发票/付款、生产商栏、原产地证；关联答卷完整性"],
        ["Bharat Rasayan / Heranba", "各62.0%", "印度→越南技术级记录", "越南后续流向、批号BN:187、生产工序"],
        ["其他印度公司", "166.2%", "Hemani、Shogun、其他技术级供应商", "是否属于列名企业/关联实体；避免错用较低税率"],
    ],
    [2650, 1200, 2500, 3010],
)
add_body(doc, "反倾销税=海关确定的计税价格×适用税率。进口增值税的计税基础包含关税和反倾销税，因此未征反倾销税还可能连带造成进口增值税差额。国家税务总局公开口径显示销售或进口农药适用9%税率；但实际仍应核对中国进口税单的商品用途、税率栏和申报编号，不能脱离税单机械写死税率。")
add_callout(doc, "税款表达边界", "本报告只把反倾销税及其引致的进口增值税差额列为条件风险，不把正常关税或基础增值税写成‘逃税额’。出口侧金额仅用于近似情景测算；认定应纳/少缴税额必须取得中国海关计税价格与税款缴款书。", fill=CALLOUT)

add_heading(doc, "三、易迅全量查询、去重和逐票分析", 1)
add_body(doc, "纳入5个已下载两年工作簿：CAS 1315501-18-8（4条）、CAS 52315-07-8（799条）、CAS 67375-30-8（203条）、CIPERMETHRIN（4条）、CYPERMETHRIN TECHNICAL（135条），原始合计1,145条。按数据源、流向、日期、描述、买卖双方、重量/数量/金额、目的国、原产国和HS精确去重后为1,029条；再标准化标点/空格得到1,024条，按同日、双方、重量/数量/金额与路线做保守商业票去重后为1,009条。")
add_caption(doc, "表3  逐票范围判定（格式镜像去重口径）")
add_table(
    doc,
    ["分类", "记录数", "折算重量kg", "处理原则"],
    [
        ["技术级/原药（措施范围候选）", "477", "3,993,734.125", "技术级、高纯、TC、工业包装等信号；进入路线核查"],
        ["CAS明确工业批次（措施范围候选）", "54", "705,738.560", "CAS列明且工业批量；仍需纯度、用途和形态"],
        ["制剂/复配", "339", "4,000,585.070", "原则上不按原药口径；核成分含量与中国申报编号"],
        ["实验室标准品/小样", "122", "0.300", "mg级或实验室用途；数量件数不按kg汇总"],
        ["商品形态待核", "32", "290.560", "信息不足或REF STD/低量边界；不直接计入应税量"],
    ],
    [2850, 1200, 1800, 3510],
)
add_body(doc, "清洗过程中专门纠正了‘STANDARD’误判：商业提单常见ISPM 15 STANDARD、TSCA STATEMENT，并不代表实验室标准品。修正后实验室类折算重量仅0.3kg，避免把约4.87万kg商业技术级货物错误排除。另将描述中明确“1 UNK=5 GRAM=0.005 KGM”的越南小票按0.005kg计，不把数量1误当1kg。")

add_heading(doc, "四、直接对华风险与税款情景", 1)
add_heading(doc, "4.1 下载表中最具体一票：Tagros 16吨", 2)
add_caption(doc, "表4  2026年2月28日印度出口侧记录")
add_table(
    doc,
    ["字段", "记录内容", "风险含义"],
    [
        ["商品", "CYPERMETHRIN TECHNICAL；CAS 52315-07-8；92%；40:60", "名称、CAS、纯度均高度匹配措施范围"],
        ["主体", "Tagros Chemicals India Private Limited→IPO LTD SHNGHAI", "Tagros为列名企业；需识别采购商完整中文注册主体和实际进口人"],
        ["路线/日期", "India→China；2026-02-28", "措施生效后直接对华"],
        ["数量/金额", "16,000kg；USD95,200；约USD5.95/kg", "可作出口侧情景测算，不能替代海关完税价格"],
        ["税号", "印度出口HS38089135", "中国执行编号2926909013；各国细分税号不同不能据此排除措施范围"],
        ["缺口", "无中国进口报关单、原产地证、税款缴款书", "A类出口侧证据；尚不能认定未缴税"],
    ],
    [1850, 3900, 3610],
)
add_caption(doc, "表5  Tagros 16吨条件性税款差额")
add_table(
    doc,
    ["税种/情景", "计算", "金额USD", "解释"],
    [
        ["反倾销税", "95,200×48.4%", "46,076.80", "若进入中国且未按列名税率征收"],
        ["进口增值税差额（9%）", "46,076.80×9%", "4,146.91", "农药政策情景；仅测AD引致的增量"],
        ["合计（9%）", "46,076.80+4,146.91", "50,223.71", "主要报告口径"],
        ["13%敏感性情景", "46,076.80×13%+46,076.80", "52,066.78", "只有中国税单实际适用13%时才使用"],
    ],
    [2350, 2500, 1600, 2910],
)

add_heading(doc, "4.2 保存页面19条：重复展示与保守批次", 2)
add_body(doc, "2026年8月8日保存的CYPERMETHRIN TECHNICAL→China页面共有19条。相同日期附近、同供应商、同数量和金额的重复记录不能直接相加；按页面重复组保守合并为8个商业批次，其中普通技术级3批、68,000kg、USD392,800；Alpha技术级1批、9,000kg、USD83,250；其余4批、72,000kg为Beta-Cypermethrin或CAS65731-84-2边界，未在终裁列明3个CAS中，不能直接纳入。")
add_caption(doc, "表6  页面保守批次税款情景")
add_table(
    doc,
    ["供应商/日期", "数量/金额", "范围判断", "AD税", "连带VAT9%差额", "合计USD"],
    [
        ["Tagros / 2026-02-28", "16,000kg / USD95,200", "普通技术级，明确CAS", "46,076.80", "4,146.91", "50,223.71"],
        ["Meghmani / 2026-02-21", "36,000kg / USD201,600", "普通技术级；买方未显示", "124,992.00", "11,249.28", "136,241.28"],
        ["Tagros / 2026-01-06", "16,000kg / USD96,000", "普通技术级；买方未显示", "46,464.00", "4,181.76", "50,645.76"],
        ["普通技术级合计", "68,000kg / USD392,800", "3个保守批次", "217,532.80", "19,577.95", "237,110.75"],
        ["Meghmani Alpha / 2026-02-09", "9,000kg / USD83,250", "终裁列明Alpha CAS，但该票未显示CAS", "51,615.00", "4,645.35", "56,260.35"],
    ],
    [1700, 1600, 2200, 1200, 1500, 1160],
    font_size=8.8,
)
add_callout(doc, "关于此前36吨/141,240.96美元口径的纠偏", "Meghmani 36吨来自保存页面，不在5个下载工作簿的直接对华记录中。USD141,240.96是13%增值税敏感性情景；按农药进口9%口径，反倾销税USD124,992、连带增值税差额USD11,249.28、合计USD136,241.28。两者都只是未缴税假设下的测算，不是已认定欠税。", fill=PALE_GREEN, title_color="176B35")

add_heading(doc, "五、第三国绕道核查：A腿充分，B腿缺失", 1)
add_body(doc, "格式镜像和保守商业票去重后，措施范围候选中印度原产流向第三国共468条、约3,764,043.495kg；主要目的地包括巴西、美国、墨西哥、俄罗斯、法国、印度尼西亚、越南、哥伦比亚、泰国等。这证明印度生产商具备向多个市场供应的能力，也给第三国再出口提供结构条件，但A腿规模不能替代第三国→中国B腿。")
add_caption(doc, "表7  亚洲潜在节点与本次证据结论")
add_table(
    doc,
    ["节点", "本次记录/主体", "证据等级", "是否证明绕道中国"],
    [
        ["越南", "进口侧措施后23票/162,325.005kg；UPL Vietnam、Fumakilla、Kien Nam等", "A腿已证实", "否；无越南→中国技术级B腿"],
        ["毛里求斯商业链", "UPL Mauritius→UPL Vietnam 18吨，原产India；官方确认其与UPL India/Shanghai关联", "B+商业/开票链", "否；没有物理经过毛里求斯或换产地证据"],
        ["新加坡商业链", "Sundat (S) Pte→Sundat Vietnam 5.25吨，原产India", "B商业链", "否；仍明确印度原产"],
        ["印度尼西亚/马来西亚/泰国", "印度技术级A腿量较大；集团网络可达", "B结构风险", "否；未匹配中国B腿"],
        ["阿联酋/孟加拉", "Tagros/UPL等存在实体或市场网络", "C至B结构风险", "否；公开源无双段提单/处罚"],
    ],
    [2000, 3900, 1700, 1760],
)

add_heading(doc, "5.1 越南进口侧23票", 2)
add_body(doc, "措施后越南进口侧印度原产技术级产品经同日、同买卖双方、同数量和金额去重后为23票、162,325.005kg、VND29,053,847,387.5。金额字段为越南盾，不能与印度出口数据中的美元直接合计。2026年5月29日Meghmani→UPL Vietnam两条18吨、金额完全相同，仅越文描述使用“制剂/原料”近义措辞，按同一商业票合并，不能计36吨。")
add_caption(doc, "表8  越南进口侧主要供应链")
add_table(
    doc,
    ["境外供应商", "越南买方", "票数/数量", "具体判断"],
    [
        ["Meghmani Organics", "UPL Vietnam", "4票/72,000kg", "列名印度生产商；原产India；是A腿，不是越南原产B腿"],
        ["Heranba Industries", "UPL Vietnam等", "2票/36,500kg", "列名企业；含BN:187可用于后续批号匹配"],
        ["UPL Mauritius Limited", "UPL Vietnam", "1票/18,000kg", "第三国关联贸易商；仍申报印度原产"],
        ["Tagros", "Kien Nam", "4票/15,000kg", "列名企业；普通/Alpha技术级"],
        ["Shogun Organics", "Fumakilla Vietnam", "9票/9,575kg", "批号序列完整，适合反查中国进口"],
        ["Sundat (S) Pte Ltd", "Sundat Crop Science", "1票/5,250kg", "新加坡商业主体；原产India"],
        ["Hemani Industries", "Kien Nam", "1票/6,000kg", "印度原产技术级"],
        ["UPL Limited", "S.C. Johnson Vietnam", "1票/0.005kg", "描述明确5克；按0.005kg，不按数量1计1kg"],
    ],
    [2400, 2300, 1500, 3160],
    font_size=9,
)

add_heading(doc, "5.2 UPL India—UPL Mauritius—UPL Shanghai", 2)
add_body(doc, "商务部终裁明确，UPL Limited与其关联公司毛里求斯联合磷化物有限公司、联磷磷品（上海）有限公司联合答卷，毛里求斯实体为贸易商。这足以建立印度生产商/出口商—毛里求斯关联贸易商—上海关联进口商的商业与开票链。UPL答卷还因关联方材料不完整、交易证明缺失及电子/纸质答卷不对应等问题被采用可得最佳信息，终裁税率166.2%，构成高优先合规核查理由。")
add_body(doc, "易迅记录进一步显示UPL Mauritius在2025年7月28日向UPL Vietnam销售18,000kg、min92%技术级产品，平台原产国明确India；2025年1月23日另有Alpha-Cypermethrin 97% 600kg；2024年对UPL Colombia还有两票各16吨、200kg桶。它们证明毛里求斯主体长期承担集团跨境销售功能，但不证明货物物理经过毛里求斯，更不证明原产地改报。")
add_callout(doc, "证据分级", "A级实体/商业关系：UPL India—UPL Mauritius—UPL Shanghai（官方确认关联）；B+商业流：UPL Mauritius→UPL Vietnam且原产India；C级/未证实：印度货物经毛里求斯或越南物理转运、再伪报非印度原产进入中国。", fill=CALLOUT)

add_heading(doc, "六、公开互联网深检：没有找到逃税闭环", 1)
add_body(doc, "公开一手来源核实了政策范围、税率、企业关系和多国网络。UPL年报列有毛里求斯、新加坡、印度尼西亚、马来西亚、菲律宾等作物保护实体；Tagros公开资料显示多个印度制造基地和迪拜、孟加拉等境外网络；Meghmani、Heranba均面向多国出口。Gharda产品资料提供93—94%纯度、25kg/225kg UN钢桶和UN3352等包装/危险品指纹。")
add_body(doc, "但公开来源中未找到中国海关查缉、法院判决、行政处罚、官方反规避调查或可闭合的双段提单，证明印度产氯氰菊酯经越南、毛里求斯、新加坡、阿联酋等地换单/换原产地后进入中国。商务部终裁没有认定第三国规避；对价格承诺的讨论也不能改写为已经发生规避。")
add_caption(doc, "表9  公开源实体线索与使用边界")
add_table(
    doc,
    ["实体", "公开事实", "可核查字段", "不能据此推定"],
    [
        ["Gharda", "CAS52315-07-8；93—94%；25/225kg桶；UN3352", "包装、UN号、批号、Saykha产能", "出现相同包装即为Gharda或即为绕道"],
        ["Tagros", "印度多基地，出口90多个国家，境外网络", "生产基地、第三国贸易商、40:60/纯度", "全球销售即规避中国措施"],
        ["Meghmani", "Cypermethrin technical；CAS52315-07-8；93%", "生产批号、UPL Vietnam订单、中国进口人", "72吨对越南一定转入中国"],
        ["Heranba", "印度Vapi/Sarigam/Saykha；出口65+国家", "BN:187、生产日期、越南买方", "跨国销售等于原产地伪报"],
        ["UPL集团", "官方关联India/Mauritius/Shanghai；多国子公司", "合同、发票、付款、原产地证、提单", "商业发票经毛里求斯即货物经毛里求斯"],
    ],
    [1700, 3050, 2650, 1960],
    font_size=9,
)

add_heading(doc, "七、重点核查实体、批号和口岸数据需求", 1)
add_caption(doc, "表10  建议按证据增益排序调取")
add_table(
    doc,
    ["优先级", "对象", "所需资料/查询条件", "目标"],
    [
        ["P1", "2026-02-28 Tagros 16吨", "IPO LTD SHNGHAI及相近拼写；CAS52315-07-8；92%；40:60；16,000kg；USD95,200；日期前后30天；中国进口报关单、税款缴款书", "确认实际进口人、口岸、2926909013、印度原产、48.4%及9%/实际VAT"],
        ["P1", "页面Meghmani 36吨", "日期2026-02-20/21；36,000kg；USD201,600；生产商Meghmani；中国买方/报关行", "去重页面记录并核62%税率；纠正13%/9%口径"],
        ["P1", "UPL关联链", "UPL India、UPL Mauritius、联磷磷品（上海）；合同、发票、付款、提单、原产地证、关联交易", "区分商业开票与物理路线；核166.2%"],
        ["P2", "越南/新加坡B腿反查", "UPL Vietnam、Fumakilla Vietnam、Kien Nam、Sundat；HS292690/380891；3个CAS；technical/TC/TG；目的国China", "发现是否存在措施后第三国→中国B腿"],
        ["P2", "批号指纹", "090126CPR、080126CPR、070825CPR、06&070825CPR、060825CPR、020425CPR、01&020425CPR、010425CPR、BN:187", "与中国报关描述、COA、装箱单和农药登记匹配"],
        ["P2", "口岸/代理", "中国报关单中的进口口岸、报关企业、运输方式、集装箱号、船名航次", "识别固定进口口岸和代理链；现有出口侧数据不足以指名高风险口岸"],
    ],
    [900, 1700, 4400, 2360],
    font_size=8.8,
)
add_body(doc, "当前不能据现有资料直接断言某一中国口岸风险最高：印度出口侧和越南进口侧记录没有公开对应中国报关口岸，终裁也只列工厂/仓库至出口港等费用而未公开港口名称。应先以中国进口报关单确认口岸，再按进口人—报关行—口岸—税号—原产地组合扩线。")

add_heading(doc, "八、最终研判与使用边界", 1)
add_caption(doc, "表11  风险矩阵")
add_table(
    doc,
    ["风险方向", "等级", "依据", "当前能否认定违法"],
    [
        ["措施后印度直接进口合规", "高", "Tagros 16吨及页面Tagros/Meghmani大票；产品、来源、列名企业匹配", "不能；需中国税单确认是否已正确缴税"],
        ["第三国商业中介/关联开票", "中高", "UPL Mauritius、Sundat Singapore；官方关联与易迅记录相互印证", "不能；商业中介本身合法且仍报印度原产"],
        ["印度→第三国A腿/供应链可达性", "中", "468条、约3,764吨措施范围候选流向第三国", "不能；A腿不是B腿"],
        ["第三国→中国物理转运", "低/未证实", "下载表B腿0；公开源无双段提单、处罚或反规避认定", "不能"],
        ["伪报原产地/逃避AD", "低/无法判断", "缺中国原产地申报、生产商、COA、缴税与货物流闭环", "不能"],
    ],
    [2350, 1350, 3900, 1760],
)
add_callout(doc, "最终结论", "当前最值得立刻核查的不是笼统的‘越南绕道’，而是措施后列名印度生产商直接对华大票是否使用正确中国商品编号、原产地、企业税率并实际缴纳反倾销税。第三国方向已识别UPL Mauritius与Sundat Singapore等具体商业链和23票越南A腿，但缺少任何进入中国的B腿，因此不能认定绕道、原产地伪报或逃税。", fill=PALE_RED, title_color=RISK_RED)

add_heading(doc, "附录：公开来源与工作底稿", 1)
sources = [
    ("商务部", "公告2025年第24号（终裁、范围、税率、实体与期限）", "https://www.mofcom.gov.cn/zcfb/zgdwjjmywg/art/2025/art_f4dda218753844ad8244cd84442c7992.html"),
    ("商务部/海关总署", "海关总署公告2025年第3号转发页（2926909013）", "https://www.mofcom.gov.cn/zcfb/zgdwjjmywg/art/2025/art_891c7244e10b4cbdbe1688e288a18e87.html"),
    ("国家税务总局", "销售或进口农药9%增值税口径", "https://www.chinatax.gov.cn/chinatax/n810356/n3255681/c5236314/content.html"),
    ("UPL", "2023—24年度报告（毛里求斯及亚洲关联网络）", "https://www.upl-ltd.com/financial_result_and_report_pdfs/Da7CGgxBEBSh8KKSrznTRJFTNU3ahra0yEy0O8R4/UPL_Annual-Report_2023-24.pdf"),
    ("Tagros", "公司制造基地与全球网络", "https://tagros.com/"),
    ("Meghmani", "Cypermethrin Technical产品页", "https://meghmani.com/product/cypermethrin-technical/"),
    ("Gharda", "Cypermethrin产品资料（纯度、包装、UN号）", "https://www.gharda.com/wp-content/uploads/2023/09/Cypermethrin.pdf"),
    ("Heranba", "企业与生产基地", "https://www.heranba.co.in/about-heranba/"),
]
table = add_table(doc, ["来源", "支持内容", "链接"], [[a, b, ""] for a, b, _ in sources], [1900, 5350, 2110], font_size=9)
for idx, (_, _, url) in enumerate(sources, start=1):
    add_hyperlink(table.cell(idx, 2).paragraphs[0], "打开一手来源", url)
add_body(doc, "工作底稿位于同目录：《氯氰菊酯_易迅逐票判定台账.xlsx》包含概览、1,029条逐票全量、页面对华19条、越南A腿23票、链路与税款、口径与来源；另保留5个原始下载工作簿、标准化CSV/JSON、全量分析JSON及页面摘录。")
add_body(doc, "限制：易迅是多国贸易情报记录集合，字段、币种、覆盖和更新日期不一致；目的国/原产国/数据源不等于中国海关申报事实；相同记录可能跨查询或跨数据源镜像显示。报告用于风险筛查和合规核查，不构成对企业违法、欠税或走私的认定。")

doc.core_properties.title = "氯氰菊酯反倾销税与第三国转运风险深度分析报告"
doc.core_properties.subject = "中国进口反倾销税合规、直接进口与第三国转运风险"
doc.core_properties.author = "反倾销税风险分析项目"
doc.core_properties.keywords = "氯氰菊酯, 反倾销税, 印度, 越南, 毛里求斯, 第三国转运, 易迅数据"
doc.core_properties.comments = "基于公开一手来源和易迅全量逐票数据形成；不构成违法认定。"
doc.save(OUT)
print(json.dumps({"output": str(OUT), "tables": len(doc.tables), "paragraphs": len(doc.paragraphs)}, ensure_ascii=False, indent=2))
