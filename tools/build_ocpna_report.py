from __future__ import annotations

from pathlib import Path
from datetime import date
import csv
import json
import zipfile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(r"D:\易迅数据\反倾销税深度分析报告\26_邻氯对硝基苯胺")
DOCX = OUT / "邻氯对硝基苯胺_反倾销税与第三国转运风险阶段审计报告.docx"
AS_OF = "2026-08-19"

# standard_business_brief token map
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(31, 41, 55)
MUTED = RGBColor(90, 98, 108)
RISK = RGBColor(155, 28, 28)
CAUTION = RGBColor(122, 90, 0)
LIGHT_GRAY = "F2F4F7"
CALLOUT = "FFF4E5"
SOFT_BLUE = "E8EEF5"
WHITE = "FFFFFF"


def set_run_font(run, size=11, bold=False, color=None, italic=False):
    run.font.name = "Calibri"
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    fonts.set(qn("w:eastAsia"), "等线")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    old = tc_pr.find(qn("w:tcMar"))
    if old is not None:
        tc_pr.remove(old)
    margins = OxmlElement("w:tcMar")
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{key}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    tc_pr.append(margins)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    if sum(widths_dxa) != 9360:
        raise ValueError(f"Table widths must sum to 9360 DXA, got {sum(widths_dxa)}")
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    for tag in ("w:tblW", "w:tblInd", "w:tblLayout"):
        node = tbl_pr.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            tbl_pr.append(node)
        if tag == "w:tblW":
            node.set(qn("w:w"), "9360")
            node.set(qn("w:type"), "dxa")
        elif tag == "w:tblInd":
            node.set(qn("w:w"), str(indent_dxa))
            node.set(qn("w:type"), "dxa")
        else:
            node.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for col in list(grid):
        grid.remove(col)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths_dxa[idx] / 1440)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths_dxa, body_size=8.7):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        shade_cell(cell, LIGHT_GRAY)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        set_run_font(p.add_run(str(header)), 9, True, DARK_BLUE)
    set_repeat_table_header(table.rows[0])
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            if idx == 0 and len(str(value)) < 12:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_run_font(p.add_run(str(value)), body_size, False, INK)
    set_table_geometry(table, widths_dxa)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)
    return table


def add_paragraph(doc, text, bold=False, color=None, italic=False, after=6, size=11):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.1
    set_run_font(p.add_run(text), size, bold, color or INK, italic)
    return p


def add_bullet(doc, label, detail):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    set_run_font(p.add_run(f"{label}："), 11, True, INK)
    set_run_font(p.add_run(detail), 11, False, INK)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run(text), {1: 16, 2: 13, 3: 12}[level], True, BLUE if level < 3 else DARK_BLUE)
    return p


def add_callout(doc, label, text, risk=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.keep_together = True
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.line_spacing = 1.12
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), CALLOUT if risk else SOFT_BLUE)
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "6")
    left.set(qn("w:color"), "9B1C1C" if risk else "2E74B5")
    borders.append(left)
    p_pr.append(borders)
    set_run_font(p.add_run(f"{label}  "), 11, True, RISK if risk else DARK_BLUE)
    set_run_font(p.add_run(text), 11, False, INK)
    return p


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "2E74B5")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    fonts = OxmlElement("w:rFonts")
    for key in ("ascii", "hAnsi"):
        fonts.set(qn(f"w:{key}"), "Calibri")
    fonts.set(qn("w:eastAsia"), "等线")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "20")
    r_pr.extend([fonts, color, underline, size])
    run.append(r_pr)
    txt = OxmlElement("w:t")
    txt.text = text
    run.append(txt)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def write_csv(path, headers, rows):
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(headers)
        writer.writerows(rows)


def build_structured_outputs():
    OUT.mkdir(parents=True, exist_ok=True)

    query_rows = [
        ["OCPNA-Q1", "中国进口主查询", "目的国China；近3年；HS292142/29214200；每页200条", "121-87-9；2-CHLORO-4-NITROANILINE；O-CHLORO-P-NITROANILINE；OCPNA", "逐页至末页；排除89-63-4/PCONA及其他异构体；保留总数、页数、末页"],
        ["OCPNA-Q2", "印度A腿", "起运/原产India；目的国全部；近3年；同HS/CAS/精确品名", "重点Indonesia、Singapore、UAE、Malaysia、Vietnam、Thailand、UK；实体Aarti/Hemani/Premier", "保存日期、批号、包装、重量、收货人、提单/柜号；不得与中国B腿机械相加"],
        ["OCPNA-Q3", "第三国B腿", "上述第三国→China；同CAS/精确品名；近3年", "按7/15/30/60/90日匹配；同批号/包装/净重/主体优先", "必须核中国报关原产国、生产商、AD税率、税款缴款书；只有平台字段不够"],
        ["OCPNA-Q4", "实体穿透", "Aarti Industries；Aarti Chemical Trading FZCO；Alchemie Europe；Alchemie Shanghai；Hawk Petroleum；Sincol", "产品关键词与实体组合检索，不限HS后再人工确认CAS", "识别生产商、贸易商、货代；不得把销售公司所在地当原产地"],
    ]
    write_csv(
        OUT / "邻氯对硝基苯胺_易迅最简查询组合.csv",
        ["查询ID", "用途", "基础条件", "关键词/实体", "完整性与判定要求"],
        query_rows,
    )

    gaps = [
        ["Y-01", "网页全页读取完成", "Codex内置浏览器已读取两个精确全称查询的全部4个结果页；62条页面记录、51条可见字段精确去重", "仅覆盖2025-08-14至2026-08-14的两个精确全称结果池，不代表易迅全库或全部同义词"],
        ["Y-02", "本地无原始下载", "D盘未发现本商品专门原始下载；本轮为登录页面只读摘录并形成结构化审计表", "数量、金额单位及币种须回查页面列头或原始申报，不可机械写成kg/USD"],
        ["Y-03", "A/B腿未闭合", "易迅确认印度→日本11,000、印度→巴西500两个数量字段A腿，并见泰国→印度453.5反向流；未见OCPNA中国B腿", "缺中国进口记录、批号、提单、柜号、第三国进口申报及中国进口报关"],
        ["Y-04", "税款未核", "没有中国海关完税价格、生产商/出口商适用税率和缴款书", "税额只允许按V做条件系数测算，不得形成实际欠税结论"],
        ["Y-05", "执法负面检索", "截至2026-08-19定向检索商务部、海关、法院公开库", "未找到本品专项反规避裁定、海关处罚、刑事判决或公开双段提单；负面结果不等于不存在"],
    ]
    write_csv(
        OUT / "邻氯对硝基苯胺_数据覆盖与缺口.csv",
        ["编号", "状态", "已覆盖", "结论边界"],
        gaps,
    )

    clues = [
        ["2025-12-18", "印度→巴西", "OCPNA；CAS 121-87-9；数量字段500；金额字段3,500", "Aarti Industries Ltd→Spice Indústria Química Ltda", "B：易迅确认A腿；目的地不是中国，无B腿闭环"],
        ["2025-09-18", "印度→日本", "OCPNA；CAS 121-87-9；数量字段11,000；金额字段67,100", "Aarti Industries Limited→Sanyo Life Material Co Limited", "B：易迅确认A腿；目的地不是中国，无B腿闭环"],
        ["2025-10-18", "泰国→印度", "OCPNA；数量字段453.5；金额字段675.71", "Thai Ambica Chemicals Co. Ltd→Colorband Dyestuff Private Limited", "反向流：不是印度→第三国，更不是第三国→中国"],
        ["2024-11-04", "印度→中国", "2-CHLORO-4-NITROANILINE OCPNA POWDER；250g样品", "Sincol Corporation Limited；上海浦东国际机场", "C：印度直达样品，不是第三国绕道；应核是否科研样品及正常缴税"],
        ["2024-10-22", "印度→印尼", "2 CHLORO 4 NITRO ANILINE；HS29214200；数量字段1,200 KGM", "公开贸易聚合页未展示完整双方", "B+：具体A腿；未见印尼→中国B腿，不能闭环"],
        ["2024-11-07", "马来西亚→巴基斯坦", "2-CLORO 4-NITRO ANILINE；200kg", "公开贸易聚合页", "C：说明马来西亚存在贸易流，不指向中国，也不证明印度原产"],
        ["2024-07-09", "德国→越南", "CAS121-87-9实验分析品；500mg", "实验室小包装", "反证：实验标准品/试剂流与工业散货不同，不得并入绕道数量"],
    ]
    write_csv(
        OUT / "邻氯对硝基苯胺_公开贸易线索.csv",
        ["日期", "路线", "产品/数量字段", "实体/口岸", "证据评价"],
        clues,
    )

    entities = [
        ["Aarti Industries Limited", "印度生产商/出口商", "官方单列AD 31.4%；官方确认既出口自产货，也采购其他印度生产商货物后对华出口；并向未披露第三国出口", "最高：调生产厂、COA plant code、批号、自产/外购台账及单列税率资格"],
        ["Hemani Global", "印度生产商线索", "官网列OCPNA CAS121-87-9、纯度99.5%", "核是否对华/对第三国出口；通常应先按其他印度公司49.9%情景核"],
        ["Premier Synthochem Industries", "印度生产商线索", "官网列OCPNA CAS121-87-9、min 99%", "核出口记录、工厂和批号"],
        ["Yashashvi Rasayan Pvt. Ltd.", "印度生产能力线索", "印度环境资料载OCPNA生产工艺", "核实际投产、产能和出口"],
        ["Aarti Chemical Trading - FZCO", "迪拜贸易节点", "Aarti官方列示的海外商业实体，未见OCPNA制造证据", "高：若为发票/发货人，穿透实际印度生产商；不能以迪拜公司推定阿联酋原产"],
        ["Alchemie Europe Ltd.", "英国贸易节点", "Aarti官方列示子公司", "核开票/付款/仓储，不等于英国生产"],
        ["Aarti Chem Trade USA Inc.", "美国贸易节点", "Aarti官方列示间接子公司", "核销售链，不等于美国生产"],
        ["Alchemie Shanghai", "中国商业节点", "官网列OCPNA，并称为Aarti授权经销商", "高：核中国进口报关、上游发票、生产商和税率"],
        ["杭州可菲克化学有限公司", "中国进口商", "2023期终复审登记并答卷", "高：调复审期及措施后报关/税款/供应链"],
        ["恒诚制药集团淮南有限公司", "中国进口商", "2023期终复审登记并答卷", "高：调用途、报关、税款和供应商"],
        ["江阴市江联工贸有限公司", "中国进口商", "登记参加复审但未提交进口商答卷", "中高：核进口历史及申报数据"],
        ["Sincol Corporation Limited", "中国收货人线索", "公开聚合页出现2024-11-04印度直达250g样品", "中：核法律实体、用途、快件/空运申报和税款"],
        ["Hawk Petroleum Pte Ltd", "新加坡贸易节点", "产品目录列OC4NA/OCPNA，未见制造证明", "中：若出现新加坡B腿，核实际生产商及库存/换单"],
        ["苏州市罗森助剂有限公司", "中国国内生产者", "原审与复审申请人", "产业/政策信息源；不是第三国绕道主体推定对象"],
    ]
    write_csv(
        OUT / "邻氯对硝基苯胺_实体核查清单.csv",
        ["实体", "角色", "公开事实", "核查重点"],
        entities,
    )

    sources = [
        ["商务部2024年第5号期终复审终裁", "现行AD范围、税率、期限、进口基线、Aarti出口及产能事实", "https://dcj.mofcom.gov.cn/article/zcfb/zcwg/202404/20240403505017.shtml", "官方"],
        ["商务部2018年第19号反倾销终裁", "原审AD税率与范围", "https://www.mofcom.gov.cn/zcfb/zgdwjjmywg/art/2018/art_ab51892b2978401b98e5317829e17d44.html", "官方"],
        ["商务部2018年第18号反补贴终裁", "历史CVD 21.2%/166.8%及计税公式", "https://www.mofcom.gov.cn/zcfb/blgg/art/2018/art_5574080e710e4821814e67dcb8ddaeb6.html", "官方"],
        ["商务部2023年第3号期终复审立案", "AD继续实施；CVD未启动复审的时间边界", "https://trb.mofcom.gov.cn/myjjdc/art/2023/art_e899d70dba2842d38f126c9012eba865.html", "官方"],
        ["Aarti OCPNA产品页", "CAS、产品、出口和包装", "https://www.aarti-industries.com/products/ortho-chloro-para-nitro-aniline", "企业一手"],
        ["Aarti全球联系页", "印度工厂与迪拜/英国/美国商业实体", "https://www.aarti-industries.com/contact", "企业一手"],
        ["Aarti 2024-25年报", "海外子公司及出口规模（非产品专属）", "https://www.aarti-industries.com/Upload/PDF/Integrated-Annual-Report-2024-25.pdf", "企业一手"],
        ["Aarti环境合规文件", "Vapi胺类产品组产能；非OCPNA单项产能", "https://www.aarti-industries.com/Upload/PDF/ec-compliance-april-23-to-sep-23-aarti-industries-ltd-amine-division-.pdf", "企业/监管文件"],
        ["印度商务部CHEMEXCIL报告", "印度产能、内需和出口背景；税种描述不作为中国现行税依据", "https://commerce.gov.in/wp-content/uploads/2020/11/MOC_637050100118245496_CHEMEXCIL.pdf", "政府资料"],
        ["Hemani Global OCPNA页", "印度生产商与纯度线索", "https://hemaniglobal.com/Ortho_Chloro_Para_Nitro_Aniline.html", "企业一手"],
        ["Premier Synthochem", "印度OCPNA产品线索", "https://www.premiersynthochem.com/", "企业一手"],
        ["Alchemie Shanghai产品页", "Aarti授权经销关系及中国商业节点", "https://www.alcshanghai.com/en/product.php?s_id=2", "企业一手"],
        ["PubChem CAS 121-87-9", "名称、分子式和同义词", "https://pubchem.ncbi.nlm.nih.gov/compound/121-87-9", "政府数据库"],
        ["Trademo Sincol公开页", "2024-11-04印度直达中国250g样品线索", "https://www.trademo.com/companies/sincol-corporation-limited/16850578", "二级贸易聚合"],
        ["NBD Sincol公开页", "同票印度原产与HS29214290交叉确认", "https://en.nbd.ltd/trader/info/NBDD3Y526468232", "二级贸易聚合"],
        ["Volza 2-nitro-aniline进口页", "印度→印尼1.2t、马来西亚→巴基斯坦200kg线索", "https://www.volza.com/p/2-nitro-aniline/import/", "二级贸易聚合"],
        ["Hawk Petroleum产品目录", "新加坡销售节点线索", "https://hawkglobal.com/wp-content/uploads/2024/05/DYES-AND-PIGMENTS-INTERMEDIATES.pdf", "企业目录"],
    ]
    write_csv(
        OUT / "邻氯对硝基苯胺_公开来源台账.csv",
        ["来源", "支持事项", "网址", "层级"],
        sources,
    )

    query_audit = [
        ["YX-OCPNA-01", "OCPNA", "商品描述", "2025-08-14至2026-08-14；全球", 0, 0, "无结果", "缩写无命中；不能单独据此判全库为0"],
        ["YX-OCPNA-02", "121-87-9", "商品描述", "2025-08-14至2026-08-14；全球", 688, 35, "未逐票纳入", "平台将连字符CAS拆词，结果以无关记录为主"],
        ["YX-OCPNA-03", "121-87-9 + HS292142", "商品描述+HS", "2025-08-14至2026-08-14；全球", 7124, 357, "未逐票纳入", "组合字段呈宽口径/并集特征，结果反而扩大，不是严格AND"],
        ["YX-OCPNA-04", "2 CHLORO 4 NITRO ANILINE", "商品描述", "2025-08-14至2026-08-14；全球", 26, 2, "全部读取", "1条OCPNA、10条PCONA、15条噪声"],
        ["YX-OCPNA-05", "ORTHO CHLORO PARA NITRO ANILINE", "商品描述", "2025-08-14至2026-08-14；全球", 36, 2, "全部读取", "3条OCPNA、33条PCONA；其中1条OCPNA与上一查询重合"],
    ]
    write_csv(
        OUT / "邻氯对硝基苯胺_易迅查询汇总.csv",
        ["query_id", "关键词", "字段", "范围", "结果条数", "页数", "审计状态", "结果说明"],
        query_audit,
    )

    first_noise = {1, 2, 3, 4, 5, 6, 7, 9, 14, 15, 16, 18, 19, 23, 24}
    first_pcona = {8, 10, 11, 12, 13, 17, 21, 22, 25, 26}
    second_ocpna = {17, 25, 28}
    shipment_audit = []
    for row_no in range(1, 27):
        if row_no == 20:
            cls, basis = "OCPNA范围候选", "全称命中；与别名查询第25行重合"
        elif row_no in first_pcona:
            cls, basis = "PCONA位置异构体", "货描指向para-chloro-ortho-nitroaniline/CAS 89-63-4，不属于本案"
        elif row_no in first_noise:
            cls, basis = "噪声/其他产品", "实验品或货描未能确认OCPNA"
        else:
            raise RuntimeError(f"Unclassified first-query row {row_no}")
        shipment_audit.append(["YX-OCPNA-04", row_no, cls, basis, "已逐条读取"])
    for row_no in range(1, 37):
        if row_no in second_ocpna:
            cls, basis = "OCPNA范围候选", "货描明确ORTHO CHLORO PARA NITRO ANILINE；其中第17/28行含CAS 121-87-9"
        else:
            cls, basis = "PCONA位置异构体", "货描指向PARA CHLORO ORTHO NITRO ANILINE/PCONA，不属于本案"
        shipment_audit.append(["YX-OCPNA-05", row_no, cls, basis, "已逐条读取"])
    write_csv(
        OUT / "邻氯对硝基苯胺_易迅近一年逐票审计摘要.csv",
        ["query_id", "query_row", "scope_class", "classification_basis", "review_status"],
        shipment_audit,
    )

    yixun_targets = [
        ["YX-OCPNA-05-17", "2025-12-18", "印度|出口", "29214212", "ORTHO CHLORO PARA NITRO ANILINE; INVOICE 5252604945; INVOICE DATE 2025-10-06; CAS 121-87-9", "SPICE INDÚSTRIA QUÍMICA LTDA", "AARTI INDUSTRIES LTD", "", "500.00", "3500.00", "Brazil", "India", "OCPNA范围候选；印度→巴西A腿"],
        ["YX-OCPNA-04-20 / YX-OCPNA-05-25", "2025-10-18", "印度|进口", "29214290", "DYES INTERMEDIATES ORTHO CHLORO PARA NITRO ANILINE 2 CHLORO 4 NITRO ANILINE", "COLORBAND DYESTUFF PRIVATE LIMITED", "THAI AMBICA CHEMICALS CO. LTD.", "", "453.50", "675.71", "India", "Thailand", "OCPNA范围候选；两查询重合；泰国→印度反向流"],
        ["YX-OCPNA-05-28", "2025-09-18", "印度|出口", "29214212", "ORTHO CHLORO PARA NITRO ANILINE; INVOICE 5252604269; CAS 121-87-9", "SANYO LIFE MATERIAL CO LIMITED", "AARTI INDUSTRIES LIMITED", "", "11000.00", "67100.00", "Japan", "India", "OCPNA范围候选；印度→日本A腿"],
    ]
    write_csv(
        OUT / "邻氯对硝基苯胺_易迅OCPNA范围候选3条.csv",
        ["record_id", "日期", "数据源/方向", "HS", "商品描述", "买方", "卖方", "重量字段", "数量字段", "金额字段", "目的国", "平台原产地", "审计结论"],
        yixun_targets,
    )

    status_row = [["已完成精确全称全页读取", AS_OF, "页面与登录正常；两个精确全称查询共62条页面记录，51条精确去重，3条OCPNA范围候选", "没有OCPNA中国B腿；不能外推为易迅全库无中国记录", "调中国进口报关及按实体/批号继续穿透"]]
    write_csv(
        OUT / "邻氯对硝基苯胺_易迅网页取数状态.csv",
        ["status", "as_of", "browser_result", "conclusion_boundary", "next_step"],
        status_row,
    )
    # Keep the legacy filename but turn it into an explicit deprecation notice.
    write_csv(
        OUT / "邻氯对硝基苯胺_易迅逐票标准化_当前0条.csv",
        ["legacy_status", "note"],
        [["勿作为零结果引用", "本文件为2026-08-13历史占位；2026-08-19已完成两个精确全称结果池全页读取。请以易迅查询汇总、近一年逐票审计摘要及OCPNA范围候选3条为准。"]],
    )

    summary = {
        "item": 26,
        "product": "邻氯对硝基苯胺",
        "english": "Ortho Chloro Para Nitro Aniline / 2-Chloro-4-nitroaniline",
        "cas": "121-87-9",
        "hs_china": "29214200",
        "taxed_origin": "India",
        "current_measure": {"type": "anti-dumping", "from": "2024-02-13", "years": 5, "expected_end": "2029-02-12", "aarti": "31.4%", "all_others": "49.9%"},
        "countervailing": {"historical": "2018-02-13 to 2023-02-12", "aarti": "21.2%", "all_others": "166.8%", "current": "terminated; do not add to current exposure"},
        "current_tax_gap_coefficients": {"aarti_ad_plus_vat_increment": "35.482% * V", "other_india_ad_plus_vat_increment": "56.387% * V", "conditional_wrong_rate_difference": "20.905% * V"},
        "verified_public_facts": [
            "Aarti exported self-produced OCPNA directly to unaffiliated Chinese customers.",
            "Aarti also bought OCPNA from other Indian producers and exported it to China as a trader.",
            "Aarti exported the subject product to an unnamed third country/region.",
            "2018-2022Q3 China subject imports were all from India in MOFCOM's review record.",
        ],
        "public_trade_clues": len(clues),
        "closed_india_third_country_china_chains": 0,
        "yixun": {
            "coverage": "2025-08-14 to 2026-08-14; global; two exact full-name queries fully read",
            "raw_query_occurrences": 62,
            "exact_12_field_unique": 51,
            "raw_classification": {"OCPNA": 4, "PCONA_isomer": 43, "noise_or_other": 15},
            "unique_classification": {"OCPNA": 3, "PCONA_isomer": 33, "noise_or_other": 15},
            "ocpna_china_b_leg": 0,
            "ocpna_routes": ["India to Japan", "India to Brazil", "Thailand to India"],
            "bridge_latency": "Yixun page and login worked; Codex browser bridge Statsig telemetry calls repeatedly waited about 10 seconds, so actions had to be split into separate calls",
        },
        "as_of": AS_OF,
    }
    (OUT / "邻氯对硝基苯胺_阶段审计摘要.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    return query_rows, gaps, clues, entities, sources, summary


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    heading_specs = {
        1: (16, 16, 8, BLUE),
        2: (13, 12, 6, BLUE),
        3: (12, 8, 4, DARK_BLUE),
    }
    for level, (size, before, after, color) in heading_specs.items():
        style = styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    bullet = styles["List Bullet"]
    bullet.font.name = "Calibri"
    bullet._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
    bullet.font.size = Pt(11)
    bullet.paragraph_format.left_indent = Inches(0.5)
    bullet.paragraph_format.first_line_indent = Inches(-0.25)
    bullet.paragraph_format.space_after = Pt(8)
    bullet.paragraph_format.line_spacing = 1.167

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    set_run_font(header.add_run("反倾销税深度分析  |  第26项"), 9, False, MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.space_after = Pt(0)
    set_run_font(footer.add_run("邻氯对硝基苯胺阶段审计  |  第 "), 9, False, MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)
    set_run_font(footer.add_run(f" 页  |  {AS_OF}"), 9, False, MUTED)


def add_title_block(doc):
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(12)
    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(4)
    set_run_font(title.add_run("邻氯对硝基苯胺"), 23, True, INK)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    set_run_font(subtitle.add_run("反倾销税与第三国转运风险阶段审计报告"), 14, False, MUTED)

    metadata = [
        ("商品", "OCPNA / 2-Chloro-4-nitroaniline；CAS 121-87-9"),
        ("中国税号", "29214200（同税号其他产品不当然属于措施范围）"),
        ("受税来源", "印度"),
        ("现行措施", "2024-02-13起续征5年；Aarti 31.4%，其他印度公司49.9%"),
        ("审计状态", "公开来源完成；易迅近一年两个精确全称结果池已全页读取"),
        ("审计日期", AS_OF),
    ]
    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.0
        set_run_font(p.add_run(f"{label}："), 10.5, True, INK)
        set_run_font(p.add_run(value), 10.5, False, INK)

    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(8)
    rule.paragraph_format.space_after = Pt(10)
    p_pr = rule._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "2E74B5")
    borders.append(bottom)
    p_pr.append(borders)


def build_doc(query_rows, gaps, clues, entities, sources, summary):
    doc = Document()
    configure_document(doc)
    add_title_block(doc)

    add_callout(
        doc,
        "核心结论",
        "易迅两个精确全称查询共读取62条页面记录，按12个可见字段去重为51条；仅3条为OCPNA范围候选，33条是PCONA位置异构体、15条为噪声。3条OCPNA分别为印度→日本、印度→巴西和泰国→印度，没有中国B腿，仍未形成“印度→第三国→中国”同批闭环，不能指认具体绕道企业、数量或实际逃税额。",
        risk=True,
    )

    add_heading(doc, "一、政策状态与税种边界")
    add_paragraph(doc, "商务部2024年第5号公告决定，自2024年2月13日起继续对原产于印度的邻氯对硝基苯胺征收反倾销税5年。措施范围以产品身份为核心：邻氯对硝基苯胺/OCPNA，CAS 121-87-9；税号29214200项下其他产品不当然属于措施范围。")
    policy_rows = [
        ["现行", "Aarti Industries Limited", "AD 31.4%", "AD+由AD增加的13%进口VAT：35.482%×V"],
        ["现行", "其他印度公司", "AD 49.9%", "AD+由AD增加的13%进口VAT：56.387%×V"],
        ["条件情景", "若不具单列资格却错用Aarti税率", "税率差18.5个百分点", "差额情景：(49.9%-31.4%)×1.13=20.905%×V"],
        ["历史至2023-02-12", "Aarti / 其他印度公司", "CVD 21.2% / 166.8%", "反补贴措施期满后已终止；当前票据不得继续叠加CVD"],
    ]
    add_table(doc, ["期间", "企业/情景", "税率", "筛查口径"], policy_rows, [1150, 2200, 1600, 4410], 8.8)
    add_paragraph(doc, "V为中国海关审定完税价格。当前税款风险只核反倾销税及其引致的进口增值税差额；关税、正常进口增值税和其他税费须按实际报关税则另核。历史票据如发生在反补贴措施有效期内，才另行复核CVD。", italic=True, color=MUTED, size=9.5)

    add_heading(doc, "二、产品识别：先排除异构体误判")
    add_bullet(doc, "纳入词", "Ortho Chloro Para Nitro Aniline、o-Chloro-p-nitroaniline、2-Chloro-4-nitroaniline、4-Nitro-2-chloroaniline、OCPNA、CAS 121-87-9。")
    add_bullet(doc, "高频误判", "PCONA / p-Chloro-o-nitroaniline / 4-Chloro-2-nitroaniline（CAS 89-63-4）是位置异构体，不能并入本案。公开贸易数据中有多笔PCONA大票，若不以CAS穿透会严重高估风险。")
    add_bullet(doc, "形态与单证", "涉案品通常为黄色结晶粉末。逐票至少核CAS、纯度、批号、包装袋规格、生产商、COA、原产地证和中国10位税号。实验室毫克/克级标准品应与工业散货分层。")

    add_heading(doc, "三、官方披露形成的最强核查逻辑")
    add_callout(doc, "官方事实一", "Aarti在复审期内既向中国出口自产OCPNA，也将从其他印度生产商采购的OCPNA作为贸易商出口中国。由此，报关“出口商=Aarti”不能替代对实际生产商、生产厂和单列税率适用条件的核验。")
    add_callout(doc, "官方事实二", "Aarti还向一个未公开名称的第三国/地区出口涉案产品，且该目的地销售占其第三国出口的主要部分并持续多年。这证明存在A腿和海外商业流，但没有公开中国B腿，不能据此认定绕道。")
    add_callout(doc, "官方事实三", "2018、2019、2020、2021及2022年1-9月，中国涉案产品进口量分别为902、1665、1054、458和157吨，且各期全部来自印度。措施后若突然出现同CAS第三国原产的大票，应优先调原产地、生产能力和批号。")
    add_paragraph(doc, "印度供给侧还具有显著出口压力：商务部复审披露2022年印度OCPNA产能约8316-9702吨、闲置约4465-5209吨；印度需求约2500吨，超过70%的产能可供出口。该数据是风险背景，不是具体绕道证据。", color=MUTED, size=10)

    add_heading(doc, "四、公开贸易线索逐条评价")
    clue_rows = [[row[0], row[1], row[2], row[4]] for row in clues]
    add_table(doc, ["日期", "路线", "货物/数量", "评价"], clue_rows, [1100, 1450, 3050, 3760], 8.4)
    add_callout(doc, "证据结论", "易迅进一步确认Aarti向日本和巴西出口OCPNA的两条A腿，并见一条泰国→印度反向流；公开来源另有印度→印尼1.2吨A腿和印度→中国250克直达样品。易迅精确全称结果池没有OCPNA中国B腿，因此闭合链数量仍为0。")

    add_heading(doc, "五、实体优先级与调查对象")
    top_entities = [
        ["最高", "Aarti Industries", "自产+外购后对华出口；向未名第三国出口", "生产厂/批号/外购台账/单列税率资格"],
        ["高", "Aarti FZCO / Alchemie Europe / Aarti USA", "Aarti海外商业实体，未见OCPNA制造证据", "发票、付款、仓储、换单；不可按公司所在地判原产"],
        ["高", "Alchemie Shanghai", "官网列OCPNA且称Aarti授权经销商", "中国报关、进口人、生产商、AD缴款书"],
        ["高", "杭州可菲克 / 恒诚制药淮南", "复审登记并提交进口商答卷", "措施后进口、用途、税率和上游供应链"],
        ["中高", "江阴市江联工贸", "登记参加复审但未交进口商答卷", "进口历史、报关行、口岸、税款"],
        ["中高", "Hemani / Premier / Yashashvi", "印度其他生产或产能线索", "对第三国/对华出口、工厂与批号；其他印度税率情景"],
        ["中", "Sincol Corporation", "2024-11-04印度直达250g样品收货人", "法律实体、空运/快件申报和税款"],
        ["中", "Hawk Petroleum Singapore", "目录列OCPNA，未见制造证明", "如出现新加坡B腿，穿透实际生产商"],
    ]
    add_table(doc, ["优先级", "实体", "现有事实", "核查重点"], top_entities, [900, 2100, 3100, 3260], 8.3)
    add_paragraph(doc, "完整14家实体及角色、公开依据和核查字段已另存《邻氯对硝基苯胺_实体核查清单.csv》。实体列名仅用于调单排序，不代表违法认定。", italic=True, color=MUTED, size=9.5)

    add_heading(doc, "六、第三国原产与绕道判定门槛")
    add_bullet(doc, "可升级为A级", "中国报关申报非印度原产，但COA、批号、厂商声明、印度出口申报或生产台账能证明实际印度生产；且中国端未按正确企业税率缴纳AD。")
    add_bullet(doc, "B+调单线索", "印度A腿后7-90日内，第三国对华B腿出现同CAS、同纯度、同包装、近似净重，并叠加共同收发货人、批号、提单/柜号或付款交叉。")
    add_bullet(doc, "合法替代解释", "第三国具备真实硝化/还原等生产工艺、原料投入、制造许可、能耗、产量收率、COA plant code及物料平衡；单纯仓储、分装、换标、换单不能证明第三国原产。")
    add_bullet(doc, "不能单独定性", "第三国公司开票、Aarti海外子公司、同数量、途经中转港、宏观贸易量增长、平台“原产地”字段冲突。")

    add_heading(doc, "七、易迅数据全页审计结果")
    add_callout(doc, "浏览器诊断", "易迅页面与登录会话正常。反复超时主要来自Codex浏览器桥接层的Statsig遥测请求，每次约等待10秒；将输入、查询、翻页拆成独立操作后可稳定读取。该技术问题不影响本轮已摘录的数据，但会显著拖慢连续交互。", risk=False)
    yixun_summary_rows = [
        ["OCPNA", "0条", "缩写无命中；不能外推全库为0"],
        ["121-87-9", "688条/35页", "连字符CAS被拆词，结果大多无关"],
        ["121-87-9 + HS292142", "7,124条/357页", "组合字段呈宽口径/并集特征，不是严格AND"],
        ["2 CHLORO 4 NITRO ANILINE", "26条/2页，全部读取", "1条OCPNA、10条PCONA、15条噪声"],
        ["ORTHO CHLORO PARA NITRO ANILINE", "36条/2页，全部读取", "3条OCPNA、33条PCONA；1条OCPNA跨查询重合"],
    ]
    add_table(doc, ["查询词", "结果", "审计结论"], yixun_summary_rows, [2750, 2050, 4560], 8.5)
    add_paragraph(doc, "两个精确全称查询合计62条页面记录；按12个可见字段精确去重后51条，其中OCPNA 3条、PCONA位置异构体33条、噪声15条。所有4个结果页均已读取，没有在发现异常后提前停止。", bold=True, color=DARK_BLUE)
    target_rows = [
        ["2025-12-18", "印度→巴西", "Aarti→Spice Indústria Química", "500 / 3,500", "A腿；无中国B腿"],
        ["2025-10-18", "泰国→印度", "Thai Ambica→Colorband", "453.5 / 675.71", "反向流；两查询重合"],
        ["2025-09-18", "印度→日本", "Aarti→Sanyo Life Material", "11,000 / 67,100", "A腿；无中国B腿"],
    ]
    add_table(doc, ["日期", "路线", "双方", "数量/金额字段", "评价"], target_rows, [1200, 1350, 2700, 1850, 2260], 8.2)
    add_paragraph(doc, "易迅摘录未给出可直接确认的数量单位、金额币种、提单号或柜号，报告保留“数量字段/金额字段”原称。平台原产地字段也不能替代中国法定原产地认定。D盘无本商品原始下载，因此本轮保存的是页面只读摘录及逐行分类审计，而不是平台原始导出。", italic=True, color=MUTED, size=9.5)
    add_heading(doc, "八、后续查询路线")
    query_table = [[r[0], r[1], r[2], r[4]] for r in query_rows]
    add_table(doc, ["查询ID", "用途", "条件", "全量要求"], query_table, [1100, 1500, 3260, 3500], 8.2)
    add_paragraph(doc, "下一轮不再用CAS+HS粗暴叠加，而应以精确全称、别名和重点实体分组检索，再对每条OCPNA中国B腿回查印度A腿。每次仍须读取全部页并保留总条数、页数、日期跨度和查询条件。", bold=True, color=DARK_BLUE)

    add_heading(doc, "九、税款风险应怎样写清")
    tax_rows = [
        ["Aarti正常单列资格、未缴AD", "V×31.4%", "AD×13%", "V×35.482%"],
        ["其他印度公司、未缴AD", "V×49.9%", "AD×13%", "V×56.387%"],
        ["条件：应适用49.9%却错用31.4%", "V×18.5%", "差额AD×13%", "V×20.905%"],
        ["第三国B腿但实际印度原产", "按实际生产商适用31.4%或49.9%", "AD引致VAT差额", "以中国完税价格和缴税书复算"],
    ]
    add_table(doc, ["场景", "反倾销税差", "进口VAT差额", "合计增量"], tax_rows, [3250, 1850, 1700, 2560], 8.6)
    add_paragraph(doc, "实际少缴税额必须取得中国海关完税价格、企业适用税率、已缴AD/CVD/VAT、报关日期和贸易方式。二级平台的金额字段、印度出口FOB和第三国申报金额都不能直接作为中国完税价格。", italic=True, color=MUTED, size=9.5)

    add_heading(doc, "十、最小闭环调证清单")
    docs_rows = [
        ["中国报关与税单", "10位税号、CAS/品名、原产国、启运国、境外发货人、实际生产商、进口人/消费使用单位、申报企业、口岸、净重、完税价格、AD/CVD/VAT缴款书", "确认范围、原产和税差"],
        ["物流", "主/分提单、空运单、柜号/封志、船名航次、装卸港、中转港、仓储、换单和分装记录", "A/B腿物理闭环"],
        ["产品", "COA、CAS、纯度、批号、包装唛头、SDS、生产厂plant code、原产地证、厂家声明", "区分OCPNA/PCONA并穿透生产厂"],
        ["第三国生产", "制造许可、BOM、原料采购、工单、能耗、产量收率、库存台账、人工设备、废水/环保记录", "验证是否真实实质生产"],
        ["商业与税率资格", "合同、发票、付款受益人、关联关系、Aarti外购台账、单列税率适用文件、贸易商/加工合同", "识别开票链和错用税率"],
    ]
    add_table(doc, ["材料", "必核字段", "目的"], docs_rows, [1650, 5550, 2160], 8.4)

    add_heading(doc, "十一、阶段结论")
    add_callout(doc, "风险评级：B+调单优先，未形成违法证据", "官方资料证明印度供给过剩、Aarti同时存在自产/外购对华出口和持续第三国销售；易迅又确认Aarti向日本11,000和巴西500两个数量字段的OCPNA A腿。但近一年两个精确全称结果池没有OCPNA中国B腿，也没有中国报关原产/税款证据。现阶段不得写成已逃税、不得给出实际逃税数量或税额。", risk=True)
    add_paragraph(doc, "下一步优先顺序：①以Alchemie Shanghai、复审登记进口商及Aarti海外商业实体做企业穿透；②调中国进口报关并锁定精确CAS/生产商；③如出现第三国B腿，再按批号、包装、提单/柜号和7—90日窗口回查印度A腿；④取得完税价格和缴款书后，才按V及正确企业税率计算差额。", bold=True, color=DARK_BLUE)

    add_heading(doc, "十二、主要来源")
    for source, support, url, level in sources:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.line_spacing = 1.05
        set_run_font(p.add_run(f"{source}（{level}）— {support}："), 9.5, False, INK)
        add_hyperlink(p, "打开来源", url)

    doc.core_properties.title = "邻氯对硝基苯胺反倾销税与第三国转运风险阶段审计报告"
    doc.core_properties.subject = "第26项商品：政策、易迅数据缺口、实体与第三国绕道风险"
    doc.core_properties.author = "反倾销税风险分析项目"
    doc.core_properties.keywords = "OCPNA, 121-87-9, 29214200, 反倾销, 印度, 第三国转运"
    doc.save(DOCX)

    # Structural read-back QA.
    reopened = Document(DOCX)
    if len(reopened.paragraphs) < 70 or len(reopened.tables) < 5:
        raise RuntimeError("DOCX structural QA failed")
    with zipfile.ZipFile(DOCX) as zf:
        document_xml = zf.read("word/document.xml").decode("utf-8")
        if "w:tblW" not in document_xml or "w:tblGrid" not in document_xml:
            raise RuntimeError("Table geometry missing")
    return {"docx": str(DOCX), "paragraphs": len(reopened.paragraphs), "tables": len(reopened.tables)}


def main():
    outputs = build_structured_outputs()
    result = build_doc(*outputs)
    qa = {
        "as_of": AS_OF,
        "preset": "standard_business_brief",
        "header_pattern": "memo_masthead",
        "page": "US Letter portrait; 1 inch margins; header/footer 0.492 inch",
        "table_width": "9360 DXA; indent 120 DXA",
        "structural_qa": result,
        "public_clues": 7,
        "closed_chains": 0,
        "yixun_status": "two exact full-name pools fully read: 62 raw query occurrences, 51 exact visible-field unique, 3 OCPNA candidates, no China B-leg",
    }
    (OUT / "邻氯对硝基苯胺_构建与结构QA.json").write_text(json.dumps(qa, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(qa, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
