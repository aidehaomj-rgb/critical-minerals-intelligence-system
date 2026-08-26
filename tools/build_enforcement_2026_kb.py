"""Build a local, evidence-linked index for the 2026 enforcement-case archive."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import sqlite3
import zipfile
from collections import Counter
from pathlib import Path, PurePosixPath

from openpyxl import load_workbook
from pypdf import PdfReader


CASE_RE = re.compile(r"(GDRM26-\d{3})", re.I)
DATE_RE = re.compile(r"20\d{2}[年./-]\s*\d{1,2}[月./-]\s*\d{0,2}[日号]?")
MODE_MAP = {"空": "空运", "陸路": "陆路", "陆路": "陆路", "郵包": "邮包", "郵寄": "邮寄", "快件": "快件"}
TYPE_TERMS = ("查詢", "查询", "情報", "情报", "回覆", "回复", "走私", "兩用物項", "两用物项")


def clean_text(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def safe_member_path(base: Path, member: str) -> Path:
    parts = [p for p in PurePosixPath(member).parts if p not in ("", ".", "..")]
    target = base.joinpath(*parts)
    if base.resolve() not in target.resolve().parents and target.resolve() != base.resolve():
        raise ValueError(f"Unsafe archive member: {member}")
    return target


def extract_nested_archive(outer: zipfile.ZipFile, member: zipfile.ZipInfo, raw_dir: Path) -> tuple[str, Path]:
    case_id = (CASE_RE.search(member.filename).group(1).upper() if CASE_RE.search(member.filename) else Path(member.filename).stem)
    case_dir = raw_dir / case_id
    case_dir.mkdir(parents=True, exist_ok=True)
    with outer.open(member) as source, zipfile.ZipFile(source) as inner:
        for entry in inner.infolist():
            if entry.is_dir():
                continue
            path = safe_member_path(case_dir, entry.filename)
            path.parent.mkdir(parents=True, exist_ok=True)
            if not path.exists():
                with inner.open(entry) as src, path.open("wb") as dst:
                    while chunk := src.read(1024 * 1024):
                        dst.write(chunk)
    return case_id, case_dir


def pdf_pages(path: Path) -> tuple[list[tuple[int, str]], str | None]:
    try:
        reader = PdfReader(str(path))
        pages = [(i + 1, clean_text(page.extract_text() or "")) for i, page in enumerate(reader.pages)]
        return pages, None
    except Exception as exc:  # Corrupt/encrypted PDFs remain discoverable in the catalog.
        return [], f"{type(exc).__name__}: {exc}"


def xlsx_text(path: Path) -> tuple[str, str | None]:
    try:
        book = load_workbook(path, read_only=True, data_only=True)
        values: list[str] = []
        for sheet in book.worksheets:
            values.append(f"[工作表] {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                line = " | ".join(clean_text(str(cell)) for cell in row if cell is not None)
                if line:
                    values.append(line)
        return "\n".join(values), None
    except Exception as exc:
        return "", f"{type(exc).__name__}: {exc}"


def classify(name: str, text: str) -> tuple[str, str, str]:
    corpus = f"{name} {text[:8000]}"
    case_type = next((term for term in TYPE_TERMS if term in corpus), "待人工分类")
    mode = next((label for term, label in MODE_MAP.items() if term in corpus), "未识别")
    date = (DATE_RE.search(corpus).group(0) if DATE_RE.search(corpus) else "")
    return case_type, mode, date


def build(archive: Path, output: Path) -> dict:
    raw_dir = output / "raw_cases"
    raw_dir.mkdir(parents=True, exist_ok=True)
    db_path = output / "enforcement_cases_2026.sqlite"
    conn = sqlite3.connect(db_path)
    conn.executescript("""
    PRAGMA journal_mode=WAL;
    DROP TABLE IF EXISTS documents;
    DROP TABLE IF EXISTS cases;
    DROP TABLE IF EXISTS document_fts;
    CREATE TABLE cases (
      case_id TEXT PRIMARY KEY, package_name TEXT NOT NULL, case_type TEXT,
      transport_mode TEXT, date_hint TEXT, raw_path TEXT NOT NULL, document_count INTEGER NOT NULL
    );
    CREATE TABLE documents (
      id INTEGER PRIMARY KEY, case_id TEXT NOT NULL, relative_path TEXT NOT NULL,
      extension TEXT NOT NULL, page_or_sheet TEXT, extracted_text TEXT,
      text_chars INTEGER NOT NULL, extraction_error TEXT,
      sha256 TEXT NOT NULL, FOREIGN KEY(case_id) REFERENCES cases(case_id)
    );
    CREATE VIRTUAL TABLE document_fts USING fts5(case_id UNINDEXED, relative_path UNINDEXED, extracted_text);
    CREATE INDEX idx_documents_case ON documents(case_id);
    """)
    stats = Counter()
    catalog: list[dict[str, str | int]] = []
    with zipfile.ZipFile(archive) as outer:
        members = [entry for entry in outer.infolist() if not entry.is_dir() and entry.filename.lower().endswith(".zip")]
        for number, member in enumerate(members, 1):
            case_id, case_dir = extract_nested_archive(outer, member, raw_dir)
            docs = [file for file in case_dir.rglob("*") if file.is_file()]
            case_text: list[str] = []
            for file in docs:
                extension = file.suffix.lower()
                rel = file.relative_to(output).as_posix()
                digest = hashlib.sha256(file.read_bytes()).hexdigest()
                if extension == ".pdf":
                    pages, error = pdf_pages(file)
                    if not pages:
                        conn.execute("INSERT INTO documents(case_id,relative_path,extension,page_or_sheet,extracted_text,text_chars,extraction_error,sha256) VALUES(?,?,?,?,?,?,?,?)", (case_id, rel, extension, "", "", 0, error, digest))
                        stats["pdf_error"] += 1
                    for page_num, text in pages:
                        conn.execute("INSERT INTO documents(case_id,relative_path,extension,page_or_sheet,extracted_text,text_chars,extraction_error,sha256) VALUES(?,?,?,?,?,?,?,?)", (case_id, rel, extension, f"page {page_num}", text, len(text), error, digest))
                        conn.execute("INSERT INTO document_fts(case_id,relative_path,extracted_text) VALUES(?,?,?)", (case_id, rel, text))
                        case_text.append(text)
                        stats["pdf_pages"] += 1
                    stats["pdf_files"] += 1
                elif extension == ".xlsx":
                    text, error = xlsx_text(file)
                    conn.execute("INSERT INTO documents(case_id,relative_path,extension,page_or_sheet,extracted_text,text_chars,extraction_error,sha256) VALUES(?,?,?,?,?,?,?,?)", (case_id, rel, extension, "workbook", text, len(text), error, digest))
                    conn.execute("INSERT INTO document_fts(case_id,relative_path,extracted_text) VALUES(?,?,?)", (case_id, rel, text))
                    case_text.append(text)
                    stats["xlsx_files"] += 1
                else:
                    conn.execute("INSERT INTO documents(case_id,relative_path,extension,page_or_sheet,extracted_text,text_chars,extraction_error,sha256) VALUES(?,?,?,?,?,?,?,?)", (case_id, rel, extension, "", "", 0, "Unsupported file type", digest))
                    stats["other_files"] += 1
                stats[f"extension:{extension or '[none]'}"] += 1
            case_type, mode, date = classify(member.filename, "\n".join(case_text))
            conn.execute("INSERT INTO cases VALUES(?,?,?,?,?,?,?)", (case_id, member.filename, case_type, mode, date, case_dir.relative_to(output).as_posix(), len(docs)))
            catalog.append({"case_id": case_id, "package_name": member.filename, "case_type": case_type, "transport_mode": mode, "date_hint": date, "document_count": len(docs), "raw_path": case_dir.relative_to(output).as_posix()})
            stats["cases"] += 1
            if number % 25 == 0:
                conn.commit()
                print(f"Processed {number}/{len(members)} cases", flush=True)
    conn.commit()
    conn.close()
    catalog.sort(key=lambda row: row["case_id"])
    with (output / "case_catalog.csv").open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(catalog[0]) if catalog else [])
        writer.writeheader()
        writer.writerows(catalog)
    report = {"archive": str(archive), "case_count": stats["cases"], "stats": dict(sorted(stats.items())), "output": str(output), "notes": ["PDF pages without extractable text are retained and should be OCR processed before content-level search.", "All extracted text can be searched in SQLite table document_fts; original material paths remain in documents.relative_path."]}
    (output / "build_summary.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return report


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--archive", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    print(json.dumps(build(args.archive, args.output), ensure_ascii=False, indent=2))
