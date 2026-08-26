"""Extract GDRM26-395 shipment facts and search the local case history for links."""

from __future__ import annotations

import json
import re
import sqlite3
import sys
from pathlib import Path

from openpyxl import load_workbook


sys.stdout.reconfigure(encoding="utf-8", errors="replace")
ROOT = Path(r"D:\codex\XG执法\2026_案例库")
DB = ROOT / "enforcement_cases_2026.sqlite"
XLSX = ROOT / "raw_cases" / "GDRM26-395" / "GDRM26-395附件三涉案報關紀錄.xlsx"
OUTPUT = ROOT / "GDRM26-395本地历史库关联分析.json"


def clean(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


book = load_workbook(XLSX, read_only=True, data_only=True)
shipments = []
for sheet in book.worksheets:
    rows = list(sheet.iter_rows(values_only=True))
    headers = [clean(value) for value in rows[0]]
    for row_number, row in enumerate(rows[1:], 2):
        record = {
            header: clean(row[index]) if index < len(row) else ""
            for index, header in enumerate(headers)
            if header
        }
        shipments.append({"sheet": sheet.title, "row": row_number, "fields": record})

searches = {
    "G Tech公司": ["G TECH SOLUCOES", "G Tech Solucoes"],
    "珠海粤诺公司": ["ZHUHAI YUENUO", "Zhuhai Yuenuo", "珠海粤诺", "珠海粵諾"],
    "Sale Centric公司": ["SALE CENTRIC LLC", "Sale Centric LLC"],
    "Extra Fortune公司": ["EXTRA FORTUNE LIMITED", "Extra Fortune Limited"],
    "马士基香港": ["MAERSK HONG KONG LIMITED", "Maersk Hong Kong"],
    "阳明香港": ["YANG MING LINE (HONG KONG) LIMITED", "Yang Ming Line"],
    "集装箱TCNU7293159": ["TCNU7293159"],
    "集装箱BEAU4890907": ["BEAU4890907"],
    "船舶San Lorenzo Maersk": ["SAN LORENZO MAERSK", "San Lorenzo Maersk"],
    "船舶YM Harmony": ["YM HARMONY", "YM Harmony"],
    "收货邮箱": ["QN.KOO13@GMAIL.COM", "QN.KOO13"],
    "再生铜原料": ["RECYCLED COPPER RAW MATERIALS", "Recycled Copper Raw Materials"],
    "HS740400": ["740400"],
    "废弃印刷电路板": ["废弃印刷电路板", "廢棄印刷電路板", "印刷电路板", "印刷電路板"],
    "废弃电气电子设备": ["废弃电气及电子设备", "廢棄電氣及電子設備", "电子废料", "電子廢料"],
    "巴西桑托斯港": ["BRSSZ", "Santos, Brazil"],
    "宁波港": ["CNNGB"],
}

conn = sqlite3.connect(DB)
conn.row_factory = sqlite3.Row
history = {}
for label, terms in searches.items():
    rows = []
    seen = set()
    for term in terms:
        for row in conn.execute(
            """SELECT case_id,relative_path,page_or_sheet,extracted_text
               FROM documents
               WHERE case_id <> 'GDRM26-395' AND upper(extracted_text) LIKE upper(?)
               ORDER BY case_id,relative_path,page_or_sheet""",
            (f"%{term}%",),
        ):
            key = (row["case_id"], row["relative_path"], row["page_or_sheet"])
            if key in seen:
                continue
            seen.add(key)
            text = clean(row["extracted_text"])
            upper_text = text.upper()
            position = upper_text.find(term.upper())
            start = max(0, position - 160) if position >= 0 else 0
            rows.append({
                "case_id": row["case_id"],
                "file": row["relative_path"],
                "page_or_sheet": row["page_or_sheet"],
                "term": term,
                "context": text[start:start + 500],
            })
    history[label] = rows

case_incidents = [
    {
        "case_id": row["case_id"],
        "incident_label": row["incident_label"],
        "file": row["relative_path"],
        "page": row["page_or_sheet"],
        "themes": json.loads(row["themes_json"]),
        "channels": json.loads(row["channels_json"]),
        "fields": json.loads(row["fields_json"]),
    }
    for row in conn.execute(
        "SELECT * FROM intel_incidents WHERE case_id='GDRM26-395' ORDER BY id"
    )
]
case_mentions = [dict(row) for row in conn.execute(
    """SELECT entity_type,display_value,normalized_value,relative_path,page_or_sheet,context
       FROM intel_mentions WHERE case_id='GDRM26-395'
       ORDER BY entity_type,display_value,page_or_sheet"""
)]
conn.close()

declared_vs_seized = [
    {
        "incident": "案件一",
        "declared_gross_kg": 21401.0,
        "declared_net_kg": 20393.0,
        "seized_total_kg": 1962.0 + 1172.0 + 14159.5,
        "seized_share_of_net": round((1962.0 + 1172.0 + 14159.5) / 20393.0, 4),
        "net_unclassified_difference_kg": round(20393.0 - (1962.0 + 1172.0 + 14159.5), 1),
    },
    {
        "incident": "案件二",
        "declared_gross_kg": 20412.0,
        "declared_net_kg": None,
        "seized_total_kg": 7100.0,
        "seized_share_of_declared_gross": round(7100.0 / 20412.0, 4),
        "gross_unclassified_difference_kg": 20412.0 - 7100.0,
    },
]

result = {
    "case_id": "GDRM26-395",
    "shipments_from_xlsx": shipments,
    "incidents_from_pdf": case_incidents,
    "entity_mentions": case_mentions,
    "declared_vs_seized": declared_vs_seized,
    "historical_searches_excluding_case_395": history,
    "historical_hit_summary": {label: len(rows) for label, rows in history.items()},
}
OUTPUT.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(result, ensure_ascii=False, indent=2))
