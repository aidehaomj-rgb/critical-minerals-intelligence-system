from __future__ import annotations

import csv
import hashlib
import json
import os
import re
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


SOURCE = Path(os.environ.get("METACRESOL_SOURCE", r"D:\易迅数据\间甲酚2023年至2026年1月15.xlsx"))
OUT = Path(os.environ.get("METACRESOL_OUT", r"D:\易迅数据\反倾销税深度分析报告\11_间甲酚"))

HEADERS = [
    "数据源", "进出口", "日期", "HS编码", "商品描述", "采购商", "供应商",
    "重量", "数量", "金额", "目的国/地区", "原产国/地区",
]
TAXED_ORIGINS = {
    "UNITED STATES", "USA", "U.S.A.", "GERMANY", "JAPAN", "SPAIN", "FRANCE",
    "BELGIUM", "UNITED KINGDOM", "UK", "ITALY", "NETHERLANDS", "AUSTRIA",
    "IRELAND", "DENMARK", "SWEDEN", "FINLAND", "POLAND", "CZECH REPUBLIC",
    "PORTUGAL", "GREECE", "LUXEMBOURG", "HUNGARY", "ROMANIA", "BULGARIA",
    "SLOVAKIA", "SLOVENIA", "CROATIA", "ESTONIA", "LATVIA", "LITHUANIA",
    "CYPRUS", "MALTA",
}


def text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def number(value):
    raw = text(value).replace(",", "")
    if not raw:
        return None
    try:
        return float(raw)
    except ValueError:
        return None


def upper(value) -> str:
    return re.sub(r"\s+", " ", text(value).upper()).strip()


def ascii_upper(value) -> str:
    raw = unicodedata.normalize("NFKD", text(value))
    raw = "".join(ch for ch in raw if not unicodedata.combining(ch))
    return re.sub(r"\s+", " ", raw.upper()).strip()


def norm_entity(value) -> str:
    return re.sub(r"[^A-Z0-9]+", "", ascii_upper(value))


def stable_hash(values) -> str:
    blob = json.dumps([text(v) for v in values], ensure_ascii=False, separators=(",", ":"))
    return hashlib.sha1(blob.encode("utf-8")).hexdigest()[:16]


def contains(pattern: str, value: str) -> bool:
    return re.search(pattern, value, flags=re.I) is not None


DERIVATIVE = re.compile(
    r"AMYL\s*[-_]?\s*META\s*[-_]?\s*CRESOL|AMYLMETACRESOL|"
    r"CHLORO\s*[-_]?\s*(?:META|M)\s*[-_]?\s*CRESOL|CHLOROCRESOL|M\s*[-_]?\s*CHLOROCRESOL|"
    r"P\s*[-_]?\s*CHLORO\s*[-_]?\s*M\s*[-_]?\s*CRESOL|PARA\s*CHLORO\s*META\s*CRESOL|"
    r"AMINO.{0,20}(?:META|M)\s*[-_]?\s*CRESOL|(?:META|M)\s*[-_]?\s*CRESOL.{0,20}AMINO|"
    r"AMMONIUM.{0,30}(?:META|M)\s*[-_]?\s*CRESOL|(?:META|M)\s*[-_]?\s*CRESOL.{0,40}SULFON|"
    r"M\s*[-_]?\s*CRESOL\s*PURPLE|META\s*[-_]?\s*CRESOL\s*PURPLE|"
    r"(?:META|M)\s*[-_]?\s*CRESOL.{0,12}(?:IMPURITY|DERIVATIVE|ESTER|ETHER|ACETATE|PHOSPHATE)|"
    r"(?:IMPURITY|DERIVATIVE|ESTER|ETHER|ACETATE|PHOSPHATE).{0,12}(?:META|M)\s*[-_]?\s*CRESOL",
    re.I,
)
MIXED = re.compile(
    r"META\s*[-_/]?\s*(?:PARA|P)\s*[-_]?\s*CRESOL|"
    r"M\s*[-_/]?\s*P\s*[-_]?\s*CRESOL|\bMP\s*[-_]?\s*CRESOL|"
    r"META\s+PARA|META/PARA|CRESOL\s+MIXTURE|MIXED\s+CRESOL|"
    r"CRESYLIC\s+ACID|ACIDO\s+CRESILICO|CRESOLES?\s+MP\s*\d|MP90",
    re.I,
)
PARA = re.compile(
    r"PARA\s*[-_]?\s*CRESOL|PARACRESOL|CRESOL\s+PARA|"
    r"(?<![A-Z])P\s*[-_]\s*CRESOL|(?<![A-Z])P\s+CRESOL|"
    r"4\s*[-_]?\s*METHYL\s*[-_]?\s*PHENOL|106\D{0,5}44\D{0,5}5",
    re.I,
)
ORTHO = re.compile(
    r"ORTHO\s*[-_]?\s*CRESOL|ORTO\s*[-_]?\s*CRESOL|"
    r"(?<![A-Z])O\s*[-_]\s*CRESOL|(?<![A-Z])O\s+CRESOL|"
    r"2\s*[-_]?\s*METHYL\s*[-_]?\s*PHENOL|95\D{0,5}48\D{0,5}7",
    re.I,
)
EXPLICIT = re.compile(
    r"(?<![A-Z])M\s*[-_]?\s*CRESOL(?![A-Z])|(?<![A-Z])MCRESOL(?![A-Z])|"
    r"(?<![A-Z])META\s*[-_]?\s*CRESOL(?![A-Z])|(?<![A-Z])METACRESOL(?![A-Z])|"
    r"3\s*[-_]?\s*METHYL\s*[-_]?\s*PHENOL|(?<!\d)108\D{0,6}39\D{0,6}4(?!\d)|"
    r"M\s*[-_]?\s*KRE[ZS]OL|METAKRE[ZS]OL",
    re.I,
)
GENERIC = re.compile(
    r"(?<![A-Z])CRESOL(?:ES|S)?(?![A-Z])|CRESOLES|CRESOLS|KRESOL|CRESILICO|"
    r"КРЕЗОЛ|КРЕЗОЛИ|КРЕЗОЛЫ",
    re.I,
)


def classify(description: str):
    raw = upper(description)
    asc = ascii_upper(description)
    both = raw + " | " + asc
    if DERIVATIVE.search(both):
        return "衍生物/指示剂", "明确范围外", "货描命中氯代、戊基、氨基、磺酸盐、紫色指示剂或其他间甲酚衍生物指纹"
    if MIXED.search(both):
        return "间对/甲酚酸混合物", "范围待成分", "货描明确为m/p-cresol、meta/para混合物或cresylic acid，不等同纯间甲酚"
    if PARA.search(both):
        return "对甲酚", "明确范围外", "货描命中para/p-cresol、4-methylphenol或CAS 106-44-5"
    if ORTHO.search(both):
        return "邻甲酚", "明确范围外", "货描命中ortho/o-cresol、2-methylphenol或CAS 95-48-7"
    if EXPLICIT.search(both) or "МЕТАКРЕЗОЛ" in raw or "М-КРЕЗОЛ" in raw:
        detail = "小包装/样品/实验室或药用级" if re.search(r"SAMPLE|LAB|STANDARD|SYNTHESIS|USP|PH\.?\s*EUR|PARENT|\bML\b|LITER|LITRE", both, re.I) else "工业/未分级"
        return "明确间甲酚", "范围候选", f"货描命中m-/meta-cresol、3-methylphenol或CAS 108-39-4；{detail}"
    if GENERIC.search(both):
        return "泛称cresol待核", "范围待CAS", "货描只写cresol/cresols及其盐，未指明异构体；须以CAS/COA确认"
    return "同HS其他", "明确范围外/信息不足", "HS 290712宽池中未命中间甲酚或泛称甲酚识别指纹"


def csv_write(path: Path, rows: list[dict], fields: list[str]):
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def uniq(rows: list[dict]) -> list[dict]:
    seen = set()
    out = []
    for row in rows:
        if row["visible_signature"] in seen:
            continue
        seen.add(row["visible_signature"])
        out.append(row)
    return out


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(SOURCE, read_only=True, data_only=True)
    ws = wb.active
    source_headers = [text(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1))]
    if source_headers != HEADERS:
        raise RuntimeError(f"字段不一致: {source_headers}")

    raw_rows = []
    for excel_row, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        original = dict(zip(HEADERS, values))
        values_text = [text(original[h]) for h in HEADERS]
        signature = stable_hash(values_text)
        desc = text(original["商品描述"])
        product_class, scope_screen, scope_reason = classify(desc)
        origin = text(original["原产国/地区"])
        destination = text(original["目的国/地区"])
        buyer = text(original["采购商"])
        seller = text(original["供应商"])
        route = "非目标/背景"
        evidence = "排除/背景"
        route_reason = scope_reason
        if product_class == "明确间甲酚":
            if destination.upper() == "CHINA":
                if origin.upper() in TAXED_ORIGINS:
                    route = f"{origin}→中国直接核税"
                    evidence = "A-直接核税"
                    route_reason = "受税来源字段直接对华；核产品范围、生产商税档和历史缴税状态"
                else:
                    route = f"{origin or '原产字段空'}→中国B腿"
                    evidence = "B-第三国对华"
                    route_reason = "第三国对华明确间甲酚记录成立；法定原产地和A腿未闭合"
            elif origin.upper() in TAXED_ORIGINS:
                route = f"{origin}→{destination or '目的地空'}A腿背景"
                evidence = "B-受税来源外流"
                route_reason = "受税来源向第三国供应明确间甲酚；尚未与中国B腿闭合"
            else:
                route = f"{origin or '原产字段空'}→{destination or '目的地空'}背景"
                evidence = "C-同品贸易背景"
                route_reason = "明确间甲酚贸易背景，未直接指向中国反倾销风险"
        elif product_class == "泛称cresol待核" and destination.upper() == "CHINA":
            route = f"{origin or '原产字段空'}→中国泛称B腿"
            evidence = "C-范围待CAS"
            route_reason = "对华货描只写cresols；须以CAS/COA确认是否间甲酚"

        raw_rows.append({
            "record_id": f"MC-{excel_row:05d}-{signature[:10]}",
            "excel_row": excel_row,
            "visible_signature": signature,
            "data_feed": text(original["数据源"]),
            "trade_direction": text(original["进出口"]),
            "date": text(original["日期"]),
            "hs_code": text(original["HS编码"]),
            "description": desc,
            "buyer_or_consignee": buyer,
            "seller_or_shipper": seller,
            "weight_field": text(original["重量"]),
            "quantity_field": text(original["数量"]),
            "amount_field": text(original["金额"]),
            "destination": destination,
            "platform_origin": origin,
            "weight_num": number(original["重量"]),
            "quantity_num": number(original["数量"]),
            "amount_num": number(original["金额"]),
            "product_class": product_class,
            "scope_screen": scope_screen,
            "scope_reason": scope_reason,
            "route": route,
            "evidence_grade": evidence,
            "route_reason": route_reason,
            "buyer_normalized": norm_entity(buyer),
            "seller_normalized": norm_entity(seller),
        })
    wb.close()

    counts = Counter(r["visible_signature"] for r in raw_rows)
    for row in raw_rows:
        row["visible_signature_count"] = counts[row["visible_signature"]]
        row["possible_exact_duplicate"] = "是" if counts[row["visible_signature"]] > 1 else "否"

    unique_rows = uniq(raw_rows)
    explicit = [r for r in raw_rows if r["product_class"] == "明确间甲酚"]
    explicit_unique = uniq(explicit)
    china_explicit = [r for r in explicit_unique if r["destination"].upper() == "CHINA"]
    china_generic_raw = [r for r in raw_rows if r["product_class"] == "泛称cresol待核" and r["destination"].upper() == "CHINA"]
    china_generic = uniq(china_generic_raw)
    taxed_to_india = [
        r for r in explicit_unique
        if r["destination"].upper() == "INDIA" and r["platform_origin"].upper() in TAXED_ORIGINS
    ]

    category_raw = Counter(r["product_class"] for r in raw_rows)
    category_unique = Counter(r["product_class"] for r in unique_rows)
    origin_to_india = Counter(r["platform_origin"] for r in taxed_to_india)
    seller_to_india = Counter(r["seller_or_shipper"] for r in taxed_to_india)
    buyer_to_india = Counter(r["buyer_or_consignee"] for r in taxed_to_india)

    summary = {
        "source_file": str(SOURCE),
        "sheet": ws.title,
        "source_rows": len(raw_rows),
        "visible_unique_rows": len(unique_rows),
        "exact_duplicate_extra_rows": len(raw_rows) - len(unique_rows),
        "date_min": min(r["date"] for r in raw_rows if r["date"]),
        "date_max": max(r["date"] for r in raw_rows if r["date"]),
        "classification_rules": {
            "candidate": "CAS108-39-4或独立m-/meta-cresol、3-methylphenol及俄文同义词",
            "priority_exclusions": "amyl/chloro/amino/ammonium sulfonate/m-cresol purple及其他衍生物；p/o/m-p异构体分流",
            "generic": "只写cresol/cresols及其盐者待CAS/COA，不作为明确间甲酚",
        },
        "category_raw": dict(category_raw),
        "category_visible_unique": dict(category_unique),
        "explicit_metacresol_raw": len(explicit),
        "explicit_metacresol_unique": len(explicit_unique),
        "china_explicit_unique": len(china_explicit),
        "china_explicit_quantity_field_sum": sum(r["quantity_num"] or 0 for r in china_explicit),
        "china_explicit_amount_field_sum": sum(r["amount_num"] or 0 for r in china_explicit),
        "china_generic_raw": len(china_generic_raw),
        "china_generic_unique": len(china_generic),
        "taxed_origin_to_india_unique": len(taxed_to_india),
        "taxed_origin_to_india_by_origin": dict(origin_to_india),
        "taxed_origin_to_india_quantity_field_sum": sum(r["quantity_num"] or 0 for r in taxed_to_india),
        "taxed_origin_to_india_amount_field_sum": sum(r["amount_num"] or 0 for r in taxed_to_india),
        "taxed_origin_to_india_top_sellers": seller_to_india.most_common(25),
        "taxed_origin_to_india_top_buyers": buyer_to_india.most_common(25),
        "a_b_closure": {
            "china_b_sellers": sorted({r["seller_or_shipper"] for r in china_explicit}),
            "same_entity_in_taxed_to_india_buyers": sorted(
                {r["seller_or_shipper"] for r in china_explicit}
                & {r["buyer_or_consignee"] for r in taxed_to_india}
            ),
            "same_entity_in_taxed_to_india_sellers": sorted(
                {r["seller_or_shipper"] for r in china_explicit}
                & {r["seller_or_shipper"] for r in taxed_to_india}
            ),
            "same_batch_or_container": 0,
            "conclusion": "没有同收发货实体、同批号、同提单/柜号或同数量闭环",
        },
        "measure": {
            "effective_from": "2021-01-15",
            "scheduled_expiry": "2026-01-14",
            "public_legal_status": "未检出期终复审立案/续征公告，公开法律记录支持2026-01-15起终止；实施中总表存在滞后冲突",
            "rates": {"US": 1.317, "LANXESS_Germany": 0.279, "other_EU_UK": 0.495, "Japan": 0.548},
            "ad_plus_vat_increment_coefficients": {"US": 1.48821, "LANXESS_Germany": 0.31527, "other_EU_UK": 0.55935, "Japan": 0.61924},
        },
        "limitations": [
            "缺2021-01-15至2022-12-31措施前半段数据",
            "易迅数量/金额单位币种不统一，不是中国海关净重/完税价格",
            "平台原产字段不是中国法定原产地申报",
            "无中国报关单、COA/CAS、原产地证、缴款书、提单/柜号",
        ],
    }

    fields = list(raw_rows[0].keys())
    csv_write(OUT / "间甲酚_易迅逐票标准化.csv", raw_rows, fields)
    (OUT / "间甲酚_易迅逐票标准化.json").write_text(json.dumps(raw_rows, ensure_ascii=False, indent=2), encoding="utf-8")
    csv_write(OUT / "间甲酚_明确范围候选.csv", explicit_unique, fields)
    csv_write(OUT / "间甲酚_对华明确记录.csv", china_explicit, fields)
    csv_write(OUT / "间甲酚_对华泛称待核记录.csv", china_generic, fields)
    csv_write(OUT / "间甲酚_受税来源至印度记录.csv", taxed_to_india, fields)
    duplicates = [r for r in raw_rows if r["visible_signature_count"] > 1]
    csv_write(OUT / "间甲酚_完全重复审计.csv", duplicates, fields)
    (OUT / "间甲酚_全量阶段审计.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
