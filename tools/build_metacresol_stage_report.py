from __future__ import annotations

import os
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(os.environ.get("METACRESOL_OUT", r"D:\易迅数据\反倾销税深度分析报告\11_间甲酚"))
REPORT = OUT / "间甲酚_反倾销税与第三国转运风险阶段审计报告.docx"

NAVY = "17324D"
BLUE = "1F4E78"
INK = "102A43"
MUTED = "66788A"
PALE_BLUE = "EAF2F8"
PALE_GOLD = "FFF4CC"
PALE_RED = "FCE8E6"
PALE_GREEN = "E6F4EA"
BORDER = "CBD5E1"
RED = "8B1E1E"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def font(run, size=9.4, color="222222", bold=False, italic=False):
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    run.bold = bold
    run.italic = italic
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    fonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trpr.append(el)


def borders(table):
    tblpr = table._tbl.tblPr
    box = tblpr.find(qn("w:tblBorders"))
    if box is None:
        box = OxmlElement("w:tblBorders")
        tblpr.append(box)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        item = OxmlElement(f"w:{edge}")
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), "4")
        item.set(qn("w:color"), BORDER)
        box.append(item)


def cell_text(cell, value, bold=False, size=8.2, color="222222"):
    cell.text = "" if value is None else str(value)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    for p in cell.paragraphs:
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.03
        for r in p.runs:
            font(r, size=size, color=color, bold=bold)


def table(doc, headers, rows, widths=None, size=8.2, header_fill=PALE_BLUE):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        cell_text(t.cell(0, i), h, True, size, INK)
        shade(t.cell(0, i), header_fill)
    repeat_header(t.rows[0])
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cell_text(cells[i], value, size=size)
    borders(t)
    t.autofit = False
    if widths:
        for row in t.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Inches(width)
    return t


def para(doc, text, size=9.6, bold=False, color="222222", italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.1
    r = p.add_run(text)
    font(r, size=size, color=color, bold=bold, italic=italic)
    return p


def bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.06
    font(p.add_run(text), size=9.2)
    return p


def number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.06
    font(p.add_run(text), size=9.2)
    return p


def callout(doc, title, text, fill=PALE_GOLD, title_color=INK):
    t = doc.add_table(rows=1, cols=1)
    c = t.cell(0, 0)
    shade(c, fill)
    p = c.paragraphs[0]
    font(p.add_run(title), size=10, color=title_color, bold=True)
    p2 = c.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.08
    font(p2.add_run(text), size=9.15)
    repeat_header(t.rows[0])
    borders(t)
    return t


def hyperlink(p, label, url):
    rid = p.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    run.append(rpr)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    link.append(run)
    p._p.append(link)


def source(doc, label, url, note=""):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    hyperlink(p, label, url)
    if note:
        font(p.add_run(f"：{note}"), size=8.7, color=MUTED)


def page_number(p):
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])
    font(run, size=8, color=MUTED)


def heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    return p


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.68)
    sec.bottom_margin = Inches(0.68)
    sec.left_margin = Inches(0.76)
    sec.right_margin = Inches(0.76)
    sec.page_width = Inches(8.27)
    sec.page_height = Inches(11.69)
    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    for name, size, color in (("Title", 25, NAVY), ("Heading 1", 15.2, NAVY), ("Heading 2", 12.1, BLUE), ("Heading 3", 10.6, INK)):
        st = styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.color.rgb = rgb(color)
        st.font.bold = True
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    font(header.add_run("反倾销税风险穿透分析｜间甲酚｜2026-08-13"), size=8, color=MUTED)
    page_number(sec.footer.paragraphs[0])

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(42)
    font(p.add_run("反倾销税风险穿透分析"), size=12, color=BLUE, bold=True)
    p = doc.add_paragraph(style="Title")
    p.add_run("间甲酚（m-Cresol）")
    p = doc.add_paragraph()
    font(p.add_run("第三国转运、原产地与历史少缴税风险阶段审计报告"), size=15, color=INK, bold=True)
    para(doc, "基于易迅已下载HS 290712全球宽池7,072行逐条审计，并结合商务部措施、到期规则、生产能力和企业关系公开资料", size=10.2, color=MUTED)
    doc.add_paragraph()
    callout(
        doc,
        "核心结论",
        "现有数据只证实6条印度发运至中国的小包装间甲酚记录，以及369条受税来源国发往印度的上游背景记录；两组在收发货主体、批号、提单、数量和时间上均未闭合。公开法律记录支持措施于2026年1月14日期满、2026年1月15日起终止，官方网页仍有滞后展示冲突。现阶段不能认定第三国绕道或少缴税，也不能用易迅金额字段计算税款。",
        PALE_RED,
        RED,
    )
    table(
        doc,
        ["数据范围", "逐条结果", "最具体核查对象", "证据结论"],
        [["2023-01-02—2026-01-15", "7,072行；6,345全字段唯一", "FINAR/ACETO—RON PHARM；Sasol—VDH", "B级线索；0条A/B同批闭环"]],
        [1.45, 1.75, 2.55, 1.55],
        8.5,
    )
    para(doc, "报告性质：执法风险筛查与调证线索，不替代海关原产地认定、税款核定或刑事/行政结论。", size=8.7, color=MUTED, italic=True)
    doc.add_page_break()

    heading(doc, "一、执行摘要", 1)
    table(doc, ["维度", "全量结果", "研判"], [
        ["逐条覆盖", "7,072原始行；6,345个全12字段唯一记录", "全部保留分类、范围理由和路线判断；未因发现异常提前停止"],
        ["明确间甲酚", "1,414原始/1,259唯一", "严格按CAS/独立化学名纳入；含工业级、药用辅料和实验室小包装"],
        ["泛称待核", "939原始/730唯一", "只写CRESOLS或盐，必须由CAS/COA区分间、邻、对及混合异构体"],
        ["其他分类", "对甲酚2,032/1,881；邻甲酚678/647；间对混合539/523；衍生物256/246；同HS其他1,214/1,059", "均不得机械纳入间甲酚反倾销风险量"],
        ["明确对华", "6条；数量字段5.90、金额字段11,073.50", "数量包含L/ML包装口径；金额币种未载，不能写成kg或人民币/美元"],
        ["受税来源→印度", "369条唯一：德国216、美国125、日本21、西班牙5、法国1、比利时1", "供应背景成立；与对华6条0主体/批号/提单闭合"],
        ["第三国绕道", "公开源及易迅均未形成同批双段链", "不存在足以指向具体企业违法的直接证据"],
    ], [1.35, 3.0, 3.0], 8.3)
    heading(doc, "优先级判断", 2)
    bullet(doc, "B｜FINAR/ACETO PHARMA INDIA→RON PHARM上海：6条全部为PARENTEX产品，2023年4条FINAR记录中批号40338C309CW一致；可直接调发票、批次COA、生产厂和中国进口申报。")
    bullet(doc, "B｜Sasol Chemicals North America→VDH Organics：受税来源至印度最大的实体通道。平台记录显示大宗纯间甲酚进入印度，VDH又有印度本地产能；应核原料投入、精馏/反应工序、增值率和后续销售，但现有对华B腿不是VDH。")
    bullet(doc, "B-｜LANXESS Deutschland→LANXESS India、Mitsui日本生产—印度/新加坡销售实体：均是可核的集团商业链，关联或区域发货不等于转运，若第三国申报纯间甲酚原产则应核实际生产厂。")
    bullet(doc, "C｜两条南非Sasol及一条印度发中国的泛称CRESOLS唯一记录：产品异构体不明，暂不计入间甲酚数量。")
    callout(doc, "定性边界", "没有中国报关单原产国字段、原产地证、生产商栏、税款缴款书和同批A/B提单，不得将上述主体写成走私或逃税企业。", PALE_GREEN)

    heading(doc, "二、措施范围、历史税率与到期状态", 1)
    para(doc, "商务部公告2021年第2号自2021年1月15日起，对原产美国、欧盟及英国、日本的进口间甲酚征收反倾销税，期限5年。商品包括m-Cresol、meta-Cresol、3-Cresol、3-Methylphenol等，分子式C7H8O，CAS 108-39-4，中国税号29071211（查询可用2907121100）。国外290712、29071200、29071290均是宽口径入口，不能替代中国归类。")
    table(doc, ["原产地/企业", "历史AD税率", "AD及其13%进口VAT增量系数", "备注"], [
        ["美国全部企业（含Sasol）", "131.7%", "完税价格×148.821%", "历史核税期适用"],
        ["LANXESS Deutschland", "27.9%", "完税价格×31.527%", "列名德国企业"],
        ["其他欧盟及英国企业", "49.5%", "完税价格×55.935%", "按生产商税档核"],
        ["日本全部企业（含Mitsui、Honshu）", "54.8%", "完税价格×61.924%", "按日本原产核"],
    ], [2.0, 1.2, 2.2, 1.9], 8.4)
    callout(doc, "公式", "反倾销税=中国海关审定完税价格×适用税率；由漏缴反倾销税引起的进口增值税差额=反倾销税×13%。上表只适用于货物实际为受税来源原产、产品落入措施范围且中国端未足额缴税的条件情景。", PALE_GOLD)
    heading(doc, "2.1 2026年1月14日到期后的法律与网页冲突", 2)
    para(doc, "商务部2025年到期通知明确：利害关系方未申请且商务部未主动发起期终复审的，措施自到期日起终止。检索2025、2026年商务部公告目录未见间甲酚期终复审立案或续征公告。因此公开法律依据支持措施于2026年1月14日期满，2026年1月15日起终止。")
    table(doc, ["层次", "公开结果", "本报告口径"], [
        ["法律期限", "2021-01-15起5年；无复审立案/续征公告", "核税主时段截至2026-01-14"],
        ["实施中案件总表", "2026-03-13表仍列项目，但到期日写20260114、复审次数空", "视为状态展示冲突，不作为继续征税法律依据"],
        ["法规页元数据", "仍显示现行有效", "网页标签不能替代续征公告"],
        ["实际核税", "需核单一窗口税费参数/税款书/海关答复", "执法定案前再次确认"],
    ], [1.35, 3.2, 2.75], 8.4)
    callout(doc, "到期后边界", "2026年1月15日后的交易原则上不再计算“逃避本项反倾销税”。本地文件中明确间甲酚最新日期为2026年1月13日，到期后明确品0条。", PALE_RED, RED)

    heading(doc, "三、数据完整性、去重与范围分类", 1)
    table(doc, ["分类", "原始", "全字段唯一", "处理"], [
        ["明确间甲酚", "1,414", "1,259", "严格规则范围候选"],
        ["泛称cresol待核", "939", "730", "CAS/COA未确认前不计风险量"],
        ["对甲酚", "2,032", "1,881", "排除"],
        ["邻甲酚", "678", "647", "排除"],
        ["间对/甲酚酸混合物", "539", "523", "与纯间甲酚区分，原则上排除/另核"],
        ["衍生物/指示剂", "256", "246", "排除"],
        ["同HS其他", "1,214", "1,059", "排除或信息不足"],
        ["合计", "7,072", "6,345", "可见字段完全重复冗余727行"],
    ], [2.3, 1.1, 1.3, 2.65], 8.4)
    callout(doc, "严格口径纠偏", "广义关键词初筛曾得到约1,450条以上命中；本次将两轮规则逐票合并，优先排除氯代/戊基/磺酸盐/指示剂，分流邻、对、间对混合物，并补识别M-CRESOL+品牌名、俄文/西语写法。最终明确间甲酚为1,414原始/1,259唯一。报告和台账以这一可复核严格口径为准。", PALE_GREEN)
    bullet(doc, "文件覆盖2023-01-02至2026-01-15，缺少措施前半段2021-01-15至2022-12-31以及2026-01-15以后完整全球数据。")
    bullet(doc, "国外HS 290712宽池混有邻甲酚CAS 95-48-7、对甲酚CAS 106-44-5、间对混合物、氯甲酚、戊基间甲酚及其他酚类。")
    bullet(doc, "实验室/药用小包装如化学上是CAS 108-39-4仍可能落入措施范围；不能仅因“样品”“100ml/500ml”而排除。")
    bullet(doc, "重量、数量、金额来自不同国家数据源，单位/币种并不统一；全量台账并列原始和去重口径，不把平台重复直接认作多票。")

    doc.add_page_break()
    heading(doc, "四、6条印度→中国明确间甲酚记录", 1)
    six = [
        ["2025-08-12", "ACETO PHARMA (INDIA)", "PARENTEX样品", "—", "1.00", "0.94", "RON PHARM上海"],
        ["2025-05-15", "ACETO PHARMA (INDIA)", "PARENTEX样品", "—", "0.20", "1.76", "RON PHARM上海"],
        ["2023-09-13", "FINAR LIMITED", "PARENTEX 500ML", "40338C309CW", "0.50", "5.00", "RON PHARM上海"],
        ["2023-09-13", "FINAR LIMITED", "PARENTEX 100ML", "40338C309CW", "0.20", "2.00", "RON PHARM上海"],
        ["2023-03-27", "FINAR LIMITED", "PARENTEX 500ML", "40338C309CW", "3.00", "8,397.30", "RON PHARM上海"],
        ["2023-03-27", "FINAR LIMITED", "PARENTEX 100ML", "40338C309CW", "1.00", "2,666.50", "RON PHARM上海"],
    ]
    table(doc, ["日期", "印度供应方", "产品/包装", "批号", "数量字段", "金额字段", "中国买方"], six, [1.05, 1.7, 1.35, 1.25, .85, 1.0, 1.15], 7.8)
    para(doc, "六条均为印度全港出口记录、目的国中国、平台原产字段India、HS 29071290；合计数量字段5.90、金额字段11,073.50。包装描述明确出现100ML/500ML，数量字段不是统一kg；金额币种和贸易条件未载。")
    callout(doc, "最具体调证线索", "2023年4条FINAR记录共享批号40338C309CW；该批号在7,072行中只出现于这4条B腿，未发现受税来源→印度同批A腿。应从中国进口报关、FINAR批次放行记录、生产厂声明和COA反查原产，不能把同批号本身解释为受税来源。", PALE_GOLD)
    heading(doc, "4.1 FINAR、ACETO与RON PHARM的关系边界", 2)
    bullet(doc, "Actylis官方产品目录确认Metacresol CDMF PARENTEX、CAS 108-39-4及100ml/500ml等包装。")
    bullet(doc, "Aceto于2021年收购FINAR，后整合为Actylis；全表PARENTEX产品由2023—2024年的FINAR逐步转为2025年的ACETO PHARMA INDIA，符合产品线/经营主体连续性。")
    bullet(doc, "Actylis官方分销网络列RON PHARM上海为中国分销商。因此该链存在正常药用辅料/实验室产品分销解释；集团或分销关系不证明美国、欧盟、英国或日本原产。")

    heading(doc, "五、受税来源→印度369条：供应背景而非闭环", 1)
    table(doc, ["平台原产", "唯一记录", "主要通道/买方", "数量口径提示"], [
        ["德国", "216", "LANXESS→LANXESS India；Hedinger→Silaris；Merck→Merck India；Novo Nordisk→Torrent", "小包装与大宗混合，不可合并解释kg"],
        ["美国", "125", "Sasol Chemicals North America→VDH Organics等", "以约20,000数量字段ISO罐大宗为主"],
        ["日本", "21", "Tokyo Chemical Industry→TCI India等", "已剔除6-tert-butyl-m-cresol衍生物；多为实验室小包装"],
        ["西班牙/法国/比利时", "7", "欧洲其他实验室/贸易通道", "需按生产商与原产证核"],
        ["合计", "369", "数量字段3,851,518.57；金额字段178,499,117.30", "跨源单位/币种混杂，禁止作为吨数/美元合计"],
    ], [1.3, 1.0, 3.5, 1.85], 8.2)
    table(doc, ["归一主体", "唯一记录", "数量字段", "研判"], [
        ["Sasol Chemicals North America", "119", "2,736,174.06", "美国受税来源大宗供应背景"],
        ["VDH Organics（印度买方）", "92", "2,310,872.03", "印度有真实产能，但原料来源/加工强度需逐批核"],
        ["LANXESS Deutschland", "58", "1,076,240", "德国生产商至印度关联公司商业链"],
        ["AUG Hedinger", "79", "1,352.25", "药用/高纯产品分销通道"],
        ["Novo Nordisk", "32", "1,908", "药用辅料通道，生产厂与销售实体需区分"],
        ["Merck Life Science", "29", "44.2", "实验室小包装通道"],
        ["Tokyo Chemical Industry", "21", "42", "日本实验室试剂通道；已剔除一条叔丁基衍生物"],
    ], [2.45, 1.0, 1.55, 2.65], 8.2)
    heading(doc, "5.1 A/B闭环审计结果", 2)
    table(doc, ["匹配维度", "结果", "含义"], [
        ["A腿印度收货人与B腿印度供应方", "FINAR/ACETO重合0", "没有同企业承接"],
        ["批号", "40338C309CW仅在4条B腿出现", "没有同批A腿"],
        ["提单/柜号", "源表无可闭合键", "不能证明同货"],
        ["数量/包装", "A腿大宗与B腿小包装不对应", "存在分装/库存/合法采购等替代解释"],
        ["中国报关原产", "缺失", "平台India不等于法定申报原产"],
        ["结论", "0条A/B同批闭环", "只能定为供应背景+B腿核查线索"],
    ], [2.25, 2.1, 3.1], 8.4)

    heading(doc, "六、实体、产能与合法替代解释", 1)
    table(doc, ["实体/链", "公开事实", "风险与反证边界"], [
        ["VDH Chemtech/Organics", "印度官方材料曾列Meta Cresol 30吨/月；企业称生产99%间甲酚", "印度不是无产能第三国；但平台又显示大宗美国Sasol料进入VDH，应核当前许可、投入产出及后续批次"],
        ["Dorf Ketal India", "印度环境许可列Meta Cresol+2,5-Xylenol合计300吨/年、captive use", "证明印度存在内部生产能力，但不是可外销纯间甲酚300吨"],
        ["Sasol美国", "美国Winnie曾生产meta-cresol；2025-11-18许可因场址关闭注销", "后续声称Winnie新生产的批次应核生产日期/库存；不等于Sasol全球产能全部关闭"],
        ["Sasol南非", "官方业务页显示当地生产phenolics及mp-cresols", "混合mp-cresols可能为真实南非产；纯MC99仍需厂址和精馏记录"],
        ["Mitsui日本—印度/新加坡", "日本Iwakuni-Ohtake生产Meta/Para-cresol；印度、新加坡官网主体为销售/营销公司", "区域销售实体作发货人可正常；若申报印度/新加坡纯间甲酚原产，应核生产厂"],
        ["LANXESS德国—亚洲", "德国有m-Cresol产品；印度公开生产业务未列间甲酚", "关联公司/贸易商不等于生产商；第三国原产纯品需工厂证据"],
        ["中国本地产能", "浙江医药公开披露间甲酚项目投产", "中国本地供给是贸易变化的合法替代解释"],
    ], [1.85, 3.0, 2.6], 8.0)

    heading(doc, "七、非优惠原产地与加工判断", 1)
    para(doc, "反倾销适用非优惠原产地规则。多国生产以最后实质性改变地为原产地；单纯储存、分装、为销售包装等微小加工不赋予新原产地，海关还可以不考虑以规避反倾销为目的的加工。第29章通常需要使用本身四位税目以外原料制成，或满足从价百分比标准。")
    table(doc, ["印度环节", "初步判断", "需要的证据"], [
        ["美国/德国/日本纯间甲酚2907→印度分装/贴标→2907", "不满足四位税目改变；高规避风险", "原料提单、分装记录、标签、原产证、批号"],
        ["受税来源纯间甲酚在印度精制，仍归2907", "需核是否达到适用实质性改变或从价门槛；不能凭“加工”自动改原产", "BOM、CIF、工厂交货价、成本账、收率、能耗、精馏记录"],
        ["印度本地原料经反应/分离生产间甲酚", "可能构成真实印度原产", "原料税目、生产许可、工单、产量、批次COA、环保/能耗"],
        ["第三国销售实体开票、货物受税来源直发中国", "商业链可经过第三国，但法定原产仍按生产地", "提单装港、生产商、CO、发票/付款链"],
    ], [2.35, 2.8, 2.35], 8.2)
    callout(doc, "不得简化为国别推定", "印度既有小规模真实产能，也有进口、加工和跨国集团销售网络。仅凭“印度→中国”既不能认定印度原产，也不能认定绕道；必须逐批审查生产工序与原料平衡。", PALE_GREEN)

    heading(doc, "八、税差：能算什么、不能算什么", 1)
    bullet(doc, "现有6条B腿金额字段没有币种、贸易条件和中国完税价格，且包括样品；不能形成实际少缴税额。")
    bullet(doc, "若历史进口期内确认受税来源原产、落入产品范围且未缴税，才可按中国海关审定完税价格乘适用综合系数。")
    bullet(doc, "若同时存在虚报优惠原产地，还需另算关税税率差及其进口增值税传导，本报告系数不含该部分。")
    bullet(doc, "2026年1月15日后原则上不再有本项反倾销税差；不得把到期后贸易量写成逃避本项税款。")
    table(doc, ["必须取得字段", "用途"], [
        ["中国进口报关单：进口日期、10位税号、商品规格、原产国、生产商", "确认历史措施时段、产品范围和税档"],
        ["海关审定完税价格及币种", "唯一可用的税差基数"],
        ["反倾销税/保证金及进口增值税缴款书", "确认是否已经足额缴税"],
        ["CO、生产商声明、COA/CAS、批号", "确认化学品范围与法定原产"],
        ["A/B提单、柜号、PO、发票和付款", "确认是否同一批货及商业/物理路径"],
    ], [3.3, 4.2], 8.5)

    heading(doc, "九、进一步核查顺序", 1)
    for text in [
        "第一优先：调2023年批号40338C309CW四条FINAR→RON PHARM的中国报关单、COA、生产商声明、原产地证和缴款书；再调FINAR批次放行及原料采购。",
        "第二优先：调2025年两条ACETO PHARMA INDIA→RON PHARM样品记录，确认其是否为收购整合后的印度制造、第三国采购转售或样品直供。",
        "第三优先：按Sasol→VDH的92条实体通道抽取大宗ISO罐记录，核VDH当前CTO许可、原料库存、精馏/生产工单、产品批号和去向；只有发现对应中国B腿才升级。",
        "第四优先：对Mitsui India/Asia Pacific、LANXESS India/Singapore/Hong Kong等区域销售实体发运的纯MC99核生产厂；销售主体名称不能替代原产证据。",
        "补数据：查询2021-01-15—2022-12-31全球宽池；查询2026-01-15后数据用于验证到期后流量但不计算反倾销税；补中国2907121100全时段及CAS/别名关键词。",
        "匹配规则：日期窗口±30/45/60天、数量容差±1%—3%，并叠加纯度、包装、品牌、批号、柜号或PO；仅同HS/同国家不升级。",
    ]:
        number(doc, text)

    heading(doc, "十、公开证据检索结论", 1)
    para(doc, "截至2026年8月13日，未检出商务部针对间甲酚的反规避调查/裁决、中国海关公开的本品原产地伪报处罚、法院判决或能连接“受税来源出口—第三国进口—第三国对华出口”的公开双段提单。已识别的是集团销售网络、印度真实/进口加工产能和具体B腿，不是违法闭环。")
    callout(doc, "证据分级", "A：同批两段提单+中国报关原产/税单或执法文书；B+：日期、数量、纯度、品牌高度匹配且第三国无合理产能；B：具体B腿或集团商业链；C：只有税号、国别或品类重合。本项当前最高为B。", PALE_BLUE)

    heading(doc, "十一、主要公开来源", 1)
    sources = [
        ("商务部公告2021年第2号", "https://policy.mofcom.gov.cn/claw/clawContent.shtml?id=87934", "措施范围、税率、期限和公式"),
        ("中国反倾销条例", "https://xzfg.moj.gov.cn/front/law/detail?LawID=1548", "措施期限与期终复审规则"),
        ("商务部2026年上半年措施到期通知", "https://trb.mofcom.gov.cn/myjjdc/art/2025/art_f05d6372c0af4559824bddda14367eae.html", "未复审则到期终止"),
        ("商务部2025年公告目录", "https://www.mofcom.gov.cn/zcfb/blgg/gg/2025/index.html", "未见间甲酚期终复审"),
        ("商务部2026年公告目录", "https://www.mofcom.gov.cn/zcfb/blgg/gg/2026/index.html", "未见续征决定"),
        ("正在实施贸易救济措施案件总表", "https://admin.cacs.mofcom.gov.cn/cacscms/article/sszaj?articleId=160267&type=", "展示状态与法律到期存在冲突"),
        ("Actylis/Finar产品目录", "https://www.actylislab.com/products?id=M", "PARENTEX、CAS和包装"),
        ("Actylis分销网络", "https://www.actylislab.com/distribution-network", "RON PHARM上海为中国分销商"),
        ("Aceto收购Finar公告", "https://actylis.com/-/media/project/aceto/acetostorefront/files/press-releases/20210517_finar-press-release-final.pdf", "实体连续性"),
        ("VDH Meta Cresol产品页", "https://www.vdhchemtech.com/Meta-Cresol.html", "印度企业自述99%间甲酚产能"),
        ("印度CPCB/NGT关于VDH报告", "https://www.greentribunal.gov.in/sites/default/files/news_updates/Report%20of%20CPCB%20in%20OA%20No.%20450%20of%202022%20%28In%20re%2012%20dead%20in%20fire%20at%20Hapur%20Factory%20that%20Produced%20firecrackers%20illegally%20Saw%20people%20with%20burns%20jumping%20.....%29.pdf", "历史许可产能"),
        ("Mitsui日本工厂", "https://jp.mitsuichemicals.com/en/corporate/ds/works/index.htm", "日本Meta/Para-cresol生产"),
        ("Mitsui Chemicals India", "https://in.mitsuichemicals.com/corporate/oversea/india/index.htm", "印度主体为销售/市场开发"),
        ("LANXESS m-Cresol产品页", "https://lanxess.com/en/products/products/m/m-cresol", "德国集团产品"),
        ("LANXESS全球网点", "https://lanxess.com/en/company/organization/locations/locations-worldwide", "亚洲销售/关联实体"),
        ("PubChem m-Cresol", "https://pubchem.ncbi.nlm.nih.gov/compound/3-Methylphenol", "CAS 108-39-4及化学身份"),
        ("中国增值税法", "https://www.npc.gov.cn/npc/c2/c30834/202412/t20241225_442038.html", "普通货物进口增值税13%"),
        ("现行非优惠原产地实质性改变规则", "https://policy.mofcom.gov.cn/claw/clawContent.shtml?id=101350", "四位税目/从价标准"),
    ]
    for label, url, note in sources:
        source(doc, label, url, note)

    heading(doc, "十二、交付文件与限制", 1)
    for text in [
        "间甲酚_易迅逐票判定台账_阶段审计.xlsx：概览、7,072全量台账、明确间甲酚、对华6条、泛称对华、受税来源至印度、重复审计、政策与调证。",
        "间甲酚_易迅逐票标准化.csv/json：每行保留原始12字段、全字段签名、分类、范围理由、路线与证据等级。",
        "间甲酚_明确对华6条.csv、间甲酚_受税来源至印度A腿_最终369条.csv、间甲酚_泛称对华去重3条.csv、分类纠偏审计CSV及摘要JSON。",
        "本报告为阶段审计：仍需补齐2021—2022时段、到期后全球数据、中国底单和同批提单后再更新总报告、总清单和前端系统。",
    ]:
        bullet(doc, text)

    doc.core_properties.title = "间甲酚反倾销税与第三国转运风险阶段审计报告"
    doc.core_properties.subject = "易迅7,072行逐票审计、印度供应链、措施到期状态与历史税差"
    doc.core_properties.author = "Codex"
    doc.core_properties.keywords = "间甲酚,m-Cresol,反倾销税,印度,第三国转运,原产地"
    doc.save(REPORT)
    print(REPORT)


if __name__ == "__main__":
    main()
