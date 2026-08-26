from __future__ import annotations

import os
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT_DIR = Path(
    os.environ.get(
        "EPDM_REPORT_DIR",
        r"D:\易迅数据\反倾销税深度分析报告\12_EPDM",
    )
)
OUTPUT_FILE = OUTPUT_DIR / "EPDM_反倾销税与第三国转运风险阶段审计报告.docx"

# Design authority: standard_business_brief + memo_masthead.
# Page geometry and component tokens are deliberately explicit.
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_TOP_BOTTOM_DXA = 80
CELL_START_END_DXA = 120

NAVY = "17324D"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "152536"
BODY = "252525"
MUTED = "66788A"
PALE_BLUE = "E8EEF5"
PALE_GRAY = "F2F4F7"
PALE_GOLD = "FFF4CC"
PALE_RED = "FCE8E6"
PALE_GREEN = "E6F4EA"
BORDER = "CBD5E1"
RED = "8B1E1E"
GOLD = "7A5A00"
GREEN = "2F6B3C"


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_run_font(
    run,
    *,
    size: float = 11,
    color: str = BODY,
    bold: bool = False,
    italic: bool = False,
    latin: str = "Calibri",
    east_asia: str = "Microsoft YaHei",
):
    run.font.name = latin
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    run.bold = bold
    run.italic = italic
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:cs"), latin)


def set_style_font(style, *, size: float, color: str, bold: bool = False):
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.color.rgb = rgb(color)
    style.font.bold = bold
    rpr = style._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), "Calibri")
    rfonts.set(qn("w:hAnsi"), "Calibri")
    rfonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def set_cell_shading(cell, fill: str):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width_dxa(cell, width: int):
    tcpr = cell._tc.get_or_add_tcPr()
    tcw = tcpr.find(qn("w:tcW"))
    if tcw is None:
        tcw = OxmlElement("w:tcW")
        tcpr.append(tcw)
    tcw.set(qn("w:type"), "dxa")
    tcw.set(qn("w:w"), str(width))


def set_table_geometry(table, widths_dxa: list[int]):
    if sum(widths_dxa) != CONTENT_WIDTH_DXA:
        raise ValueError(f"Table widths must sum to {CONTENT_WIDTH_DXA}: {widths_dxa}")
    table.autofit = False
    tbl = table._tbl
    tblpr = tbl.tblPr

    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:type"), "dxa")
    tblw.set(qn("w:w"), str(CONTENT_WIDTH_DXA))

    tblind = tblpr.find(qn("w:tblInd"))
    if tblind is None:
        tblind = OxmlElement("w:tblInd")
        tblpr.append(tblind)
    tblind.set(qn("w:type"), "dxa")
    tblind.set(qn("w:w"), str(TABLE_INDENT_DXA))

    layout = tblpr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblpr.append(layout)
    layout.set(qn("w:type"), "fixed")

    margins = tblpr.find(qn("w:tblCellMar"))
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        tblpr.append(margins)
    for tag, val in (
        ("top", CELL_TOP_BOTTOM_DXA),
        ("bottom", CELL_TOP_BOTTOM_DXA),
        ("start", CELL_START_END_DXA),
        ("end", CELL_START_END_DXA),
    ):
        child = margins.find(qn(f"w:{tag}"))
        if child is None:
            child = OxmlElement(f"w:{tag}")
            margins.append(child)
        child.set(qn("w:type"), "dxa")
        child.set(qn("w:w"), str(val))

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for cell, width in zip(row.cells, widths_dxa):
            set_cell_width_dxa(cell, width)


def set_table_borders(table, color: str = BORDER, size: str = "4"):
    tblpr = table._tbl.tblPr
    borders = tblpr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblpr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        old = borders.find(qn(f"w:{edge}"))
        if old is not None:
            borders.remove(old)
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)


def mark_header_row(row):
    trpr = row._tr.get_or_add_trPr()
    header = trpr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        trpr.append(header)
    header.set(qn("w:val"), "true")


def prevent_row_split(row):
    trpr = row._tr.get_or_add_trPr()
    if trpr.find(qn("w:cantSplit")) is None:
        trpr.append(OxmlElement("w:cantSplit"))


def set_cell_text(
    cell,
    value,
    *,
    size: float = 9.2,
    color: str = BODY,
    bold: bool = False,
    align=WD_ALIGN_PARAGRAPH.LEFT,
):
    cell.text = "" if value is None else str(value)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for p in cell.paragraphs:
        p.alignment = align
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.08
        p.paragraph_format.keep_together = True
        for run in p.runs:
            set_run_font(run, size=size, color=color, bold=bold)


def add_table(
    doc,
    headers: list[str],
    rows: list[list[str]],
    widths_dxa: list[int],
    *,
    header_fill: str = PALE_GRAY,
    font_size: float = 9.0,
    aligns=None,
):
    table = doc.add_table(rows=1, cols=len(headers))
    if aligns is None:
        aligns = [WD_ALIGN_PARAGRAPH.LEFT] * len(headers)
    for idx, header in enumerate(headers):
        set_cell_text(
            table.rows[0].cells[idx],
            header,
            size=font_size,
            color=INK,
            bold=True,
            align=aligns[idx],
        )
        set_cell_shading(table.rows[0].cells[idx], header_fill)
    mark_header_row(table.rows[0])
    prevent_row_split(table.rows[0])
    for row_values in rows:
        row = table.add_row()
        prevent_row_split(row)
        for idx, value in enumerate(row_values):
            set_cell_text(
                row.cells[idx],
                value,
                size=font_size,
                align=aligns[idx],
            )
    set_table_geometry(table, widths_dxa)
    set_table_borders(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_paragraph(
    doc,
    text: str = "",
    *,
    size: float = 11,
    bold: bool = False,
    color: str = BODY,
    italic: bool = False,
    after: float = 6,
    before: float = 0,
    line: float = 1.10,
    align=WD_ALIGN_PARAGRAPH.LEFT,
    keep=False,
):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    p.paragraph_format.keep_together = keep
    run = p.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold, italic=italic)
    return p


def add_mixed_paragraph(doc, segments, *, after=6, line=1.10, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    p.paragraph_format.keep_together = keep
    for text, attrs in segments:
        r = p.add_run(text)
        set_run_font(r, **attrs)
    return p


def add_callout(doc, title: str, text: str, *, fill=PALE_BLUE, title_color=INK):
    # A labeled shaded paragraph is semantically preferable to a one-cell
    # layout table and avoids an artificial table-header accessibility issue.
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.10)
    p.paragraph_format.right_indent = Inches(0.10)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.line_spacing = 1.10
    p.paragraph_format.keep_together = True
    ppr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)
    pbdr = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "5")
        border.set(qn("w:space"), "4")
        border.set(qn("w:color"), BORDER)
        pbdr.append(border)
    ppr.append(pbdr)
    set_run_font(p.add_run(title), size=11.2, color=title_color, bold=True)
    p.add_run().add_break()
    set_run_font(p.add_run(text), size=10.2, color=BODY)
    return p


def add_heading(doc, text: str, level: int = 1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.keep_together = True
    return p


def add_hyperlink(paragraph, label: str, url: str):
    rel_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), "Calibri")
    rfonts.set(qn("w:hAnsi"), "Calibri")
    rfonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    rpr.append(rfonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "20")
    rpr.append(size)
    run.append(rpr)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def create_numbering(doc, *, bullet: bool):
    numbering = doc.part.numbering_part.element
    existing_abs = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    existing_num = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abs_id = max(existing_abs or [0]) + 1
    num_id = max(existing_num or [0]) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abs_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    numfmt = OxmlElement("w:numFmt")
    numfmt.set(qn("w:val"), "bullet" if bullet else "decimal")
    lvl.append(numfmt)
    lvltext = OxmlElement("w:lvlText")
    lvltext.set(qn("w:val"), "•" if bullet else "%1.")
    lvl.append(lvltext)
    jc = OxmlElement("w:lvlJc")
    jc.set(qn("w:val"), "left")
    lvl.append(jc)
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    ppr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    ppr.append(ind)
    lvl.append(ppr)
    if bullet:
        rpr = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:ascii"), "Calibri")
        fonts.set(qn("w:hAnsi"), "Calibri")
        rpr.append(fonts)
        lvl.append(rpr)
    abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    absref = OxmlElement("w:abstractNumId")
    absref.set(qn("w:val"), str(abs_id))
    num.append(absref)
    numbering.append(num)
    return num_id


def add_list_item(doc, text: str, num_id: int, *, bold_prefix: str | None = None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    ppr = p._p.get_or_add_pPr()
    numpr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    numid = OxmlElement("w:numId")
    numid.set(qn("w:val"), str(num_id))
    numpr.extend([ilvl, numid])
    ppr.append(numpr)
    if bold_prefix and text.startswith(bold_prefix):
        set_run_font(p.add_run(bold_prefix), size=10.6, color=INK, bold=True)
        set_run_font(p.add_run(text[len(bold_prefix):]), size=10.6, color=BODY)
    else:
        set_run_font(p.add_run(text), size=10.6, color=BODY)
    return p


def add_page_number_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(paragraph.add_run("EPDM阶段审计  |  "), size=8.5, color=MUTED)
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    display = OxmlElement("w:t")
    display.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, display, end])
    set_run_font(run, size=8.5, color=MUTED)


def configure_document(doc: Document):
    doc.settings.odd_and_even_pages_header_footer = False
    section = doc.sections[0]
    section.different_first_page_header_footer = False
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    set_style_font(normal, size=11, color=BODY)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    normal.paragraph_format.widow_control = True

    title = doc.styles["Title"]
    set_style_font(title, size=26, color=NAVY, bold=True)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(5)
    title.paragraph_format.line_spacing = 1.0
    title_ppr = title._element.get_or_add_pPr()
    old_border = title_ppr.find(qn("w:pBdr"))
    if old_border is not None:
        title_ppr.remove(old_border)

    specs = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in specs.items():
        style = doc.styles[name]
        set_style_font(style, size=size, color=color, bold=True)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.space_after = Pt(0)
    set_run_font(
        header.add_run("反倾销税风险穿透分析｜三元乙丙橡胶（EPDM）｜2026-08-13"),
        size=8.5,
        color=MUTED,
    )
    footer = section.footer.paragraphs[0]
    footer.paragraph_format.space_after = Pt(0)
    add_page_number_field(footer)

    props = doc.core_properties
    props.title = "EPDM反倾销税与第三国转运风险阶段审计报告"
    props.subject = "易迅逐票全量审计、公开证据核验、第三国绕道与税款风险"
    props.author = "反倾销税风险分析项目"
    props.keywords = "EPDM; 反倾销税; 第三国转运; 原产地; 易迅数据; 风险审计"


def add_source(doc, bullet_id: int, code: str, title: str, url: str, note: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.10
    ppr = p._p.get_or_add_pPr()
    numpr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    numid = OxmlElement("w:numId")
    numid.set(qn("w:val"), str(bullet_id))
    numpr.extend([ilvl, numid])
    ppr.append(numpr)
    set_run_font(p.add_run(f"[{code}] "), size=10, color=INK, bold=True)
    add_hyperlink(p, title, url)
    set_run_font(p.add_run(f"。{note}"), size=10, color=MUTED)


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)
    bullet_id = create_numbering(doc, bullet=True)
    decimal_policy_id = create_numbering(doc, bullet=False)
    decimal_hexpol_id = create_numbering(doc, bullet=False)

    # Opening block: memo_masthead without decorative rule.
    add_paragraph(
        doc,
        "阶段审计｜线索核查工作稿",
        size=10.5,
        bold=True,
        color=BLUE,
        after=7,
        before=8,
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.0
    p.add_run("三元乙丙橡胶（EPDM）")
    for r in p.runs:
        set_run_font(r, size=26, color=NAVY, bold=True)
    add_paragraph(
        doc,
        "反倾销税、第三国转运与原产地风险阶段审计报告",
        size=15,
        bold=True,
        color=INK,
        after=18,
        line=1.05,
    )
    metadata = [
        ("审计对象", "HS 400270全球贸易宽池；措施税号为40027010、40027090"),
        ("数据期间", "2025-09-01至2026-08-05"),
        ("审计规模", "9,356条原始记录；9,160条12字段完全去重记录"),
        ("基准日期", "2026-08-13"),
        ("结论属性", "风险筛查与调证建议，不替代海关归类、原产地审定或案件定性"),
    ]
    for label, value in metadata:
        add_mixed_paragraph(
            doc,
            [
                (f"{label}：", {"size": 10.5, "color": INK, "bold": True}),
                (value, {"size": 10.5, "color": BODY}),
            ],
            after=4,
            line=1.08,
        )

    add_paragraph(doc, "", after=7)
    add_callout(
        doc,
        "核心结论",
        "公开材料已确认存在一宗三元乙丙橡胶反倾销产品走私系列案，涉及4,900余吨、案值2.87亿元，属于A档“产品级走私确证”；但公开材料未披露来源国、第三国、进口口岸、涉案企业、报关行、具体手法或实际税损，因此不能把该案与本次易迅记录中的任何企业或路线相连。易迅全量逐票审计中，最强的具体第三国线索是HEXPOL墨西哥同一实体的418条受税来源A腿与4条墨西哥原产对华B腿，评为B+；其真实橡胶混炼产能又构成重要替代解释，现阶段仍无批次、提单、柜号、牌号或生产工单闭环。",
        fill=PALE_RED,
        title_color=RED,
    )
    add_paragraph(
        doc,
        "一句话判断：EPDM具备真实、已被执法验证的少缴反倾销税风险，但本样本尚未形成任何一条“受税来源原料—第三国未实质加工—同货对华—错误申报原产地—少缴税款”的完整证据链。",
        size=10.6,
        bold=True,
        color=NAVY,
        after=0,
        keep=True,
    )

    doc.add_page_break()

    add_heading(doc, "执行摘要", 1)
    for item, prefix in [
        (
            "数据完整性：已对工作簿全部9,356行逐条分类，按12个原始字段完全相等去重为9,160条；196条为重复展开记录。所有源行均保留记录ID、规范记录ID、重复组和判定理由。",
            "数据完整性：",
        ),
        (
            "对华记录：共55条，均为完全去重记录。其中受税来源直接对华6条、重量合计226,987千克；第三国来源对华49条。6条欧盟记录是直接核税对象，不是绕道证据。",
            "对华记录：",
        ),
        (
            "第三国结构：沙特40条、6,419,911.08千克；加拿大4条、81,320千克；墨西哥4条，数量字段618.2、金额字段3,378.88；印度1条，数量字段106、金额字段271.36。后两组的数量单位和金额币种未统一，不得擅自写成千克或美元。",
            "第三国结构：",
        ),
        (
            "最强路线：HEXPOL COMPOUNDING S.A. de C.V.在样本期内接收418条美国、韩国、荷兰、法国来源EPDM，并在2025-12-23以墨西哥原产字段向中国发出4条。实体、方向和时间关联成立；同货、原产地虚假和少缴税款均未成立。",
            "最强路线：",
        ),
        (
            "执法事实：“0716”案证明EPDM反倾销产品已经发生大规模走私，但仅能支持产品层面的高风险，不支持把金陵海关侦办等同于南京口岸，也不支持推断本次样本企业涉案。",
            "执法事实：",
        ),
        (
            "税种性质：若原产地或生产商申报导致反倾销税未足额缴纳，核心差额为反倾销税及由该反倾销税增加的进口环节增值税；实际金额必须以中国海关审定完税价格、最终原产地、具体生产商税率和税款缴款书计算。",
            "税种性质：",
        ),
    ]:
        add_list_item(doc, item, bullet_id, bold_prefix=prefix)

    add_callout(
        doc,
        "证据边界",
        "易迅数据属于贸易情报/提单与各国申报字段汇编，不等同于中国进口报关单、海关原产地审定、税款缴款书或司法裁判。平台“原产国/地区”字段是核查入口，不是最终法律结论；企业名称出现在记录中也不等于违法。",
        fill=PALE_GOLD,
        title_color=GOLD,
    )

    add_heading(doc, "审计范围与逐票方法", 1)
    add_heading(doc, "数据池、期间与去重", 2)
    add_paragraph(
        doc,
        "本次使用已下载的《EPDM_400270.xlsx》全球宽池，数据期间为2025年9月1日至2026年8月5日，共9,356条、12个原始字段。逐票标准化台账为每一行附加原始Excel行号、记录ID、12字段哈希、规范记录ID、重复序号、范围分类、路线分类、证据等级、判定理由和决定性缺口。完全去重仅在12个原始字段逐值完全相等时成立，不做模糊合并。结果为9,160条规范记录、196条重复展开记录、138个重复组，最大重复组11条。",
    )
    add_paragraph(
        doc,
        "该宽池的优势是覆盖不同国家和多语种描述，适合发现实体级双腿；局限是仅近一年、各数据源字段口径不同、部分重量/数量/金额缺失或单位币种未统一，并且不含完整中国报关单、税款缴款书、进口口岸、报关行、提单号和集装箱号。",
    )

    add_heading(doc, "全部9,160条规范记录的互斥分类", 2)
    add_table(
        doc,
        ["逐票范围分类", "规范记录", "占比", "审计含义"],
        [
            ["原胶/通用EPDM候选", "6,306", "68.84%", "描述为通用EPDM、合成橡胶或牌号；仍需成分和形态确认"],
            ["税目标准表述（初步纳入）", "504", "5.50%", "描述对应初级形状或板、片、带等标准语句"],
            ["板/片/带/卷/粒料等形态待核", "1,163", "12.70%", "形态可能落入措施范围，也可能为后续制品，需实物与申报要素"],
            ["混炼改性/TPV/TPE/功能聚合物待核", "300", "3.28%", "存在复合、改性或热塑性弹性体信号，不能仅凭HS 400270纳入"],
            ["疑似制品/型材/密封件等范围外", "887", "9.68%", "描述更接近制品、汽车零件、密封件或型材"],
            ["合计", "9,160", "100.00%", "五类互斥，全部规范记录已归类"],
        ],
        [3150, 1250, 1050, 3910],
        header_fill=PALE_BLUE,
        font_size=9.0,
        aligns=[
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.LEFT,
        ],
    )
    add_paragraph(
        doc,
        "上述分类是审计筛查口径，不是海关商品归类或商务部范围裁定。尤其是“混炼改性”和“形态待核”记录，需要通过化学组成、是否硫化、主聚合物比例、产品形态、用途、技术数据表和实物查验确定。",
        size=10.2,
        color=MUTED,
    )

    add_heading(doc, "证据等级与定性约束", 2)
    add_table(
        doc,
        ["等级", "本报告含义", "可以证明", "不能单独证明"],
        [
            ["A", "直接事实，但必须写明对象层级", "公开执法案可证产品级走私；贸易记录可证平台字段所载路线", "不能把产品级案件投射到具体企业；不能由平台字段直接认定违法"],
            ["B+", "同一实体双腿、方向与时间关联较强", "值得优先调取同批/生产/原产地资料", "同货转运、未实质加工、虚假原产地、少缴税款"],
            ["C+ / C", "范围、品牌、方向或弱实体线索", "形成抽样核查名单或替代解释", "具体绕道链条"],
            ["D", "行业或全球背景", "理解产能、品牌和供应网络", "本次对华风险"],
        ],
        [900, 2300, 2580, 3580],
        font_size=8.8,
        aligns=[
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.LEFT,
        ],
    )

    add_heading(doc, "现行反倾销措施与税款机制", 1)
    add_heading(doc, "政策时间轴和当前状态", 2)
    add_list_item(
        doc,
        "2020-12-20：商务部2020年第60号公告开始对原产于美国、韩国和欧盟的进口EPDM征收反倾销税，原期限5年。[S2]",
        decimal_policy_id,
    )
    add_list_item(
        doc,
        "2025-12-20：商务部2025年第81号公告启动期终复审；复审期间对美国、韩国和欧盟继续按原范围和税率征收。英国因无人申请复审，自当日起措施到期终止。[S1]",
        decimal_policy_id,
    )
    add_list_item(
        doc,
        "当前基准日2026-08-13：期终复审仍在公告确定的调查期内，应于2026-12-20前（不含当日）结束；因此美国、韩国、欧盟仍是当前受税来源，英国不是当前受税来源。[S1]",
        decimal_policy_id,
    )
    add_paragraph(
        doc,
        "措施产品名称为三元乙丙橡胶（Ethylene-Propylene-non-conjugated Diene Rubber / Ethylene Propylene Diene Monomer，EPDM），公告列入中国税则号40027010和40027090。HS编码只是申报入口，最终范围仍取决于公告产品描述和货物实际属性。[S1][S2]",
    )

    add_heading(doc, "现行公司税率与“反倾销税×1.13”增量系数", 2)
    add_paragraph(
        doc,
        "商务部公告明确：反倾销税额＝海关审定完税价格×反倾销税率；进口环节增值税的计税价格包含反倾销税。[S2] 2026年《增值税法》对通常进口货物适用13%税率。[S3] 因此，仅计算“反倾销税及其引起的进口增值税增量”时，系数为反倾销税率×1.13。该系数不含关税本身，也不代表案件的最终应补税额。",
    )
    add_table(
        doc,
        ["来源", "生产商/适用档", "反倾销税率", "AD及其VAT增量系数"],
        [
            ["美国", "陶氏化学；其他美国公司", "222.0%", "250.860%"],
            ["美国", "埃克森美孚", "214.9%", "242.837%"],
            ["美国", "阿朗新科美国；美国狮子弹性体", "219.8%", "248.374%"],
            ["韩国", "锦湖POLYCHEM", "12.5%", "14.125%"],
            ["韩国", "乐天玮萨黎司弹性体", "21.1%", "23.843%"],
            ["韩国", "其他韩国公司", "24.5%", "27.685%"],
            ["欧盟", "阿朗新科荷兰", "18.1%", "20.453%"],
            ["欧盟", "埃克森美孚化工法国", "14.7%", "16.611%"],
            ["欧盟", "意大利玮萨黎司", "16.5%", "18.645%"],
            ["欧盟", "其他欧盟公司", "31.7%", "35.821%"],
        ],
        [1250, 3800, 1710, 2600],
        font_size=8.8,
        aligns=[
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.CENTER,
        ],
    )

    add_callout(
        doc,
        "税额计算的硬约束",
        "易迅的“金额”字段不是中国海关审定完税价格，且不同数据源币种、贸易术语和统计口径不一。没有中国进口报关单、价格资料、原产地审定、具体生产商和税款缴款书，不得把平台金额乘税率写成实际逃税额。",
        fill=PALE_GOLD,
        title_color=GOLD,
    )

    add_heading(doc, "对华55条的全量结构", 1)
    add_paragraph(
        doc,
        "9,160条规范记录中，目的国/地区字段为中国的记录共55条，且55条均为12字段完全去重记录。按原产国字段分为受税来源直接对华6条和第三国来源对华49条。",
    )
    add_table(
        doc,
        ["路线组", "记录数", "重量字段", "数量/金额字段", "阶段判断"],
        [
            ["欧盟→中国（比利时5、意大利1）", "6", "226,987 kg", "数量字段212；金额空", "直接核税对象；不是绕道"],
            ["沙特→中国", "40", "6,419,911.08 kg", "数量字段191,090.00；金额空", "真实产能强反证；低绕道风险"],
            ["加拿大→中国", "4", "81,320 kg", "数量字段152；金额空", "FUSABOND N302范围/产地核验"],
            ["墨西哥→中国", "4", "重量空", "数量618.2；金额3,378.88", "HEXPOL实体级双腿，B+"],
            ["印度→中国", "1", "重量空", "数量106；金额271.36", "KEI B腿早于A腿，时序不支持"],
            ["合计", "55", "6,728,218.08 kg（50条有值）", "单位/币种不统一", "逐组采取不同核查策略"],
        ],
        [2470, 900, 1800, 2060, 2130],
        font_size=8.5,
        aligns=[
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.LEFT,
        ],
    )
    add_paragraph(
        doc,
        "重量字段在上述50条记录中表现为千克并按任务口径汇总；墨西哥、印度记录重量为空，其数量和金额字段未提供统一单位/币种。本报告保留字段属性，不做跨字段换算。",
        size=10.2,
        color=MUTED,
    )

    doc.add_page_break()
    add_heading(doc, "欧盟直接对华6条：应核税，不是第三国绕道", 2)
    add_table(
        doc,
        ["日期", "原产地", "供应商→中国采购商", "重量", "审计判断"],
        [
            ["2026-03-17", "比利时", "VESPER POLYMERS AG / CH → DEZHOU TIANDING", "84,600 kg", "直接受税来源"],
            ["2026-03-17", "比利时", "VESPER POLYMERS AG / CH → DEZHOU TIANDING", "28,200 kg", "直接受税来源"],
            ["2026-02-04", "比利时", "VESPER POLYMERS AG / CH → DEZHOU TIANDING", "23,180 kg", "直接受税来源"],
            ["2026-01-20", "意大利", "VERSALIS S.P.A. → SHANGHAI HAO SHEN", "17,747 kg", "直接受税来源"],
            ["2025-11-21", "比利时", "VESPER POLYMERS AG / CH → DEZHOU TIANDING", "46,360 kg", "直接受税来源"],
            ["2025-11-10", "比利时", "VESPER POLYMERS AG / CH → DEZHOU TIANDING", "26,900 kg", "直接受税来源"],
        ],
        [1400, 1000, 4140, 1280, 1540],
        font_size=8.3,
        aligns=[
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.CENTER,
        ],
    )
    add_paragraph(
        doc,
        "风险点在税率适用和缴纳闭环，而非路线规避：应核对中国报关单的原产国、生产商、40027010/40027090归类、征免性质、反倾销税率、完税价格和税款缴款书。贸易商注册地（例如“/CH”）不能替代聚合生产地。",
    )

    add_heading(doc, "HEXPOL墨西哥线：当前最强的具体第三国线索", 1)
    add_heading(doc, "双腿事实和时间关联", 2)
    add_paragraph(
        doc,
        "A腿：HEXPOL COMPOUNDING S.A. de C.V.在样本期内作为采购商接收418条受税来源EPDM记录，其中美国389条、韩国19条、荷兰6条、法国4条；数量字段合计2,987,820.95，金额字段合计8,477,644.31，单位和币种不可由源表统一确认。B腿：同一实体于2025年12月23日以墨西哥原产字段、HS 4002700100、通用西语EPDM描述向TE CONNECTIVITY CORP（目的国中国）发出4条。",
    )
    add_table(
        doc,
        ["B腿日期", "商品描述", "数量字段", "金额字段", "字段原产地"],
        [
            ["2025-12-23", "CAUCHO ETILENO PROPILENO DIENO NO CONJUGADO EPDM", "124.70", "569.90", "Mexico"],
            ["2025-12-23", "同上", "119.90", "541.97", "Mexico"],
            ["2025-12-23", "同上", "119.40", "538.47", "Mexico"],
            ["2025-12-23", "同上", "254.20", "1,728.54", "Mexico"],
            ["合计", "4条；重量字段均为空", "618.20", "3,378.88", "Mexico"],
        ],
        [1420, 3970, 1180, 1300, 1490],
        font_size=8.6,
        aligns=[
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.CENTER,
        ],
    )
    add_paragraph(
        doc,
        "截至B腿日期，样本中已有171条受税来源A腿；B腿前30日有36条，前7日有3条。A腿中270条与B腿具有相同的规范化通用描述，其中78条发生在B腿之前。离B腿最近的2025年12月19日A腿含两条美国来源记录，数量字段均为8,250.14，描述分别含ROYALENE 547以及ROYALENE 525/556。上述事实构成实体、方向和时间层面的强关联。",
        keep=True,
    )

    add_heading(doc, "为何只能评为B+，不能评为闭环", 2)
    for text in [
        "B腿4条没有重量字段，且数量618.2无法与任何A腿数量建立同批对应；记录未提供统一单价、币种和贸易术语。",
        "A腿出现ROYALENE、NORDEL、VISTALON、KEP等牌号，B腿仅写通用EPDM；没有牌号、批次、采购订单、生产批号或配方对应。",
        "现有表不含中国进口报关单号、墨西哥出口报关单号、提单号、集装箱号、船名航次和中转节点，无法做物流闭合。",
        "HEXPOL官方资料明确其Querétaro工厂从事橡胶混炼产品的开发和生产，涉及材料、配方和混炼工艺。[S7] 因而存在“受税来源原胶在墨西哥经真实混炼后形成新产品”的合法替代解释。",
        "是否发生原产地实质性改变必须依据最终产品归类、增值比例、主要加工工序以及产品特定规则判定；不能从“同一实体既进口又出口”自动推导。",
    ]:
        add_list_item(doc, text, bullet_id)

    add_callout(
        doc,
        "B+结论",
        "HEXPOL线应列为第一优先调证对象，但报告不指控HEXPOL、TE Connectivity或其他相关方违法。决定性问题是：4条B腿究竟是原胶简单转运，还是在墨西哥完成足以改变税则归类、基本特征或适用规则的混炼生产；以及中国进口申报是否如实反映最终产品和原产地。",
        fill=PALE_GOLD,
        title_color=GOLD,
    )

    add_heading(doc, "其他第三国路线：强反证、范围问题与时序否定", 1)
    add_heading(doc, "沙特40条：真实聚合产能是强反证", 2)
    add_paragraph(
        doc,
        "沙特来源对华记录共40条、6,419,911.08千克，是第三国对华数量主体。供应链字段分布为：ExxonMobil相关29条、5,207,019.56千克；SABIC相关7条、660,000千克；Sumitomo相关3条、491,091.52千克；Ravago 1条、61,800千克。样本未发现受税来源→沙特的对应实体级A腿。",
    )
    add_paragraph(
        doc,
        "公开一手资料提供了生产能力反证：SABIC公布的KEMYA Al-Jubail工厂认证范围明确包含EPDM制造；ExxonMobil的Vistalon储存指南明确列出美国Baton Rouge和沙特Al Jubail生产地；Petro Rabigh官方产品页列有EPR/EPDM相关聚合产品。[S8][S9][S11] 因此，Vistalon、Sumitomo等品牌或销售主体名称不能被当成美国、日本或其他受税来源的原产地证据。",
        keep=True,
    )
    add_callout(
        doc,
        "阶段判断：C（低风险背景）",
        "现有数据更支持沙特真实生产后正常对华供应。仅当具体牌号、包装、生产批号或原产地证与沙特工厂能力矛盾，或发现受税来源A腿与对华B腿同批闭合时，才升级。",
        fill=PALE_GREEN,
        title_color=GREEN,
    )

    add_heading(doc, "加拿大4条：FUSABOND N302首先是范围与生产地问题", 2)
    add_paragraph(
        doc,
        "4条记录均为20,330千克，合计81,320千克，商品描述含FUSABOND N302 FUNCTIONAL POLYMER，供应商为DOW EUROPE GMBH，采购商为Celanese（Shanghai）相关名称，字段原产地为加拿大。2026年5月20日存在两条采购商名称略异的记录；虽不属于12字段完全重复，但可能对应同一物理批次，需提单/柜号去重。",
    )
    add_paragraph(
        doc,
        "Dow官方把FUSABOND N302描述为酸酐改性乙烯弹性体，主要用作聚酰胺共混物的抗冲改性剂。[S10] 这与通用EPDM原胶并非同一清晰表述。应先核对聚合物组成、接枝工艺、是否属于EPDM措施产品、实际生产地点及400270归类；“欧洲供应商+加拿大原产字段”本身不构成美国经加拿大绕道。",
    )
    add_callout(
        doc,
        "阶段判断：C+（范围/产地核验）",
        "调取技术数据表、SDS、聚合/接枝工序说明、生产厂证明、原产地证、进口申报要素、提单和柜号。在范围未确认前，不计算反倾销税差。",
        fill=PALE_BLUE,
        title_color=DARK_BLUE,
    )

    add_heading(doc, "印度1条：KEI时序不支持本次B腿", 2)
    add_paragraph(
        doc,
        "2025年12月7日，KEI INDUSTRIES LIMITED向中国江苏CENMEN相关采购商发出1条“RUBBER COMPOUND EPDM”，数量字段106、金额字段271.36、原产地字段印度。样本中KEI接收受税来源EPDM的第一批记录出现在次日2025年12月8日（韩国），之后还有12月15日美国等记录；因此，现有样本不支持用这些A腿解释12月7日B腿。",
    )
    add_paragraph(
        doc,
        "这不是对更早期库存的绝对排除，但数据池自2025年9月1日起已覆盖B腿前约三个月，仍未发现对应A腿；同时商品描述为橡胶混炼胶，措施范围和印度加工是否实质改变也需先核验。阶段评级为C/C+，优先级低于HEXPOL。",
    )

    add_heading(doc, "公开执法证据：“0716”案的证明力与边界", 1)
    add_heading(doc, "A档产品级走私确证", 2)
    add_paragraph(
        doc,
        "南京晨报刊载南京海关2024年打击走私十大典型案例，其中金陵海关缉私分局侦办的“0716”特大走私三元乙丙橡胶系列案载明：涉案团伙为逃避海关监管、偷逃巨额税款，将4,900余吨反倾销产品三元乙丙橡胶走私进境，案值2.87亿元。[S6] 这直接证明EPDM反倾销产品存在已被执法查证的大规模走私风险，因此对“商品层级”给予A档。",
    )
    add_heading(doc, "公开材料不能支持的推断", 2)
    for text in [
        "未披露原产于美国、韩国还是欧盟，也未披露是否经过墨西哥、印度、沙特、加拿大或其他第三国。",
        "未披露进口口岸、启运港、途经港、船名航次、集装箱、报关单、涉案企业、境外收发货人、报关行或货代。",
        "未披露具体走私手法，不能自行写成伪报品名、伪报原产地、低报价格、夹藏或绕关。",
        "“金陵海关缉私分局侦办”说明侦办机关，不等于货物经南京口岸进境；公开材料不足以给任何口岸贴高风险标签。",
        "案值2.87亿元不是法定意义上的海关审定完税价格，也不是税款损失；不能据此直接反推实际逃税额。",
    ]:
        add_list_item(doc, text, bullet_id)

    add_heading(doc, "2.87亿元案值的反事实情景区间", 2)
    add_paragraph(
        doc,
        "仅为展示税率敏感性，假设（并非事实）2.87亿元全部等于海关审定完税价格，且全部货物分别适用当前最低12.5%或最高222.0%反倾销税率，则“反倾销税＋其引起的13%进口增值税增量”理论区间为：",
    )
    add_table(
        doc,
        ["反事实情景", "增量系数", "机械计算", "结果"],
        [
            ["最低：韩国锦湖12.5%", "14.125%", "287,000,000 × 14.125%", "40,538,750元"],
            ["最高：美国222.0%", "250.860%", "287,000,000 × 250.860%", "719,968,200元"],
        ],
        [2540, 1660, 2730, 2430],
        font_size=9.0,
        aligns=[
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.RIGHT,
        ],
    )
    add_callout(
        doc,
        "不得误读",
        "上述区间不是“0716”案实际税损，也不是估计值。真实结果还取决于案值与完税价格的关系、货物来源国、生产商、不同批次、关税、征免、计税汇率和已缴税额。该情景的唯一用途，是说明EPDM税率差异极大，错误适用来源国/生产商可能产生显著税收后果。",
        fill=PALE_RED,
        title_color=RED,
    )

    add_heading(doc, "原产地法律规则：转运、包装与混炼应区别", 1)
    add_paragraph(
        doc,
        "《进出口货物原产地条例》明确适用于反倾销等非优惠贸易措施；两个以上国家参与生产的，以最后完成实质性改变的国家为原产地。为运输、贮存、装卸或销售进行的保存、包装等微小处理不影响原产地。[S4] 2024修正的非优惠原产地规则进一步规定，以四位税目归类改变为基本标准，必要时结合从价百分比和制造/加工工序；列明货物以具体清单规则为准。[S5]",
    )
    add_table(
        doc,
        ["情形", "通常审计含义", "EPDM所需证明"],
        [
            ["仓储、换单、转船、换包装、贴标", "通常不构成实质性改变；原产地原则上不因第三国中转改变", "全程提单、入出库、包装变更、原产地证、未再加工证明"],
            ["简单切片、分装或混合", "不能自动认定改变原产地；需看税目与基本特征", "工艺说明、成分比例、税则归类、增值资料"],
            ["配方混炼、加入填料/助剂并形成橡胶混炼胶", "可能形成不同产品，但须按具体规则审定，不能仅凭“加工”二字", "BOM、工单、领料、能耗、设备、产量、质量检验、预裁定/归类意见"],
            ["第三国重新聚合生产EPDM", "聚合工序通常是强原产地事实，仍须与法律规则和实际货物对应", "聚合釜批次、生产记录、原料、牌号、工厂证书、原产地核查"],
        ],
        [2200, 3200, 3960],
        font_size=8.8,
        aligns=[
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.LEFT,
        ],
    )
    add_paragraph(
        doc,
        "品牌不等于原产地，销售公司所在地也不等于原产地。对于Vistalon、Keltan、Dutral、KEP、NORDEL、ROYALENE等牌号，应以具体批次生产厂和聚合/加工工序为准；同一品牌可以存在多个生产地区和区域销售实体。",
    )

    add_heading(doc, "具体核查对象与调证路径", 1)
    add_heading(doc, "优先级清单", 2)
    add_table(
        doc,
        ["优先级", "对象/路线", "当前证据", "下一步必须取得"],
        [
            ["P1", "HEXPOL墨西哥→TE Connectivity（中国）4条", "B+同实体双腿、方向和时间；有真实混炼产能替代解释", "4条中国报关单、墨西哥出口申报、原产地证、BOM/工单/批次、提单柜号、入出库和税款"],
            ["P1", "“0716”案产品级执法材料", "A档产品级走私确证；公开信息缺路线和主体", "如依法可得：案号、扣押清单、报关单、口岸、报关行、境外链路、鉴定与计税材料"],
            ["P1", "比利时/意大利直接对华6条", "A档直接贸易字段；226,987千克", "生产商与税率匹配、完税价格、反倾销税和进口增值税缴款书"],
            ["P2", "FUSABOND N302加拿大4条", "C+范围/产地线索；可能有物理重复", "TDS/SDS、组成、接枝工序地点、归类意见、提单柜号、原产地证"],
            ["P3", "KEI印度1条", "B腿早于样本A腿；时序不支持", "B腿报关单、混炼配方/工单、此前库存和生产批次；未新增证据前不升级"],
            ["P3", "沙特40条", "真实EPDM产能强反证、无A腿闭合", "按牌号抽样核对生产批次与原产地证即可"],
        ],
        [900, 2310, 2740, 3410],
        font_size=8.4,
        aligns=[
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.LEFT,
        ],
    )

    add_heading(doc, "中国进口申报与口岸侧字段", 2)
    add_paragraph(
        doc,
        "本次易迅工作簿没有“进口口岸/主管海关、报关行、报关单号、提运单号、集装箱号”等字段，不能据此给出口岸风险排名。若开展行政核查或企业内审，应至少取得以下字段并按报关单逐票关联：",
    )
    for text in [
        "中国进口报关单号、申报日期、运输方式、进境关别/口岸代码、申报地海关、提运单号、船名航次/航班、集装箱号。",
        "境内收货人、消费使用单位、申报单位（报关行）、委托协议、经办人员、境外发货人、贸易国、启运国、经停地、原产国。",
        "10位商品编号、商品名称及规格型号、申报要素、成交方式、单价、总价、币制、运保费、完税价格、计税汇率。",
        "生产商、生产厂地址、品牌牌号、批次、第三单体、门尼粘度、充油量、是否混炼/改性/硫化、包装和净重。",
        "原产地证、非优惠原产地声明、未再加工证明、工厂证明、采购订单、发票、装箱单、全程提单、换单记录和保税仓进出记录。",
        "反倾销税税率、关税、进口增值税、担保/保证金、税款缴款书、补税或退税记录、海关估价/归类/原产地预裁定。",
    ]:
        add_list_item(doc, text, bullet_id)

    add_heading(doc, "HEXPOL双腿专门闭环测试", 2)
    for text in [
        "身份闭环：统一HEXPOL墨西哥法人名称、税号、工厂地址、出口主体和实际生产主体；排除集团不同实体误并。",
        "货物闭环：A腿原料牌号/批号/数量与B腿成品批号/数量、损耗、配方和库存移动一致。",
        "物流闭环：A腿进口提单/柜号、墨西哥仓库入库、生产领料、B腿出口提单/柜号、中国进口提单顺序可追溯。",
        "生产闭环：设备、班次、工单、能耗、质量检验、产量和废料能够证明实际混炼；若仅换包装或仓储，应保留相反证据。",
        "法律闭环：确认混炼后商品是否仍属商务部EPDM措施范围、是否发生四位税目改变或满足具体加工/增值规则，并视需要申请海关预裁定。",
        "税款闭环：以中国完税价格和最终生产商税率重算反倾销税及关联进口增值税，与缴款书逐票核对。",
    ]:
        add_list_item(doc, text, decimal_hexpol_id)

    add_heading(doc, "最终风险判断", 1)
    add_table(
        doc,
        ["问题", "结论", "置信度"],
        [
            ["EPDM是否存在真实走私/少缴反倾销税风险？", "是。公开“0716”案已在产品层级直接确认。", "高（A，产品级）"],
            ["是否已找到具体第三国绕道闭环？", "否。HEXPOL为最强B+线索，但缺同货、加工、原产地和税款闭环。", "中高（线索）/不足定性"],
            ["哪个第三国优先？", "墨西哥优先；加拿大以范围核查为主；印度时序不支持；沙特有强产能反证。", "中"],
            ["哪个口岸风险最高？", "无法判断。现有字段不含口岸；侦办机关不能替代进境口岸。", "高（数据缺口明确）"],
            ["涉及逃什么税？", "可能少缴反倾销税及由其引起的进口环节增值税；实际税额尚不能计算。", "高（税种）/不足量化"],
            ["能否指控具体企业违法？", "不能。现有企业名称仅用于核查排序，不构成违法认定。", "高"],
        ],
        [2950, 4880, 1530],
        font_size=8.8,
        aligns=[
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.CENTER,
        ],
    )
    add_callout(
        doc,
        "审计意见",
        "立即对HEXPOL墨西哥4条B腿和欧盟直达6条开展报关单—原产地—生产商—税款四项闭环；同步通过合法渠道获取“0716”案可公开或可共享的结构化字段，验证其是否涉及第三国原产地规避。加拿大FUSABOND先做范围裁定式技术审查；沙特路线在无新矛盾证据前维持低风险。",
        fill=PALE_GREEN,
        title_color=GREEN,
    )

    add_heading(doc, "数据附件索引", 1)
    add_paragraph(
        doc,
        "本报告结论可回溯至同目录结构化交付文件；逐票台账保留全部9,356条原始展开行，子集文件均从同一规范记录体系导出：",
    )
    for text in [
        "EPDM_易迅逐票标准化.csv / .json：9,356条全量逐票判定。",
        "EPDM_对华55条.csv / .json：全部对华规范记录。",
        "EPDM_受税来源直达中国6条.csv / .json：欧盟直接对华核税对象。",
        "EPDM_第三国对华49条.csv / .json：沙特、加拿大、墨西哥、印度对华记录。",
        "EPDM_HEXPOL_A腿418条.csv / .json 与 EPDM_HEXPOL_B腿4条.csv / .json：B+实体级双腿。",
        "EPDM_全12字段重复审计.csv / .json：重复组展开审计。",
        "EPDM_交付摘要.json：计数、口径、字段和自检结果。",
    ]:
        add_list_item(doc, text, bullet_id)

    add_heading(doc, "公开资料来源", 1)
    add_paragraph(
        doc,
        "以下链接均为本报告政策、执法、原产地、产品或产能事实的公开依据；网页状态以基准日为准。",
        size=10.2,
        color=MUTED,
    )
    sources = [
        (
            "S1",
            "商务部公告2025年第81号：EPDM反倾销措施期终复审",
            "https://policy.mofcom.gov.cn/claw/clawContent.shtml?id=104427",
            "确认当前受税来源、公司税率、英国终止、措施范围及调查期限",
        ),
        (
            "S2",
            "商务部公告2020年第60号：EPDM反倾销终裁",
            "https://dcj.mofcom.gov.cn/article/zcfb/gpmy/202012/20201203024350.shtml",
            "确认终裁范围、原税率、完税价格计征公式及进口增值税计税基础",
        ),
        (
            "S3",
            "中华人民共和国增值税法（国家税务总局政策法规库）",
            "https://fgk.chinatax.gov.cn/zcfgk/c100009/c5237365/content.html",
            "自2026年1月1日起施行；通常进口货物税率为13%",
        ),
        (
            "S4",
            "中华人民共和国进出口货物原产地条例（国家行政法规库）",
            "https://xzfg.moj.gov.cn/front/law/detail?LawID=1523",
            "适用于反倾销；规定最后实质性改变和微小加工规则",
        ),
        (
            "S5",
            "海关总署关于非优惠原产地规则中实质性改变标准的规定（2024修正）",
            "https://policy.mofcom.gov.cn/claw/clawContent.shtml?id=101350",
            "确认税目改变、从价百分比及制造加工工序判定框架",
        ),
        (
            "S6",
            "南京晨报：南京海关2024年打击走私十大典型案例",
            "https://doss.xhby.net/zpaper/njcb/pc/att/202502/18/8c924621-8910-48c9-968f-0441fdc593ca.pdf",
            "刊载“0716”案4,900余吨EPDM反倾销产品、案值2.87亿元",
        ),
        (
            "S7",
            "HEXPOL Querétaro工厂官方页面",
            "https://www.hexpol.com/plant-locations/hexpol-compounding-qro/",
            "确认墨西哥工厂开展橡胶混炼产品开发与生产",
        ),
        (
            "S8",
            "SABIC/KEMYA Al-Jubail工厂认证",
            "https://www.sabic.com/en/Images/Saudi%20Arabia-KEMYA%20RC-ISO14001_tcm1010-42778.pdf",
            "工厂认证范围明确包含EPDM制造",
        ),
        (
            "S9",
            "ExxonMobil Vistalon EPDM Storage Guide",
            "https://www.exxonmobilchemical.com/-/media/media-assets/media-library-assets/10/vistalon_storage_guide_iha_en.pdf",
            "列出美国Baton Rouge与沙特Al Jubail生产地点，说明品牌不唯一对应美国原产",
        ),
        (
            "S10",
            "Dow FUSABOND N302 Functional Polymer",
            "https://www.dow.com/en-us/pdp.fusabond-n302-functional-polymer.1892752z.html",
            "官方描述为酸酐改性乙烯弹性体和聚酰胺改性用途",
        ),
        (
            "S11",
            "Petro Rabigh官方聚合物产品页面",
            "https://www.petrorabigh.com/en/OurProducts/Polymer",
            "确认沙特本地生产乙烯丙烯橡胶等聚合产品",
        ),
    ]
    for code, title, url, note in sources:
        add_source(doc, bullet_id, code, title, url, note)

    add_paragraph(
        doc,
        "报告终止点：本报告在2026-08-13可得易迅记录和公开资料基础上形成。任何新增中国报关单、海关原产地审定、生产记录、完整提单链或司法材料均可能改变证据等级。",
        size=9.5,
        color=MUTED,
        italic=True,
        after=0,
    )

    doc.save(OUTPUT_FILE)
    print(str(OUTPUT_FILE))


if __name__ == "__main__":
    main()
