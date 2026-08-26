from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import load_workbook


def clean(value):
    if value is None:
        return None
    if hasattr(value, "isoformat"):
        try:
            return value.isoformat()
        except Exception:
            pass
    return str(value)


def main() -> None:
    source = Path(sys.argv[1])
    workbook = load_workbook(source, read_only=False, data_only=False)
    result = {
        "source": str(source),
        "size_bytes": source.stat().st_size,
        "sheets": [],
    }
    for sheet in workbook.worksheets:
        rows = []
        for row in sheet.iter_rows(min_row=1, max_row=min(12, sheet.max_row), values_only=True):
            rows.append([clean(value) for value in row])
        result["sheets"].append(
            {
                "title": sheet.title,
                "state": sheet.sheet_state,
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
                "freeze_panes": clean(sheet.freeze_panes),
                "auto_filter": sheet.auto_filter.ref,
                "merged_cells": [str(item) for item in list(sheet.merged_cells.ranges)[:30]],
                "hidden_rows": [index for index, dim in sheet.row_dimensions.items() if dim.hidden],
                "hidden_columns": [index for index, dim in sheet.column_dimensions.items() if dim.hidden],
                "first_rows": rows,
            }
        )
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
