from __future__ import annotations

from collections import Counter, defaultdict
from pathlib import Path
import csv
import json

from openpyxl import load_workbook


SOURCE = Path(r"D:\易迅数据\EPDM_400270.xlsx")
OUT = Path(r"D:\易迅数据\反倾销税深度分析报告\12_EPDM")
FIELDS = [
    "数据源", "进出口", "日期", "HS编码", "商品描述", "采购商", "供应商",
    "重量", "数量", "金额", "目的国/地区", "原产国/地区",
]


def load_json(name: str):
    return json.loads((OUT / name).read_text(encoding="utf-8"))


def csv_count(name: str) -> int:
    with (OUT / name).open("r", encoding="utf-8-sig", newline="") as handle:
        return sum(1 for _ in csv.DictReader(handle))


def main() -> None:
    wb = load_workbook(SOURCE, read_only=True, data_only=True)
    ws = wb["数据列表"]
    iterator = ws.iter_rows(values_only=True)
    assert list(next(iterator)) == FIELDS
    source_rows = [tuple(values) for values in iterator if any(v not in (None, "") for v in values)]
    assert len(source_rows) == 9356

    full = load_json("EPDM_易迅逐票标准化.json")
    assert len(full) == 9356
    assert csv_count("EPDM_易迅逐票标准化.csv") == 9356
    for expected_excel_row, (source_values, delivered) in enumerate(zip(source_rows, full), start=2):
        assert delivered["excel_row"] == expected_excel_row
        assert tuple(delivered[field] for field in FIELDS) == source_values
        for field in (
            "canonical_record_id", "scope_class", "scope_reason", "route_class",
            "evidence_grade", "evidence_reason", "decisive_data_gap",
        ):
            assert delivered[field] not in (None, "")

    canonical = [row for row in full if row["dedup_keep"] is True]
    duplicates = [row for row in full if row["is_exact_duplicate"] is True]
    assert len(canonical) == 9160
    assert len(duplicates) == 196

    groups = defaultdict(list)
    for row in full:
        groups[row["duplicate_group_id"]].append(row)
    duplicate_groups = [rows for rows in groups.values() if len(rows) > 1]
    assert len(duplicate_groups) == 138
    assert max(map(len, duplicate_groups)) == 11
    duplicate_audit = load_json("EPDM_全12字段重复审计.json")
    assert len(duplicate_audit) == sum(map(len, duplicate_groups)) == 334
    assert csv_count("EPDM_全12字段重复审计.csv") == 334

    expected_subsets = {
        "EPDM_对华55条": 55,
        "EPDM_受税来源直达中国6条": 6,
        "EPDM_HEXPOL_A腿418条": 418,
        "EPDM_HEXPOL_B腿4条": 4,
        "EPDM_第三国对华49条": 49,
    }
    for stem, expected in expected_subsets.items():
        assert len(load_json(f"{stem}.json")) == expected
        assert csv_count(f"{stem}.csv") == expected

    a_rows = load_json("EPDM_HEXPOL_A腿418条.json")
    assert Counter(row["原产国/地区"] for row in a_rows) == Counter(
        {"United States": 389, "South Korea": 19, "Netherlands": 6, "France": 4}
    )
    b_rows = load_json("EPDM_HEXPOL_B腿4条.json")
    assert {row["日期"] for row in b_rows} == {"2025-12-23"}
    assert {row["原产国/地区"] for row in b_rows} == {"Mexico"}
    assert all(row["route_class"] == "HEXPOL墨西哥→中国（B腿）" for row in b_rows)

    summary = load_json("EPDM_交付摘要.json")
    assert all(summary["self_checks"].values())
    assert summary["hexpol_relationship_metrics"]["a_same_normalized_description_as_b_before_b_unique_rows"] == 78
    assert sum(summary["scope_counts_exact_unique"].values()) == 9160
    assert summary["subsets"]["taxed_origin_direct_china_6"]["重量_raw_field_sum"] == "226987"
    assert summary["subsets"]["hexpol_a_418"]["数量_raw_field_sum"] == "2987820.95"
    assert summary["subsets"]["hexpol_b_4"]["数量_raw_field_sum"] == "618.2"

    result = {
        "source_rows_exactly_reconciled": len(full),
        "exact_unique": len(canonical),
        "duplicate_extra_rows": len(duplicates),
        "duplicate_audit_rows": len(duplicate_audit),
        "subsets": expected_subsets,
        "scope_counts": summary["scope_counts_exact_unique"],
        "all_checks_passed": True,
    }
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
