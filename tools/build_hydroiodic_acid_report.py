from __future__ import annotations

import csv
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from build_epdm_stage_report import (
    BLUE, GOLD, GREEN, MUTED, NAVY, PALE_BLUE, PALE_GOLD, PALE_GRAY, PALE_GREEN, PALE_RED, RED,
    add_callout, add_heading, add_list_item, add_page_number_field, add_paragraph, add_source,
    add_table, configure_document, create_numbering, set_run_font,
)


OUT = Path(r"D:\易迅数据\反倾销税深度分析报告\21_氢碘酸")
DOCX = OUT / "氢碘酸_反倾销税与第三国转运风险阶段深度分析报告.docx"


def caption(table, text="氢碘酸反倾销税与第三国转运风险阶段审计表"):
    cap = OxmlElement("w:tblCaption"); cap.set(qn("w:val"), text); table._tbl.tblPr.append(cap)
    if table.rows:
        trp = table.rows[0]._tr.get_or_add_trPr(); flag = OxmlElement("w:tblHeader"); flag.set(qn("w:val"), "true"); trp.append(flag)
        for c in table.rows[0].cells:
            for p in c.paragraphs: p.paragraph_format.keep_with_next = True


def table(doc, headers, rows, widths, **kwargs):
    t = add_table(doc, headers, rows, widths, **kwargs); caption(t); return t


def header_footer(doc):
    sec = doc.sections[0]
    h = sec.header.paragraphs[0]; h.clear(); h.alignment = WD_ALIGN_PARAGRAPH.RIGHT; h.paragraph_format.space_after = Pt(0)
    set_run_font(h.add_run("反倾销税风险穿透分析｜氢碘酸｜2026-08-13"), size=8.5, color=MUTED)
    f = sec.footer.paragraphs[0]; f.clear(); f.paragraph_format.space_after = Pt(0); add_page_number_field(f)
    if f.runs: f.runs[0].text = "氢碘酸阶段深度分析报告  |  "
    doc.core_properties.title = "氢碘酸反倾销税与第三国转运风险阶段深度分析报告"
    doc.core_properties.subject = "政策税率、易迅数据盘点、第三国绕道证据、实体与调证"
    doc.core_properties.author = "反倾销税风险分析项目"
    doc.core_properties.keywords = "氢碘酸; Hydroiodic Acid; 10034-85-2; 反倾销; 第三国转运"


def main():
    doc = Document(); configure_document(doc); header_footer(doc)
    bullets = create_numbering(doc, bullet=True)

    # standard_business_brief + memo_masthead
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(18); p.paragraph_format.space_after = Pt(4)
    set_run_font(p.add_run("第21项商品核查"), size=11, color=GOLD, bold=True)
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4)
    set_run_font(p.add_run("氢碘酸（HI）"), size=27, color=NAVY, bold=True)
    add_paragraph(doc, "反倾销税与第三国转运风险｜阶段深度分析报告", size=15.5, color=BLUE, bold=True, after=14)
    table(doc, ["基准日","措施状态","本地逐票记录","第三国闭环"], [["2026-08-13","续征至2029-10-15","0条","0条"]], [1700,3000,2200,2460], header_fill=PALE_BLUE, font_size=9.5, aligns=[WD_ALIGN_PARAGRAPH.CENTER]*4)
    add_callout(doc,"核心结论","现有D盘343个数据文件经内容级预筛与逐条复核，未发现可确认的氢碘酸贸易票据。114处字符串命中由113处“HI”缩写噪声和主清单第21行构成，均非易迅贸易记录。当前没有美国/日本→第三国→中国的A/B腿闭环，没有中国进口报关单和税款书，也就没有可确认的少缴税数量或税额。",fill=PALE_GOLD,title_color=GOLD)
    table(doc,["证据层","现有事实","等级","结论边界"],[
        ["政策","范围、税号、税率、期限已核","A","可据中国报关底单核税"],
        ["本地数据","114处命中全部复核，贸易票据0","排除/缺口","不等于易迅全库0条"],
        ["公开实体","美日受税生产/供应网络与印度真实产能","B/C","结构性线索与合法反证并存"],
        ["绕道执法","未检出中国处罚、判决、反规避或双提单闭环","未形成","负面检索不等于现实无风险"],
    ],[1500,3400,1300,3160],header_fill=PALE_GRAY,font_size=8.8)
    add_paragraph(doc,"完整性声明：已登录易迅页面并识别海关全球搜索界面；专项关键词查询连续超时，未取得可验证的总记录数和页数。因此本报告将网页全页查询列为待补，不把超时解释为“查无数据”。",size=9.2,color=RED,bold=True,after=0)

    doc.add_page_break()
    add_heading(doc,"1. 现行政策、产品范围和税率",1)
    add_paragraph(doc,"商务部2024年第43号公告决定，自2024年10月16日起，对原产于美国和日本的进口氢碘酸继续征收反倾销税5年。产品是碘化氢的水溶液，中文名称氢碘酸/碘化氢，英文名称Hydriodic Acid、Hydroiodic Acid或Hydrogen Iodide，化学式HI，CAS 10034-85-2。中国税则号为28111990；该税号下除氢碘酸以外的其他产品明确不在措施范围。现行细分查询可叠加2811199010，但最终以进口申报时税则版本为准。")
    table(doc,["原产地","适用企业","AD税率","AD+AD引致13% VAT增量"],[
        ["美国","Iofina Chemical及其他美国公司","123.4%","完税价格×139.442%"],
        ["日本","所有日本公司","41.1%","完税价格×46.443%"],
    ],[1500,3500,1700,2660],header_fill=PALE_BLUE,font_size=9)
    add_callout(doc,"税差口径","条件系数只包含反倾销税和反倾销税增加所引致的13%进口增值税增量，不包括正常关税与正常进口增值税。实际少缴额必须使用中国海关审定完税价格、法定原产地、生产商和已缴AD/VAT税单。",fill=PALE_GREEN,title_color=GREEN)
    add_heading(doc,"1.1 官方基线",2)
    for txt in [
        "美国：商务部复审裁定引用材料显示，2019-2022年产能均为2,000吨，2022年产量230吨；2023年上半年产能1,000吨、产量125吨，闲置产能比例87.5%。",
        "日本：2019-2022年产能均为1,150吨，2022年产量268吨；2023年上半年产能575吨、产量135吨，闲置产能比例76.52%。",
        "中国自日本进口：2019-2022年分别111.90、116.00、170.50和25.03吨；2023年上半年仅0.04吨。中国自美国进口同期一直很少，2019-2022年分别0.15、0.25、0.23和0.06吨。",
        "这些官方宏观数据证明受税来源仍有出口能力，不证明任何货物已绕道第三国。"
    ]: add_list_item(doc,txt,bullets)

    add_heading(doc,"2. 本地数据全盘审计",1)
    table(doc,["环节","结果","质量控制"],[
        ["文件盘点","343个XLSX/XLS/CSV/JSON","项目派生输出只登记，不重复展开计票"],
        ["原始候选文件","23个","按HS、全称、CAS、HI化学上下文预筛"],
        ["字符串命中","114处","保留源文件、工作表/JSON路径和原始文本"],
        ["逐条复核","114/114完成","113处HI噪声；1处为主清单"],
        ["贸易票据/中国B腿","0/0","不虚构重量、金额、企业或口岸"],
    ],[1800,2200,5360],header_fill=PALE_GRAY,font_size=9)
    add_paragraph(doc,"“HI”可能出现在企业名称、材料型号、地址、英文单词片段或其他商品长货描中。只有明确出现Hydriodic/Hydroiodic Acid、Hydrogen Iodide水溶液、CAS 10034-85-2，或税号28111990与上述化学标识共同出现，才能进入真正氢碘酸逐票池。仅税号也不足，因为公告明确同税号有范围外产品。")
    add_callout(doc,"数据结论","现有D盘能支持的结论是“未发现可确认氢碘酸贸易票据”，不是“易迅数据库没有氢碘酸贸易”，也不是“现实中不存在第三国转运”。",fill=PALE_RED,title_color=RED)

    add_heading(doc,"3. 第三国绕道证据评估",1)
    table(doc,["闭环要件","现有证据","缺失证据","判断"],[
        ["美/日→第三国A腿","0条逐票记录","出口申报、生产商、批号、UN1787、提单/柜号","未形成"],
        ["第三国→中国B腿","0条逐票记录","中国进口人、发货人、品名、浓度、数量、口岸","未形成"],
        ["同货匹配","无","30/60/90日、同浓度/包装/批号/柜号/主体","未形成"],
        ["原产改变","无第三国加工记录","反应/生产工艺、BOM、能耗、收率、CO底稿","无法判断"],
        ["少缴税款","无中国底单","完税价、原产国、生产商、AD/VAT缴款书","无法测算"],
    ],[1550,2300,3550,1960],header_fill=PALE_GRAY,font_size=8.6)
    add_paragraph(doc,"截至2026-08-13，公开定向检索未发现中国海关行政处罚、缉私通报、法院判决、商务部反规避裁定或公开双段提单，能够证明美国或日本氢碘酸经第三国换单、换原产地后进入中国逃避反倾销税。不得把“受税来源产能大”“第三国有分销商”或“品牌来自美国/日本”改写成既成绕道事实。")
    add_heading(doc,"3.1 原产地判断",2)
    add_paragraph(doc,"氢碘酸是碘化氢水溶液。第三国若只仓储、换桶、分装、贴标、开票或换单，通常不能仅凭这些操作主张原产地改变；但若第三国以碘等原料真实制造碘化氢并配制水溶液，可能构成当地生产。应按照中国非优惠原产地规则核最后实质性改变，并核实是否存在为规避反倾销而实施的加工。")

    add_heading(doc,"4. 实体与路线优先级",1)
    table(doc,["优先级","实体/路线","核查理由","合法替代解释"],[
        ["A","Iofina Chemical（美国）→第三国→中国","公告列名，123.4%税率；美国官网确认碘及卤素衍生物生产","第三国可能仅为正常销售/仓储，必须看生产批次"],
        ["A","日本HI生产/供应网络→第三国→中国","日本统一41.1%；Kishida、Junsei、Wako、Kanto、Nippoh、Godo等公开有HI产品","部分主体可能是经销/试剂分装，未必是生产厂"],
        ["B+","印度HI制造实体→中国","Calibre、Samrat、Eskay、Infinium、M.M. Arochem公开有50%-57%HI产品/COA/产能","印度存在真实制造能力，不能见印度发货即推定美日绕道"],
        ["B","新加坡/阿联酋/香港/韩国/越南/马来西亚分销节点","适合做美日A腿与中国B腿主体交叉","开票国、转运港、品牌国均不等于原产国"],
    ],[900,2600,3300,2560],header_fill=PALE_BLUE,font_size=8.45)
    add_heading(doc,"4.1 具体可调实体",2)
    for txt in [
        "美国：Iofina Chemical, Inc.；重点调其Kentucky制造批次、第三国经销合同、发票、付款受益人和对华最终客户。",
        "日本：Nippoh Chemical、Godo Shigen、Kishida Chemical、Junsei Chemical、Fujifilm Wako Pure Chemical、Kanto Chemical、Tokyo Chemical Industry；先区分制造商、品牌商与试剂经销商。",
        "印度：Calibre Chemicals、Samrat Pharmachem、Eskay Iodine、Infinium Pharmachem、M.M. Arochem；调工厂许可、原料碘采购、反应/配制记录、产能、批号和CO原产资格。",
        "中国端：凡28111990/2811199010申报为印度、韩国、新加坡、越南、马来西亚、阿联酋等非受税来源的HI，优先核境外生产商栏、COA/厂号和非优惠原产地证，而不是只看发货国。",
    ]: add_list_item(doc,txt,bullets)

    add_heading(doc,"5. 易迅专项查询方案",1)
    add_paragraph(doc,"已生成三组最简条件，后续页面恢复后应切换200条/页并读取全部页。每个查询结果都需逐票做产品范围、原产地和A/B腿判断，不能发现一条异常后停止。")
    table(doc,["查询","核心条件","完成标准"],[
        ["HI-Q1","中国进口；HS281119/28111990；全称+CAS；全部来源","全页逐票排除同税号其他酸，确认中国收货人/发货人/原产字段"],
        ["HI-Q2","美国/日本→重点第三国；同关键词+受税实体","记录浓度、包装、批号、UN1787、提单/柜号和第三国收货人"],
        ["HI-Q3","第三国→中国；同关键词+A腿实体/批号","按30/60/90日及同货特征闭合，并回查中国原产申报和税单"],
    ],[1100,5250,3010],header_fill=PALE_GRAY,font_size=8.7)
    add_callout(doc,"网页状态","Chrome中的易迅账号仍显示已登录，页面可识别为海关全球搜索；但本轮专项关键词提交和结果读取连续两次超过页面控制时限，未获得可核验的总数、页数或逐票列表。因此网页专项检索保持待补状态。",fill=PALE_RED,title_color=RED)

    add_heading(doc,"6. 必调单证与条件税差",1)
    for txt in [
        "中国报关：品名、CAS、浓度、10位税号、原产国、启运国、境外发货人、境内进口人、消费使用单位、申报企业、口岸、净重、完税价格、生产商、AD税率及税款书。",
        "物流单证：合同、商业发票、原厂发票、付款受益人、提单、柜号、封志、船名航次、中转港、UN1787危险品申报、包装唛头和照片。",
        "产品穿透：COA/SDS、批号、浓度（常见55%-58%）、生产日期、plant code、原料碘来源、反应/配制工单、设备、能耗、产量/收率和库存。",
        "原产资格：第三国非优惠原产地证及申请底稿、生产成本/BOM、原料进口申报、签证机构核查；仅分装或换标不得当作实质制造。",
    ]: add_list_item(doc,txt,bullets)
    add_callout(doc,"条件测算","若中国海关完税价格为V，实际确认美国原产且未缴AD，则AD及其引致VAT增量为V×139.442%；日本原产为V×46.443%。在没有中国完税价格和缴款书前，不应使用境外出口金额字段直接写实际欠税额。",fill=PALE_GREEN,title_color=GREEN)

    add_heading(doc,"7. 阶段结论",1)
    add_paragraph(doc,"当前证据评级为C/未形成具体绕道证据。本地逐票池为0，公开源没有中国执法或双段提单闭环，因此不能指向具体中国进口企业、报关行或口岸，也不能写涉及数量和逃税额。结构性风险仍需保持中高优先级：美国税率123.4%，与非受税第三国价差极大；日本有大量闲置能力；印度等第三国又确有真实生产，导致“合法第三国产能”和“受税成品分装转售”必须靠批号、生产记录和中国原产申报区分。最有效的下一步是完成HI-Q1至HI-Q3全页查询，并对每一条中国B腿调底单。")

    add_heading(doc,"来源",1)
    sources = [
        ("S1","商务部公告2018年第80号：氢碘酸反倾销终裁","https://policy.mofcom.gov.cn/claw/clawContent.shtml?id=66459","产品范围、税号、初始税率和计税方法。"),
        ("S2","商务部公告2024年第43号及期终复审裁定","https://cacs.mofcom.gov.cn/cacscms/article/jkdc?articleId=182123&type=1","续征期限、现行税率、官方产能和进口基线。"),
        ("S3","Iofina官网：美国碘和卤素衍生物生产","https://iofina.com/iofina-chemical/","列名美国实体的生产基地与业务边界。"),
        ("S4","Nippoh Chemical氢碘酸产品页","https://www.npckk.co.jp/cms/product/106/","日本氢碘酸产品与CAS线索。"),
        ("S5","Godo Shigen氢碘酸水溶液产品页","https://www.godoshigen.co.jp/products/%E3%83%A8%E3%82%A6%E5%8C%96%E6%B0%B4%E7%B4%A0%E9%85%B8%E6%B0%B4%E6%BA%B6%E6%B6%B2/","日本供应/制造核查线索。"),
        ("S6","Samrat Pharmachem氢碘酸COA","https://www.samratpharmachem.com/wp-content/uploads/2025/05/11.-Hydroiodic-Acid.pdf","印度工厂、CAS、HS、57%浓度和批次反证。"),
        ("S7","Eskay Iodine氢碘酸57%产品页","https://eskayiodine.com/hydriodic-acid-57/","印度产品、20kg包装和应用。"),
        ("S8","中国进出口货物原产地条例","https://xzfg.moj.gov.cn/front/law/detail?LawID=1523&Query=%E8%B4%A7%E7%89%A9%E5%8E%9F%E4%BA%A7%E5%9C%B0","反倾销非优惠原产地及最后实质性改变规则。"),
    ]
    for code,title,url,note in sources: add_source(doc,bullets,code,title,url,note)
    add_paragraph(doc,"本报告是基于现有本地文件、易迅页面可验证状态与公开资料的风险筛查，不替代海关归类、原产地核定、税款稽核或司法认定。",size=8.8,color=MUTED,after=0)

    doc.save(DOCX)
    assert DOCX.exists() and DOCX.stat().st_size > 30000
    print(DOCX)


if __name__ == "__main__":
    main()
