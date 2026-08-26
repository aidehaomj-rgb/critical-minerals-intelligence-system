from __future__ import annotations

import collections
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook


PATTERNS = {
    "cas_108_39_4": re.compile(r"(?<!\d)108\D{0,6}39\D{0,6}4(?!\d)", re.I),
    "meta_cresol": re.compile(r"\b(?:META[\s_-]*CRESOL|METACRESOL|M[\s_-]+CRESOL|M[\s_-]*KRESOL|METAKRESOL)\b", re.I),
    "3_methylphenol": re.compile(r"\b3\s*[-_]?\s*METHYL\s*[-_]?\s*PHENOL\b", re.I),
    "meta_any": re.compile(r"META|M[-\s]CRES|M[-\s]KRES|3[-\s]?METHYL|108[-\s/]?39", re.I),
}


def norm(value) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def main() -> None:
    source = Path(sys.argv[1])
    output_path = Path(sys.argv[2]) if len(sys.argv) > 2 else None
    sheet = load_workbook(source, read_only=True, data_only=True).active
    headers = [norm(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    index = {name: position for position, name in enumerate(headers)}
    rows = []
    distributions = {key: collections.Counter() for key in ("数据源", "进出口", "HS编码", "目的国/地区", "原产国/地区")}
    dates = []
    for excel_row, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        rec = {name: values[pos] for name, pos in index.items()}
        desc = norm(rec.get("商品描述"))
        matches = [name for name, pattern in PATTERNS.items() if pattern.search(desc)]
        for key in distributions:
            distributions[key][norm(rec.get(key))] += 1
        if rec.get("日期") is not None:
            dates.append(str(rec["日期"])[:10])
        if matches:
            rows.append(
                {
                    "excel_row": excel_row,
                    "matches": matches,
                    "数据源": norm(rec.get("数据源")),
                    "进出口": norm(rec.get("进出口")),
                    "日期": str(rec.get("日期"))[:10],
                    "HS编码": norm(rec.get("HS编码")),
                    "商品描述": desc,
                    "采购商": norm(rec.get("采购商")),
                    "供应商": norm(rec.get("供应商")),
                    "重量": norm(rec.get("重量")),
                    "数量": norm(rec.get("数量")),
                    "金额": norm(rec.get("金额")),
                    "目的国": norm(rec.get("目的国/地区")),
                    "原产国": norm(rec.get("原产国/地区")),
                }
            )
    by_pattern = {name: sum(name in row["matches"] for row in rows) for name in PATTERNS}
    distinct_descriptions = collections.Counter(row["商品描述"] for row in rows)
    output = {
        "headers": headers,
        "raw_rows": sheet.max_row - 1,
        "date_min": min(dates),
        "date_max": max(dates),
        "by_pattern": by_pattern,
        "distribution_top": {key: value.most_common(30) for key, value in distributions.items()},
        "matched_raw": len(rows),
        "matched_distinct_descriptions": len(distinct_descriptions),
        "distinct_descriptions": [
            {"count": count, "description": description}
            for description, count in distinct_descriptions.most_common()
        ],
        "matched_rows": rows,
    }
    rendered = json.dumps(output, ensure_ascii=False, indent=2, default=str)
    if output_path:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text(rendered, encoding="utf-8")
        print(
            json.dumps(
                {
                    key: output[key]
                    for key in (
                        "headers",
                        "raw_rows",
                        "date_min",
                        "date_max",
                        "by_pattern",
                        "matched_raw",
                        "matched_distinct_descriptions",
                    )
                },
                ensure_ascii=False,
                indent=2,
            )
        )
    else:
        print(rendered)


if __name__ == "__main__":
    main()
