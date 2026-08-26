from __future__ import annotations

import json
from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(r"D:\易迅数据\反倾销税深度分析报告\10_聚苯醚")
AUDIT = json.loads((OUT / "聚苯醚_全量阶段审计.json").read_text(encoding="utf-8"))
RECORDS = json.loads((OUT / "聚苯醚_易迅逐票标准化.json").read_text(encoding="utf-8"))
REPORT = OUT / "聚苯醚_反倾销税与第三国转运风险阶段审计报告.docx"

NAVY, BLUE, INK = "17324D", "1F4E78", "102A43"
MUTED, BORDER = "66788A", "CBD5E1"
LIGHT_BLUE, GOLD, RED_FILL, GREEN = "EAF2F8", "FFF4CC", "FCE8E6", "E6F4EA"


def rgb(value):
    return RGBColor.from_string(value)


def fmt(value, digits=2):
    if value in (None, ""):
        return "—"
    value = float(value)
    return f"{value:,.{digits}f}" if digits else f"{value:,.0f}"


def font(run, size=9.5, color="222222", bold=False, italic=False):
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    run.bold = bold
    run.italic = italic
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    fonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    trpr.append(node)


def borders(table):
    tblpr = table._tbl.tblPr
    box = tblpr.find(qn("w:tblBorders"))
    if box is None:
        box = OxmlElement("w:tblBorders")
        tblpr.append(box)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        item = OxmlElement(f"w:{edge}")
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), "4")
        item.set(qn("w:color"), BORDER)
        box.append(item)


def cell_text(cell, value, bold=False, size=8.3, color="222222"):
    cell.text = "" if value is None else str(value)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    for p in cell.paragraphs:
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.03
        for r in p.runs:
            font(r, size=size, color=color, bold=bold)


def table(doc, headers, rows, widths=None, size=8.3, header_fill=LIGHT_BLUE):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        cell_text(t.cell(0, i), h, bold=True, size=size, color=INK)
        shade(t.cell(0, i), header_fill)
    repeat_header(t.rows[0])
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cell_text(cells[i], value, size=size)
    borders(t)
    t.autofit = False
    if widths:
        for row in t.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Inches(width)
    return t


def para(doc, text, size=9.7, bold=False, color="222222", italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.1
    r = p.add_run(text)
    font(r, size=size, color=color, bold=bold, italic=italic)
    return p


def bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.06
    r = p.add_run(text)
    font(r, size=9.35)
    return p


def callout(doc, title, text, fill=GOLD, title_color=INK):
    t = doc.add_table(rows=1, cols=1)
    c = t.cell(0, 0)
    shade(c, fill)
    p = c.paragraphs[0]
    r = p.add_run(title)
    font(r, size=10, color=title_color, bold=True)
    p2 = c.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.08
    r2 = p2.add_run(text)
    font(r2, size=9.25)
    repeat_header(t.rows[0])
    borders(t)
    return t


def hyperlink(p, label, url):
    rid = p.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), BLUE); rpr.append(color)
    underline = OxmlElement("w:u"); underline.set(qn("w:val"), "single"); rpr.append(underline)
    run.append(rpr)
    text = OxmlElement("w:t"); text.text = label; run.append(text)
    link.append(run); p._p.append(link)


def source(doc, label, url, note=""):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    hyperlink(p, label, url)
    if note:
        r = p.add_run(f"：{note}")
        font(r, size=8.7, color=MUTED)


def page_number(p):
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end]); font(run, size=8, color=MUTED)


def heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    return p


def uniq(rows):
    seen, out = set(), []
    for r in rows:
        if r["visible_signature"] not in seen:
            seen.add(r["visible_signature"]); out.append(r)
    return out


doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.68); sec.bottom_margin = Inches(0.68)
sec.left_margin = Inches(0.78); sec.right_margin = Inches(0.78)
sec.page_width = Inches(8.27); sec.page_height = Inches(11.69)
styles = doc.styles
styles["Normal"].font.name = "Calibri"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
for name, size, color in (("Title", 25, NAVY), ("Heading 1", 15.5, NAVY), ("Heading 2", 12.4, BLUE), ("Heading 3", 10.8, INK)):
    st = styles[name]; st.font.name = "Calibri"; st.font.size = Pt(size); st.font.color.rgb = rgb(color); st.font.bold = True
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
header = sec.header.paragraphs[0]; header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
font(header.add_run("反倾销税风险穿透分析｜聚苯醚（PPE/PPO）｜2026-08-13"), size=8, color=MUTED)
page_number(sec.footer.paragraphs[0])

# Cover
p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(42)
font(p.add_run("反倾销税风险穿透分析"), size=12, color=BLUE, bold=True)
p = doc.add_paragraph(style="Title"); p.add_run("聚苯醚（PPE/PPO）")
p = doc.add_paragraph(); font(p.add_run("第三国转运、原产地与少缴税风险阶段审计报告"), size=15, color=INK, bold=True)
para(doc, "基于易迅两份已下载工作簿10,347行逐条审计，并结合中国现行贸易救济措施、境外制造能力与企业关联公开资料", size=10.5, color=MUTED)
doc.add_paragraph()
callout(doc, "核心判断", "发现HPP Mexico两组同日同实体同货描、Motores相邻1日同实体同货描等可进一步调证的具体链路；但尚无同一柜号、批号、PO或中国报关原产地闭环。HPP Mexico虽有官方确认的热塑性配混工厂，但若美国投入料与墨西哥产出均归3907，真实配混也未必满足非优惠原产地四位税目改变。54条再生PPO/PPE候选应先做成分和原产地鉴别。", RED_FILL, "8B1E1E")
table(doc, ["数据范围", "原始/去重", "最具体线索", "结论等级"], [["中国进口池＋美国出口池", "10,347行；范围候选665/654", "HPP、Motores、Flextronics；54条再生料", "B+调证线索；0条违法闭环"]], [1.55,1.65,2.8,1.25], 8.6)
para(doc, "报告性质：执法风险筛查与调证线索，不替代海关原产地认定、税款核定或违法结论。数据截止以两份工作簿实际覆盖为准。", size=8.9, color=MUTED, italic=True)
doc.add_page_break()

# 1 Executive summary
heading(doc, "一、执行摘要", 1)
table(doc, ["维度", "全量结果", "风险解释"], [
    ["逐条覆盖", "10,347原始行；10,157可见字段唯一；10,142商业签名唯一", "每条均留判定字段，未因发现异常提前停止"],
    ["范围候选", "665原始/654唯一；明确范围600唯一、再生料待成分54唯一", "其余9,682行为PPS、PEG/PEO、聚醚多元醇等范围外或信息不足"],
    ["中国端", "68原始/67唯一；越南42、印尼8、菲律宾8、墨西哥7、原产字段空2", "证明对华路线，不证明法定原产或未缴税"],
    ["美国来源第三国供给", "597原始/587唯一；墨西哥489、越南66、印度17等", "平台来源字段和工作簿名不能自动替代制造原产"],
    ["机械牌号闭合", "0条", "GTX973、X552H/Z552H、PCN2615等均无同牌号A腿"],
    ["实体＋货描链", "4条中国端记录、134个历史A腿候选", "HPP三条、Motores一条；具体性提高但仍非同批闭环"],
], [1.35,3.0,3.0], 8.4)

heading(doc, "优先级结论", 2)
bullet(doc, "B+｜HPP Mexico：2026-04-09和2026-02-03均存在同日美国来源A腿与对华B腿。4月9日B腿货描将FENILENO高度疑似误录为FELINO，数量字段25；同日A腿包括通用货描18,360及GFN3F牌号40。2月3日A腿1,180与B腿680货描逐字相同。优先调取两日提单、批号、配混工单和中国报关原产地。")
bullet(doc, "B｜Motores：2026-02-27美国来源字段A腿29,157/180,773.38，次日出现Motores→中国100/6,199.99，同实体、同HS、货描逐字相同；但数量相差约291倍、A腿供应商为空、中国收货人缺失。")
bullet(doc, "C+/B-｜Flextronics：同集团两个墨西哥实体、同PPE品类和15日窗口成立，法律实体、规格、数量均不闭合。")
bullet(doc, "C｜54条再生PPO/PPE：候选物理量868,069kg（越南594,069、菲律宾243,150、印尼30,850，单位仍须原单确认）；无美国→主要再生料供应商或同再生货描A腿。")
bullet(doc, "B-反证｜印尼/泰国：Asahi Kasei确认Nippisun Indonesia受托生产XYRON；SABIC公开资料显示泰国Rayong有NORYL产能，且越南一票明确#&TH。")
callout(doc, "最终定性", "公开资料和易迅现有字段尚不足以证明美国原产聚苯醚经第三国换证入华，也不能证明少缴反倾销税。HPP与Motores已达到“应立即调单”的具体程度；只有中国进口报关原产地、反倾销税缴款书与A/B两程同批证据闭合后，才能升级。", GREEN)

# 2 policy
heading(doc, "二、措施范围、税率与少缴税口径", 1)
para(doc, "商务部公告2022年第1号自2022年1月7日起对原产于美国的聚苯醚征收反倾销税，期限5年，预计至2027年1月6日。范围包括PPE/PPO及改性产品，以及与聚苯乙烯、氢化苯乙烯—丁二烯嵌段共聚物、尼龙、填料或添加剂混合的组合物；税则号为39072990，该税号项下其他产品不在措施范围。")
table(doc, ["税档", "AD税率", "AD引致VAT差额", "合计条件增量", "适用前提"], [
    ["SHPP US LLC", "17.3%", "AD×13%", "完税价×19.549%", "确认美国原产且生产商为SHPP US"],
    ["其他美国公司", "48.6%", "AD×13%", "完税价×54.918%", "确认美国原产且不适用列名税档"],
], [1.55,1.0,1.4,1.4,2.0], 8.4)
para(doc, "商务部公告2022年第20号规定SHPP US LLC承继SABIC Innovative Plastics US LLC的17.3%税率；继续以旧公司名报关时适用其他美国公司48.6%。公告2022年第2号认定补贴为微量并终止反补贴调查，因此本案不得再加征反补贴税。", size=9.4)
callout(doc, "税额边界", "平台金额字段没有统一币种，且不是中国海关审定完税价格。只有“实际美国原产＋中国申报第三国原产＋未按适用企业税档缴AD＋取得中国完税价格”同时成立，才能计算少缴情景。本报告不把平台金额直接写成实际欠税。", GOLD)
para(doc, "截至2026年8月13日，商务部实施中措施表仍列本案预计终止日为2027年1月6日；本次定向检索未发现期终复审立案公告。此状态具有时效性，后续应持续复核。", size=9.3)

# 3 method
heading(doc, "三、数据完整性与逐条分析方法", 1)
table(doc, ["文件", "原始行", "日期覆盖", "审计方法"], [
    ["聚苯醚中国进口.xlsx", "467", "2025-09-03—2026-07-31", "逐条商品范围、实体、来源字段、尾标、数量金额和B腿判断"],
    ["聚苯醚美国出口.xlsx", "9,880", "2025-09-01—2026-07-29", "逐条范围、目的国、A腿供给、主体/货描/日期匹配"],
], [2.0,0.8,1.9,2.9], 8.3)
para(doc, "筛查同时覆盖英文PPE/PPO/POLYPHENYLENE ETHER/OXIDE、NORYL、XYRON、MPPE，以及西语POLIOXIDO DE FENILENO、ETER DE POLIFENILENO等。PPS、PEG/PEO和其他聚醚逐条排除。完全同值记录不直接删除：报告并列原始数和精确去重数，逐票台账保留全部行。")
table(doc, ["层级", "数量", "说明"], [
    ["全部原始", "10,347", "不遗漏任何已下载记录"],
    ["范围候选原始/唯一", "665 / 654", "含明确化学范围、144条西语/实验室变体、1条WYRON品牌误拼补漏及再生料待成分"],
    ["中国端原始/唯一", "68 / 67", "一组菲律宾53,200kg/2,120包完全同值重复"],
    ["美国来源原始/唯一", "597 / 587", "第二轮假阴性审计补入144条西语连写/倒装/实验室品名及1条WYRON≈XYRON 540Z"],
], [2.3,1.3,4.0], 8.5)

# 4 China B legs
heading(doc, "四、中国端67条候选的结构", 1)
origin_rows = AUDIT["china_b_legs"]["origin_split"]
table(doc, ["平台来源/原产字段", "唯一条数", "重量字段合计", "数量字段合计", "金额字段合计", "解释"], [
    [x["platform_origin"], x["records"], fmt(x["weight_field_sum"]), fmt(x["quantity_field_sum"]), fmt(x["amount_field_sum"]), "字段单位/币种按数据源不同，禁止跨国直接相加"]
    for x in origin_rows
], [1.35,0.75,1.2,1.25,1.45,2.0], 7.9)
para(doc, "越南42条包括39条再生料及3条明确混配/改性PPE；印尼8条包括5条再生料与3条Nippisun XYRON；墨西哥7条均为明确或高度疑似PPO/PPE/NORYL；菲律宾及原产空白合计10条均为GBW Acritech再生料。", size=9.3)
table(doc, ["日期/路线", "货物", "字段量", "境外→中国主体", "研判"], [
    ["2026-04-09 墨→中", "改性PPO/PPE（FELINO疑为FENILENO）", "25 / 304.98", "HPP Mexico→SHPP Shanghai", "B+同日链；优先核GFN3F 40单位A腿"],
    ["2026-02-03 墨→中", "改性PPO/PPE组合物", "680 / 8,296.02", "HPP Mexico→SHPP Shanghai", "B+同日同实体同货描链"],
    ["2026-01-19 墨→中", "同上", "607 / 10,045.84", "HPP Mexico→SHPP Shanghai", "B，历史同实体同货描"],
    ["2026-02-28 墨→中", "改性聚苯醚", "100 / 6,199.99", "Motores→字段同名主体", "B，相邻1日链"],
    ["2026-06-16 墨→中", "改性PPO/PPE", "24.54 / 3,446.85", "Flextronics Plastics→买方空", "C+/B-关联实体窗口"],
    ["2026-04-13 墨→中", "NORYL GTX973", "1,800 / 1,800.01", "TB&C→Suzhou Hillion", "B，无同牌号A腿"],
    ["2026-05-27 墨→中", "PPO初级形状", "7,440 / 3,348.02", "Stache Lion→Shining Way", "B，制造地/单位待核"],
    ["2026-01-15 越→中", "NORYL PCN2615", "1,000 / 101,689,819", "Nagase Vietnam→上海华昌", "货描#&TH，泰国产能反证"],
    ["2026-01-12 越→中", "MPPE/GF_FR EFE-8108", "150 / 7,824,300", "越南记录", "带CAS配比、#&KR，原产待核"],
    ["2025-11-19 越→中", "PA6-PPE40D（50/50）", "450 / 53,420,179.14", "越南记录", "#&CN，可能中国料返运"],
], [1.15,1.45,1.25,2.2,2.0], 7.6)

# 5 exact chains
heading(doc, "五、四条最具体A/B链及合法替代解释", 1)
heading(doc, "5.1 HPP Mexico：两组B+同日链", 2)
para(doc, "美国来源字段A腿中，HIGH PERFORMANCE PLASTICS MANUFACTURING MEXICO S DE RL DE CV精确去重199条，数量字段2,213,683、金额字段9,884,456.18；其中181条供应商为SHPP US LLC，数量字段1,957,343、金额字段8,991,591.18。中国端共有3条HPP Mexico→SHPP Shanghai候选；在将4月9日B腿的FELINO按高度疑似录入错误规范为FENILENO后，3条B腿对应130个HPP历史候选对，另有Motores 4个，共134个。")
table(doc, ["腿段", "日期", "主体", "数量/金额字段", "货描匹配", "差距"], [
    ["A", "2026-04-09", "SHPP US→HPP Mexico", "18,360 / 84,086.60", "规范化后同通用货描", "同日；数量差18,335"],
    ["A牌号", "2026-04-09", "SHPP US→HPP Mexico", "GFN3F 40 / 557.32", "同品类、带牌号", "同日；数量差15"],
    ["B", "2026-04-09", "HPP Mexico→SHPP Shanghai", "25 / 304.98", "FELINO疑为FENILENO", "—"],
    ["A", "2026-02-03", "SHPP US→HPP Mexico", "1,180 / 5,745.77", "逐字相同", "同日；数量差500"],
    ["B", "2026-02-03", "HPP Mexico→SHPP Shanghai", "680 / 8,296.02", "逐字相同", "—"],
    ["A近量", "2025-12-16", "SHPP US→HPP Mexico", "600 / 2,805.98", "逐字相同", "距1月19日34天；数量差7"],
    ["B", "2026-01-19", "HPP Mexico→SHPP Shanghai", "607 / 10,045.84", "逐字相同", "—"],
], [0.8,1.1,2.35,1.35,1.05,1.35], 8.0)
callout(doc, "真实配混不等于当然取得墨西哥原产", "SABIC证书确认HPP Mexico为真实热塑性配混工厂，这是反对“纯换包装”的强证据；但现行非优惠原产地规则原则上以四位税目改变为基本标准，39章特定加工清单未见3907。若美国投入料和墨西哥产出均归3907，真实配混仍可能不构成原产地改变。必须调配方、投入/产出税号、工单、批号和中国原产申报，且对以规避反倾销为目的的加工海关可不予考虑。", GOLD)
para(doc, "三条HPP B腿的数量字段合计1,312，原始金额字段合计18,646.84。若仅作机械敏感性分析，并假设该金额字段与中国海关完税价格同币种同口径、货物实际美国原产且中国端完全未缴AD，则SHPP US 17.3%档的AD及其引致13%VAT差额为3,645.27；其他美国公司48.6%档为10,240.47。该计算不是应补税额：币种、成交条件、关联交易调整和中国完税价格均未核。", size=8.9, color=MUTED, italic=True)
para(doc, "另据NSF认证库，GFN3F并非只能由美国站点制造，上海及泰国Rayong相关SHPP站点亦有该牌号认证记录。因此4月9日A腿GFN3F 40与B腿25的价格/数量指纹只能用于调批号，不能仅凭牌号认定美国原产。", size=9.2)

heading(doc, "5.2 Motores：B级相邻1日链", 2)
para(doc, "A腿同一实体8条/数量字段147,814/金额字段916,446.75，货描及HS一致、供应商字段均为空。2026年2月27日A腿29,157/180,773.38，次日B腿100/6,199.99，B记录中收发货人字段均为Motores。该链满足实体、货描和相邻日期，但数量相差约291倍，无法证明同货。")
heading(doc, "5.3 Flextronics：C+/B-集团品类链", 2)
para(doc, "美国来源字段到Flextronics Technologies San Luis的同类PPE/PS共34条/数量字段1,612,907；2026年6月1日最后一批34,880后，2026年6月16日Flextronics Plastics向中国发24.54。SEC历史子公司清单证明两实体在2018年同属Flex集团；但法律实体、货描规格和数量不一致，且不能无证外推2026年持股关系，故仅作集团品类窗口线索。")

# 6 recycled
heading(doc, "六、54条再生PPO/PPE候选", 1)
table(doc, ["路线", "唯一条数", "候选物理量", "主要境外主体", "主要中国主体", "当前等级"], [
    ["越南→中国", "39", "594,069（数量字段；描述为kg/袋）", "VIỆT-CAM、Huangying、Thủy Anh、Houseware等", "HETE、HK Haolang、Ningbo Langcong等", "C"],
    ["菲律宾→中国", "10", "243,150kg重量字段", "GBW Acritech Plastic Products", "Xiamen Xiefuhao等", "C"],
    ["印尼→中国", "5", "30,850kg重量字段", "PT Wahana、PT Alam", "NINGBO ENJOYLIFE", "C"],
], [1.0,0.75,1.55,2.2,1.75,0.55], 7.8)
para(doc, "54条均明确写PPO或PPE，并非仅有泛称recycled plastic；但“再生PPO/PPE”不自动证明聚苯醚比例达到措施范围，也不自动证明原料来自美国。美国出口表未发现美国→上述主要再生供应商，亦无同“PPO/PPE recycled”货描A腿。")
bullet(doc, "先做产品范围：FTIR确认PPE/PPO骨架，DSC/TGA及灰分/填料分析，配方与COA核含量；核中国10位税号和申报品名。")
bullet(doc, "再做原产地：废料来源、进口申报、采购发票、批次库存、分选/清洗/造粒工序、能耗、设备、员工、投入产出与非优惠CO申请底稿。")
bullet(doc, "若仅分拣、换包、贴标或简单造粒，原产地改变基础弱；若有真实配混/改性，还须结合四位税目变化、从价比例及主要工序逐批认定。")

# 7 counterevidence and origin
heading(doc, "七、第三国产能反证与原产地判断", 1)
table(doc, ["实体/地区", "公开事实", "对本案影响", "仍需核查"], [
    ["P.T. Nippisun Indonesia", "Asahi Kasei官方确认其受托生产XYRON改性PPE", "3条印尼XYRON更可能是合法印尼配混", "CO、配方、工单、增值率"],
    ["SABIC Rayong Thailand", "官方资料列有NORYL产能；越南PCN2615货描#&TH", "支持泰国产经越南发华解释", "泰国CO、批号、越南再出口与中国申报"],
    ["HPP Mexico", "官方证书明确热塑性配混工厂", "同日链也可能是库存/配混后出口", "主要工序、配方、税则与≥30%增值计算"],
    ["Flex Mexico", "官方披露墨西哥塑料运营；历史SEC清单列两实体", "集团内制造是替代解释", "2026关联关系与具体货物工单"],
], [1.6,2.45,2.25,1.55], 8.0)
para(doc, "依据《进出口货物原产地条例》，多国加工以最后完成实质性改变的国家为原产地；简单仓储、装卸、销售包装等不赋予新原产地；为规避反倾销等措施而进行的加工，海关在原产地判断中可以不予考虑。现行实质性改变规则以四位税目改变为基本标准，特定货物再适用清单中的制造工序/从价标准；初核39章清单未见3907专门标准，故HPP投入和产出若均为3907是首要风险点。最终仍须以海关适用的现行官方附件和逐票资料复核。")

# 8 investigation plan
heading(doc, "八、调证方案与税款核验", 1)
table(doc, ["优先级", "对象", "立即调取", "升级标准"], [
    ["1", "HPP 2026-04-09及02-03同日链", "A/B两端申报、提单、柜号、批号、PO；配混工单/BOM；投入/产出四位税目；中国报关单、CO、税款书", "同批标识＋投入/产出均3907＋中国申报非美原产＋未缴AD"],
    ["2", "Motores 2026-02-27/28", "A腿供应商、两端提单、库存流水、样品/销售发票、中国收货人", "小批B腿可回溯至具体A腿批次"],
    ["3", "54条再生料", "FTIR/DSC/TGA、COA、废料来源、工单、能耗、投入产出、原产申请底稿", "确认在范围且加工不足以改变原产"],
    ["4", "GTX973、越南#&TH/#&KR/#&CN", "制造商声明、批号、原产证、第三国进口/再出口、中国申报", "牌号/批次和法定原产冲突"],
], [0.7,1.65,3.2,2.3], 7.9)
heading(doc, "中国端最低证据包", 2)
bullet(doc, "进口报关单：10位税号、法定原产国、境外生产商、贸易国、成交方式、完税价格及币种。")
bullet(doc, "非优惠原产地证、制造商声明、COA/SDS/TDS、批号和生产日期。")
bullet(doc, "反倾销税缴款书/保证金、进口增值税缴款书；核适用17.3%或48.6%税档。")
bullet(doc, "美国A腿、第三国进口/加工/再出口和中国进口四端单证；柜号、封志、批号、PO、净重及日期平衡。")
callout(doc, "条件税差公式", "若最终确认实际美国原产且中国端未缴税：SHPP US税档的AD及其引致VAT差额＝中国海关完税价格×19.549%；其他美国公司税档＝完税价格×54.918%。如另有优惠原产地虚假申报，普通关税差额应另案计算，不能混入本报告AD口径。", GOLD)

# 9 conclusion
heading(doc, "九、结论", 1)
para(doc, "本轮完整读取并逐条分析了现有两份数据的10,347行。最终形成67条中国端范围候选、587条美国来源字段第三国供给记录、4条同实体同货描具体链和134个历史A腿候选对。与初步只看牌号的8条窄口径相比，本报告补入了54条再生PPO/PPE、144条西语连写/倒装/实验室变体、1条WYRON≈XYRON 540Z品牌误拼，以及MPPE、PA6-PPE、FELINO≈FENILENO等模式。")
para(doc, "证据上，HPP Mexico两组同日链是当前最具体线索，Motores相邻日链次之，Flextronics为关联实体窗口线索；但均没有同批/同柜/同PO与中国申报闭环。HPP、Nippisun和SABIC泰国存在真实制造能力，却不当然意味着非优惠原产地已经改变。因此最稳妥结论是：存在应优先调证的第三国供应链风险，但尚无证据足以认定美国原产聚苯醚经第三国绕道入华或实际少缴反倾销税。")
para(doc, "下一轮应优先完成HPP 2026-04-09、2026-02-03四端单证闭合和54条再生料的成分/原产鉴别；在易迅恢复可查询后，补做HS39072990、PPE/PPO、NORYL、XYRON及重点牌号各自独立全页，避免仅依赖两份宽池数据。")

# Sources
heading(doc, "十、主要来源", 1)
source(doc, "商务部公告2022年第1号", "https://www.mofcom.gov.cn/zcfb/zgdwjjmywg/art/2022/art_8314783c9fa74c169f85b7eda373c230.html", "聚苯醚反倾销终裁、范围、税率与期限")
source(doc, "商务部公告2022年第2号", "https://www.mofcom.gov.cn/zcfb/blgg/art/2022/art_a528612c0a4541a2a0b500854f935fe9.html", "微量补贴并终止反补贴调查")
source(doc, "商务部公告2022年第20号", "https://cacs.mofcom.gov.cn/cacscms/article/jkdc?articleId=173788&type=11", "SHPP US LLC税率承继")
source(doc, "商务部实施中措施表", "https://admin.cacs.mofcom.gov.cn/cacscms/article/sszaj?articleId=160267&type=", "预计终止日与现行状态")
source(doc, "Asahi Kasei：Nippisun Indonesia受托生产XYRON", "https://www.asahi-kasei.co.jp/asahi/en/news/2013/e130808.html", "印尼本地制造反证")
source(doc, "Asahi Kasei：XYRON官方产品页", "https://www.asahi-kasei-plastics.com/en/products/xyron/", "官方列有540Z，支持将WYRON 540Z按品牌误拼纳入候选")
source(doc, "SABIC：NORYL GTX产品页", "https://www.sabic.com/en/products/specialties/noryl-resins/noryl-gtx-resin", "PPE+PA产品性质")
source(doc, "SABIC：HPP Mexico 2025 ISCC PLUS证书", "https://www.sabic.com/ar/Images/ISCC-PLUS-certificate-SABIC-SHPP-SLP-Mexico-2025_tcm12-47041.pdf", "站点类型为Compounding plant")
source(doc, "SABIC/SGS：HPP Mexico管理体系证书", "https://www.sabic.com/en/Images/San_Luis_Potosi%20MexicoCertificate_Final_tcm1010-6341.pdf", "活动为热塑性组合物制造")
source(doc, "SEC：Flex 2018子公司清单", "https://www.sec.gov/Archives/edgar/data/866374/000086637418000007/flex-exx210133118.htm", "历史上列Flextronics Plastics与Technologies San Luis")
source(doc, "Flex：塑料制造能力", "https://flex.com/solutions-and-services/manufacturing/plastics", "集团塑料加工能力")
source(doc, "NSF认证库：NORYL GFN3F多站点记录", "https://info.nsf.org/Certified/PwsComponents/Listings.asp?ProductType=Potable+Water+Materials", "提示GFN3F并非美国站点独有，须以批号与制造声明定原产")
source(doc, "中国进出口货物原产地条例", "https://xzfg.moj.gov.cn/front/law/detail?LawID=1523&Query=%E8%B4%A7%E7%89%A9", "实质性改变、微小加工与规避条款")
source(doc, "商务部关于非优惠原产地实质性改变标准的规则", "https://policy.mofcom.gov.cn/claw/clawContent.shtml?id=101350", "四位税目改变、从价比例及特定加工清单机制")

heading(doc, "附件索引", 2)
table(doc, ["附件", "内容"], [
    ["聚苯醚_易迅逐票判定台账_阶段审计.xlsx", "全部10,347行、665范围候选、67中国端候选、597美国来源记录、A/B匹配与政策调证"],
    ["聚苯醚_中国进口范围候选全量.csv", "67条精确去重中国端候选"],
    ["聚苯醚_AB同实体同货描候选明细.csv", "134个A/B候选对"],
    ["聚苯醚_AB同实体同货描链路摘要.csv", "HPP三条、Motores一条"],
    ["聚苯醚_易迅逐票标准化.csv/json", "10,347行逐条标准化与判定"],
], [3.0,4.7], 8.3)

for p in doc.paragraphs:
    if p.style.name.startswith("Heading"):
        p.paragraph_format.keep_with_next = True

doc.core_properties.title = "聚苯醚（PPE/PPO）反倾销税与第三国转运风险阶段审计报告"
doc.core_properties.subject = "易迅全量逐条审计、A/B实体链、原产地与条件税差"
doc.core_properties.author = "Codex"
doc.save(REPORT)
print(json.dumps({"report": str(REPORT), "paragraphs": len(doc.paragraphs), "tables": len(doc.tables)}, ensure_ascii=False, indent=2))
