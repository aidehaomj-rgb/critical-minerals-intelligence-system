from __future__ import annotations
import re, sys, collections
from openpyxl import load_workbook

p=sys.argv[1]
ws=load_workbook(p,read_only=True,data_only=True).active
head=[str(c.value or '').strip() for c in next(ws.iter_rows(min_row=1,max_row=1))]
ix={x:i for i,x in enumerate(head)}

def n(v): return re.sub(r'\s+',' ',str(v or '')).strip().upper()
def has(p,s): return re.search(p,s,re.I) is not None

explicit_re=r'(?<![A-Z])META[\s_-]*CRESOL|(?<![A-Z])METACRESOL|(?<![A-Z])M[\s_-]+CRESOL|(?<![A-Z])M[\s_-]*KRE[ZS]OL|(?<![A-Z])3\s*[-_]?\s*METHYL\s*[-_]?\s*PHENOL|(?<!\d)108\D{0,6}39\D{0,6}4(?!\d)|МЕТАКРЕЗОЛ|МЕТА[-\s]*КРЕЗОЛ'
deriv_re=r'CHLORO|CHLORO[-\s]*M[-\s]*CRES|CHLOROCRESOL|AMYL[\s_-]*META|AMYLMETA|IMPURITY|DERIVATIVE|ESTER|ETHER|ACETATE|PHOSPHATE|SULFON|METHOXY|NITRO|BROMO|IODO|BENZYL|CARVACROL'
mixed_re=r'META[\s_-]*(?:PARA|P)[\s_-]*CRESOL|M[\s_-]*P[\s_-]*CRESOL|MP[\s_-]*CRESOL|M/P[\s_-]*CRESOL|META PARA|META/PARA|MIXED CRESOL|CRESOL MIXTURE|CRESYLIC ACID'
para_re=r'PARA[\s_-]*CRESOL|P[\s_-]+CRESOL|4[\s_-]*METHYL[\s_-]*PHENOL|106\D{0,5}44\D{0,5}5'
ortho_re=r'ORTHO[\s_-]*CRESOL|O[\s_-]+CRESOL|2[\s_-]*METHYL[\s_-]*PHENOL|95\D{0,5}48\D{0,5}7'
generic_re=r'(?<![A-Z])CRESOL(?:ES|S)?(?![A-Z])|(?<![A-Z])CRESOLES(?![A-Z])|КРЕЗОЛ|KRESOL|CRESYLIC'

c=collections.Counter(); desc=collections.Counter(); rows=[]
for er,v in enumerate(ws.iter_rows(min_row=2,values_only=True),2):
 s=n(v[ix['商品描述']]); explicit=has(explicit_re,s); deriv=has(deriv_re,s); mixed=has(mixed_re,s); para=has(para_re,s); ortho=has(ortho_re,s); generic=has(generic_re,s)
 if explicit and not deriv and not mixed: cat='explicit'
 elif generic and not deriv and not mixed and not para and not ortho: cat='generic'
 elif deriv: cat='derivative'
 elif mixed: cat='mixed'
 elif para: cat='para'
 elif ortho: cat='ortho'
 else: cat='other'
 c[cat]+=1; desc[(cat,s)]+=1
 if explicit and cat!='explicit': rows.append((er,cat,s))
print(c)
cy=collections.Counter()
for (cat,s),nn in desc.items():
 if cat=='generic' and re.search(r'М[- ]*КРЕЗОЛ',s,re.I):
  cy['all']+=nn
  if re.match(r'^(?:1\.\s*)?М[- ]*КРЕЗОЛ',s,re.I): cy['starts']+=nn
  if 'ЛАБОРАТ' in s: cy['lab']+=nn
  if 'СТАНДАРТ' in s: cy['std']+=nn
  if re.match(r'^М[- ]*КРЕЗОЛ,?\s*(?:ЛАБОРАТ|ДЛЯ СИНТЕЗА|99|ЧИСТ|СВЕРХЧИСТ)',s,re.I): cy['product_start']+=nn
print('cy',cy)
print('explicit excluded',len(rows))
for (cat,s),nn in sorted(desc.items(),key=lambda x:-x[1])[:0]: print(nn,cat,s)
print('top generic')
for (cat,s),nn in sorted(((k,v) for k,v in desc.items() if k[0]=='generic'),key=lambda x:-x[1])[:100]:print(nn,cat,s[:240])
print('top explicit exclusions')
for (cat,s),nn in sorted(((k,v) for k,v in desc.items() if k[0]!='explicit' and has(explicit_re,k[1])),key=lambda x:-x[1])[:80]:print(nn,cat,s[:240])
