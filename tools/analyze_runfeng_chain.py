"""Extract and summarize every structured row involving the Runfeng entities."""
from __future__ import annotations

import json
import re
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(r"D:\codex\XG执法\2026_案例库")
TARGETS = {
    "深圳润丰成贸易发展有限公司": ("深圳润丰成贸易发展有限公司",),
    "香港潤豐貿易發展有限公司": (
        "香港潤豐貿易發展有限公司",
        "香港润丰贸易发展有限公司",
        "HONGKONG RUNFENG TRADE DEVELOPMENT",
    ),
}
CASE_IDS = ("GDRM26-037", "GDRM26-049", "GDRM26-127", "GDRM26-253", "GDRM26-269", "GDRM26-319")
INTELLIGENCE_CASE_IDS = {"GDRM26-049", "GDRM26-253", "GDRM26-269"}
HEADER_MARKERS = (
    "发货人", "發貨人", "收货人", "收貨人", "申报人", "申報人",
    "报关", "報關", "车牌", "車牌", "车辆", "車輛", "货物", "貨物",
)


def clean(value):
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def identify_target(value: str):
    for canonical, aliases in TARGETS.items():
        if any(alias in value for alias in aliases):
            return canonical
    return None


rows_out = []
for case_id in CASE_IDS:
    for path in (ROOT / "raw_cases" / case_id).rglob("*.xlsx"):
        try:
            book = load_workbook(path, read_only=True, data_only=True)
        except Exception:
            continue
        for sheet in book.worksheets:
            header = None
            for row_num, row in enumerate(sheet.iter_rows(values_only=True), 1):
                values = [clean(value) for value in row]
                marker_count = sum(any(marker in value for marker in HEADER_MARKERS) for value in values)
                if marker_count >= 2:
                    header = values
                    continue
                if header is None:
                    continue
                targets_in_row = sorted({identify_target(value) for value in values if identify_target(value)})
                if not targets_in_row:
                    continue
                record = {}
                for index, value in enumerate(values):
                    if not value:
                        continue
                    key = header[index] if index < len(header) and header[index] else f"列{index + 1}"
                    if key in record:
                        key = f"{key}_{index + 1}"
                    record[key] = value
                rows_out.append({
                    "case_id": case_id,
                    "file": path.relative_to(ROOT).as_posix(),
                    "sheet": sheet.title,
                    "row": row_num,
                    "targets": targets_in_row,
                    "record": record,
                })


def values_for(row, patterns):
    values = []
    for key, value in row["record"].items():
        if any(pattern in key.casefold() for pattern in patterns):
            values.append(value)
    return values


def first_value(row, patterns):
    values = values_for(row, patterns)
    return values[0] if values else ""


def transaction_fingerprint(row):
    record = row["record"]
    date = first_value(row, ("运输日期", "運輸日期", "时间", "時間", "申报日期", "申報日期"))
    vehicle = first_value(row, ("车牌", "車牌", "车辆登记", "車輛登記"))
    shipper = first_value(row, ("发货人名称", "發貨人名稱"))
    consignee = first_value(row, ("收货人名称", "收貨人名稱"))
    declarant = "".join(values_for(row, ("申报人名称", "申報人名稱")))
    goods = first_value(row, ("报关货物名称", "報關貨物名稱"))
    value = first_value(row, ("货物价格", "貨物價格"))
    weight = first_value(row, ("重量",))
    if not date and not vehicle:
        return None
    return (
        row["case_id"], date, vehicle, shipper, consignee,
        declarant, goods, value, weight,
    )


unique_transactions = []
seen_transactions = set()
for row in rows_out:
    fingerprint = transaction_fingerprint(row)
    if fingerprint is None or fingerprint in seen_transactions:
        continue
    seen_transactions.add(fingerprint)
    unique_transactions.append(row)


summary = {}
for target in TARGETS:
    target_rows = [row for row in unique_transactions if target in row["targets"]]
    case_counts = Counter(row["case_id"] for row in target_rows)
    role_counts = Counter()
    goods = Counter()
    vehicles = Counter()
    declarants = Counter()
    counterparties = Counter()
    dates = []
    for row in target_rows:
        for key, value in row["record"].items():
            folded = key.casefold()
            if ("发货人" in key or "發貨人" in key) and target in row["targets"] and identify_target(value) == target:
                role_counts["发货人"] += 1
            if ("收货人" in key or "收貨人" in key) and identify_target(value) == target:
                role_counts["收货人"] += 1
            if ("申报人" in key or "申報人" in key) and identify_target(value) == target:
                role_counts["申报人"] += 1
        goods.update(values_for(row, ("货物名称", "貨物名稱", "报关货物", "報關貨物")))
        vehicles.update(values_for(row, ("车牌", "車牌", "车辆登记", "車輛登記")))
        declarants.update(values_for(row, ("申报人名称", "申報人名稱")))
        dates.extend(values_for(
            row,
            ("运输日期", "運輸日期", "时间", "時間", "申报日期", "申報日期", "进出境日期", "進出境日期"),
        ))
        for value in values_for(row, ("发货人名称", "發貨人名稱", "收货人名称", "收貨人名稱")):
            if identify_target(value) != target:
                counterparties[value] += 1
    summary[target] = {
        "row_count": len(target_rows),
        "case_counts": case_counts.most_common(),
        "roles": role_counts.most_common(),
        "date_min": min(dates) if dates else "",
        "date_max": max(dates) if dates else "",
        "goods_top": goods.most_common(20),
        "vehicles_top": vehicles.most_common(20),
        "declarants_top": declarants.most_common(20),
        "counterparties_top": counterparties.most_common(20),
    }

cooccurrence = [
    row for row in unique_transactions
    if set(row["targets"]) == set(TARGETS)
]
intelligence_transactions = [
    row for row in unique_transactions if row["case_id"] in INTELLIGENCE_CASE_IDS
]
intelligence_cooccurrence = [
    row for row in intelligence_transactions
    if set(row["targets"]) == set(TARGETS)
]
output = {
    "targets": list(TARGETS),
    "summary": summary,
    "unique_transaction_count": len(unique_transactions),
    "unique_transactions": unique_transactions,
    "same_transaction_cooccurrence_count": len(cooccurrence),
    "same_transaction_cooccurrence": cooccurrence,
    "intelligence_transaction_count": len(intelligence_transactions),
    "intelligence_transactions": intelligence_transactions,
    "intelligence_cooccurrence_count": len(intelligence_cooccurrence),
    "intelligence_cooccurrence": intelligence_cooccurrence,
    "all_matching_rows": rows_out,
}
(ROOT / "润丰成潤豐链路逐票核查.json").write_text(
    json.dumps(output, ensure_ascii=False, indent=2), encoding="utf-8"
)
print(json.dumps({
    "summary": summary,
    "unique_transaction_count": len(unique_transactions),
    "same_transaction_cooccurrence_count": len(cooccurrence),
    "intelligence_transaction_count": len(intelligence_transactions),
    "intelligence_cooccurrence_count": len(intelligence_cooccurrence),
    "same_transaction_examples": cooccurrence[:10],
}, ensure_ascii=True, indent=2))
