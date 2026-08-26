from __future__ import annotations

from collections import Counter
from pathlib import Path
import json
import re

from openpyxl import load_workbook


SOURCE = Path(r"D:\易迅数据\EPDM_400270.xlsx")


def norm(v: object) -> str:
    return "" if v is None else str(v).strip()


def main() -> None:
    wb = load_workbook(SOURCE, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    header = next(it)
    seen = set()
    rows = []
    for excel_row, vals in enumerate(it, start=2):
        vals = tuple(norm(v) for v in vals)
        if vals in seen:
            continue
        seen.add(vals)
        rows.append((excel_row, vals))

    pats = [
        "PRIMARY FORMS", "PRIMARY FORM", "PLATES", "SHEETS", "STRIP", "ROLL",
        "PELLET", "GRANULE", "POWDER", "BALE", "BAG", "BLOCK", "LUMP",
        "COMPOUND", "COMPOUNDED", "MODIFIED", "FUNCTIONAL POLYMER", "TPV", "TPE",
        "SANTOPRENE", "VULCANIZED", "VULCANISED", "MASTERBATCH", "BLEND",
        "PROFILE", "GASKET", "SEAL", "O-RING", "ORING", "HOSE", "TUBE", "PIPE",
        "AUTOMOTIVE", "AUTO PART", "PARTS", "MEMBRANE", "ROOFING", "COVER",
        "FOAM", "SPONGE", "EXTRUDED", "MOLDED", "MOULDED", "WIRE", "CABLE",
        "ARTICLE", "PRODUCT", "SCRAP", "WASTE", "RECYCLED", "RECLAIMED",
    ]
    counts = {}
    for p in pats:
        counts[p] = sum(1 for _, vals in rows if p in vals[4].upper())

    common = Counter(vals[4] for _, vals in rows)
    form_terms = ("FORM", "FORMAS", "PLATES", "PLACAS", "SHEETS", "HOJAS", "STRIP", "TIRAS", "FEUILLES", "BANDES", "CHAPAS", "FOLHAS", "BENTUK", "LEMBAR")
    form_common = Counter(
        vals[4] for _, vals in rows if any(term in vals[4].upper() for term in form_terms)
    )
    payload = {
        "unique": len(rows),
        "pattern_counts": counts,
        "top_descriptions": common.most_common(120),
        "top_form_descriptions": form_common.most_common(220),
    }
    print(json.dumps(payload, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
