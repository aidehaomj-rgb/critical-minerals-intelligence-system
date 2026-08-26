from __future__ import annotations

from collections import Counter
from pathlib import Path
import json

from openpyxl import load_workbook


SOURCE = Path(r"D:\易迅数据\EPDM_400270.xlsx")


def main() -> None:
    wb = load_workbook(SOURCE, read_only=True, data_only=True)
    out: dict[str, object] = {"source": str(SOURCE), "sheets": []}
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        header = next(rows)
        samples = []
        nonempty_counts = Counter()
        types = {str(h): Counter() for h in header}
        first_date = None
        last_date = None
        n = 0
        destination_counts = Counter()
        origin_counts = Counter()
        source_counts = Counter()
        china_rows = []
        for excel_row, values in enumerate(rows, start=2):
            if not any(v not in (None, "") for v in values):
                continue
            n += 1
            if len(samples) < 8:
                samples.append(list(values))
            for h, v in zip(header, values):
                key = str(h)
                if v not in (None, ""):
                    nonempty_counts[key] += 1
                    types[key][type(v).__name__] += 1
            d = values[0] if values else None
            if d is not None:
                if first_date is None or d < first_date:
                    first_date = d
                if last_date is None or d > last_date:
                    last_date = d
            source_counts[str(values[0])] += 1
            destination_counts[str(values[10])] += 1
            origin_counts[str(values[11])] += 1
            if str(values[10]).strip().casefold() == "china":
                china_rows.append([excel_row, *values])
        out["sheets"].append(
            {
                "name": ws.title,
                "max_row": ws.max_row,
                "max_column": ws.max_column,
                "data_rows_nonempty": n,
                "header": list(header),
                "nonempty_counts": dict(nonempty_counts),
                "value_types": {k: dict(v) for k, v in types.items()},
                "first_value_col1": str(first_date),
                "last_value_col1": str(last_date),
                "samples": samples,
                "top_sources": source_counts.most_common(30),
                "destinations_containing_china": [
                    [k, v] for k, v in destination_counts.items() if "china" in k.casefold()
                ],
                "origins_for_china": Counter(str(r[-1]) for r in china_rows),
                "china_rows": china_rows,
            }
        )
    print(json.dumps(out, ensure_ascii=False, indent=2, default=str))


if __name__ == "__main__":
    main()
