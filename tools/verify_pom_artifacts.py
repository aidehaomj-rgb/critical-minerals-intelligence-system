from pathlib import Path
from openpyxl import load_workbook
from docx import Document
folder=Path(r'D:\易迅数据\反倾销税深度分析报告\06_共聚聚甲醛（POM）')
x=folder/'POM_易迅全量逐票判定台账.xlsx'
wb=load_workbook(x,read_only=True,data_only=False)
print('xlsx_sheets',wb.sheetnames)
for ws in wb.worksheets: print(ws.title,ws.max_row,ws.max_column)
doc=Document(folder/'共聚聚甲醛（POM）_反倾销税与第三国转运风险深度分析报告.docx')
print('doc_paragraphs',len(doc.paragraphs),'tables',len(doc.tables),'sections',len(doc.sections))
print('headings',[(p.style.name,p.text[:60]) for p in doc.paragraphs if p.style.name.startswith('Heading')][:8])
print('files',x.stat().st_size,(folder/'共聚聚甲醛（POM）_反倾销税与第三国转运风险深度分析报告.docx').stat().st_size)
