#!/usr/bin/env python3
"""Build item 17 ODCB antidumping / rerouting stage report."""

from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(r"D:\易迅数据\反倾销税深度分析报告\17_邻二氯苯")
DOCX = OUT / "邻二氯苯_反倾销与第三国转运风险_数据缺口阶段报告.docx"
NAVY = RGBColor(11, 37, 69); BLUE = RGBColor(46, 116, 181); GRAY = RGBColor(90, 98, 108); RED = RGBColor(155, 28, 28)


def font(run, size=11, bold=False, color=None):
    run.font.name = "Calibri"; run._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
    run.font.size = Pt(size); run.bold = bold
    if color: run.font.color.rgb = color


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr(); node = OxmlElement("w:shd"); node.set(qn("w:fill"), fill); tcpr.append(node)


def margins(cell, top=75, start=110, bottom=75, end=110):
    tcpr = cell._tc.get_or_add_tcPr(); tcmar = tcpr.first_child_found_in("w:tcMar")
    if tcmar is None: tcmar = OxmlElement("w:tcMar"); tcpr.append(tcmar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        x = tcmar.find(qn(f"w:{name}"))
        if x is None: x = OxmlElement(f"w:{name}"); tcmar.append(x)
        x.set(qn("w:w"), str(value)); x.set(qn("w:type"), "dxa")


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    trpr = t.rows[0]._tr.get_or_add_trPr(); rep = OxmlElement("w:tblHeader"); rep.set(qn("w:val"), "true"); trpr.append(rep)
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; c.width = Inches(widths[i]); shade(c, "F2F4F7"); margins(c)
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; font(p.add_run(h), 9.5, True, NAVY)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].width = Inches(widths[i]); cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; margins(cells[i])
            p = cells[i].paragraphs[0]; p.paragraph_format.space_after = Pt(0); font(p.add_run(str(v)), 9)
    return t


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after = Pt(3); font(p.add_run(text), 10.2); return p


def link(doc, label, url):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2); font(p.add_run(label + "："), 8.8, True, GRAY); font(p.add_run(url), 8.8, False, BLUE)


def main():
    OUT.mkdir(parents=True, exist_ok=True); doc = Document(); sec = doc.sections[0]
    sec.page_width = Inches(8.5); sec.page_height = Inches(11); sec.top_margin = sec.bottom_margin = Inches(0.9); sec.left_margin = sec.right_margin = Inches(0.9)
    normal = doc.styles["Normal"]; normal.font.name = "Calibri"; normal._element.rPr.rFonts.set(qn("w:eastAsia"), "等线"); normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5); normal.paragraph_format.line_spacing = 1.08
    for name, size, before, after, color in (("Heading 1", 15, 14, 7, BLUE), ("Heading 2", 12.5, 10, 5, BLUE)):
        s = doc.styles[name]; s.font.name = "Calibri"; s._element.rPr.rFonts.set(qn("w:eastAsia"), "等线"); s.font.size = Pt(size); s.font.bold = True; s.font.color.rgb = color
        s.paragraph_format.space_before = Pt(before); s.paragraph_format.space_after = Pt(after)
    h = sec.header.paragraphs[0]; font(h.add_run("反倾销税深度分析｜第17项"), 9, True, GRAY)
    f = sec.footer.paragraphs[0]; f.alignment = WD_ALIGN_PARAGRAPH.RIGHT; font(f.add_run("阶段审计报告｜2026-08-13"), 9, False, GRAY)

    p = doc.add_paragraph(); font(p.add_run("风险核查备忘录"), 10, True, BLUE)
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(3); font(p.add_run("邻二氯苯（ODCB）"), 23, True, NAVY)
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(12); font(p.add_run("中国反倾销措施、第三国转运风险与易迅数据缺口阶段审计"), 12.5, False, GRAY)
    for label, value in (("受税来源", "日本、印度"), ("中国税号", "29039110；CAS 95-50-1；UN 1591"), ("审计日期", "2026-08-13"),
                         ("本地范围", "D:\易迅数据内294个XLSX/XLS/CSV/JSON文件内容级预筛"),
                         ("阶段结论", "未取得可确认的邻二氯苯逐票贸易数据；公开源亦未闭合日本/印度→第三国→中国链")):
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(1); font(p.add_run(label + "："), 10.2, True); font(p.add_run(value), 10.2)

    doc.add_heading("一、结论先行", level=1)
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6); font(p.add_run("当前没有证据证明邻二氯苯经第三国进入中国并逃避反倾销税。"), 12, True, RED)
    bullet(doc, "现行措施有效：2025年1月23日起续征5年，预计至2030年1月22日；日本全部公司70.4%，印度全部公司31.9%。")
    bullet(doc, "本地294个文件预筛中8个文件出现项目名称或数字片段；逐条复核后只有主清单和进度台账两处项目元数据，可确认贸易记录为0。")
    bullet(doc, "商务部复审基线显示，2019年至2023年前9个月中国进口几乎全部来自日本和印度。措施后若29039110突然以缺乏已核实产能的第三国原产大票进入中国，应优先穿透生产厂、原产证和批次。")
    bullet(doc, "公开检索未找到本品中国海关处罚、法院判决、反规避裁定或可闭合双段提单；第三国商业节点只能作为筛选器，不能写成既成逃税。")

    doc.add_heading("二、现行政策、范围与税差口径", level=1)
    table(doc, ["来源/企业", "反倾销税率", "AD及AD引致VAT增量*"], [
        ("KUREHA CORPORATION", "70.4%", "79.552% × 中国完税价格"),
        ("其他日本公司", "70.4%", "79.552% × 中国完税价格"),
        ("所有印度公司", "31.9%", "36.047% × 中国完税价格"),
    ], [2.35, 1.25, 2.9])
    p = doc.add_paragraph("* 按进口增值税13%作风险筛查：完税价格×AD税率×1.13。实际税款只能以中国进口报关完税价格、法定原产地和税款缴款书计算；不含正常关税和正常进口增值税。")
    p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(6); font(p.runs[0], 8.5, False, GRAY)
    p = doc.add_paragraph("范围：邻二氯苯（Ortho Dichlorobenzene、1,2-Dichlorobenzene、O-Dichlorobenzene，简称ODCB），分子式C₆H₄Cl₂，通常为无色易挥发液体，用作有机溶剂及农药、医药、染料中间体。中国税号29039110。")
    p = doc.add_paragraph("排除边界：对二氯苯/PDCB（1,4-）、间二氯苯/MDCB（1,3-）、二氯硝基苯、二氯苯酚、二氯苯胺及其他下游衍生物，不得仅凭“DICHLOROBENZENE”子串或六位HS机械纳入。")

    doc.add_heading("三、官方贸易基线与风险含义", level=1)
    table(doc, ["官方复审指标", "事实", "风险含义"], [
        ("中国进口来源", "2019—2023年前9个月，日印合计份额约99.85%、99.97%、99.41%、79.36%、99.97%", "措施后第三国大票偏离历史结构，但异常不等于规避"),
        ("日本产能", "2019—2022年约13,000吨/年；出口量长期占产量较高比例", "日本生产端可外销，需匹配Kureha/贸易商A腿"),
        ("印度产能", "2019—2022年约42,000吨/年；复审期对华直接出口很少", "印度具备显著产能；直接出口少不能排除转售，也不能反推转运"),
        ("复审配合", "日印生产商/出口商均未登记答卷", "官方采用可得最佳信息；不配合不等于违法"),
    ], [1.3, 3.0, 2.2])
    p = doc.add_paragraph("公开可核生产端：日本列名生产商KUREHA CORPORATION；印度环境合规资料可见Aarti Industries等企业生产或使用ODCB。上述实体用于精确检索和生产厂核验，不代表其涉嫌规避。")

    doc.add_heading("四、本地易迅数据全盘盘点", level=1)
    table(doc, ["审计项", "结果", "说明"], [
        ("文件总数", "294", "XLSX/XLS/CSV/JSON；排除本项输出目录自身"),
        ("预筛候选文件", "8", "主清单、进度台账及POM/正丙醇项目派生文件"),
        ("逐条重点命中", "2", "均为商品清单/台账中的“邻二氯苯”项目名称"),
        ("可确认贸易记录", "0", "无日期—商品—主体—数量—路线的逐票记录"),
        ("A腿/B腿/闭合链", "0 / 0 / 0", "不能统计数量、企业、口岸或税差"),
        ("解析错误", "0", "本地文件内容级预筛完成"),
    ], [1.55, 1.2, 3.75])
    p = doc.add_paragraph("审计边界：该结果只证明现有D盘缺少邻二氯苯原始易迅数据，不代表易迅全库无数据，也不代表中国没有进口或没有第三国风险。")

    doc.add_heading("五、第三国风险判定框架", level=1)
    table(doc, ["证据等级", "满足条件", "当前状态"], [
        ("A｜可认定链", "A/B腿同批号、同ISO罐/柜号或同提单，第三国无实质生产，且中国申报非日印原产/未缴AD", "0条"),
        ("B+｜优先调单", "同实体、同品名/CAS/纯度、0—90日、净重±2%，且第三国生产能力存疑", "现有数据无法构建"),
        ("B/C｜路线异常", "仅第三国发货、集团贸易商、宏观流量或品牌/供应商相同", "只能作查询条件，不能定性"),
        ("反证", "第三国真实氯化/精馏产能、库存分拨、换单但法定原产仍如实申报并正常缴税", "必须与风险线索同步核查"),
    ], [1.2, 3.7, 1.6])
    bullet(doc, "原产地判断适用中国非优惠原产地规则。仓储、分装、换包装、换单通常不改变原产；同一ODCB仅作简单精馏/净化是否构成实质性改变，必须看投入产出税目、工序、成本和海关核定。")
    bullet(doc, "即便货物从第三国物理发运，只要中国报关仍如实申报日本/印度原产并缴纳相应反倾销税，也不构成逃税。")

    doc.add_heading("六、下一轮易迅全页查询矩阵", level=1)
    table(doc, ["查询", "条件", "目标"], [
        ("Q1 中国B腿", "目的国中国；HS29039110/290391；2024-01-23至最新并补历史；7组品名/CAS/UN关键词分别查", "每页200条读至末页；跨查询去重并逐票范围判定"),
        ("Q2 日印A腿", "日本、印度→全球；同税号/关键词；精确补查KUREHA、AARTI及Q1出现主体", "识别实际第三国收货人、仓储商、贸易商和生产厂"),
        ("Q3 第三国B腿", "Q2第三国→中国；同品名、同主体精确补查", "按0—30/60/90/180日、净重±2%和运输指纹闭合"),
    ], [1.25, 3.85, 1.4])
    p = doc.add_paragraph("关键词：ORTHO DICHLOROBENZENE、1,2-DICHLOROBENZENE、O-DICHLOROBENZENE、ODCB、CAS 95-50-1、UN1591、邻二氯苯。重点第三国先覆盖越南、新加坡、阿联酋、马来西亚、韩国、台湾、泰国、印尼，再按Q2真实收货国扩展。")

    doc.add_heading("七、调证字段与可计算税差", level=1)
    bullet(doc, "贸易单证：两段提单、ISO罐号/柜号、封志、船名航次、中转港、合同、商业发票、付款受益人、仓储/分装记录。")
    bullet(doc, "货物指纹：CAS、纯度、水分、色度、批号、COA、SDS、UN1591、包装/罐型、净重、生产厂代码、生产日期。")
    bullet(doc, "原产证据：原产地证、厂家声明、第三国生产许可、氯化/精馏装置、BOM、能耗、投入产出和库存平衡；仅贸易公司地址不足以证明产能。")
    bullet(doc, "中国端：进口人、消费使用单位、报关企业、申报口岸/隶属关、启运国、原产国、境外生产商、完税价格、反倾销税及进口增值税缴款书。")
    bullet(doc, "若中国完税价格为V，且确认实际日本原产而未缴AD，筛查税差为0.79552V；实际印度原产为0.36047V。没有中国完税价格和税单时不得把出口侧金额写成欠税。")

    doc.add_heading("八、阶段审计结论", level=1)
    bullet(doc, "已核清政策、税率、产品范围、官方来源结构和本地数据缺口。")
    bullet(doc, "没有发现可指向中国进口人、境外第三国收发货人、报关行或具体口岸的邻二氯苯风险票据。")
    bullet(doc, "没有形成日本/印度→第三国→中国的具体证据；后续取证优先级是Q1全量中国B腿，其次按B腿主体反查日印A腿。")
    bullet(doc, "在取得全页逐票数据前，本项评级为“数据不足/无法判断”，不是“低风险”或“无风险”。")

    doc.add_heading("九、主要公开来源", level=1)
    link(doc, "商务部2025年第4号期终复审公告", "https://www.mofcom.gov.cn/zcfb/blgg/gg/2025/art/2025/art_c97ea609cb4644be9fb3de7b5fd3c95b.html")
    link(doc, "商务部2019年第1号终裁公告", "https://www.mofcom.gov.cn/zcfb/dwmygl/art/2019/art_5a43250926634a0f81a74c3777fc2acf.html")
    link(doc, "期终复审裁定公开PDF", "https://swt.fujian.gov.cn/xxgk/tzgg/202502/P020250207378636987637.pdf")
    link(doc, "中国非优惠原产地实质性改变标准", "https://policy.mofcom.gov.cn/claw/clawContent.shtml?id=101350")
    link(doc, "中华人民共和国进出口货物原产地条例", "https://xzfg.moj.gov.cn/front/law/detail?LawID=1523")
    doc.save(DOCX); print(DOCX)


if __name__ == "__main__":
    main()
