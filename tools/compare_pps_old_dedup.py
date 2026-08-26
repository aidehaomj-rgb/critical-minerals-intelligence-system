from __future__ import annotations

import json
import re
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

import openpyxl


FILES = [
    Path("D:/\u6613\u8fc5\u6570\u636e/PPS_391190_1.xlsx"),
    Path("D:/\u6613\u8fc5\u6570\u636e/PPS_391190_2.xlsx"),
]

PATTERN = re.compile(
    r"POLYPHENYLENE\s*SULFIDE|POLYPHENYLENE\s*SULPHIDE|"
    r"\bPPS\s*(?:RESIN|COMPOUND|NEAT|GF|[A-Z0-9-])|"
    r"RYTON|DURAFIDE|FORTRON|TORELINA",
    flags=re.I,
)


def norm_text(value: object) -> str:
    return "" if value is None else " ".join(str(value).strip().split())


def old_norm(value: object) -> str:
    return re.sub(r"[^A-Z0-9]+", "", str(value).upper())


def old_num(value: object) -> str:
    # pandas numeric coercion in the old script yielded floats.
    if value in (None, ""):
        return "nan"
    try:
        return str(float(str(value).replace(",", "")))
    except ValueError:
        return "nan"


def old_date(value: object) -> str:
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    return norm_text(value)[:10]


def exact_key(values: tuple[object, ...]) -> tuple[str, ...]:
    return tuple(norm_text(v) for v in values)


def trim_key(values: tuple[object, ...]) -> tuple[str, ...]:
    return tuple("" if v is None else str(v).strip() for v in values)


def old_key(values: tuple[object, ...]) -> str:
    # Column order: source, direction, date, HS, description, buyer, supplier,
    # weight, quantity, amount, destination, origin.
    return "|".join(
        [
            old_date(values[2]),
            old_norm(values[4])[:250],
            old_norm(values[5]),
            old_norm(values[6]),
            old_num(values[7]),
            old_num(values[8]),
            old_num(values[9]),
            old_norm(values[10]),
            old_norm(values[11]),
        ]
    )


rows: list[dict[str, object]] = []
for path in FILES:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    iterator = ws.iter_rows(values_only=True)
    headers = [norm_text(v) for v in next(iterator)]
    for excel_row, values in enumerate(iterator, start=2):
        if not any(v not in (None, "") for v in values):
            continue
        vals = tuple(values[:12])
        rows.append(
            {
                "file": path.name,
                "row": excel_row,
                "values": vals,
                "desc": norm_text(vals[4]),
                "exact": exact_key(vals),
                "trim": trim_key(vals),
                "old": old_key(vals),
                "matched": bool(PATTERN.search(norm_text(vals[4]))),
            }
        )
    wb.close()

matched = [r for r in rows if r["matched"]]
old_groups: dict[str, list[dict[str, object]]] = defaultdict(list)
for row in matched:
    old_groups[str(row["old"])].append(row)

collisions = []
for key, group in old_groups.items():
    exacts = {r["exact"] for r in group}
    if len(exacts) <= 1:
        continue
    collisions.append(
        {
            "old_key": key,
            "raw_rows": len(group),
            "exact_distinct": len(exacts),
            "locations": [{"file": r["file"], "row": r["row"]} for r in group],
            "records": [list(exact) for exact in exacts],
        }
    )

unmatched = [r for r in rows if not r["matched"]]
result = {
    "raw": len(rows),
    "regex_matched_raw": len(matched),
    "regex_unmatched_raw": len(unmatched),
    "regex_unmatched_exact": len({r["exact"] for r in unmatched}),
    "regex_unmatched": [
        {
            "file": r["file"],
            "row": r["row"],
            "record": dict(zip(headers, r["values"])),
        }
        for r in unmatched
    ],
    "old_key_unique": len(old_groups),
    "exact_unique_all": len({r["exact"] for r in rows}),
    "trim_unique_all": len({r["trim"] for r in rows}),
    "exact_unique_matched": len({r["exact"] for r in matched}),
    "old_key_collision_groups": len(collisions),
    "old_key_collision_exact_excess": sum(c["exact_distinct"] - 1 for c in collisions),
    "collisions": collisions,
    "whitespace_equivalence_collisions": [
        {
            "normalized_key": list(key),
            "trim_variants": [list(x) for x in {r["trim"] for r in group}],
            "locations": [{"file": r["file"], "row": r["row"]} for r in group],
        }
        for key, group in (
            (key, [r for r in rows if r["exact"] == key]) for key in {r["exact"] for r in rows}
        )
        if len({r["trim"] for r in group}) > 1
    ],
}
print(json.dumps(result, ensure_ascii=False, indent=2, default=str))
