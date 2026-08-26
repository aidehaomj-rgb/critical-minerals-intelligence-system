from __future__ import annotations

from collections import Counter
from pathlib import Path
import json
import re

from openpyxl import load_workbook

SOURCE = Path(r"D:\易迅数据\EPDM_400270.xlsx")


def raw(v):
    return v


def text(v):
    return "" if v is None else str(v)


def has(pat: str, s: str) -> bool:
    return re.search(pat, s, flags=re.I) is not None


STANDARD = re.compile(
    r"(?:"
    r"PRIMARY\s+FORMS?.{0,80}(?:PLATES?|SHEETS?|STRIPS?)|"
    r"FORMAS?\s+PRIMARIAS?.{0,100}(?:PLACAS?|HOJAS?|TIRAS?)|"
    r"FORMES?\s+PRIMAIRES?.{0,100}(?:PLAQUES?|FEUILLES?|BANDES?)|"
    r"FORMAS?\s+PRIM[ÁA]RIAS?.{0,100}(?:CHAPAS?|FOLHAS?|TIRAS?)|"
    r"BENTUK\s+(?:ASAL|PRIMER).{0,100}(?:PELAT|LEMBARAN|JALUR)|"
    r"ПЕРВИЧН.{0,120}(?:ПЛАСТИН|ЛИСТ|ПОЛОС)|"
    r"(?:PRIM[ÄA]RFORM|PRIMÆRFORM).{0,100}(?:PLATT|SKIV|REMS)"
    r")", re.I | re.S
)

COMPOUND = re.compile(
    r"(?:COMPOUND|COMPOUNDED|COMPUEST[OA]|COMPOST[OA]|MISTURA|MIXED\s+RUBBER|"
    r"RUBBER\s+MIX|CAUCHO\s+MEZCLADO|FUNCTIONAL\s+POLYMER|FUNCTIONALIZED|"
    r"\bTPV\b|\bTPE\b|SANTOPRENE|TERMOPLAST|THERMOPLAST|"
    r"MASTERBATCH|POLYMER\s+BLEND|EPDM\s+BLEND|MODIFIED\s+EPDM|"
    r"HỖN\s*HỢP|HOP\s*CHAT|PHOI\s*TRON)", re.I
)

OUTSIDE = re.compile(
    r"(?:"
    r"PROFILES?|PERFILES?|PROFILE\s+AUTO|AUTO\s+PARTS?|AUTOMOTIVE\s+PARTS?|"
    r"GASKETS?|O[- ]?RINGS?|\bSEALS?\b|SEALING|WEATHERSTRIP|WEATHER\s+STRIP|"
    r"HOSES?|MANGUERAS?|\bTUBES?\b|PIPES?|BELTS?|CONVEYOR|"
    r"MEMBRANES?|ROOFING|FLOOR(?:ING)?|MAT(?:S|TING)?\b|"
    r"ADHESIVE|SELF[- ]?ADHESIVE|TAPE\b|\bFOAM\b|\bSPONGE\b|"
    r"RUBBER\s+ARTICLE|RUBBER\s+PRODUCT|FINISHED\s+PRODUCT|"
    r"MOLDED|MOULDED|EXTRUDED|INJECTION\s+MOLD|"
    r"GIO[ĂA]NG|MIẾNG\s+ĐỆM|MIE?NG\s+DEM|ĐỆM\s+BẢO\s+VỆ|"
    r"DẢI\s+CAO\s+SU|DAI\s+CAO\s+SU|PHỤ\s+TÙNG|PHU\s+TUNG|"
    r"TẤM\s+CAO\s+SU\s+XỐP|TAM\s+CAO\s+SU\s+XOP|"
    r"JUNTAS?|SELLOS?|SOPORTES?\s+DE\s+CAUCHO|PIEZAS?\s+DE|"
    r"TIRAS?\s+DE\s+(?:DE\s+)?(?:CAUCHO|HULE).{0,60}(?:ADHES|AUTO|PUERTA)|"
    r"PLACAS?\s+DE\s+CAUCHO\s+(?:CELULAR|ESPONJOSO)|"
    r"ESPUMA|ESPONJA|SPONGE\s+(?:SEAL|STRIP|SHEET)|FOAM\s+(?:SEAL|STRIP|SHEET)|"
    r"CABLE|WIRE\s+SEAL|COVER|BUMPER|MOUNT|PAD\b|PROTECTIVE\s+PAD|"
    r"SCRAP|WASTE|RECYCLED\s+ARTICLE"
    r")", re.I
)

FORM = re.compile(
    r"(?:"
    r"\bSHEETS?\b|\bPLATES?\b|\bSTRIPS?\b|\bROLLS?\b|\bPELLETS?\b|\bGRANULES?\b|"
    r"\bTIRAS?\b|\bPLACAS?\b|\bHOJAS?\b|\bROLLOS?\b|\bGRANULOS?\b|\bGRANULADO\b|"
    r"\bCHAPAS?\b|\bFOLHAS?\b|\bBANDES?\b|\bPLAQUES?\b|\bFEUILLES?\b|"
    r"DẠNG\s+(?:TẤM|TỜ|HẠT|CUỘN)|DANG\s+(?:TAM|TO|HAT|CUON)|"
    r"HÌNH\s+HẠT|HINH\s+HAT|颗粒|片材|板材|卷材|条状"
    r")", re.I
)


def classify(desc: str) -> str:
    s = desc.upper()
    if COMPOUND.search(s):
        return "compound"
    if OUTSIDE.search(s):
        return "outside"
    if STANDARD.search(s):
        return "standard"
    if FORM.search(s):
        return "form"
    return "raw"


def main():
    wb = load_workbook(SOURCE, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    next(it)
    seen = set()
    rows = []
    for idx, vals in enumerate(it, start=2):
        key = tuple(vals)
        if key in seen:
            continue
        seen.add(key)
        desc = text(vals[4])
        rows.append((idx, text(vals[0]), desc, classify(desc)))
    counts = Counter(cat for _, _, _, cat in rows)
    by_cat = {}
    for cat in counts:
        c = Counter(desc for _, _, desc, ccat in rows if ccat == cat)
        by_cat[cat] = c.most_common(35)
    source_by_cat = {
        cat: Counter(source for _, source, _, ccat in rows if ccat == cat).most_common()
        for cat in counts
    }
    standard_india = [
        [idx, desc] for idx, source, desc, cat in rows if cat == "standard" and source == "印度全港"
    ]
    cue_terms = ["ADHES", "CELLULAR", "FOAM", "SPONGE", "GASKET", "SEAL", "AUTO", "PART", "PROFILE", "KÍCH THƯỚC", "KICH THUOC", "SỬ DỤNG CHO", "SU DUNG CHO"]
    uncaught_cues = {
        cat: {
            cue: sum(1 for _, _, desc, ccat in rows if ccat == cat and cue in desc.upper())
            for cue in cue_terms
        }
        for cat in ("raw", "form", "standard")
    }
    form_token_patterns = [
        ("granule_pellet", re.compile(r"\b(?:GRANULES?|PELLETS?|GRANULOS?|GRANULADO)\b|DẠNG\s+HẠT|DANG\s+HAT|HÌNH\s+HẠT|HINH\s+HAT|颗粒", re.I)),
        ("sheet_plate", re.compile(r"\b(?:SHEETS?|PLATES?|PLACAS?|HOJAS?|CHAPAS?|FOLHAS?|PLAQUES?|FEUILLES?)\b|DẠNG\s+(?:TẤM|TỜ)|DANG\s+(?:TAM|TO)|片材|板材", re.I)),
        ("strip", re.compile(r"\b(?:STRIPS?|TIRAS?|BANDES?)\b|条状", re.I)),
        ("roll", re.compile(r"\b(?:ROLLS?|ROLLOS?)\b|DẠNG\s+CUỘN|DANG\s+CUON|卷材", re.I)),
    ]
    form_tokens = Counter()
    granule_detail = Counter()
    for _, _, desc, cat in rows:
        if cat != "form":
            continue
        labels = [name for name, pat in form_token_patterns if pat.search(desc)]
        form_tokens["+".join(labels) or "other"] += 1
        if "granule_pellet" in labels:
            u = desc.upper()
            if re.search(r"\bGRANULES?\b", u): granule_detail["english_granule"] += 1
            elif re.search(r"\bPELLETS?\b", u): granule_detail["english_pellet"] += 1
            elif re.search(r"\b(?:GRANULOS?|GRANULADO)\b", u): granule_detail["spanish_portuguese"] += 1
            elif re.search(r"DẠNG\s+HẠT|HÌNH\s+HẠT", u): granule_detail["vietnamese_diacritic"] += 1
            elif re.search(r"DANG\s+HAT|HINH\s+HAT", u): granule_detail["vietnamese_ascii"] += 1
            elif "颗粒" in desc: granule_detail["chinese"] += 1
            else: granule_detail["other"] += 1
    print(json.dumps({"counts": counts, "source_by_cat": source_by_cat, "standard_india": standard_india, "uncaught_cues": uncaught_cues, "form_tokens": form_tokens, "granule_detail": granule_detail, "top": by_cat}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
