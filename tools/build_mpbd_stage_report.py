from pathlib import Path
import csv, json
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path(r"D:\易迅数据\反倾销税深度分析报告\24_间苯氧基苯甲醛")
DOCX = OUT / "间苯氧基苯甲醛_反倾销税与第三国转运风险阶段审计报告.docx"
BLUE = RGBColor(46,116,181); DARK = RGBColor(31,77,120); GRAY = RGBColor(90,98,108); RED = RGBColor(155,28,28)

def set_font(run, size=11, bold=False, color=None, italic=False):
    run.font.name = "Calibri"; run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "等线")
    run.font.size = Pt(size); run.bold=bold; run.italic=italic
    if color: run.font.color.rgb=color

def shade(cell, fill):
    tcpr=cell._tc.get_or_add_tcPr(); shd=OxmlElement("w:shd"); shd.set(qn("w:fill"),fill); tcpr.append(shd)

def margins(cell, top=80, start=120, bottom=80, end=120):
    tcpr=cell._tc.get_or_add_tcPr(); node=tcpr.first_child_found_in("w:tcMar")
    if node is None: node=OxmlElement("w:tcMar"); tcpr.append(node)
    for k,v in (("top",top),("start",start),("bottom",bottom),("end",end)):
        x=node.find(qn("w:"+k))
        if x is None: x=OxmlElement("w:"+k); node.append(x)
        x.set(qn("w:w"),str(v)); x.set(qn("w:type"),"dxa")

def table(doc, headers, rows, widths):
    t=doc.add_table(rows=1,cols=len(headers)); t.style="Table Grid"; t.autofit=False
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.width=Inches(widths[i]); shade(c,"F2F4F7"); margins(c)
        c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(h),9.2,True,DARK)
    tr=t.rows[0]._tr.get_or_add_trPr(); x=OxmlElement("w:tblHeader"); x.set(qn("w:val"),"true"); tr.append(x)
    for row in rows:
        cs=t.add_row().cells
        for i,v in enumerate(row):
            cs[i].width=Inches(widths[i]); margins(cs[i]); cs[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cs[i].paragraphs[0]; p.paragraph_format.space_after=Pt(0)
            if i==0: p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            set_font(p.add_run(str(v)),8.9)
    return t

def heading(doc,text,level=1):
    p=doc.add_paragraph(style=f"Heading {level}"); p.paragraph_format.keep_with_next=True
    set_font(p.add_run(text),16 if level==1 else 13,True,BLUE if level<3 else DARK); return p

def para(doc,text,bold=False,color=None):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(6); p.paragraph_format.line_spacing=1.10
    set_font(p.add_run(text),11,bold,color); return p

def bullet(doc,text):
    p=doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after=Pt(6); p.paragraph_format.line_spacing=1.167
    set_font(p.add_run(text)); return p

def write_supporting_files():
    OUT.mkdir(parents=True,exist_ok=True)
    queries=[
      ["MPBD-Q1","中国进口主查询","HS 29124990；目的国China；近1/2/3年分别查；关键词3-PHENOXY BENZALDEHYDE、META PHENOXY BENZALDEHYDE、M-PHENOXY BENZALDEHYDE、MPBD、MPB、39515-51-0","逐页200条；记录总数、页数、末页、最新日期；同税号其他醛类排除"],
      ["MPBD-Q2","印度A腿","原产印度；目的国Vietnam/Singapore/UAE/Thailand/Malaysia/Indonesia分别查；同税号+同义词+5家印度实体","保存生产商、第三国收货人、数量、金额、日期、批号、225kg桶包装、提单/柜号"],
      ["MPBD-Q3","第三国B腿","上述第三国至China；同税号+同义词；按7/15/30/60/90日匹配A腿","同牌号/纯度/批号/包装/数量/主体/提单闭合后才升级为具体绕道证据"]]
    with (OUT/"间苯氧基苯甲醛_易迅最简查询组合.csv").open("w",encoding="utf-8-sig",newline="") as f:
        w=csv.writer(f); w.writerow(["查询ID","用途","条件","完整性要求"]); w.writerows(queries)
    gaps=[
      ["Y-01","已完成","3-PHENOXYBENZALDEHYDE、全球、近一年","页面可核验返回‘暂无数据’；仅代表该拼写和时间窗"],
      ["Y-02","未完成","CAS 39515-51-0、MPBD、META PHENOXY BENZALDEHYDE","CAS提交后页面超时，未取得可验证总数；不能写零记录"],
      ["Y-03","未完成","HS29124990中国进口宽池","需逐条穿透货描/CAS；同税号其他产品排除"],
      ["Y-04","未完成","印度至重点第三国A腿及第三国至中国B腿","当前无可闭合双段记录、提单号或柜号"]]
    with (OUT/"间苯氧基苯甲醛_数据覆盖与缺口.csv").open("w",encoding="utf-8-sig",newline="") as f:
        w=csv.writer(f); w.writerow(["编号","状态","覆盖条件","结论边界"]); w.writerows(gaps)
    sources=[
      ["商务部2024年第20号","现行措施、范围、税率、印度产能与复审参加方","https://www.mofcom.gov.cn/zcfb/zgdwjjmywg/art/2024/art_a53a9ebe96574af5bf1939e8a159d924.html"],
      ["中国贸易救济信息网案件页","案件节点、产品名称与HS","https://cacs.mofcom.gov.cn/cacscms/articleDetail/jkdc?articleId=180770&id=53d8a6e28e5ecfe20190051dc7f62024"],
      ["Heranba产品页","CAS、99%规格、225kg桶包装","https://www.heranba.co.in/intermediates/metaphenoxy-benzaldehyde/"],
      ["Heranba制造设施","印度Vapi/Sarigam生产网络","https://www.heranba.co.in/manufacturing-facilities/"],
      ["Bharat产品页","MPBD 99%/99.5%、CAS 39515-51-0","https://www.bharatgroup.co.in/bharat-rasayan/technical-intermediates.php"]]
    with (OUT/"间苯氧基苯甲醛_公开来源台账.csv").open("w",encoding="utf-8-sig",newline="") as f:
        w=csv.writer(f); w.writerow(["来源","用途","网址"]); w.writerows(sources)
    headers=["record_id","date","hs","description","buyer","supplier","weight","quantity","amount","destination","platform_origin","scope","route","evidence_grade","gap"]
    with (OUT/"间苯氧基苯甲醛_易迅逐票标准化_当前0条.csv").open("w",encoding="utf-8-sig",newline="") as f: csv.writer(f).writerow(headers)
    summary={"item":24,"product":"间苯氧基苯甲醛","hs":"29124990","cas":"39515-51-0","taxed_origin":"印度","measure_end":"2029-06-07","verified_yixun_query":{"keyword":"3-PHENOXYBENZALDEHYDE","date_window":"近一年（页面默认2025-08-11至2026-08-11）","scope":"全球","result":"暂无数据"},"other_yixun_queries":"未完成，不能视为0","closed_rerouting_chains":0,"risk_conclusion":"无印度—第三国—中国闭环证据；印度过剩产能与实体网络形成核查必要性","qa":{"no_claim_full_yixun":True,"no_tax_amount_without_customs_value":True,"no_entity_illegality_claim":True}}
    (OUT/"间苯氧基苯甲醛_阶段审计摘要.json").write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding="utf-8")

def main():
    write_supporting_files(); doc=Document(); sec=doc.sections[0]
    sec.page_width=Inches(8.5); sec.page_height=Inches(11); sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
    sec.header_distance=sec.footer_distance=Inches(.492)
    normal=doc.styles["Normal"]; normal.font.name="Calibri"; normal.font.size=Pt(11); normal._element.rPr.rFonts.set(qn("w:eastAsia"),"等线"); normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.10
    for i,(size,before,after,col) in enumerate(((16,16,8,BLUE),(13,12,6,BLUE),(12,8,4,DARK)),1):
        s=doc.styles[f"Heading {i}"]; s.font.name="Calibri"; s._element.rPr.rFonts.set(qn("w:eastAsia"),"等线"); s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=col; s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after); s.paragraph_format.keep_with_next=True
    h=sec.header.paragraphs[0]; set_font(h.add_run("反倾销税深度分析 | 第24项"),9,False,GRAY)
    f=sec.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.RIGHT; set_font(f.add_run("间苯氧基苯甲醛阶段审计 | 2026-08-13"),9,False,GRAY)
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(12); p.paragraph_format.space_after=Pt(4); set_font(p.add_run("间苯氧基苯甲醛反倾销税与第三国转运风险"),23,True)
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(14); set_font(p.add_run("第24项｜易迅数据与公开来源阶段审计报告"),14,False,GRAY)
    meta=[("商品","间苯氧基苯甲醛（MPB/MPBD）"),("CAS / 中国税号","39515-51-0 / 29124990"),("受税来源","印度"),("措施期限","2024-06-08起续征5年，预计至2029-06-07"),("审计状态","公开政策已核；易迅仅完成1组可验证查询，别名/税号及A-B腿未全量完成")]
    for k,v in meta:
        p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(2); set_font(p.add_run(k+"："),11,True); set_font(p.add_run(v))
    para(doc,"核心结论：当前没有取得能够闭合‘印度生产/出口—第三国收货或加工—中国进口’的双段贸易、提单或报关证据，不能认定具体企业绕道或逃税。易迅对英文全称的近一年全球查询可核验显示‘暂无数据’，但CAS、别名和HS宽池未全部完成，因此不能把单一零结果外推为全库无记录。",True,RED)

    heading(doc,"一、现行反倾销措施",1)
    para(doc,"商务部2024年第20号公告决定，自2024年6月8日起继续对原产于印度的进口间苯氧基苯甲醛征收反倾销税5年。产品英文名包括Meta Phenoxy Benzaldehyde、M-Phenoxy Benzaldehyde、3-Phenoxy Benzaldehyde，简称MPB或MPBD；归入税则号29124990，但该税号项下其他产品不在措施范围。")
    table(doc,["印度生产/出口企业","AD税率","AD及其13%进口增值税增量系数"],[
      ("Hemani Industries Limited（赫曼尼）","36.4%","41.132% × V"),("Gujarat Insecticides Limited","52.0%","58.760% × V"),("Bharat Rasayan Limited","56.4%","63.732% × V"),("其他印度公司（含未单列企业）","56.9%","64.297% × V")],[3.4,1.2,1.9])
    para(doc,"注：V仅指中国海关审定完税价格。综合系数=AD税率×1.13，只用于筛查‘反倾销税+由反倾销税增加的进口增值税差额’；不等于整票总税负，也不能用易迅出口侧金额直接计算实际欠税。",False,GRAY)

    heading(doc,"二、商品识别与范围边界",1)
    bullet(doc,"关键识别：CAS 39515-51-0；分子式C13H10O2；通常为无色或淡黄色透明液体；Heranba公开规格为MPBD含量≥99%。")
    bullet(doc,"包装指纹：Heranba公开技术级包装为225kg UN认可漆面钢桶或HMHDPE桶，可用于匹配A/B腿重量倍数、桶数和批次。")
    bullet(doc,"用途：是氯氰菊酯、高效氯氰菊酯、功夫菊酯、氰戊菊酯、甲氰菊酯等拟除虫菊酯原药的中间体。下游农药制剂、MPB醇、其他同税号醛类不能直接并入。")

    heading(doc,"三、易迅数据覆盖与完整性",1)
    table(doc,["查询","状态","可下结论"],[
      ("3-PHENOXYBENZALDEHYDE｜全球｜近一年","已完成","页面明确显示‘暂无数据’；仅限该拼写和时间窗"),("39515-51-0 / MPBD / META PHENOXY BENZALDEHYDE","未完成","CAS查询提交后页面超时，未取得总数或末页"),("HS29124990中国进口宽池","未完成","必须逐条核CAS/货描，不能将同税号全算MPBD"),("印度→第三国A腿、第三国→中国B腿","未完成","当前无同批次、提单、柜号或报关闭环")],[2.5,1.0,3.0])
    para(doc,"完整性判断：本地D盘未发现该商品专门原始下载文件；当前逐票标准化表为0条，只表示‘现有可用逐票数据为0’，不表示易迅全库、全球贸易或中国进口为0。")

    heading(doc,"四、印度供给压力与可核查实体",1)
    para(doc,"商务部复审数据显示，印度MPBD产能从2018年的17,540吨增至2022年的29,980吨；2022年产量17,089吨、闲置产能12,891吨、可供出口产能15,344吨。官方据此认定印度存在显著过剩产能和出口依赖。供给压力说明需要持续核查，但不是第三国绕道证据。")
    table(doc,["优先级","实体/线索","证据与核查方向"],[
      ("高","Hemani Industries Limited","复审唯一提交完整答卷并在调查期直接向中国非关联客户销售；优先调中国进口人、税单和原产申报"),("高","Bharat Rasayan Limited","官网列MPBD 99%/99.5%及CAS；适用56.4%；核印度出口、第三国收货人与225kg桶包装"),("高","Gujarat Insecticides Limited","列名52.0%；参加复审并主张调查期无对华出口；措施后出现第三国B腿时优先穿透生产商"),("中高","Heranba Industries Limited","官网确认在印度制造设施生产MPBD；未单列，原则上适用其他印度公司56.9%"),("中高","Tagros Chemicals India / 万民利有机物","参加复审并主张调查期无对华出口；属于实体调单线索，不能据此认定转运")],[.7,2.3,3.5])

    heading(doc,"五、第三国绕道证据评估",1)
    para(doc,"公开检索未发现商务部反规避裁定、海关处罚、法院判决或公开双段提单，能够证明印度产MPBD经越南、新加坡、阿联酋、泰国或马来西亚换单、换标或伪报原产后进入中国。当前证据等级为‘结构性高关注、具体链路未证实’。")
    bullet(doc,"应升级为B+线索的条件：印度列名/已知生产商向第三国发运MPBD，随后7—90日内第三国主体向中国发运同CAS、同纯度、同包装倍数，且收发货人、批号或提单存在交叉。")
    bullet(doc,"应升级为A级具体证据的条件：中国进口报关原产申报非印度，但COA/厂家声明/批号或印度A腿证明实际印度生产，同时税款书未按相应印度企业税率征税。")
    bullet(doc,"不能单独作为证据：第三国开票、经中转港、印度企业具有海外客户、同类重量相近、企业未参加复审、贸易量异常增长。")
    bullet(doc,"若第三国声称当地生产，必须核设备、工艺、BOM、能耗、环保许可、工单、收率与库存物料平衡；仅仓储、分装、换桶或换标不改变真实原产地。")

    heading(doc,"六、下一轮易迅全页查询方案",1)
    para(doc,"MPBD-Q1：中国进口主查询。HS29124990、目的国China，分别按近1年/2年/3年；逐一查询英文全称、CAS 39515-51-0、MPBD、MPB。每页200条，必须读取全部页并保留总数、页数、末页和最新日期。")
    para(doc,"MPBD-Q2：印度A腿。印度至越南、新加坡、阿联酋、泰国、马来西亚、印度尼西亚分别查；叠加Hemani、Bharat Rasayan、Gujarat Insecticides、Heranba、Tagros实体名。")
    para(doc,"MPBD-Q3：第三国B腿。上述第三国至中国按同关键词/税号查询；以7/15/30/60/90日、225kg桶数、纯度、批次、收发货人、提单/柜号和付款受益人进行双腿匹配。不能在发现个别异常后停止，必须覆盖全部结果。")

    heading(doc,"七、需要调取的中国端材料",1)
    table(doc,["材料","必核字段"],[
      ("进口报关单","中国10位商品编号、原产国/地区、启运国、生产商、境外发货人、境内收货人、消费使用单位、申报企业、口岸、净重、完税价格"),("税款单","适用AD企业税率、反倾销税、进口增值税及是否已正常征收"),("货运与产地","主/分提单、箱号/封志、船名航次、转运港、原产地证、COA、厂家声明、批号、包装照片"),("第三国加工","进口原料、BOM、设备/工艺、工单、能耗、产量、库存、销售、增值和投入产出平衡")],[1.4,5.1])

    heading(doc,"八、阶段结论",1)
    para(doc,"本项风险来自三方面：印度显著过剩产能、MPBD高税率、下游拟除虫菊酯供应链对该中间体的持续需求。最值得优先核查的是Hemani直接对华链以及Bharat、Gujarat、Heranba、Tagros可能经第三国销售的A/B腿。但是，截至本报告，没有任何可认定的第三国绕道数量、实际少缴税额、国内进口企业、报关行或高风险口岸。待取得中国报关完税价格和税款书后，才能按企业税率测算。",True)

    heading(doc,"九、主要来源",1)
    for s in [
      "商务部公告2024年第20号：https://www.mofcom.gov.cn/zcfb/zgdwjjmywg/art/2024/art_a53a9ebe96574af5bf1939e8a159d924.html",
      "中国贸易救济信息网案件页：https://cacs.mofcom.gov.cn/cacscms/articleDetail/jkdc?articleId=180770&id=53d8a6e28e5ecfe20190051dc7f62024",
      "Heranba MPBD产品页：https://www.heranba.co.in/intermediates/metaphenoxy-benzaldehyde/",
      "Heranba制造设施：https://www.heranba.co.in/manufacturing-facilities/",
      "Bharat Rasayan中间体产品页：https://www.bharatgroup.co.in/bharat-rasayan/technical-intermediates.php"]: para(doc,s,False,GRAY)
    doc.save(DOCX); print(DOCX)

if __name__=="__main__": main()
