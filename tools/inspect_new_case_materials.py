"""Read-only inspection of the GDRM26-394/395 source files and case DB."""

from __future__ import annotations

import hashlib
import json
import sqlite3
import sys
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader


sys.stdout.reconfigure(encoding="utf-8", errors="replace")
PASSWORD = "GDRM@2026"
ROOT = Path(r"D:\codex\XG执法\2026_案例库")
DB = ROOT / "enforcement_cases_2026.sqlite"
FILES = [
    Path(r"E:\ZMJ\GDRM26-395情報159號兩宗海路貨運走私廢金屬案件.pdf"),
    Path(r"E:\ZMJ\GDRM26-394附件一及二.pdf"),
    Path(r"E:\ZMJ\GDRM26-394情報158號一宗河路貨運走私未列艙單貨物案件.pdf"),
    Path(r"E:\ZMJ\GDRM26-395附件三涉案報關紀錄.xlsx"),
    Path(r"E:\ZMJ\GDRM26-395附件一及二.pdf"),
]


def digest(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        while chunk := stream.read(1024 * 1024):
            h.update(chunk)
    return h.hexdigest()


results = []
for path in FILES:
    item = {
        "file": str(path),
        "exists": path.exists(),
        "bytes": path.stat().st_size if path.exists() else None,
        "sha256": digest(path) if path.exists() else None,
    }
    if not path.exists():
        results.append(item)
        continue
    if path.suffix.lower() == ".pdf":
        reader = PdfReader(str(path))
        item["encrypted"] = reader.is_encrypted
        if reader.is_encrypted:
            item["decrypt_result"] = int(reader.decrypt(PASSWORD))
        texts = [page.extract_text() or "" for page in reader.pages]
        item["pages"] = len(texts)
        item["characters"] = sum(map(len, texts))
        item["preview"] = " ".join(texts[:2])[:800]
    else:
        book = load_workbook(path, read_only=True, data_only=True)
        item["sheets"] = []
        for sheet in book.worksheets:
            nonempty = 0
            preview = []
            for row in sheet.iter_rows(values_only=True):
                values = [str(value) for value in row if value is not None]
                if values:
                    nonempty += 1
                    if len(preview) < 4:
                        preview.append(values)
            item["sheets"].append({
                "title": sheet.title,
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
                "nonempty_rows": nonempty,
                "preview": preview,
            })
    results.append(item)

conn = sqlite3.connect(DB)
schema = {
    "tables": [row[0] for row in conn.execute(
        "SELECT name FROM sqlite_master WHERE type='table' ORDER BY name"
    )],
    "cases_columns": conn.execute("PRAGMA table_info(cases)").fetchall(),
    "documents_columns": conn.execute("PRAGMA table_info(documents)").fetchall(),
    "existing_cases": conn.execute(
        "SELECT * FROM cases WHERE case_id IN ('GDRM26-394','GDRM26-395')"
    ).fetchall(),
    "existing_hashes": conn.execute(
        "SELECT case_id,relative_path,sha256 FROM documents WHERE sha256 IN (?,?,?,?,?)",
        tuple(item["sha256"] for item in results),
    ).fetchall(),
    "counts": {
        "cases": conn.execute("SELECT count(*) FROM cases").fetchone()[0],
        "documents": conn.execute("SELECT count(*) FROM documents").fetchone()[0],
    },
}
conn.close()
print(json.dumps({"files": results, "database": schema}, ensure_ascii=False, indent=2))
