from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(r"D:\易迅数据\反倾销税深度分析报告\23_苯乙烯")
DOCX = OUT / "苯乙烯_反倾销税与第三国转运风险_数据缺口阶段报告.docx"
BLUE = RGBColor(46, 116, 181)
DARK = RGBColor(31, 77, 120)
GRAY = RGBColor(90, 98, 108)
RED = RGBColor(155, 28, 28)


def font(run, size=11, bold=False, color=None, italic=False):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "等线")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def cell_margin(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v)); node.set(qn("w:type"), "dxa")


def add_table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"; t.autofit = False
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; c.width = Inches(widths[i]); shade(c, "F2F4F7")
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h); font(r, 9.5, True, DARK)
    trPr = t.rows[0]._tr.get_or_add_trPr(); rep = OxmlElement("w:tblHeader"); rep.set(qn("w:val"), "true"); trPr.append(rep)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].width = Inches(widths[i]); cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; cell_margin(cells[i])
            p = cells[i].paragraphs[0]; p.paragraph_format.space_after = Pt(0)
            if i == 0: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(v)); font(r, 9)
    return t


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text); font(r, 16 if level == 1 else 13, True, BLUE if level < 3 else DARK)
    return p


def para(doc, text, bold_lead=None, color=None):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6); p.paragraph_format.line_spacing = 1.10
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead); font(r, 11, True, color)
        r = p.add_run(text[len(bold_lead):]); font(r, 11, False, color)
    else:
        r = p.add_run(text); font(r, 11, False, color)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after = Pt(5); p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text); font(r)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5); sec.page_height = Inches(11)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
    sec.header_distance = sec.footer_distance = Inches(0.492)
    styles = doc.styles
    normal = styles["Normal"]; normal.font.name = "Calibri"; normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
    normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.10
    for i, (size, before, after, col) in enumerate(((16,16,8,BLUE),(13,12,6,BLUE),(12,8,4,DARK)),1):
        s=styles[f"Heading {i}"]; s.font.name="Calibri"; s._element.rPr.rFonts.set(qn("w:eastAsia"),"等线"); s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=col
        s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after); s.paragraph_format.keep_with_next=True
    header = sec.header.paragraphs[0]; header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r=header.add_run("反倾销税深度分析 | 第23项"); font(r,9,False,GRAY)
    footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r=footer.add_run("苯乙烯风险阶段审计 | 2026-08-13"); font(r,9,False,GRAY)

    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(12); p.paragraph_format.space_after=Pt(4)
    r=p.add_run("苯乙烯反倾销税与第三国转运风险"); font(r,23,True,RGBColor(0,0,0))
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(14)
    r=p.add_run("第23项｜数据缺口阶段审计报告"); font(r,14,False,GRAY)
    for k,v in [("商品","苯乙烯（Styrene / Styrene Monomer / SM）"),("中国税号","29025000；CAS 100-42-5"),("受税来源","韩国、台湾地区、美国"),("审计日期","2026年8月13日"),("状态","公开政策已核；易迅逐票数据未完成，严禁据此认定无风险")]:
        p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(2)
        r=p.add_run(k+"："); font(r,11,True); r=p.add_run(v); font(r)
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(10); p.paragraph_format.space_after=Pt(10)
    r=p.add_run("核心结论：截至本次审计，没有取得可闭合的“受税来源—第三国—中国”双段贸易证据。网页检索持续超时，当前只能交付政策、公开产能反证、核查路径与数据缺口，不能把未查到写成不存在。新加坡和印尼均有真实苯乙烯生产，第三国发货本身不构成逃税证据。"); font(r,11,True,RED)

    heading(doc,"一、现行反倾销措施",1)
    para(doc,"商务部公告2024年第24号自2024年6月23日起继续对原产于韩国、台湾地区和美国的进口苯乙烯征收反倾销税，期限5年，预计至2029年6月22日。产品包括Styrene、Styrene Monomer（SM）、Phenylethylene等，归入29025000。反倾销税按海关审定完税价格从价计征；进口增值税计税基础包含反倾销税。")
    add_table(doc,["来源/企业","AD税率","AD+增值税差额系数"],[
        ("韩国：韩华道达尔、丽川NCC","6.2%","7.006%V"),("韩国：LG、SK","6.6%","7.458%V"),("韩国：乐天/其他","7.5%","8.475%V"),
        ("台湾化纤","3.8%","4.294%V"),("台湾其他","4.2%","4.746%V"),
        ("美国华美/Westlake","13.7%","15.481%V"),("美国列名13.9%企业","13.9%","15.707%V"),("美国其他","55.7%","62.941%V")],[3.6,1.2,1.7])
    para(doc,"注：V仅指中国海关审定完税价格；综合系数=AD税率×1.13，仅用于筛查潜在少缴反倾销税及其引致的13%进口增值税差额，不是整票进口总税负，也不是实际欠税额。", color=GRAY)

    heading(doc,"二、易迅数据覆盖与完整性",1)
    para(doc,"本地盘点：D盘未发现苯乙烯专项原始页面、XLSX、CSV或JSON逐票数据。")
    para(doc,"网页状态：已登录并打开易迅全球搜索页，已尝试HS290250+目的国China检索；搜索提交或结果读取持续超时，未能取得可靠总条数、页数、最新日期、末页或任何逐票记录。")
    para(doc,"完整性结论：本项当前确认易迅逐票记录0条仅表示“本地可用数据为0”，不表示易迅网页为0，更不表示贸易为0。报告不得使用“全量已查”“未发现第三国记录”等措辞。", color=RED)

    heading(doc,"三、公开第三国风险与合法替代解释",1)
    bullet(doc,"新加坡：Shell官方确认Jurong Island生产苯乙烯单体；BASF公开资料显示ELLBA Eastern苯乙烯年产能55万吨。因此新加坡生产商、发货人或装港均可能对应真实新加坡原产。")
    bullet(doc,"印度尼西亚：Chandra Asri/子公司PT Styrindo Mono Indonesia公开确认在印尼生产苯乙烯单体。印尼对华发货不能仅凭路线判为换产地。")
    bullet(doc,"品牌与销售实体不是生产地。INEOS、Shell、LyondellBasell等跨国销售公司可能从不同国家工厂供货，必须穿透COA的plant code、生产批次和原产地证。")
    bullet(doc,"苯乙烯为散装危险液体，UN2055，常经化学品船舶与罐区调拨。普通换单、分提单、仓储或商业转售不改变原产地；若原货仅在第三国储运、分装或换标，仍需按真实生产地核反倾销税。")
    para(doc,"公开负面检索：未检出中国商务部反规避裁定、海关处罚、法院判决或公开双段提单，能够证明现行措施下某一韩国、台湾或美国苯乙烯批次经第三国伪报后进入中国。1995年韩国苯乙烯涉嫌走私的旧合同案早于现行措施，不能作为现行第三国绕道证据。")

    heading(doc,"四、应优先核查的路线与实体",1)
    add_table(doc,["优先级","路线/实体","核查重点"],[
        ("高","韩国/台湾/美国生产商→新加坡、越南、马来西亚、印度→中国","同船名航次、7—30日、同量级、共同收货人/开票人；第三国是否仅罐区换单"),
        ("高","受税生产商的亚洲销售公司→中国","销售主体与COA生产厂是否不一致；中国报关是否申报真实原产"),
        ("中","新加坡Shell/ELLBA、印尼Styrindo→中国","先按真实第三国产能反证核验；取得plant code、COA、原产证"),
        ("中","大型化学品码头/罐区分拨","船对船转运、混批、换提单、罐区库存流水与数量平衡")],[.8,2.6,3.1])
    para(doc,"受税实体名称检索池：HANWHA TOTAL PETROCHEMICAL、YECHUN NCC、LOTTE CHEMICAL、LG CHEM、SK global chemical/SK geo centric、FORMOSA CHEMICALS & FIBRE、LYONDELL CHEMICAL、WESTLAKE STYRENE、INEOS STYROLUTION AMERICA、AMERICAS STYRENICS。名称命中只用于调单，不等同原产认定。")

    heading(doc,"五、三组易迅查询与逐票判定标准",1)
    para(doc,"STY-Q1 中国进口：HS290250/29025000；目的国China；关键词STYRENE MONOMER、STYRENE、PHENYLETHYLENE、VINYLBENZENE、ETHENYLBENZENE、100-42-5；200条/页，读取全部页。")
    para(doc,"STY-Q2 A腿：韩国、台湾地区、美国至新加坡、印尼、马来西亚、越南、印度、泰国、阿联酋、荷兰；同税号、关键词及列名生产商；记录COA、UN2055、船名航次、装港、卸港、储罐和收货人。")
    para(doc,"STY-Q3 B腿：上述第三国至China；按7/15/30/60日、同船/驳船、数量、纯度、抑制剂、批次、共同主体匹配。每个查询均需保存总数、页数、末页和覆盖日期，不能发现一条异常后停止。")
    add_table(doc,["闭环字段","需要的证据"],[
        ("中国申报","原产国/地区、生产商、境外发货人、进口人、报关企业、口岸、贸易方式、净重、完税价格、AD税率和缴款书"),
        ("货运","船名航次、装卸港、母船/驳船、主/分提单、舱单、罐区进出库、换单和船对船记录"),
        ("产品","COA、CAS100-42-5、纯度、抑制剂、UN2055、批号、plant code、SDS和厂家声明"),
        ("商业","合同、发票、付款受益人、第三国贸易商、保险、信用证、关联交易和价格链")],[1.3,5.2])

    heading(doc,"六、阶段风险判断",1)
    para(doc,"风险等级：数据缺口阶段，暂不作商品实体风险评分。结构上，韩国/台湾/美国至亚洲罐区再对华的散装液体链值得高优先核查；但新加坡和印尼真实产能是强反证。")
    para(doc,"证据边界：当前没有A/B腿闭环、没有中国进口报关原产字段、没有税款书，因此不能写“涉及逃税数量”“实际少缴税额”“具体企业走私”或“具体口岸高风险”。只有取得中国报关完税价格和适用企业税率后，才可按V×AD率及AD×13%计算税差。")

    heading(doc,"七、来源",1)
    for txt in [
        "商务部公告2024年第24号：https://www.mofcom.gov.cn/zcfb/blgg/gg/2024/art/2024/art_37cbd5e452754ce9876065eaa25f9468.html",
        "台湾化纤苯乙烯产品页：https://en.fcfc.com.tw/styrene-monomer",
        "Shell Jurong Island：https://www.shell.com.sg/about-us/what-we-do/projects-and-sites/shell-jurong-island.html",
        "BASF ELLBA Eastern公告：https://www.basf.com/dam/jcr%3A74fc24af-55c7-3221-a40b-82025714dfaa/basf/www/global/documents/en/investor-relations/basf-at-a-glance/strategy/portfolio-optimization/press-releases/P437e_BASF-to-sell-shares-in-ELLBA-Eastern.pdf",
        "Chandra Asri苯乙烯产品页：https://chandra-asri.com/en/our-business/chemical-solutions/styrene-monomer",
    ]:
        para(doc,txt,color=GRAY)
    doc.save(DOCX)
    print(DOCX)


if __name__ == "__main__":
    main()
