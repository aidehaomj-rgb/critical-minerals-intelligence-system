from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


DIR = Path(r"D:\易迅数据\反倾销税深度分析报告\06_共聚聚甲醛（POM）")
SUMMARY = json.loads((DIR / "POM_全量分析结果.json").read_text("utf-8"))
OUT = DIR / "共聚聚甲醛（POM）_反倾销税与第三国转运风险深度分析报告.docx"

NAVY = "17324D"
BLUE = "1F4E78"
INK = "102A43"
MUTED = "66788A"
LIGHT_BLUE = "EAF2F8"
LIGHT_GRAY = "F2F4F7"
PALE_GOLD = "FFF4CC"
PALE_RED = "FCE8E6"
PALE_GREEN = "E6F4EA"
BORDER = "CBD5E1"
RED = "9B1C1C"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_run_font(run, size=None, color=None, bold=None, italic=None, name="Calibri"):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill: str):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    node = trpr.find(qn("w:tblHeader"))
    if node is None:
        node = OxmlElement("w:tblHeader")
        trpr.append(node)
    node.set(qn("w:val"), "true")


def set_cell_text(cell, value, bold=False, color="000000", size=9.2):
    cell.text = "" if value is None else str(value)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.05
        for run in paragraph.runs:
            set_run_font(run, size=size, color=color, bold=bold)


def set_table_borders(table, color=BORDER, size="4"):
    tblpr = table._tbl.tblPr
    borders = tblpr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblpr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa):
    table.autofit = False
    tbl = table._tbl
    tblpr = tbl.tblPr
    for tag in ("w:tblW", "w:tblInd", "w:tblLayout", "w:tblCellMar"):
        old = tblpr.find(qn(tag))
        if old is not None:
            tblpr.remove(old)
    tblw = OxmlElement("w:tblW")
    tblw.set(qn("w:w"), "9360")
    tblw.set(qn("w:type"), "dxa")
    tblpr.append(tblw)
    tblind = OxmlElement("w:tblInd")
    tblind.set(qn("w:w"), "120")
    tblind.set(qn("w:type"), "dxa")
    tblpr.append(tblind)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr.append(layout)
    margins = OxmlElement("w:tblCellMar")
    for tag, val in (("top", 70), ("bottom", 70), ("start", 100), ("end", 100)):
        node = OxmlElement(f"w:{tag}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    tblpr.append(margins)
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths_dxa):
            tcpr = cell._tc.get_or_add_tcPr()
            tcw = tcpr.find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                tcpr.append(tcw)
            tcw.set(qn("w:w"), str(width))
            tcw.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)


def add_table(doc, headers, rows, widths, header_fill=LIGHT_BLUE, size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for idx, text in enumerate(headers):
        set_cell_text(table.cell(0, idx), text, bold=True, color=INK, size=size)
        set_cell_shading(table.cell(0, idx), header_fill)
    set_repeat_header(table.rows[0])
    for vals in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(vals):
            set_cell_text(cells[idx], value, size=size)
    set_table_geometry(table, widths)
    set_table_borders(table)
    return table


def add_callout(doc, title, text, fill=PALE_GOLD, title_color=INK):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, size=10.5, color=title_color, bold=True)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.12
    r2 = p2.add_run(text)
    set_run_font(r2, size=10.2, color="333333")
    set_table_geometry(table, [9360])
    set_table_borders(table, color=BORDER, size="6")
    set_repeat_header(table.rows[0])
    return table


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), "Calibri")
    rfonts.set(qn("w:hAnsi"), "Calibri")
    rfonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    rpr.extend([rfonts, color, underline])
    run.append(rpr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_source(doc, label, url, note=""):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    add_hyperlink(p, label, url)
    if note:
        r = p.add_run(f"：{note}")
        set_run_font(r, size=9.3, color=MUTED)


def add_bullet(doc, text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.08
    r = p.add_run(text)
    set_run_font(r, size=10.2, color="222222")
    return p


def add_number(doc, text, level=0):
    style = "List Number" if level == 0 else "List Number 2"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.08
    r = p.add_run(text)
    set_run_font(r, size=10.2, color="222222")
    return p


def add_para(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.13
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, size=10.3, color=INK, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, size=10.3, color="222222")
    else:
        r = p.add_run(text)
        set_run_font(r, size=10.3, color="222222")
    return p


doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.72)
sec.bottom_margin = Inches(0.72)
sec.left_margin = Inches(0.82)
sec.right_margin = Inches(0.82)

styles = doc.styles
styles["Normal"].font.name = "Calibri"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
styles["Normal"].font.size = Pt(10.3)
for name, size, color in (("Title", 27, NAVY), ("Heading 1", 17, NAVY), ("Heading 2", 13.5, BLUE), ("Heading 3", 11.5, INK)):
    style = styles[name]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(size)
    style.font.color.rgb = rgb(color)
    style.font.bold = True

# 页眉页脚
header = sec.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hr = header.add_run("反倾销税风险穿透分析｜POM｜2026-08-13")
set_run_font(hr, size=8.5, color=MUTED)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = footer.add_run("内部研判材料｜交易数据为核查线索，不构成违法认定")
set_run_font(fr, size=8.5, color=MUTED)

# 封面
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(36)
r = p.add_run("反倾销税风险穿透分析")
set_run_font(r, size=12, color=BLUE, bold=True)
p = doc.add_paragraph()
p.style = styles["Title"]
p.paragraph_format.space_before = Pt(14)
p.paragraph_format.space_after = Pt(8)
r = p.add_run("共聚聚甲醛（POM）")
set_run_font(r, size=27, color=NAVY, bold=True)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(22)
r = p.add_run("中国反倾销税、第三国转运与原产地申报风险深度分析报告")
set_run_font(r, size=17, color=INK, bold=True)

add_table(doc, ["项目", "内容"], [
    ["措施来源", "美国、欧盟、台湾地区、日本"],
    ["临时措施", "2025年1月24日起"],
    ["终裁实施", "2025年5月19日起，期限5年"],
    ["易迅窗口", "2024年8月6日至2026年8月6日"],
    ["逐票覆盖", "6,162个页面行；跨查询精确唯一4,179条；中国大陆目的地3,685条"],
    ["报告日期", "2026年8月13日"],
], [2300, 7060], header_fill=LIGHT_GRAY, size=10)

doc.add_paragraph()
add_callout(
    doc,
    "结论先行",
    "当前未发现能闭合证明受税来源POM经第三国物理转运、改报第三国原产并逃缴中国反倾销税的证据。发现两条需优先核查的具体交易线索：一是Acumen同批HOSTAFORM C52021的Germany/Philippines原产字段切换；二是TENAC-C EX352由越南发往中国但货描保留#&JP。另有沙特来源措施后集中放量、欧盟/美国受税来源直接对华核税等高优先级问题。LW15EWX 25kg因蜡改性和熔点约173℃，初步不属于本案范围，必须从“疑似逃税量”中剔除。",
    fill=PALE_GOLD,
)

doc.add_page_break()

doc.add_heading("一、执行摘要", level=1)
add_para(doc, "本次排查同时覆盖易迅税号、通用商品词、核心品牌和共聚描述词。所有页面均切换为每页200条后逐页提取到底，未在发现异常后提前停止。6,162个页面行经跨查询精确去重后形成4,179条唯一记录，并在逐票台账中逐条标注商品范围、措施阶段、路线风险、证据级别、条件税款和下一步动作。")

add_table(doc, ["结论", "数量/事实", "证据强度", "判断"], [
    ["第三国逃税闭环", "0条", "未证实", "无同批A腿+B腿、无中国改报第三国原产和税单缺口闭环"],
    ["Acumen字段切换", "保守1个商业事件，5,000kg；高值情景2票/10,000kg", "B+", "同货物关键字段一致，仅日期/原产Germany与Philippines不同"],
    ["TENAC-C日本标记B腿", "50kg", "B+", "越南出口货描#&JP，平台原产Vietnam；共聚范围待COA"],
    ["沙特异常放量", "HS口径11行/959,650kg；POM口径另见7行/201,320kg", "B", "需核真实生产厂；不同查询口径可能重叠，不合并加总"],
    ["直接受税来源", "欧盟POM PULVER终裁后47行/约3,511.5t；美国CELCON标准料约716.61t", "A-/B+", "核心为范围和已缴税核验，不是绕道"],
    ["LW15EWX纠偏", "25kg（同票另有25kg CELCON及4kg色母）", "范围反证", "LW15EWX初步范围外，不计算POM反倾销税差"],
], [1900, 2500, 1200, 3760], size=9.1)

add_callout(doc, "表述边界", "本报告中的“高风险”“异常”“线索”均指向进一步核验优先级，不是对企业违法、走私或逃税的认定。只有取得中国进口报关单、法定原产地、生产商、COA/批号、原产地证和税款缴款书，才能核实是否少缴税。", fill=PALE_RED, title_color=RED)

doc.add_heading("二、措施范围、税率与税种", level=1)
add_para(doc, "商务部公告2025年第25号自2025年5月19日起对原产于美国、欧盟、台湾地区和日本的进口共聚聚甲醛征收反倾销税，期限5年。2025年1月24日至5月18日为临时措施保证金期，保证金按终裁范围及税率转为反倾销税。")
add_para(doc, "HS 39071010、39071090只是检索入口。商品还必须同时满足公告中的化学结构和性能指标，包括熔融温度160≤T＜170℃、密度1.38—1.43g/cm³等；均聚POM、改性POM及同税号其他产品不在范围内。公告没有规定CAS号，24969-26-4只能作为补充检索词，不能单独认定范围。")

add_table(doc, ["来源/企业", "终裁税率", "每100万元计税价格：AD+VAT增量", "核验重点"], [
    ["美国：Ticona及其他美国公司", "74.9%", "84.637万元", "生产商、规格和是否已按74.9%缴税"],
    ["欧盟：Celanese Germany及其他欧盟公司", "34.5%", "38.985万元", "欧盟生产厂、关联贸易链与税款缴款书"],
    ["台湾：Polyplastics Taiwan", "3.8%", "4.294万元", "真实生产商与工厂"],
    ["台湾：Formosa Plastics", "4.0%", "4.520万元", "真实生产商与工厂"],
    ["台湾：其他公司", "32.6%", "36.838万元", "是否误用个别低税率"],
    ["日本：Polyplastics及其他日本公司", "35.5%", "40.115万元", "日本/马来西亚/台湾真实生产地"],
    ["日本：Asahi Kasei", "24.5%", "27.685万元", "TENAC-C生产批号与工厂"],
], [2650, 1350, 2400, 2960], size=9.1)
add_para(doc, "税款情景只计算反倾销税及其导致的进口增值税增量：反倾销税=海关完税价格×适用税率；VAT增量=反倾销税×13%。正常关税和本来就应缴的基础进口增值税不能写成“逃税额”。易迅金额字段未标币种时，本报告仅写“同金额字段币种”，不擅自换算成美元或人民币。")

doc.add_heading("三、易迅全量检索与数据质量", level=1)
add_table(doc, ["查询", "原始行", "页数/状态", "关键清洗结果"], [
    ["HS 390710→China", "1,398", "7页，全页完成", "完全重复276；唯一1,122；中国大陆唯一1,080；终裁后363条"],
    ["POM→China", "2,920", "15页，全页完成", "台湾误纳291；中国大陆精确唯一2,127；折叠提单重量版本后2,094"],
    ["HOSTAFORM→China", "443", "3页，全页完成", "精确重复冗余280；精确唯一163；保守100个商业记录组"],
    ["DELRIN→China", "1,028", "6页，全页完成", "用于识别并排除均聚POM及追踪第三国库存链"],
    ["DURACON/TENAC/CELCON", "75/3/294", "1/1/2页，全页完成", "覆盖日本宝理、旭化成、Celanese共聚牌号"],
    ["共聚描述词/CAS", "1/0/0", "均完成", "POLYACETAL COPOLYMER 1条；POLYOXYMETHYLENE COPOLYMER、24969-26-4均0条"],
    ["跨查询总计", "6,162", "全部页面已保存", "精确唯一4,179；中国大陆3,685；查询内出现次数保留在台账"],
], [2350, 1100, 2050, 3860], size=9.1)

add_bullet(doc, "易迅“目的国China”会纳入台湾目的地记录；本报告逐票按目的地字段剔除台湾及其他非中国大陆记录。")
add_bullet(doc, "越南数据货描中的#&DE/#&JP/#&CN等通常是出口申报原产代码，而平台“原产国”常接近报告国/发运国；两字段不一致是核单线索，不等于中国进口端伪报。")
add_bullet(doc, "环球提单同一长货描可能按集装箱重量生成多个展示版本；直接相加会严重虚增，报告对重点数量采用商业组、内文净重或较低保守值。")
add_bullet(doc, "数据源更新截止日期不一致；“未发现”只表示已检索窗口和平台覆盖内未发现，不能外推为全市场不存在。")

doc.add_heading("四、最强线索：Acumen C52021原产字段切换", level=1)
add_table(doc, ["字段", "2025-07-21记录", "2025-07-23记录"], [
    ["数据源", "菲律宾新版", "菲律宾新版"],
    ["商品", "HOSTAFORM C52021 NATURAL", "HOSTAFORM C52021 NATURAL"],
    ["买方", "ACUMEN ENGINEERING SHANGHAI CO LTD", "同左"],
    ["卖方", "ACUMEN ENGINEERING PTE LTD", "同左"],
    ["重量/数量", "5,000kg / 200", "同左"],
    ["金额字段", "9,000", "9,000"],
    ["平台原产字段", "Philippines", "Germany"],
    ["范围", "C52021为Standard Unfilled，公开熔点约166℃；初步在范围，仍须本批COA", "同左"],
], [2200, 3580, 3580], size=9.2)

add_para(doc, "行22、23属于2025-07-23同一记录的完全重复展示，必须合并。2025-07-21记录除日期和原产字段外，其余关键字段完全相同，更像同一票修订、重报或数据映射版本。保守口径按一个商业事件5,000kg；只有确认两票分别完成实际出运，才按10,000kg。")
add_para(doc, "条件税款：若金额字段9,000可作为完税价格代理、实际为德国原产且中国按非受税原产申报并未缴税，则反倾销税=9,000×34.5%=3,105；VAT增量=403.65；合计3,508.65（币种同平台字段）。若确认为两票，合计7,017.30。该测算不是正式追税额。")
add_callout(doc, "为什么仍不能定性", "PTE LTD只说明新加坡法人名称，菲律宾数据源也不等于货物实际在菲律宾换柜。当前没有中国进口18位申报编号、申报原产地、提单号、集装箱、批号、CO或税单，无法确认中国端是否仍按德国原产申报并已缴税。", fill=PALE_RED, title_color=RED)

doc.add_heading("五、第二条具体线索：TENAC-C EX352经越南对华", level=1)
add_table(doc, ["日期", "路线/实体", "商品", "数量", "金额字段", "字段异常"], [[
    "2026-01-14",
    "YAMATO INDUSTRIES VIETNAM→SUZHOU INDUSTRIAL PARK ORIENTAL IMP.&EXP.CO.,LTD.",
    "POM, TENAC-C EX352；#&JP",
    "50kg（数量字段）",
    "10,732,331.50；币种未标",
    "平台原产Vietnam，货描明确#&JP",
]], [1200, 2800, 1700, 1200, 1300, 1160], size=8.9)
add_para(doc, "旭化成官方资料确认TENAC-C为共聚POM。EX352货描未出现玻纤、PTFE、弹性体等明显改性词，但仍须当前TDS/COA核实熔点、密度、共聚结构和添加剂。若最终确认为旭化成日本原产、落入范围且中国端改报越南并漏税，按24.5%和VAT13%计算，条件合计为2,971,245.98（与金额字段同币种）。")
add_para(doc, "同日另有TENAC MC050 1.25kg、#&JP。TENAC（非TENAC-C）倾向均聚物，初步排除，不与EX352合并计算反倾销税风险。")

doc.add_heading("六、其他第三国B腿、反证与数量纠偏", level=1)
add_table(doc, ["路线/日期", "商品与数量", "事实", "范围/风险结论"], [
    ["Hamakyu越南→上海恒久｜2026-05-26", "LW15EWX 25kg #&DE；CELCON LW90-S2 25kg #&US；越南色母4kg", "同一收发货组合4行，合计54kg；金额字段合计8,293,037", "具体混合原产试料批次；两受税来源牌号均有改性/范围外指征，暂不计本案税差"],
    ["Morimura越南→Morimura上海｜2026-02-04", "IUPITAL FV-30 NATURAL 50kg #&JP", "同批另有中国原产POM+TISMO/OA30B 50kg #&CN", "FV-30疑增强，需TDS；当前不能计入案涉量"],
    ["Providence越南→Providence｜2025-08-26", "DELRIN 500AL NC010 2,000kg #&NL", "货描引用越南前序申报107322453120/E11和107331742230", "最接近可追溯库存链，但Delrin为均聚POM，本案排除"],
    ["Sonion越南→Polygon｜2025-09-19", "Hostaform/Ultraform 520 Black 10kg #&DK", "丹麦标记；着色牌号", "范围待TDS；小量试料，尚无A腿"],
    ["Shaily印度→深圳德美创｜2026-02-26", "Hostaform MT12R01 25kg、MT12U03 150kg、MT24F01 50kg", "均为FOC材料验证包", "滑动/PTFE改性或熔点170℃边界外；不计案涉量"],
], [2100, 2450, 2500, 2310], size=8.9)

doc.add_heading("七、直接受税来源：高优先级核税，不是绕道", level=1)
doc.add_heading("7.1 欧盟POM PULVER与Leschaco节点", level=2)
add_para(doc, "POM关键词结果中，终裁后出现47条荷兰/欧盟POM PULVER记录，明确或推算毛重约3,511,503.4kg；主要为09034 MB800、27045 MB800、13014 MB800等，链路集中在LESCHACO NEDERLAND BV→LESCHACO (CHINA) LTD. NANJING BRANCH。多数平台HS为空或0，收发货人均呈物流角色，真实生产商、进口人和报关主体被遮蔽。")
add_para(doc, "“PULVER”“MB800”是否代表基础POM粉料、色母/母粒或其他改性配方，必须逐牌号TDS确认。若落入措施范围，应核实是否按欧盟34.5%缴税；这属于受税来源直接对华税款合规，不是第三国绕道。")

doc.add_heading("7.2 HOSTAFORM欧盟直达", level=2)
add_para(doc, "HOSTAFORM查询中，终裁后Leschaco Nederland→Leschaco China Nanjing共6个中国大陆商业组，平台重量合计242.816吨。2025-09-19长提单内可见C9021 ECO-B NATURAL至少1,420袋，按25kg/袋为35.5吨净重；该基础牌号公开熔点约166℃、ECO-B为质量平衡选项，初步在范围。C13031熔点约170℃、C9021 SW/TF、GV、XGC、AW等具有边界或改性指征，应排除或待TDS。")

doc.add_heading("7.3 美国CELCON", level=2)
add_para(doc, "HS390710结果中，终裁后美国Ticona/Celcon直达记录可拆出CELCON M25/M90天然基础料约716,609.595kg，另有CF2001改性料约10,650.276kg。基础标准料是应税候选，应查中国报关单和税款缴款书；无金额时只能写公式：条件差额=完税价格×74.9%×1.13。")
add_para(doc, "保证金期另有2025-04-15 Ticona→中国HOSTAFORM C9021 NATURAL，毛重19.548吨、约760袋（推定净重约19吨），应核实临时保证金及转税，但仍是美国直达，不是绕道。")

doc.add_heading("八、沙特来源措施后集中放量", level=1)
add_para(doc, "HS390710检索在2025-07-01至07-10集中出现HIZAM MOHAMMED AL-QAHTANI TRADING→BEIJING KANG JIE KONG INTL CARGO/INTERNATIONAL的POLYACETALS IN PRIMARY FORMS，10行、重量字段合计931,000kg；2026-01另有SABIC (CHINA) HOLDING→XIAMEN GOLDEN CHEMICALS 28,650kg。HS口径终裁后共11行、959,650kg。")
add_para(doc, "POM关键词另出现SABIC ASIA PACIFIC PTE LIMITED两端的POM 90S/140S 7行、201,320kg。两个查询口径的货描、双方和重量字段不同，但仍可能是同一供应体系的不同展示，报告不将两组简单相加。")
add_callout(doc, "研判", "沙特确有大型石化产业与SABIC销售体系，措施后来源增长可能是合法替代供应；也可能需要验证POM实际生产厂、聚合工序和原产规则。现有数据没有美国/欧盟/日本/台湾→沙特的同批A腿，因此只能列为中高优先级流量线索。", fill=PALE_GOLD)

doc.add_heading("九、公开来源中的商业链与结构风险", level=1)
doc.add_heading("9.1 日本宝理/大赛璐", level=2)
add_para(doc, "商务部终裁文书确认日本宝理存在经第三国（地区）关联贸易商向中国销售、经第三国关联贸易商再经第三国非关联贸易商向中国销售等渠道，并对该第三国关联贸易商实施视频核查。这证明第三国商业中间环节真实存在，但调查期为2023年，且文书没有认定货物在第三国换柜、换标或改变原产地。")
add_para(doc, "大赛璐2024年11月投资者问答公开表示，中国需求通常由马来西亚或台湾工厂供货，因反倾销税应根据税率切换出口国。宝理在日本、马来西亚、台湾和中国均有真实生产厂，因此“切换真实生产国”可以合法；只有日本制成品仅经第三国转售后改报第三国原产，才形成规避风险。")

doc.add_heading("9.2 德国Celanese", level=2)
add_para(doc, "商务部终裁文书确认德国Celanese全部对华出口经过关联贸易商，并存在德国生产商→欧盟关联贸易商→第三国关联贸易商→中国客户的商业链。Celanese在新加坡、香港、上海等地有贸易实体，而公开POM生产地主要包括德国法兰克福和美国得州Bishop。第三国关联贸易网络提供转售条件，但不等于物理转运或伪报。")

doc.add_heading("9.3 旭化成同牌号多产地", level=2)
add_para(doc, "旭化成公开资料显示部分TENAC-C牌号可在日本和中国生产。同品牌、同牌号不能反推原产地；必须看包装厂码、批次号和COA。旭化成越南实体公开业务为工程塑料部件开发/CAE支持而非POM聚合生产，因此若申报为越南原产，需重点核实生产工序和原产地证。")

doc.add_heading("十、是否存在第三国绕道的具体证据", level=1)
add_table(doc, ["证据层级", "已经掌握", "缺口", "结论"], [
    ["A2：官方商业链", "商务部核实宝理/Celanese经第三国关联贸易商对华销售", "未披露物理运输、批号、箱号及中国申报原产地", "证明商业中间环节，不证明绕道"],
    ["B+：交易字段异常", "Acumen C52021同批Germany/Philippines切换；TENAC-C #&JP/Vietnam", "中国报关单、CO、税单、批号和A腿", "强核单线索，尚非违法证明"],
    ["B：流量/结构异常", "沙特措施后集中放量；越南存在受税来源尾码B腿", "生产厂、上游受税来源A腿、实质加工", "需进一步核查"],
    ["反证/范围排除", "LW15EWX、Delrin、MT改性/边界牌号", "当前批次TDS/COA", "不得纳入本案逃税数量"],
    ["执法闭环", "未检出中国海关处罚、法院判决、商务部反规避认定或闭合双段提单", "正式执法/申报记录", "当前未证实第三国逃避反倾销税"],
], [1550, 3000, 2600, 2210], size=9.0)

doc.add_heading("十一、实体与口岸核查优先级", level=1)
add_table(doc, ["优先级", "实体/节点", "风险原因", "建议调取"], [
    ["1", "ACUMEN ENGINEERING PTE LTD / ACUMEN ENGINEERING SHANGHAI CO LTD", "同批C52021原产字段切换，产品初步在范围", "两日期修撤单、中国申报、CO、税单、200袋包装、批号"],
    ["2", "YAMATO INDUSTRIES VIETNAM / SUZHOU INDUSTRIAL PARK ORIENTAL I/E", "TENAC-C EX352 #&JP经越南对华", "越南出口单、中国报关单、Asahi批号、COA"],
    ["3", "LESCHACO NEDERLAND BV / LESCHACO CHINA NANJING", "POM PULVER/Hostaform大体量、真实货主被物流节点遮蔽", "House/Master B/L、实际进口人、生产商、南京报关信息"],
    ["4", "Ticona/Celcon美国节点及中国收货人", "标准共聚料直接受税来源", "中国进口人、适用74.9%税率、完税价格和缴款书"],
    ["5", "Hizam Al-Qahtani / Beijing Kang Jie Kong / SABIC链", "沙特来源措施后集中放量", "沙特工厂、POM产能、聚合工序、COA、口岸和真实货主"],
    ["6", "Hamakyu / 上海恒久百传动；Morimura；Providence", "第三国试料、前序申报和混合原产批次", "同票四行申报、越南前序进口单、批号与中国税单"],
], [900, 2650, 2850, 2960], size=9.0)
add_para(doc, "易迅JSON没有可靠的中国进口口岸、报关行和18位申报编号，企业名中的Shanghai/Nanjing或航运代码不能替代正式口岸字段。因此本报告不虚构“高风险口岸”；应以中国进口报关单补齐申报口岸、报关行和境内收货人后再排序。")

doc.add_heading("十二、下一步决定性数据清单", level=1)
for text in [
    "Acumen C52021两日期对应的中国进口报关单、18位申报编号、修撤单记录、申报原产国、10位税号、境内收货人和申报口岸。",
    "原产地证、商业发票、装箱单、COA/TDS、生产批号、包装唛头和工厂代码；特别核实C52021的200单位是否为200袋×25kg。",
    "反倾销税、保证金和进口增值税缴款书，包括完税价格、适用企业税率和税种明细。",
    "Master/House B/L、集装箱号、封志号、船名航次、装卸港、换单记录和通知方。",
    "TENAC-C EX352的越南出口申报、中国进口申报、Asahi日本工厂批号及当前批次COA。",
    "沙特931,000kg集群的真实货主、生产商、工厂地址、POM聚合产能、牌号、净重、箱号和中国进口口岸。",
    "POM PULVER各牌号TDS、真实生产商和House B/L，穿透Leschaco物流节点后的实际进口人。",
    "按30—90天时滞反查德国/日本/美国→新加坡、菲律宾、越南或沙特A腿，要求同牌号、批号、数量包装及箱号一致。",
]:
    add_number(doc, text)

doc.add_heading("十三、结论", level=1)
add_para(doc, "截至2026年8月13日，本次全量易迅排查和公开源深检没有形成“受税来源→第三国→中国、同批实物闭合、进入中国时改报第三国原产且未缴反倾销税”的完整证据链。Acumen C52021是当前最强的具体原产字段异常，TENAC-C EX352是最具体的日本标记经越南B腿；两者均值得优先调取中国进口端单证。")
add_para(doc, "风险工作的主次应调整：第一优先是受税来源直接对华范围和缴税核验，尤其欧盟POM PULVER、HOSTAFORM基础料和美国CELCON标准料；第二优先是Acumen/TENAC-C等具体字段异常；第三优先是沙特来源替代供应的生产实质。LW15EWX、Delrin及明确改性牌号必须从疑似逃税量中剔除，避免因品牌、平台“原产国”或第三国贸易商名称而误判。")

doc.add_heading("附录A：主要公开来源", level=1)
add_source(doc, "商务部公告2025年第25号（POM终裁）", "https://www.mofcom.gov.cn/zcfb/zgdwjjmywg/art/2025/art_09eb6be1f50f4cdaa6a36dfbb09bb529.html", "措施范围、企业税率、生效日和期限")
add_source(doc, "商务部公告2025年第5号（POM初裁）", "https://www.mofcom.gov.cn/zwgk/zcfb/art/2025/art_980edf9a52a4464bbe38582874eca3ee.html", "保证金期及初裁税率")
add_source(doc, "Celanese Hostaform产品手册", "https://www.celanese.com/-/media/Engineered%20Materials/Files/Product%20Sell%20Sheets/POM-062_HostaformProductManual_EU_EN_0614.pdf", "C52021、LW15EWX、C9021、C13031等牌号性能")
add_source(doc, "宝理全球生产网络", "https://www.polyplastics-global.com/cn/aboutus/network/production.html", "日本、马来西亚、台湾和中国生产基地")
add_source(doc, "大赛璐2024年11月投资者问答", "https://www.daicel.com/en/ir/pdf/q%26a_summary_25e-2q.pdf", "按税率调整供货生产国的公开说明")
add_source(doc, "旭化成TENAC/TENAC-C官方产品说明", "https://www.asahi-kasei-plastics.com/en/products/tenac/", "TENAC-C共聚体系")
add_source(doc, "旭化成TENAC-C多产地说明", "https://www.asahi-kasei.com/news/2015-2021/2019/e190522", "同牌号可能有日本/中国不同产地")
add_source(doc, "旭化成越南实体业务说明", "https://www.asahi-kasei.com/news/2015-2021/2016/e160425-2", "越南实体为技术/部件开发，不是已公开POM聚合厂")
add_source(doc, "Celanese美国SEC子公司清单", "https://www.sec.gov/Archives/edgar/data/1306830/000130683025000027/ex211-10k123124.htm", "新加坡、香港、上海等贸易实体网络")

doc.add_heading("附录B：交付文件", level=1)
for text in [
    "POM_易迅全量逐票判定台账.xlsx：4,179条跨查询唯一记录，含全部查询命中、重复次数、范围/路线/税款判定。",
    "POM_易迅跨查询逐票标准化.csv / .json：可供后续程序分析和系统导入。",
    "各查询原始全量JSON：HS390710、POM、HOSTAFORM、DELRIN、DURACON、TENAC、CELCON及补查词。",
    "POM_全量分析结果.json：统计口径、关键线索和程序化摘要。",
]:
    add_bullet(doc, text)

doc.save(OUT)
print(str(OUT))
