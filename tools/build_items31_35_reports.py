from __future__ import annotations

"""Build item 31-35 anti-dumping and rerouting risk reports.

The reports follow the documents skill's standard_business_brief preset and
memo_masthead opening pattern.  All platform quantities/amounts are described
as fields unless the shipment text itself states a unit.
"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_ROOT = ROOT / "outputs" / "反倾销税深度分析报告"
AS_OF = "2026年8月21日"

# standard_business_brief tokens
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "1C2733"
MUTED = "5C6773"
TABLE_FILL = "F2F4F7"
CALLOUT_FILL = "F4F6F9"
WHITE = "FFFFFF"
PALE_RED = "FDECEC"
PALE_GOLD = "FFF5D9"
PALE_GREEN = "EAF7F1"
RED = "9B1C1C"
GOLD = "7A5A00"
GREEN = "176B4B"
TOTAL_DXA = 9360
TABLE_INDENT_DXA = 120


def set_run_font(run, size=None, color=INK, bold=None, italic=None):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    return run


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, bottom=80, start=120, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def set_table_geometry(table, widths, caption):
    # Treat the first row of every table as its semantic header/entry row.
    # This also covers the one-cell masthead and callout tables.
    if table.rows:
        set_repeat_header(table.rows[0])
    assert sum(widths) == TOTAL_DXA, (widths, sum(widths))
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(TOTAL_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")
    cap = OxmlElement("w:tblCaption")
    cap.set(qn("w:val"), caption)
    tbl_pr.append(cap)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def create_numbering(doc, bullet=True):
    numbering = doc.part.numbering_part.element
    abs_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abs_id = max(abs_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abs_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), "bullet" if bullet else "decimal")
    lvl.append(fmt)
    text = OxmlElement("w:lvlText")
    text.set(qn("w:val"), "•" if bullet else "%1.")
    lvl.append(text)
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    lvl.append(suff)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    p_pr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:before"), "0")
    spacing.set(qn("w:after"), "160")
    spacing.set(qn("w:line"), "280")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    lvl.append(p_pr)
    abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abs_ref = OxmlElement("w:abstractNumId")
    abs_ref.set(qn("w:val"), str(abs_id))
    num.append(abs_ref)
    numbering.append(num)
    return num_id


def apply_num(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    nid = OxmlElement("w:numId")
    nid.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, nid])
    p_pr.append(num_pr)


def add_hyperlink(paragraph, text, url):
    rid = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Calibri")
    r_fonts.set(qn("w:hAnsi"), "Calibri")
    r_fonts.set(qn("w:eastAsia"), "微软雅黑")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend([r_fonts, color, underline])
    run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def configure_doc(doc, spec):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    title = doc.styles["Title"]
    title.font.name = "Calibri"
    title._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
    title.font.size = Pt(23)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(4)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.space_after = Pt(0)
    set_run_font(header.add_run(f"反倾销税深度核查｜ITEM {spec['num']}｜{spec['short']}"), 8.5, MUTED, True)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.space_after = Pt(0)
    set_run_font(footer.add_run("内部核查材料  |  "), 8, MUTED)
    add_page_field(footer)

    props = doc.core_properties
    props.title = spec["title"]
    props.subject = "反倾销税、易迅逐票数据、第三国转运与原产地核查"
    props.author = "反倾销税风险分析项目"
    props.keywords = f"{spec['short']};反倾销;易迅数据;第三国转运;原产地"


def add_para(doc, text="", bold_prefix=None, size=11, color=INK, after=None, italic=False):
    p = doc.add_paragraph()
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    if bold_prefix and text.startswith(bold_prefix):
        set_run_font(p.add_run(bold_prefix), size, color, True)
        set_run_font(p.add_run(text[len(bold_prefix):]), size, color, False, italic)
    else:
        set_run_font(p.add_run(text), size, color, False, italic)
    return p


def add_heading(doc, text, level=1):
    return doc.add_paragraph(text, style=f"Heading {level}")


def add_bullets(doc, items, num_id):
    for item in items:
        p = doc.add_paragraph()
        apply_num(p, num_id)
        set_run_font(p.add_run(item), 11, INK)


def add_table(doc, headers, rows, widths, caption, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_repeat_header(table.rows[0])
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade(cell, TABLE_FILL)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(str(header)), 9, DARK_BLUE, True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if i == 0 and len(headers) <= 3:
                set_run_font(p.add_run(str(value)), font_size, INK, True)
            else:
                set_run_font(p.add_run(str(value)), font_size, INK)
    set_table_geometry(table, widths, caption)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return table


def add_callout(doc, title, text, fill=CALLOUT_FILL, title_color=DARK_BLUE):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    shade(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    set_run_font(p.add_run(title + "\n"), 11, title_color, True)
    set_run_font(p.add_run(text), 10.5, INK)
    set_table_geometry(table, [TOTAL_DXA], f"{title}提示框")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_source_list(doc, sources, bullet_id):
    for label, url, note in sources:
        p = doc.add_paragraph()
        apply_num(p, bullet_id)
        set_run_font(p.add_run(label + "："), 9.5, INK, True)
        add_hyperlink(p, "官方/公开原文", url)
        if note:
            set_run_font(p.add_run("。" + note), 9.5, MUTED)


def build_report(spec):
    folder = OUT_ROOT / f"{spec['num']}_{spec['short']}"
    folder.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_doc(doc, spec)
    bullets = create_numbering(doc, True)
    decimals = create_numbering(doc, False)

    # memo_masthead opening
    kicker = add_para(doc, f"ITEM {spec['num']}  /  TRADE REMEDY RISK REVIEW", size=9, color=DARK_BLUE, after=4)
    kicker.runs[0].bold = True
    doc.add_paragraph(spec["title"], style="Title")
    add_para(doc, spec["subtitle"], size=13, color=MUTED, after=12)
    for label, value in (
        ("报告状态", spec["status"]),
        ("风险评级", spec["rating"]),
        ("核查基准日", AS_OF),
        ("分析口径", spec["method"]),
    ):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        set_run_font(p.add_run(label + "："), 10.5, INK, True)
        set_run_font(p.add_run(value), 10.5, INK)
    add_para(doc, "", after=4)
    add_callout(doc, "结论先行", spec["summary"], spec.get("summary_fill", PALE_GOLD), spec.get("summary_color", GOLD))
    add_table(
        doc,
        ["判断维度", "证据等级", "阶段结论"],
        spec["decision_rows"],
        [1900, 1400, 6060],
        f"ITEM {spec['num']}阶段判断",
        9.2,
    )
    add_para(doc, "说明：A=官方查处或同批单证闭环；B+=多项独立证据高度一致但缺决定性单证；B/C=核查线索；反证=存在合法来源或加工解释。", size=9, color=MUTED, after=0)

    doc.add_page_break()
    add_heading(doc, "1. 措施范围、期限与税率", 1)
    add_table(doc, ["项目", "核定内容"], spec["policy_rows"], [2100, 7260], f"ITEM {spec['num']}政策口径", 9.4)
    add_heading(doc, "1.1 条件性税差系数", 2)
    add_table(doc, ["情景", "反倾销税率", "AD引致增值税差额", "综合增量系数"], spec["tax_rows"], [3300, 1600, 1960, 2500], f"ITEM {spec['num']}税差系数", 8.8)
    add_para(doc, "注：综合增量系数按进口增值税13%作筛查，公式为反倾销税率×1.13。它不含正常关税和正常进口增值税，也不等于整票总税负；正式追税必须使用中国海关完税价格、生产商税档和税款缴款书。", size=9, color=MUTED)

    add_heading(doc, "2. 易迅查询覆盖与逐票审计", 1)
    add_table(doc, ["口径", "结果"], spec["query_rows"], [2500, 6860], f"ITEM {spec['num']}易迅查询覆盖", 9.4)
    add_para(doc, spec["query_note"], size=9.3, color=MUTED)
    add_heading(doc, "2.1 关键记录与实体", 2)
    add_table(doc, spec["lead_headers"], spec["lead_rows"], spec["lead_widths"], f"ITEM {spec['num']}关键记录", 8.5)
    add_heading(doc, "2.2 逐票文件与复核边界", 2)
    add_bullets(doc, spec["data_files"], bullets)

    if spec.get("page_break_before_chain"):
        doc.add_page_break()
    add_heading(doc, "3. 第三国路径、支持事实与反证", 1)
    add_table(doc, ["线索/路径", "等级", "支持事实", "反证或缺口"], spec["chain_rows"], [2100, 900, 3100, 3260], f"ITEM {spec['num']}第三国风险矩阵", 8.3)
    add_heading(doc, "3.1 合法第三国产能与供应链反证", 2)
    add_bullets(doc, spec["counterevidence"], bullets)
    add_callout(doc, "证据边界", spec["evidence_boundary"], PALE_RED, RED)

    if spec.get("page_break_before_entities"):
        doc.add_page_break()
    add_heading(doc, "4. 优先核查实体与调证清单", 1)
    add_table(doc, ["优先级", "实体/记录", "核查资料与目的"], spec["entity_rows"], [900, 3100, 5360], f"ITEM {spec['num']}调证清单", 8.5)
    add_heading(doc, "4.1 建议核查顺序", 2)
    for step in spec["steps"]:
        p = doc.add_paragraph()
        apply_num(p, decimals)
        set_run_font(p.add_run(step), 11, INK)

    add_heading(doc, "5. 结论", 1)
    add_para(doc, spec["conclusion"])
    add_bullets(doc, spec["closing_points"], bullets)

    add_heading(doc, "6. 公开来源", 1)
    add_source_list(doc, spec["sources"], bullets)
    add_para(doc, "公开检索的负面结果仅说明未在可索引的官方/公开来源中找到相应案件，不代表不存在未公开调查、行政处理或刑事案件。", size=9, color=MUTED)

    out = folder / spec["filename"]
    doc.save(out)
    return out


ITEMS = [
    {
        "num": 31,
        "short": "腈纶",
        "title": "腈纶反倾销税与第三国转运风险深度核查报告",
        "subtitle": "日本、韩国、土耳其原产｜易迅全页逐票审计与生产实体穿透",
        "status": "易迅两年关键词结果全部读取；公开来源深检完成；待中国底单闭环",
        "rating": "中高：土耳其直达核税票明确，第三国绕道未闭合",
        "method": "ACRYLIC FIBER × China × 两年，200条/页，2页全读；逐票范围初筛、可见字段去重和公开产能反证",
        "summary": "两年易迅结果278条、可见字段精确去重274条。仅发现1条受税来源直达：2026-01-26，Ak-Pa向杭州永芳纺织供应AK200腈纶短纤，重量字段40,974，平台原产Turkey。公开披露可将Ak-Pa与Aksa出口体系关联，因此该票应优先核Aksa 8.2%税档、原产地和缴款书；它不是第三国绕道。秘鲁来源213条、重量字段合计4,362,888.31，供应商高度集中于Sudamericana de Fibras，但秘鲁政府和企业资料证明当地具有真实腈纶聚合/纺丝能力，构成强反证。没有形成日本/韩国/土耳其→第三国→中国的A/B腿闭环。",
        "decision_rows": [
            ["土耳其直达核税", "B", "1条AK200短纤，受税来源与产品范围高度吻合，需核实际生产商和税款"],
            ["秘鲁来源规模", "反证/B", "213条规模显著，但当地36,000吨/年真实生产能力可合法解释"],
            ["第三国绕道", "未形成", "现有数据无同牌号、同批次、同柜号或受税A腿闭合"],
        ],
        "policy_rows": [
            ["措施与期限", "商务部2022年第21号：自2022-07-14起继续征收5年，预计至2027-07-13；受税来源为日本、韩国、土耳其。"],
            ["税号", "55013000、55033000、55063000。"],
            ["技术范围", "聚丙烯腈或丙烯腈共聚物中丙烯腈重复单元不低于85%的长丝束、短纤维和毛条。"],
            ["主要排除", "变性腈纶；PAN基碳纤维原丝（丙烯腈≥95%、连续无卷曲丝束、强度/伸长及用途同时满足公告条件）。"],
            ["日本税率", "日本爱克斯兰16.1%；三菱化学15.8%；东丽16.0%；其他日本公司16.1%。"],
            ["韩国税率", "泰光产业8.6%；其他韩国公司21.7%。"],
            ["土耳其税率", "Aksa 8.2%；其他土耳其公司16.1%。"],
        ],
        "tax_rows": [
            ["Aksa土耳其", "8.2%", "1.066%×完税价", "9.266%×完税价"],
            ["其他土耳其/日本爱克斯兰", "16.1%", "2.093%×完税价", "18.193%×完税价"],
            ["韩国泰光", "8.6%", "1.118%×完税价", "9.718%×完税价"],
            ["其他韩国公司", "21.7%", "2.821%×完税价", "24.521%×完税价"],
        ],
        "query_rows": [
            ["查询", "ACRYLIC FIBER + 目的国China + 2024-08-14至2026-08-14；结果278条；200条/页，共2页（200+78），逐页读取。"],
            ["去重与日期", "原始278条；13个可见字段精确去重274条；重复多余4条；记录日期2024-08-14至2026-07-15。"],
            ["原产分布", "Peru 213、United States 23、Indonesia 16、Vietnam 11、China 6、India 3、Turkey 1、Philippines 1。"],
            ["范围初筛", "218条为措施范围候选；56条需按税号、成分、形态或是否成衣/制品进一步复核。"],
            ["受税来源", "Turkey 1条；Japan 0；Korea 0。"],
        ],
        "query_note": "平台的“原产国地区”并非中国海关法定原产地字段；重量/数量/金额字段在不同数据源口径不一。除货描明确单位外，本报告不擅自把字段解释为中国申报净重、美元或人民币。",
        "lead_headers": ["日期/路线", "产品与主体", "可见规模", "判断"],
        "lead_widths": [1600, 3900, 1500, 2360],
        "lead_rows": [
            ["2026-01-26 土耳其→中国", "AK200 R01 BRIGHT 1.1DTEX 38MM；Ak-Pa→杭州永芳纺织", "重量字段40,974；金额空", "范围候选；受税直达；优先核Aksa生产商与8.2%税档"],
            ["两年 秘鲁→中国", "Sudamericana de Fibras为供应商的213条", "重量字段合计4,362,888.31", "非受税来源大宗；真实秘鲁产能构成强反证"],
            ["美国等→中国", "23条美国平台原产及其他国家记录", "多含成衣/纺织品或税号空", "逐票范围待核；美国不是本措施受税来源"],
        ],
        "data_files": [
            "原始页面摘录：腈纶_易迅_ACRYLIC_FIBER_China_两年.json.gz.b64。",
            "逐票标准化：31_keyword_易迅逐票标准化.csv，保留13个可见字段并给出每条范围初筛。",
            "完整性边界：本轮完成中国B腿关键词查询；未取得受税来源→秘鲁/其他第三国的同产品A腿和提单号。",
        ],
        "chain_rows": [
            ["Ak-Pa/Aksa土耳其直达", "B", "受税来源、牌号、形态、实体关系一致", "是否Aksa生产、是否正常缴AD、中国进口人底单"],
            ["秘鲁Sudamericana→中国", "反证", "213条、单一真实制造商、当地政府资料确认腈纶工厂", "仍需COA/厂批号核每票是否范围内，但无受税A腿"],
            ["受税国→第三国→中国", "未形成", "未发现同牌号、同批次或同柜号双腿", "缺A腿、生产批号、提单、原产地证"],
        ],
        "counterevidence": [
            "秘鲁Sudamericana de Fibras公开资料及秘鲁政府文件确认Callao腈纶工厂和约36,000吨/年产能，秘鲁来源不能仅因非传统来源而推定绕道。",
            "Aksa官方披露Yalova约355,000吨/年腈纶能力；Ak-Pa是Akkök体系的出口/营销主体，贸易商开票不改变实际土耳其生产的可能性。",
            "韩国Taekwang、土耳其Aksa、日本Toray等均有本国真实产能，品牌、集团国籍和发货国必须与具体厂批号区分。",
        ],
        "evidence_boundary": "目前没有证据证明秘鲁货物实际由日本、韩国或土耳其生产，也没有官方查处、法院判决或同柜双段提单。对任何企业只能表述为核查对象，不能写成伪报原产地、走私或逃税主体。",
        "entity_rows": [
            ["1", "Ak-Pa Tekstil / Aksa / 杭州永芳纺织", "调2026-01-26报关单、生产商栏、COA、原产地证、反倾销税缴款书和完税价格，确认8.2%或16.1%税档。"],
            ["2", "Sudamericana de Fibras", "按批号核聚合/纺丝工厂、原料账、COA和原产证；如出现受税牌号，回查90日A腿。"],
            ["3", "中国端隐藏买方232条", "向承运人/数据源补提单和中国收货人，排除成衣/制品后再定位进口人和口岸。"],
        ],
        "steps": [
            "先调土耳其直达票的中国报关单、税款书与生产商证明；金额为空，不能测实际少缴税额。",
            "再以AK200、Aksa、Ak-Pa和批号为键，回查土耳其出口及中国进口是否镜像/重复。",
            "对秘鲁大宗只在出现受税牌号、厂批号冲突或受税A腿后升级；否则保留真实生产反证。",
        ],
        "conclusion": "第31项已完成易迅中国B腿全页审计。最具体风险不是第三国绕道，而是一条土耳其受税来源直接对华的AK200短纤票；应核是否按Aksa税率缴纳反倾销税。秘鲁来源数量大，但真实制造能力证据充分，不能由贸易规模反推规避。当前无A等级或B+等级的第三国绕道闭环。",
        "closing_points": ["立即核税：土耳其直达票。", "条件升级：只有受税A腿+秘鲁/第三国B腿在牌号、批号、重量或柜号闭合时才提升风险。", "实际税额：取得中国完税价格和缴款书后计算。"],
        "sources": [
            ["商务部公告2022年第21号", "https://www.mofcom.gov.cn/zcfb/zgdwjjmywg/art/2022/art_c2935f35d87647e586b5a17a6b4eaf0a.html", "范围、税号、税率与期限。"],
            ["Aksa官方About Us", "https://aksa.com/en/corporate/about-us", "土耳其腈纶产能与全球销售。"],
            ["Akkök官方Ak-Pa介绍", "https://akkok.com.tr/en/sektorler/akpa/", "Ak-Pa为Aksa产品的出口/营销主体。"],
            ["Sudamericana de Fibras公司历史", "https://sdef.com/historia", "秘鲁Callao真实腈纶生产。"],
            ["秘鲁政府产业文件", "https://spij.minjus.gob.pe/Normas/textos/220312T.pdf", "Drytex腈纶约36,000吨/年。"],
            ["中国进出口货物原产地条例", "https://xzfg.moj.gov.cn/front/law/detail?LawID=1523", "反倾销非优惠原产地法律边界。"],
        ],
        "filename": "31_腈纶_反倾销税与第三国转运风险深度核查报告.docx",
    },
    {
        "num": 32,
        "short": "未漂白纸袋纸",
        "title": "未漂白纸袋纸反倾销税与第三国转运风险深度核查报告",
        "subtitle": "美国、欧盟、日本原产｜物性门槛、加拿大贸易商与直达税源核查",
        "status": "易迅两年精准关键词结果全部读取；商业记录归并完成；待七项物性和原厂闭环",
        "rating": "中高：受税来源直达明确，加拿大贸易商链需穿透",
        "method": "UNBLEACHED SACK KRAFT PAPER × China × 两年；51条全读；可见字段去重、CERS/日期商业归并、物性与生产商核查",
        "summary": "易迅结果51条、可见字段精确去重32条。按日期、供应商、重量及相同CERS归并后，受税来源明确的五个候选物理组约346,097kg：美国260,963kg，欧盟85,134kg，日本0；均需补七项物性和税款书。加拿大Fortis Trading链最值得核原产：归并为7个商业组、约642,742kg，供应商是贸易商且货描未列纸厂和七项指标。但公开贸易资料显示Fortis从巴西、德国、俄罗斯等多源采购，加拿大发货本身不能证明产地，也没有受税国A腿闭合。Bay Hill四组83,525kg写明“rejected silicone release paper”，存在范围外反证。",
        "decision_rows": [
            ["受税来源直达", "B", "5个候选商业组约346,097kg；核七项物性、生产商税档及缴款"],
            ["加拿大Fortis链", "B", "7个商业组约642,742kg；贸易商/无纸厂字段，需穿透原产"],
            ["第三国绕道", "未形成", "无受税国→Fortis/加拿大同批A腿、卷号、柜号或CO闭合"],
        ],
        "policy_rows": [
            ["措施与期限", "商务部2022年第10号：自2022-04-10起继续征收5年，预计至2027-04-09；受税来源美国、欧盟、日本。英国措施已于2021-04-10终止。"],
            ["税号", "48042100、48043100。"],
            ["七项物性", "定量≤115g/m²；抗张指数纵横向之和≥69N·m/g；纵向撕裂指数≥10mN·m²/g；横向/纵向抗张能量吸收指数分别≥1.0/0.8J/g；透气度≥3.4μm/(Pa·s)；纵向伸长率≥2%。"],
            ["美国税率", "所有美国公司14.9%。"],
            ["欧盟税率", "Billerud Sweden 23.5%；Mondi Stambolijski 29.0%；Billerud Finland及Mondi Frantschach/Steti/Dynas 26.2%；其他欧盟29.0%。"],
            ["日本税率", "所有日本公司20.5%。"],
        ],
        "tax_rows": [
            ["美国", "14.9%", "1.937%×完税价", "16.837%×完税价"],
            ["Billerud Sweden", "23.5%", "3.055%×完税价", "26.555%×完税价"],
            ["欧盟26.2%档", "26.2%", "3.406%×完税价", "29.606%×完税价"],
            ["其他欧盟/部分Mondi", "29.0%", "3.770%×完税价", "32.770%×完税价"],
            ["日本", "20.5%", "2.665%×完税价", "23.165%×完税价"],
        ],
        "query_rows": [
            ["查询", "UNBLEACHED SACK KRAFT PAPER + 目的国China + 2024-08-14至2026-08-14；结果51条，单页全读。"],
            ["去重", "原始51条；可见字段精确去重32条；重复多余19条；日期2024-08-28至2026-05-27。"],
            ["可见原产", "Canada 21、United States 7、Sweden 2、Poland 1、Germany 1。"],
            ["商业归并", "同日/同供应商/同重量并结合相同CERS跨日归并；Fortis加拿大7组约642,742kg；清晰受税来源5组约346,097kg。"],
            ["物性完整性", "所有32条都缺至少一项公告物性；货描命中不等于最终落入措施范围。"],
        ],
        "query_note": "同一可见字段相同的行可能是平台重复，也可能是多票同值；商业归并仅用于筛查，正式票数须以提单号、卷号、箱号和中国报关单确认。",
        "lead_headers": ["路线/实体", "归并规模", "具体事实", "核查判断"],
        "lead_widths": [1900, 1600, 3300, 2560],
        "lead_rows": [
            ["美国→中国", "2组/260,963kg", "Six Continents 27,900kg；FMS/10×40柜233,063kg", "受税直达；七项物性和14.9%缴税待核"],
            ["欧盟→中国", "3组/85,134kg", "Mondi德国29,166；Billerud瑞典28,368；Stora Enso波兰27,600", "受税直达；需对应企业税档与缴款书"],
            ["加拿大Fortis→中国", "7组/642,742kg", "上海Yearich、青岛Funderief；部分记录有CERS编号", "贸易商链B级；调纸厂、卷号、原产证及A腿"],
            ["美国Bay Hill→中国", "4组/83,525kg", "货描同时写rejected silicone release paper", "范围外反证强；需涂层/硅油、定量和用途"],
        ],
        "data_files": [
            "原始页面摘录：未漂白纸袋纸_易迅关键词_China_两年.json.gz.b64。",
            "逐票标准化：32_keyword_易迅逐票标准化.csv，51条均保留并逐条标为七项物性待核。",
            "商业归并口径不替代海关票数；报告并列原始51、精确去重32和归并后的物理/商业组。",
        ],
        "chain_rows": [
            ["Fortis Canada→中国", "B", "贸易商、规模642.742t、买方集中、部分CERS可调", "缺原厂/卷号/七项物性；Fortis多源采购可合法解释"],
            ["美国/欧盟直达", "B", "来源国、供应商和纸类货描吻合", "是否范围内及是否已缴税未知；不是绕道"],
            ["受税国→加拿大→中国", "未形成", "公开历史显示Fortis曾从多国采购", "无同批A腿、时间/重量/柜号闭合"],
            ["Bay Hill release paper", "反证", "货描明确硅离型纸/退货纸", "可能不满足未漂白纸袋纸物性与用途"],
        ],
        "counterevidence": [
            "加拿大及巴西等地存在真实牛皮纸/纸袋纸产能；Fortis作为贸易商可合法采购非受税来源。",
            "Mondi在Steti、Frantschach、Dynas等地有真实纸袋纸工厂，欧盟来源直达更像正常应税进口而非第三国链。",
            "Billerud Karlsborg等工厂具有真实牛皮纸产能；品牌或贸易商国籍不能替代卷号对应的生产厂。",
        ],
        "evidence_boundary": "Fortis是高优先原产地核查对象，但“加拿大贸易商+中国买方+大批量”仍不足以证明其货物由美国、欧盟或日本生产。只有纸厂、卷号、A/B提单和中国原产申报闭合，才能升级为第三国规避证据。",
        "entity_rows": [
            ["1", "Fortis Trading / Shanghai Yearich / Qingdao Funderief", "调CERS、提单、纸厂发票、卷号、原产证、七项物性和中国报关单；按30—120日回查上游采购。"],
            ["2", "FMS Seaways / 中国收货人隐藏", "以合同8931854、船名YM MOBILITY 084W和Long Beach装船信息定位进口人、报关行与税款。"],
            ["3", "Mondi、Billerud、Stora Enso直达票", "核生产厂、列名税档、完税价格和AD缴款；不得用出口商名自动套最低税率。"],
            ["4", "Bay Hill rejected release paper", "调TDS/涂层成分、定量、硅油处理和用途，先做范围排除。"],
        ],
        "steps": [
            "先对五个受税来源候选组核物性、生产商和税款书，判断是否已正常缴AD。",
            "再按Fortis每个CERS/日期/重量调纸厂和卷号，回查受税来源A腿；只凭Vancouver装港不定产地。",
            "对离型纸、涂层纸和joblot逐票做产品范围复核，避免把范围外废纸/特种纸计入风险量。",
        ],
        "conclusion": "第32项已完成精准关键词全页审计。已有受税来源直达核税对象约346,097kg；加拿大Fortis链规模更大、实体和CERS可操作性强，是最值得原产地穿透的第三国线索。但没有双段提单或原厂闭环，当前证据等级为B，不得写成已逃反倾销税。",
        "closing_points": ["核税优先：美国、德国、瑞典、波兰直达候选。", "原产优先：Fortis加拿大7个商业组。", "范围优先：Bay Hill硅离型纸及所有缺七项物性的记录。"],
        "sources": [
            ["商务部公告2022年第10号", "https://trb.mofcom.gov.cn/myjjdc/art/2022/art_2b2552a348664bae9dc5c92554c9a8aa.html", "范围、物性、税率与期限。"],
            ["Mondi 2024年度报告", "https://www.mondigroup.com/globalassets/mondigroup.com/investors/results-reports-and-presentations/2024/integrated-report-and-financial-statements/mondi-group-integrated-report-and-financial-statements-2024.pdf", "欧洲纸袋纸工厂与产能反证。"],
            ["Mondi Steti工厂", "https://www.mondigroup.com/de/standorte/tschechien/mondi-steti/", "Steti纸袋纸生产。"],
            ["Billerud Karlsborg产能公告", "https://www.billerud.com/press--news/press-releases/2017/billerudkorsnas-unit-in-karlsborg-now-back-to-full-production", "真实生产能力。"],
            ["ImportGenius Fortis历史公开页", "https://www.importgenius.cn/importers/fortis-trading-ltd", "二级聚合来源，仅用于说明贸易商多源采购；不证明当前批次原产。"],
            ["中国进出口货物原产地条例", "https://xzfg.moj.gov.cn/front/law/detail?LawID=1523", "非优惠原产地核定。"],
        ],
        "filename": "32_未漂白纸袋纸_反倾销税与第三国转运风险深度核查报告.docx",
    },
    {
        "num": 33,
        "short": "光纤预制棒",
        "title": "光纤预制棒反倾销税与第三国转运风险深度核查报告",
        "subtitle": "日本、美国原产｜HS700220全页审计、直径门槛与印度返修链",
        "status": "关键词0条、HS宽池34条全部读取；预制棒候选4条；待直径、生产批次和返修单证",
        "rating": "中低：未见受税来源或闭合绕道，印度返修链可核",
        "method": "OPTICAL FIBER PREFORM与HS700220分别查询；China × 两年；关键词0、HS34条全读并逐票分类",
        "summary": "精准关键词结果0；HS700220结果34条、可见字段精确去重33条。只有4条明确写PREFORM，均为ZTT India→Zhongtian Technology Advanced，货描为“FOR REPAIR AND RETURN—PREFORM OF SILICA G652D”，数量字段合计2,159.28、金额字段135,461.76；平台原产India。另有Corning India→Corning Hainan的2条低值/样品式记录，但货描不足。未见Japan或United States平台原产，也未形成受税来源→印度/越南→中国双腿。ZTT India官方宣称具有从preform到fiber的产业链能力，Corning India也有真实光纤工厂，均构成第三国合法生产/返修反证。",
        "decision_rows": [
            ["印度ZTT返修预制棒", "B-/C+", "4条同一集团、G652D、返修退运；先核货物原始来源与直径"],
            ["Corning India内部记录", "C", "2条低数量/低金额，描述为blanks roots/LUT bond；范围不明"],
            ["日本/美国绕道", "未形成", "B池没有受税原产；无同批A腿、柜号、厂批号"],
        ],
        "policy_rows": [
            ["措施与期限", "商务部2024年第27号：自2024-07-11起继续征收5年，预计至2029-07-10；受税来源日本、美国。"],
            ["税号", "70022010。"],
            ["产品范围", "用于制造光导纤维、具有特定折射率剖面的石英玻璃棒；英文Optical Fiber Preform或Fiber Preform。"],
            ["主要排除", "直径小于60毫米的进口产品不在措施范围。普通石英棒、玻璃棒、设备备件和光纤/光缆制品不能仅凭HS700220并入。"],
            ["日本税率", "Shin-Etsu 17.0%；Fujikura 14.4%；Sumitomo 31.2%；Furukawa 31.2%；其他日本公司31.2%。"],
            ["美国税率", "Corning 41.7%；OFS Fitel 17.4%；其他美国公司41.7%。"],
        ],
        "tax_rows": [
            ["Fujikura日本", "14.4%", "1.872%×完税价", "16.272%×完税价"],
            ["Shin-Etsu日本", "17.0%", "2.210%×完税价", "19.210%×完税价"],
            ["其他/部分日本", "31.2%", "4.056%×完税价", "35.256%×完税价"],
            ["OFS Fitel美国", "17.4%", "2.262%×完税价", "19.662%×完税价"],
            ["Corning/其他美国", "41.7%", "5.421%×完税价", "47.121%×完税价"],
        ],
        "query_rows": [
            ["精准关键词", "OPTICAL FIBER PREFORM + China + 两年：0条。"],
            ["HS宽池", "HS700220 + China + 两年：34条；单页全读；可见字段精确去重33条；日期2024-08-23至2026-05-27。"],
            ["原产分布", "India 24、Vietnam 6、Indonesia 1、Kazakhstan 1、Germany 1。"],
            ["范围分层", "明确预制棒候选4；石英/玻璃棒用途与直径待核16；同HS其他/待核13。"],
            ["受税来源", "Japan 0；United States 0。"],
        ],
        "query_note": "HS700220在外国来源数据中包含普通熔融石英棒、设备部件和实验/返修品；中国措施税号70022010还要求用途和折射率剖面，且直径必须不小于60毫米。",
        "lead_headers": ["日期/路线", "实体与货描", "数量/金额字段", "判断"],
        "lead_widths": [1600, 4100, 1700, 1960],
        "lead_rows": [
            ["2024-09-26至2025-11-12 印度→中国", "ZTT India→Zhongtian Technology Advanced；G652D silica preform；repair and return", "4条；2,159.28 / 135,461.76", "范围候选；直径、原始出口/返修性质待核"],
            ["2025-01-18 印度→中国", "Corning Technologies India→Corning Hainan；LUT bond / blanks roots", "2条；数量4；金额150", "无法确认预制棒；可能样品/工艺坯/集团内寄送"],
            ["其他第三国→中国", "印度普通石英棒、越南玻璃棒、印尼OVD sparepart等", "27条唯一", "多数非预制棒或范围待核；不能并入风险量"],
        ],
        "data_files": [
            "原始页面摘录：光纤预制棒_易迅_HS700220_China_两年.json.gz.b64。",
            "逐票标准化：33_hs_易迅逐票标准化.csv，34条全部保留并逐条分类。",
            "关键词0条与HS34条的差异说明平台对“preform”拼写/语种覆盖有限；不能以关键词0推断无进口。",
        ],
        "chain_rows": [
            ["ZTT India返修→中国", "B-/C+", "同集团、G652D、连续4条、返修退运措辞", "缺原始中国→印度出境单、直径、批号、柜号；印度可真实制造"],
            ["Corning India→Hainan", "C", "集团实体、HS700220、低值样品式记录", "货描无preform/直径；Corning印度有光纤制造能力"],
            ["日本/美国→第三国→中国", "未形成", "无受税原产B腿和同批A腿", "缺日本/美国出口侧、生产厂代码和中国原产申报"],
        ],
        "counterevidence": [
            "ZTT India官方产品页称其产业链覆盖Preform-Fiber-Cable并在印度制造光纤/光缆；印度来源可由当地真实工艺解释。",
            "Corning官方确认Pune光纤制造设施；但公开页面未必证明该厂本身生产预制棒，因此仍须看具体plant code和COA。",
            "“repair and return”通常对应返修/退运或集团工艺往返；若货物原本中国原产、返修未改变原产地，则更不是日本/美国规避。",
        ],
        "evidence_boundary": "ZTT的公开“从预制棒到光纤”表述是产业能力证据，不等于每条对华记录均由印度聚合/沉积形成；反过来也不能把集团技术来源当成日本或美国原产。原产判断必须落到每支预制棒的沉积厂、直径、批号和返修单证。",
        "entity_rows": [
            ["1", "ZTT India / Zhongtian Technology Advanced", "调四条返修单、原始出境报关、设备/工艺维修说明、直径、重量单位、厂批号、原产地证和中国复进口税单。"],
            ["2", "Corning Technologies India / Corning Hainan", "调LUT bond对应发票、样品说明、blanks roots技术定义、直径、plant code与COA。"],
            ["3", "普通石英棒买卖方", "先按折射率剖面、用途、直径和中国10位税号排除，再决定是否调原产。"],
        ],
        "steps": [
            "先核四条ZTT返修票的原始出境和复进口链，确认是否为中国货返修。",
            "对所有700220记录取得产品图纸/TDS、折射率剖面和直径；小于60毫米直接排除。",
            "只有出现日本/美国生产厂代码或A腿、且第三国未发生实质制造时，才升级绕道风险。",
        ],
        "conclusion": "第33项没有发现受税来源直达，也没有日本/美国经第三国进入中国的闭合证据。四条ZTT India返修预制棒记录是最具体的核查线索，但印度真实制造能力和“repair and return”措辞提供了强合法解释。当前评级中低，重点是产品直径和返修原始出境闭环。",
        "closing_points": ["产品范围门槛优先于路线判断。", "返修/复进口单证可快速验证ZTT链。", "金额字段不明币种，不能测税额。"],
        "sources": [
            ["商务部公告2024年第27号", "https://www.mofcom.gov.cn/zcfb/zgdwjjmywg/art/2024/art_f4d088624b074201ae0bd516170d751a.html", "范围、直径排除、税率与期限。"],
            ["ZTT India光纤产品页", "https://www.zttindia.com/optical-fiber/", "公开宣称Preform-Fiber-Cable产业链。"],
            ["ZTT India公司介绍", "https://www.zttindia.com/about-us/", "印度工厂与产品网络。"],
            ["Corning India公司介绍", "https://www.corning.com/in/en/about-us/company-profile/Corning-In-India.html", "Pune光纤制造设施。"],
            ["Corning全球光纤制造", "https://www.corning.com/optical-communications/in/en/home/products/fiber/manufacturing-excellence.html", "全球制造网络反证。"],
            ["中国进出口货物原产地条例", "https://xzfg.moj.gov.cn/front/law/detail?LawID=1523", "返修/多国生产原产地边界。"],
        ],
        "filename": "33_光纤预制棒_反倾销税与第三国转运风险深度核查报告.docx",
    },
    {
        "num": 34,
        "short": "太阳能级多晶硅",
        "title": "太阳能级多晶硅反倾销税与第三国转运风险深度核查报告",
        "subtitle": "美国、韩国原产｜HSC直达、马来西亚ST600与东南亚返料链",
        "status": "关键词与HS查询全部读取并跨查询去重；23条唯一；待牌号、等级、完税价格和A/B腿",
        "rating": "高优先核税：美国直达明确；第三国链有具体牌号但反证较强",
        "method": "POLYSILICON与HS280461分别查询China × 两年；18+18条全读；跨查询23条唯一，逐票范围与路线复核",
        "summary": "两组查询各18条，跨查询可见字段合并去重23条。可归并为3个美国原产直接对华物理批次，净重代理各约17,100kg、合计约51,300kg；其中2024-10-30货描含HSC ST740/ST260，HSC高度指向Hemlock。该线索应优先核太阳能级/电子级、Hemlock 53.3%税档及缴款，但不是绕道。第三国最具体的是：2024-12-12 Runergy Vietnam向上海发100,800数量字段多晶硅，货描#&MY并注明出口退回；2026-06-26 PT Nusa Solar Indonesia向DMEGC发ST600 2,880kg。公开贸易资料将ST600稳定关联OCI/OCIM马来西亚太阳能级料，OCI官方也确认马来西亚为太阳能级生产基地，故更可能是马来西亚料在东南亚再出口，而非美韩规避。没有闭合美国/韩国→第三国→中国链。",
        "decision_rows": [
            ["美国HSC/Sea Trade直达", "B+核税", "3个物理批次、约51.3t净重代理；需核太阳能级与53.3%税档"],
            ["ST600东南亚B腿", "B/反证", "印尼→中国2.88t；品牌/牌号指向OCI马来西亚合法产能"],
            ["美韩第三国绕道", "未形成", "无同批A腿、生产批号、柜号或中国原产申报闭合"],
        ],
        "policy_rows": [
            ["措施与期限", "商务部2026年第3号：对美国、韩国原产太阳能级多晶硅继续征收5年，预计至2031-01-13。"],
            ["税号", "28046190。"],
            ["产品范围", "以氯硅烷为原料，通过改良西门子法或硅烷法生产的太阳能级多晶硅棒状、块状或颗粒状产品。"],
            ["主要排除", "电子级多晶硅不在措施范围；单晶硅棒/锭、硅片、太阳能电池/组件及含硅分散液等下游/其他产品不应并入。"],
            ["美国税率", "REC Solar Grade/REC Advanced/AE及其他57.0%；Hemlock 53.3%；MEMC Pasadena 53.6%。"],
            ["韩国税率", "OCI 4.4%；Hankook 9.5%；Hanwha 8.9%；SMP及其他88.7%；Woongjin等113.8%。"],
        ],
        "tax_rows": [
            ["Hemlock美国", "53.3%", "6.929%×完税价", "60.229%×完税价"],
            ["REC/其他美国", "57.0%", "7.410%×完税价", "64.410%×完税价"],
            ["OCI韩国", "4.4%", "0.572%×完税价", "4.972%×完税价"],
            ["其他韩国88.7%档", "88.7%", "11.531%×完税价", "100.231%×完税价"],
            ["韩国最高档", "113.8%", "14.794%×完税价", "128.594%×完税价"],
        ],
        "query_rows": [
            ["关键词", "POLYSILICON + China + 两年：18条；可见字段精确去重13条；重复多余5。"],
            ["HS宽池", "HS280461 + China + 两年：18条；可见字段精确去重14条；重复多余4。"],
            ["跨查询", "交集4条；并集23条唯一；关键词独有9、HS独有10。"],
            ["主要类别", "美国多晶硅直达/镜像、越南/印尼太阳能链、单晶硅锭、半导体硅、Nouryon含硅分散液等。"],
            ["数据日期", "2024-08-20至2026-06-26；平台记录不是中国海关全量。"],
        ],
        "query_note": "美国出口视角与环球提单可能是同一物理票的净重/毛重镜像。本报告以日期、主体和约17.1t净重代理归并为3批，不把6条可见行机械相加。",
        "lead_headers": ["日期/路线", "货描与实体", "规模", "判断"],
        "lead_widths": [1600, 4100, 1600, 2060],
        "lead_rows": [
            ["2024-08-25、10-30、2025-01-04 美国→中国", "Sea Trade体系；POLYSILICON；10-30含HSC ST740/ST260", "3批；净重代理约51,300kg；金额空", "受税直达B+；太阳能/电子级和Hemlock税档待核"],
            ["2024-12-12 马来西亚标记→越南→中国", "Runergy Vietnam→Shanghai Mingyao；多晶硅块10-50mm；#&MY；出口退回", "数量字段100,800；金额65,296,070,784", "明确保留MY来源；更像合法马来西亚料返运"],
            ["2026-06-26 印尼→中国", "PT Nusa Solar Indonesia→DMEGC；HIGH PURITY ST600", "2,880kg；47,520金额字段", "ST600与OCI/OCIM马来西亚高度对应；核再出口和原产申报"],
            ["2025-12-17 越南→中国", "Trina Vietnam→Yichang CSG；太阳能多晶硅；#&CN", "数量字段3,710", "中国来源标记的返料/退运，非美韩绕道线索"],
        ],
        "data_files": [
            "原始摘录：太阳能级多晶硅_易迅_POLYSILICON_China_两年.json.gz.b64与太阳能级多晶硅_易迅_HS280461_China_两年.json.gz.b64。",
            "逐票标准化：34_keyword_易迅逐票标准化.csv、34_hs_易迅逐票标准化.csv。",
            "跨查询去重：太阳能级多晶硅_跨查询合并去重.csv，共23条唯一。",
        ],
        "chain_rows": [
            ["美国HSC→中国", "B+核税", "HSC牌号、美国平台原产、3个大宗物理批次", "电子级可能排除；生产商/税款/完税价格未知"],
            ["马来西亚料→越南→中国", "反证/B", "#&MY、退回申报、100.8t规模；OCI马来西亚太阳能级产能", "缺上游A腿/品牌；中国端是否申报Malaysia未知"],
            ["ST600印尼→中国", "B", "具体牌号、2.88t、PT Nusa非多晶硅生产商", "ST600指向马来西亚OCI；无Malaysia→Indonesia同批A腿"],
            ["美国/韩国→第三国→中国", "未形成", "无同柜/批次/重量闭合", "需A腿和中国原产申报/缴款书"],
        ],
        "counterevidence": [
            "OCI官方资料显示韩国Gunsan以电子级为主，马来西亚Sarawak拥有太阳能级多晶硅大规模生产；Malaysia来源可合法不适用美韩AD。",
            "公开贸易聚合页反复把ST600列作OCI/OCIM Malaysia太阳能级牌号，因此印尼ST600首先应核马来西亚原产，而非推定韩国。",
            "HSC同时服务半导体和太阳能产业；HSC/ST牌号本身不足以确认措施范围，必须取得COA、等级、用途和产品规格。",
            "Jinko/Trina越南的单晶硅锭/棒是下游产品，不能因HS280461误归或高纯硅词样并入太阳能级多晶硅措施。",
        ],
        "evidence_boundary": "本报告识别了可调单的美国直达和东南亚再出口记录，但没有证明中国进口申报把美国/韩国原产改成第三国，也没有证实未缴税。平台金额多为出口侧字段且币种/成交条件未核，不能据此计算实际欠税。",
        "entity_rows": [
            ["1", "Sea Trade International / Sea Trade Canada / 实际中国进口人", "按2024-08-25、10-30和2025-01-04调提单、生产商、HSC牌号COA、太阳能/电子级、完税价格和Hemlock税款书。"],
            ["2", "Runergy Vietnam / Shanghai Mingyao", "调越南进口申报106538226910、#&MY原产证、出口退运/返料原因、中国原产申报和税单。"],
            ["3", "PT Nusa Solar Indonesia / DMEGC", "调ST600原厂包装、OCI/OCIM发票、Malaysia→Indonesia A腿、库存领退料和中国报关单。"],
            ["4", "Trina/Jinko越南链", "先按单晶锭/棒和#&CN排除；仅对块状多晶硅原料调原产与税款。"],
        ],
        "steps": [
            "先核3批美国直达的具体生产商和太阳能级属性；若为Hemlock且未缴税，综合增量系数为60.229%×完税价格。",
            "再调Runergy越南退回票的马来西亚原产证及中国申报，验证是否正常按非受税Malaysia申报。",
            "对印尼ST600按包装袋、批号、发票和上游提单追至OCI/OCIM Malaysia；若实际Malaysia原产，风险降级。",
        ],
        "conclusion": "第34项的最高优先级是美国原产多晶硅直达核税，而不是第三国绕道。东南亚记录提供具体ST600/退运线索，但马来西亚太阳能级真实产能和货描#&MY/#&CN构成强反证。当前未形成美韩→第三国→中国的A等级闭环。",
        "closing_points": ["直接核税：3批美国原产多晶硅。", "原产穿透：Runergy马来西亚标记和PT Nusa ST600。", "范围排除：电子级、单晶锭/棒、硅片/组件和含硅分散液。"],
        "sources": [
            ["商务部公告2026年第3号", "https://cacs.mofcom.gov.cn/cacscms/article/jkdc?articleId=186918&type=1", "现行范围、税率与期限。"],
            ["OCI Polysilicon业务", "https://ftp.oci.co.kr/eng/sub/business/poly.asp", "韩国电子级与马来西亚太阳能级生产结构。"],
            ["Hemlock 2024可持续发展报告", "https://www.hscpoly.com/wp-content/uploads/2025/08/2024-Sustainability-Report-Digital.pdf", "美国高纯多晶硅及太阳能/半导体双用途。"],
            ["REC Silicon 2024年报", "https://recsilicon.com/wp-content/uploads/2025/03/rec-silicon-annual-report-2024.pdf", "美国生产状态与库存核查背景。"],
            ["DMEGC/PT Nusa Solar公开产品资料", "https://www.dmegcsolar.com/upload/img/2026-03/69cb79812de7c.pdf", "印尼Batam实体和太阳能制造网络。"],
            ["Volza ST600公开聚合页", "https://www.volza.com/p/polycrystalline/import/hsn-code-28046100/", "二级来源，仅用于回查ST600与Malaysia贸易指纹，不作为违法证据。"],
            ["中国进出口货物原产地条例", "https://xzfg.moj.gov.cn/front/law/detail?LawID=1523", "原产核定。"],
        ],
        "filename": "34_太阳能级多晶硅_反倾销税与第三国转运风险深度核查报告.docx",
    },
    {
        "num": 35,
        "short": "碳钢紧固件",
        "title": "碳钢紧固件反倾销税与第三国转运风险深度核查报告",
        "subtitle": "欧盟、英国原产｜精准货描、列名生产商与超大HS宽池阶段审计",
        "status": "精准关键词43条全部读取；列名实体/英国精准查询为0；HS731815宽池23,608条未全量摘录",
        "rating": "阶段中等：精准池未见受税来源，宽池仍有数据缺口",
        "method": "CARBON STEEL BOLT/SCREW/WASHER × China × 两年全页；KAMAX/NEDSCHROEF/UK实体精准查询；HS731815宽池仅完成规模与首屏排除审计",
        "summary": "三组精准货描共43条、可见字段精确去重42条：India 30、Vietnam 7、Indonesia 3、China 1、United States 1；无EU或United Kingdom。15条印度M8/M10/M12碳钢六角螺栓和5条越南垫圈是范围候选但来自非受税来源；印尼3条螺钉及多条印度小螺钉因杆径≤6mm存在明确范围外理由。KAMAX GMBH、KAMAX S.L.U、NEDSCHROEF FASTENERS及英国原产HS731815精准结果均为0。与此同时，HS731815宽池有23,608条，混入不锈钢、阀门/总成、螺母和小直径件，本轮受页面读取成本限制未全量摘录，故报告只能作为阶段审计，不能宣称“全部碳钢紧固件无风险”。",
        "decision_rows": [
            ["精准货描池", "B-/C", "42条唯一，无EU/UK；20条范围候选来自India/Vietnam"],
            ["列名实体/英国", "负面查询", "KAMAX、NEDSCHROEF和UK+731815精准结果均0；仅说明本查询未见"],
            ["HS宽池完整性", "缺口", "731815共23,608条未全量读取，需下载/API或分国分实体补齐"],
        ],
        "policy_rows": [
            ["措施与期限", "商务部2022年第17号：自2022-06-29起继续征收5年，预计至2027-06-28；受税来源欧盟、英国。"],
            ["税号", "73181200、73181400、73181510、73181590、73182100、73182200、90211000、90212900。"],
            ["产品范围", "木螺钉、自攻螺钉、螺钉和螺栓（可带螺母/垫圈）、垫圈等碳钢紧固件。"],
            ["主要排除", "铁道用螺钉；杆径不超过6毫米的螺钉和螺栓；螺母；民用航空器维护维修用紧固件；以及不属于碳钢或下游总成的产品。"],
            ["欧盟税率", "KAMAX 6.1%；Nedschroef集团5.5%；其他欧盟公司26.0%。"],
            ["英国税率", "所有英国公司26.0%。"],
        ],
        "tax_rows": [
            ["KAMAX", "6.1%", "0.793%×完税价", "6.893%×完税价"],
            ["Nedschroef", "5.5%", "0.715%×完税价", "6.215%×完税价"],
            ["其他欧盟/英国", "26.0%", "3.380%×完税价", "29.380%×完税价"],
        ],
        "query_rows": [
            ["碳钢螺栓", "CARBON STEEL BOLT + China + 两年：20条，全部读取。"],
            ["碳钢螺钉", "CARBON STEEL SCREW + China + 两年：11条，全部读取。"],
            ["碳钢垫圈", "CARBON STEEL WASHER + China + 两年：12条，全部读取。"],
            ["合并去重", "原始43条；可见字段精确去重42条；重复多余1；20条措施范围候选、8条小直径排除候选、5条其他总成排除、1条螺母排除、8条待核。"],
            ["列名实体/英国", "KAMAX GMBH、KAMAX S.L.U、NEDSCHROEF FASTENERS对华查询为0；United Kingdom+HS731815为0。"],
            ["HS宽池", "HS731815 + China + 两年显示23,608条；首屏已确认大量不锈钢、小直径和其他范围外记录；未完成119页@200的全量摘录。"],
        ],
        "query_note": "第35项不能用HS731815总数代表案涉碳钢紧固件；产品材质、杆径、用途、是否螺母/总成均是法定范围门槛。精准查询全页已完成，但HS宽池完整性仍是本报告的主要数据缺口。",
        "lead_headers": ["路线/记录", "规模", "产品判断", "风险判断"],
        "lead_widths": [2300, 1600, 3100, 2360],
        "lead_rows": [
            ["India Krishna→Wenzhou Fengding", "15条；M8/M10/M12", "DIN933碳钢六角螺栓，初步在范围", "非受税印度来源；无EU/UK A腿"],
            ["Vietnam Sunluen→Joy Prosper", "4条；数量字段4,450", "碳弹簧钢retaining washer，HS731822", "#&VN；需核实际制造和原产，但非受税标记"],
            ["Vietnam Techtronic→TTI Partners", "1条；数量字段5", "13×6.5×1mm washer，货描#&CN", "中国来源返运/内部流转线索，非EU/UK绕道"],
            ["Indonesia Meiloon→Dongguan", "3条", "TF3.5/MF6螺钉，杆径≤6mm", "公告排除候选"],
            ["US ECU混载", "1条；整柜毛重18,726", "长货描仅夹带retainer washers，无法分摊", "美国非受税；范围和数量均无法确认"],
        ],
        "data_files": [
            "原始精准摘录：碳钢紧固件_易迅精准查询_China_两年.json及其base64固化文件。",
            "逐票标准化：碳钢紧固件_易迅精准查询逐票标准化.csv（43条）与可见字段去重.csv（42条）。",
            "未覆盖：HS731815宽池23,608条及其他7个措施税号的全量逐页数据；后续须拆分查询或取得页面导出/API。",
        ],
        "chain_rows": [
            ["India碳钢螺栓→中国", "B-/C", "产品范围较清楚、15条同日同实体", "印度非受税；无EU/UK生产商/牌号/上游A腿"],
            ["Vietnam垫圈→中国", "C", "HS731822、规格可见、#&VN/#&CN", "当地加工/中国返料可合法解释；无受税A腿"],
            ["KAMAX/Nedschroef→中国", "未见", "精准实体查询0", "平台名称变体/货代/隐藏实体及其他HS仍可能漏报"],
            ["EU/UK→第三国→中国", "未形成", "无同批牌号/柜号/数量闭合", "HS宽池未全量、缺A腿和中国原产申报"],
        ],
        "counterevidence": [
            "KAMAX官方披露除欧盟外还在中国、美国、墨西哥、巴西、印度等地布局生产；KAMAX品牌不等于欧盟原产。",
            "Nedschroef公开历史显示保加利亚、中国等制造/运营节点；集团名同样不能替代具体生产厂和原产地。",
            "越南、印度、印尼存在真实紧固件和零部件加工能力；只有材质/尺寸/工序、投入原料和生产记录才能判断原产。",
        ],
        "evidence_boundary": "本轮没有发现欧盟或英国原产的精准货描记录，更没有绕道闭环；但HS宽池未全量读取，不能把“精准条件0”外推为平台全库0或无风险。任何企业违法判断都必须等到受税A腿、中国B腿和中国报关/税单闭合。",
        "entity_rows": [
            ["1", "HS731815宽池中的#&DE/#&EU/#&GB及受税生产商变体", "按原产、供应商、品牌、M7以上杆径、材质carbon steel分批查询；每批200条/页全读。"],
            ["2", "Krishna / Wenzhou Fengding", "调15条同日螺栓的提单/发票/原产证和印度工厂证明，排除贸易商采购欧盟货。"],
            ["3", "Sunluen Vietnam / Joy Prosper", "调垫圈原料、冲压/热处理工单、BOM、设备/能耗、原产证和中国报关单。"],
            ["4", "KAMAX与Nedschroef名称变体", "检索S.L.U/S.R.O/Schrozberg/Fraulautern/Helmond/Herentals及货代字段，覆盖全部8个税号。"],
        ],
        "steps": [
            "先把23,608条731815宽池按原产EU/UK、供应商实体和材质/杆径拆成可全页读取的小查询。",
            "对其他7个措施税号执行相同矩阵，逐票排除螺母、小直径、航空维修和非碳钢。",
            "仅对出现受税生产商/原产标记的第三国B腿回查30—180日A腿和柜号，随后调中国税单。",
        ],
        "conclusion": "第35项精准货描审计未见EU/UK来源或列名生产商对华记录，现有42条唯一记录主要是印度非受税螺栓、越南垫圈和范围外小螺钉；不支持第三国绕道结论。但HS731815宽池23,608条尚未全量审计，报告必须保持“阶段性、数据缺口未关闭”的结论。",
        "closing_points": ["不得把23,608条宽池当案涉量。", "不得把KAMAX/Nedschroef精准0外推为全库0。", "下一步以8税号×受税来源/实体×材质/杆径拆分补齐。"],
        "sources": [
            ["商务部公告2022年第17号", "https://www.mofcom.gov.cn/zcfb/dwmygl/art/2022/art_e31a129b09594fdcb734a620a0d06f14.html", "范围、排除、税率与期限。"],
            ["KAMAX全球生产网络/Fontana交易", "https://www.kamax.com/en/story/kamax-will-join-fontana-gruppo/", "多国生产反证。"],
            ["KAMAX中国新工厂", "https://www.kamax.com/en/story/kamax-eroeffnet-zweites-produktionswerk-in-china/", "中国本地产能。"],
            ["Nedschroef历史", "https://www.nedschroef.com/en/company/history", "保加利亚、中国等节点。"],
            ["Nedschroef全球网络", "https://www.nedschroef.com/en/company", "集团多地运营。"],
            ["中国进出口货物原产地条例", "https://xzfg.moj.gov.cn/front/law/detail?LawID=1523", "多国加工原产核定。"],
        ],
        "filename": "35_碳钢紧固件_反倾销税与第三国转运风险阶段核查报告.docx",
    },
]


def main():
    outputs = []
    for spec in ITEMS:
        outputs.append(build_report(spec))
    for path in outputs:
        print(path)


if __name__ == "__main__":
    main()
