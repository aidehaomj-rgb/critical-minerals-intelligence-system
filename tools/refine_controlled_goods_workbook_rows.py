"""Refine workbook-level leads to individual spreadsheet-row evidence."""

from __future__ import annotations

import csv
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

import analyze_controlled_goods_communications as base


sys.stdout.reconfigure(encoding="utf-8", errors="replace")

INPUT_CSV = base.OUTPUT_DIR / "候选通讯与规避线索_逐文档.csv"
OUTPUT_CSV = base.OUTPUT_DIR / "候选通讯与规避线索_逐交易行复核.csv"
SUMMARY_JSON = base.OUTPUT_DIR / "逐交易行复核摘要.json"

# Explicit TLD endings prevent OCR/cell-concatenation artifacts such as
# ``name@example.comTEL`` from being treated as a new domain.
STRICT_EMAIL_RE = re.compile(
    r"(?<![A-Z0-9._%+\-])"
    r"[A-Z0-9._%+\-]{1,64}@"
    r"(?:[A-Z0-9\-]{1,63}\.)+"
    r"(?:COM\.BR|COM\.CN|COM\.HK|CO\.UK|COM\.AU|CO\.JP|CO\.KR|"
    r"COM|NET|ORG|EDU|GOV|CN|HK|UK|DE|FR|GR|EU|VN|TH|JP|KR|AU|NZ|BR|CO)"
    r"(?![A-Z])",
    re.I,
)


def extract_emails(text: str) -> list[str]:
    results: list[str] = []
    for raw in STRICT_EMAIL_RE.findall(text):
        value = raw.lower().rstrip(".")
        local, domain = value.rsplit("@", 1)
        # A tax/registration number is sometimes concatenated immediately before
        # an address in flattened customs text.  Trim only a clearly artificial
        # 8+ digit prefix followed by an alphabetic local part.
        local = re.sub(r"^\d{8,}(?=[a-z])", "", local)
        value = f"{local}@{domain}"
        if value not in results:
            results.append(value)
    return results


def load_document_candidates() -> list[dict[str, str]]:
    with INPUT_CSV.open("r", encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def score_row(matched: dict[str, list[str]], emails: list[str]) -> int:
    has_control = "管制或两用物项" in matched
    has_evasion = any(
        name in matched
        for name in ("明确申报不实或改写", "明确藏匿或未申报", "第三国转口或主体替换")
    )
    has_parameter = "型号参数可核查字段" in matched
    has_communication = bool(emails) or "通讯字段" in matched
    return 3 * has_control + 3 * has_evasion + 2 * has_communication + has_parameter


def main() -> None:
    candidates = load_document_candidates()
    document_map: dict[tuple[str, str], dict[str, str]] = {}
    for row in candidates:
        relative = row["relative_path"]
        if not relative.lower().endswith(".xlsx"):
            continue
        document_map[(row["case_id"], relative)] = row

    output_rows: list[dict] = []
    workbook_errors: list[dict] = []
    scanned_rows = 0
    scanned_workbooks = 0
    case_high_rows: defaultdict[str, int] = defaultdict(int)
    group_counts: Counter[str] = Counter()

    for (case_id, relative), document in sorted(document_map.items()):
        path = base.CASE_ROOT / Path(relative)
        if not path.exists():
            workbook_errors.append({"case_id": case_id, "path": str(path), "error": "missing"})
            continue
        try:
            book = load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:
            workbook_errors.append({"case_id": case_id, "path": str(path), "error": str(exc)})
            continue
        scanned_workbooks += 1
        try:
            for sheet in book.worksheets:
                for row_number, values in enumerate(sheet.iter_rows(values_only=True), 1):
                    cells = [base.clean(value) for value in values if base.clean(value)]
                    if not cells:
                        continue
                    scanned_rows += 1
                    row_text = " | ".join(cells)
                    matched = {
                        name: base.terms(row_text, pattern)
                        for name, pattern in base.PATTERNS.items()
                    }
                    matched = {name: found for name, found in matched.items() if found}
                    emails = extract_emails(row_text)
                    score = score_row(matched, emails)

                    # Keep rows that contain an address, or strong row-local
                    # convergence.  This excludes workbook-wide false joins.
                    if not emails and score < 6:
                        continue
                    if score >= 8:
                        case_high_rows[case_id] += 1
                    group_counts.update(matched.keys())
                    output_rows.append(
                        {
                            "risk_score": score,
                            "case_id": case_id,
                            "package_name": document["package_name"],
                            "relative_path": relative,
                            "sheet": sheet.title,
                            "row_number": row_number,
                            "emails": "; ".join(emails),
                            "matched_groups": "; ".join(matched),
                            "matched_terms": json.dumps(matched, ensure_ascii=False),
                            "row_evidence": row_text[:6000],
                        }
                    )
        finally:
            book.close()

    output_rows.sort(
        key=lambda row: (
            -int(row["risk_score"]), row["case_id"], row["relative_path"],
            row["sheet"], int(row["row_number"]),
        )
    )
    base.write_csv(
        OUTPUT_CSV,
        output_rows,
        [
            "risk_score", "case_id", "package_name", "relative_path", "sheet",
            "row_number", "emails", "matched_groups", "matched_terms", "row_evidence",
        ],
    )

    summary = {
        "candidate_workbooks": len(document_map),
        "scanned_workbooks": scanned_workbooks,
        "scanned_nonempty_rows": scanned_rows,
        "retained_rows": len(output_rows),
        "rows_with_email": sum(bool(row["emails"]) for row in output_rows),
        "unique_normalized_emails": len(
            {
                email.strip()
                for row in output_rows
                for email in row["emails"].split(";")
                if email.strip()
            }
        ),
        "high_priority_rows_score_8_or_9": sum(int(row["risk_score"]) >= 8 for row in output_rows),
        "high_priority_cases": dict(sorted(case_high_rows.items())),
        "matched_group_retained_rows": dict(group_counts),
        "workbook_errors": workbook_errors,
        "method_note": "All risk groups on a retained row co-occur in the same spreadsheet row.",
    }
    SUMMARY_JSON.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
