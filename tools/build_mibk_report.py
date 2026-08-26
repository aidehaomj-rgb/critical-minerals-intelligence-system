"""Build the item-25 MIBK deep-audit report and supporting registers."""

from __future__ import annotations

import csv
import json
from pathlib import Path
from urllib.parse import quote

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor


OUT = Path(r"D:\易迅数据\反倾销税深度分析报告\25_甲基异丁基酮")
DOCX = OUT / "甲基异丁基（甲）酮_反倾销税与第三国转运风险深度审计报告.docx"
SUMMARY_PATH = OUT / "MIBK_易迅审计摘要.json"
DETAIL_PATH = OUT / "MIBK_重点AB腿_易迅详情.json"
HIGH_CONF_A = OUT / "MIBK_受税来源至第三国A腿_高置信226条.csv"

FONT = "Calibri"
EAST_ASIA = "等线"
NAVY = RGBColor(11, 37, 69)
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(91, 105, 119)
RISK = RGBColor(155, 28, 28)
AMBER = RGBColor(122, 90, 0)
WHITE = RGBColor(255, 255, 255)
LIGHT_GRAY = "F2F4F7"
LIGHT_BLUE = "E8EEF5"
PALE_BLUE = "F4F6F9"
PALE_AMBER = "FFF8E8"
PALE_RED = "FDF2F2"
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def set_run_font(run, size=11, color=None, bold=None, italic=None):
    run.font.name = FONT
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), FONT)
    rfonts.set(qn("w:hAnsi"), FONT)
    rfonts.set(qn("w:eastAsia"), EAST_ASIA)
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_fill(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, bottom=80, start=120, end=120):
    tcpr = cell._tc.get_or_add_tcPr()
    old = tcpr.find(qn("w:tcMar"))
    if old is not None:
        tcpr.remove(old)
    mar = OxmlElement("w:tcMar")
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{tag}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tcpr.append(mar)


def set_table_geometry(table, widths_dxa):
    if sum(widths_dxa) != TABLE_WIDTH_DXA:
        raise ValueError(f"table widths must sum to {TABLE_WIDTH_DXA}: {widths_dxa}")
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tblpr = table._tbl.tblPr
    for tag in ("w:tblW", "w:tblInd", "w:tblLayout", "w:tblCellMar"):
        node = tblpr.find(qn(tag))
        if node is not None:
            tblpr.remove(node)
    tblw = OxmlElement("w:tblW")
    tblw.set(qn("w:w"), str(TABLE_WIDTH_DXA))
    tblw.set(qn("w:type"), "dxa")
    tblpr.append(tblw)
    tblind = OxmlElement("w:tblInd")
    tblind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tblind.set(qn("w:type"), "dxa")
    tblpr.append(tblind)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr.append(layout)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths_dxa[idx] / 1440)
            tcpr = cell._tc.get_or_add_tcPr()
            tcw = tcpr.find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                tcpr.append(tcw)
            tcw.set(qn("w:w"), str(widths_dxa[idx]))
            tcw.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def mark_repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    trpr.append(header)


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])
    set_run_font(run, 9, MUTED)


def add_hyperlink(paragraph, text, url, color="2E74B5"):
    rid = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.extend([c, underline])
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.extend([rpr, text_node])
    link.append(run)
    paragraph._p.append(link)


def create_numbering(doc, bullet=False):
    numbering = doc.part.numbering_part.element
    abs_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abs_id = max(abs_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abs_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), "bullet" if bullet else "decimal")
    txt = OxmlElement("w:lvlText")
    txt.set(qn("w:val"), "•" if bullet else "%1.")
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "160")
    spacing.set(qn("w:line"), "280")
    spacing.set(qn("w:lineRule"), "auto")
    ppr.extend([tabs, ind, spacing])
    level.extend([start, fmt, txt, suff, ppr])
    abstract.append(level)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    absref = OxmlElement("w:abstractNumId")
    absref.set(qn("w:val"), str(abs_id))
    num.append(absref)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id):
    ppr = paragraph._p.get_or_add_pPr()
    numpr = ppr.find(qn("w:numPr"))
    if numpr is None:
        numpr = OxmlElement("w:numPr")
        ppr.append(numpr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    numpr.extend([ilvl, num])


def add_paragraph(doc, text="", bold=False, color=None, size=11, italic=False, after=6, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.10
    p.paragraph_format.keep_with_next = keep
    set_run_font(p.add_run(text), size, color, bold, italic)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.add_run(text)
    return p


def add_list_item(doc, text, num_id, bold_prefix=None):
    p = doc.add_paragraph()
    apply_numbering(p, num_id)
    if bold_prefix and text.startswith(bold_prefix):
        set_run_font(p.add_run(bold_prefix), 11, None, True)
        set_run_font(p.add_run(text[len(bold_prefix):]), 11)
    else:
        set_run_font(p.add_run(text), 11)
    return p


def add_callout(doc, label, text, fill=PALE_BLUE, color=NAVY):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    set_cell_fill(cell, fill)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    set_run_font(p.add_run(f"{label}  "), 11, color, True)
    set_run_font(p.add_run(text), 11, color)
    mark_repeat_header(table.rows[0])
    set_table_geometry(table, [TABLE_WIDTH_DXA])
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return table


def add_table(doc, headers, rows, widths_dxa, font_size=9.2, center_cols=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    center_cols = set(center_cols or [])
    for idx, head in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_fill(cell, LIGHT_GRAY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(str(head)), font_size, DARK_BLUE, True)
    mark_repeat_header(table.rows[0])
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx in center_cols else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            set_run_font(p.add_run(str(value)), font_size)
    set_table_geometry(table, widths_dxa)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return table


def configure_document(doc):
    section = doc.sections[0]
    doc.settings.odd_and_even_pages_header_footer = False
    section.different_first_page_header_footer = False
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.right_margin = section.bottom_margin = section.left_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.492)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), EAST_ASIA)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), EAST_ASIA)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    if "Report Title" not in styles:
        title_style = styles.add_style("Report Title", WD_STYLE_TYPE.PARAGRAPH)
    else:
        title_style = styles["Report Title"]
    title_style.font.name = FONT
    title_style._element.rPr.rFonts.set(qn("w:eastAsia"), EAST_ASIA)
    title_style.font.size = Pt(23)
    title_style.font.bold = True
    title_style.font.color.rgb = NAVY
    title_style.paragraph_format.space_before = Pt(12)
    title_style.paragraph_format.space_after = Pt(4)
    title_style.paragraph_format.keep_with_next = True
    hp = section.header.paragraphs[0]
    hp.paragraph_format.space_after = Pt(0)
    set_run_font(hp.add_run("反倾销税深度分析  |  第25项"), 9, MUTED, True)
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(fp.add_run("甲基异丁基（甲）酮深度审计  ·  "), 9, MUTED)
    add_page_field(fp)


def read_json(path):
    return json.loads(path.read_text(encoding="utf-8"))


def write_supporting(summary):
    queries = [
        ["MIBK-Q1", "METHYL ISOBUTYL KETONE", "全球；2025-08-14—2026-08-14", "4,036条 / 21页 / 已至末页"],
        ["MIBK-Q2", "HS 291413", "全球；同期间", "1,314条 / 7页 / 已至末页"],
        ["MIBK-Q3", "4-METHYL-2-PENTANONE", "全球；同期间", "1,250条 / 7页 / 已至末页"],
        ["MIBK-Q4", "MIBK", "全球；同期间", "991条 / 5页 / 已至末页"],
        ["MIBK-Q5", "CAS108-10-1 + HS291413", "目的国中国；同期间", "136条 / 1页 / 已读取全部"],
        ["MIBK-Q6", "CAS108-10-1", "全球；同期间", "平台拆词100,897条/505页且超过4万展示上限；不宣称全量"],
    ]
    with (OUT / "甲基异丁基酮_易迅查询覆盖清单.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["查询ID", "关键词/税号", "范围", "完整性"])
        writer.writerows(queries)

    gaps = [
        ["D-01", "已完成", "5个可控查询逐页到末页", "合并7,727条次、去重6,321条"],
        ["D-02", "受平台限制", "CAS108-10-1全球检索", "拆词噪声且超过展示上限；已用名称、结构名、税号及中国定向CAS补漏"],
        ["D-03", "待调取", "重点A/B链的共同柜号、批号、PO、发票", "现有两腿提单号不同，B腿无柜号和中国收货人"],
        ["D-04", "待调取", "中国报关单、原产地证、COA、税款缴款书", "决定是否实际韩国原产、是否已缴AD及正式税差"],
        ["D-05", "待核", "B腿MIBK实际净重", "整票14,309kg货描并列MIBK与苯酚，不得全部计作MIBK"],
    ]
    with (OUT / "甲基异丁基酮_数据覆盖与缺口.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["编号", "状态", "项目", "结论边界"])
        writer.writerows(gaps)

    sources = [
        ["商务部公告2024年第8号", "措施期限、产品范围、税号、税率和征收公式", "https://cacs.mofcom.gov.cn/cacscms/article/jkdc?articleId=179929&type="],
        ["Kumho P&B MIBK", "CAS、纯度、包装、韩国生产支持联系信息", "https://www.kpb.co.kr/eng/product/mibk"],
        ["Sasol MIBK", "99.5%纯度、CAS及全球供应区域", "https://products.sasol.com/pic/products/home/grades/ZA/5mibk/index.html"],
        ["Mitsubishi Chemical MIBK", "日本企业MIBK产品与用途", "https://www.m-chemical.co.jp/products/departments/mcc/phl/product/1200360_88690.html"],
        ["Mitsui Chemicals AP MIBK", "MIBK产品及销售口径", "https://ap.mitsuichemicals.com/service/product/methylisobutylketone/index.htm"],
        ["GLOVIS America", "第三方物流、报关、仓储及分拨业务属性", "https://www.glovisusa.com/"],
        ["Cole International", "Houston地址及报关/货代属性", "https://www.coleintl.com/contact-us/"],
        ["DuPont DDP legal entities", "DDP Specialty Electronic Materials US 5, LLC及Collegeville地址", "https://www.dupont.com/content/dam/dupont/amer/us/en/corporate/supplier-center/documents/rh-legal-entity-changes/Apr2024_RH-LegalEntityNameChanges-Supplier-APPENDIX_EN.pdf"],
    ]
    with (OUT / "甲基异丁基酮_公开来源台账.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["来源", "用途", "网址"])
        writer.writerows(sources)

    stage = {
        "item": 25,
        "product": "甲基异丁基（甲）酮 MIBK",
        "cas": "108-10-1",
        "hs": "29141300",
        "taxed_origins": ["韩国", "日本", "南非"],
        "measure_end": "2029-03-19",
        "yixun": summary,
        "closed_rerouting_chains": 0,
        "priority_chain": {
            "grade": "B+",
            "route": "韩国→美国长滩→中国上海",
            "a_leg": "2025-11-14，14,480kg，80 PKG，韩国原产MIBK",
            "b_leg": "2025-11-19，14,309kg，69 PK，货描并列MIBK和苯酚",
            "weight_similarity": "98.82%",
            "conclusion": "高度优先调单线索；未形成同柜/同批/同主体和中国税单闭环",
        },
    }
    (OUT / "甲基异丁基酮_阶段审计摘要.json").write_text(
        json.dumps(stage, ensure_ascii=False, indent=2), encoding="utf-8"
    )


def write_high_conf_a():
    source = OUT / "MIBK_受税来源至第三国A腿.csv"
    rows = list(csv.DictReader(source.open(encoding="utf-8-sig", newline="")))
    high = [r for r in rows if r.get("scope_class") in {"范围内明确", "范围内化学品-试剂/样品"}]
    fields = list(rows[0].keys()) if rows else []
    with HIGH_CONF_A.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(high)
    if len(high) != 226:
        raise AssertionError(f"expected 226 high-confidence A-leg rows, got {len(high)}")


def build_report():
    summary = read_json(SUMMARY_PATH)
    details = read_json(DETAIL_PATH)
    write_supporting(summary)
    write_high_conf_a()

    doc = Document()
    configure_document(doc)
    bullet_id = create_numbering(doc, bullet=True)
    decimal_id = create_numbering(doc, bullet=False)

    title = doc.add_paragraph(style="Report Title")
    title.add_run("甲基异丁基（甲）酮反倾销税\n与第三国转运风险")
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    set_run_font(subtitle.add_run("第25项  |  易迅全页逐票审计与公开证据深度核查"), 14, MUTED)

    metadata = [
        ("审计商品", "MIBK / 4-Methyl-2-Pentanone / CAS 108-10-1"),
        ("中国税号", "29141300"),
        ("受税来源", "韩国、日本、南非"),
        ("数据窗口", "易迅近一年：2025-08-14—2026-08-14"),
        ("报告日期", "2026-08-20"),
        ("证据结论", "B+高优先调单线索；尚无A级绕道/逃税闭环"),
    ]
    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        set_run_font(p.add_run(f"{label}："), 11, NAVY, True)
        set_run_font(p.add_run(value), 11)

    add_callout(
        doc,
        "核心结论",
        "易迅检出1条最具体的“韩国→美国长滩→上海”近时点、近等重链路：A腿14,480kg，5天后B腿14,309kg，重量相似度98.82%。但B腿货描同时列MIBK与苯酚，主体、提单、柜号和批号均未闭合，故只能列B+调单线索，不能认定绕道、伪报原产或少缴税。",
        PALE_RED,
        RISK,
    )

    add_heading(doc, "一、结论与处置优先级")
    add_list_item(doc, "风险判断：商品本身反倾销税率梯度高，具体交易线索取证价值高；违法事实尚未证实。", bullet_id, "风险判断：")
    add_list_item(doc, "数量口径：B腿整票重量为14,309kg，但因货描并列MIBK和苯酚，MIBK实际净重未知，14,309kg只能作为整票筛查上限。", bullet_id, "数量口径：")
    add_list_item(doc, "税种边界：若最终证实货物为韩国原产且未征税，风险为反倾销税及其引致的进口增值税差额；现无中国完税价格、生产商税档和税单，不能给出实际欠税。", bullet_id, "税种边界：")
    add_list_item(doc, "优先动作：先调中国进口报关单和税款缴款书，再调两腿提单/柜号/批号及DDP采购、库存和出库记录；未取得这些材料前不得升级为A级。", bullet_id, "优先动作：")

    section_two = add_heading(doc, "二、现行反倾销措施与税差公式")
    section_two.paragraph_format.page_break_before = True
    add_paragraph(
        doc,
        "商务部公告2024年第8号决定，自2024年3月20日起继续对原产于韩国、日本和南非的MIBK征收反倾销税5年，预计至2029年3月19日。范围为4-甲基-2-戊酮，英文Methyl Isobutyl Ketone / 4-Methyl-2-Pentanone，税号29141300。",
    )
    add_table(
        doc,
        ["受税来源/企业", "AD税率", "AD+13%VAT增量系数"],
        [
            ("韩国：Kumho P&B", "18.5%", "20.905% × V"),
            ("韩国：其他公司", "32.3%", "36.499% × V"),
            ("日本：Mitsui", "45.0%", "50.850% × V"),
            ("日本：Mitsubishi", "47.8%", "54.014% × V"),
            ("日本：其他公司", "190.4%", "215.152% × V"),
            ("南非：Sasol", "15.9%", "17.967% × V"),
            ("南非：其他公司", "34.1%", "38.533% × V"),
        ],
        [3500, 1800, 4060],
        center_cols={1, 2},
    )
    add_callout(
        doc,
        "公式边界",
        "V仅指中国海关审定完税价格。综合系数=反倾销税率×1.13，只表示反倾销税及由其增加的13%进口增值税；实际增值税率以税单为准。A腿美国进口申报价值USD289,600不是中国完税价格，不用于正式追税测算。",
        PALE_AMBER,
        AMBER,
    )

    add_heading(doc, "三、易迅全页查询与逐票审计")
    add_paragraph(doc, "本轮通过当前登录会话的只读查询接口逐页读取页面数据，不触发下载中心，不消耗下载额度。5组可控查询均核对平台总数、页数和末页；CAS全球查询因拆词噪声超过平台4万条展示上限，单独列为限制，不冒充全量。")
    add_table(
        doc,
        ["查询口径", "平台总数", "页数", "完整性"],
        [
            ("METHYL ISOBUTYL KETONE｜全球", "4,036", "21", "全部读取"),
            ("HS 291413｜全球", "1,314", "7", "全部读取"),
            ("4-METHYL-2-PENTANONE｜全球", "1,250", "7", "全部读取"),
            ("MIBK｜全球", "991", "5", "全部读取"),
            ("CAS+HS｜目的国中国", "136", "1", "全部读取"),
            ("CAS 108-10-1｜全球", "100,897", "505", "拆词噪声且超4万上限"),
        ],
        [4100, 1500, 1200, 2560],
        center_cols={1, 2, 3},
    )
    add_paragraph(doc, "前5组共形成7,727条查询出现次数，按15项可见贸易字段联合去重后为6,321条唯一记录，重叠/重复1,406条次。完整性QA全部通过。", True, NAVY)
    add_table(
        doc,
        ["逐票范围分类", "唯一记录", "审计解释"],
        [
            ("范围内明确", "587", "名称/CAS与291413同时命中"),
            ("试剂/样品", "90", "化学品明确，规模和用途需区分"),
            ("范围待核/待归类", "1,155", "税号空、名称不明、税号冲突或组合货描"),
            ("含MIBK配方/混合物", "3,998", "涂料、胶黏剂、清洗剂等，不能按MIBK单体计"),
            ("相邻化学品/无MIBK配方", "280", "双丙酮醇、MIBK-free等"),
            ("过氧化物", "20", "MIBK过氧化物/引发剂，排除"),
            ("其他假阳性", "191", "未满足名称/CAS/税号范围条件"),
        ],
        [3350, 1450, 4560],
        center_cols={1},
    )
    add_list_item(doc, "受税来源→第三国A腿：561条范围候选，其中范围内明确225条、试剂/样品1条，其余335条待成分/归类核查。", bullet_id)
    add_list_item(doc, "A腿来源结构：韩国284条、南非205条、日本72条；主要目的地为印度270条、越南116条、美国29条。", bullet_id)
    add_list_item(doc, "受税来源直达中国：0条；第三国/非受税来源对华B腿：1条；A/B候选匹配：1条。这里的“0条”仅指本轮近一年、当前查询口径未检出。", bullet_id)

    add_heading(doc, "四、重点A/B链逐字段核查")
    add_table(
        doc,
        ["字段", "A腿：韩国→美国", "B腿：美国→中国"],
        [
            ("易迅记录ID", "USIMP0294139159", "USEXP0029742362"),
            ("关键日期", "2025-11-14 抵长滩", "2025-11-19 离长滩；间隔5天"),
            ("货描/税号", "MIBK；HS291413", "[MIBK, PHENOL；税号空"),
            ("平台原产/目的", "South Korea → United States", "United States → China"),
            ("重量/包装", "14,480kg；80 PKG", "14,309kg；69 PK"),
            ("重量关系", "两腿相差171kg", "相似度98.82%"),
            ("可见主体", "GLOVIS America → Cole International USA", "DDP Specialty Electronic Materials → 中国收货人空"),
            ("港口", "Busan/Pusan → Long Beach", "Long Beach → Shanghai"),
            ("运输", "HMM EMERALD / 0010E", "COSCO BELGIUM / 081W；OOCL"),
            ("提单/柜封", "B/L HDMUSELM82373400；BSIU3218631；封志25H0300616", "B/L 2163395920；AES 6004202511AES0000025235；柜号未载"),
            ("金额", "USD289,600（A腿申报值）", "空；不得以A腿值替代中国完税价格"),
        ],
        [1900, 3730, 3730],
        font_size=8.5,
    )

    add_heading(doc, "五、为何是B+，而不是已证实绕道")
    add_heading(doc, "5.1 支持升级调单的事实", level=2)
    add_list_item(doc, "时序紧密：A腿抵港后5天即出现从同一港口发往上海的B腿。", bullet_id)
    add_list_item(doc, "重量高度接近：14,480kg与14,309kg，相差171kg，相似度98.82%。", bullet_id)
    add_list_item(doc, "商品关键词重合：两腿均明确出现MIBK；A腿还载明HS291413、UN1245和危险品3类。", bullet_id)
    add_list_item(doc, "节点可追踪：A腿具有主提单、柜号、封志和船名航次；B腿具有提单、AES编号、船名航次和上海到港信息。", bullet_id)
    add_heading(doc, "5.2 阻止定性的关键缺口", level=2)
    add_list_item(doc, "无共同主体：A腿收发货人为Cole/GLOVIS，B腿为DDP，算法未发现实体精确重合。", bullet_id)
    add_list_item(doc, "无同柜/同提单：两腿提单号不同，B腿未披露柜号、封志或批号，不能证明是同一货物。", bullet_id)
    add_list_item(doc, "B腿范围不清：货描同时列MIBK和苯酚，整票重量不能直接当作MIBK净重。", bullet_id)
    add_list_item(doc, "法定原产和纳税未知：平台“United States”字段不等于中国进口报关原产国；缺中国报关单、原产地证和税款缴款书。", bullet_id)
    add_heading(doc, "5.3 合法替代解释", level=2)
    add_list_item(doc, "GLOVIS America与Cole International公开资料均显示物流/报关/货代属性，A腿可见主体不是韩国生产商。", bullet_id)
    add_list_item(doc, "DDP Specialty Electronic Materials US 5, LLC为DuPont体系美国法律实体，公开地址与B腿Collegeville地址一致；这能解释美国供应主体，但不能证明MIBK制造地。", bullet_id)
    add_list_item(doc, "B腿可能是美国库存、独立采购、电子材料相关多品项合并货载或其他非A腿来源；只有库存批次、采购发票和出库记录才能排除这些解释。", bullet_id)

    add_heading(doc, "六、涉税数量与税额边界")
    add_table(
        doc,
        ["项目", "当前可写", "当前不可写"],
        [
            ("涉税数量", "整票筛查上限14,309kg；MIBK净重未知", "不能写14,309kg全部为MIBK"),
            ("真实原产", "A腿平台显示South Korea", "不能据此认定B腿或中国申报仍为韩国原产"),
            ("生产商税档", "若韩国原产，可能为Kumho 18.5%或其他32.3%", "GLOVIS/Cole不是生产商，不能直接适用Kumho税率"),
            ("税额", "条件系数20.905%×V或36.499%×V", "无中国完税价格V和税单，不得写实际少缴税额"),
        ],
        [1700, 3830, 3830],
        font_size=9.0,
    )
    add_callout(doc, "条件情景", "只有在确认B腿货物实际为韩国原产MIBK、属于中国进口涉案范围且未缴反倾销税后，才按相应韩国生产商税率计算。A腿USD289,600即使可见，也不能替代中国海关完税价格。", PALE_AMBER, AMBER)

    add_heading(doc, "七、实体和单证核查顺序")
    steps = [
        "中国端：调取B/L 2163395920对应的中国进口报关单、境内收货人/消费使用单位、申报企业、入境口岸、原产国、生产商、净重、完税价格、AD税率和税款缴款书。",
        "美国B腿：向DDP Specialty Electronic Materials US 5, LLC核2025-11-19出货的采购合同、供应商发票、库存批次、出库单、SDS/COA、MIBK实际分项重量及AES底单。",
        "韩国A腿：向承运人/货代调主提单HDMUSELM82373400、柜BSIU3218631、封志25H0300616、实际托运人/生产商、商业发票、COA和韩国原产证明。",
        "链路比对：核两腿批号、PO、装桶/罐箱号、危化品申报、仓库入出库、付款受益人及Long Beach短期换装/拆拼记录。",
        "实体穿透：核GLOVIS、Cole、Radix、DDP与实际买卖双方的代理关系；不要把报关行、NVOCC或通知方误写为生产商。",
    ]
    for step in steps:
        add_list_item(doc, step, decimal_id)

    add_heading(doc, "八、公开互联网深检结论")
    add_list_item(doc, "官方政策证据充分：商务部明确产品范围、税号、全部企业税率和征税公式。", bullet_id)
    add_list_item(doc, "生产/产品证据充分：Kumho P&B、Sasol、Mitsubishi和Mitsui均有公开MIBK产品信息；这支持实体与牌号核验，但不证明本票来源。", bullet_id)
    add_list_item(doc, "物流属性得到公开反证：GLOVIS和Cole公开资料显示第三方物流、报关、货代或仓储职能，不能把其国籍当成货物原产地。", bullet_id)
    add_list_item(doc, "负面检索：截至2026-08-20，针对海关总署、最高人民法院和商务部官网的定向检索，未发现本商品公开的中国反规避裁定、海关处罚、法院判决或已披露双段提单案件。该负面结果不等于现实中不存在案件。", bullet_id)

    add_heading(doc, "九、最终判断")
    add_callout(
        doc,
        "证据等级：B+",
        "已形成具体日期、港口、重量、货描和提单字段可核的“韩国→美国长滩→上海”链路，足以进入高优先调单；但未形成同柜、同批、同主体、法定原产和未税事实闭环。当前最准确表述是“存在近时点近等重转贸风险线索，第三国绕道和逃避反倾销税尚未证实”。",
        PALE_RED,
        RISK,
    )

    add_heading(doc, "十、主要来源")
    sources = [
        ("商务部公告2024年第8号：现行措施、税率、范围与公式", "https://cacs.mofcom.gov.cn/cacscms/article/jkdc?articleId=179929&type="),
        ("Kumho P&B：MIBK产品、CAS、纯度和包装", "https://www.kpb.co.kr/eng/product/mibk"),
        ("Sasol：MIBK 99.5%及供应区域", "https://products.sasol.com/pic/products/home/grades/ZA/5mibk/index.html"),
        ("Mitsubishi Chemical：MIBK产品与用途", "https://www.m-chemical.co.jp/products/departments/mcc/phl/product/1200360_88690.html"),
        ("Mitsui Chemicals AP：MIBK产品与销售口径", "https://ap.mitsuichemicals.com/service/product/methylisobutylketone/index.htm"),
        ("GLOVIS America：物流、报关和仓储业务", "https://www.glovisusa.com/"),
        ("Cole International：Houston地址及报关/货代业务", "https://www.coleintl.com/contact-us/"),
        ("DuPont：DDP Specialty Electronic Materials法律实体与地址", "https://www.dupont.com/content/dam/dupont/amer/us/en/corporate/supplier-center/documents/rh-legal-entity-changes/Apr2024_RH-LegalEntityNameChanges-Supplier-APPENDIX_EN.pdf"),
    ]
    for label, url in sources:
        p = doc.add_paragraph()
        apply_numbering(p, bullet_id)
        add_hyperlink(p, label, url)

    add_paragraph(doc, "数据附件：MIBK_易迅逐票标准化_联合去重.csv、MIBK_受税来源至第三国A腿.csv、MIBK_受税来源至第三国A腿_高置信226条.csv、MIBK_第三国对华B腿.csv、MIBK_AB腿候选匹配.csv、MIBK_重点AB腿_易迅详情.json。", False, MUTED, 9.5, False, 0)
    doc.save(DOCX)
    return DOCX


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    print(build_report())
