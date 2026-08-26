"""Inspect incremental GDRM archives without extracting or mutating the case library."""

from __future__ import annotations

import hashlib
import io
import json
import os
import re
import sqlite3
import sys
import zipfile
from pathlib import Path, PurePosixPath

sys.path.insert(0, str(Path(__file__).resolve().parent / ".vendor"))

from openpyxl import load_workbook
from pypdf import PdfReader

try:
    import msoffcrypto
except ImportError:  # pragma: no cover - reported clearly when encountered
    msoffcrypto = None


sys.stdout.reconfigure(encoding="utf-8", errors="replace")

PASSWORD = "GDRM@2026"
DB = Path(os.environ["CASE_DB"])
ARCHIVES = [Path(value) for value in json.loads(os.environ["CASE_ARCHIVES_JSON"])]


def normalize_text(value: str) -> str:
    lines: list[str] = []
    for line in (value or "").replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        cleaned = re.sub(r"[ \t]+", " ", line).strip()
        if cleaned:
            lines.append(cleaned)
    return "\n".join(lines)


def safe_entry(name: str) -> bool:
    path = PurePosixPath(name.replace("\\", "/"))
    return not path.is_absolute() and ".." not in path.parts and path.name != ""


def pdf_info(data: bytes) -> dict:
    reader = PdfReader(io.BytesIO(data))
    encrypted = reader.is_encrypted
    decrypt_result = None
    if encrypted:
        decrypt_result = int(reader.decrypt(PASSWORD))
        if decrypt_result == 0:
            raise RuntimeError("PDF password failed")
    texts = [normalize_text(page.extract_text() or "") for page in reader.pages]
    return {
        "encrypted": encrypted,
        "decrypt_result": decrypt_result,
        "pages": len(reader.pages),
        "characters": sum(map(len, texts)),
        "first_text": texts[0][:900] if texts else "",
    }


def xlsx_info(data: bytes) -> dict:
    encrypted = not data.startswith(b"PK")
    decrypt_result = None
    workbook_data = data
    if encrypted:
        if msoffcrypto is None:
            raise RuntimeError("Encrypted Office file found but msoffcrypto is unavailable")
        source = io.BytesIO(data)
        office = msoffcrypto.OfficeFile(source)
        office.load_key(password=PASSWORD)
        target = io.BytesIO()
        office.decrypt(target)
        workbook_data = target.getvalue()
        decrypt_result = True
    book = load_workbook(io.BytesIO(workbook_data), read_only=True, data_only=True)
    lines: list[str] = []
    for sheet in book.worksheets:
        lines.append(f"[工作表] {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            line = " | ".join(normalize_text(str(cell)) for cell in row if cell is not None)
            if line:
                lines.append(line)
    text = "\n".join(lines)
    return {
        "encrypted": encrypted,
        "decrypt_result": decrypt_result,
        "sheets": len(book.worksheets),
        "characters": len(text),
        "first_text": text[:900],
    }


def main() -> None:
    conn = sqlite3.connect(DB)
    existing_hashes = {row[0] for row in conn.execute("SELECT DISTINCT sha256 FROM documents")}
    existing_cases = {row[0] for row in conn.execute("SELECT case_id FROM cases")}
    output = []
    for archive in ARCHIVES:
        archive_row = {"archive": str(archive), "entries": []}
        with zipfile.ZipFile(archive) as zf:
            for entry in zf.infolist():
                if entry.is_dir():
                    continue
                if not safe_entry(entry.filename):
                    raise RuntimeError(f"Unsafe archive entry: {entry.filename}")
                data = zf.read(entry)
                digest = hashlib.sha256(data).hexdigest()
                match = re.search(r"GDRM26-(\d{3})", PurePosixPath(entry.filename).name, re.I)
                if not match:
                    raise RuntimeError(f"No case ID in entry: {entry.filename}")
                case_id = f"GDRM26-{match.group(1)}"
                suffix = PurePosixPath(entry.filename).suffix.lower()
                row = {
                    "entry": entry.filename,
                    "case_id": case_id,
                    "case_exists": case_id in existing_cases,
                    "bytes": len(data),
                    "sha256": digest,
                    "hash_exists": digest in existing_hashes,
                    "extension": suffix,
                }
                if suffix == ".pdf":
                    row.update(pdf_info(data))
                elif suffix == ".xlsx":
                    row.update(xlsx_info(data))
                else:
                    raise RuntimeError(f"Unsupported file type: {entry.filename}")
                archive_row["entries"].append(row)
        output.append(archive_row)
    conn.close()
    print(json.dumps(output, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
