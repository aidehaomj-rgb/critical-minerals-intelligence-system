"""Analyze all intelligence case packs after encrypted PDFs have been ingested."""

from __future__ import annotations

import csv
import json
import re
import sqlite3
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


sys.stdout.reconfigure(encoding="utf-8", errors="replace")

ROOT = Path(r"D:\codex\XG执法\2026_案例库")
DB = ROOT / "enforcement_cases_2026.sqlite"
CATALOG = ROOT / "case_catalog.csv"
OUTPUT = ROOT / "情报案件解密后全量分析.json"

INTEL = re.compile(r"情報|情报")

THEMES = {
    "香烟及烟草制品": r"香煙|香烟|加熱煙|加热烟|電子煙|电子烟|雪茄|煙草|烟草",
    "毒品及受管制药物": r"毒品|可卡因|海洛因|冰毒|大麻|氯胺酮|藥物|药物|毒藥|毒药",
    "未列舱单货物": r"未列艙單|未列舱单",
    "濒危物种及动植物": r"瀕危物種|濒危物种|活貓|活猫|動物|动物|蜥蜴|穿山甲|魚翅|鱼翅|象牙|種子|种子",
    "冒牌物品": r"冒牌|假冒",
    "燃油": r"汽油|柴油|燃油|加油站",
    "金属及贵金属": r"金屬|金属|貴金屬|贵金属|銀錠|银锭|白銀|白银",
    "酒类": r"烈酒|酒類|酒类",
    "武器": r"武器|槍械|枪械|彈藥|弹药",
    "电子废料": r"電子廢料|电子废料",
}

CHANNELS = {
    "空路客运": r"空路客運|空路客运",
    "空运货物/邮包": r"空運貨物|空运货物|空運郵包|空运邮包|郵包|邮包",
    "陆路客运": r"陸路客運|陆路客运",
    "陆路货车/货运": r"陸路貨車|陆路货车|陸路貨運|陆路货运",
    "陆路小车": r"陸路小車|陆路小车",
    "海河路货运": r"海路貨運|海路货运|河路貨運|河路货运",
    "海路客运/渔船": r"海路客運|海路客运|漁船|渔船",
    "仓储/场所": r"儲存倉|储存仓|加油站",
}

FIELD_LABELS = [
    ("过关日期及时间", re.compile(r"^(?:過關|过关)日期及時間")),
    ("案件日期", re.compile(r"^案件日期")),
    ("案发地点", re.compile(r"^(?:案發|案发)地點")),
    ("运输方式", re.compile(r"^(?:運輸|运输)(?:方式|途徑|途径)")),
    ("抵港日期", re.compile(r"^抵港日期")),
    ("路线", re.compile(r"^(?:路線|路线)")),
    ("车牌号码", re.compile(r"^(?:車牌|车牌)號碼")),
    ("航班资料", re.compile(r"^(?:航班資料|航班资料|航班編號|航班编号)")),
    ("船舶资料", re.compile(r"^(?:船名|船舶名稱|船舶名称|貨船名稱|货船名称|船隻|船只)")),
    ("集装箱编号", re.compile(r"^(?:集裝箱|集装箱)(?:編號|编号)")),
    ("发货人资料", re.compile(r"^(?:發貨人|发货人)(?:資料|资料)")),
    ("收货人资料", re.compile(r"^(?:收貨人|收货人)(?:資料|资料)")),
    ("申报货物(舱单)", re.compile(r"^(?:申報|申报)(?:貨物|货物)\s*[（(](?:艙單|舱单)[）)]")),
    ("申报货物(报关)", re.compile(r"^(?:申報|申报)(?:貨物|货物)\s*[（(](?:報關|报关)[）)]")),
    ("申报货物", re.compile(r"^(?:申報|申报)(?:貨物|货物)")),
    ("被捕人资料", re.compile(r"^被捕人(?:資料|资料)(?:[（(][一二三四五六七八九十\d]+[）)])?")),
    ("检获物品", re.compile(r"^(?:檢獲|检获)物品")),
    ("价值(港元)", re.compile(r"^(?:價值|价值)\s*[（(]?\s*(?:港元)?\$?\s*[）)]?")),
    ("总值(港元)", re.compile(r"^(?:總值|总值)\s*[（(]?\s*(?:港元)?\$?\s*[）)]?")),
    ("应课税值(港元)", re.compile(r"^(?:應課稅值|应课税值)\s*[（(]?\s*(?:港元)?\$?\s*[）)]?")),
    ("收藏手法", re.compile(r"^(?:收藏|藏匿)手法")),
]

COMPANY_SUFFIX = re.compile(
    r"[\u3400-\u9fffA-Za-z0-9&'（）()·.,，\- ]{2,100}"
    r"(?:有限責任公司|有限责任公司|股份有限公司|有限公司|集團|集团|公司|企業|企业)"
)
EN_COMPANY = re.compile(
    r"\b[A-Z][A-Z0-9&'.,()\- ]{3,100}"
    r"(?:CO\.?\s*,?\s*LTD\.?|CO\.?\s*,?\s*LIMITED|LIMITED|LTDA|LTD\.?|LLC|INC\.?|CORPORATION|COMPANY)\b",
    re.I,
)
GENERIC_COMPANY_KEYS = {
    "运输企业",
    "运输企业名称",
    "COMPANY",
    "LIMITED",
    "COMPANYLIMITED",
    "COLTD",
    "COLIMITED",
    "PTELTD",
    "PRIVATE LIMITED",
    "PRIVATELIMITED",
}
PERSON = re.compile(
    r"姓名\s*[:：]\s*([\u3400-\u9fff·]{2,8})"
    r"(?:\s+([A-Z][A-Z ]{2,50}))?",
    re.I,
)
DOCUMENT_ID = re.compile(r"(?:證件|证件)號碼\s*[:：]\s*([A-Z0-9()\-]{5,30})", re.I)
CONTAINER = re.compile(r"\b[A-Z]{4}\s?\d{7}\b")
HK_PLATE = re.compile(r"\b[A-Z]{1,3}\s*\d{1,4}\b")
CN_CROSS_BORDER_PLATE = re.compile(
    r"(?:粤|粵)\s*[A-Z]\s*[A-Z0-9]{4,6}\s*(?:港|澳)?",
    re.I,
)

HEADER_GROUPS = {
    "企业": (
        "發貨人名稱", "发货人名称", "收貨人名稱", "收货人名称", "申報人名稱",
        "申报人名称", "寄件人", "收件人", "承運人", "承运人", "公司名稱",
        "公司名称", "出口商", "进口商", "進口商", "consignee", "shipper",
        "declarant", "company",
    ),
    "人员": ("姓名", "司機", "司机", "駕駛員", "驾驶员", "聯絡人", "联系人", "負責人", "负责人"),
    "车辆": ("車輛登記", "车辆登记", "車牌", "车牌", "拖架", "拖頭", "拖头"),
    "集装箱": ("集裝箱", "集装箱", "貨櫃", "货柜", "container"),
    "电话": ("電話", "电话", "tel", "mobile"),
    "地址": ("地址", "address"),
    "提运单": ("提單", "提单", "運單", "运单", "awb", "waybill"),
}


def clean(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def compact(value: str) -> str:
    return re.sub(r"\s+", "", value or "")


def normalized_key(value: str) -> str:
    text = clean(value).upper()
    text = re.sub(r"[，,。.;；:：'\"()（）\[\]【】/\\]", "", text)
    return re.sub(r"\s+", "", text)


def parse_fields(text: str) -> dict[str, str]:
    fields: dict[str, list[str]] = defaultdict(list)
    current = None
    for raw_line in (text or "").splitlines():
        line = clean(raw_line)
        if not line:
            continue
        matched = False
        for canonical, pattern in FIELD_LABELS:
            match = pattern.match(line)
            if not match:
                continue
            current = canonical
            value = line[match.end():].lstrip(" :：")
            if value:
                fields[current].append(value)
            matched = True
            break
        if not matched and current:
            fields[current].append(line)
    return {key: clean(" ".join(values)) for key, values in fields.items()}


def company_candidates(value: str, explicit_company_field: bool = False) -> list[str]:
    raw = clean(value)
    if not raw or len(raw) > 500:
        return []
    found = COMPANY_SUFFIX.findall(raw) + EN_COMPANY.findall(raw.upper())
    if explicit_company_field and not found and 3 <= len(raw) <= 120:
        if any(char.isalpha() for char in raw) or re.search(r"[\u3400-\u9fff]", raw):
            found.append(raw)
    results = []
    for item in found:
        candidate = clean(item)
        candidate = re.sub(r"^(?:C\s*/?\s*O|O)\s+", "", candidate, flags=re.I)
        candidate = re.sub(
            r"^(?:發貨人名稱|发货人名称|收貨人名稱|收货人名称|申報人名稱|申报人名称|"
            r"寄件人|收件人|承運人|承运人|公司名稱|公司名称)\s*[:：]?\s*",
            "",
            candidate,
        )
        if candidate in {
            "有限公司", "有限责任公司", "有限責任公司", "股份有限公司",
            "公司", "企业", "企業", "COMPANY", "LIMITED", "CO LTD", "CO LIMITED",
        }:
            continue
        if normalized_key(candidate) in {normalized_key(value) for value in GENERIC_COMPANY_KEYS}:
            continue
        if len(candidate) >= 3:
            results.append(candidate)
    return sorted(set(results))


def vehicle_candidates(value: str) -> list[str]:
    raw = clean(value).upper()
    found = CN_CROSS_BORDER_PLATE.findall(raw) + HK_PLATE.findall(raw)
    results = []
    for item in found:
        token = re.sub(r"\s+", "", item).upper()
        if re.fullmatch(r"[A-Z]\d{5,}", token):
            continue
        if 3 <= len(token) <= 16:
            results.append(token)
    return sorted(set(results))


def group_for_header(header: str) -> str | None:
    folded = clean(header).casefold()
    for group, tokens in HEADER_GROUPS.items():
        if any(token.casefold() in folded for token in tokens):
            return group
    return None


def workbook_mentions(path: Path):
    try:
        book = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return
    for sheet in book.worksheets:
        header = None
        groups = {}
        for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
            values = [clean(value) for value in row]
            candidate_groups = {
                index: group_for_header(value)
                for index, value in enumerate(values)
                if value
            }
            candidate_groups = {index: group for index, group in candidate_groups.items() if group}
            if row_number <= 25 and len(candidate_groups) >= 1 and sum(bool(value) for value in values) >= 3:
                header = values
                groups = candidate_groups
                continue
            if header is None or not any(values):
                continue
            for index, group in groups.items():
                if index >= len(values):
                    continue
                value = values[index]
                if not value:
                    continue
                field = header[index]
                if group == "企业":
                    for candidate in company_candidates(value, explicit_company_field=True):
                        yield group, candidate, sheet.title, row_number, field, value
                elif group == "人员":
                    chinese = re.findall(r"[\u3400-\u9fff·]{2,8}", value)
                    if chinese:
                        for candidate in chinese[:3]:
                            yield group, candidate, sheet.title, row_number, field, value
                    elif 2 <= len(value) <= 80:
                        yield group, value, sheet.title, row_number, field, value
                elif group == "车辆":
                    for candidate in vehicle_candidates(value):
                        yield group, candidate, sheet.title, row_number, field, value
                elif group == "集装箱":
                    for candidate in CONTAINER.findall(value.upper()):
                        yield group, compact(candidate).upper(), sheet.title, row_number, field, value
                elif group == "电话":
                    digits = re.sub(r"\D", "", value)
                    if 7 <= len(digits) <= 18:
                        yield group, digits, sheet.title, row_number, field, value
                elif group == "地址":
                    if 8 <= len(value) <= 240:
                        yield group, value, sheet.title, row_number, field, value
                elif group == "提运单":
                    for candidate in re.findall(r"\b[A-Z0-9][A-Z0-9-]{7,30}\b", value.upper()):
                        if any(char.isdigit() for char in candidate):
                            yield group, candidate, sheet.title, row_number, field, value


def theme_tags(text: str) -> list[str]:
    return [name for name, pattern in THEMES.items() if re.search(pattern, text, re.I)]


def channel_tags(text: str) -> list[str]:
    return [name for name, pattern in CHANNELS.items() if re.search(pattern, text, re.I)]


def concealment_category(value: str) -> str:
    text = value or ""
    if re.search(r"行李|手提袋|背包|旅行袋|衣物", text):
        return "行李/随身物品夹藏"
    if re.search(r"身上|體內|体内|腰|腿|衣服內|衣服内", text):
        return "人体藏匿"
    if re.search(r"油缸|輪胎|轮胎|底盤|底盘|暗格|車身|车身|車廂|车厢|司機艙|司机舱", text):
        return "车辆结构藏匿"
    if re.search(r"包裝箱|包装箱|紙箱|纸箱|貨物|货物|木箱", text):
        return "申报货物/包装夹藏"
    if re.search(r"貨櫃|货柜|集裝箱|集装箱", text):
        return "集装箱藏匿"
    if re.search(r"倉|仓|店|住宅|單位|单位|場地|场地", text):
        return "仓库/场所储存"
    if re.search(r"船|艙|舱", text):
        return "船舶藏匿"
    return "其他/未明确"


with CATALOG.open("r", encoding="utf-8-sig", newline="") as stream:
    catalog = list(csv.DictReader(stream))

intel_rows = [row for row in catalog if INTEL.search(row["case_type"]) or INTEL.search(row["package_name"])]
intel_ids = {row["case_id"] for row in intel_rows}
title_by_case = {row["case_id"]: row["package_name"] for row in intel_rows}

conn = sqlite3.connect(DB)
conn.row_factory = sqlite3.Row
placeholders = ",".join("?" for _ in intel_ids)
document_rows = conn.execute(
    f"""SELECT case_id,relative_path,page_or_sheet,extracted_text,text_chars,extraction_error
        FROM documents WHERE case_id IN ({placeholders})""",
    tuple(sorted(intel_ids)),
).fetchall()

mentions: dict[tuple[str, str], dict] = {}
mention_rows = []


def add_mention(
    entity_type: str,
    display: str,
    case_id: str,
    relative_path: str,
    page_or_sheet: str,
    context: str,
):
    display = clean(display)
    key = normalized_key(display)
    if not key or key in {"NA", "NONE", "UNKNOWN", "不詳", "不详"}:
        return
    record = mentions.setdefault(
        (entity_type, key),
        {
            "entity_type": entity_type,
            "normalized_value": key,
            "aliases": Counter(),
            "case_ids": set(),
            "occurrences": 0,
            "evidence": [],
        },
    )
    record["aliases"][display] += 1
    record["case_ids"].add(case_id)
    record["occurrences"] += 1
    evidence = {
        "case_id": case_id,
        "file": relative_path,
        "page_or_sheet": page_or_sheet,
        "context": clean(context)[:300],
    }
    if len(record["evidence"]) < 10:
        record["evidence"].append(evidence)
    mention_rows.append((
        case_id,
        relative_path,
        page_or_sheet,
        entity_type,
        key,
        display,
        clean(context)[:500],
    ))


incidents = []
theme_packages = Counter()
channel_packages = Counter()
theme_incidents = Counter()
channel_incidents = Counter()
concealment_counts = Counter()
route_counts = Counter()
gender_counts = Counter()
nationality_counts = Counter()

for row in intel_rows:
    title = row["package_name"]
    tags = theme_tags(title) or ["其他/综合情报"]
    channels = channel_tags(title) or ["未明确/综合情报"]
    theme_packages.update(tags)
    channel_packages.update(channels)

for row in document_rows:
    case_id = row["case_id"]
    relative_path = row["relative_path"]
    page_or_sheet = row["page_or_sheet"]
    text = row["extracted_text"] or ""
    if not text:
        continue

    for match in PERSON.finditer(text):
        chinese_name = match.group(1)
        romanized = clean(match.group(2) or "")
        context = match.group(0)
        add_mention("人员", chinese_name, case_id, relative_path, page_or_sheet, context)
        if romanized:
            add_mention("人员英文名", romanized, case_id, relative_path, page_or_sheet, context)

    for match in DOCUMENT_ID.finditer(text):
        add_mention("证件号", match.group(1), case_id, relative_path, page_or_sheet, match.group(0))

    for line in text.splitlines():
        cleaned_line = clean(line)
        if not cleaned_line:
            continue
        if re.match(r"^(?:車牌|车牌)號碼", cleaned_line):
            for plate in vehicle_candidates(cleaned_line):
                add_mention("车辆", plate, case_id, relative_path, page_or_sheet, cleaned_line)
        if re.match(r"^(?:航班資料|航班资料|航班編號|航班编号)", cleaned_line):
            value = re.sub(r"^(?:航班資料|航班资料|航班編號|航班编号)\s*[:：]?\s*", "", cleaned_line)
            for token in re.findall(r"\b[A-Z]{1,3}\s*\d{2,5}\b", value.upper()):
                add_mention("航班", compact(token), case_id, relative_path, page_or_sheet, cleaned_line)
        if re.match(r"^(?:船名|船舶名稱|船舶名称|貨船名稱|货船名称)", cleaned_line):
            value = re.sub(r"^(?:船名|船舶名稱|船舶名称|貨船名稱|货船名称)\s*[:：]?\s*", "", cleaned_line)
            if 2 <= len(value) <= 80:
                add_mention("船舶", value, case_id, relative_path, page_or_sheet, cleaned_line)
        # Flattened XLSX text is also scanned because a few trade-record
        # workbooks use merged or multi-row headers that defeat column parsing.
        # company_candidates() filters generic header/suffix-only artifacts.
        for company in company_candidates(cleaned_line):
            add_mention("企业", company, case_id, relative_path, page_or_sheet, cleaned_line)

    for container in CONTAINER.findall(text.upper()):
        add_mention("集装箱", compact(container).upper(), case_id, relative_path, page_or_sheet, container)

    if re.search(r"案件(?:[一二三四五六七八九十\d]+)?(?:資料|资料)", text) and re.search(r"檢獲物品|检获物品", text):
        fields = parse_fields(text)
        for field_name in ("发货人资料", "收货人资料"):
            value = fields.get(field_name, "")
            if not value:
                continue
            parts = re.split(r"[（(][一二三四五六七八九十\d]+[）)]", value)
            for part in parts:
                for company in company_candidates(part, explicit_company_field=True):
                    add_mention(
                        "企业", company, case_id, relative_path, page_or_sheet,
                        f"{field_name}: {clean(part)}",
                    )
        title = title_by_case.get(case_id, "")
        seized = fields.get("检获物品", "")
        tags = theme_tags(f"{title} {seized}") or ["其他/综合情报"]
        channels = channel_tags(f"{title} {fields.get('运输方式', '')}") or ["未明确/综合情报"]
        theme_incidents.update(tags)
        channel_incidents.update(channels)
        concealment = fields.get("收藏手法", "")
        concealment_counts[concealment_category(concealment)] += 1
        route = fields.get("路线", "")
        if route:
            route_counts[route] += 1
        person_block = fields.get("被捕人资料", "")
        gender_match = re.search(r"性別\s*[:：]\s*(男|女)|性别\s*[:：]\s*(男|女)", person_block)
        if gender_match:
            gender_counts[gender_match.group(1) or gender_match.group(2)] += 1
        nationality_match = re.search(r"國籍\s*[:：]\s*([^\s]+)|国籍\s*[:：]\s*([^\s]+)", person_block)
        if nationality_match:
            nationality_counts[nationality_match.group(1) or nationality_match.group(2)] += 1
        incident_match = re.search(r"案件([一二三四五六七八九十\d]*)(?:資料|资料)", text)
        incidents.append({
            "case_id": case_id,
            "incident_label": incident_match.group(1) if incident_match else page_or_sheet,
            "file": relative_path,
            "page": page_or_sheet,
            "themes": tags,
            "channels": channels,
            "fields": fields,
        })

for row in intel_rows:
    case_id = row["case_id"]
    case_root = ROOT / row["raw_path"]
    for path in case_root.rglob("*.xlsx"):
        relative_path = path.relative_to(ROOT).as_posix()
        for group, display, sheet, row_number, field, context in workbook_mentions(path) or []:
            add_mention(
                group,
                display,
                case_id,
                relative_path,
                f"{sheet}!{row_number}",
                f"{field}: {context}",
            )

entities = []
for record in mentions.values():
    display, _ = record["aliases"].most_common(1)[0]
    entities.append({
        "entity_type": record["entity_type"],
        "normalized_value": record["normalized_value"],
        "display_value": display,
        "aliases": [alias for alias, _ in record["aliases"].most_common()],
        "case_count": len(record["case_ids"]),
        "case_ids": sorted(record["case_ids"]),
        "occurrences": record["occurrences"],
        "evidence": record["evidence"],
    })
entities.sort(key=lambda item: (-item["case_count"], -item["occurrences"], item["entity_type"], item["display_value"]))

repeated = [
    item for item in entities
    if item["case_count"] >= 2
]
frequent_within_case = [
    item for item in entities
    if item["case_count"] == 1 and item["occurrences"] >= 3
]

conn.executescript("""
DROP TABLE IF EXISTS intel_incidents;
DROP TABLE IF EXISTS intel_mentions;
DROP TABLE IF EXISTS intel_entities;
CREATE TABLE intel_incidents (
    id INTEGER PRIMARY KEY,
    case_id TEXT NOT NULL,
    incident_label TEXT,
    relative_path TEXT NOT NULL,
    page_or_sheet TEXT,
    themes_json TEXT,
    channels_json TEXT,
    fields_json TEXT
);
CREATE TABLE intel_mentions (
    id INTEGER PRIMARY KEY,
    case_id TEXT NOT NULL,
    relative_path TEXT NOT NULL,
    page_or_sheet TEXT,
    entity_type TEXT NOT NULL,
    normalized_value TEXT NOT NULL,
    display_value TEXT NOT NULL,
    context TEXT
);
CREATE TABLE intel_entities (
    entity_type TEXT NOT NULL,
    normalized_value TEXT NOT NULL,
    display_value TEXT NOT NULL,
    case_count INTEGER NOT NULL,
    occurrence_count INTEGER NOT NULL,
    case_ids_json TEXT NOT NULL,
    aliases_json TEXT NOT NULL,
    PRIMARY KEY(entity_type, normalized_value)
);
CREATE INDEX idx_intel_mentions_type_value ON intel_mentions(entity_type, normalized_value);
CREATE INDEX idx_intel_mentions_case ON intel_mentions(case_id);
""")
for incident in incidents:
    conn.execute(
        """INSERT INTO intel_incidents(
               case_id,incident_label,relative_path,page_or_sheet,
               themes_json,channels_json,fields_json
           ) VALUES(?,?,?,?,?,?,?)""",
        (
            incident["case_id"],
            incident["incident_label"],
            incident["file"],
            incident["page"],
            json.dumps(incident["themes"], ensure_ascii=False),
            json.dumps(incident["channels"], ensure_ascii=False),
            json.dumps(incident["fields"], ensure_ascii=False),
        ),
    )
conn.executemany(
    """INSERT INTO intel_mentions(
           case_id,relative_path,page_or_sheet,entity_type,
           normalized_value,display_value,context
       ) VALUES(?,?,?,?,?,?,?)""",
    mention_rows,
)
for item in entities:
    conn.execute(
        "INSERT INTO intel_entities VALUES(?,?,?,?,?,?,?)",
        (
            item["entity_type"],
            item["normalized_value"],
            item["display_value"],
            item["case_count"],
            item["occurrences"],
            json.dumps(item["case_ids"], ensure_ascii=False),
            json.dumps(item["aliases"], ensure_ascii=False),
        ),
    )
conn.commit()

document_stats = conn.execute(
    f"""SELECT extension,count(DISTINCT relative_path),count(*),
               sum(CASE WHEN text_chars>0 THEN 1 ELSE 0 END),sum(text_chars)
        FROM documents WHERE case_id IN ({placeholders})
        GROUP BY extension ORDER BY extension""",
    tuple(sorted(intel_ids)),
).fetchall()
remaining_encrypted_errors = conn.execute(
    """SELECT count(*) FROM documents
       WHERE extraction_error LIKE '%FileNotDecryptedError%'"""
).fetchone()[0]
conn.close()

ingestion = json.loads((ROOT / "加密PDF入库结果.json").read_text(encoding="utf-8"))
result = {
    "scope": {
        "intelligence_case_packages": len(intel_rows),
        "ingested_encrypted_pdf_files": ingestion["stats"]["files_ingested"],
        "ingested_encrypted_pdf_pages": ingestion["stats"]["pages"],
        "ingested_encrypted_pdf_characters": ingestion["stats"]["characters"],
        "remaining_encrypted_errors": remaining_encrypted_errors,
        "parsed_incident_pages": len(incidents),
        "document_stats": [list(row) for row in document_stats],
    },
    "theme_packages": theme_packages.most_common(),
    "theme_parsed_incidents": theme_incidents.most_common(),
    "channel_packages": channel_packages.most_common(),
    "channel_parsed_incidents": channel_incidents.most_common(),
    "concealment_categories": concealment_counts.most_common(),
    "routes": route_counts.most_common(30),
    "gender_counts": gender_counts.most_common(),
    "nationality_counts": nationality_counts.most_common(),
    "entity_counts": Counter(item["entity_type"] for item in entities).most_common(),
    "repeated_entities": repeated,
    "frequent_within_case_entities": frequent_within_case,
    "incidents": incidents,
}
OUTPUT.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")

print(json.dumps({
    "scope": result["scope"],
    "theme_packages": result["theme_packages"],
    "theme_parsed_incidents": result["theme_parsed_incidents"],
    "channel_parsed_incidents": result["channel_parsed_incidents"],
    "concealment_categories": result["concealment_categories"],
    "entity_counts": result["entity_counts"],
    "repeated_top": repeated[:80],
}, ensure_ascii=False, indent=2))
