from __future__ import annotations

import csv
import json
import os
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from build_epdm_stage_report import (
    BLUE, GOLD, GREEN, INK, MUTED, NAVY, PALE_BLUE, PALE_GOLD, PALE_GRAY,
    PALE_GREEN, PALE_RED, RED, add_callout, add_heading, add_hyperlink,
    add_list_item, add_page_number_field, add_paragraph, add_source,
    add_table, configure_document, create_numbering, set_run_font,
)


OUT = Path(r"D:\易迅数据\反倾销税深度分析报告\19_丁腈橡胶")
DOCX = OUT / "丁腈橡胶_反倾销税与第三国转运风险阶段深度分析报告.docx"
LEDGER = OUT / "丁腈橡胶_易迅逐票判定台账_现有10条.csv"
SUMMARY = OUT / "丁腈橡胶_交付摘要.json"


def table_caption(table, text: str) -> None:
    props = table._tbl.tblPr
    caption = OxmlElement("w:tblCaption")
    caption.set(qn("w:val"), text)
    props.append(caption)
    if table.rows:
        for cell in table.rows[0].cells:
            tr_pr = table.rows[0]._tr.get_or_add_trPr()
            if tr_pr.find(qn("w:tblHeader")) is None:
                flag = OxmlElement("w:tblHeader")
                flag.set(qn("w:val"), "true")
                tr_pr.append(flag)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True


def report_table(doc, headers, rows, widths, **kwargs):
    table = add_table(doc, headers, rows, widths, **kwargs)
    table_caption(table, "丁腈橡胶反倾销税与第三国转运风险阶段审计表")
    return table


def header_footer(doc: Document) -> None:
    section = doc.sections[0]
    header = section.header.paragraphs[0]
    header.clear()
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.space_after = Pt(0)
    set_run_font(header.add_run("反倾销税风险穿透分析｜丁腈橡胶（NBR）｜2026-08-13"), size=8.5, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.clear()
    footer.paragraph_format.space_after = Pt(0)
    add_page_number_field(footer)
    if footer.runs:
        footer.runs[0].text = "丁腈橡胶阶段深度分析报告  |  "
    props = doc.core_properties
    props.title = "丁腈橡胶反倾销税与第三国转运风险阶段深度分析报告"
    props.subject = "易迅逐票审计、政策税率、第三国链路、实体与调证清单"
    props.author = "反倾销税风险分析项目"
    props.keywords = "丁腈橡胶; NBR; KNB-35L; 反倾销; 第三国转运; 易迅数据"


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    rows = list(csv.DictReader(LEDGER.open(encoding="utf-8-sig", newline="")))
    summary = json.loads(SUMMARY.read_text(encoding="utf-8"))
    assert len(rows) == 10

    doc = Document()
    configure_document(doc)
    header_footer(doc)
    bullets = create_numbering(doc, bullet=True)
    decimals = create_numbering(doc, bullet=False)

    # memo_masthead + standard_business_brief
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(4)
    set_run_font(p.add_run("第19项商品核查"), size=11, color=GOLD, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    set_run_font(p.add_run("丁腈橡胶（NBR）"), size=27, color=NAVY, bold=True)
    add_paragraph(doc, "反倾销税与第三国转运风险｜阶段深度分析报告", size=15.5, color=BLUE, bold=True, after=14)
    report_table(
        doc,
        ["基准日", "措施状态", "易迅现有逐票", "第三国闭环"],
        [["2026-08-13", "实施中；续征至2029-11-08", "10条可审计贸易记录", "0条"]],
        [1700, 2700, 2400, 2560],
        header_fill=PALE_BLUE,
        font_size=9.6,
        aligns=[WD_ALIGN_PARAGRAPH.CENTER] * 4,
    )
    add_callout(
        doc,
        "核心结论",
        "现有易迅原始页面交叉命中可以证明韩国KUMHO KNB-35L向印度、印尼发运的9条A腿，"
        "其中1条由青岛前湾保税港区贸易主体向印度供货；但没有发现印度或印尼主体再向中国发运同牌号、同批次或同柜货物的B腿。"
        "因此，当前结论是“存在可核第三国供应链与保税贸易线索，但绕道中国和逃税均未证实”。",
        fill=PALE_GOLD,
        title_color=GOLD,
    )
    report_table(
        doc,
        ["维度", "当前证据", "等级", "不能越过的边界"],
        [
            ["产品/牌号", "9条KUMHO KNB-35L；韩国平台原产字段", "B", "平台字段与外国HS不等于中国法定原产或归类"],
            ["A腿", "韩国来源→印尼6条、→印度3条", "B", "A腿存在不等于货物随后进入中国"],
            ["实体链", "锦湖牌号—贸易/物流主体—印尼鞋厂或印度买方", "B-/C+", "需提单、批号、柜号和生产用途闭环"],
            ["B腿/税款", "未发现第三国→中国同货线；无中国报关单和税单", "未形成", "不能认定伪报、走私、逃税或测算实际税额"],
        ],
        [1500, 3200, 1200, 3460],
        header_fill=PALE_GRAY,
        font_size=8.8,
    )
    add_paragraph(doc, "报告性质：阶段深度审计。易迅NBR专项网页全库查询因长查询页面超时尚未补齐；本报告不把D盘现有10条冒充全库结果。", size=9.2, color=RED, bold=True, after=0)

    doc.add_page_break()
    add_heading(doc, "1. 政策、产品范围与税率", 1)
    add_heading(doc, "1.1 现行措施", 2)
    add_paragraph(doc, "商务部2024年第48号公告决定，自2024年11月9日起继续对原产于韩国和日本的进口丁腈橡胶征收反倾销税，实施期限5年，预计至2029年11月8日。涉案产品为丙烯腈和丁二烯的共聚物，通常呈灰白至淡黄色块状或粉末状固体；中国税则号为40025910、40025990。")
    report_table(
        doc,
        ["受税来源", "生产商/出口商", "AD税率", "AD+由AD增加的13%VAT系数"],
        [
            ["韩国", "锦湖石油化学株式会社", "12.0%", "完税价×13.560%"],
            ["韩国", "LG化学株式会社", "15.0%", "完税价×16.950%"],
            ["韩国", "其他韩国公司", "37.3%", "完税价×42.149%"],
            ["日本", "日本瑞翁株式会社", "28.1%", "完税价×31.753%"],
            ["日本", "ENEOS Materials Corporation", "16.0%", "完税价×18.080%"],
            ["日本", "其他日本公司", "56.4%", "完税价×63.732%"],
        ],
        [1300, 3500, 1400, 3160],
        header_fill=PALE_BLUE,
        font_size=8.8,
    )
    add_callout(doc, "主体名称风险", "ENEOS Materials继承JSR株式会社16.0%的税率；若继续以JSR名义报关，则适用“其他日本公司”56.4%。中国进口底单必须同时核生产商、出口商、发票抬头及适用税率。", fill=PALE_RED, title_color=RED)
    add_heading(doc, "1.2 产品范围边界", 2)
    for text in [
        "中国措施税号是40025910、40025990。现有9条KNB-35L在印度、印尼数据中使用40023900/10/90，形成“货描—外国税号错位”。这可能是境外归类、平台映射或申报问题，不能直接推定中国逃税。",
        "NBR乳胶通常归入400251；HNBR、XNBR以及手套、胶管、密封件、鞋材等制品不能仅凭出现NBR字样并入本措施。必须核CAS、聚合物组成、形态、用途和中国进口税号。",
        "KNB-35L公开产品资料指向丙烯腈—丁二烯橡胶、韩国蔚山生产及35千克/包包装，支持其为NBR牌号；但正式认定仍应以该票COA、TDS、生产批号和中国归类为准。",
    ]:
        add_list_item(doc, text, bullets)

    add_heading(doc, "2. 数据覆盖与逐票方法", 1)
    add_paragraph(doc, "本项先对D盘可解析的XLSX、CSV、JSON做内容级预筛，再对45个候选文件逐行穿透。预筛覆盖313个文件；深扫原始命中、派生报告与元数据后，最终回溯至易迅原始页面采集并确认10条贸易记录。主清单中的商品名称记录不计贸易票；400259税号下6条实际为溴化丁基橡胶的记录被剔除。")
    report_table(
        doc,
        ["环节", "结果", "质量控制"],
        [
            ["全盘预筛", "313文件；45候选", "文件级税号、名称、品牌宽抓"],
            ["逐行穿透", "45文件；0解析错误", "区分原始页面、派生输出和项目元数据"],
            ["贸易记录", "10条", "回溯原始易迅页面位置；逐票人工判断"],
            ["产品构成", "KNB-35L 9条；混合BUTYL/NBR 1条", "混合货描不强行拆分数量"],
            ["中国B腿", "0条", "无同牌号/同批次/同柜第三国→中国记录"],
        ],
        [1800, 2500, 5060],
        header_fill=PALE_GRAY,
        font_size=9,
    )
    add_callout(doc, "完整性声明", "专项网页条件“HS 40025910/40025990 + NBR关键词 + 目的国中国”的全页查询尚未成功完成。现有10条来自既有22页HS400239宽池中的交叉命中，仅代表当前D盘可回溯记录；不能据此声称易迅全库没有韩国/日本直达中国或第三国B腿。", fill=PALE_RED, title_color=RED)

    doc.add_page_break()
    add_heading(doc, "3. 现有10条贸易记录全量分析", 1)
    add_heading(doc, "3.1 韩国KNB-35L→印度尼西亚：6条", 2)
    id_rows = [r for r in rows if r["目的国地区"] == "Indonesia"]
    report_table(
        doc,
        ["日期", "买方", "境外供货/物流方", "重量字段", "金额字段"],
        [[r["日期"], "PARKLAND WORLD INDONESIA", r["供应商"].replace("EUNSAN SHIPPINGANDAIRCARGO CO LTD O B OF SINHWA CO LTD", "EUNSAN O/B/O SINHWA"), r["重量字段"], r["金额字段"]] for r in id_rows],
        [1400, 2500, 2650, 1300, 1510],
        header_fill=PALE_BLUE,
        font_size=8.3,
    )
    add_paragraph(doc, "六条记录的重量字段均为16,800，合计100,800；平台未在导出文件中显式给出单位，但字段结构高度类似千克。若按公开35千克/包包装推算，每票对应480包，是很强的包装指纹，但这一推算必须由发票/装箱单确认。金额字段合计225,120，币种同样未显式载明。", after=5)
    add_paragraph(doc, "风险判断：韩国牌号和印尼收货人稳定，A腿可信度较高；但PARKLAND公开信息显示其为鞋类制造企业，真实消耗NBR用于鞋材是合理替代解释。现有D盘未发现PARKLAND、SINHWA或EUNSAN向中国发运KNB-35L，故绕道风险仅列C+/待核。")

    add_heading(doc, "3.2 韩国KNB-35L→印度：3条", 2)
    in_rows = [r for r in rows if r["目的国地区"] == "India" and "KNB" in r["商品描述"].upper()]
    report_table(
        doc,
        ["日期", "买方", "供货方", "数量字段", "金额字段"],
        [[r["日期"], r["采购商"], "青岛中联融创（前湾保税区）" if "QINGDAO" in r["供应商"] else "PIONEER COMMODITY FZCO", r["数量字段"], r["金额字段"]] for r in in_rows],
        [1400, 2100, 2800, 1400, 1660],
        header_fill=PALE_BLUE,
        font_size=8.6,
    )
    add_paragraph(doc, "三条数量字段合计126，金额字段合计241,047.24。50.4、50.4和25.2若单位为吨，分别可整除为35千克包装的1,440、1,440和720包，且金额/数量比在大宗合成橡胶价格区间内；但这只是结构性推断，不能替代平台单位和商业发票。")
    add_callout(doc, "最具体调证线索", "2025-08-26一票由青岛中联融创新材料科技有限公司（地址指向青岛前湾保税港区）向印度S.M. Associates供货，平台原产字段为韩国，数量字段25.2、金额字段49,779.24。该票证明中国保税贸易主体参与韩国牌号转出印度；方向是中国→印度，不是印度→中国。应调中国保税进口、仓储及复出口底单，确认是否原状转口、批号和生产商；若未来发现同批回流中国，才可能升级。", fill=PALE_GOLD, title_color=GOLD)
    add_paragraph(doc, "印度本身拥有Apcotex Valia NBR/NVC/粉末橡胶约21,000吨/年干胶产能。因此，未来即使出现印度→中国NBR，也必须先穿透生产厂和批号，不能仅凭启运国印度认定韩国货洗产地。")

    add_heading(doc, "3.3 比利时→印度混合货描：1条", 2)
    mixed = [r for r in rows if r["平台原产国地区"] == "Belgium"][0]
    report_table(doc, ["日期", "买方", "卖方", "货描", "数量/金额字段"], [[mixed["日期"], mixed["采购商"], "RAVAGO DISTRIBUTION CENTER NV", "BUTYL LUMPS + NBR LUMPS", f'{mixed["数量字段"]} / {mixed["金额字段"]}']], [1200, 1800, 2500, 2300, 1560], header_fill=PALE_GRAY, font_size=8.5)
    add_paragraph(doc, "该票无法从可见字段确认NBR与丁基橡胶的各自数量、具体生产商、CAS或牌号；比利时也不是本NBR措施的受税来源。只能作为货描拆分和供应链核单线索，不计入韩国/日本绕道数量，更不能计算中国反倾销税差。")

    doc.add_page_break()
    add_heading(doc, "4. 第三国绕道证据评估", 1)
    report_table(
        doc,
        ["证据问题", "已取得", "仍缺失", "结论"],
        [
            ["是否存在受税来源A腿", "是：韩国KNB-35L 9条", "原始韩国出口报关、CO、厂家批号", "A腿成立，B级"],
            ["是否存在第三国B腿", "否：现有D盘未见印度/印尼→中国同牌号", "专项易迅全页结果、中国进口明细", "未形成"],
            ["是否同一货物", "无", "提单/柜号/封志、批号、包装唛头、数量平衡", "未形成"],
            ["是否在第三国实质加工", "印尼鞋厂消费和印度本地产能均为合理替代解释", "BOM、领料、能耗、产销存、工单", "不能否定合法加工/消费"],
            ["是否少缴中国税款", "无中国报关单、完税价、原产申报或税单", "进口报关单、CO、AD税单、VAT税单", "不能测算实际逃税额"],
        ],
        [1600, 2450, 3150, 2160],
        header_fill=PALE_GRAY,
        font_size=8.6,
    )
    add_heading(doc, "4.1 当前可以确认的风险", 2)
    for item in [
        "查询隐蔽风险：KNB-35L货描明确指向NBR，但境外数据税号为400239，而中国措施税号为40025910/90。只用中国措施税号或只用境外HS会漏票，必须税号+牌号双轨查询。",
        "贸易主体风险：Pioneer Commodity FZCO、SINHWA、EUNSAN以及青岛保税贸易主体均可能位于生产商与最终用户之间。贸易/物流主体可以改变发票或发货链，但不改变货物法定原产地。",
        "税率识别风险：KNB-35L若确由锦湖生产，适用12.0%；若生产商身份不明并被认定为其他韩国公司，税率可达37.3%。企业名称、牌号和生产厂必须穿透。",
    ]:
        add_list_item(doc, item, bullets)
    add_heading(doc, "4.2 公开来源没有形成的证据", 2)
    add_paragraph(doc, "截至基准日，定向检索未发现中国海关、法院或商务部公开材料披露本措施下“韩国/日本NBR经某第三国进入中国逃避反倾销税”的具体处罚、判决、反规避裁定或双段提单闭环。商务部期终复审也未认定第三国绕道；裁定中“没有证据表明其他国家/地区进口产品存在倾销”不能被改写为已发生规避。负面检索结果只说明公开证据未见，不证明现实中绝对不存在。")

    add_heading(doc, "5. 真实第三国产能与误判反证", 1)
    report_table(
        doc,
        ["国家/地区", "公开生产能力或制造证据", "对风险判断的影响"],
        [
            ["印度", "Apcotex Valia生产NBR/NVC/粉末橡胶约21,000吨/年（干）", "印度→中国可能是真实印度生产；须查生产厂/批号"],
            ["台湾", "Nantex高雄及中国镇江生产NANCAR NBR；公开制造流程含聚合、凝聚、干燥、打包", "NANCAR品牌或台湾发货不应自动归为日韩原产"],
            ["俄罗斯", "SIBUR公开资料确认Krasnoyarsk NBR并供应中国及东南亚", "俄罗斯来源具有真实产能反证"],
            ["法国/欧盟", "ARLANXEO法国生产基地公开列有NBR", "比利时/欧盟贸易链可能由真实欧洲生产支撑"],
        ],
        [1400, 4200, 3760],
        header_fill=PALE_GREEN,
        font_size=8.7,
    )
    add_callout(doc, "判断规则", "品牌国、集团国籍、发票国、中转港和平台原产字段都不能单独决定反倾销原产地。只有生产厂、聚合工序、批号/plant code、原产证明和中国海关认定能够形成法定结论。", fill=PALE_BLUE, title_color=INK)

    doc.add_page_break()
    add_heading(doc, "6. 税差测算口径", 1)
    add_paragraph(doc, "对本措施，反倾销税=海关审定完税价格×企业适用税率；进口增值税计税基础包含反倾销税。若仅筛查“未缴AD及其引起的13%进口VAT增量”，综合系数为AD税率×1.13。")
    report_table(
        doc,
        ["情景", "条件", "综合税差系数", "可否用于现有9条A腿"],
        [
            ["锦湖KNB-35L", "中国进口确认韩国锦湖原产且未缴AD", "中国完税价×13.560%", "否；尚无中国B腿/完税价"],
            ["韩国其他", "生产商不明或适用其他韩国公司37.3%", "中国完税价×42.149%", "否；只能做风险上限情景"],
            ["ENEOS", "中国进口确认ENEOS原产并正确以ENEOS主体申报", "中国完税价×18.080%", "本地未见相关票"],
            ["日本其他", "日本生产商不明或旧JSR名义触发56.4%", "中国完税价×63.732%", "本地未见相关票"],
        ],
        [1700, 3750, 1900, 2010],
        header_fill=PALE_BLUE,
        font_size=8.7,
    )
    add_callout(doc, "为什么不写实际逃税额", "易迅境外数据中的金额字段没有明确币种、贸易术语或中国完税价格；这些票的目的地是印度、印尼，不是中国。把境外金额直接乘中国反倾销税率会产生伪精确结论。本报告仅给税差系数，待取得中国进口报关完税价后再计算。", fill=PALE_RED, title_color=RED)

    add_heading(doc, "7. 优先实体与调证清单", 1)
    report_table(
        doc,
        ["优先级", "实体/链路", "最小调证包", "升级条件"],
        [
            ["1", "青岛中联融创→S.M. Associates（印度）", "保税进口/仓储/复出口报关；发票；原产证；COA；批号；提单柜号", "发现同批或同柜由印度回流中国"],
            ["2", "Pioneer Commodity FZCO→S.M. Associates", "阿联酋/印度合同、付款、发票；韩国生产商；B/L；包装唛头", "与中国进口同牌号、数量、时间和批号匹配"],
            ["3", "SINHWA/EUNSAN→Parkland Indonesia", "韩国出口申报；承运提单；Parkland收货、领料、BOM和产成品记录", "货物未被生产消耗并短期转销中国"],
            ["4", "Ravago Belgium→Balaji India", "发票行项目拆分；NBR/丁基各自数量；CAS、牌号和生产厂", "确认其中NBR实际为日韩原产并对华转销"],
        ],
        [1000, 2750, 3600, 2010],
        header_fill=PALE_GRAY,
        font_size=8.4,
    )
    add_heading(doc, "7.1 中国端闭环必需字段", 2)
    for item in [
        "进口报关单：40025910/40025990、原产国、启运国、生产商、品牌/牌号、境外发货人、境内收货人/消费使用单位、申报企业、口岸/隶属关、净重、完税价格、贸易方式。",
        "税款单证：企业适用AD税率、反倾销税缴款书、进口VAT缴款书；ENEOS/JSR主体名称必须核对。",
        "货运与原产：B/L、集装箱号、封志、船名航次、中转港、CO、厂家声明、COA、批号、plant code、包装唛头照片。",
        "第三国加工：原料进口、BOM、工单、设备、能耗、人工、产量、损耗、库存和销售质量平衡；仅换标、分装、仓储或转卖不能证明实质生产。",
    ]:
        add_list_item(doc, item, bullets)

    add_heading(doc, "8. 易迅补查矩阵与完成标准", 1)
    report_table(
        doc,
        ["查询", "最简条件", "完成标准"],
        [
            ["Q1 中国进口", "HS40025910/40025990；China；2018-11-09至最新；全部来源；200条/页", "所有页逐条读取，记录总数/页数/最新日期"],
            ["Q2 关键词补漏", "NBR、NITRILE BUTADIENE、9003-18-3、KNB、NIPOL、ZEON、ENEOS、JSR", "每个关键词独立查询并跨查询去重"],
            ["Q3 双腿匹配", "韩国/日本→印/印尼/越/马/新/阿联酋，再查这些国家→China；30/60/90日", "同牌号+生产商+数量+批号/柜号；逐条写替代解释"],
        ],
        [1600, 5050, 2710],
        header_fill=PALE_BLUE,
        font_size=8.6,
    )
    add_paragraph(doc, "完成标准不是“发现一条异常即停止”，而是查询结果每一页、每一条均完成范围判定、去重、路线归类、实体穿透、税号检查和证据缺口记录。只有全部页面核对后，才能把“尚未发现B腿”升级为该查询窗口内的完整结论。", bold=True, color=INK)

    doc.add_page_break()
    add_heading(doc, "9. 逐票附录（10条）", 1)
    add_heading(doc, "9.1 逐票可见字段与判定", 2)
    for r in sorted(rows, key=lambda x: x["日期"], reverse=True):
        add_heading(doc, f'{r["记录ID"]}｜{r["日期"]}｜{r["平台原产国地区"]}→{r["目的国地区"]}', 3)
        add_paragraph(doc, f'货描：{r["商品描述"]}', size=9.2, after=2)
        add_paragraph(doc, f'买方：{r["采购商"]}；卖方/物流方：{r["供应商"]}', size=9.2, after=2)
        add_paragraph(doc, f'外国HS：{r["HS编码_外国数据"]}；重量字段：{r["重量字段"] or "空"}；数量字段：{r["数量字段"] or "空"}；金额字段：{r["金额字段"] or "空"}', size=9.2, after=2)
        add_paragraph(doc, f'判定：{r["证据等级"]}。{r["逐票结论"]}', size=9.2, color=INK, after=7)

    doc.add_page_break()
    add_heading(doc, "10. 公开来源", 1)
    add_source(doc, bullets, "S1", "商务部公告2024年第48号：丁腈橡胶期终复审裁定公告", "https://policy.mofcom.gov.cn/claw/clawContent.shtml?id=101445", "现行措施、税号、企业税率、期限和计税公式。")
    add_source(doc, bullets, "S2", "商务部2024年第48号期终复审裁定附件", "https://picpolicy.mofcom.gov.cn/file/20250109/13121736385538059.pdf", "市场、产能、进口量、倾销继续或再度发生的可能性；本地已留存PDF。")
    add_source(doc, bullets, "S3", "商务部2022年第18号：ENEOS继承JSR税率", "https://interview.mofcom.gov.cn/mofcom_interview/front/opdata/downlodePdf?id=20220603323067", "ENEOS 16.0%及继续以JSR名义申报的税率风险。")
    add_source(doc, bullets, "S4", "Apcotex Business Overview / Corporate History", "https://apcotex.com/our-business", "印度Valia NBR/NVC/粉末橡胶21,000吨/年干胶真实产能反证。")
    add_source(doc, bullets, "S5", "Nantex NANCAR NBR产品与制造信息", "https://www.nantex.com.tw/index.php?id=71&index=2&lang=en&option=product&task=showlist", "台湾/中国真实NBR制造、产品范围及包装。")
    add_source(doc, bullets, "S6", "SIBUR：Krasnoyarsk NBR面向中国与东南亚", "https://www.sibur.ru/en/press-center/articles-interviews/siburprimarilyfocusesonthedomesticmarketwithsomeexportstoasiaandeuropeolegmakarovmanagingdirector/", "俄罗斯真实NBR供应能力反证。")
    add_source(doc, bullets, "S7", "KUMHO KNB-35L产品资料", "https://www.hbchemical.com/wp-content/uploads/2022/09/KUMHO-KNB-35L.pdf", "KNB-35L化学品身份、韩国制造主体和产品用途的辅助核验。")
    add_paragraph(doc, "证据说明：公开来源用于核政策、产品、产能和主体网络；易迅页面记录用于识别可调证贸易线索。任何违法、原产地伪报或税款结论均须回到中国海关报关单、法定原产地认定和税款缴款书。", size=9.2, color=MUTED, after=0)

    doc.save(DOCX)
    assert DOCX.exists() and DOCX.stat().st_size > 50_000


if __name__ == "__main__":
    main()
