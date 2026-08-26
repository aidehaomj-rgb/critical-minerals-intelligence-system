from __future__ import annotations

import csv
import json
import os
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT_DIR = Path(os.environ.get(
    "PHT_REPORT_DIR",
    r"D:\易迅数据\反倾销税深度分析报告\08_酞菁类颜料",
))
AUDIT_PATH = OUT_DIR / "酞菁类颜料_全量阶段审计.json"
LEADS_PATH = OUT_DIR / "酞菁类颜料_越南对华31票重点线索.csv"
REPORT_PATH = OUT_DIR / "酞菁类颜料_反倾销税与第三国转运风险阶段审计报告.docx"

AUDIT = json.loads(AUDIT_PATH.read_text(encoding="utf-8"))
with LEADS_PATH.open("r", encoding="utf-8-sig", newline="") as fh:
    LEADS = list(csv.DictReader(fh))

NAVY = "17324D"
BLUE = "1F4E78"
INK = "102A43"
MUTED = "66788A"
LIGHT_BLUE = "EAF2F8"
PALE_GOLD = "FFF4CC"
PALE_RED = "FCE8E6"
PALE_GREEN = "E6F4EA"
BORDER = "CBD5E1"
RED = "9B1C1C"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def fmt(value, digits=2):
    if value in (None, ""):
        return "—"
    number = float(value)
    if digits == 0:
        return f"{number:,.0f}"
    return f"{number:,.{digits}f}"


def set_run_font(run, size=None, color=None, bold=None, italic=None, name="Calibri"):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    node = trpr.find(qn("w:tblHeader"))
    if node is None:
        node = OxmlElement("w:tblHeader")
        trpr.append(node)
    node.set(qn("w:val"), "true")


def set_cell_text(cell, value, bold=False, color="000000", size=8.7):
    cell.text = "" if value is None else str(value)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.03
        for run in paragraph.runs:
            set_run_font(run, size=size, color=color, bold=bold)


def set_table_borders(table, color=BORDER, size="4"):
    tblpr = table._tbl.tblPr
    borders = tblpr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblpr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa, total=9360):
    table.autofit = False
    tbl = table._tbl
    tblpr = tbl.tblPr
    for tag in ("w:tblW", "w:tblInd", "w:tblLayout", "w:tblCellMar"):
        old = tblpr.find(qn(tag))
        if old is not None:
            tblpr.remove(old)
    tblw = OxmlElement("w:tblW")
    tblw.set(qn("w:w"), str(total))
    tblw.set(qn("w:type"), "dxa")
    tblpr.append(tblw)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr.append(layout)
    margins = OxmlElement("w:tblCellMar")
    for tag, val in (("top", 55), ("bottom", 55), ("start", 75), ("end", 75)):
        node = OxmlElement(f"w:{tag}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    tblpr.append(margins)
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths_dxa):
            tcpr = cell._tc.get_or_add_tcPr()
            tcw = tcpr.find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                tcpr.append(tcw)
            tcw.set(qn("w:w"), str(width))
            tcw.set(qn("w:type"), "dxa")


def add_table(doc, headers, rows, widths, header_fill=LIGHT_BLUE, size=8.7, shades=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for idx, text in enumerate(headers):
        set_cell_text(table.cell(0, idx), text, bold=True, color=INK, size=size)
        set_cell_shading(table.cell(0, idx), header_fill)
    set_repeat_header(table.rows[0])
    for ridx, vals in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(vals):
            set_cell_text(cells[idx], value, size=size)
            if shades and ridx in shades:
                set_cell_shading(cells[idx], shades[ridx])
    set_table_geometry(table, widths)
    set_table_borders(table)
    return table


def add_callout(doc, title, text, fill=PALE_GOLD, title_color=INK):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, size=10.2, color=title_color, bold=True)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.1
    r2 = p2.add_run(text)
    set_run_font(r2, size=9.7, color="333333")
    set_table_geometry(table, [9360])
    set_table_borders(table, color=BORDER, size="6")
    set_repeat_header(table.rows[0])
    return table


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.07
    r = p.add_run(text)
    set_run_font(r, size=9.9, color="222222")
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.07
    r = p.add_run(text)
    set_run_font(r, size=9.9, color="222222")
    return p


def add_para(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(text)
    set_run_font(r, size=10.0, color="222222")
    return p


def add_hyperlink(paragraph, text, url):
    rid = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    for node_name, attrs in (
        ("w:rFonts", {"w:ascii": "Calibri", "w:hAnsi": "Calibri", "w:eastAsia": "Microsoft YaHei"}),
        ("w:color", {"w:val": BLUE}),
        ("w:u", {"w:val": "single"}),
    ):
        node = OxmlElement(node_name)
        for key, val in attrs.items():
            node.set(qn(key), val)
        rpr.append(node)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_source(doc, label, url, note=""):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    add_hyperlink(p, label, url)
    if note:
        r = p.add_run(f"：{note}")
        set_run_font(r, size=9.0, color=MUTED)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_1)
    run._r.append(instr)
    run._r.append(fld_char_2)
    set_run_font(run, size=8.2, color=MUTED)


def page_break(doc):
    doc.add_page_break()


doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.72)
sec.bottom_margin = Inches(0.72)
sec.left_margin = Inches(0.82)
sec.right_margin = Inches(0.82)
sec.page_width = Inches(8.27)
sec.page_height = Inches(11.69)

styles = doc.styles
styles["Normal"].font.name = "Calibri"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
styles["Normal"].font.size = Pt(10.0)
for name, size, color in (
    ("Title", 26, NAVY),
    ("Heading 1", 16.2, NAVY),
    ("Heading 2", 13.0, BLUE),
    ("Heading 3", 11.2, INK),
):
    style = styles[name]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(size)
    style.font.color.rgb = rgb(color)
    style.font.bold = True

header = sec.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hr = header.add_run("反倾销税风险穿透分析｜酞菁类颜料｜2026-08-13")
set_run_font(hr, size=8.2, color=MUTED)
footer = sec.footer.paragraphs[0]
add_page_number(footer)

# 封面
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(38)
r = p.add_run("反倾销税风险穿透分析")
set_run_font(r, size=12, color=BLUE, bold=True)
p = doc.add_paragraph(style="Title")
p.paragraph_format.space_before = Pt(14)
p.paragraph_format.space_after = Pt(7)
r = p.add_run("酞菁类颜料")
set_run_font(r, size=26, color=NAVY, bold=True)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(18)
r = p.add_run("印度来源经越南供应链、非优惠原产地及中国反倾销税风险阶段审计报告")
set_run_font(r, size=15.5, color=INK, bold=True)

add_table(doc, ["项目", "内容"], [
    ["受税来源", "印度"],
    ["实施期间", "2023年2月27日起，期限5年，预计至2028年2月26日"],
    ["税号入口", "32041700、32129000；最终仍按产品描述及成分判定"],
    ["易迅已审计查询", "PHTHALOCYANINE→中国129条；印度→越南215条（近一年）"],
    ["逐条审计覆盖", "344次原始出现；336个可见字段唯一签名；全量进入台账"],
    ["阶段状态", "尚缺两项税号及扩展关键词的全页补查，故本文件不是最终结案报告"],
    ["报告日期", "2026年8月13日"],
], [1700, 7660], size=9.2)

add_callout(
    doc,
    "核心结论",
    "现有数据证实越南向中国发运31条货描尾码为#&VN的酞菁颜料记录，数量字段合计903,600，涉及Vinh Gia、Yicai及两家中国接收方；同期也存在印度向越南供应同类酞菁的211条范围候选记录。但尚无同一批货的A腿—越南加工/库存—B腿箱号、批号或申报单闭环，也没有中国进口原产国和税款缴纳证据。因此可以认定“越南对华B腿及印度对越供给背景成立”，不能认定“印度原产货已绕道逃税”。",
)
add_callout(
    doc,
    "计量与币种警示",
    "本地JSON没有单列数量单位和金额币种。报告统一使用“数量字段”“金额字段”。越南货描中的25kg/包可帮助理解数量口径，但正式数量、净重、币种和完税价格必须回查原始申报；条件税差只能作为同币种风险情景，绝不能直接写成人民币欠税或实际追税额。",
    fill=PALE_RED,
    title_color=RED,
)

page_break(doc)
doc.add_heading("一、核心判断", level=1)
add_bullet(doc, "已证实B腿：31条越南→中国记录均保留#&VN标记，数量字段903,600、金额字段129,537,974,760；其中酞菁绿29条/867,600，酞菁蓝2条/36,000。")
add_bullet(doc, "已证实供给背景：印度→越南查询中有211条范围候选、204个可见字段唯一签名；但印度出口侧与越南进口侧可能是同一货运的镜像，绝不能相加成独立贸易量。")
add_bullet(doc, "最集中实体：Vinh Gia承担29条B腿、数量字段807,600；Yicai承担2条、96,000；中国端东莞成翰20条/495,600，湖南亿高11条/408,000。")
add_bullet(doc, "最接近的A/B候选仍不闭合：V. G. CO. LTD.从印度接收8条酞菁蓝、数量字段14,500，但其名称不能直接等同Vinh Gia，且牌号、数量和时间均未一一对应。")
add_bullet(doc, "原产地不是只看发货国。第32章产品在越南继续归入3204时，四位税目改变路径通常不成立；但若越南增值达到30%，技术上仍可能取得越南非优惠原产地。为规避反倾销而实施的加工，海关确定原产地时可以不予考虑。")
add_bullet(doc, "公开来源未发现本措施项下已公开的反规避调查、海关处罚、法院判决或双段提单闭环。当前证据适合定向调单，不足以公开指称有关企业违法。")

doc.add_heading("二、政策范围、税率与税种", level=1)
add_para(doc, "商务部公告2023年第8号自2023年2月27日起，对原产于印度的酞菁类颜料征收反倾销税，期限5年。产品包括酞菁颜料，无论是否经过精制和/或颜料化处理；8位税号入口为32041700、32129000，海关10位商品编号为3204170020、3212900010。同税号下其他产品不当然属于措施范围。")
add_table(doc, ["印度生产商/税档", "反倾销税率", "含13%进口VAT增量的条件总系数"], [
    ["Ramdev Chemical Industries", "11.9%", "完税价格×13.447%"],
    ["Dhanveen Pigments", "14.1%", "完税价格×15.933%"],
    ["Meghmani Organics", "18.7%", "完税价格×21.131%"],
    ["其他合作印度公司", "16.0%", "完税价格×18.080%"],
    ["其他印度公司", "30.7%", "完税价格×34.691%"],
], [4200, 1700, 3460], size=9.0)
add_callout(doc, "税种口径", "反倾销税=海关审定完税价格×适用税率；如果反倾销税漏征，它还会使进口增值税计税基础减少，连带VAT差额按反倾销税×13%估算。报告中的“合计风险”只包括反倾销税及其连带进口VAT差额，不把正常关税或基础进口VAT整体写成逃税。", fill=LIGHT_BLUE)

page_break(doc)
doc.add_heading("三、数据完整性与逐条审计", level=1)
q = AUDIT["queries"]
add_table(doc, ["查询", "原始出现", "可见字段唯一", "完整性判断", "实际日期覆盖"], [
    ["PHTHALOCYANINE→中国", "129", "128", "1次完全相同可见记录重复", f"{q['downstream']['date_min']}—{q['downstream']['date_max']}"],
    ["印度→越南 PHTHALOCYANINE", "215", "208", "第1页200+第2页15；无跨页交叠；7次可见重复", f"{q['upstream']['date_min']}—{q['upstream']['date_max']}"],
    ["合计", "344", "336", "全部逐行标准化并进入工作簿", "近一年窗口"],
], [2550, 1100, 1250, 2860, 1600], size=8.8)
add_para(doc, "此前摘要中曾出现“印度→越南490条、越南→中国35票/738,620”的口径，但本地没有这490/35条的逐票原始文件，且与当前可审计数据明显冲突。本报告不混用该旧摘要，也不声称其已经逐票复核。")
add_table(doc, ["当前缺口", "影响", "补查条件"], [
    ["HS32041700全页", "关键词可能漏掉未写PHTHALOCYANINE的记录", "中国方向、印度→越南、越南→中国"],
    ["HS32129000全页", "可能漏掉分散染料/着色料项下记录", "同上"],
    ["同义词/CAS/色号", "可能漏掉只写CI号、CAS或商品色号的记录", "PIGMENT BLUE 15、PIGMENT GREEN 7、CAS147-14-8、CAS1328-53-6、COPPER PHTHALOCYANINE CRUDE"],
    ["A/B票级字段", "不能把同类供给背景升级为同批转运", "提单号、箱号、批号、越南进口申报号、库存核销、中国报关单"],
], [2400, 2800, 4160], size=8.6)

doc.add_heading("四、越南对华31条B腿", level=1)
ds = AUDIT["downstream_findings"]["vietnam_to_china_hash_vn"]
add_table(doc, ["维度", "主体/商品", "记录", "数量字段", "金额字段"], [
    ["商品", "酞菁绿", "29", "867,600", "124,555,829,160"],
    ["商品", "酞菁蓝", "2", "36,000", "4,982,145,600"],
    ["越南出口方", "Vinh Gia", "29", "807,600", "116,207,160,360"],
    ["越南出口方", "Yicai", "2", "96,000", "13,330,814,400"],
    ["中国方", "DONGGUAN CHENGHAN", "20", "495,600", "71,623,166,040"],
    ["中国方", "HUNAN YIGAO", "11", "408,000", "57,914,808,720"],
], [1600, 2800, 900, 1700, 2360], size=8.7)
add_para(doc, "另有1条越南→中国记录数量字段45、金额字段31,500,000，货描尾码为#&DE。该票只能说明德国标记货物经越南发运，不属于本项目所讨论的印度绕道数量，已从31条及税差情景中剔除。")

doc.add_heading("4.1 Yicai两条可定向核查记录", level=2)
add_table(doc, ["日期", "商品", "中国接收方", "数量字段", "金额字段", "研判"], [
    ["2025-08-29", "PIGMENT GREEN G070 / CAS1328-53-6 / #&VN", "HUNAN YIGAO", "48,000", "6,669,350,400", "B腿成立；无同名A腿"],
    ["2025-09-26", "PIGMENT GREEN G070 / CAS1328-53-6 / #&VN", "HUNAN YIGAO", "48,000", "6,661,464,000", "B腿成立；存在弱时间/倍数候选"],
], [1150, 2800, 1600, 1050, 1600, 1160], size=8.3)
add_callout(doc, "弱候选不等于闭环", "印度A腿中最接近的是2025-09-09 Riverside→越南24,000的绿色RF704，至9月26日相隔17天，但越南收货人隐藏、规格不同、数量仅为B腿的一半且没有第二条匹配A腿。因此只能列C级待核线索，不能写成Yicai使用该批印度货。", fill=PALE_GOLD)

doc.add_heading("五、印度→越南A腿及V.G.简称候选", level=1)
up = AUDIT["upstream_findings"]
add_table(doc, ["观察视角", "记录", "数量字段", "金额字段", "解释限制"], [
    ["印度出口侧", "95", fmt(up["india_export_perspective"]["quantity_field_sum"], 2), fmt(up["india_export_perspective"]["amount_field_sum"], 2), "可能与越南进口侧为同一货运镜像"],
    ["越南进口侧", "116", fmt(up["vietnam_import_perspective"]["quantity_field_sum"], 2), fmt(up["vietnam_import_perspective"]["amount_field_sum"], 2), "不得与印度出口侧直接相加"],
    ["范围候选合计", "211", "—", "—", "仅表示逐行筛查池，不是独立票数或吨位合计"],
], [1800, 900, 1500, 1700, 3460], size=8.7)
add_para(doc, "V. G. CO. LTD.在印度出口侧出现8条、数量字段14,500，均为SUYOG DYE CHEMIE供应的VEEFAST BLUE 5300/5300PL/5310/5311（Pigment Blue 15:0/15:1/15:3），日期为2025年10月14日和12月19日。Vinh Gia的B腿则以G070酞菁绿为主，只有两条BLUE150030，合计36,000。")
add_table(doc, ["比对项", "V.G. A腿", "Vinh Gia B腿", "结论"], [
    ["名称", "V. G. CO. LTD.", "CÔNG TY TNHH XUẤT NHẬP KHẨU VĨNH GIA", "非精确同名；需地址/税号"],
    ["主要产品", "VEEFAST BLUE 5300系列", "G070绿色为主；蓝色为BLUE150030", "牌号不一致"],
    ["数量", "14,500", "蓝色36,000；全部807,600", "无法一一平衡"],
    ["最近时间", "2025-12-19 9,500", "2026-03-06 蓝色24,000", "相隔77天，但数量和牌号不合"],
    ["企业识别", "本地记录无越南税号", "公开信息税号3703186307", "必须取得A腿原始收货人税号才可升级"],
], [1400, 2350, 2550, 3060], size=8.5)
add_callout(doc, "证据等级", "V.G.只能作为“简称候选”列入反查清单，不能在报告中直接写成Vinh Gia，更不能据此计算其印度输入与对华输出的货量平衡。", fill=PALE_RED, title_color=RED)

doc.add_heading("5.1 Riverside与Vinh Gia的四组数量/时间候选", level=2)
add_table(doc, ["印度A腿", "越南对华B腿", "间隔", "为何不能闭合"], [
    ["2025-10-07｜绿色RF707｜24,000", "2025-10-14｜G070｜24,000", "7天", "A腿越南收货人隐藏，牌号不同"],
    ["2026-01-31｜RF707/G030｜24,000", "2026-02-09｜G070｜24,000", "9天", "主体与批号不明，规格不同"],
    ["2026-02-27｜RF707/G030｜24,000", "2026-03-05｜G070｜24,000；3月6日另有绿色24,000", "6—7天", "一条A腿无法解释两条B腿"],
    ["2025-09-09｜绿色RF704｜24,000", "2025-10-02｜两条G070各24,000", "23天", "一条A腿无法解释两条B腿；主体隐藏"],
], [2450, 3150, 900, 2860], size=8.2)
add_para(doc, "四组记录只能构成C级“数量字段相同＋短时间窗口”筛选线索。A腿没有出现G070/BLUE150030同牌号，且缺少越南进口申报号、箱号和批号；真实加工、境内采购或不同批货物均是合理替代解释。")

doc.add_heading("六、实体网络与替代解释", level=1)
add_table(doc, ["实体", "公开/贸易数据事实", "风险意义", "不得越界"], [
    ["Vinh Gia", "2024年设立的进出口贸易公司；官网关联Dongguan Licai及Yeong Shing Vietnam，并描述母粒配色/混炼能力", "贸易—加工—对华销售网络可定向核查", "不证明采购印度酞菁，也不证明31条粉末经过实质加工"],
    ["Yicai", "公开页面称塑料色母/颜色相关制造主体；易迅有两条G070对华B腿", "需查工厂产能、BOM和进口申报", "主体规模与工艺信息主要来自二手源"],
    ["Dongguan Chenghan", "二手贸易情报显示多次从越南进口，供应方含Vinh Gia/Yicai", "是中国端最集中调单对象", "平台聚合数量不能替代中国报关单"],
    ["Hunan Yigao", "公开登记/贸易页面显示新材料、颜料销售及进出口活动", "接收Yicai与Vinh Gia记录", "不能仅凭经营范围认定实际用途或纳税状态"],
], [1600, 3050, 2350, 2360], size=8.2)
add_para(doc, "Vinh Gia集团具备母粒混炼能力，为“越南真实加工并形成新商品/本地增值”提供合理替代解释；但31条B腿货描是酞菁颜料粉末而非塑料母粒。是否发生足以改变原产地的精制、颜料化或表面处理，仍须逐批核对BOM、工单、投入产出、能耗和成本账。")

doc.add_heading("6.1 Heubach/Sudarshan跨境关联商业链", level=2)
add_para(doc, "印度证券交易所审计披露显示，终裁16%税档列名实体Heubach Colour Private Limited在2024年10月至2025年3月与多国关联方发生“货物或服务销售”。这是目前公开来源中最具体的跨境商业/开票网络证据，但披露未列产品、HS、发票或提单，不能认定这些交易就是酞菁颜料，更不能认定物理转运。")
add_table(doc, ["关联交易对手", "披露金额（印度卢比）", "节点性质/核查意义"], [
    ["Heubach Colorants (Shanghai) Ltd.", "30,865,000", "最具体的中国关联交易对手；查产品、发票、进口申报"],
    ["Heubach Colorants Singapore Pte. Ltd.", "27,906,000", "ISO范围以销售/供应链为主，优先核制造厂证明"],
    ["Heubach Colorants (Thailand) Ltd.", "5,833,000", "泰国另有制造节点，不能仅凭关联推定转运"],
    ["Heubach日本两实体", "32,366,000", "日本存在制造节点；须落具体牌号和批次"],
    ["P.T. Heubach Colorants Coatings Indonesia", "2,871,000", "印尼存在制造能力；当地生产是合理替代解释"],
], [3600, 2000, 3760], size=8.3)
add_para(doc, "2025年Sudarshan完成对Heubach集团收购，后续检索需同时覆盖Heubach Colour Private Limited、Sudarshan Gujarat MFG、Heubach/Sudarshan Singapore及Thailand等新旧别名。公开资料尚不足以确认名称变更后中国海关实际适用何种税档，应调公司注册证明和进口税单。")

doc.add_heading("6.2 越南渠道与宏观量异常", level=2)
add_para(doc, "印度CHEMEXCIL 2022年越南涂料展官方报告列出Ramdev携PB15系列/PG7、Unilex携320417有机颜料在越南营销，证明列名受税企业在越南具有正常商业渠道；渠道在终裁前即已存在，也构成“正常市场开发”反证。UN Comtrade显示中国自越南进口HS320417由2022年12,792千克升至2023年965,026千克、2024年1,541,270千克，但HS6涵盖多种合成有机颜料，只能作为C级宏观筛查，不能证明本案产品绕道。")

doc.add_heading("七、非优惠原产地与“绕道”判定门槛", level=1)
add_para(doc, "反倾销措施适用的是中国非优惠原产地规则，不能把RCEP优惠原产地证直接当成反倾销原产地结论。涉及两个以上国家生产时，一般以最后完成实质性改变的国家（地区）为原产地。")
add_table(doc, ["测试", "第32章适用口径", "本项目判断"], [
    ["四位税目改变", "使用货物本身四位税目以外的原料制成", "若印度输入和越南输出都归3204，税目改变路径不成立"],
    ["从价百分比", "（工厂交货价－非越南原料CIF）÷工厂交货价≥30%", "真实精制/颜料化并达到30%时，技术上可能取得越南非优惠原产地"],
    ["微小加工", "为运输、储存、销售所作的包装、分装等通常不足以改变原产地", "简单换包、换标、仓储再出口不能支持越南原产"],
    ["反规避条款", "为规避反倾销等措施而实施的加工或处理，海关确定原产地时可以不予考虑", "即使形式上有加工，也须审查目的、经济实质和时间数量对应"],
], [1800, 3500, 4060], size=8.5)
add_callout(doc, "正确表述", "现阶段不能笼统说“印度粗品在越南加工后仍必然是印度原产”，也不能因#&VN就认定越南原产。应逐批计算30%增值、核实四位税目、生产工艺和反规避因素，再以中国海关最终认定为准。", fill=PALE_GOLD)

doc.add_heading("八、条件税差情景", level=1)
tax_rows = []
labels = {
    0.119: "Ramdev 11.9%",
    0.141: "Dhanveen 14.1%",
    0.160: "其他合作公司 16.0%",
    0.187: "Meghmani 18.7%",
    0.307: "其他印度公司 30.7%",
}
for item in AUDIT["conditional_tax_scenarios_for_31_hash_vn_records"]:
    rate = round(item["ad_rate"], 3)
    tax_rows.append([
        labels[rate],
        fmt(item["amount_field_proxy"]),
        fmt(item["conditional_antidumping_duty"]),
        fmt(item["conditional_import_vat_delta"]),
        fmt(item["conditional_total_increment"]),
    ])
add_table(doc, ["生产商/税率情景", "金额字段代理", "条件反倾销税", "条件VAT差额", "条件合计"], tax_rows, [2000, 1800, 1850, 1750, 1960], size=8.5)
add_callout(doc, "五项严格前提", "上述数字同时假设：①31条实际为印度原产；②产品属于公告范围；③印度生产商及税档判断正确；④易迅金额字段币种和数值可作为中国海关完税价格代理；⑤中国端完全未征反倾销税。任一前提不成立，数字都必须调整或归零。因此表中是同币种风险暴露情景，不是实际欠税额。", fill=PALE_RED, title_color=RED)

doc.add_heading("九、证据分级与总体评级", level=1)
add_table(doc, ["等级", "当前证据", "可支持结论", "尚缺"], [
    ["B+", "31条#&VN越南→中国具体记录", "越南对华B腿、主体和货描成立", "中国报关原产国、税单"],
    ["B", "211条印度→越南范围候选", "同期开存在印度对越同类供给", "与31条B腿同批匹配"],
    ["C", "V.G.简称候选、Yicai弱时间/数量候选", "可定向调越南进口申报", "税号、提单、批号、工单"],
    ["未闭合", "无箱号/批号/双段申报和中国税单", "逃税、伪报原产均未证实", "A腿—加工—B腿—中国申报全链"],
], [1000, 3100, 2500, 2760], size=8.7)
add_callout(doc, "总体风险评级", "第三国再出口与原产地核查风险：中高；已证实逃避反倾销税：未发现。优先顺序为Vinh Gia→Dongguan Chenghan、Vinh Gia/Yicai→Hunan Yigao，再反查V.G. CO. LTD.和Riverside的越南进口收货人。", fill=PALE_GOLD)

doc.add_heading("十、进一步核查清单", level=1)
for text in [
    "对31条中国进口记录调取报关单、原产地证、税款缴款书，核对原产国、启运国、生产商、反倾销税率和缴税金额。",
    "取得越南A腿原始进口申报及提单，优先核验V. G. CO. LTD.收货地址/税号是否为3703186307，并核Riverside 2025-09-09记录的隐藏收货人。",
    "按30—180天窗口做A/B匹配；重量容差设±3%—5%，同时比对25kg包装、完整牌号、CAS、批号、箱号和封志。",
    "对Vinh Gia、Yicai及Yeong Shing Vietnam调取BOM、工单、原料领用、成品入库、产量收率、能耗、设备和人员记录，判断是否真实完成精制/颜料化。",
    "按第32章规则逐批核算越南增值：（工厂交货价－非越南原料CIF）÷工厂交货价；取得Form B申请底稿及签证机构核验材料。",
    "补做HS32041700、32129000及蓝15/绿7、两项CAS、COPPER PHTHALOCYANINE CRUDE的全页查询，保留条件截图、总数和分页完整性证据。",
]:
    add_number(doc, text)

doc.add_heading("十一、主要公开来源", level=1)
sources = [
    ("商务部公告2023年第8号", "https://www.mofcom.gov.cn/zcfb/gpmy/art/2023/art_9703e40264f342ea924f8fbb0db7a14a.html", "产品范围、实施期、税率与计税公式"),
    ("商务部税率附件", "https://images.mofcom.gov.cn/trb/202302/20230225183421737.pdf", "列名印度生产商和税率"),
    ("海关总署公告2022年第104号", "http://www.customs.gov.cn/customs/302249/2480148/4658463/index.html", "10位商品编号3204170020、3212900010"),
    ("《中华人民共和国进出口货物原产地条例》", "https://xzfg.moj.gov.cn/front/law/detail?LawID=1523&Query=%E4%BB%A5%E8%B4%A7%E7%89%A9", "实质性改变、微小加工及反规避条款"),
    ("海关总署令第273号原产地规则", "https://policy.mofcom.gov.cn/claw/clawContent.shtml?id=101350", "四位税目改变、30%从价公式与清单机制"),
    ("原122号规则及制造/加工工序清单", "https://www.moj.gov.cn/pub/sfbgw/flfggz/flfggzbmgz/200504/t20050412_143920.html", "第32章判定标准索引"),
    ("Vinh Gia集团网站", "https://vietmasterbatch.com/", "贸易公司、关联企业与母粒混炼能力；仅作主体网络自述"),
    ("Vinh Gia企业登记聚合页", "https://thuvienphapluat.vn/ma-so-thue/cong-ty-tnhh-xuat-nhap-khau-vinh-gia-mst-3703186307.html?hl=en", "税号3703186307及地址；二手登记聚合"),
    ("Yicai企业信息页", "https://www.emis.com/php/company-profile/VN/Yicai_Color_Plastic_Company_Limited__C%C3%B4ng_Ty_Tnhh_Nh%E1%BB%B1a_M%C3%A0u_Yicai___C%C3%B4ng_Ty_Tnhh_Nh%E1%BB%B1a_M%C3%A0u_Yicai__Yicai_Color_Plastic_Company_Limited___en_14827786.html", "主体经营信息；二手来源"),
    ("Dongguan Chenghan贸易聚合页", "https://www.volza.com/company-profile/dongguan-chenghan-plastic-trading-co.-ltd.-37328082/", "越南供应方关系样本；非官方、不用于认定票数"),
    ("NSE Heubach关联交易披露", "https://nsearchives.nseindia.com/corporate/ixbrl/INTEGRATED_FILING_INDAS_105990_26072025224150_iXBRL_WEB.html", "印度实体与上海、新加坡、泰国、日本、印尼关联交易金额"),
    ("CHEMEXCIL越南展报告", "https://chemexcil.in/uploads/events/DETAILED_REPORT_COATING_EXPO_VIETNAM_2022.pdf", "Ramdev、Unilex在越南的PB15/PG7和320417渠道"),
    ("UN Comtrade 2024进口筛查", "https://comtradeapi.un.org/public/v1/preview/C/A/HS?period=2024&reporterCode=156&cmdCode=320417&flowCode=M&partnerCode=356,704,702,458,410,764,392,360&maxRecords=500", "HS6宏观异常；不等于酞菁专属数据"),
]
for label, url, note in sources:
    add_source(doc, label, url, note)

page_break(doc)
doc.add_heading("附录一：31条越南对华重点记录", level=1)
appendix_rows = []
for row in LEADS:
    desc = row["description"]
    product = row["product_class"]
    cas = "1328-53-6" if "1328-53-6" in desc else ("147-14-8" if "147-14-8" in desc else "—")
    seller = row["seller_or_shipper"]
    if "VĩNH GIA" in seller.upper() or "VĨNH GIA" in seller.upper():
        seller = "Vinh Gia"
    elif "YICAI" in seller.upper():
        seller = "Yicai"
    buyer = row["buyer_or_consignee"].replace(" CO.,LTD", "").replace(" CO., LTD", "")
    appendix_rows.append([
        row["date"],
        product,
        cas,
        seller,
        buyer,
        fmt(row["quantity_num"], 0),
        fmt(row["amount_num"], 0),
    ])
add_table(doc, ["日期", "品类", "CAS", "越南方", "中国方", "数量字段", "金额字段"], appendix_rows, [1050, 950, 1100, 1300, 2500, 1050, 1410], size=7.5)

doc.add_heading("附录二：交付文件", level=1)
for text in [
    "酞菁类颜料_易迅逐票判定台账_阶段审计.xlsx：含概览、31条B腿、211条A腿候选、344条全部记录、重复审计和政策取证口径。",
    "酞菁类颜料_易迅逐票标准化.csv/json：344条逐行标准化底稿。",
    "酞菁类颜料_全量阶段审计.json：统计、证据分级、条件税差和查询缺口。",
    "两份易迅原始读取JSON：保留129条中国方向和215条印度→越南的原始页面记录。",
]:
    add_bullet(doc, text)
add_callout(doc, "报告状态", "本文件为阶段审计报告。完成两项税号、扩展关键词、企业反查及中国端报关/税单核验后，应覆盖更新为最终深度分析报告，并同步总报告、总清单和前端系统。", fill=PALE_RED, title_color=RED)

doc.core_properties.title = "酞菁类颜料反倾销税与第三国转运风险阶段审计报告"
doc.core_properties.subject = "易迅数据逐条审计、印度—越南—中国供应链、非优惠原产地和条件税差"
doc.core_properties.author = "Codex"
doc.core_properties.keywords = "酞菁类颜料, 反倾销税, 印度, 越南, 第三国转运, 原产地"
doc.save(REPORT_PATH)
print(REPORT_PATH)
