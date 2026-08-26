"""Incrementally ingest GDRM26-394/395 source files into the case knowledge base."""

from __future__ import annotations

import csv
import hashlib
import json
import re
import shutil
import sqlite3
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader


sys.stdout.reconfigure(encoding="utf-8", errors="replace")

ROOT = Path(r"D:\codex\XG执法\2026_案例库")
DB = ROOT / "enforcement_cases_2026.sqlite"
CATALOG = ROOT / "case_catalog.csv"
AUDIT = ROOT / "新增案件394_395入库结果.json"

CASES = {
    "GDRM26-394": {
        "package_name": "新增/GDRM26-394情報158號一宗河路貨運走私未列艙單貨物案件",
        "case_type": "情報",
        "transport_mode": "河路",
        "date_hint": "2026-08-19",
        "files": [
            Path(r"E:\ZMJ\GDRM26-394情報158號一宗河路貨運走私未列艙單貨物案件.pdf"),
            Path(r"E:\ZMJ\GDRM26-394附件一及二.pdf"),
        ],
    },
    "GDRM26-395": {
        "package_name": "新增/GDRM26-395情報159號兩宗海路貨運走私廢金屬案件",
        "case_type": "情報",
        "transport_mode": "海路",
        "date_hint": "2026-08-19",
        "files": [
            Path(r"E:\ZMJ\GDRM26-395情報159號兩宗海路貨運走私廢金屬案件.pdf"),
            Path(r"E:\ZMJ\GDRM26-395附件一及二.pdf"),
            Path(r"E:\ZMJ\GDRM26-395附件三涉案報關紀錄.xlsx"),
        ],
    },
}


def normalize_text(value: str) -> str:
    lines = []
    for line in (value or "").replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        cleaned = re.sub(r"[ \t]+", " ", line).strip()
        if cleaned:
            lines.append(cleaned)
    return "\n".join(lines)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        while chunk := stream.read(1024 * 1024):
            digest.update(chunk)
    return digest.hexdigest()


def xlsx_text(path: Path) -> str:
    book = load_workbook(path, read_only=True, data_only=True)
    values = []
    for sheet in book.worksheets:
        values.append(f"[工作表] {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            line = " | ".join(
                normalize_text(str(cell)) for cell in row if cell is not None
            )
            if line:
                values.append(line)
    return "\n".join(values)


def load_catalog() -> list[dict[str, str]]:
    with CATALOG.open("r", newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle))


def write_catalog(rows: list[dict[str, str]]) -> None:
    fields = [
        "case_id", "package_name", "case_type", "transport_mode",
        "date_hint", "document_count", "raw_path",
    ]
    rows.sort(key=lambda row: row["case_id"])
    with CATALOG.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    catalog_rows = load_catalog()
    catalog_ids = {row["case_id"] for row in catalog_rows}

    for case_id, meta in CASES.items():
        if case_id in catalog_ids:
            raise RuntimeError(f"Catalog already contains {case_id}")
        for source in meta["files"]:
            if not source.is_file():
                raise FileNotFoundError(source)
            destination = ROOT / "raw_cases" / case_id / source.name
            if destination.exists():
                raise RuntimeError(f"Destination already exists: {destination}")

    conn = sqlite3.connect(DB)
    conn.execute("PRAGMA foreign_keys=ON")
    case_ids = tuple(CASES)
    placeholders = ",".join("?" for _ in case_ids)
    if conn.execute(
        f"SELECT count(*) FROM cases WHERE case_id IN ({placeholders})", case_ids
    ).fetchone()[0]:
        raise RuntimeError("Database already contains one of the requested case IDs")
    all_hashes = [sha256(path) for meta in CASES.values() for path in meta["files"]]
    hash_placeholders = ",".join("?" for _ in all_hashes)
    if conn.execute(
        f"SELECT count(*) FROM documents WHERE sha256 IN ({hash_placeholders})",
        all_hashes,
    ).fetchone()[0]:
        raise RuntimeError("Database already contains one of the requested file hashes")

    backup_path = ROOT / f"enforcement_cases_2026.pre_394_395_{timestamp}.sqlite"
    backup = sqlite3.connect(backup_path)
    conn.backup(backup)
    backup.close()
    catalog_backup = ROOT / f"case_catalog.pre_394_395_{timestamp}.csv"
    shutil.copy2(CATALOG, catalog_backup)

    copied = []
    inserted_files = []
    try:
        for case_id, meta in CASES.items():
            case_dir = ROOT / "raw_cases" / case_id
            case_dir.mkdir(parents=True, exist_ok=False)
            for source in meta["files"]:
                destination = case_dir / source.name
                shutil.copy2(source, destination)
                copied.append(destination)

            conn.execute(
                "INSERT INTO cases VALUES(?,?,?,?,?,?,?)",
                (
                    case_id,
                    meta["package_name"],
                    meta["case_type"],
                    meta["transport_mode"],
                    meta["date_hint"],
                    case_dir.relative_to(ROOT).as_posix(),
                    len(meta["files"]),
                ),
            )

            for destination in copied[-len(meta["files"]):]:
                relative_path = destination.relative_to(ROOT).as_posix()
                digest = sha256(destination)
                extension = destination.suffix.lower()
                file_result = {
                    "case_id": case_id,
                    "file": relative_path,
                    "extension": extension,
                    "sha256": digest,
                }
                if extension == ".pdf":
                    reader = PdfReader(str(destination))
                    if reader.is_encrypted:
                        raise RuntimeError(f"Unexpected encrypted PDF: {destination}")
                    page_results = []
                    for page_number, page in enumerate(reader.pages, 1):
                        text = normalize_text(page.extract_text() or "")
                        conn.execute(
                            """INSERT INTO documents(
                                   case_id,relative_path,extension,page_or_sheet,
                                   extracted_text,text_chars,extraction_error,sha256
                               ) VALUES(?,?,?,?,?,?,?,?)""",
                            (
                                case_id, relative_path, extension, f"page {page_number}",
                                text, len(text), None, digest,
                            ),
                        )
                        conn.execute(
                            "INSERT INTO document_fts(case_id,relative_path,extracted_text) VALUES(?,?,?)",
                            (case_id, relative_path, text),
                        )
                        page_results.append({"page": page_number, "characters": len(text)})
                    file_result["pages"] = page_results
                    file_result["characters"] = sum(row["characters"] for row in page_results)
                elif extension == ".xlsx":
                    text = xlsx_text(destination)
                    conn.execute(
                        """INSERT INTO documents(
                               case_id,relative_path,extension,page_or_sheet,
                               extracted_text,text_chars,extraction_error,sha256
                           ) VALUES(?,?,?,?,?,?,?,?)""",
                        (case_id, relative_path, extension, "workbook", text, len(text), None, digest),
                    )
                    conn.execute(
                        "INSERT INTO document_fts(case_id,relative_path,extracted_text) VALUES(?,?,?)",
                        (case_id, relative_path, text),
                    )
                    file_result["characters"] = len(text)
                else:
                    raise RuntimeError(f"Unsupported file type: {destination}")
                inserted_files.append(file_result)

            catalog_rows.append({
                "case_id": case_id,
                "package_name": meta["package_name"],
                "case_type": meta["case_type"],
                "transport_mode": meta["transport_mode"],
                "date_hint": meta["date_hint"],
                "document_count": str(len(meta["files"])),
                "raw_path": case_dir.relative_to(ROOT).as_posix(),
            })

        conn.commit()
        write_catalog(catalog_rows)
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()

    result = {
        "ingested_at": datetime.now().astimezone().isoformat(),
        "database": str(DB),
        "database_backup": str(backup_path),
        "catalog_backup": str(catalog_backup),
        "cases_added": list(CASES),
        "files_added": inserted_files,
        "stats": {
            "cases": len(CASES),
            "files": len(inserted_files),
            "pdf_pages": sum(len(row.get("pages", [])) for row in inserted_files),
            "characters": sum(row.get("characters", 0) for row in inserted_files),
        },
    }
    AUDIT.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
