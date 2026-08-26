from __future__ import annotations

import csv
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from build_epdm_stage_report import (
    BLUE, GOLD, GREEN, INK, MUTED, NAVY, PALE_BLUE, PALE_GOLD, PALE_GRAY,
    PALE_GREEN, PALE_RED, RED, add_callout, add_heading, add_hyperlink,
    add_list_item, add_page_number_field, add_paragraph, add_source,
    add_table, configure_document, create_numbering, set_run_font,
)


OUT = Path(r"D:\易迅数据\反倾销税深度分析报告\20_乙醇胺")
DOCX = OUT / "乙醇胺_反倾销税与第三国转运风险阶段深度分析报告.docx"
LEDGER = OUT / "乙醇胺_易迅逐票判定台账_现有3条.csv"
SUMMARY = OUT / "乙醇胺_交付摘要.json"


def caption(table, text="乙醇胺反倾销税与第三国转运风险阶段审计表"):
    cap=OxmlElement("w:tblCaption"); cap.set(qn("w:val"), text); table._tbl.tblPr.append(cap)
    if table.rows:
        tr_pr=table.rows[0]._tr.get_or_add_trPr(); flag=OxmlElement("w:tblHeader"); flag.set(qn("w:val"),"true"); tr_pr.append(flag)
        for c in table.rows[0].cells:
            for p in c.paragraphs: p.paragraph_format.keep_with_next=True


def rtable(doc, headers, rows, widths, **kwargs):
    t=add_table(doc, headers, rows, widths, **kwargs); caption(t); return t


def header_footer(doc):
    sec=doc.sections[0]
    h=sec.header.paragraphs[0]; h.clear(); h.alignment=WD_ALIGN_PARAGRAPH.RIGHT; h.paragraph_format.space_after=Pt(0)
    set_run_font(h.add_run("反倾销税风险穿透分析｜乙醇胺（EA）｜2026-08-13"), size=8.5, color=MUTED)
    f=sec.footer.paragraphs[0]; f.clear(); f.paragraph_format.space_after=Pt(0); add_page_number_field(f)
    if f.runs: f.runs[0].text="乙醇胺阶段深度分析报告  |  "
    doc.core_properties.title="乙醇胺反倾销税与第三国转运风险阶段深度分析报告"
    doc.core_properties.subject="政策税率、易迅逐票审计、第三国链路、实体与调证"
    doc.core_properties.author="反倾销税风险分析项目"
    doc.core_properties.keywords="乙醇胺; MEA; DEA; TEA; 反倾销; 第三国转运; 易迅数据"


def main():
    rows=list(csv.DictReader(LEDGER.open(encoding="utf-8-sig", newline="")))
    summary=json.loads(SUMMARY.read_text(encoding="utf-8")); assert len(rows)==3
    doc=Document(); configure_document(doc); header_footer(doc)
    bullets=create_numbering(doc, bullet=True); decimals=create_numbering(doc, bullet=False)

    # standard_business_brief + memo_masthead
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(18); p.paragraph_format.space_after=Pt(4)
    set_run_font(p.add_run("第20项商品核查"), size=11, color=GOLD, bold=True)
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(4)
    set_run_font(p.add_run("乙醇胺（EA）"), size=27, color=NAVY, bold=True)
    add_paragraph(doc,"反倾销税与第三国转运风险｜阶段深度分析报告",size=15.5,color=BLUE,bold=True,after=14)
    rtable(doc,["基准日","措施状态","现有易迅逐票","闭合绕道链"],[["2026-08-13","实施中；续征至2029-10-29","3条，均为下游切削油","0条"]],[1700,2800,2700,2160],header_fill=PALE_BLUE,font_size=9.5,aligns=[WD_ALIGN_PARAGRAPH.CENTER]*4)
    add_callout(doc,"核心结论","D盘329个有效文件内容级盘点只回溯到3条贸易记录：均为韩国平台原产字段的PROCUT 3800切削油进入越南，三乙醇胺仅占12%–15%，境外HS为34031919。现有记录既非美国、沙特、马来西亚、泰国纯乙醇胺A腿，也非第三国进入中国的B腿；未形成第三国绕道或逃税证据。",fill=PALE_GOLD,title_color=GOLD)
    rtable(doc,["证据层","现有事实","等级","结论边界"],[
        ["政策","受税来源、税号、企业税率和计税公式已核","A","可据中国报关底单核税"],
        ["易迅贸易","3条韩国→越南含TEA切削油","C/排除","仅说明下游制剂贸易，不是涉案纯品"],
        ["第三国A/B腿","0条可闭合","未形成","无同品种、同批次、同柜或同主体双腿"],
        ["税款","无中国报关单/完税价/税单","未形成","不能计算实际逃税额"],
    ],[1500,3300,1300,3260],header_fill=PALE_GRAY,font_size=8.8)
    add_paragraph(doc,"完整性声明：乙醇胺专项易迅网页全页查询因页面读取超时未完成；本报告只代表D盘现有329个有效文件和可回溯页面采集，不把本地0条涉案纯品冒充易迅全库无记录。",size=9.2,color=RED,bold=True,after=0)

    doc.add_page_break()
    add_heading(doc,"1. 现行政策、产品范围和税率",1)
    add_paragraph(doc,"商务部2024年第44号公告决定，自2024年10月30日起继续对原产于美国、沙特阿拉伯、马来西亚和泰国的进口乙醇胺征收反倾销税，实施期限5年，预计至2029年10月29日。产品包括一乙醇胺（MEA）、二乙醇胺（DEA）和三乙醇胺（TEA），中国税号29221100、29221200、29221500；29221100项下一乙醇胺盐、29221200项下二乙醇胺盐明确排除。")
    rtable(doc,["来源","列名生产商/销售链","AD税率","AD+由AD增加的13%VAT"],[
        ["美国","Dow Chemical","76.0%","完税价×85.880%"],
        ["美国","INEOS Americas / Huntsman / 其他","97.1%","完税价×109.723%"],
        ["沙特","SABIC","10.1%","完税价×11.413%"],
        ["沙特","其他（含未获列名主体）","27.9%","完税价×31.527%"],
        ["马来西亚","PC Derivatives生产+PC Marketing (Labuan)出口","18.3%","完税价×20.679%"],
        ["马来西亚","其他任何组合","20.3%","完税价×22.939%"],
        ["泰国","TOC Glycol / 其他","37.6%","完税价×42.488%"],
    ],[1150,3900,1300,3010],header_fill=PALE_BLUE,font_size=8.35)
    add_callout(doc,"计税口径","反倾销税=中国海关审定完税价格×企业适用税率；进口增值税计税基础含反倾销税。本报告的综合系数只用于筛查“少缴AD及其引致的13%进口VAT增量”，不是整票总税负。",fill=PALE_GREEN,title_color=GREEN)
    add_heading(doc,"1.1 最容易发生的税率错配",2)
    for text in [
        "美国：若美国其他公司货误套Dow 76.0%，条件税差为中国完税价×(97.1%−76.0%)×1.13=完税价×23.843%。",
        "沙特：Sadara有真实沙特乙醇胺装置，但未在公告中单列；若误套SABIC 10.1%而非其他沙特27.9%，条件税差为完税价×20.114%。",
        "马来西亚：18.3%只适用于“PETRONAS Chemicals Derivatives生产+PETRONAS Chemicals Marketing (Labuan)出口”的指定组合；其他生产商、其他出口商或组合不闭合时为20.3%，差额系数为完税价×2.260%。",
        "品牌不能替代生产厂：INEOS官方同时列美国Plaquemine和法国Lavera乙醇胺制造基地。法国真实生产不属于本措施；由美国生产、经欧洲销售才仍需按美国原产核税。",
    ]: add_list_item(doc,text,bullets)

    add_heading(doc,"2. 数据覆盖与逐票方法",1)
    rtable(doc,["环节","结果","质量控制"],[
        ["文件盘点","329个XLSX/CSV/JSON","派生报告只登记不重复计票"],
        ["原始候选","3个文件","全称、CAS、措施税号及化学上下文缩写宽抓"],
        ["逐行深扫","0解析错误","逐条回溯原始页面位置"],
        ["贸易记录","3条","项目清单1行未计贸易票"],
        ["涉案纯品/中国B腿","0条/0条","下游混合物不强行按TEA征税"],
    ],[1750,2300,5310],header_fill=PALE_GRAY,font_size=9)
    add_paragraph(doc,"缩写控制：MEA、DEA、TEA可指其他化学品、企业或普通英文词。只有全称、CAS、措施税号，或缩写与AMINE/ETHANOL/CHEMICAL等化学上下文共同出现，才进入逐票池。Triethylamine（三乙胺）、MDEA、DEIPA、酰胺/表活衍生物、乙醇胺盐和下游制剂须分开判断。")

    add_heading(doc,"3. 现有3条易迅记录全量逐票结论",1)
    rtable(doc,["日期","境外HS","双方主体","数量字段","金额字段"],[[r["日期"],r["境外HS"],"SHIN HO HIGH-TECH→SSEP Vietnam",r["数量字段"],r["金额字段"]] for r in rows],[1300,1300,3200,1400,2160],header_fill=PALE_BLUE,font_size=8.7)
    add_paragraph(doc,"三条数量字段合计10,000，金额字段合计695,701,699；原始页面未给单位和币种，因此只能保留为“数量字段/金额字段”，不能直接写千克或越南盾。货描均为PROCUT 3800金属切削油：石蜡油35%–43%、三乙醇胺12%–15%、蓖麻油酸10%–15%、菜籽油10%–17%及其他组分。若数量字段可等同成品重量，内含TEA的机械区间是1,200–1,500个同单位，但这不是涉案纯乙醇胺贸易数量。")
    add_callout(doc,"范围排除结论","三票均为HS34031919下游润滑/切削油制剂，目的国越南，平台原产字段韩国。它们不构成美国、沙特、马来西亚、泰国乙醇胺A腿，也不是第三国→中国B腿；不应据此测算中国反倾销税。",fill=PALE_RED,title_color=RED)
    add_heading(doc,"3.1 可保留的供应链线索",2)
    add_paragraph(doc,"SHIN HO HIGH-TECH向SSEP Vietnam三次稳定供应同配方PROCUT 3800，日期分别为2026-02-12、04-15、06-03，数量字段2,000/4,000/4,000。该链只说明韩国含TEA切削油在越南的连续贸易。若未来发现SSEP或关联方将纯MEA/DEA/TEA或相同制剂发往中国，应先做产品范围和归类，而不能从现有三票倒推其上游TEA原产。")
    add_heading(doc,"4. 第三国绕道证据评估",1)
    rtable(doc,["闭环要件","现有证据","缺失证据","判断"],[
        ["受税来源A腿","0条纯MEA/DEA/TEA","出口报关、生产商、批号、提单","未形成"],
        ["第三国B腿","0条进入中国","中国进口明细、境外发货人","未形成"],
        ["同货匹配","无","同品种/等级/包装/批号/柜号/30–90日","未形成"],
        ["原产改变","无第三国加工记录","BOM、工单、能耗、产销存、原产证底稿","无法判断"],
        ["少缴税款","无中国底单","完税价格、企业税率、AD/VAT缴款书","无法测算"],
    ],[1550,2400,3400,2010],header_fill=PALE_GRAY,font_size=8.6)
    add_paragraph(doc,"公开检索截至2026-08-13未发现中国海关、法院或商务部公开披露本措施下“受税来源乙醇胺经第三国进入中国”的处罚、判决、反规避裁定或双段提单闭环。这个负面检索结果不等于现实中不存在风险，只表示当前公开证据未达到可定性标准。")

    add_heading(doc,"5. 实体与路线优先级",1)
    rtable(doc,["优先级","实体/路线","为什么值得核","合法替代解释"],[
        ["A","PETRONAS Derivatives—Marketing (Labuan)—中国","18.3%取决于指定生产+销售组合","马来西亚Kertih有75kt真实产能"],
        ["A","Sadara/SABIC—第三国—中国","沙特生产商税率10.1%与27.9%差异大","Sadara与SABIC均有沙特真实生产"],
        ["A","INEOS Americas/Europe—中国","美国97.1%与法国非受税原产差异巨大","Plaquemine和Lavera均真实制造"],
        ["B","Dow/Huntsman美国—新加坡/欧洲—中国","销售/开票节点可能遮蔽美国生产厂","第三国销售不等于第三国原产"],
        ["C","SHIN HO→SSEP Vietnam","现有3条连续含TEA制剂链","下游切削油真实消费/加工"],
    ],[900,2800,3000,2660],header_fill=PALE_BLUE,font_size=8.45)
    add_heading(doc,"5.1 必调字段",2)
    for text in [
        "中国进口报关：29221100/29221200/29221500、品名/CAS/纯度/等级、原产国、启运国、生产商、境外发货人、进口人、申报企业、口岸、净重、完税价格、适用AD税率和税款书。",
        "物流单证：商业发票、原厂发票、合同、付款受益人、提单、柜号、封志、船名航次、中转港、包装唛头、COA批号和plant code。",
        "第三国加工：原料进口报关、BOM、领料、工单、设备/能耗、产量/收率、库存、销售发票、原产地证申请底稿；仅分装、换标、仓储或换单不构成充分的原产改变证据。",
        "马来西亚特殊税率：逐票同时核生产商是否PC Derivatives、出口销售是否PC Marketing (Labuan)，任一不符即不能机械套18.3%。",
    ]: add_list_item(doc,text,bullets)

    add_heading(doc,"6. 下一轮易迅全页查询",1)
    add_paragraph(doc,"已生成三组最简条件，避免表格过多：Q1中国进口主查询；Q2四个受税来源至重点第三国A腿；Q3候选第三国至中国B腿。每次切换200条/页，必须读完全部页，再按可见全字段去重，并保留疑似重复而不擅自删除。")
    rtable(doc,["查询","核心条件","完成标准"],[
        ["EA-Q1","中国进口；HS292211/292212/292215；全称+3个CAS","全页逐票分类纯品/盐/衍生物/制剂"],
        ["EA-Q2","美/沙/马/泰→重点第三国；叠加列名实体","记录生产商、销售方、等级、批号、数量"],
        ["EA-Q3","第三国→中国；同HS/关键词/实体","30/60/90日匹配并回中国原产申报与税单"],
    ],[1200,5200,2960],header_fill=PALE_GRAY,font_size=8.8)
    add_callout(doc,"网页状态","Chrome中的易迅页面仍保留登录状态，但本轮读取页面结构超时，未取得可验证的页数和总记录数。因此报告将网页专项查询列为待补，不把超时解释成0条。",fill=PALE_RED,title_color=RED)

    add_heading(doc,"7. 阶段结论",1)
    add_paragraph(doc,"风险评级：当前第三国绕道证据为C/未形成。现有三票均已排除为韩国→越南含TEA下游制剂；0条涉案纯品、0条中国B腿、0条双段闭环、0元可确认逃税额。结构性核查优先级仍然高，原因是美国最高97.1%、沙特最低10.1%，且存在列名生产商、其他生产商、跨国双产地和指定销售链导致的巨大税率差。下一步最有效动作不是继续扩大品牌猜测，而是完成Q1–Q3全页并调中国报关底单。")

    add_heading(doc,"来源",1)
    sources=[
        ("S1","商务部公告2024年第44号及期终复审裁定","https://cacs.mofcom.gov.cn/cacscms/article/jkdc?articleId=182280&type=1","范围、税号、税率、期限、公式及官方进口基线。"),
        ("S2","PETRONAS乙醇胺产品资料","https://www.petronas.com/pcg/sites/pcg/files/download/ethanolamines-brochure.pdf","Kertih生产厂、75kt产能、MEA/DEA/TEA和等级。"),
        ("S3","INEOS Oxide产品页","https://www.ineos.com/businesses/ineos-oxide/products/","Plaquemine美国和Lavera法国双生产基地及产品范围。"),
        ("S4","Sadara Amines Facility投产公告","https://sadara.com/-/media/News%20Articles/2017/07/Sadara%20announces%20startup%20of%20its%20Amines%20Facility.ashx?la=en","Jubail乙醇胺/乙烯胺真实制造能力。"),
        ("S5","国务院进出口货物原产地条例","https://xzfg.moj.gov.cn/front/law/detail?LawID=1523&Query=%E8%B4%A7%E7%89%A9%E5%8E%9F%E4%BA%A7%E5%9C%B0","反倾销适用非优惠原产地及最后实质性改变规则。"),
    ]
    for sid,title,url,note in sources: add_source(doc,bullets,sid,title,url,note)
    add_paragraph(doc,"本报告是基于现有易迅页面采集、本地文件盘点和公开资料的风险筛查，不替代海关归类、原产地核定、税款稽核或司法认定。",size=8.8,color=MUTED,after=0)

    doc.save(DOCX)
    assert DOCX.exists() and DOCX.stat().st_size>30000
    print(DOCX)


if __name__=="__main__": main()
