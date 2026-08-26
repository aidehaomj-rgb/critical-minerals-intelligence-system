from __future__ import annotations

import json
import os
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


DIR = Path(os.environ.get("PC_OUT_DIR", r"D:\易迅数据\反倾销税深度分析报告\07_聚碳酸酯"))
AUDIT = json.loads((DIR / "聚碳酸酯_易迅_HS390740_全量分析结果.json").read_text("utf-8"))
REVIEW = json.loads((DIR / "聚碳酸酯_TW线索人工复核.json").read_text("utf-8"))
SUMMARY = REVIEW["summary"]
LEADS = REVIEW["records"]
OUT = DIR / "聚碳酸酯_反倾销税与第三国转运风险阶段审计报告.docx"

NAVY = "17324D"
BLUE = "1F4E78"
INK = "102A43"
MUTED = "66788A"
LIGHT_BLUE = "EAF2F8"
PALE_GOLD = "FFF4CC"
PALE_RED = "FCE8E6"
PALE_GREEN = "E6F4EA"
BORDER = "CBD5E1"
RED = "9B1C1C"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def money(value):
    if value in (None, ""):
        return "—"
    return f"{float(value):,.2f}"


def qty(value):
    if value in (None, ""):
        return "—"
    return f"{float(value):,.2f}"


def set_run_font(run, size=None, color=None, bold=None, italic=None, name="Calibri"):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    node = trpr.find(qn("w:tblHeader"))
    if node is None:
        node = OxmlElement("w:tblHeader")
        trpr.append(node)
    node.set(qn("w:val"), "true")


def set_cell_text(cell, value, bold=False, color="000000", size=8.8):
    cell.text = "" if value is None else str(value)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.03
        for run in paragraph.runs:
            set_run_font(run, size=size, color=color, bold=bold)


def set_table_borders(table, color=BORDER, size="4"):
    tblpr = table._tbl.tblPr
    borders = tblpr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblpr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa, total=9360):
    table.autofit = False
    tbl = table._tbl
    tblpr = tbl.tblPr
    for tag in ("w:tblW", "w:tblInd", "w:tblLayout", "w:tblCellMar"):
        old = tblpr.find(qn(tag))
        if old is not None:
            tblpr.remove(old)
    tblw = OxmlElement("w:tblW")
    tblw.set(qn("w:w"), str(total))
    tblw.set(qn("w:type"), "dxa")
    tblpr.append(tblw)
    tblind = OxmlElement("w:tblInd")
    tblind.set(qn("w:w"), "120")
    tblind.set(qn("w:type"), "dxa")
    tblpr.append(tblind)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr.append(layout)
    margins = OxmlElement("w:tblCellMar")
    for tag, val in (("top", 65), ("bottom", 65), ("start", 85), ("end", 85)):
        node = OxmlElement(f"w:{tag}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    tblpr.append(margins)
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths_dxa):
            tcpr = cell._tc.get_or_add_tcPr()
            tcw = tcpr.find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                tcpr.append(tcw)
            tcw.set(qn("w:w"), str(width))
            tcw.set(qn("w:type"), "dxa")


def add_table(doc, headers, rows, widths, header_fill=LIGHT_BLUE, size=8.8, shades=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for idx, text in enumerate(headers):
        set_cell_text(table.cell(0, idx), text, bold=True, color=INK, size=size)
        set_cell_shading(table.cell(0, idx), header_fill)
    set_repeat_header(table.rows[0])
    for ridx, vals in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(vals):
            set_cell_text(cells[idx], value, size=size)
            if shades and ridx in shades:
                set_cell_shading(cells[idx], shades[ridx])
    set_table_geometry(table, widths)
    set_table_borders(table)
    return table


def add_callout(doc, title, text, fill=PALE_GOLD, title_color=INK):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, size=10.3, color=title_color, bold=True)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.1
    r2 = p2.add_run(text)
    set_run_font(r2, size=9.8, color="333333")
    set_table_geometry(table, [9360])
    set_table_borders(table, color=BORDER, size="6")
    set_repeat_header(table.rows[0])
    return table


def add_hyperlink(paragraph, text, url):
    rid = paragraph.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    for node in (
        ("w:rFonts", {"w:ascii": "Calibri", "w:hAnsi": "Calibri", "w:eastAsia": "Microsoft YaHei"}),
        ("w:color", {"w:val": BLUE}),
        ("w:u", {"w:val": "single"}),
    ):
        e = OxmlElement(node[0])
        for key, val in node[1].items():
            e.set(qn(key), val)
        rpr.append(e)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_source(doc, label, url, note=""):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    add_hyperlink(p, label, url)
    if note:
        r = p.add_run(f"：{note}")
        set_run_font(r, size=9.1, color=MUTED)


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.07
    r = p.add_run(text)
    set_run_font(r, size=10.0, color="222222")
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.07
    r = p.add_run(text)
    set_run_font(r, size=10.0, color="222222")
    return p


def add_para(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.12
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, size=10.1, color=INK, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, size=10.1, color="222222")
    else:
        r = p.add_run(text)
        set_run_font(r, size=10.1, color="222222")
    return p


def page_break(doc):
    doc.add_page_break()


doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.72)
sec.bottom_margin = Inches(0.72)
sec.left_margin = Inches(0.82)
sec.right_margin = Inches(0.82)

styles = doc.styles
styles["Normal"].font.name = "Calibri"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
styles["Normal"].font.size = Pt(10.1)
for name, size, color in (("Title", 26, NAVY), ("Heading 1", 16.5, NAVY), ("Heading 2", 13.2, BLUE), ("Heading 3", 11.3, INK)):
    style = styles[name]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(size)
    style.font.color.rgb = rgb(color)
    style.font.bold = True

header = sec.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hr = header.add_run("反倾销税风险穿透分析｜聚碳酸酯｜2026-08-13")
set_run_font(hr, size=8.3, color=MUTED)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = footer.add_run("阶段审计材料｜交易数据为核查线索，不构成违法认定")
set_run_font(fr, size=8.3, color=MUTED)

# 封面
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(36)
r = p.add_run("反倾销税风险穿透分析")
set_run_font(r, size=12, color=BLUE, bold=True)
p = doc.add_paragraph(style="Title")
p.paragraph_format.space_before = Pt(14)
p.paragraph_format.space_after = Pt(7)
r = p.add_run("聚碳酸酯（PC）")
set_run_font(r, size=26, color=NAVY, bold=True)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(18)
r = p.add_run("中国反倾销税、台湾来源经第三国转运及原产地申报风险阶段审计报告")
set_run_font(r, size=16, color=INK, bold=True)

add_table(doc, ["项目", "内容"], [
    ["受税来源", "台湾地区"],
    ["终裁实施", "2024年4月20日起，预计至2029年4月19日"],
    ["税则号", "39074000；措施范围还须满足双酚A型PC含量≥99%"],
    ["易迅窗口", "2024年8月6日至2026年6月30日（现有本地数据实际覆盖）"],
    ["逐票覆盖", "页面显示2,629次；剔除整页复制后2,429个槽位；精确唯一2,324条"],
    ["阶段状态", "未完成：缺约1页、关键词查询及台湾→越南A腿待补"],
    ["报告日期", "2026年8月13日"],
], [1700, 7660], size=9.3)

add_callout(doc, "结论摘要", "现有证据已能证明台湾来源聚碳酸酯/PC合金经越南向中国发运的具体B腿：17条记录货描保留#&TW，数量字段合计48,094；其中4条、数量字段合计3,294还引用越南前序进口申报或再出口信息。但越南侧主动保留台湾标记，既是跨境链线索，也可能是合规申报反证。由于尚未取得中国进口报关单和税款缴款书，不能认定中国端改报越南原产或逃缴反倾销税。", fill=PALE_GOLD)
add_callout(doc, "完整性警示", "旧抓取第6页与第7页200行逐行完全相同，日期从2025-08-16直接跳到2025-06-16，约缺一页。旧报告关于“2,629条全部读取且每页首行变化校验通过”的表述应废止。本报告只对当前可见2,324条唯一记录和232条备用缺页候选负责，且明确标为阶段审计。", fill=PALE_RED, title_color=RED)

page_break(doc)
doc.add_heading("一、核心判断", level=1)
add_bullet(doc, "第三国B腿已证实：越南对华记录中17条货描明确标注#&TW，存在可定位实体、牌号、日期和数量的台湾来源货物流向。")
add_bullet(doc, "逃税尚未证实：没有中国进口申报原产国、生产商栏、反倾销税税单，无法判断中国端是否按台湾原产正常纳税。")
add_bullet(doc, "产品范围不能按HS或“PC”名称整体判断：终裁只覆盖双酚A型PC按重量计含量达到99%的产品；PC/ABS、含较高比例颜料/填料的牌号通常排除。")
add_bullet(doc, "10条、数量字段合计41,044目前仍待TDS/COA；其中EMERGE 38,425、U415 2,319、PC-6715VT 300。平台JSON未单列数量单位，结合越南货描/包装高度可能为千克，但正式口径应以原始申报计量单位为准。")
add_bullet(doc, "同色号IC8800624和IC8800412合计1,300已有公开范围外反证：聚合页货描列PC 90—98%+TiO₂ 2—10%；该证据为非一手C级，需原始越南进口申报或TDS确认。")
add_bullet(doc, "公开官方来源没有发现措施实施后的聚碳酸酯专门反规避调查、原产地伪报处罚、走私判决或闭合双段提单案件。")

doc.add_heading("二、措施范围、税率和税种", level=1)
add_para(doc, "商务部公告2024年第13号自2024年4月20日起，对原产于台湾地区的聚碳酸酯征收反倾销税，期限5年。产品归入HS 39074000，但税号只作为入口：被调查产品须为分子主链含双酚A型碳酸酯基结构的高分子聚合物，且双酚A型PC按重量计含量达到99%。低于99%的产品排除。")
add_para(doc, "终裁还明确，医疗级、阻燃、着色等产品不能仅因“改性”而排除；添加剂比例很小且PC含量仍达到99%时仍可落入范围。相反，PC/ABS、PC/PET、PC/PBT、玻纤增强PC等多数会因PC比例不足99%而排除，最终仍以该票TDS、COA和配方为准。")
add_table(doc, ["生产商/税档", "反倾销税率", "含13%VAT增量的条件总系数", "本项目应用"], [
    ["台湾化纤、台湾出光", "9.0%", "10.17%×完税价格", "仅生产商栏确认后适用"],
    ["奇美实业、奇菱科技", "12.2%", "13.786%×完税价格", "PC-6715VT若为台湾奇美且≥99%"],
    ["其他台湾地区公司", "22.4%", "25.312%×完税价格", "Trinseo Taiwan符合范围的产品暂按此情景"],
], [2350, 1500, 2450, 3060], size=9.0)
add_callout(doc, "税种口径", "反倾销税=海关审定完税价格×企业税率；反倾销税进入进口增值税计税基础，因此漏征AD会连带产生AD×13%的进口VAT差额。不能把正常关税或基础进口VAT整体写成“逃税额”。", fill=LIGHT_BLUE)

page_break(doc)
doc.add_heading("三、易迅数据完整性和逐票方法", level=1)
add_table(doc, ["审计项目", "结果", "风险影响"], [
    ["旧页面显示", "2,629次，14页", "不能等同2,629条不同且完整的记录"],
    ["整页复制", "第6页与第7页200行逐行相同", "缺约2025-06-17至2025-08-15的一页"],
    ["有效页面槽位", "2,429", "剔除复制页后的可见出现数"],
    ["精确唯一", "2,324", "全部逐票进入阶段台账"],
    ["备用恢复候选", "232条，其中#&TW 3条/数量字段7,050", "只作重抓校验，不并入主计数"],
    ["关键词查询", "尚未完成", "需补POLYCARBONATE、PC RESIN、EMERGE、LEXAN、MAKROLON等"],
    ["台湾A腿", "尚未完成", "无法闭合台湾→越南→中国的同批双段链"],
], [2500, 2600, 4260], size=9.0)
add_para(doc, "2,324条唯一记录的目的地字段全部为China，日期均可解析，HS均为390740扩展码。当前分类是风险筛查而非海关归类：再生PC、改性/复合PC也不能自动排除，核心仍是双酚A型PC含量是否达到99%。")
scope = AUDIT["counts"]["scope_category"]
add_table(doc, ["逐票筛查分类", "唯一记录数", "判定边界"], [
    ["再生PC（含量与原产待核）", f"{scope.get('再生PC（含量与原产待核）',0):,}", "再生不等于范围外；须核PC含量和生产加工原产地"],
    ["改性/复合PC（含量待核）", f"{scope.get('改性/复合PC（含量待核）',0):,}", "添加剂<1%仍可能在范围"],
    ["基础/原生PC待核", f"{scope.get('基础/原生PC待核',0):,}", "优先取TDS/COA"],
    ["税号命中但货描不足", f"{scope.get('税号命中但货描不足',0):,}", "不可仅凭390740定性"],
    ["PC合金/共混料（倾向范围外）", f"{scope.get('PC合金/共混料（倾向范围外）',0):,}", "通常PC<99%，仍需配方"],
    ["疑似误匹配/其他", f"{scope.get('疑似误匹配/其他',0):,}", "货描与PC不符"],
], [3650, 1500, 4210], size=9.0)

doc.add_heading("四、台湾来源经越南对华B腿", level=1)
add_para(doc, "17条越南对华记录的货描均保留#&TW，平台产地/来源字段则显示Vietnam。两字段确有冲突，但不能据此写成“越南报关申报越南原产”：越南货描反而明确保留台湾标记。可安全表述为“平台字段冲突及台湾来源货物经越南对华B腿”。")
add_table(doc, ["越南主体", "中国方向主体", "记录", "数量字段", "主要牌号/研判"], [
    ["SIL-MORE Vietnam", "SIL-MORE INDUSTRIAL LIMITED", "7", "38,750", "EMERGE；其中1,300已有范围外反证，其余逐色号核TDS"],
    ["Goertek Vina", "GOERTEK (HONGKONG)", "2", "2,319", "U415；引用越南进口申报106821820600"],
    ["FortechVN", "KINGTECH (DONGGUAN)", "1", "925", "EMERGE 8130-6；引用越南进口申报106298628810"],
    ["Gu Vina", "HUIZHOU GALAXIA", "2", "4,500", "PC-6715VT 300；PC/ABS 4,200倾向范围外"],
    ["Nishoku Vietnam", "NISHOKU TECHNOLOGY", "5", "1,600", "仅EMERGE 50待核；其余PC/ABS或PPA范围外"],
], [2300, 2350, 900, 1250, 2560], size=8.7)

doc.add_heading("4.1 四条可按前序申报追A腿的记录", level=2)
add_table(doc, ["日期", "链路", "商品", "数量字段", "越南前序申报", "取证意义"], [
    ["2026-06-15", "Nishoku Vietnam→Nishoku", "EMERGE PC 8130ECO-13 IC7800746", "50", "107014946400", "可从进口申报追台湾发货人、批号和配方"],
    ["2026-01-27", "Goertek Vina→Goertek HK", "U415", "1,809", "106821820600", "同一进口申报项1清算/再出口"],
    ["2026-01-27", "Goertek Vina→Goertek HK", "U415", "510", "106821820600", "同一进口申报项2清算/再出口"],
    ["2025-10-14", "FortechVN→Kingtech Dongguan", "EMERGE 8130-6 IC1300102E", "925", "106298628810", "可反查2024-05-22进口批次"],
], [1250, 2100, 2500, 950, 1500, 1060], size=8.4)
add_callout(doc, "证据含义", "这4条、数量字段合计3,294已具备越南A段申报号，可核实台湾→越南的入境货物与越南→中国再出口是否为同批。但仍须叠加中国进口报关单和税款缴款书，才能判断中国端原产地和纳税。", fill=PALE_GOLD)

page_break(doc)
doc.add_heading("五、牌号范围复核", level=1)
add_table(doc, ["牌号组", "记录/数量字段", "现有证据", "当前结论"], [
    ["EMERGE 8130/8230", "9条/39,725", "Trinseo牌号明确；台湾新竹具着色/混配能力；同一基础牌号存在不同色号和添加剂", "不能按系列整体认定≥99%"],
    ["IC8800624、IC8800412", "2条/1,300", "非一手聚合页的完全同色号货描列PC90–98%+TiO₂2–10%", "已有范围外反证；调原申报/TDS确认"],
    ["其余EMERGE完整色号", "7条/38,425", "尚无逐色号一手成分", "范围待TDS/COA"],
    ["U415", "2条/2,319", "台湾标记及越南进口申报号明确；生产商/配方未锁定", "范围和税档均待查"],
    ["WONDERLITE PC-6715VT", "1条/300", "奇美官网确认透明阻燃PC牌号", "若台湾奇美且PC≥99%，适用12.2%"],
    ["PC/ABS合金", "4条/5,350", "货描明确PC/ABS", "通常PC<99%，倾向范围外"],
    ["PPA+50%GF", "1条/400", "非PC材料", "明确不属于本案商品"],
], [2300, 1400, 3450, 2210], size=8.8, shades={1: PALE_GREEN, 5: PALE_GREEN, 6: PALE_GREEN})
add_para(doc, "Trinseo官方资料证明EMERGE Advanced Resins可由PC与ABS、PET、色料或其他添加剂构成；“ECO”通常描述再生含量方案，并不直接等同非PC比例。因此必须匹配“完整牌号＋色号＋物料号”，再用该票TDS/COA确认99%门槛。")
add_para(doc, "奇美官方资料可把PC-6715VT较强地回连至WONDERLITE阻燃PC系列，但公开网页未直接给出双酚A型PC重量含量。其300数量字段仍应列为待核，而不是直接认定应税。")

doc.add_heading("六、条件税款暴露", level=1)
producer = SUMMARY["producer_specific_tax_scenarios"]
trinseo = next(x for x in producer if x["producer_scenario"].startswith("Trinseo"))
chimei = next(x for x in producer if x["producer_scenario"].startswith("奇美"))
umin = next(x for x in producer if x["producer_scenario"] == "待查" and x["ad_rate"] == 0.09)
umax = next(x for x in producer if x["producer_scenario"] == "待查" and x["ad_rate"] == 0.224)
low = trinseo["conditional_total"] + chimei["conditional_total"] + umin["conditional_total"]
high = trinseo["conditional_total"] + chimei["conditional_total"] + umax["conditional_total"]
add_table(doc, ["牌号/生产商情景", "数量字段", "金额字段代理", "AD率", "条件AD", "VAT增量", "条件合计"], [
    ["EMERGE/Trinseo Taiwan", qty(trinseo["quantity_field_sum"]), money(trinseo["amount_field_proxy"]), "22.4%", money(trinseo["conditional_ad"]), money(trinseo["conditional_vat_delta"]), money(trinseo["conditional_total"])],
    ["PC-6715VT/奇美", qty(chimei["quantity_field_sum"]), money(chimei["amount_field_proxy"]), "12.2%", money(chimei["conditional_ad"]), money(chimei["conditional_vat_delta"]), money(chimei["conditional_total"])],
    ["U415—最低税档情景", qty(umin["quantity_field_sum"]), money(umin["amount_field_proxy"]), "9.0%", money(umin["conditional_ad"]), money(umin["conditional_vat_delta"]), money(umin["conditional_total"])],
    ["U415—最高税档情景", qty(umax["quantity_field_sum"]), money(umax["amount_field_proxy"]), "22.4%", money(umax["conditional_ad"]), money(umax["conditional_vat_delta"]), money(umax["conditional_total"])],
    ["10条范围待核合计", "41,044", "5,999,245,916.15", "按牌号推定", "—", "—", f"{money(low)}—{money(high)}"],
], [2050, 1100, 1650, 1000, 1200, 1160, 1200], size=8.1)
add_callout(doc, "不可省略的前提", "上述只是同金额字段币种的风险上限情景，不是欠税额。它同时假设：10条均满足PC≥99%、生产商/税档推定正确、金额字段等同中国海关完税价格、且中国端完全未征反倾销税。任一条件不成立，数值都须调整或归零。平台本地JSON未单列币种，报告不得把这些金额写成人民币或美元。", fill=PALE_RED, title_color=RED)

page_break(doc)
doc.add_heading("七、公开来源中的第三国链和实体", level=1)
doc.add_heading("7.1 商务部确认的多层商业链", level=2)
add_para(doc, "终裁记载，台湾化纤—台湾出光体系在调查期存在直接对大陆销售、经台湾贸易商、经第三国关联贸易商、经两层第三国关联贸易商及再经上海关联贸易商销售等多种路径；调查机关还对上海出光化学有限公司开展实地核查。这是A级的商业、开票和关联贸易网络证据。")
add_para(doc, "该事实发生在措施前调查期并已向商务部披露、纳入倾销计算，不能推导为货物物理经过第三国，也不能推导为2024年后的逃税。它的价值在于锁定应优先核查的关联主体、合同、发票、付款和转售路径。")
doc.add_heading("7.2 Trinseo Taiwan—EMERGE", level=2)
add_para(doc, "Trinseo官方目录、认证页面和年报确认台湾新竹存在工程塑料着色、混配和共混业务，并涉及EMERGE PC、PC/ABS及含再生料牌号。终裁同时明确，即便PC基料并非台湾聚合，只要在台湾制成且符合被调查产品描述，仍可能属于措施范围。")
add_para(doc, "因此，“台湾没有本地PC聚合基地”不能自动排除台湾原产风险；但混配牌号是否达到99%仍须逐票配方证明。")

doc.add_heading("八、证据分级和风险结论", level=1)
add_table(doc, ["层级", "当前证据", "可写结论", "不可越界"], [
    ["A级结构证据", "商务部终裁确认多层第三国关联贸易及上海转售路径", "相关集团具成熟跨境商业链", "不能说物理转运或违法"],
    ["B/B+具体B腿", "17条#&TW；4条有越南前序进口申报号", "台湾来源PC经越南向中国发运线索已成立", "不能说中国端改报产地"],
    ["C/C+范围反证", "同色号公开货描PC90–98%+TiO₂", "1,300数量字段当前倾向范围外", "非一手资料不能直接定案"],
    ["未闭合", "缺中国报关、税单、COA、箱号和完整A腿", "逃税、少缴税均未证实", "不得写“已查实绕道逃税”"],
], [1700, 2850, 2450, 2360], size=8.8)
add_callout(doc, "总体风险评级", "第三国再出口/原产地核查风险：中高；反倾销税实际少缴风险：待中国端税单核定；公开违法证据：未发现。最优先对象是30,000数量字段的EMERGE PC 8230ECO-20 IC8800778、U415 2,319、IC7800746合计3,500、IC1300102E合计2,925及PC-6715VT 300。", fill=PALE_GOLD)

doc.add_heading("九、下一步查询与取证顺序", level=1)
for text in [
    "重新抓取易迅HS390740实时第7页，核实2025-06-17至2025-08-15缺口，并用232条备用记录逐行校验。",
    "对中国方向完成POLYCARBONATE、PC RESIN、EMERGE、LEXAN、MAKROLON等关键词全页查询；裸关键词PC误匹配过多，不单独使用。",
    "围绕台湾→越南反查HS390740、EMERGE/Trinseo、完整色号、SIL-MORE及越南E11申报号，建立台湾A腿。",
    "优先调中国进口报关单：原产国、启运国、生产商、境外发货人、规格型号、贸易方式、口岸和境内收货人。",
    "调税款缴款书核是否征收反倾销税及采用9%、12.2%或22.4%税率。",
    "调完整牌号TDS/COA，确认双酚A型PC重量含量是否达到99%。",
    "调越南原进口及再出口申报、库存核销、提单/箱号/封志，按日期、数量、色号和批号闭合物理链。",
]:
    add_number(doc, text)

page_break(doc)
doc.add_heading("附录一：20条#&TW主线及缺页备用记录", level=1)
lead_rows = []
for row in LEADS:
    status = "主" if row["dataset_status"].startswith("主") else "备用"
    desc = row["description"]
    if len(desc) > 92:
        desc = desc[:89] + "…"
    pre = ";".join(row["pre_entry_numbers"]) if isinstance(row["pre_entry_numbers"], list) else str(row["pre_entry_numbers"] or "")
    lead_rows.append([
        status,
        row.get("date", ""),
        desc,
        row.get("foreign_party", ""),
        qty(row.get("quantity")),
        row.get("final_scope_screen", ""),
        pre or "—",
    ])
add_table(doc, ["状态", "日期", "商品/色号", "越南主体", "数量字段", "范围筛查", "前序申报"], lead_rows, [650, 1000, 3150, 1750, 900, 1100, 810], size=7.6)

doc.add_heading("附录二：主要公开来源", level=1)
sources = [
    ("商务部公告2024年第13号及最终裁定", "https://www.mofcom.gov.cn/zcfb/zgdwjjmywg/art/2024/art_ead5e0f4810c4a6a9fd02a86e72421dd.html", "措施范围、税率、期限、调查实体和第三国商业销售路径"),
    ("中国贸易救济信息网措施状态", "https://cacs.mofcom.gov.cn/cacscms/article/sszaj?articleId=160267&type=", "实施状态与案件节点"),
    ("Trinseo全球目录", "https://www.trinseo.com/global-directory", "台湾主体及工厂"),
    ("Trinseo新竹认证页面", "https://www.trinseo.com/company/quality-environment-health-safety-and-certifications/trinseo-management-system-standards-and-iso-certifications", "EMERGE PC/PC-ABS和再生含量牌号验证"),
    ("Trinseo 2024 Form 10-K", "https://www.sec.gov/Archives/edgar/data/1519061/000155837025001736/tse-20241231x10k.htm", "新竹厂compounds & blends职能"),
    ("Trinseo EMERGE官方说明", "https://www.trinseo.com/solutions/polycarbonate/emerge", "EMERGE可由PC与其他聚合物/添加剂组成"),
    ("奇美PC-6715VT官方产品页", "https://www.chimeicorp.com/zh-CN/?feature=T10503&material=T10500&page=products&series=T10000", "WONDERLITE阻燃PC牌号"),
    ("Volza同色号公开货描", "https://www.volza.com/p/titanium-dioxide/export/hsn-code-39074000/", "IC8800624/IC8800412的C级范围反证；非一手，需回查原始申报"),
]
for label, url, note in sources:
    add_source(doc, label, url, note)

doc.add_heading("附录三：交付文件", level=1)
for filename, note in [
    ("聚碳酸酯_易迅逐票判定台账_阶段审计.xlsx", "含概览、20条TW线索、2,324条唯一记录、232条缺页恢复候选和政策取证口径"),
    ("聚碳酸酯_易迅_HS390740_逐票标准化.csv/json", "当前可见唯一记录逐票底稿"),
    ("聚碳酸酯_TW线索人工复核.csv/json", "主线与缺页备用记录的范围、实体、税档和取证建议"),
    ("聚碳酸酯_易迅_HS390740_全量分析结果.json", "分页完整性、分类和路线统计"),
]:
    add_bullet(doc, f"{filename}：{note}")

add_callout(doc, "报告状态", "本文件是阶段审计报告，不是本商品最终深度报告。完成缺页重抓、商品关键词全页查询、台湾A腿检索和逐票合并后，应覆盖更新为最终报告，并同步更新总报告、总清单和前端系统。", fill=PALE_RED, title_color=RED)

doc.core_properties.title = "聚碳酸酯反倾销税与第三国转运风险阶段审计报告"
doc.core_properties.subject = "易迅数据逐票审计、公开来源核验、台湾来源经越南B腿"
doc.core_properties.author = "Codex"
doc.core_properties.keywords = "聚碳酸酯, 反倾销税, 台湾, 越南, 第三国转运, 原产地"
doc.save(OUT)
print(OUT)
