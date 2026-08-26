"""Enrich repeated intelligence vehicle leads with GDRM26-269 trade history."""
from __future__ import annotations

import json
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(r"D:\codex\XG执法\2026_案例库")
analysis = json.loads((ROOT / "情报案件风险分析.json").read_text(encoding="utf-8"))
titles = {item["case_id"]: item["title"] for item in analysis["case_details"]}
vehicle_items = [
    item for item in analysis["repeated_identifiers"]
    if item["group"] == "车辆/车牌" and "GDRM26-269" in item["case_ids"]
]
targets = {item["identifier"] for item in vehicle_items}

workbooks = list((ROOT / "raw_cases" / "GDRM26-269").rglob("*.xlsx"))
if not workbooks:
    raise FileNotFoundError("GDRM26-269 trade workbook not found")
book = load_workbook(workbooks[0], read_only=True, data_only=True)

records = defaultdict(list)
for sheet in book.worksheets:
    rows = sheet.iter_rows(values_only=True)
    header = None
    index = {}
    for row in rows:
        values = [str(value).strip() if value is not None else "" for value in row]
        if header is None:
            if any("车牌号码" in value or "車牌號碼" in value for value in values):
                header = values
                for col, value in enumerate(values):
                    if "车牌号码" in value or "車牌號碼" in value:
                        index["vehicle"] = col
                    elif "申报日期" in value or "申報日期" in value:
                        index["date"] = col
                    elif "进口" in value and "出口" in value or "進口" in value and "出口" in value:
                        index["direction"] = col
                    elif "发货人名称" in value or "發貨人名稱" in value:
                        index["shipper"] = col
                    elif "收货人名称" in value or "收貨人名稱" in value:
                        index["consignee"] = col
                    elif "申报人名称" in value or "申報人名稱" in value:
                        index["declarant"] = col
                    elif "报关货物名称" in value or "報關貨物名稱" in value:
                        index["goods"] = col
                    elif "目的地国家" in value or "目的地國家" in value:
                        index["destination_country"] = col
                    elif "来源地国家" in value or "來源地國家" in value:
                        index["origin_country"] = col
                continue
            continue
        if "vehicle" not in index or index["vehicle"] >= len(values):
            continue
        vehicle = values[index["vehicle"]].replace(" ", "").upper()
        if vehicle not in targets:
            continue
        record = {}
        for key, col in index.items():
            record[key] = values[col] if col < len(values) else ""
        records[vehicle].append(record)


def top_values(rows, key, limit=5):
    counts = Counter(row.get(key, "") for row in rows if row.get(key, ""))
    return [{"value": value, "count": count} for value, count in counts.most_common(limit)]


def dates(rows):
    found = []
    for row in rows:
        value = row.get("date", "")
        if not value:
            continue
        found.append(value)
    return min(found) if found else "", max(found) if found else ""


leads = []
for item in vehicle_items:
    vehicle = item["identifier"]
    rows = records.get(vehicle, [])
    start, end = dates(rows)
    incident_ids = [case_id for case_id in item["case_ids"] if case_id != "GDRM26-269"]
    leads.append({
        "vehicle": vehicle,
        "priority": "A" if len(incident_ids) >= 2 else "B",
        "linked_intelligence_cases": incident_ids,
        "linked_case_titles": [titles.get(case_id, "") for case_id in incident_ids],
        "silver_trade_record_count": len(rows),
        "silver_trade_date_min": start,
        "silver_trade_date_max": end,
        "directions": top_values(rows, "direction"),
        "goods_top": top_values(rows, "goods"),
        "shippers_top": top_values(rows, "shipper"),
        "consignees_top": top_values(rows, "consignee"),
        "declarants_top": top_values(rows, "declarant"),
        "origin_countries_top": top_values(rows, "origin_country"),
        "destination_countries_top": top_values(rows, "destination_country"),
        "source_evidence": item["evidence"],
    })
leads.sort(key=lambda item: (item["priority"], -len(item["linked_intelligence_cases"]), -item["silver_trade_record_count"], item["vehicle"]))
(ROOT / "情报车辆核查对象.json").write_text(json.dumps(leads, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(leads, ensure_ascii=True, indent=2))
