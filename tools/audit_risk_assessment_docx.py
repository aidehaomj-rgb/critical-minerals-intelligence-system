"""Structural preset audit for the generated risk-assessment DOCX."""

import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


DOCX = Path(r"D:\codex\XG执法\2026_案例库\情报案件解密后详细风险研判与布控验证清单.docx")
doc = Document(DOCX)
errors = []

section = doc.sections[0]
expected = {
    "page_width": 7772400,
    "page_height": 10058400,
    "left_margin": 914400,
    "right_margin": 914400,
    "top_margin": 914400,
    "bottom_margin": 914400,
}
for key, value in expected.items():
    actual = int(getattr(section, key))
    if actual != value:
        errors.append(f"{key}: expected {value}, got {actual}")

heading_counts = {
    style: sum(p.style.name == style for p in doc.paragraphs)
    for style in ("Heading 1", "Heading 2", "Heading 3")
}

numbered = 0
fake_lists = []
for paragraph in doc.paragraphs:
    num_pr = paragraph._p.pPr.numPr if paragraph._p.pPr is not None else None
    if num_pr is not None:
        numbered += 1
    text = paragraph.text.strip()
    if paragraph.style.name == "Normal" and (
        text.startswith(("- ", "• ", "● ")) or re.match(r"^\d+\.\s", text)
    ):
        fake_lists.append(text[:80])

for index, table in enumerate(doc.tables, 1):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    grid_widths = [
        int(col.get(qn("w:w"))) for col in table._tbl.tblGrid.findall(qn("w:gridCol"))
    ]
    if tbl_w is None or tbl_w.get(qn("w:w")) != "9360" or tbl_w.get(qn("w:type")) != "dxa":
        errors.append(f"table {index}: invalid tblW")
    if tbl_ind is None or tbl_ind.get(qn("w:w")) != "120":
        errors.append(f"table {index}: invalid tblInd")
    if sum(grid_widths) != 9360:
        errors.append(f"table {index}: grid sum {sum(grid_widths)}")
    for row_index, row in enumerate(table.rows, 1):
        if row._tr.trPr is not None and row._tr.trPr.find(qn("w:trHeight")) is not None:
            errors.append(f"table {index} row {row_index}: fixed/explicit height present")
        for col_index, cell in enumerate(row.cells):
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None or int(tc_w.get(qn("w:w"))) != grid_widths[col_index]:
                errors.append(f"table {index} row {row_index} cell {col_index + 1}: tcW mismatch")

all_text = "\n".join(p.text for p in doc.paragraphs)
for anchor in (
    "吴子俊 / NG TSZ CHUN",
    "XR4740",
    "深圳润丰成—香港潤豐链路",
    "LINEHAUL EXPRESS LLC",
    "建议的布控验证规则",
    "使用边界",
):
    if anchor not in all_text:
        errors.append(f"missing anchor: {anchor}")

print(f"file={DOCX}")
print(f"paragraphs={len(doc.paragraphs)}")
print(f"tables={len(doc.tables)}")
print(f"heading_counts={heading_counts}")
print(f"real_numbered_list_paragraphs={numbered}")
print(f"fake_lists={fake_lists}")
print(f"errors={errors}")
if errors or fake_lists:
    raise SystemExit(1)
