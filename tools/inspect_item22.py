from pathlib import Path
from docx import Document
import csv, json
b=Path(r'D:\易迅数据\反倾销税深度分析报告\01_卤化丁基橡胶')
d=Document(b/'卤化丁基橡胶_反倾销税与第三国转运风险深度分析报告.docx')
for p in d.paragraphs:
    if p.text.strip(): print(p.text)
print('\nTABLE HEADERS')
for i,t in enumerate(d.tables): print(i,[c.text for c in t.rows[0].cells],len(t.rows))
with (b/'卤化丁基橡胶_易迅逐票判定台账_合并去重.csv').open(encoding='utf-8-sig',newline='') as f:
 r=csv.reader(f); h=next(r); print('\nCSV fields',len(h),h)
 for _ in range(2): print(next(r))
