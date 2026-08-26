"""Mine repeat people/company candidates from case-library spreadsheets.

This deliberately produces review leads, not findings of liability.
"""
from __future__ import annotations

import json
import re
import sqlite3
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(r"D:\codex\XG执法\2026_案例库")
DB = ROOT / "enforcement_cases_2026.sqlite"
CN_COMPANY = re.compile(r"[\u3400-\u9fff]{3,40}(?:有限责任公司|股份有限公司|有限公司|公司|企业)")
EN_COMPANY = re.compile(r"\b[A-Z][A-Z0-9&'.,\- ]{3,100}(?:CO\.?\s*,?\s*LTD\.?|CO\.?\s*,?\s*LIMITED|LIMITED|LTD\.?|LLC|INC\.?|CORP\.?|CORPORATION|COMPANY)\b", re.I)
EXCLUDE = re.compile(r"^(?:N/?A|NA|NONE|UNKNOWN|ADDRESS|CHINA|HONG KONG|HK|CN|I|O)$", re.I)
ROLE_TOKENS = {
    "发货/出口方": ("寄件", "發貨", "发货", "出口", "shipper", "sender", "consignor"),
    "收货/进口方": ("收件", "收貨", "收货", "进口", "import", "consignee", "receiver"),
    "申报/报关方": ("報關", "报关", "申報", "申报", "declar"),
    "承运/货代": ("承運", "承运", "船公司", "航空", "shipping", "carrier", "forwarder"),
    "其他主体": (),
}

def norm(value: str) -> str:
    text = str(value or "").upper()
    text = re.sub(r"\s+", " ", text).strip()
    text = re.sub(r"[，,。.;；:：'\"()（）\[\]【】]", "", text)
    return text

def role_for_header(header: str) -> str:
    folded = header.casefold()
    for role, tokens in ROLE_TOKENS.items():
        if any(token.casefold() in folded for token in tokens):
            return role
    return "其他主体"

def candidates(value: object) -> list[str]:
    raw = str(value or "").replace("\n", " ").strip()
    if not raw or len(raw) < 3 or EXCLUDE.match(raw):
        return []
    found = []
    found.extend(CN_COMPANY.findall(raw))
    found.extend(EN_COMPANY.findall(raw.upper()))
    # Company names in many shipping sheets precede an address. Preserve an obvious first line.
    if not found and re.match(r"^[A-Z][A-Z0-9 &'.,\-]{3,80}$", raw.upper()) and any(t in raw.upper() for t in ("LOGISTICS", "TRADING", "SHIPPING", "FREIGHT", "INTERNATIONAL")):
        found.append(raw)
    clean = []
    for item in found:
        candidate = norm(item)
        if candidate in {"有限公司", "有限责任公司", "股份有限公司", "COMPANY", "LIMITED", "CO LTD", "CO LIMITED", "COMPANY LIMITED", "COMPANY LTD", "PTE LTD", "PVT LTD", "COLTD"}:
            continue
        clean.append(candidate)
    return clean

def sheet_rows(path: Path):
    try:
        book = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return
    for sheet in book.worksheets:
        rows = sheet.iter_rows(values_only=True)
        header = None
        for index, row in enumerate(rows, 1):
            values = [str(x).strip() if x is not None else "" for x in row]
            joined = " ".join(values)
            if header is None and index <= 12 and sum(bool(x) for x in values) >= 4 and any(token in joined for token in ("公司", "公司", "寄件", "收件", "发货", "發貨", "运输", "運輸", "申报", "申報", "地址", "货物", "貨物")):
                header = values
                continue
            if header is None or not any(values):
                continue
            yield sheet.title, header, values

conn = sqlite3.connect(DB)
xlsx_rows = conn.execute("SELECT DISTINCT case_id, relative_path FROM documents WHERE extension='.xlsx'").fetchall()
conn.close()

entities: dict[str, dict] = {}
for case_id, relative_path in xlsx_rows:
    path = ROOT / relative_path
    for sheet, header, values in sheet_rows(path) or []:
        for col, value in enumerate(values):
            if col >= len(header):
                continue
            header_name = header[col]
            role = role_for_header(header_name)
            for name in candidates(value):
                record = entities.setdefault(name, {"case_ids": set(), "roles": set(), "evidence": []})
                record["case_ids"].add(case_id)
                record["roles"].add(role)
                if len(record["evidence"]) < 8:
                    record["evidence"].append({"case_id": case_id, "file": relative_path, "sheet": sheet, "field": header_name, "value": str(value)[:260]})

results = []
for name, record in entities.items():
    # Exact normalized-name matches across distinct case packs only.
    if len(record["case_ids"]) >= 2:
        results.append({"entity": name, "case_count": len(record["case_ids"]), "case_ids": sorted(record["case_ids"]), "roles": sorted(record["roles"]), "evidence": record["evidence"]})
results.sort(key=lambda item: (-item["case_count"], item["entity"]))

# A repeated match is more useful for follow-up when it occurs in a party/declarant column,
# rather than only as a carrier/terminal mentioned in a shipment record.
operational = [item for item in results if set(item["roles"]) & {"发货/出口方", "收货/进口方", "申报/报关方"}]
logistics = [item for item in results if item not in operational]

output = {
    "method": "Exact normalized organization-name matches in XLSX fields, counted across distinct case packs. Entries are investigative leads only.",
    "candidate_count": len(results),
    "operational_subjects": operational[:100],
    "logistics_or_other_entities": logistics[:100],
}
(ROOT / "重复主体候选.json").write_text(json.dumps(output, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(output, ensure_ascii=True, indent=2))
