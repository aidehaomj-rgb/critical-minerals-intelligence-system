#!/usr/bin/env python3
"""Build a stage report for item 16 stainless billet / hot-rolled products."""

from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(r"D:\易迅数据\反倾销税深度分析报告\16_不锈钢钢坯和热轧板卷")
DOCX = OUT / "不锈钢钢坯和热轧板卷_反倾销与第三国转运风险_数据缺口阶段报告.docx"
NAVY = RGBColor(11, 37, 69); BLUE = RGBColor(46, 116, 181); GRAY = RGBColor(90, 98, 108); RED = RGBColor(155, 28, 28)


def set_font(run, size=11, bold=False, color=None):
    run.font.name = "Calibri"; run._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
    run.font.size = Pt(size); run.bold = bold
    if color: run.font.color.rgb = color


def cell_shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); tcPr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None: tcMar = OxmlElement("w:tcMar"); tcPr.append(tcMar)
    for m, v in (("top",top),("start",start),("bottom",bottom),("end",end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None: node=OxmlElement(f"w:{m}"); tcMar.append(node)
        node.set(qn("w:w"),str(v)); node.set(qn("w:type"),"dxa")


def add_table(doc, headers, rows, widths):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    trPr = t.rows[0]._tr.get_or_add_trPr(); tbl_header = OxmlElement("w:tblHeader"); tbl_header.set(qn("w:val"), "true"); trPr.append(tbl_header)
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.width=Inches(widths[i]); cell_shade(c,"F2F4F7"); set_cell_margins(c)
        p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; set_font(p.add_run(h),10,bold=True,color=NAVY)
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row):
            cells[i].width=Inches(widths[i]); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_cell_margins(cells[i])
            p=cells[i].paragraphs[0]; p.paragraph_format.space_after=Pt(0); set_font(p.add_run(str(v)),9.5)
    return t


def bullet(doc, text):
    p=doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after=Pt(4); set_font(p.add_run(text),10.5); return p


def link_para(doc, label, url):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(3); set_font(p.add_run(label+": "),9,bold=True,color=GRAY); set_font(p.add_run(url),9,color=BLUE)


def main():
    OUT.mkdir(parents=True, exist_ok=True); doc=Document(); sec=doc.sections[0]
    sec.page_width=Inches(8.5); sec.page_height=Inches(11); sec.top_margin=sec.bottom_margin=Inches(1); sec.left_margin=sec.right_margin=Inches(1)
    styles=doc.styles; normal=styles["Normal"]; normal.font.name="Calibri"; normal._element.rPr.rFonts.set(qn("w:eastAsia"),"等线"); normal.font.size=Pt(11)
    normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.10
    for name,size,before,after,color in (("Heading 1",16,16,8,BLUE),("Heading 2",13,12,6,BLUE),("Heading 3",12,8,4,NAVY)):
        s=styles[name]; s.font.name="Calibri"; s._element.rPr.rFonts.set(qn("w:eastAsia"),"等线"); s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=color; s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after)
    h=sec.header.paragraphs[0]; set_font(h.add_run("反倾销税深度分析 | 第16项"),9,bold=True,color=GRAY)
    f=sec.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.RIGHT; set_font(f.add_run("阶段审计报告 | 2026-08-13"),9,color=GRAY)
    p=doc.add_paragraph(); set_font(p.add_run("风险核查备忘录"),10,bold=True,color=BLUE)
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(4); set_font(p.add_run("不锈钢钢坯和不锈钢热轧板/卷"),23,bold=True,color=NAVY)
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(14); set_font(p.add_run("中国反倾销措施、第三国加工链与易迅数据缺口阶段审计"),13,color=GRAY)
    for label,value in (("对象","原产欧盟、英国、韩国、印度尼西亚的涉案产品"),("审计日期","2026-08-13"),("本地数据范围","D:\\易迅数据中288个XLSX/XLS/CSV/JSON文件内容级预筛"),("阶段结论","本地未发现可确认的第16项原始易迅底表；公开资料有具体第三国加工链，但未指向中国逃税")):
        p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(2); set_font(p.add_run(label+"："),10.5,bold=True); set_font(p.add_run(value),10.5)

    doc.add_heading("一、结论先行", level=1)
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(8); set_font(p.add_run("当前没有证据证明受税来源产品经第三国进入中国并逃避反倾销税。"),12,bold=True,color=RED)
    bullet(doc,"本地全盘扫描288个可解析文件，0个文件命中22个涉案税号；27个宽词候选主要为其他商品货描中的不锈钢包装、紧固件、丝材、带材或项目元数据。")
    bullet(doc,"公开最具体链条是“印度尼西亚不锈钢板坯—土耳其Çolakoğlu代工热轧—欧盟”，物流、数量和实体由欧盟调查资料支持，但该链没有公开证据显示进入中国。")
    bullet(doc,"土耳其真热轧可能造成四位税目由7218变为7219/7220，具备实质加工的合法解释；仅酸洗、退火、纵剪、切板且税目不变则风险更高，必须以中国非优惠原产地核定和生产底账判断。")
    bullet(doc,"易迅页面会话已登录，但当前查询页持续卡在“正在搜索中”，筛选条件未能稳定验证；本报告不得把页面“暂无数据”记录为零结果。")

    doc.add_heading("二、现行政策与税率", level=1)
    p=doc.add_paragraph("商务部2025年第33号公告自2025年7月1日起继续实施5年，正常到期日2030年6月30日。日本措施已于2024年7月23日届满，现行受税来源为欧盟、英国、韩国和印度尼西亚。")
    add_table(doc,["来源/企业","反倾销税率","潜在少缴系数*"],[
        ("欧盟全部","43.0%","48.590% × 完税价格"),("英国全部","43.0%","48.590% × 完税价格"),("韩国POSCO","23.1%","26.103% × 完税价格"),("韩国其他","103.1%","116.503% × 完税价格"),("印度尼西亚全部","20.2%","22.826% × 完税价格")],[2.55,1.35,2.60])
    p=doc.add_paragraph("* 仅测反倾销税及其引致的13%进口增值税增量：V×AD税率×1.13；不含正常关税、正常进口增值税、滞纳金及罚款。", style=None); p.paragraph_format.space_before=Pt(4); p.paragraph_format.space_after=Pt(8); set_font(p.runs[0],8.5,color=GRAY)
    p=doc.add_paragraph("范围：非冷轧，碳含量不超过1.2%、铬含量不低于10.5%的不锈钢钢坯及不锈钢热轧板/卷。钢坯包括矩形（不含正方形）及其他半成品；同税号下其他产品不当然纳入。")
    hs="72189100、72189900、72191100、72191210、72191290、72191312、72191319、72191322、72191329、72191412、72191419、72191422、72191429、72192100、72192200、72192300、72192410、72192420、72192430、72201100、72201200、72223000。"
    p=doc.add_paragraph(); set_font(p.add_run("22个中国税号："),10.5,bold=True); set_font(p.add_run(hs),10.5)

    doc.add_heading("三、POSCO价格承诺：不能按企业名称直接免税", level=1)
    bullet(doc,"公开承诺产品仅限POSCO对华出口的特定400系涉案产品；300系产品不能仅凭“POSCO”认定适用承诺。")
    bullet(doc,"加工贸易手册进口、进口至保税区或保税仓库的涉案产品，公开承诺明确不适用。")
    bullet(doc,"免税前提包括有效商业发票、POSCO承诺证明信、最低进口限价合规；关联中国公司转售还需核首次转售净价。")
    bullet(doc,"承诺明确禁止经第三国转运或再出口隐瞒品名、出口商，以及故意误报归类、原产地或出口商身份；违反时按23.1%征税。")

    doc.add_heading("四、公开第三国链：事实、法律状态与边界", level=1)
    add_table(doc,["层级","已核事实","中国风险含义"],[
        ("物流/实体事实","欧委会2023/825认定印尼板坯由Marcegaglia Specialties控制采购，交Çolakoğlu在土耳其热轧后进入欧盟；印尼→土耳其板坯由2018年0吨升至调查期约40,513吨。","A等级链路事实，但只证明进入欧盟，不证明进入中国或逃中国税。"),
        ("司法状态","欧盟普通法院2026-03-04就Çolakoğlu撤销条例；欧委会2026-05-13上诉，至审计日待决。","判决未否定物流/代工事实，争议在欧盟法反规避构成。"),
        ("中国原产地","多国生产以最后实质性改变为准；规避贸易救济而实施的加工可不予考虑。","7218板坯真热轧为7219/7220可能改变税目；需中国海关按配方、工序和目的核定。")
    ],[1.15,3.10,2.25])
    doc.add_heading("五、本地易迅数据盘点", level=1)
    add_table(doc,["审计项","结果","解释"],[
        ("文件总数","288","XLSX/XLS/CSV/JSON；排除本项输出目录自身。"),("命中涉案税号的文件","0","未发现可直接作为第16项底表的本地下载。"),("宽词候选文件","27","25个仅宽词、2个为主清单/进度台账商品名。"),("解析错误","0","文件级预筛完成。"),("易迅网页","未形成有效零结果","页面持续“正在搜索中”，筛选条件与结果总数未稳定确认。")
    ],[1.65,1.40,3.45])
    p=doc.add_paragraph("审计边界：上述结果只说明当前D盘缺少该商品的可确认原始数据，不表示易迅全库没有贸易，更不能表示没有第三国风险。")

    doc.add_heading("六、下一轮易迅最简查询矩阵", level=1)
    add_table(doc,["查询","条件","用途"],[
        ("Q1 中国进口底表","目的国中国；2024-07-23至最新；22税号分7218/7219/7220三组；关键词STAINLESS STEEL BILLET/SLAB/BLOOM/SEMI-FINISHED/HOT ROLLED/SSHR/HRC/HR COIL/HOT ROLLED PLATE/NO.1/1D及中文同义词。","建立全部B腿；冷热轧、宽窄、制品逐票分层。"),
        ("Q2 受税国→第三国A腿","印尼→土耳其/越南/台湾/马来西亚；韩国→越南/泰国/马来西亚；欧英→土耳其/马来西亚；同税号与关键词。","寻找原始板坯/热轧料输入。"),
        ("Q3 第三国→中国B腿","上述第三国→中国；同税号/关键词；企业名精确补查。","按15—180日、炉号/卷号/规格重量/提单闭环。")
    ],[1.35,3.55,1.60])
    doc.add_heading("七、重点实体与逐票闭环字段", level=1)
    bullet(doc,"首查实体：Marcegaglia Specialties、Çolakoğlu Metalurji；PT Indonesia Guang Ching、PT Obsidian/OSS、Indonesia Tsingshan/ITSS、SMI、Indonesia Ruipu；POSCO/POSCO International、POSCO-VST、POSCO Assan；Acerinox/Bahru Stainless；Outokumpu、Aperam、Industeel；Saritas、Üças、AST Turkey Metal、Yongjin Vietnam、Lam Khang。")
    bullet(doc,"货物指纹：钢种304/304L/316L/430/duplex，炉号、板坯号、卷号，厚度×宽度×长度，净重，表面状态No.1/1D/black/white，MTC化学成分与生产厂代码。")
    bullet(doc,"运输与商业链：提单号、船名/IMO/航次、中转港、集装箱/散杂货批次、买卖方、通知方、加工合同、发票、付款受益人、第三国加工费。")
    bullet(doc,"生产与原产：热轧线工单、加热炉/轧机/退火酸洗线记录、能耗、投入产出平衡、四位税目变化、中国海关原产地预裁定/核定、原产地证申请底稿。")
    bullet(doc,"中国核税：进口报关单、生产商、贸易方式、保税状态、POSCO商业发票/证明信/MIP、反倾销税缴款书、进口增值税缴款书、申报口岸和报关企业。")

    doc.add_heading("八、证据分级与阶段判断", level=1)
    bullet(doc,"A级（链路事实）：印尼板坯—土耳其Çolakoğlu代工—欧盟，属于境外已核实体链；不等于中国违法证据。")
    bullet(doc,"B级：中国端尚无同炉号/卷号、同提单或同批A/B腿。")
    bullet(doc,"C级：第三国加工/集团网络/宏观流量只作筛选器；真实热轧与本地产能必须作为反证同步核查。")
    bullet(doc,"结论：本项当前为“公开具体链存在、本地易迅底表缺失、对华绕道无法判断”。待Q1—Q3全页数据取得后再给数量、企业、口岸及条件税差。")

    doc.add_heading("九、主要公开来源", level=1)
    link_para(doc,"商务部2025年第33号公告","https://www.mofcom.gov.cn/zcfb/zc/art/2025/art_3c3f7a5f2bb14c1b900a3931f9d2212b.html")
    link_para(doc,"欧委会2023/825","https://eur-lex.europa.eu/legal-content/EN/TXT/?uri=CELEX:32023R0825")
    link_para(doc,"欧盟普通法院T-379/23","https://eur-lex.europa.eu/legal-content/EN/TXT/PDF/?uri=celex:62023TJ0379")
    link_para(doc,"欧委会上诉C-485/26 P","https://juris.curia.europa.eu/juris/document/document.jsf?docid=313369&doclang=en")
    link_para(doc,"中国实质性改变标准规定","https://policy.mofcom.gov.cn/claw/clawContent.shtml?id=101350")
    link_para(doc,"中国进出口货物原产地条例","https://xzfg.moj.gov.cn/front/law/detail?LawID=1523")
    link_para(doc,"POSCO价格承诺公开文本","https://www.tid.gov.hk/english/aboutus/tradecircular/cic/asia/2019/files/ci2019558a.pdf")
    doc.save(DOCX); print(DOCX)

if __name__ == "__main__": main()
