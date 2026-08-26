from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(r"D:\codex\管制反制\汇报大纲_完善版.docx")

FONT = "Microsoft YaHei"
NAVY = RGBColor(11, 37, 69)
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
CYAN = RGBColor(0, 142, 204)
MUTED = RGBColor(91, 111, 130)
LIGHT = "E8EEF5"
PALE = "F4F6F9"
WHITE = RGBColor(255, 255, 255)


PAGES = [
    {
        "name": "封面",
        "title": "AI大模型应用探索：关键矿产情报采集分析系统",
        "lead": "以开源情报为基础，以证据化分析为核心，以决策输出为目标，构建面向实战的关键矿产情报作战平台。",
        "items": [
            ("汇报定位", "围绕关键矿产出口管制场景，展示AI大模型在情报采集、风险研判和成果输出中的应用探索。"),
            ("视觉方向", "全球矿产供应链、跨境数据流与AI网络动态汇聚，形成系统名称和平台核心标识。"),
        ],
        "visual": "深蓝科技场景；全球地图、矿产节点、数据链路与AI光束形成视觉中心。",
        "talk": "开场不急于介绍功能，先明确平台解决的是关键矿产情报工作从信息获取到决策支撑的完整问题。",
    },
    {
        "name": "信息情报工作",
        "title": "信息情报工作的核心，是把分散信息转化为行动依据",
        "lead": "平台围绕真实业务过程组织能力，不是简单堆叠数据和功能，而是打通“发现信息—验证事实—形成判断—服务决策”的工作链条。",
        "items": [
            ("定向采集", "持续获取政策公告、贸易数据、执法案例、企业动态和境外供应链信息。"),
            ("结构化整理", "统一来源、时间、物项、主体、税号和流向等关键字段，形成可查询、可复用的情报资产。"),
            ("风险研判", "基于政策边界、交易特征和多源证据识别疑点，形成分级核查线索。"),
            ("决策转化", "将事实、证据和判断组织为处置建议、专题研究和正式报告。"),
        ],
        "visual": "横向业务链路图：定向采集 → 结构化整理 → 风险研判 → 决策转化。",
        "talk": "本页讲业务要完成什么；下一页再讲平台用什么架构支撑，避免业务流程与系统架构混为一谈。",
    },
    {
        "name": "总体架构",
        "title": "三层能力闭环：感知、研判、输出",
        "lead": "情报工作的价值，不止于“看见信息”，更在于发现风险、支撑核查、服务决策。",
        "items": [
            ("情报感知层", "汇聚政策公告、矿产态势、开源情报、执法案例、贸易数据和研究资料，形成持续更新的可信底座。"),
            ("分析研判层", "通过政策时间轴、AI智能分析和专题研究组织证据、识别风险。"),
            ("决策输出层", "形成核查建议、专题研究、海关要情、综合信息呈报和课题研究报告。"),
            ("两类基础支撑", "数据底座支撑政策、情报、文件、贸易和实体关系管理；AI能力支撑检索增强、证据提取、风险分析和报告编排。"),
        ],
        "visual": "“感知—研判—输出”三层架构图，下方呈现数据底座与AI能力两类基础支撑。",
        "talk": "突出三层之间的数据回流：分析结果沉淀为新情报资产，报告成果反向指导下一轮采集和核查。",
    },
    {
        "name": "第一板块导入：情报工作台",
        "title": "情报工作台：构建面向日常业务的统一作业入口",
        "lead": "情报工作台不是简单的信息展示界面，而是将政策、市场、案例和文件转化为可持续维护的业务资产。",
        "items": [
            ("全局态势总览", "集中呈现政策状态、矿产品变化、最新情报、执法案例和重点预警，帮助快速掌握整体情况。"),
            ("政策与矿产专题监测", "围绕管制清单、政策状态、出口指数、海外价格和替代供应链进展，持续跟踪重点矿产变化。"),
            ("开源证据与知识资产管理", "将网页信息、执法案例、研究资料和业务文件结构化归档，保留来源、时间和原文依据。"),
        ],
        "visual": "以情报工作台首页为中心，向外连接“全局态势、专题监测、知识资产”三个功能模块。",
        "talk": "把“到处查资料”转变为“在一个工作台持续掌握情况、积累证据和管理专题”。",
    },
    {
        "name": "情报总览",
        "title": "一张态势图串联政策、市场、执法与研究成果",
        "lead": "一屏统揽整体态势，让政策变化、市场异动、执法关联和研究成果在同一视角下联动呈现。",
        "items": [
            ("核心指标", "11组政策动态监测；22种重点矿产品指数展示；6条情报快照、4条执法案例；3份深度研究、2份战略报告。"),
            ("政策态势雷达", "现行、暂停及到期窗口一目了然，快速识别需要持续关注的政策节点。"),
            ("最新动态流", "政策、情报、案例和研究成果按时间统一排序，减少跨模块查找。"),
            ("风险预警面板", "同步提示政策到期、价格异动和执法关联三类重点风险信号。"),
        ],
        "visual": "情报总览驾驶舱全屏截图，突出核心指标、全球态势图、动态流和预警面板。",
        "talk": "本页强调“一屏掌握”，不展开单项功能细节，后续页面再逐项说明。",
    },
    {
        "name": "关键矿产清单",
        "title": "构建统一、可核验的出口管制业务规则底座",
        "lead": "把复杂政策转化为可查询、可核验、可执行的业务规则，降低不同人员对政策理解不一致的风险。",
        "items": [
            ("政策完整归集", "归集镓锗、石墨、锑、钨、中重稀土、超硬材料等11组政策。"),
            ("商品逐项拆分", "将公告中的组合物项拆分为26项独立记录，支持逐项查询和关联分析。"),
            ("多维关联查询", "关联物项名称、管制编码、参考税号、政策来源、许可要求和技术条件。"),
            ("状态动态管理", "现行与暂停分类展示，纳入暂停公告、到期时间和政策恢复逻辑。"),
        ],
        "visual": "政策趋势图、管制清单表格和公告参考税号清单组合展示。",
        "talk": "强调税号只是参考维度，是否受控仍需结合成分、纯度、形态、技术参数、最终用户和生效时间综合判断。",
    },
    {
        "name": "关键矿产态势",
        "title": "从政策变化延伸到全球替代供应链竞争",
        "lead": "不仅关注“管了什么”，还持续观察境外“如何应对、何时形成替代能力”。",
        "items": [
            ("出口量指数", "以2023年为统一基期，观察22种重点矿产品出口变化方向。"),
            ("海外参考价格", "展示代表性产品价格及月度变动，辅助识别市场异常。"),
            ("国别政策态势", "跟踪中国、美国、欧盟、日本、韩国、澳大利亚等主要经济体。"),
            ("替代进展", "区分政策宣布、投资建厂、项目建设、试产和商业交付等不同阶段。"),
            ("供应链研判", "把矿山、加工、认证、长期承购和战略储备放在同一链条中分析。"),
        ],
        "visual": "全球态势地图、出口指数图、海外价格趋势和替代项目进度组合展示。",
        "talk": "口径提示：当前出口指数和价格主要用于功能展示，后续需接入权威数据源。",
    },
    {
        "name": "开源信息情报",
        "title": "让每一条公开信息都成为可追溯的证据线索",
        "lead": "把网页信息转化为带来源、带时间、带标签、可全文阅读的结构化情报。",
        "items": [
            ("情报快照", "保存采集时间、原始链接、来源机构、分类标签和关键事实，避免网页变化造成证据丢失。"),
            ("外文翻译", "外文资料先完成中文转换，再进入分析和归档流程，并保留原文对照。"),
            ("执法案例库", "按国家、执法机构、矿种、案件状态和流向组织案例。"),
            ("AI风险信号", "从公开案例中提炼伪报品名、税号不符、夹藏、拆分出口和异常转运等风险模式。"),
        ],
        "visual": "情报快照详情、外文对照、案例库和AI风险信号四个局部截图。",
        "talk": "当前已完成多矿种筛选框架，钨专题率先形成“采集—案例—风险分析”闭环。",
    },
    {
        "name": "文件库",
        "title": "让情报资产统一沉淀，检索不再依赖个人记忆",
        "lead": "把政策、执法、研究、采集快照和AI生成成果统一归档，形成平台知识底仓。",
        "items": [
            ("多格式管理", "支持DOCX、PDF、PPT、Excel、CSV等文件分类管理。"),
            ("搜索与筛选", "按来源、文件类型、矿产和资料类别快速定位。"),
            ("全文预览与分页", "支持内容预览、下载及每页10条的分页浏览。"),
            ("能力复用", "为AI分析、深度研究和战略报告提供可重复调用的资料。"),
        ],
        "visual": "文件中心列表、分类筛选、全文预览和资料详情组合截图。",
        "talk": "强调约70项文件资产已纳入统一管理，平台知识不会随着人员变化而流失。",
    },
    {
        "name": "第二板块导入：研究和分析工具",
        "title": "研究和分析工具：从信息发现走向证据化研判",
        "lead": "依托情报底座，把分散信息进一步转化为有边界、有依据、可复核的风险判断。",
        "items": [
            ("政策演化与规则理解", "通过时间轴还原政策发布、生效、调整、暂停和恢复的连续监管逻辑。"),
            ("证据化AI研判", "围绕时间、物项、主体和来源组织证据，排除错误样本并形成分级核查线索。"),
            ("专题分析任务", "把政策规则、公开案例和贸易明细纳入同一分析过程，形成可复核的专题研判成果。"),
        ],
        "visual": "以研究任务为中心，连接政策时间轴、证据边界、AI分析结果和核查建议。",
        "talk": "让结论说得清、证据追得到、风险定位准、核查建议可执行。",
    },
    {
        "name": "政策时间轴",
        "title": "让散落公告进入统一的时间坐标",
        "lead": "政策不是孤立文件，而是一条包含发布、生效、调整、暂停和恢复的连续监管逻辑线。",
        "items": [
            ("年份筛选", "还原政策演变顺序，快速定位关键时间节点。"),
            ("类型筛选", "区分战略金属、中重稀土、技术和超硬材料等类别。"),
            ("状态筛选", "突出管制生效、暂停及到期窗口。"),
            ("分析价值", "把交易发生时间与政策状态联动，避免跨越政策边界作出错误判断。"),
        ],
        "visual": "横向政策时间轴，突出公告发布、生效、调整、暂停和到期节点。",
        "talk": "避免将政策生效前交易错误纳入违规分析，也避免忽略暂停政策对统计结果的影响。",
    },
    {
        "name": "AI智能分析",
        "title": "智能研判风险：让每条线索更清晰、更可追溯",
        "lead": "模型不是替代业务判断，而是帮助组织证据、排除错误样本并发现优先核查线索。",
        "items": [
            ("证据边界", "设定2025年2月4日钨相关政策生效时间边界，将生效前交易排除在违规研判之外。"),
            ("双重校验", "对时间与物项进行联合校验，并关联企业、商品、税号、路线和单证。"),
            ("线索分层", "识别1起公开处罚案例，形成4条许可核查优先线索，新增2条高优先核查链。"),
            ("贸易匹配", "定向匹配338条商业数据库交易明细记录，并明确其不等同于338票独立出口。"),
            ("可追溯输出", "事实、推断和待核实事项分开展示，每条线索附可核验编号和建议动作。"),
        ],
        "visual": "证据边界、风险指标、重点主体、交易明细和核查建议联动页面。",
        "talk": "重点展示AI如何减少误判、提高核查排序效率，而不是强调模型自动得出最终执法结论。",
    },
    {
        "name": "第三板块导入：战略报告生成",
        "title": "战略报告生成：让研判成果进入决策链条",
        "lead": "将情报、证据、观点和图表组织为规范化决策产品，使研究过程可审核、成果可复用、报告可上报。",
        "items": [
            ("研究任务组织", "围绕复杂问题开展拆解、检索、核验和观点形成。"),
            ("报告产品化", "按照海关要情、综合信息呈报和课题研究报告等规范编排成果。"),
            ("审核与归档", "AI生成初稿，人工复核事实和结论，最终成果统一归档并进入知识底仓。"),
        ],
        "visual": "从研究问题、证据卡片、观点图表到正式报告的纵向生成链路。",
        "talk": "让每一次采集、每一次核查和每一轮研判，最终沉淀为可以复用、审核和上报的正式成果。",
    },
    {
        "name": "AI深度研究",
        "title": "从复杂问题到可核验的研究结论",
        "lead": "AI Agent围绕研究问题完成任务拆解、多轮检索、证据分级和交叉核验，人工负责关键证据审核与最终结论把关。",
        "items": [
            ("问题拆解", "自动拆解研究问题，识别政策、主体、商品、税号和时间等子任务。"),
            ("检索计划", "制定多轮检索计划，检索政府、海关、法院和企业资料。"),
            ("证据提取", "提取政策、企业、商品、税号和关键事件信息，并形成证据索引。"),
            ("边界表达", "区分事实、推断与待核实事项，避免把推测包装为确定性结论。"),
            ("成果生成", "生成主要观点、图表、脚注和待核实清单。"),
            ("当前成果", "已形成3份AI深度研究成果，覆盖反管制体系、境外供应链重构、监管难点和风险防控等专题。"),
        ],
        "visual": "研究任务拆解树、多轮检索过程、证据分级卡片和研究成果预览。",
        "talk": "深度研究的核心价值是把复杂问题拆成可验证步骤，并让每个观点都能回到证据来源。",
    },
    {
        "name": "AI战略报告",
        "title": "从研究成果到正式决策产品",
        "lead": "将情报、证据、研判结论和数据图表自动编排为规范化报告，由业务人员审核定稿，推动研究成果进入决策链条。",
        "items": [
            ("三类报告产品", "支持海关要情、海关综合信息呈报和课题研究报告。"),
            ("自动提炼编排", "提炼核心观点和决策摘要，按正式报告结构生成标题、摘要、正文、图表和脚注。"),
            ("人工复核审批", "支持业务人员修改、复核和审批，确保事实、表述和建议符合业务要求。"),
            ("多格式输出", "支持Word、PDF格式查看与下载。"),
            ("归档复用", "报告成果统一归档，为后续问答、研究和专题更新提供可复用材料。"),
            ("当前成果", "已形成2份AI战略报告，覆盖供应链重构、监管风险和防控建议等专题。"),
        ],
        "visual": "左侧正式报告封面与目录，右侧展示“编排—审核—输出—归档”流程。",
        "talk": "强调AI负责提高材料组织和初稿生成效率，正式报告仍坚持人工审核和最终定稿。",
    },
    {
        "name": "价值总结与下一步",
        "title": "让信息情报工作从“资料支撑”走向“风险牵引”",
        "lead": "平台把开源信息、业务规则、风险线索和研究成果连接起来，使信息情报工作能够更早发现风险、更准支撑核查、更专业服务决策。",
        "items": [
            ("情报更系统", "政策、市场、执法、案例和文件统一归集，减少信息碎片化与重复劳动。"),
            ("研判更精准", "通过证据边界、AI分析和多源交叉核验提高风险识别质量。"),
            ("输出更专业", "把情报成果转化为核查建议、专题研究和正式报告。"),
            ("下一步一：夯实数据", "接入权威出口量、价格、贸易和许可证核验数据，统一数据更新口径。"),
            ("下一步二：扩展专题", "扩展钨以外矿产专题，持续沉淀政策规则、案例和专题研究成果。"),
            ("下一步三：完善治理", "对接内网大模型及海关业务系统，增加权限、审计、人工复核和报告审批流程。"),
        ],
        "visual": "左侧三项价值，右侧三条建设路线，底部以收束语形成汇报结束画面。",
        "talk": "情报不是终点，研判也不是终点。真正的目标，是把信息优势转化为风险防控优势，把分析成果转化为决策支撑能力。",
    },
]


def set_font(run, size=None, color=None, bold=None, italic=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade_paragraph(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_bottom_border(paragraph, color="2E74B5", size="12"):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])
    set_font(run, 9, MUTED)


def make_numbering(doc):
    numbering = doc.part.numbering_part.element
    existing_abs = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    existing_num = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(existing_abs, default=-1) + 1
    num_id = max(existing_num, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal")
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "%1.")
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.extend([tabs, ind, spacing])
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "2E74B5")
    bold = OxmlElement("w:b")
    r_pr.extend([color, bold])
    lvl.extend([start, num_fmt, lvl_text, suff, p_pr, r_pr])
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num])


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.right_margin = Inches(0.85)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.85)
    section.header_distance = Inches(0.38)
    section.footer_distance = Inches(0.38)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    if "Page Title" not in styles:
        page_title = styles.add_style("Page Title", WD_STYLE_TYPE.PARAGRAPH)
    else:
        page_title = styles["Page Title"]
    page_title.font.name = FONT
    page_title._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    page_title.font.size = Pt(21)
    page_title.font.bold = True
    page_title.font.color.rgb = NAVY
    page_title.paragraph_format.space_before = Pt(2)
    page_title.paragraph_format.space_after = Pt(11)
    page_title.paragraph_format.keep_with_next = True

    for style_name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ]:
        style = styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(2)
    hr = hp.add_run("关键矿产情报采集分析系统  |  汇报大纲（完善版）")
    set_font(hr, 8.5, MUTED, True)
    add_bottom_border(hp, "D8E1EA", "6")

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fr = fp.add_run("内部汇报材料  ·  ")
    set_font(fr, 8.5, MUTED)
    add_page_field(fp)

    num_id = make_numbering(doc)

    for page_idx, page in enumerate(PAGES, 1):
        if page_idx > 1:
            doc.add_page_break()

        kicker = doc.add_paragraph()
        kicker.paragraph_format.space_before = Pt(5)
        kicker.paragraph_format.space_after = Pt(6)
        kr = kicker.add_run(f"第 {page_idx:02d} 页  |  {page['name']}")
        set_font(kr, 9.5, CYAN, True)

        title = doc.add_paragraph(style="Page Title")
        tr = title.add_run(page["title"])
        set_font(tr, 21, NAVY, True)
        add_bottom_border(title, "2E74B5", "14")

        label = doc.add_paragraph()
        label.paragraph_format.space_before = Pt(4)
        label.paragraph_format.space_after = Pt(4)
        lr = label.add_run("页面主旨")
        set_font(lr, 10, BLUE, True)

        lead = doc.add_paragraph()
        lead.paragraph_format.left_indent = Inches(0.12)
        lead.paragraph_format.right_indent = Inches(0.12)
        lead.paragraph_format.space_before = Pt(0)
        lead.paragraph_format.space_after = Pt(9)
        lead.paragraph_format.line_spacing = 1.3
        shade_paragraph(lead, PALE)
        lead_r = lead.add_run(page["lead"])
        set_font(lead_r, 11, NAVY, True)

        h = doc.add_paragraph("核心内容", style="Heading 2")
        for item_title, item_desc in page["items"]:
            p = doc.add_paragraph()
            apply_numbering(p, num_id)
            p.paragraph_format.keep_together = True
            p.paragraph_format.space_after = Pt(4)
            r1 = p.add_run(item_title + "：")
            set_font(r1, 10.5, DARK_BLUE, True)
            r2 = p.add_run(item_desc)
            set_font(r2, 10.5, RGBColor(37, 49, 61))

        vh = doc.add_paragraph("画面 / 截图建议", style="Heading 2")
        vp = doc.add_paragraph()
        vp.paragraph_format.left_indent = Inches(0.12)
        vp.paragraph_format.right_indent = Inches(0.12)
        vp.paragraph_format.space_after = Pt(7)
        shade_paragraph(vp, LIGHT)
        vr = vp.add_run(page["visual"])
        set_font(vr, 10, DARK_BLUE)

        th = doc.add_paragraph("讲述重点", style="Heading 2")
        tp = doc.add_paragraph()
        tp.paragraph_format.left_indent = Inches(0.12)
        tp.paragraph_format.right_indent = Inches(0.12)
        tp.paragraph_format.space_after = Pt(0)
        trun = tp.add_run(page["talk"])
        set_font(trun, 10.5, RGBColor(55, 66, 78), False, True)

    doc.core_properties.title = "AI大模型应用探索：关键矿产情报采集分析系统｜汇报大纲（完善版）"
    doc.core_properties.subject = "关键矿产情报采集分析系统H5汇报大纲"
    doc.core_properties.author = "Codex"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
