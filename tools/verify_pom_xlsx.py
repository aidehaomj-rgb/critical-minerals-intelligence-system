from pathlib import Path
from openpyxl import load_workbook
x=Path(r'D:\易迅数据\反倾销税深度分析报告\06_共聚聚甲醛（POM）\POM_易迅全量逐票判定台账.xlsx')
wb=load_workbook(x,read_only=False,data_only=False)
for ws in wb.worksheets:
 print(ws.title,ws.max_row,ws.max_column,ws['A1'].value)
