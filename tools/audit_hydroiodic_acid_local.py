from __future__ import annotations

import csv
import json
import re
import sys
import zipfile
from collections import Counter
from datetime import datetime, timezone, timedelta
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


ROOT = Path(r"D:\易迅数据")
OUT = ROOT / "反倾销税深度分析报告" / "21_氢碘酸"
OUT.mkdir(parents=True, exist_ok=True)
EXTS = {".xlsx", ".xls", ".csv", ".json"}

HS_RE = re.compile(r"(?<!\d)28111990(?!\d)", re.I)
CAS_RE = re.compile(r"(?<!\d)10034[- ]?85[- ]?2(?!\d)", re.I)
NAME_RE = re.compile(
    r"HYDRIODIC\s+ACID|HYDROIODIC\s+ACID|HYDROGEN\s+IODIDE(?:\s+(?:AQUEOUS|SOLUTION))?|"
    r"氢碘酸|碘化氢(?:水溶液)?",
    re.I,
)
FORMULA_CONTEXT_RE = re.compile(
    r"(?<![A-Z0-9])HI(?![A-Z0-9]).{0,80}(?:ACID|IODIDE|AQUEOUS|SOLUTION|CAS|28111990|55\s*%|57\s*%)|"
    r"(?:ACID|IODIDE|AQUEOUS|SOLUTION|CAS|28111990|55\s*%|57\s*%).{0,80}(?<![A-Z0-9])HI(?![A-Z0-9])",
    re.I,
)
OTHER_ACID_RE = re.compile(
    r"HYDROCHLORIC|HYDROBROMIC|HYDROFLUORIC|PHOSPHOROUS|HYPOPHOSPHOROUS|SULFAMIC|"
    r"BORIC|PERCHLORIC|CHLOROSULFONIC|OTHER\s+INORGANIC\s+ACID|盐酸|氢溴酸|氢氟酸|亚磷酸|次磷酸|硼酸",
    re.I,
)
IODIDE_PRODUCT_RE = re.compile(
    r"POTASSIUM\s+IODIDE|SODIUM\s+IODIDE|AMMONIUM\s+IODIDE|METHYL\s+IODIDE|"
    r"IODATE|IODOPHOR|IODIN(?:E|ATED)|碘化钾|碘化钠|碘酸盐|碘伏|甲基碘",
    re.I,
)
DERIVATIVE_RE = re.compile(r"(?:HYDRIODIDE|HYDROIODIDE)\s+(?:SALT|COMPOUND)|氢碘酸盐|碘化氢加成物", re.I)
META_RE = re.compile(r"反倾销|报告状态|任务ID|受税来源|商务部公告|查询组合|调查产品|税率", re.I)
DATE_RE = re.compile(r"20\d{2}[-/]\d{1,2}[-/]\d{1,2}")


def clean(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def classify(text: str) -> tuple[str, str, list[str]]:
    hs = bool(HS_RE.search(text))
    cas = bool(CAS_RE.search(text))
    name = bool(NAME_RE.search(text))
    formula = bool(FORMULA_CONTEXT_RE.search(text))
    basis = [label for label, flag in (("措施HS", hs), ("CAS", cas), ("化学名称", name), ("HI上下文", formula)) if flag]
    if not basis:
        return "不命中", "", []
    if META_RE.search(text) and not DATE_RE.search(text):
        return "项目元数据_政策引用", "非贸易记录", basis
    if OTHER_ACID_RE.search(text) and not (name or cas):
        return "同税号其他无机酸_明确排除", "公告明确同税号除氢碘酸外其他产品不在范围", basis
    if IODIDE_PRODUCT_RE.search(text) and not (name or cas):
        return "其他碘化物或含碘产品_明确排除", "并非碘化氢水溶液", basis
    if DERIVATIVE_RE.search(text) and not cas:
        return "氢碘酸衍生物_倾向排除", "需CAS和成分表确认是否为游离氢碘酸水溶液", basis
    if name or cas:
        return "氢碘酸候选", "需逐票核浓度、形态、原产国、生产商、中国进口申报HS及税款", basis
    if hs:
        return "仅措施HS待货描", "28111990还包含其他无机酸，不能仅凭税号纳入", basis
    return "仅HI上下文待核", "HI噪声高，需化学名称/CAS/浓度/水溶液描述闭合", basis


def source_class(path: Path) -> str:
    s = str(path)
    if "_易迅页面采集" in s:
        return "易迅原始页面采集"
    if "反倾销税深度分析报告" in s:
        return "项目派生输出/QA"
    if path.parent == ROOT:
        return "用户下载原始文件"
    return "其他本地文件"


def make_record(path: Path, location: str, row_no: int, values: list[Any], headers: list[str] | None) -> dict[str, Any]:
    vals = [clean(x) for x in values]
    text = " | ".join(v for v in vals if v)
    classification, reason, basis = classify(text)
    mapped = {clean(k): v for k, v in zip(headers or [], vals)} if headers and len(headers) == len(vals) else {}
    return {
        "source_file": str(path), "source_class": source_class(path), "location": str(location),
        "row_no": row_no, "classification": classification, "classification_reason": reason,
        "match_basis": ";".join(basis), "record_text": text[:30000],
        "record_json": json.dumps(mapped, ensure_ascii=False) if mapped else "",
    }


def walk_json(obj: Any, loc: str = "$"):
    if isinstance(obj, dict):
        keys = " ".join(map(str, obj.keys()))
        if re.search(r"日期|HS|描述|商品|采购|供应|买方|卖方|date|description|buyer|seller|origin", keys, re.I):
            yield loc, obj
        for k, v in obj.items():
            if isinstance(v, (dict, list)):
                yield from walk_json(v, f"{loc}.{k}")
    elif isinstance(obj, list):
        if obj and all(not isinstance(x, (dict, list)) for x in obj):
            yield loc, obj
        else:
            for i, v in enumerate(obj):
                yield from walk_json(v, f"{loc}[{i}]")


def iter_rows(path: Path) -> Iterable[dict[str, Any]]:
    if path.suffix.lower() == ".xlsx":
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            for ws in wb.worksheets:
                it = ws.iter_rows(values_only=True)
                try:
                    first = list(next(it))
                except StopIteration:
                    continue
                headers = [clean(x) for x in first]
                header_like = any(re.search(r"日期|HS|描述|商品|采购|供应|买方|卖方|date|description", h, re.I) for h in headers)
                if not header_like:
                    yield make_record(path, ws.title, 1, first, None)
                    headers = None
                for idx, row in enumerate(it, 2):
                    yield make_record(path, ws.title, idx, list(row), headers)
        finally:
            wb.close()
    elif path.suffix.lower() == ".csv":
        csv.field_size_limit(min(sys.maxsize, 2_000_000_000))
        for enc in ("utf-8-sig", "utf-8", "gb18030"):
            try:
                with path.open(encoding=enc, newline="") as f:
                    reader = csv.reader(f)
                    try:
                        headers = next(reader)
                    except StopIteration:
                        return
                    for idx, row in enumerate(reader, 2):
                        yield make_record(path, "CSV", idx, row, headers if len(headers) == len(row) else None)
                return
            except UnicodeDecodeError:
                continue
    elif path.suffix.lower() == ".json":
        data = json.loads(path.read_text(encoding="utf-8-sig"))
        for loc, obj in walk_json(data):
            if isinstance(obj, dict):
                yield make_record(path, loc, 1, list(obj.values()), list(obj.keys()))
            else:
                yield make_record(path, loc, 1, obj, None)


def prefilter(path: Path) -> tuple[bool, str, str]:
    try:
        if path.suffix.lower() == ".xlsx":
            with zipfile.ZipFile(path) as zf:
                names = [n for n in zf.namelist() if n == "xl/sharedStrings.xml" or (n.startswith("xl/worksheets/") and n.endswith(".xml"))]
                text = " ".join(zf.read(n).decode("utf-8", errors="ignore") for n in names)
        else:
            raw = path.read_bytes()
            text = ""
            for enc in ("utf-8", "utf-8-sig", "gb18030", "latin-1"):
                try:
                    text = raw.decode(enc)
                    break
                except UnicodeDecodeError:
                    continue
        basis = [label for label, regex in (("HS", HS_RE), ("NAME", NAME_RE), ("CAS", CAS_RE), ("HI_CONTEXT", FORMULA_CONTEXT_RE)) if regex.search(text)]
        return bool(basis), "", ";".join(basis)
    except Exception as exc:
        return False, f"{type(exc).__name__}: {exc}", ""


def write_csv(path: Path, rows: list[dict[str, Any]], fields: list[str]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        w.writeheader(); w.writerows(rows)


def main() -> None:
    files = sorted(p for p in ROOT.rglob("*") if p.is_file() and p.suffix.lower() in EXTS and OUT not in p.parents and not p.name.startswith("~$"))
    inventory, candidates = [], []
    for p in files:
        derived = "反倾销税深度分析报告" in str(p) and "_易迅页面采集" not in str(p)
        hit, error, basis = (False, "", "派生输出仅盘点") if derived else prefilter(p)
        inventory.append({"file": str(p), "extension": p.suffix.lower(), "size_bytes": p.stat().st_size,
                          "source_class": source_class(p), "prefilter_hit": hit, "match_basis": basis, "error": error})
        if hit and not error and not derived:
            candidates.append(p)
    write_csv(OUT / "氢碘酸_本地预筛文件台账.csv", inventory, list(inventory[0]))

    hits, audit = [], []
    for p in candidates:
        before, error = len(hits), ""
        try:
            for row in iter_rows(p):
                if row["classification"] != "不命中":
                    hits.append(row)
        except Exception as exc:
            error = f"{type(exc).__name__}: {exc}"
        audit.append({"file": str(p), "hit_rows": len(hits)-before, "error": error})
    hit_fields = ["source_file","source_class","location","row_no","classification","classification_reason","match_basis","record_text","record_json"]
    write_csv(OUT / "氢碘酸_本地深扫全部命中.csv", hits, hit_fields)
    (OUT / "氢碘酸_本地深扫全部命中.json").write_text(json.dumps(hits, ensure_ascii=False, indent=2), encoding="utf-8")
    write_csv(OUT / "氢碘酸_本地深扫文件审计.csv", audit, ["file","hit_rows","error"])

    dedup: dict[str, dict[str, Any]] = {}
    for row in hits:
        key = row["record_text"].strip()
        if key not in dedup:
            dedup[key] = dict(row, occurrence_count=1, provenance=[])
        else:
            dedup[key]["occurrence_count"] += 1
        dedup[key]["provenance"].append({"file": row["source_file"], "location": row["location"], "row_no": row["row_no"]})
    unique = []
    for i, row in enumerate(dedup.values(), 1):
        r = dict(row, dedup_id=f"HI-{i:05d}")
        r["provenance"] = json.dumps(r["provenance"], ensure_ascii=False)
        unique.append(r)
    unique_fields = ["dedup_id","occurrence_count"] + hit_fields + ["provenance"]
    write_csv(OUT / "氢碘酸_本地深扫去重记录.csv", unique, unique_fields)
    (OUT / "氢碘酸_本地深扫去重记录.json").write_text(json.dumps(unique, ensure_ascii=False, indent=2), encoding="utf-8")

    summary = {
        "updated_at": datetime.now(timezone(timedelta(hours=8))).isoformat(timespec="seconds"),
        "scan_scope": "D盘易迅数据下xlsx/xls/csv/json；项目派生输出只盘点、不重复展开；易迅原始页面采集保留",
        "files_total": len(files), "candidate_files": len(candidates),
        "prefilter_errors": sum(bool(x["error"]) for x in inventory),
        "deep_scan_errors": sum(bool(x["error"]) for x in audit),
        "hit_occurrences": len(hits), "visible_unique": len(unique),
        "classification_raw": dict(Counter(x["classification"] for x in hits)),
        "classification_unique": dict(Counter(x["classification"] for x in unique)),
        "candidate_source_files": [str(x) for x in candidates],
        "important_limit": "本地盘点不等于易迅网页全库查询；0条候选只能表述为现有D盘数据未发现。",
    }
    (OUT / "氢碘酸_本地盘点摘要.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
