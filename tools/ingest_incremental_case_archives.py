"""Safely ingest selected GDRM ZIP archives into the 2026 case library.

Archive contents are treated only as source data.  The originals are preserved
byte-for-byte; encrypted PDF/XLSX content is decrypted only in memory so that
searchable text can be added to SQLite.
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import os
import re
import shutil
import sqlite3
import sys
import zipfile
from collections import defaultdict
from datetime import datetime
from pathlib import Path, PurePosixPath

sys.path.insert(0, str(Path(__file__).resolve().parent / ".vendor"))

import msoffcrypto
from openpyxl import load_workbook
from pypdf import PdfReader


sys.stdout.reconfigure(encoding="utf-8", errors="replace")

ROOT = Path(os.environ.get("CASE_LIBRARY_ROOT", r"D:\codex\XG执法\2026_案例库"))
DB = ROOT / "enforcement_cases_2026.sqlite"
CATALOG = ROOT / "case_catalog.csv"
AUDIT = ROOT / "新增案件302_373入库结果.json"
PASSWORD = os.environ["CASE_ARCHIVE_PASSWORD"]
ARCHIVES = [Path(value) for value in json.loads(os.environ["CASE_ARCHIVES_JSON"])]

INFO_NUMBERS = {
    "GDRM26-344": 132,
    "GDRM26-345": 133,
    "GDRM26-348": 135,
    "GDRM26-349": 136,
    "GDRM26-351": 137,
    "GDRM26-353": 138,
    "GDRM26-355": 140,
    "GDRM26-356": 141,
    "GDRM26-359": 143,
    "GDRM26-360": 144,
    "GDRM26-363": 147,
    "GDRM26-364": 148,
    "GDRM26-365": 149,
    "GDRM26-366": 150,
    "GDRM26-372": 151,
    "GDRM26-373": 152,
}

TRANSPORT_MODES = {
    "GDRM26-344": "陆路",
    "GDRM26-345": "陆路",
    "GDRM26-348": "陆路",
    "GDRM26-349": "空运",
    "GDRM26-351": "海路",
    "GDRM26-353": "空运",
    "GDRM26-355": "空运",
    "GDRM26-356": "空运",
    "GDRM26-359": "陆路",
    "GDRM26-360": "陆路",
    "GDRM26-363": "空运",
    "GDRM26-364": "陆路",
    "GDRM26-365": "河路",
    "GDRM26-366": "陆路",
    "GDRM26-372": "陆路",
    "GDRM26-373": "陆路",
}

CATALOG_FIELDS = [
    "case_id",
    "package_name",
    "case_type",
    "transport_mode",
    "date_hint",
    "document_count",
    "raw_path",
]


def normalize_text(value: str) -> str:
    lines: list[str] = []
    for line in (value or "").replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        cleaned = re.sub(r"[ \t]+", " ", line).strip()
        if cleaned:
            lines.append(cleaned)
    return "\n".join(lines)


def digest_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def safe_entry_name(name: str) -> str:
    path = PurePosixPath(name.replace("\\", "/"))
    if path.is_absolute() or ".." in path.parts or not path.name:
        raise RuntimeError(f"Unsafe archive entry: {name}")
    return path.name


def case_id_from_name(name: str) -> str:
    match = re.search(r"GDRM26-(\d{3})", name, re.I)
    if not match:
        raise RuntimeError(f"No case ID in archive entry: {name}")
    return f"GDRM26-{match.group(1)}"


def read_pdf(data: bytes) -> tuple[list[str], bool]:
    reader = PdfReader(io.BytesIO(data))
    encrypted = reader.is_encrypted
    if encrypted and int(reader.decrypt(PASSWORD)) == 0:
        raise RuntimeError("PDF password failed")
    return [normalize_text(page.extract_text() or "") for page in reader.pages], encrypted


def read_xlsx(data: bytes) -> tuple[str, bool, int]:
    encrypted = not data.startswith(b"PK")
    workbook_data = data
    if encrypted:
        office = msoffcrypto.OfficeFile(io.BytesIO(data))
        office.load_key(password=PASSWORD)
        decrypted = io.BytesIO()
        office.decrypt(decrypted)
        workbook_data = decrypted.getvalue()
    book = load_workbook(io.BytesIO(workbook_data), read_only=True, data_only=True)
    lines: list[str] = []
    for sheet in book.worksheets:
        lines.append(f"[工作表] {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            line = " | ".join(normalize_text(str(cell)) for cell in row if cell is not None)
            if line:
                lines.append(line)
    return "\n".join(lines), encrypted, len(book.worksheets)


def load_catalog() -> list[dict[str, str]]:
    with CATALOG.open("r", newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle))


def write_catalog(path: Path, rows: list[dict[str, str]]) -> None:
    rows.sort(key=lambda row: row["case_id"])
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=CATALOG_FIELDS)
        writer.writeheader()
        writer.writerows(rows)


def collect_entries() -> dict[str, list[dict]]:
    grouped: dict[str, list[dict]] = defaultdict(list)
    seen_pairs: set[tuple[str, str]] = set()
    for archive in ARCHIVES:
        if not archive.is_file():
            raise FileNotFoundError(archive)
        with zipfile.ZipFile(archive) as zf:
            for item in zf.infolist():
                if item.is_dir():
                    continue
                filename = safe_entry_name(item.filename)
                case_id = case_id_from_name(filename)
                data = zf.read(item)
                digest = digest_bytes(data)
                pair = (case_id, digest)
                if pair in seen_pairs:
                    raise RuntimeError(f"Duplicate entry across supplied archives: {case_id} {filename}")
                seen_pairs.add(pair)
                suffix = Path(filename).suffix.lower()
                if suffix not in {".pdf", ".xlsx"}:
                    raise RuntimeError(f"Unsupported source type: {filename}")
                grouped[case_id].append(
                    {
                        "archive": str(archive),
                        "filename": filename,
                        "extension": suffix,
                        "data": data,
                        "sha256": digest,
                    }
                )
    return dict(grouped)


def cleanup_new_directories(paths: list[Path]) -> None:
    raw_root = (ROOT / "raw_cases").resolve()
    for path in reversed(paths):
        resolved = path.resolve()
        if resolved.parent != raw_root:
            raise RuntimeError(f"Refusing cleanup outside raw case root: {resolved}")
        if resolved.exists():
            shutil.rmtree(resolved)


def main() -> None:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    grouped = collect_entries()
    catalog_rows = load_catalog()
    catalog_by_id = {row["case_id"]: row for row in catalog_rows}

    conn = sqlite3.connect(DB)
    conn.execute("PRAGMA foreign_keys=ON")
    existing_cases = {row[0] for row in conn.execute("SELECT case_id FROM cases")}
    existing_hashes = {
        row[0]: set()
        for row in conn.execute("SELECT DISTINCT sha256 FROM documents WHERE sha256 IS NOT NULL")
    }
    # Populate hash ownership separately, retaining the ability to audit duplicates.
    hash_cases: dict[str, set[str]] = defaultdict(set)
    for digest, case_id in conn.execute(
        "SELECT DISTINCT sha256, case_id FROM documents WHERE sha256 IS NOT NULL"
    ):
        hash_cases[digest].add(case_id)

    skipped: list[dict] = []
    pending: dict[str, list[dict]] = {}
    for case_id, entries in grouped.items():
        duplicate_entries = [entry for entry in entries if entry["sha256"] in hash_cases]
        if case_id in existing_cases:
            if len(duplicate_entries) != len(entries):
                raise RuntimeError(
                    f"Existing case {case_id} has non-duplicate supplied material; "
                    "supplemental merge requires explicit review"
                )
            wrong_owner = [
                entry for entry in entries if case_id not in hash_cases[entry["sha256"]]
            ]
            if wrong_owner:
                raise RuntimeError(f"Duplicate hash belongs to a different case: {case_id}")
            skipped.append(
                {
                    "case_id": case_id,
                    "reason": "all supplied files already exist with identical SHA-256",
                    "files": [entry["filename"] for entry in entries],
                }
            )
            continue
        if case_id in catalog_by_id:
            raise RuntimeError(f"Catalog contains {case_id}, but cases table does not")
        if duplicate_entries:
            raise RuntimeError(f"New case {case_id} contains a hash already present in the database")
        if case_id not in INFO_NUMBERS or case_id not in TRANSPORT_MODES:
            raise RuntimeError(f"Missing reviewed metadata for {case_id}")
        case_dir = ROOT / "raw_cases" / case_id
        if case_dir.exists():
            raise RuntimeError(f"Destination already exists: {case_dir}")
        names = [entry["filename"].casefold() for entry in entries]
        if len(names) != len(set(names)):
            raise RuntimeError(f"Filename collision inside {case_id}")
        pending[case_id] = entries

    if not pending:
        conn.close()
        raise RuntimeError("No new case material remains after duplicate filtering")

    backup_path = ROOT / f"enforcement_cases_2026.pre_302_373_{timestamp}.sqlite"
    backup_conn = sqlite3.connect(backup_path)
    conn.backup(backup_conn)
    backup_conn.close()
    catalog_backup = ROOT / f"case_catalog.pre_302_373_{timestamp}.csv"
    shutil.copy2(CATALOG, catalog_backup)

    created_dirs: list[Path] = []
    inserted_files: list[dict] = []
    temp_catalog = ROOT / f".case_catalog.{timestamp}.tmp"
    try:
        conn.execute("BEGIN")
        for case_id in sorted(pending):
            entries = pending[case_id]
            case_dir = ROOT / "raw_cases" / case_id
            case_dir.mkdir(parents=True, exist_ok=False)
            created_dirs.append(case_dir)
            relative_dir = case_dir.relative_to(ROOT).as_posix()
            info_number = INFO_NUMBERS[case_id]
            package_name = f"新增附件/{case_id} 情報{info_number}號案件资料"
            conn.execute(
                "INSERT INTO cases VALUES(?,?,?,?,?,?,?)",
                (
                    case_id,
                    package_name,
                    "情報",
                    TRANSPORT_MODES[case_id],
                    "",
                    relative_dir,
                    len(entries),
                ),
            )

            for entry in entries:
                destination = case_dir / entry["filename"]
                destination.write_bytes(entry["data"])
                relative_path = destination.relative_to(ROOT).as_posix()
                result = {
                    "case_id": case_id,
                    "source_archive": entry["archive"],
                    "file": relative_path,
                    "extension": entry["extension"],
                    "sha256": entry["sha256"],
                    "original_bytes": len(entry["data"]),
                }
                if entry["extension"] == ".pdf":
                    pages, encrypted = read_pdf(entry["data"])
                    result["encrypted"] = encrypted
                    result["pages"] = len(pages)
                    result["characters"] = sum(len(text) for text in pages)
                    for page_number, text in enumerate(pages, 1):
                        conn.execute(
                            """INSERT INTO documents(
                                   case_id,relative_path,extension,page_or_sheet,
                                   extracted_text,text_chars,extraction_error,sha256
                               ) VALUES(?,?,?,?,?,?,?,?)""",
                            (
                                case_id,
                                relative_path,
                                ".pdf",
                                f"page {page_number}",
                                text,
                                len(text),
                                None,
                                entry["sha256"],
                            ),
                        )
                        conn.execute(
                            "INSERT INTO document_fts(case_id,relative_path,extracted_text) VALUES(?,?,?)",
                            (case_id, relative_path, text),
                        )
                else:
                    text, encrypted, sheets = read_xlsx(entry["data"])
                    result["encrypted"] = encrypted
                    result["sheets"] = sheets
                    result["characters"] = len(text)
                    conn.execute(
                        """INSERT INTO documents(
                               case_id,relative_path,extension,page_or_sheet,
                               extracted_text,text_chars,extraction_error,sha256
                           ) VALUES(?,?,?,?,?,?,?,?)""",
                        (
                            case_id,
                            relative_path,
                            ".xlsx",
                            "workbook",
                            text,
                            len(text),
                            None,
                            entry["sha256"],
                        ),
                    )
                    conn.execute(
                        "INSERT INTO document_fts(case_id,relative_path,extracted_text) VALUES(?,?,?)",
                        (case_id, relative_path, text),
                    )
                inserted_files.append(result)

            catalog_rows.append(
                {
                    "case_id": case_id,
                    "package_name": package_name,
                    "case_type": "情報",
                    "transport_mode": TRANSPORT_MODES[case_id],
                    "date_hint": "",
                    "document_count": str(len(entries)),
                    "raw_path": relative_dir,
                }
            )

        write_catalog(temp_catalog, catalog_rows)
        conn.commit()
        os.replace(temp_catalog, CATALOG)
    except Exception:
        conn.rollback()
        if temp_catalog.exists():
            temp_catalog.unlink()
        cleanup_new_directories(created_dirs)
        raise
    finally:
        conn.close()

    result = {
        "ingested_at": datetime.now().astimezone().isoformat(),
        "database": str(DB),
        "database_backup": str(backup_path),
        "catalog_backup": str(catalog_backup),
        "archives_processed": [str(path) for path in ARCHIVES],
        "cases_added": sorted(pending),
        "duplicates_skipped": skipped,
        "files_added": inserted_files,
        "stats": {
            "cases_added": len(pending),
            "cases_skipped_as_exact_duplicates": len(skipped),
            "files_added": len(inserted_files),
            "pdf_files": sum(row["extension"] == ".pdf" for row in inserted_files),
            "pdf_pages": sum(row.get("pages", 0) for row in inserted_files),
            "xlsx_files": sum(row["extension"] == ".xlsx" for row in inserted_files),
            "characters": sum(row.get("characters", 0) for row in inserted_files),
            "encrypted_files_parsed": sum(bool(row.get("encrypted")) for row in inserted_files),
        },
    }
    AUDIT.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
