from __future__ import annotations

from collections import Counter, defaultdict
from datetime import datetime
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from pathlib import Path
import csv
import json
import re
import unicodedata

from openpyxl import load_workbook


SOURCE = Path(r"D:\易迅数据\EPDM_400270.xlsx")
OUTPUT_DIR = Path(r"D:\易迅数据\反倾销税深度分析报告\12_EPDM")
SHEET_NAME = "数据列表"

ORIGINAL_FIELDS = [
    "数据源",
    "进出口",
    "日期",
    "HS编码",
    "商品描述",
    "采购商",
    "供应商",
    "重量",
    "数量",
    "金额",
    "目的国/地区",
    "原产国/地区",
]

META_FIELDS = [
    "record_id",
    "excel_row",
    "canonical_record_id",
    "canonical_excel_row",
    "duplicate_group_id",
    "duplicate_group_size",
    "duplicate_ordinal",
    "dedup_keep",
    "is_exact_duplicate",
    "scope_class",
    "scope_rule_code",
    "scope_matched_terms",
    "scope_reason",
    "scope_disposition",
    "route_class",
    "taxed_origin_flag",
    "china_destination_flag",
    "evidence_grade",
    "evidence_reason",
    "decisive_data_gap",
]


TAXED_ORIGIN_NAMES = {
    # 美国、韩国
    "united states",
    "united states of america",
    "south korea",
    "korea, republic of",
    "republic of korea",
    # 欧盟27国（本数据窗口内对欧盟措施仍在复审有效期）
    "austria",
    "belgium",
    "bulgaria",
    "croatia",
    "cyprus",
    "czech republic",
    "czechia",
    "denmark",
    "estonia",
    "finland",
    "france",
    "germany",
    "greece",
    "hungary",
    "ireland",
    "italy",
    "latvia",
    "lithuania",
    "luxembourg",
    "malta",
    "netherlands",
    "the netherlands",
    "poland",
    "portugal",
    "romania",
    "slovakia",
    "slovenia",
    "spain",
    "sweden",
}


STANDARD_PATTERN = re.compile(
    r"(?:"
    r"PRIMARY\s+FORMS?.{0,80}(?:PLATES?|SHEETS?|STRIPS?)|"
    r"FORMAS?\s+PRIMARIAS?.{0,100}(?:PLACAS?|HOJAS?|TIRAS?)|"
    r"FORMES?\s+PRIMAIRES?.{0,100}(?:PLAQUES?|FEUILLES?|BANDES?)|"
    r"FORMAS?\s+PRIM[ÁA]RIAS?.{0,100}(?:CHAPAS?|FOLHAS?|TIRAS?)|"
    r"BENTUK\s+(?:ASAL|PRIMER).{0,100}(?:PELAT|LEMBARAN|JALUR)|"
    r"ПЕРВИЧН.{0,120}(?:ПЛАСТИН|ЛИСТ|ПОЛОС)|"
    r"(?:PRIM[ÄA]RFORM|PRIMÆRFORM).{0,100}(?:PLATT|SKIV|REMS)"
    r")",
    re.I | re.S,
)

COMPOUND_PATTERN = re.compile(
    r"(?:COMPOUND|COMPOUNDED|COMPUEST[OA]|COMPOST[OA]|MISTURA|"
    r"MIXED\s+RUBBER|RUBBER\s+MIX|CAUCHO\s+MEZCLADO|"
    r"FUNCTIONAL\s+POLYMER|FUNCTIONALIZED|\bTPV\b|\bTPE\b|"
    r"SANTOPRENE|TERMOPLAST|THERMOPLAST|MASTERBATCH|POLYMER\s+BLEND|"
    r"EPDM\s+BLEND|MODIFIED\s+EPDM|HỖN\s*HỢP|HOP\s*CHAT|PHOI\s*TRON)",
    re.I,
)

OUTSIDE_PATTERN = re.compile(
    r"(?:"
    r"PROFILES?|PERFILES?|PROFILE\s+AUTO|AUTO\s+PARTS?|AUTOMOTIVE\s+PARTS?|"
    r"GASKETS?|O[- ]?RINGS?|\bSEALS?\b|SEALING|WEATHERSTRIP|WEATHER\s+STRIP|"
    r"HOSES?|MANGUERAS?|\bTUBES?\b|PIPES?|BELTS?|CONVEYOR|"
    r"MEMBRANES?|ROOFING|FLOOR(?:ING)?|MAT(?:S|TING)?\b|"
    r"ADHESIVE|SELF[- ]?ADHESIVE|TAPE\b|\bFOAM\b|\bSPONGE\b|"
    r"RUBBER\s+ARTICLE|RUBBER\s+PRODUCT|FINISHED\s+PRODUCT|"
    r"MOLDED|MOULDED|EXTRUDED|INJECTION\s+MOLD|"
    r"GIO[ĂA]NG|MIẾNG\s+ĐỆM|MIE?NG\s+DEM|ĐỆM\s+BẢO\s+VỆ|"
    r"DẢI\s+CAO\s+SU|DAI\s+CAO\s+SU|PHỤ\s+TÙNG|PHU\s+TUNG|"
    r"TẤM\s+CAO\s+SU\s+XỐP|TAM\s+CAO\s+SU\s+XOP|"
    r"JUNTAS?|SELLOS?|SOPORTES?\s+DE\s+CAUCHO|PIEZAS?\s+DE|"
    r"TIRAS?\s+DE\s+(?:DE\s+)?(?:CAUCHO|HULE).{0,60}(?:ADHES|AUTO|PUERTA)|"
    r"PLACAS?\s+DE\s+CAUCHO\s+(?:CELULAR|ESPONJOSO)|"
    r"ESPUMA|ESPONJA|CABLE|WIRE\s+SEAL|COVER|BUMPER|MOUNT|PAD\b|"
    r"PROTECTIVE\s+PAD|SCRAP|WASTE|RECYCLED\s+ARTICLE"
    r")",
    re.I,
)

FORM_PATTERN = re.compile(
    r"(?:"
    r"\bSHEETS?\b|\bPLATES?\b|\bSTRIPS?\b|\bROLLS?\b|\bPELLETS?\b|"
    r"\bGRANULES?\b|\bTIRAS?\b|\bPLACAS?\b|\bHOJAS?\b|\bROLLOS?\b|"
    r"\bGRANULOS?\b|\bGRANULADO\b|\bCHAPAS?\b|\bFOLHAS?\b|"
    r"\bBANDES?\b|\bPLAQUES?\b|\bFEUILLES?\b|"
    r"DẠNG\s+(?:TẤM|TỜ|HẠT|CUỘN)|DANG\s+(?:TAM|TO|HAT|CUON)|"
    r"HÌNH\s+HẠT|HINH\s+HAT|颗粒|片材|板材|卷材|条状"
    r")",
    re.I,
)


def clean_for_match(value: object) -> str:
    if value is None:
        return ""
    normalized = unicodedata.normalize("NFKC", str(value)).upper()
    return re.sub(r"\s+", " ", normalized).strip()


def norm_country(value: object) -> str:
    return clean_for_match(value).casefold()


def entity_key(value: object) -> str:
    return re.sub(r"[^A-Z0-9]", "", clean_for_match(value))


def matched(pattern: re.Pattern[str], value: str) -> str:
    m = pattern.search(value)
    return "" if m is None else re.sub(r"\s+", " ", m.group(0)).strip()


def classify_scope(description: object) -> dict[str, str]:
    desc = clean_for_match(description)

    term = matched(COMPOUND_PATTERN, desc)
    if term:
        return {
            "scope_class": "混炼改性/TPV/TPE/功能聚合物待核",
            "scope_rule_code": "S1",
            "scope_matched_terms": term,
            "scope_reason": "描述出现混炼、复合、热塑性弹性体或功能化聚合物信号；须核对成分、是否硫化、实际形态及税则归类，不能仅凭HS 400270认定属于措施范围。",
            "scope_disposition": "范围待核",
        }

    term = matched(OUTSIDE_PATTERN, desc)
    if term:
        return {
            "scope_class": "疑似制品/型材/密封件等范围外",
            "scope_rule_code": "S2",
            "scope_matched_terms": term,
            "scope_reason": "描述出现型材、密封件、垫片、汽车零件、泡棉/海绵、胶粘成品等制品信号；初步倾向不属于EPDM原胶措施范围，仍须以申报要素、硫化状态和实物确认。",
            "scope_disposition": "范围外倾向",
        }

    term = matched(STANDARD_PATTERN, desc)
    if term:
        return {
            "scope_class": "税目标准表述（初步纳入）",
            "scope_rule_code": "S3",
            "scope_matched_terms": term,
            "scope_reason": "描述直接复述“初级形状或板、片、带”等HS 400270税目范围语言；在未出现制品或改性信号时初步纳入，但仍须核对原产地和具体牌号。",
            "scope_disposition": "初步纳入",
        }

    term = matched(FORM_PATTERN, desc)
    if term:
        return {
            "scope_class": "板/片/带/卷/粒料等形态待核",
            "scope_rule_code": "S4",
            "scope_matched_terms": term,
            "scope_reason": "描述出现板、片、带、卷或粒料等形态；是否为未硫化原胶税目形态、混炼料或已经制成品，仅凭平台文本不能闭合。",
            "scope_disposition": "形态待核",
        }

    return {
        "scope_class": "原胶/通用EPDM候选",
        "scope_rule_code": "S5",
        "scope_matched_terms": "EPDM/HS 400270基线",
        "scope_reason": "描述为通用EPDM、合成橡胶、牌号或原胶表述，未发现明确制品、改性或特殊形态信号；作为措施范围候选，最终仍需成分、形态、原产地证据。",
        "scope_disposition": "初步纳入候选",
    }


def classify_route(raw_row: dict[str, object]) -> dict[str, object]:
    destination = norm_country(raw_row["目的国/地区"])
    origin = norm_country(raw_row["原产国/地区"])
    buyer = entity_key(raw_row["采购商"])
    supplier = entity_key(raw_row["供应商"])
    description = clean_for_match(raw_row["商品描述"])

    china = destination == "china"
    mexico = destination == "mexico"
    taxed = origin in TAXED_ORIGIN_NAMES
    hexpol = "HEXPOLCOMPOUNDINGSADECV"
    is_hexpol_a = mexico and hexpol in buyer and taxed
    is_hexpol_b = china and origin == "mexico" and hexpol in supplier

    if is_hexpol_a:
        return {
            "route_class": "受税来源→HEXPOL墨西哥（A腿）",
            "taxed_origin_flag": True,
            "china_destination_flag": False,
            "evidence_grade": "B+（实体级双腿线索）",
            "evidence_reason": "同一HEXPOL墨西哥实体在样本期内接收美国、韩国及欧盟来源EPDM；可与其对华B腿构成实体级双腿线索，但本行本身不证明同批货物被转运。",
            "decisive_data_gap": "需批次/牌号、采购订单、入出库、集装箱或提单号、原产地证书及墨西哥加工记录闭合。",
        }

    if is_hexpol_b:
        return {
            "route_class": "HEXPOL墨西哥→中国（B腿）",
            "taxed_origin_flag": False,
            "china_destination_flag": True,
            "evidence_grade": "B+（实体级双腿线索）",
            "evidence_reason": "HEXPOL COMPOUNDING SA DE CV以墨西哥原产字段向中国同日发运4条EPDM，且该实体同期存在大量受税来源A腿；但缺少批次、数量和提单闭合，且墨西哥存在真实混炼加工能力。",
            "decisive_data_gap": "需本票报关单、原产地证、配方/生产工单、BOM、库存台账、原料领用及集装箱/提单关联。",
        }

    if china and taxed:
        return {
            "route_class": "受税来源直接对华",
            "taxed_origin_flag": True,
            "china_destination_flag": True,
            "evidence_grade": "A（直接贸易事实）",
            "evidence_reason": "平台记录的目的国为中国、原产国字段为当前受税来源；这是直接贸易/征税核对对象，不是第三国绕道证据。",
            "decisive_data_gap": "需中国报关单、征免性质、原产地证明、生产商名称、价格承诺/税率适用及税款缴款书。",
        }

    if china and origin == "saudi arabia":
        return {
            "route_class": "沙特来源→中国（第三国对华）",
            "taxed_origin_flag": False,
            "china_destination_flag": True,
            "evidence_grade": "C（低风险背景）",
            "evidence_reason": "第三国来源对华记录成立，但样本未发现受税来源→沙特的实体级A腿；沙特存在真实EPDM产能，现有数据更支持正常产地供应。",
            "decisive_data_gap": "若继续核查，应调取具体生产厂、牌号、聚合工序地点和原产地证，不能仅凭品牌或贸易商判断原产地。",
        }

    if china and origin == "canada" and (
        "FUSABOND" in entity_key(description) or ("FUSABO" in description and "N302" in description)
    ):
        return {
            "route_class": "加拿大FUSABOND→中国（第三国对华）",
            "taxed_origin_flag": False,
            "china_destination_flag": True,
            "evidence_grade": "C+（范围/产地核验线索）",
            "evidence_reason": "4条记录使用HS 400270和加拿大原产字段，但商品描述为FUSABOND N302功能聚合物；重点是措施范围、实际生产地和重复记录核验，尚无受税来源A腿闭合。",
            "decisive_data_gap": "需技术数据表、聚合/接枝工序地点、报关申报要素、原产地证及本票提单去重。",
        }

    if china and origin == "india" and "KEIINDUSTRIES" in supplier:
        return {
            "route_class": "KEI印度→中国（第三国对华）",
            "taxed_origin_flag": False,
            "china_destination_flag": True,
            "evidence_grade": "C+（时序反证后保留）",
            "evidence_reason": "该B腿描述为RUBBER COMPOUND EPDM；样本内KEI的受税来源输入均晚于本票日期，现有时序不能支持此前A腿供应本票，仅保留成分/原产地核验。",
            "decisive_data_gap": "需本票生产工单、配方、原料批次、库存期初余额和原产地证。",
        }

    if china:
        return {
            "route_class": "其他第三国来源→中国",
            "taxed_origin_flag": False,
            "china_destination_flag": True,
            "evidence_grade": "C（方向性线索）",
            "evidence_reason": "目的国为中国、原产国字段为非受税来源；单一B腿不足以证明绕道，须寻找同一实体、时间、数量、牌号或提单层面的受税来源A腿。",
            "decisive_data_gap": "需报关单、原产地证、生产能力、入出库和双腿提单/集装箱关联。",
        }

    if taxed:
        return {
            "route_class": "受税来源→其他国家（全球A腿背景）",
            "taxed_origin_flag": True,
            "china_destination_flag": False,
            "evidence_grade": "C（背景记录）",
            "evidence_reason": "记录显示受税来源流向非中国市场，可用于建立潜在A腿基线；未与对华B腿形成实体/批次关联时不得认定转运。",
            "decisive_data_gap": "需与第三国对华记录在实体、牌号、日期、数量、提单或集装箱层面匹配。",
        }

    return {
        "route_class": "其他全球基线记录",
        "taxed_origin_flag": False,
        "china_destination_flag": False,
        "evidence_grade": "D（背景/不适用）",
        "evidence_reason": "本行不属于对华B腿，也未显示当前受税来源原产字段；保留作全球贸易基线，不据此作绕道判断。",
        "decisive_data_gap": "如与重点实体或B腿产生关联，再补充提单、原产地、生产和库存证据。",
    }


def stable_row_key(values: tuple[object, ...]) -> str:
    payload = json.dumps(values, ensure_ascii=False, separators=(",", ":"), default=str)
    return sha256(payload.encode("utf-8")).hexdigest()


def parse_decimal(value: object) -> Decimal | None:
    if value in (None, ""):
        return None
    try:
        return Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, ValueError):
        return None


def decimal_text(value: Decimal) -> str:
    result = format(value, "f")
    if "." in result:
        result = result.rstrip("0").rstrip(".")
    return result or "0"


def subset_stats(rows: list[dict[str, object]]) -> dict[str, object]:
    result: dict[str, object] = {
        "record_count": len(rows),
        "origin_counts": dict(Counter(str(row.get("原产国/地区") or "（空）") for row in rows)),
        "date_min": min((str(row.get("日期")) for row in rows), default=""),
        "date_max": max((str(row.get("日期")) for row in rows), default=""),
        "raw_field_sum_note": "仅对可解析数字作字段内机械求和；重量/数量/金额的单位和币种不得据此推定或互相换算。",
    }
    for field in ("重量", "数量", "金额"):
        values = [parse_decimal(row.get(field)) for row in rows]
        parsed = [v for v in values if v is not None]
        result[f"{field}_nonblank_numeric_count"] = len(parsed)
        result[f"{field}_raw_field_sum"] = decimal_text(sum(parsed, Decimal("0")))
    return result


def write_csv(path: Path, rows: list[dict[str, object]], fields: list[str]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def write_json(path: Path, value: object) -> None:
    with path.open("w", encoding="utf-8") as handle:
        json.dump(value, handle, ensure_ascii=False, indent=2, default=str)


def build() -> dict[str, object]:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    wb = load_workbook(SOURCE, read_only=True, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise AssertionError(f"缺少工作表：{SHEET_NAME}; 实际={wb.sheetnames}")
    ws = wb[SHEET_NAME]
    iterator = ws.iter_rows(values_only=True)
    headers = list(next(iterator))
    if headers != ORIGINAL_FIELDS:
        raise AssertionError(f"源表列名变化：{headers}")

    raw_records: list[tuple[int, tuple[object, ...]]] = []
    for excel_row, values in enumerate(iterator, start=2):
        if not any(v not in (None, "") for v in values):
            continue
        raw_records.append((excel_row, tuple(values)))

    if len(raw_records) != 9356:
        raise AssertionError(f"源表非空记录数异常：{len(raw_records)}")

    key_members: dict[tuple[object, ...], list[int]] = defaultdict(list)
    for excel_row, values in raw_records:
        key_members[values].append(excel_row)

    canonical_order = {key: idx for idx, key in enumerate(key_members, start=1)}
    standardized: list[dict[str, object]] = []
    duplicate_audit: list[dict[str, object]] = []

    for excel_row, values in raw_records:
        original = dict(zip(ORIGINAL_FIELDS, values))
        members = key_members[values]
        canonical_excel_row = members[0]
        ordinal = members.index(excel_row) + 1
        canonical_seq = canonical_order[values]
        digest = stable_row_key(values)
        canonical_id = f"EPDM-CAN-{canonical_seq:05d}-{digest[:12].upper()}"
        scope = classify_scope(original["商品描述"])
        route = classify_route(original)
        row = {
            "record_id": f"EPDM-XLSX-{excel_row:05d}",
            "excel_row": excel_row,
            "canonical_record_id": canonical_id,
            "canonical_excel_row": canonical_excel_row,
            "duplicate_group_id": canonical_id,
            "duplicate_group_size": len(members),
            "duplicate_ordinal": ordinal,
            "dedup_keep": ordinal == 1,
            "is_exact_duplicate": ordinal > 1,
            **scope,
            **route,
            **original,
        }
        standardized.append(row)
        if len(members) > 1:
            duplicate_audit.append(row)

    canonical = [row for row in standardized if row["dedup_keep"] is True]
    duplicate_groups = [members for members in key_members.values() if len(members) > 1]
    if len(canonical) != 9160:
        raise AssertionError(f"全12字段去重数异常：{len(canonical)}")
    if len(standardized) - len(canonical) != 196:
        raise AssertionError("重复展开数不是196")
    if len(duplicate_groups) != 138 or max(map(len, duplicate_groups), default=0) != 11:
        raise AssertionError(
            f"重复组审计异常：groups={len(duplicate_groups)}, max={max(map(len, duplicate_groups), default=0)}"
        )

    china_rows = [row for row in canonical if row["china_destination_flag"] is True]
    taxed_direct = [row for row in china_rows if row["route_class"] == "受税来源直接对华"]
    third_country_china = [row for row in china_rows if row["taxed_origin_flag"] is False]
    hexpol_a = [row for row in canonical if row["route_class"] == "受税来源→HEXPOL墨西哥（A腿）"]
    hexpol_b = [row for row in canonical if row["route_class"] == "HEXPOL墨西哥→中国（B腿）"]

    assert len(china_rows) == 55, len(china_rows)
    assert len(taxed_direct) == 6, len(taxed_direct)
    assert len(third_country_china) == 49, len(third_country_china)
    assert len(hexpol_a) == 418, len(hexpol_a)
    assert len(hexpol_b) == 4, len(hexpol_b)
    expected_a_origins = {"United States": 389, "South Korea": 19, "Netherlands": 6, "France": 4}
    actual_a_origins = dict(Counter(str(row["原产国/地区"]) for row in hexpol_a))
    assert actual_a_origins == expected_a_origins, actual_a_origins

    full_fields = META_FIELDS + ORIGINAL_FIELDS
    duplicate_fields = [
        "duplicate_group_id",
        "duplicate_group_size",
        "canonical_excel_row",
        "excel_row",
        "duplicate_ordinal",
        "dedup_keep",
        "is_exact_duplicate",
    ] + [f for f in full_fields if f not in {
        "duplicate_group_id", "duplicate_group_size", "canonical_excel_row", "excel_row",
        "duplicate_ordinal", "dedup_keep", "is_exact_duplicate"
    }]

    deliverables = {
        "EPDM_易迅逐票标准化": standardized,
        "EPDM_对华55条": china_rows,
        "EPDM_受税来源直达中国6条": taxed_direct,
        "EPDM_HEXPOL_A腿418条": hexpol_a,
        "EPDM_HEXPOL_B腿4条": hexpol_b,
        "EPDM_第三国对华49条": third_country_china,
    }
    for stem, rows in deliverables.items():
        write_csv(OUTPUT_DIR / f"{stem}.csv", rows, full_fields)
        write_json(OUTPUT_DIR / f"{stem}.json", rows)
    write_csv(OUTPUT_DIR / "EPDM_全12字段重复审计.csv", duplicate_audit, duplicate_fields)
    write_json(OUTPUT_DIR / "EPDM_全12字段重复审计.json", duplicate_audit)

    scope_unique_counts = Counter(str(row["scope_class"]) for row in canonical)
    scope_raw_counts = Counter(str(row["scope_class"]) for row in standardized)
    route_unique_counts = Counter(str(row["route_class"]) for row in canonical)

    b_date = datetime.strptime("2025-12-23", "%Y-%m-%d").date()
    a_dated = []
    for row in hexpol_a:
        try:
            parsed = datetime.strptime(str(row["日期"]), "%Y-%m-%d").date()
        except ValueError:
            continue
        a_dated.append((parsed, row))
    a_before = [row for date, row in a_dated if date < b_date]
    a_prior_30 = [row for date, row in a_dated if 0 < (b_date - date).days <= 30]
    a_prior_7 = [row for date, row in a_dated if 0 < (b_date - date).days <= 7]
    b_descriptions = {clean_for_match(row["商品描述"]) for row in hexpol_b}
    a_same_description = [row for row in hexpol_a if clean_for_match(row["商品描述"]) in b_descriptions]
    a_same_description_before = [
        row
        for date, row in a_dated
        if date < b_date and clean_for_match(row["商品描述"]) in b_descriptions
    ]

    source_hash = sha256(SOURCE.read_bytes()).hexdigest()
    generated_at = datetime.now().astimezone().isoformat(timespec="seconds")
    summary: dict[str, object] = {
        "delivery_id": "EPDM-400270-20250901-20260805-STAGE-AUDIT",
        "generated_at": generated_at,
        "source_file": str(SOURCE),
        "source_sha256": source_hash,
        "source_sheet": SHEET_NAME,
        "source_headers": ORIGINAL_FIELDS,
        "source_data_rows": len(standardized),
        "source_date_min": min(str(row["日期"]) for row in standardized),
        "source_date_max": max(str(row["日期"]) for row in standardized),
        "dedup_method": "12个原始字段逐值完全相等；首条excel_row作为规范记录，其余仅标记重复，不删除原始展开行。",
        "exact_unique_rows": len(canonical),
        "duplicate_extra_rows": len(standardized) - len(canonical),
        "duplicate_group_count": len(duplicate_groups),
        "largest_duplicate_group_size": max(map(len, duplicate_groups), default=0),
        "scope_priority": [
            "S1 混炼/改性/TPV/TPE/功能聚合物待核",
            "S2 型材/密封件/垫片/汽车零件/泡棉制品等范围外倾向",
            "S3 税目标准语句‘初级形状或板、片、带’初步纳入",
            "S4 其余板/片/带/卷/粒料等形态待核",
            "S5 其余原胶/通用EPDM候选",
        ],
        "scope_counts_exact_unique": dict(scope_unique_counts),
        "scope_counts_all_source_rows": dict(scope_raw_counts),
        "scope_reconciliation_note": (
            "本次结构化复核采用多语种字面规则，规范记录计数为原胶/通用EPDM候选6306、"
            "税目标准表述504、形态待核1163、混炼改性待核300、疑似制品范围外887。"
            "此前阶段人工/较窄规则口径为6381/488/1104/300/887；差异集中在多语种‘初级形状/板片带/卷/粒’描述，"
            "不影响对华55条、直达6条、第三国49条和HEXPOL双腿计数，报告引用时须注明所用口径。"
        ),
        "route_counts_exact_unique": dict(route_unique_counts),
        "subsets": {
            "china_55": subset_stats(china_rows),
            "taxed_origin_direct_china_6": subset_stats(taxed_direct),
            "third_country_to_china_49": subset_stats(third_country_china),
            "hexpol_a_418": subset_stats(hexpol_a),
            "hexpol_b_4": subset_stats(hexpol_b),
        },
        "hexpol_relationship_metrics": {
            "a_origin_counts": actual_a_origins,
            "a_before_b_unique_rows": len(a_before),
            "a_within_30_days_before_b_unique_rows": len(a_prior_30),
            "a_within_7_days_before_b_unique_rows": len(a_prior_7),
            "a_same_normalized_description_as_b_unique_rows": len(a_same_description),
            "a_same_normalized_description_as_b_before_b_unique_rows": len(a_same_description_before),
            "interpretation": "实体级、方向级和时间级关联成立；尚无牌号/批次/数量/提单/集装箱闭合，不能据此认定同货转运或原产地虚假。",
        },
        "units_and_currency_caveat": (
            "重量、数量、金额均按源表原字段保留；源表未提供统一单位/币种，任何字段求和仅为机械汇总，"
            "不得写成千克、吨、美元、人民币或税款损失。"
        ),
        "evidence_scale": {
            "A": "记录直接支持贸易方向/原产字段等事实；不等于违法事实。",
            "B+": "同一实体双腿与时间关联较强，但缺决定性批次/提单/原产地证据。",
            "C/C+": "方向性或范围核验线索，存在重要替代解释。",
            "D": "全球背景或与本次对华绕道判断不直接相关。",
        },
        "generated_files": sorted(
            [f"{stem}.{suffix}" for stem in deliverables for suffix in ("csv", "json")]
            + ["EPDM_全12字段重复审计.csv", "EPDM_全12字段重复审计.json", "EPDM_交付摘要.json"]
        ),
        "self_checks": {
            "all_source_rows_9356": len(standardized) == 9356,
            "unique_rows_9160": len(canonical) == 9160,
            "duplicate_extras_196": len(standardized) - len(canonical) == 196,
            "duplicate_groups_138": len(duplicate_groups) == 138,
            "largest_duplicate_group_11": max(map(len, duplicate_groups), default=0) == 11,
            "china_rows_55": len(china_rows) == 55,
            "direct_taxed_origin_china_6": len(taxed_direct) == 6,
            "third_country_china_49": len(third_country_china) == 49,
            "hexpol_a_418": len(hexpol_a) == 418,
            "hexpol_b_4": len(hexpol_b) == 4,
            "hexpol_a_origins_match": actual_a_origins == expected_a_origins,
            "all_scope_rows_reconciled": sum(scope_unique_counts.values()) == len(canonical),
        },
    }
    write_json(OUTPUT_DIR / "EPDM_交付摘要.json", summary)
    return summary


if __name__ == "__main__":
    result = build()
    print(json.dumps(result, ensure_ascii=False, indent=2))
