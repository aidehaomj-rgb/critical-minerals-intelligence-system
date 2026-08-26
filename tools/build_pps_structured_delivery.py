from __future__ import annotations

import csv
import json
import re
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from pathlib import Path
from typing import Any

import openpyxl


SOURCES = [
    Path(r"D:\易迅数据\PPS_391190_1.xlsx"),
    Path(r"D:\易迅数据\PPS_391190_2.xlsx"),
]
OUTPUT_DIR = Path(r"D:\易迅数据\反倾销税深度分析报告\13_PPS")
SHEET_NAME = "数据列表"

ORIGINAL_FIELDS = [
    "数据源",
    "进出口",
    "日期",
    "HS编码",
    "商品描述",
    "采购商",
    "供应商",
    "重量",
    "数量",
    "金额",
    "目的国/地区",
    "原产国/地区",
]

TAXED_ORIGINS = {"JAPAN", "UNITED STATES", "SOUTH KOREA", "MALAYSIA"}
HDC = "HDCPOLYALLCOLTD"
HWASEUNG = "CONGTYTNHHHWASEUNGCHEMICALVIETNAM"
CHAO_JU = "CHAOJUNEWMATERIALTECHNOLOGYCOLTD"

CHINA_CITY_TOKENS = [
    "SHANGHAI",
    "DONGGUAN",
    "DONGUAN",
    "FOSHAN",
    "YANCHENG",
    "CHENZHOU",
    "TIANJIN",
    "GUANGZHOU",
    "HANGZHOU",
    "DANDONG",
    "ZHONGSHAN",
    "ZHONG SHAN",
    "NINGBO",
    "SUZHOU",
]

META_FIELDS = [
    "record_id",
    "源文件",
    "源工作表",
    "源Excel行号",
    "trim_canonical_record_id",
    "trim_canonical_source_location",
    "trim_duplicate_group_id",
    "trim_duplicate_group_size",
    "trim_duplicate_ordinal",
    "trim_dedup_keep",
    "trim_is_duplicate_extra",
    "whitespace_equivalent_record_id",
    "whitespace_equivalent_source_location",
    "whitespace_equivalent_group_size",
    "whitespace_equivalent_ordinal",
    "whitespace_equivalent_keep",
    "cross_file_exact_group",
    "product_layer",
    "product_scope_disposition",
    "product_layer_reason",
    "route_layer",
    "buyer_geo_class",
    "platform_china_flag",
    "taxed_origin_flag",
    "direct_taxed_origin_china_flag",
    "hdc_destination_conflict_flag",
    "corrected_non_hdc_china_flag",
    "mainland_name_china_flag",
    "hwaseung_non_hdc_china_flag",
    "chaoju_china_flag",
    "hdc_to_hwaseung_a_flag",
    "recycled_pps_china_flag",
    "recycled_pps_mainland_name_flag",
    "hpp_shpp_sample_flag",
    "destination_correction",
    "origin_issue",
    "evidence_grade",
    "evidence_reason",
    "counterevidence_or_alternative",
    "decisive_data_gap",
]


def text_trim(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def text_ws(value: Any) -> str:
    return " ".join(text_trim(value).split())


def fold(value: Any) -> str:
    value_nfkd = unicodedata.normalize("NFKD", text_ws(value))
    no_marks = "".join(ch for ch in value_nfkd if not unicodedata.combining(ch))
    return re.sub(r"\s+", " ", no_marks).strip().upper()


def entity_key(value: Any) -> str:
    return re.sub(r"[^A-Z0-9]", "", fold(value))


def stable_digest(values: tuple[str, ...]) -> str:
    payload = json.dumps(values, ensure_ascii=False, separators=(",", ":"))
    return sha256(payload.encode("utf-8")).hexdigest().upper()


def parse_decimal(value: Any) -> Decimal | None:
    if value in (None, ""):
        return None
    try:
        return Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, ValueError):
        return None


def decimal_text(value: Decimal) -> str:
    result = format(value, "f")
    if "." in result:
        result = result.rstrip("0").rstrip(".")
    return result or "0"


def product_layer(description: Any) -> tuple[str, str, str]:
    desc = fold(description)
    modified = (
        re.search(
            r"\b(COMPOUND|COMPOUNDED|GF\s*\d|GLASS FIB(?:ER|RE)|CALCIUM CARBONATE|"
            r"MINERAL|ALUMINA|CARBON BLACK|COPPER|ZINC OXIDE|INORGANIC SALT)\b",
            desc,
        )
        or any(x in desc for x in ["65997-17-3", "1317-65-3", "471-34-1", "1333-86-4", "7440-50-8"])
        or "THANH PHAN" in desc
        or "TP:" in desc
    )
    if modified:
        return (
            "改性/复合PPS",
            "措施范围候选",
            "出现复合、玻纤/矿物/炭黑/金属填料或多组分CAS信号；中国措施范围明确包括改性、混合和加填料的PPS组合物。",
        )
    if re.search(r"\b(NEAT|UNFILLED|PURE)\b", desc) or "NGUYEN LIEU SAN XUAT SAN PHAM HAT NHUA" in desc:
        return (
            "纯/基础PPS树脂明确",
            "措施范围候选",
            "描述明确使用纯、未填充、neat或基础树脂生产原料表述。",
        )
    if (
        re.search(r"\b(RESIN|GRANULES?|PELLETS?|POWDER|POLYMER RESIN|HAT NHUA)\b", desc)
        or "NGUYEN SINH" in desc
        or "PRIMARY FORM" in desc
    ):
        return (
            "PPS树脂/初级形态明确、填充状态不详",
            "措施范围候选",
            "描述明确为树脂、粒料、粉末或初级形态，但仅凭文本不能完整判断填充状态。",
        )
    return (
        "PPS明确但形态待核",
        "措施范围待核",
        "描述出现PPS全称、独立缩写、品牌或牌号，但缺少足够的树脂/组合物形态词；需申报要素、TDS或实物确认。",
    )


def buyer_geo_class(buyer: Any) -> str:
    value = fold(buyer)
    key = entity_key(buyer)
    if key == HDC:
        return "韩国HDC实体—平台目的国冲突"
    if any(token in value for token in CHINA_CITY_TOKENS):
        return "中国内地实体名线索（需地址/提单终核）"
    if "HONG KONG" in value or "HONGKONG" in value or "(HK)" in value:
        return "香港实体名线索"
    return "买方国别待核"


def is_recycled(description: Any) -> bool:
    desc = fold(description)
    return bool(
        re.search(r"\bRECYC(?:LE|LED|LING)\b|\bRECYCLE\s+PELLETS?\b|\bREPROCESSED\b", desc)
        or "TAI SINH" in desc
    )


def route_and_evidence(raw: dict[str, Any]) -> dict[str, Any]:
    destination = fold(raw["目的国/地区"])
    origin = fold(raw["原产国/地区"])
    buyer = entity_key(raw["采购商"])
    supplier = entity_key(raw["供应商"])
    description = fold(raw["商品描述"])
    geo = buyer_geo_class(raw["采购商"])

    china = destination == "CHINA"
    taxed = origin in TAXED_ORIGINS
    hdc_conflict = china and buyer == HDC
    corrected_non_hdc = china and buyer != HDC
    mainland = corrected_non_hdc and geo.startswith("中国内地实体名线索")
    hw_non_hdc = corrected_non_hdc and supplier == HWASEUNG
    chaoju = hw_non_hdc and buyer == CHAO_JU
    hdc_hw_a = (
        destination == "VIETNAM"
        and origin == "SOUTH KOREA"
        and buyer == HWASEUNG
        and supplier == HDC
    )
    recycled = china and is_recycled(raw["商品描述"])
    recycled_mainland = recycled and mainland
    hpp_shpp = (
        corrected_non_hdc
        and "HIGHPERFORMANCEPLASTICSINDIAPRIVATELIMITED" in supplier
        and "SHPP" in buyer
        and "SHANGHAI" in buyer
    )
    direct = china and taxed

    if hdc_conflict:
        route = "平台中国→HDC韩国买方（目的国冲突/疑似返韩）"
        grade = "B+（同票公开数据纠偏线索）"
        reason = (
            "平台目的国字段为China，但买方HDC POLYALL为韩国实体；本地X200P NC 2024-08-29数量字段48,000"
            "与公开同日同品同量韩国进口记录吻合，不能作为中国B腿。"
        )
        alternative = "HDC韩国基础料赴越南真实配混后返韩，与本地HDC→Hwaseung输入链和HDC公开业务能力一致。"
        gap = "原始提单、卸货港、收货地址、越南出口报关单及韩国进口申报；平台目的国字段须逐票回源。"
        correction = "优先校正为韩国/疑似返韩；不得计入对华规模，待原提单终核。"
        origin_issue = "越南出口侧#&VN/平台Vietnam原产字段不构成中国非优惠原产地核定。"
    elif chaoju:
        route = "HDC韩国→Hwaseung越南→Chao Ju（疑似上海超聚）B腿"
        grade = "B+（实体+产品+时序，局部同牌号闭合）"
        reason = (
            "Hwaseung Vietnam向Chao Ju发送PPS；Chao Ju英文名高度对应上海超聚新材料，且其公开页面显示PPS改性造粒能力。"
            "其中2026-05-15 E5060G 12,000与6日前HDC→Hwaseung同牌号9,000形成局部闭合。"
        )
        alternative = "越南存在真实配混加工；其余J200/E1040S/E1040ST在各票前90日无同牌号HDC输入，不能认定同货转运。"
        gap = "Chao Ju地址/统一信用代码、中国进口报关单、原产地申报核定、税款缴款书、双腿提单/柜号、越南BOM和库存台账。"
        correction = "平台目的国暂保留中国，但买方实体须地址和中国报关单确认。"
        origin_issue = "若韩国3911基础料在越南配混后仍归3911，通常未发生四位税目改变；#&VN不能替代中国海关原产核定。"
    elif hpp_shpp:
        route = "HPP India→SHPP Shanghai集团内样品/小票"
        grade = "C（低量集团内线索）"
        reason = "4条合计数量字段28，HPP India与SHPP Shanghai均可映射至SABIC集团法律实体。"
        alternative = "数量极小且无受税来源→HPP India A腿，更符合集团内样品、研发或测试用途。"
        gap = "中国报关单、用途、样品/销售属性、印度生产BOM及受税来源原料台账。"
        correction = "平台目的国暂保留中国，按低量样品线索处理。"
        origin_issue = "无受税来源A腿，现阶段不能推定印度原产申报异常。"
    elif recycled:
        route = "越南回收PPS→中国（回收料路线）"
        grade = "C+（方向性线索）"
        reason = "描述明确为recycle/recycled PPS，对华B腿成立；部分买方名称可定位佛山或东莞。"
        alternative = "未发现受税来源→同一越南回收料供应商A腿；真实本地回收、废料再生是强替代解释。"
        gap = "越南废料/回收料采购、再生生产记录、能耗、实验室聚合物/灰分检测、中国申报品名、原产地核定及提单。"
        correction = "平台目的国暂保留中国；须区分再生料正常贸易与受税PPS简单转运。"
        origin_issue = "再生加工是否构成实质性改变须结合原料形态、工序、税则变化及中国海关核定。"
    elif hw_non_hdc:
        route = "Hwaseung越南→中国非HDC买方"
        grade = "C+（产品链线索，多为样品）"
        reason = "供应商为Hwaseung Vietnam且平台目的国为China；除Chao Ju外多为25/50/100等小数量字段记录。"
        alternative = "样品、客户验证、真实越南改性加工均可解释；未形成同票/同牌号/同数量双腿闭合。"
        gap = "中国报关单、收货地址、申报原产地、牌号BOM、税款缴款书及双腿提单/柜号。"
        correction = "平台目的国暂保留中国，买方国别和最终收货地址需逐票确认。"
        origin_issue = "越南出口原产字段不能替代中国非优惠原产地判断。"
    elif hdc_hw_a:
        route = "HDC韩国→Hwaseung越南（重点A腿）"
        grade = "B（实体级输入链）"
        reason = "同一HDC韩国实体向Hwaseung Vietnam持续供应韩国原产PPS基础料，可作为Chao Ju等B腿的输入池。"
        alternative = "HDC与Hwaseung存在真实配混/返韩供应链，A腿本身不证明绕道中国。"
        gap = "生产批次、BOM、原料领用、库存、双腿提单/柜号和出货客户订单。"
        correction = "非对华记录；仅作为潜在A腿输入池。"
        origin_issue = "本腿韩国原产字段与HDC韩国实体一致。"
    elif direct:
        route = "受税来源直接对华"
        grade = "A（直接贸易事实）"
        reason = "平台目的国为中国且原产国字段为受税来源。"
        alternative = "直接贸易不等于少缴税，须核实际税率和缴税文件。"
        gap = "中国报关单、生产商、原产地证、反倾销税率适用及税款缴款书。"
        correction = "无目的国校正。"
        origin_issue = "需核申报原产地和生产商税率。"
    elif china:
        route = "其他第三国来源→中国（非HDC）"
        grade = "C（方向性线索）"
        reason = "平台目的国为中国、原产字段为非受税来源，但单一B腿不足以证明绕道。"
        alternative = "第三国真实生产、改性、回收或贸易均可能。"
        gap = "买方地址、中国报关单、原产地核定、生产能力、BOM、库存和双腿提单。"
        correction = "平台目的国暂保留中国，仍需原提单和中国申报终核。"
        origin_issue = "第三国原产字段不等于中国海关已接受该原产地。"
    elif taxed:
        route = "受税来源→第三国（全球A腿背景）"
        grade = "C（背景记录）"
        reason = "受税来源PPS流向非中国市场，可用于建立第三国输入池。"
        alternative = "未与对华B腿在实体、牌号、数量、日期或提单层面闭合时，不得认定转运。"
        gap = "第三国实体映射、牌号、批次、数量单位、双腿提单/柜号、库存及加工记录。"
        correction = "非对华记录。"
        origin_issue = "本腿仅按平台原产字段归类。"
    else:
        route = "其他全球PPS基线记录"
        grade = "D（背景/不直接适用）"
        reason = "本行既非平台对华记录，也非受税来源输入记录，保留作全量贸易基线。"
        alternative = "正常第三国生产、消费或区域贸易。"
        gap = "仅在与重点实体或B腿形成进一步关联时补调提单、生产和原产数据。"
        correction = "无目的国校正。"
        origin_issue = "无当前直接原产争议结论。"

    return {
        "route_layer": route,
        "buyer_geo_class": geo,
        "platform_china_flag": china,
        "taxed_origin_flag": taxed,
        "direct_taxed_origin_china_flag": direct,
        "hdc_destination_conflict_flag": hdc_conflict,
        "corrected_non_hdc_china_flag": corrected_non_hdc,
        "mainland_name_china_flag": mainland,
        "hwaseung_non_hdc_china_flag": hw_non_hdc,
        "chaoju_china_flag": chaoju,
        "hdc_to_hwaseung_a_flag": hdc_hw_a,
        "recycled_pps_china_flag": recycled,
        "recycled_pps_mainland_name_flag": recycled_mainland,
        "hpp_shpp_sample_flag": hpp_shpp,
        "destination_correction": correction,
        "origin_issue": origin_issue,
        "evidence_grade": grade,
        "evidence_reason": reason,
        "counterevidence_or_alternative": alternative,
        "decisive_data_gap": gap,
    }


def write_csv(path: Path, rows: list[dict[str, Any]], fields: list[str]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def write_json(path: Path, value: Any) -> None:
    with path.open("w", encoding="utf-8") as handle:
        json.dump(value, handle, ensure_ascii=False, indent=2, default=str)


def subset_stats(rows: list[dict[str, Any]]) -> dict[str, Any]:
    result: dict[str, Any] = {
        "record_count": len(rows),
        "date_min": min((str(row["日期"]) for row in rows), default=""),
        "date_max": max((str(row["日期"]) for row in rows), default=""),
        "origin_counts": dict(Counter(str(row["原产国/地区"] or "（空）") for row in rows)),
        "destination_counts": dict(Counter(str(row["目的国/地区"] or "（空）") for row in rows)),
        "units_currency_note": "重量、数量、金额仅为源字段；源表未给统一单位/币种，不得直接写kg、吨、美元、人民币或实际税损。",
    }
    for field in ("重量", "数量", "金额"):
        parsed = [parse_decimal(row[field]) for row in rows]
        values = [value for value in parsed if value is not None]
        result[f"{field}_numeric_count"] = len(values)
        result[f"{field}_raw_field_sum"] = decimal_text(sum(values, Decimal("0")))
    return result


def read_sources() -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    rows: list[dict[str, Any]] = []
    profiles: list[dict[str, Any]] = []
    for source in SOURCES:
        wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
        if SHEET_NAME not in wb.sheetnames:
            raise AssertionError(f"{source.name}缺少工作表{SHEET_NAME}，实际={wb.sheetnames}")
        ws = wb[SHEET_NAME]
        iterator = ws.iter_rows(values_only=True)
        headers = [text_trim(value) for value in next(iterator)]
        if headers != ORIGINAL_FIELDS:
            raise AssertionError(f"{source.name}表头变化：{headers}")
        source_rows = 0
        dates: list[str] = []
        for excel_row, values in enumerate(iterator, start=2):
            if not any(value not in (None, "") for value in values):
                continue
            if len(values) != 12:
                raise AssertionError(f"{source.name}!{SHEET_NAME}!{excel_row}列数不是12：{len(values)}")
            original = {field: value for field, value in zip(ORIGINAL_FIELDS, values)}
            original["日期"] = text_trim(original["日期"])
            rows.append(
                {
                    "_source_file": source.name,
                    "_source_sheet": SHEET_NAME,
                    "_excel_row": excel_row,
                    "_raw_values": tuple(values),
                    "_trim_key": tuple(text_trim(value) for value in values),
                    "_ws_key": tuple(text_ws(value) for value in values),
                    **original,
                }
            )
            source_rows += 1
            dates.append(text_trim(original["日期"]))
        wb.close()
        profiles.append(
            {
                "source_file": str(source),
                "source_sha256": sha256(source.read_bytes()).hexdigest(),
                "source_sheet": SHEET_NAME,
                "headers": headers,
                "data_rows": source_rows,
                "date_min": min(dates),
                "date_max": max(dates),
            }
        )
    return rows, profiles


def build() -> dict[str, Any]:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    raw_rows, profiles = read_sources()
    if [profile["data_rows"] for profile in profiles] != [6311, 1547]:
        raise AssertionError(f"源表行数异常：{profiles}")
    if len(raw_rows) != 7858:
        raise AssertionError(f"合计非空记录不是7,858：{len(raw_rows)}")

    trim_groups: dict[tuple[str, ...], list[dict[str, Any]]] = defaultdict(list)
    ws_groups: dict[tuple[str, ...], list[dict[str, Any]]] = defaultdict(list)
    for row in raw_rows:
        trim_groups[row["_trim_key"]].append(row)
        ws_groups[row["_ws_key"]].append(row)

    trim_order = {key: index for index, key in enumerate(trim_groups, start=1)}
    ws_order = {key: index for index, key in enumerate(ws_groups, start=1)}
    standardized: list[dict[str, Any]] = []

    for raw_index, raw in enumerate(raw_rows, start=1):
        trim_members = trim_groups[raw["_trim_key"]]
        ws_members = ws_groups[raw["_ws_key"]]
        trim_ordinal = trim_members.index(raw) + 1
        ws_ordinal = ws_members.index(raw) + 1
        trim_id = f"PPS-TRIM-{trim_order[raw['_trim_key']]:05d}-{stable_digest(raw['_trim_key'])[:12]}"
        ws_id = f"PPS-WS-{ws_order[raw['_ws_key']]:05d}-{stable_digest(raw['_ws_key'])[:12]}"
        trim_first = trim_members[0]
        ws_first = ws_members[0]
        product_name, disposition, product_reason = product_layer(raw["商品描述"])
        original = {field: raw[field] for field in ORIGINAL_FIELDS}
        row = {
            "record_id": f"PPS-RAW-{raw_index:05d}",
            "源文件": raw["_source_file"],
            "源工作表": raw["_source_sheet"],
            "源Excel行号": raw["_excel_row"],
            "trim_canonical_record_id": trim_id,
            "trim_canonical_source_location": f"{trim_first['_source_file']}!{trim_first['_source_sheet']}!{trim_first['_excel_row']}",
            "trim_duplicate_group_id": trim_id,
            "trim_duplicate_group_size": len(trim_members),
            "trim_duplicate_ordinal": trim_ordinal,
            "trim_dedup_keep": trim_ordinal == 1,
            "trim_is_duplicate_extra": trim_ordinal > 1,
            "whitespace_equivalent_record_id": ws_id,
            "whitespace_equivalent_source_location": f"{ws_first['_source_file']}!{ws_first['_source_sheet']}!{ws_first['_excel_row']}",
            "whitespace_equivalent_group_size": len(ws_members),
            "whitespace_equivalent_ordinal": ws_ordinal,
            "whitespace_equivalent_keep": ws_ordinal == 1,
            "cross_file_exact_group": len({member["_source_file"] for member in trim_members}) > 1,
            "product_layer": product_name,
            "product_scope_disposition": disposition,
            "product_layer_reason": product_reason,
            **route_and_evidence(original),
            **original,
        }
        standardized.append(row)

    canonical = [row for row in standardized if row["trim_dedup_keep"] is True]
    whitespace_canonical = [row for row in standardized if row["whitespace_equivalent_keep"] is True]
    duplicate_audit = [row for row in standardized if int(row["trim_duplicate_group_size"]) > 1]
    whitespace_collision = [
        row
        for row in standardized
        if len({member["_trim_key"] for member in ws_groups[raw_rows[int(str(row["record_id"]).split("-")[-1]) - 1]["_ws_key"]]}) > 1
    ]

    if len(canonical) != 6726:
        raise AssertionError(f"12字段trim精确唯一数不是6,726：{len(canonical)}")
    if len(whitespace_canonical) != 6725:
        raise AssertionError(f"空白规范化等价唯一数不是6,725：{len(whitespace_canonical)}")

    subsets: dict[str, list[dict[str, Any]]] = {
        "PPS_平台目的国中国420条": [row for row in canonical if row["platform_china_flag"] is True],
        "PPS_HDC目的国冲突314条": [row for row in canonical if row["hdc_destination_conflict_flag"] is True],
        "PPS_纠偏后非HDC平台中国106条": [row for row in canonical if row["corrected_non_hdc_china_flag"] is True],
        "PPS_中国内地实体名49条": [row for row in canonical if row["mainland_name_china_flag"] is True],
        "PPS_Hwaseung非HDC对华39条": [row for row in canonical if row["hwaseung_non_hdc_china_flag"] is True],
        "PPS_ChaoJu重点B腿9条": [row for row in canonical if row["chaoju_china_flag"] is True],
        "PPS_HDC至Hwaseung重点A腿130条": [row for row in canonical if row["hdc_to_hwaseung_a_flag"] is True],
        "PPS_回收料对华全部": [row for row in canonical if row["recycled_pps_china_flag"] is True],
        "PPS_回收料中国内地实体名": [row for row in canonical if row["recycled_pps_mainland_name_flag"] is True],
    }
    expected_counts = {
        "PPS_平台目的国中国420条": 420,
        "PPS_HDC目的国冲突314条": 314,
        "PPS_纠偏后非HDC平台中国106条": 106,
        "PPS_中国内地实体名49条": 49,
        "PPS_Hwaseung非HDC对华39条": 39,
        "PPS_ChaoJu重点B腿9条": 9,
        "PPS_HDC至Hwaseung重点A腿130条": 130,
        "PPS_回收料中国内地实体名": 16,
    }
    for name, expected in expected_counts.items():
        if len(subsets[name]) != expected:
            raise AssertionError(f"{name}计数异常：{len(subsets[name])} != {expected}")

    full_fields = META_FIELDS + ORIGINAL_FIELDS
    duplicate_fields = [
        "trim_duplicate_group_id",
        "trim_duplicate_group_size",
        "trim_canonical_source_location",
        "源文件",
        "源工作表",
        "源Excel行号",
        "trim_duplicate_ordinal",
        "trim_dedup_keep",
        "whitespace_equivalent_record_id",
        "whitespace_equivalent_group_size",
    ] + [field for field in full_fields if field not in {
        "trim_duplicate_group_id", "trim_duplicate_group_size", "trim_canonical_source_location",
        "源文件", "源工作表", "源Excel行号", "trim_duplicate_ordinal", "trim_dedup_keep",
        "whitespace_equivalent_record_id", "whitespace_equivalent_group_size"
    }]

    deliverables = {"PPS_易迅逐票标准化_全量7858条": standardized, **subsets}
    generated_pairs: list[tuple[str, list[dict[str, Any]]]] = []
    for stem, rows in deliverables.items():
        stem_with_count = stem
        if stem in {"PPS_回收料对华全部", "PPS_回收料中国内地实体名"}:
            stem_with_count = f"{stem}{len(rows)}条"
        write_csv(OUTPUT_DIR / f"{stem_with_count}.csv", rows, full_fields)
        write_json(OUTPUT_DIR / f"{stem_with_count}.json", rows)
        generated_pairs.append((stem_with_count, rows))

    write_csv(OUTPUT_DIR / "PPS_全12字段trim重复审计.csv", duplicate_audit, duplicate_fields)
    write_json(OUTPUT_DIR / "PPS_全12字段trim重复审计.json", duplicate_audit)
    write_csv(OUTPUT_DIR / "PPS_空白规范化等价差异审计4行.csv", whitespace_collision, duplicate_fields)
    write_json(OUTPUT_DIR / "PPS_空白规范化等价差异审计4行.json", whitespace_collision)

    trim_duplicate_groups = [members for members in trim_groups.values() if len(members) > 1]
    ws_duplicate_groups = [members for members in ws_groups.values() if len(members) > 1]
    cross_file_trim_groups = [members for members in trim_duplicate_groups if len({member["_source_file"] for member in members}) > 1]
    product_unique_counts = Counter(str(row["product_layer"]) for row in canonical)
    product_raw_counts = Counter(str(row["product_layer"]) for row in standardized)
    route_unique_counts = Counter(str(row["route_layer"]) for row in canonical)

    chaoju = subsets["PPS_ChaoJu重点B腿9条"]
    chaoju_amount = sum((parse_decimal(row["金额"]) or Decimal("0") for row in chaoju), Decimal("0"))
    hdc_rate = Decimal("0.327")
    vat_rate = Decimal("0.13")
    conditional_coefficient = hdc_rate * (Decimal("1") + vat_rate)

    generated_at = datetime.now().astimezone().isoformat(timespec="seconds")
    summary: dict[str, Any] = {
        "delivery_id": "PPS-39119000-20240806-20260716-LOCAL-FULL-AUDIT",
        "generated_at": generated_at,
        "output_directory": str(OUTPUT_DIR),
        "source_profiles": profiles,
        "source_query_metadata": "源文件未附检索条件元数据；据描述分布推断表1接近POLYPHENYLENE SULFIDE全称检索、表2接近PPS缩写检索，仅为推断。",
        "source_raw_rows": len(standardized),
        "date_min": min(str(row["日期"]) for row in standardized),
        "date_max": max(str(row["日期"]) for row in standardized),
        "dedup_policy": {
            "primary": "完整12个原字段逐值仅去首尾空白后完全相等；首条源位置作trim canonical，原始7,858行全部保留、不删除。",
            "primary_unique_rows": len(canonical),
            "primary_duplicate_extra_rows": len(standardized) - len(canonical),
            "primary_duplicate_group_count": len(trim_duplicate_groups),
            "primary_largest_group_size": max(map(len, trim_duplicate_groups), default=0),
            "secondary": "在trim基础上把字段内连续空白折叠为单空格，仅作等价审计，不覆盖主口径。",
            "secondary_equivalent_unique_rows": len(whitespace_canonical),
            "secondary_duplicate_extra_rows": len(standardized) - len(whitespace_canonical),
            "secondary_duplicate_group_count": len(ws_duplicate_groups),
            "whitespace_variant_difference": len(canonical) - len(whitespace_canonical),
            "whitespace_collision_expanded_rows": len(whitespace_collision),
            "cross_file_exact_overlap_group_count": len(cross_file_trim_groups),
        },
        "old_json_correction": {
            "old_raw_rows": 7858,
            "old_regex_matched_raw_rows": 7856,
            "old_unique_rows_after_regex_and_old_dedup": 6723,
            "old_excluded_same_hs_rows_label_was_misleading": 1135,
            "actual_equivalent_duplicate_extras_under_secondary_method": 1133,
            "actual_regex_misses": 2,
            "regex_miss_locations": ["PPS_391190_2.xlsx!数据列表!7", "PPS_391190_2.xlsx!数据列表!326"],
        },
        "product_layer_counts_trim_unique": dict(product_unique_counts),
        "product_layer_counts_all_source_rows": dict(product_raw_counts),
        "route_layer_counts_trim_unique": dict(route_unique_counts),
        "subsets": {name: subset_stats(rows) for name, rows in subsets.items()},
        "key_findings": {
            "direct_taxed_origin_to_china_trim_unique": sum(row["direct_taxed_origin_china_flag"] is True for row in canonical),
            "platform_china_trim_unique": len(subsets["PPS_平台目的国中国420条"]),
            "hdc_destination_conflict_trim_unique": len(subsets["PPS_HDC目的国冲突314条"]),
            "corrected_non_hdc_platform_china_trim_unique": len(subsets["PPS_纠偏后非HDC平台中国106条"]),
            "chaoju_trim_unique": len(chaoju),
            "chaoju_quantity_raw_field_sum": subset_stats(chaoju)["数量_raw_field_sum"],
            "chaoju_amount_raw_field_sum": decimal_text(chaoju_amount),
            "chaoju_conditional_tax_scenario": {
                "assumptions": "仅在Chao Ju为中国实际进口人、非优惠原产仍为韩国、适用HDC 32.7%、未缴反倾销税、金额字段可代理完税价时成立。",
                "anti_dumping_rate": decimal_text(hdc_rate),
                "incremental_vat_rate_on_ad": decimal_text(hdc_rate * vat_rate),
                "combined_conditional_coefficient": decimal_text(conditional_coefficient),
                "conditional_amount_same_raw_currency_units": decimal_text(chaoju_amount * conditional_coefficient),
                "warning": "不是实际欠税或税损；平台金额币种和中国完税价格未知。",
            },
            "strongest_specific_chain": "2026-05-09 HDC→Hwaseung E5060G BK 6,000 + NC 3,000；2026-05-15 Hwaseung→Chao Ju E5060G BR 12,000，6日间隔、同牌号局部闭合但差3,000且颜色后缀不同。",
            "strongest_counterevidence": "HDC平台中国314条中的同票公开记录指向韩国；HDC→Hwaseung 130条输入与返韩输出更符合真实越南配混/返韩供应链。",
        },
        "units_and_currency_caveat": "所有重量、数量、金额均保留源字段；未给统一单位和币种，字段求和仅作机械对账，不得改写为kg、吨、美元、人民币或实际税损。",
        "evidence_conclusion": "现有数据形成Chao Ju B+核查线索，不构成绕道、原产地虚假或逃避反倾销税的定案证据。HDC 314条旧对华结论须撤回/纠偏。",
        "generated_files": sorted(
            [f"{stem}.{suffix}" for stem, _rows in generated_pairs for suffix in ("csv", "json")]
            + [
                "PPS_全12字段trim重复审计.csv",
                "PPS_全12字段trim重复审计.json",
                "PPS_空白规范化等价差异审计4行.csv",
                "PPS_空白规范化等价差异审计4行.json",
                "PPS_交付摘要.json",
                "PPS_QA校验.json",
            ]
        ),
        "self_checks": {
            "source_rows_7858": len(standardized) == 7858,
            "trim_unique_6726": len(canonical) == 6726,
            "whitespace_equivalent_unique_6725": len(whitespace_canonical) == 6725,
            "product_layers_reconcile": sum(product_unique_counts.values()) == len(canonical),
            "routes_reconcile": sum(route_unique_counts.values()) == len(canonical),
            **{f"subset_{name}": len(subsets[name]) == count for name, count in expected_counts.items()},
        },
    }
    if not all(summary["self_checks"].values()):
        raise AssertionError(f"摘要自检失败：{summary['self_checks']}")
    write_json(OUTPUT_DIR / "PPS_交付摘要.json", summary)

    qa_files: dict[str, Any] = {}
    for stem, expected_rows in generated_pairs:
        csv_path = OUTPUT_DIR / f"{stem}.csv"
        json_path = OUTPUT_DIR / f"{stem}.json"
        with csv_path.open("r", encoding="utf-8-sig", newline="") as handle:
            csv_count = sum(1 for _ in csv.DictReader(handle))
        with json_path.open("r", encoding="utf-8") as handle:
            json_count = len(json.load(handle))
        qa_files[stem] = {
            "expected_rows": len(expected_rows),
            "csv_rows": csv_count,
            "json_rows": json_count,
            "pass": csv_count == json_count == len(expected_rows),
            "csv_sha256": sha256(csv_path.read_bytes()).hexdigest(),
            "json_sha256": sha256(json_path.read_bytes()).hexdigest(),
        }
    qa = {
        "generated_at": generated_at,
        "all_pass": all(item["pass"] for item in qa_files.values()),
        "file_row_checks": qa_files,
        "summary_self_checks": summary["self_checks"],
        "duplicate_audit_rows": len(duplicate_audit),
        "whitespace_collision_rows": len(whitespace_collision),
    }
    if not qa["all_pass"]:
        raise AssertionError(f"CSV/JSON行数复读校验失败：{qa_files}")
    write_json(OUTPUT_DIR / "PPS_QA校验.json", qa)
    return summary


if __name__ == "__main__":
    result = build()
    print(json.dumps({
        "output_directory": result["output_directory"],
        "source_raw_rows": result["source_raw_rows"],
        "dedup_policy": result["dedup_policy"],
        "subset_counts": {name: item["record_count"] for name, item in result["subsets"].items()},
        "self_checks": result["self_checks"],
    }, ensure_ascii=False, indent=2))
