from __future__ import annotations

import csv
import json
import re
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


BASE = Path(r"D:\易迅数据\反倾销税深度分析报告\19_丁腈橡胶")
LEDGER = BASE / "丁腈橡胶_本地预筛文件台账.csv"

HS_RE = re.compile(r"(?<!\d)400259(?:10|90|00)?(?!\d)", re.I)
CHEM_RE = re.compile(
    r"ACRYLONITRILE\s*[-–—]?\s*BUTADIENE\s+(?:RUBBER|COPOLYMER)|"
    r"BUTADIENE\s*[-–—]?\s*ACRYLONITRILE\s+(?:RUBBER|COPOLYMER)|"
    r"丁腈橡胶|丁二烯\s*[-–—]?\s*丙烯腈.*橡胶",
    re.I,
)
NBR_RE = re.compile(r"(?<![A-Z0-9])NBR(?![A-Z0-9])", re.I)
KNB_RE = re.compile(r"(?<![A-Z0-9])K(?:UMHO\s*)?NBL?[-\s]?[0-9]{2,4}[A-Z0-9-]*", re.I)
NIPOL_RE = re.compile(r"(?<![A-Z0-9])NIPOL(?:\s+[A-Z0-9-]+)?", re.I)
RUBBER_RE = re.compile(r"RUBBER|CAUCHO|KAUTSCHUK|GOMMA|CAO\s+SU|橡胶", re.I)

LATEX_RE = re.compile(r"LATEX|胶乳|乳胶", re.I)
HNBR_RE = re.compile(r"(?<![A-Z0-9])HNBR(?![A-Z0-9])|HYDROGENATED\s+NITRILE|氢化丁腈", re.I)
XNBR_RE = re.compile(r"(?<![A-Z0-9])XNBR(?![A-Z0-9])|CARBOXYL(?:ATED)?\s+NITRILE|羧基丁腈", re.I)
PRODUCT_RE = re.compile(
    r"GLOVE|手套|SEAL(?:ING)?|O[- ]?RING|GASKET|HOSE|TUBE|BELT|FOAM|"
    r"密封|胶管|软管|输送带|鞋|SOLE|ADHESIVE|涂层|COATED|ARTICLE|PARTS?",
    re.I,
)
META_RE = re.compile(r"反倾销|税率|报告|政策|查询条件|调查期|商务部|审计|证据等级")

DATE_RE = re.compile(r"20\d{2}[-/]\d{1,2}[-/]\d{1,2}")


def clean(v: Any) -> str:
    if v is None:
        return ""
    return re.sub(r"\s+", " ", str(v)).strip()


def classify(text: str) -> tuple[str, str, list[str]]:
    hits: list[str] = []
    hs = bool(HS_RE.search(text))
    chem = bool(CHEM_RE.search(text))
    nbr = bool(NBR_RE.search(text) and RUBBER_RE.search(text))
    knb = bool(KNB_RE.search(text))
    nipol = bool(NIPOL_RE.search(text) and RUBBER_RE.search(text))
    if hs:
        hits.append("HS400259")
    if chem:
        hits.append("化学名称")
    if nbr:
        hits.append("NBR+橡胶")
    if knb:
        hits.append("KNB牌号")
    if nipol:
        hits.append("NIPOL+橡胶")

    has_product = chem or nbr or knb or nipol
    if not (hs or has_product):
        return "不命中", "", hits
    if META_RE.search(text) and not DATE_RE.search(text):
        return "项目元数据/政策引用", "非逐票贸易记录", hits
    if HNBR_RE.search(text):
        return "氢化丁腈HNBR", "与涉案普通NBR分开核定范围及税号", hits
    if XNBR_RE.search(text):
        return "羧基丁腈XNBR", "与涉案普通NBR分开核定范围及税号", hits
    if LATEX_RE.search(text):
        return "丁腈胶乳", "通常归入400251，现行措施税号为40025910/90", hits
    if PRODUCT_RE.search(text) and not (hs and has_product and re.search(r"BALE|BLOCK|LUMP|POWDER|RAW|PRIMARY|原胶|块|粉", text, re.I)):
        return "下游制品/混合用途待排除", "不能仅因含NBR字样并入原胶措施范围", hits
    if has_product:
        return "NBR原胶/牌号候选", "需以中国进口HS、COA、成分和生产商确认", hits
    return "仅税号待货描", "HS命中但货描不足以确认NBR", hits


def source_kind(path: Path) -> str:
    s = str(path)
    if "_易迅页面采集" in s:
        return "易迅原始页面采集"
    if path.parent == Path(r"D:\易迅数据"):
        return "用户下载原始文件"
    if path.name in {"中国反倾销税商品清单_2026-08-11.xlsx", "00_全商品查询与报告进度台账.csv"}:
        return "项目元数据"
    return "项目派生输出/QA"


def row_record(path: Path, location: str, row_no: str, values: list[Any], headers: list[str] | None = None):
    vals = [clean(x) for x in values]
    text = " | ".join(x for x in vals if x)
    cls, reason, hits = classify(text)
    if cls == "不命中":
        return None
    mapped = {}
    if headers and len(headers) == len(vals):
        mapped = {clean(k): v for k, v in zip(headers, vals)}
    return {
        "source_file": str(path),
        "source_kind": source_kind(path),
        "location": location,
        "row_no": row_no,
        "classification": cls,
        "classification_reason": reason,
        "match_basis": ";".join(hits),
        "record_text": text[:12000],
        "record_json": json.dumps(mapped, ensure_ascii=False) if mapped else "",
    }


def iter_xlsx(path: Path) -> Iterable[dict[str, str]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        it = ws.iter_rows(values_only=True)
        try:
            first = list(next(it))
        except StopIteration:
            continue
        headers = [clean(x) for x in first]
        # Header-like if it contains common trade column labels.
        header_like = any(re.search(r"日期|HS|描述|商品|买方|卖方|采购|供应|date|description", x, re.I) for x in headers)
        if not header_like:
            rec = row_record(path, ws.title, "1", first)
            if rec:
                yield rec
            headers = None
        for idx, row in enumerate(it, start=2):
            rec = row_record(path, ws.title, str(idx), list(row), headers)
            if rec:
                yield rec
    wb.close()


def iter_csv(path: Path) -> Iterable[dict[str, str]]:
    for enc in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            with path.open("r", encoding=enc, newline="", errors="strict") as f:
                reader = csv.reader(f)
                rows = iter(reader)
                try:
                    headers = next(rows)
                except StopIteration:
                    return
                for idx, row in enumerate(rows, start=2):
                    rec = row_record(path, "CSV", str(idx), row, headers if len(headers) == len(row) else None)
                    if rec:
                        yield rec
            return
        except UnicodeDecodeError:
            continue


def walk_json_rows(obj: Any, path: str = "$") -> Iterable[tuple[str, list[Any], list[str] | None]]:
    if isinstance(obj, dict):
        # EasyXun page capture.
        if isinstance(obj.get("pages"), list):
            global_headers = obj.get("headers") if isinstance(obj.get("headers"), list) else None
            for page in obj["pages"]:
                if not isinstance(page, dict):
                    continue
                pnum = page.get("page", "?")
                headers = page.get("headers") if isinstance(page.get("headers"), list) else global_headers
                for i, row in enumerate(page.get("rows", []), start=1):
                    if isinstance(row, list):
                        yield f"{path}.pages[page={pnum}].rows[{i}]", row, headers
            return
        # A dict containing a trade-like record.
        joined_keys = " ".join(map(str, obj.keys()))
        if re.search(r"日期|HS|描述|商品|买方|卖方|采购|供应|date|description|buyer|seller", joined_keys, re.I):
            yield path, list(obj.values()), list(obj.keys())
        for k, v in obj.items():
            if isinstance(v, (dict, list)):
                yield from walk_json_rows(v, f"{path}.{k}")
    elif isinstance(obj, list):
        if obj and all(not isinstance(x, (dict, list)) for x in obj):
            yield path, obj, None
        else:
            for i, v in enumerate(obj):
                yield from walk_json_rows(v, f"{path}[{i}]")


def iter_json(path: Path) -> Iterable[dict[str, str]]:
    obj = json.loads(path.read_text(encoding="utf-8-sig", errors="replace"))
    for idx, (loc, row, headers) in enumerate(walk_json_rows(obj), start=1):
        rec = row_record(path, loc, str(idx), row, headers)
        if rec:
            yield rec


def main() -> None:
    BASE.mkdir(parents=True, exist_ok=True)
    with LEDGER.open("r", encoding="utf-8-sig", newline="") as f:
        candidate_files = [Path(r["file"]) for r in csv.DictReader(f) if r["prefilter_class"] != "no_hit"]

    hits: list[dict[str, str]] = []
    audits: list[dict[str, Any]] = []
    for p in candidate_files:
        before = len(hits)
        error = ""
        try:
            if p.suffix.lower() == ".xlsx":
                hits.extend(iter_xlsx(p))
            elif p.suffix.lower() == ".csv":
                hits.extend(iter_csv(p))
            elif p.suffix.lower() == ".json":
                hits.extend(iter_json(p))
        except Exception as exc:
            error = f"{type(exc).__name__}: {exc}"
        subset = hits[before:]
        audits.append({
            "source_file": str(p),
            "source_kind": source_kind(p),
            "match_records": len(subset),
            "class_counts": json.dumps(Counter(x["classification"] for x in subset), ensure_ascii=False, sort_keys=True),
            "error": error,
        })

    # Deduplicate identical visible records, preserving provenance separately.
    groups: dict[tuple[str, str], list[dict[str, str]]] = defaultdict(list)
    for h in hits:
        groups[(h["classification"], h["record_text"])].append(h)
    unique: list[dict[str, str]] = []
    for n, ((cls, text), group) in enumerate(groups.items(), start=1):
        original_group = [x for x in group if x["source_kind"] in {"易迅原始页面采集", "用户下载原始文件"}]
        primary = (original_group or group)[0].copy()
        primary.update({
            "dedup_id": f"NBR-{n:05d}",
            "occurrence_count": len(group),
            "source_files_count": len({x["source_file"] for x in group}),
            "all_provenance": json.dumps(
                [{"file": x["source_file"], "location": x["location"], "row_no": x["row_no"], "kind": x["source_kind"]} for x in group],
                ensure_ascii=False,
            ),
        })
        unique.append(primary)

    fields_hits = ["source_file", "source_kind", "location", "row_no", "classification", "classification_reason", "match_basis", "record_text", "record_json"]
    fields_unique = ["dedup_id", "occurrence_count", "source_files_count"] + fields_hits + ["all_provenance"]
    fields_audit = ["source_file", "source_kind", "match_records", "class_counts", "error"]

    def write_csv(name: str, rows: list[dict[str, Any]], fields: list[str]):
        with (BASE / name).open("w", encoding="utf-8-sig", newline="") as f:
            w = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
            w.writeheader(); w.writerows(rows)

    write_csv("丁腈橡胶_本地深扫全部命中.csv", hits, fields_hits)
    write_csv("丁腈橡胶_本地深扫去重记录.csv", unique, fields_unique)
    write_csv("丁腈橡胶_本地深扫文件审计.csv", audits, fields_audit)
    (BASE / "丁腈橡胶_本地深扫全部命中.json").write_text(json.dumps(hits, ensure_ascii=False, indent=2), encoding="utf-8")
    (BASE / "丁腈橡胶_本地深扫去重记录.json").write_text(json.dumps(unique, ensure_ascii=False, indent=2), encoding="utf-8")

    original_hits = [x for x in hits if x["source_kind"] in {"易迅原始页面采集", "用户下载原始文件"}]
    original_unique = []
    seen = set()
    for x in original_hits:
        key = (x["classification"], x["record_text"])
        if key not in seen:
            seen.add(key); original_unique.append(x)

    summary = {
        "candidate_files": len(candidate_files),
        "files_scanned": len(audits),
        "files_with_errors": sum(bool(x["error"]) for x in audits),
        "all_hit_occurrences": len(hits),
        "all_visible_unique": len(unique),
        "all_class_counts": dict(Counter(x["classification"] for x in hits)),
        "unique_class_counts": dict(Counter(x["classification"] for x in unique)),
        "original_hit_occurrences": len(original_hits),
        "original_visible_unique": len(original_unique),
        "original_unique_class_counts": dict(Counter(x["classification"] for x in original_unique)),
        "qa": {
            "candidate_files_equal_audit_rows": len(candidate_files) == len(audits),
            "unique_not_greater_than_occurrences": len(unique) <= len(hits),
            "no_file_errors": not any(x["error"] for x in audits),
        },
    }
    (BASE / "丁腈橡胶_本地深扫摘要.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
