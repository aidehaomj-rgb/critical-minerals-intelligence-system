#!/usr/bin/env python3
"""Read-only, content-level inventory for local phenol trade data.

The script scans XLSX/XLS/CSV/JSON files under ``D:\易迅数据``.  It preserves
the source path plus worksheet/row or JSON path for every content hit, separates
single phenol from common derivatives, and writes only derived audit artifacts
under product folder 15.  Source exports are never modified.
"""

from __future__ import annotations

import csv
import io
import json
import re
import sys
import zipfile
from collections import Counter
from datetime import datetime, timezone, timedelta
from pathlib import Path
from typing import Any, Iterable

import openpyxl


ROOT = Path(r"D:\易迅数据")
OUT = ROOT / "反倾销税深度分析报告" / "15_苯酚"
EXTENSIONS = {".xlsx", ".xls", ".csv", ".json"}

# Identifiers with high product specificity.  HS 290711 covers phenol and its
# salts, so it is strong scope evidence but still needs the complete goods text.
STRONG_PATTERNS: list[tuple[str, re.Pattern[str]]] = [
    ("CAS 108-95-2", re.compile(
        r"(?<!\d)(?:108\s*[-./_]\s*95\s*[-./_]\s*2|108\s+95\s+2|108952)(?!\d)", re.I)),
    ("HS 290711", re.compile(r"(?<!\d)290711(?:00)?(?!\d)", re.I)),
]

NAME_PATTERNS: list[tuple[str, re.Pattern[str]]] = [
    ("苯酚", re.compile(r"苯酚", re.I)),
    ("PHENOL", re.compile(r"\bphenol\b", re.I)),
    ("HYDROXYBENZENE", re.compile(r"\bhydroxy[\s_-]*benzene\b", re.I)),
    ("CARBOLIC ACID", re.compile(r"\bcarbolic[\s_-]*acid\b", re.I)),
    ("BENZENOL", re.compile(r"\bbenzenol\b", re.I)),
    ("PHENIC ACID", re.compile(r"\bphenic[\s_-]*acid\b", re.I)),
]

# These indicate that a PHENOL/苯酚 text hit may instead be a derivative,
# mixture, resin or another antidumping product.  They are retained for manual
# review rather than silently discarded.
EXCLUSION_PATTERNS: list[tuple[str, re.Pattern[str]]] = [
    ("酚醛树脂", re.compile(r"酚醛|phenol[\s_-]*formaldehyde|phenolic[\s_-]*resin", re.I)),
    ("双酚", re.compile(r"双酚|bis[\s_-]*phenol|bisphenol", re.I)),
    ("氯酚", re.compile(r"氯(?:代)?苯?酚|chloro[\s_-]*phenol|chlorophenol", re.I)),
    ("溴酚", re.compile(r"溴(?:代)?苯?酚|bromo[\s_-]*phenol|bromophenol", re.I)),
    ("硝基酚", re.compile(r"硝基苯?酚|nitro[\s_-]*phenol|nitrophenol", re.I)),
    ("氨基酚", re.compile(r"氨基苯?酚|amino[\s_-]*phenol|aminophenol", re.I)),
    ("甲酚", re.compile(r"甲酚|cresol|methyl[\s_-]*phenol", re.I)),
    ("二甲酚", re.compile(r"二甲酚|xylenol", re.I)),
    ("烷基酚", re.compile(
        r"烷基苯?酚|alkyl[\s_-]*phenol|叔丁基苯?酚|tert[\s_-]*butyl[\s_-]*phenol|"
        r"辛基苯?酚|octyl[\s_-]*phenol|壬基苯?酚|nonyl[\s_-]*phenol", re.I)),
    ("多羟基苯", re.compile(
        r"邻苯二酚|间苯二酚|对苯二酚|catechol|resorcinol|hydroquinone|"
        r"dihydroxy[\s_-]*benzene", re.I)),
    ("农药/药物衍生物", re.compile(
        r"acetamido[\s_-]*phenol|acetaminophen|paracetamol|phenoxy|"
        r"phenolphthalein|phenol[\s_-]*red", re.I)),
]

PREFILTER_TOKENS = (
    "苯酚", "phenol", "hydroxybenzene", "hydroxy benzene", "carbolic acid",
    "carbolic-acid", "benzenol", "phenic acid", "108-95-2", "108 95 2",
    "108/95/2", "108.95.2", "108952", "290711",
)
PREFILTER_RX = re.compile("|".join(re.escape(x) for x in PREFILTER_TOKENS), re.I)

PROJECT_METADATA_NAMES = {
    "中国反倾销税商品清单_2026-08-11.xlsx",
    "00_全商品查询与报告进度台账.csv",
}


def norm_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat()
    return re.sub(r"\s+", " ", str(value)).strip()


def any_positive_hit(text: str) -> bool:
    return PREFILTER_RX.search(text) is not None


def match_names(text: str, patterns: list[tuple[str, re.Pattern[str]]]) -> list[str]:
    return [name for name, rx in patterns if rx.search(text)]


def classify_hit(strong: list[str], names: list[str], exclusions: list[str]) -> str:
    if strong and exclusions:
        return "strong_identifier_with_derivative_review"
    if strong:
        return "strong_target_candidate"
    if names and exclusions:
        return "excluded_or_derivative_review"
    if names:
        return "name_only_target_candidate"
    return "no_hit"


def source_class(path: Path) -> str:
    if path.name in PROJECT_METADATA_NAMES:
        return "project_metadata"
    if "反倾销税深度分析报告" in path.parts:
        if "_易迅页面采集" in path.parts:
            return "easyxun_page_capture_other_product"
        if "_分析中间数据" in path.parts:
            return "derived_analysis_other_product"
        return "project_output_or_qa"
    return "root_download_or_user_file"


def safe_decode(raw: bytes) -> tuple[str, str]:
    for enc in ("utf-8-sig", "utf-8", "gb18030", "utf-16", "cp1252"):
        try:
            return raw.decode(enc), enc
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace"), "utf-8-replace"


def xlsx_has_any_hit(path: Path) -> bool:
    try:
        with zipfile.ZipFile(path) as zf:
            for name in zf.namelist():
                if not (name.startswith("xl/") and name.endswith(".xml")):
                    continue
                text, _ = safe_decode(zf.read(name))
                if any_positive_hit(text):
                    return True
    except (zipfile.BadZipFile, OSError):
        # Let openpyxl record the concrete error below.
        return True
    return False


def row_excerpt(values: list[Any], max_chars: int = 8000) -> str:
    return " | ".join(norm_text(v) for v in values if norm_text(v))[:max_chars]


def make_hit(path: Path, fmt: str, location: str, text: str) -> dict[str, Any] | None:
    if not any_positive_hit(text):
        return None
    strong = match_names(text, STRONG_PATTERNS)
    names = match_names(text, NAME_PATTERNS)
    exclusions = match_names(text, EXCLUSION_PATTERNS)
    category = classify_hit(strong, names, exclusions)
    if category == "no_hit":
        return None
    return {
        "file": str(path),
        "source_class": source_class(path),
        "format": fmt,
        "location": location,
        "match_class": category,
        "strong_terms": strong,
        "name_terms": names,
        "exclusion_terms": exclusions,
        "text_excerpt": text,
    }


def scan_xlsx(path: Path) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    info: dict[str, Any] = {"sheets": 0, "rows_scanned": 0, "cells_scanned": 0, "errors": []}
    hits: list[dict[str, Any]] = []
    if not xlsx_has_any_hit(path):
        info["fast_scan_no_hit"] = True
        return info, hits
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
        info["sheets"] = len(wb.sheetnames)
        for ws in wb.worksheets:
            for row_no, row in enumerate(ws.iter_rows(values_only=True), start=1):
                info["rows_scanned"] += 1
                values = list(row)
                info["cells_scanned"] += len(values)
                text = row_excerpt(values)
                hit = make_hit(path, "xlsx", f"{ws.title}!row{row_no}", text)
                if hit:
                    hits.append(hit)
        wb.close()
    except Exception as exc:
        info["errors"].append(f"{type(exc).__name__}: {exc}")
    return info, hits


def scan_csv(path: Path) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    raw = path.read_bytes()
    text, enc = safe_decode(raw)
    info: dict[str, Any] = {"encoding": enc, "rows_scanned": 0, "errors": []}
    hits: list[dict[str, Any]] = []
    try:
        try:
            dialect = csv.Sniffer().sniff(text[:65536], delimiters=",\t;|")
        except csv.Error:
            dialect = csv.excel
        for row_no, row in enumerate(csv.reader(io.StringIO(text), dialect), start=1):
            info["rows_scanned"] += 1
            hit = make_hit(path, "csv", f"row{row_no}", row_excerpt(row))
            if hit:
                hits.append(hit)
    except Exception as exc:
        info["errors"].append(f"{type(exc).__name__}: {exc}")
    return info, hits


def scalar_items(obj: dict[str, Any]) -> dict[str, Any]:
    return {str(k): v for k, v in obj.items() if not isinstance(v, (dict, list))}


def json_records(obj: Any, loc: str = "$") -> Iterable[tuple[str, dict[str, Any]]]:
    if isinstance(obj, dict):
        scalars = scalar_items(obj)
        if scalars:
            yield loc, scalars
        for key, value in obj.items():
            if isinstance(value, (dict, list)):
                yield from json_records(value, f"{loc}.{key}")
    elif isinstance(obj, list):
        for idx, value in enumerate(obj):
            yield from json_records(value, f"{loc}[{idx}]")
    else:
        yield loc, {"value": obj}


def scan_json(path: Path) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    raw = path.read_bytes()
    text, enc = safe_decode(raw)
    info: dict[str, Any] = {"encoding": enc, "records_scanned": 0, "errors": []}
    hits: list[dict[str, Any]] = []
    if not any_positive_hit(text):
        info["fast_scan_no_hit"] = True
        return info, hits
    try:
        obj = json.loads(text)
        for loc, record in json_records(obj):
            info["records_scanned"] += 1
            combined = " | ".join(f"{key}={norm_text(value)}" for key, value in record.items())[:8000]
            hit = make_hit(path, "json", loc, combined)
            if hit:
                hits.append(hit)
    except Exception as exc:
        info["errors"].append(f"{type(exc).__name__}: {exc}")
    return info, hits


def scan_one(path: Path) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    base = {
        "file": str(path), "name": path.name, "extension": path.suffix.lower(),
        "size_bytes": path.stat().st_size, "modified": path.stat().st_mtime,
        "source_class": source_class(path),
    }
    if path.suffix.lower() == ".xlsx":
        detail, hits = scan_xlsx(path)
    elif path.suffix.lower() == ".csv":
        detail, hits = scan_csv(path)
    elif path.suffix.lower() == ".json":
        detail, hits = scan_json(path)
    else:
        detail, hits = {"errors": ["legacy .xls is not parsed by openpyxl"]}, []
    base.update(detail)
    base["hit_rows_or_records"] = len(hits)
    base["hit_class_counts"] = dict(Counter(hit["match_class"] for hit in hits))
    return base, hits


def write_csv(path: Path, rows: list[dict[str, Any]], fieldnames: list[str]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            clean = dict(row)
            for key, value in clean.items():
                if isinstance(value, (list, dict)):
                    clean[key] = json.dumps(value, ensure_ascii=False, separators=(",", ":"))
            writer.writerow(clean)


def main() -> int:
    files = sorted(
        path for path in ROOT.rglob("*")
        if path.is_file() and path.suffix.lower() in EXTENSIONS
        and OUT not in path.parents and not path.name.startswith("~$")
    )
    inventory: list[dict[str, Any]] = []
    hits: list[dict[str, Any]] = []
    for idx, path in enumerate(files, start=1):
        if idx == 1 or idx % 25 == 0 or idx == len(files):
            print(f"scan {idx}/{len(files)}: {path}", flush=True)
        info, file_hits = scan_one(path)
        inventory.append(info)
        hits.extend(file_hits)

    OUT.mkdir(parents=True, exist_ok=True)
    now_cn = datetime.now(timezone(timedelta(hours=8))).isoformat(timespec="seconds")
    source_candidate_classes = {"root_download_or_user_file", "easyxun_page_capture_other_product"}
    source_candidates = [
        hit for hit in hits
        if hit["source_class"] in source_candidate_classes
        and hit["match_class"] in {"strong_target_candidate", "name_only_target_candidate",
                                   "strong_identifier_with_derivative_review"}
        and Path(hit["file"]).name not in PROJECT_METADATA_NAMES
    ]
    summary = {
        "scanned_at_asia_shanghai": now_cn,
        "root": str(ROOT),
        "files_scanned": len(files),
        "extension_counts": dict(Counter(path.suffix.lower() for path in files)),
        "source_class_counts": dict(Counter(item["source_class"] for item in inventory)),
        "files_with_content_hits": sum(1 for item in inventory if item["hit_rows_or_records"]),
        "hit_rows_or_records": len(hits),
        "hit_class_counts": dict(Counter(hit["match_class"] for hit in hits)),
        "hit_source_class_counts": dict(Counter(hit["source_class"] for hit in hits)),
        "source_trade_candidate_records_before_manual_review": len(source_candidates),
        "source_trade_candidate_files": sorted({hit["file"] for hit in source_candidates}),
        "scan_errors": sum(bool(item.get("errors")) for item in inventory),
        "limitations": [
            "本地文件内容盘点不等同于易迅网页全库查询。",
            "PHENOL/苯酚名称命中仍须排除酚醛树脂、双酚、氯酚等衍生物。",
            "HS 290711包括苯酚及其盐，须结合完整货描确认措施范围。",
            "项目清单、报告正文和派生台账不作为新增独立贸易票证。",
        ],
    }

    (OUT / "苯酚_本地全盘内容审计.json").write_text(
        json.dumps({"summary": summary, "inventory": inventory, "hits": hits}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    write_csv(
        OUT / "苯酚_本地文件盘点台账.csv", inventory,
        ["file", "name", "extension", "size_bytes", "modified", "source_class",
         "sheets", "rows_scanned", "cells_scanned", "records_scanned", "encoding",
         "fast_scan_no_hit", "hit_rows_or_records", "hit_class_counts", "errors"],
    )
    write_csv(
        OUT / "苯酚_本地内容命中明细.csv", hits,
        ["file", "source_class", "format", "location", "match_class",
         "strong_terms", "name_terms", "exclusion_terms", "text_excerpt"],
    )
    (OUT / "苯酚_本地盘点摘要.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
