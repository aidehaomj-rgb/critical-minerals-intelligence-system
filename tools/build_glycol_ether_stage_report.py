from __future__ import annotations

import csv
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT_DIR = Path(r"D:\易迅数据\反倾销税深度分析报告\09_乙二醇和丙二醇的单烷基醚")
AUDIT_PATH = OUT_DIR / "单烷基醚_全量阶段审计.json"
LEADS_PATH = OUT_DIR / "单烷基醚_明确范围内重点记录.csv"
REPORT_PATH = OUT_DIR / "乙二醇和丙二醇单烷基醚_反倾销税与第三国转运风险阶段审计报告.docx"

AUDIT = json.loads(AUDIT_PATH.read_text(encoding="utf-8"))
with LEADS_PATH.open("r", encoding="utf-8-sig", newline="") as fh:
    ALL_LEADS = list(csv.DictReader(fh))

# 可见字段精确去重；重复与疑似同票另保留在工作簿中。
seen = set()
LEADS = []
for item in ALL_LEADS:
    sig = item["visible_signature"]
    if sig not in seen:
        seen.add(sig)
        LEADS.append(item)

NAVY = "17324D"
BLUE = "1F4E78"
INK = "102A43"
MUTED = "66788A"
LIGHT_BLUE = "EAF2F8"
PALE_GOLD = "FFF4CC"
PALE_RED = "FCE8E6"
PALE_GREEN = "E6F4EA"
BORDER = "CBD5E1"
RED = "9B1C1C"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def fmt(value, digits=2):
    if value in (None, ""):
        return "—"
    number = float(value)
    if digits == 0:
        return f"{number:,.0f}"
    return f"{number:,.{digits}f}"


def set_run_font(run, size=None, color=None, bold=None, italic=None, name="Calibri"):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    node = trpr.find(qn("w:tblHeader"))
    if node is None:
        node = OxmlElement("w:tblHeader")
        trpr.append(node)
    node.set(qn("w:val"), "true")


def set_cell_text(cell, value, bold=False, color="000000", size=8.6):
    cell.text = "" if value is None else str(value)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.03
        for run in paragraph.runs:
            set_run_font(run, size=size, color=color, bold=bold)


def set_table_borders(table, color=BORDER, size="4"):
    tblpr = table._tbl.tblPr
    borders = tblpr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblpr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa, total=9360):
    table.autofit = False
    tbl = table._tbl
    tblpr = tbl.tblPr
    for tag in ("w:tblW", "w:tblInd", "w:tblLayout", "w:tblCellMar"):
        old = tblpr.find(qn(tag))
        if old is not None:
            tblpr.remove(old)
    tblw = OxmlElement("w:tblW")
    tblw.set(qn("w:w"), str(total))
    tblw.set(qn("w:type"), "dxa")
    tblpr.append(tblw)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr.append(layout)
    margins = OxmlElement("w:tblCellMar")
    for tag, val in (("top", 50), ("bottom", 50), ("start", 70), ("end", 70)):
        node = OxmlElement(f"w:{tag}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    tblpr.append(margins)
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths_dxa):
            tcpr = cell._tc.get_or_add_tcPr()
            tcw = tcpr.find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                tcpr.append(tcw)
            tcw.set(qn("w:w"), str(width))
            tcw.set(qn("w:type"), "dxa")


def add_table(doc, headers, rows, widths, header_fill=LIGHT_BLUE, size=8.6, shades=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for idx, text in enumerate(headers):
        set_cell_text(table.cell(0, idx), text, bold=True, color=INK, size=size)
        set_cell_shading(table.cell(0, idx), header_fill)
    set_repeat_header(table.rows[0])
    for ridx, vals in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(vals):
            set_cell_text(cells[idx], value, size=size)
            if shades and ridx in shades:
                set_cell_shading(cells[idx], shades[ridx])
    set_table_geometry(table, widths)
    set_table_borders(table)
    return table


def add_callout(doc, title, text, fill=PALE_GOLD, title_color=INK):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, size=10.2, color=title_color, bold=True)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.1
    r2 = p2.add_run(text)
    set_run_font(r2, size=9.5, color="333333")
    set_table_geometry(table, [9360])
    set_table_borders(table, color=BORDER, size="6")
    set_repeat_header(table.rows[0])
    return table


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.07
    r = p.add_run(text)
    set_run_font(r, size=9.8, color="222222")
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.07
    r = p.add_run(text)
    set_run_font(r, size=9.8, color="222222")
    return p


def add_para(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(text)
    set_run_font(r, size=9.9, color="222222")
    return p


def add_hyperlink(paragraph, text, url):
    rid = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    for node_name, attrs in (
        ("w:rFonts", {"w:ascii": "Calibri", "w:hAnsi": "Calibri", "w:eastAsia": "Microsoft YaHei"}),
        ("w:color", {"w:val": BLUE}),
        ("w:u", {"w:val": "single"}),
    ):
        node = OxmlElement(node_name)
        for key, val in attrs.items():
            node.set(qn(key), val)
        rpr.append(node)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_source(doc, label, url, note=""):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    add_hyperlink(p, label, url)
    if note:
        r = p.add_run(f"：{note}")
        set_run_font(r, size=8.9, color=MUTED)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_1)
    run._r.append(instr)
    run._r.append(fld_char_2)
    set_run_font(run, size=8.2, color=MUTED)


def page_break(doc):
    doc.add_page_break()


doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.72)
sec.bottom_margin = Inches(0.72)
sec.left_margin = Inches(0.82)
sec.right_margin = Inches(0.82)
sec.page_width = Inches(8.27)
sec.page_height = Inches(11.69)

styles = doc.styles
styles["Normal"].font.name = "Calibri"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
styles["Normal"].font.size = Pt(10.0)
for name, size, color in (
    ("Title", 25.5, NAVY),
    ("Heading 1", 16.0, NAVY),
    ("Heading 2", 12.8, BLUE),
    ("Heading 3", 11.2, INK),
):
    style = styles[name]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(size)
    style.font.color.rgb = rgb(color)
    style.font.bold = True

header = sec.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hr = header.add_run("反倾销税风险穿透分析｜单烷基醚｜2026-08-13")
set_run_font(hr, size=8.2, color=MUTED)
add_page_number(sec.footer.paragraphs[0])

# 封面
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(38)
r = p.add_run("反倾销税风险穿透分析")
set_run_font(r, size=12, color=BLUE, bold=True)
p = doc.add_paragraph(style="Title")
p.paragraph_format.space_before = Pt(14)
p.paragraph_format.space_after = Pt(7)
r = p.add_run("乙二醇和丙二醇的单烷基醚")
set_run_font(r, size=25.5, color=NAVY, bold=True)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(18)
r = p.add_run("美国来源、德国供应链与中国反倾销税风险阶段审计报告")
set_run_font(r, size=15.2, color=INK, bold=True)

add_table(doc, ["项目", "内容"], [
    ["受税来源", "美国"],
    ["实施期间", "2022年1月11日起5年，预计至2027年1月10日；截至2026年8月13日公开信息仍显示实施中"],
    ["税号入口", "29094400、29094990；必须命中公告列明的20项具体产品"],
    ["易迅已审计查询", "GLYCOL ETHER→中国，近一年，5页共84条"],
    ["逐条审计覆盖", "84次原始出现；81个可见字段唯一签名；约78个保守商业记录组"],
    ["阶段状态", "尚缺两项税号、20项具体品名/CAS/品牌及美国→第三国A腿全页补查"],
    ["报告日期", "2026年8月13日"],
], [1700, 7660], size=9.1)

add_callout(
    doc,
    "核心结论",
    "84条逐行筛查后，只有20次原始出现、18个可见字段唯一记录可明确落入措施产品范围：美国直达中国5条，德国→中国13条。德国B腿存在具体企业、集装箱、LC/PO和订单键值，但未找到美国→德国同批A腿，也没有中国进口申报原产国和税款书；公开资料同时证明Dow在德国Stade生产DOWANOL、BASF在德国供应N-Hexyl Glycol。因此目前可以定向核单，不能认定经德国绕道或少缴税。",
)
add_callout(
    doc,
    "重大范围纠偏",
    "沙特15条、重量字段489,457.68的乙二醇/二乙二醇单丁醚不在公告20项清单内；越南28条PTMEG是聚合物，不是本案小分子单烷基醚；美国8条EPh/PPh为苯基醚而非烷基醚。上述记录全部不得计入反倾销税风险数量。",
    fill=PALE_RED,
    title_color=RED,
)

page_break(doc)
doc.add_heading("一、核心判断", level=1)
add_bullet(doc, "现有查询已完整读取5页：20+20+20+20+4=84条，日期覆盖2025年8月6日至2026年5月14日，目的国均为中国。")
add_bullet(doc, "可见字段精确去重后81条；另有2组高概率同票镜像、2组中等概率近似记录，故保守约78个商业记录组。正式数量仍须提单号/柜号去重。")
add_bullet(doc, "明确范围内18个唯一记录，重量字段373,748.02；其中美国直达5条/114,366.02，德国B腿13条/259,382。")
add_bullet(doc, "美国直达5条属于优先核税对象，不是第三国绕道证据；买方字段均隐藏、金额为空，无法计算实际税额。")
add_bullet(doc, "德国13条只证明德国发货或平台原产字段，不证明美国原产。Dow集团/品牌关系、数据源国和货代名称均不能替代原产地证据。")
add_bullet(doc, "公开官方及企业一手资料未发现本品已公开的中国反规避调查、海关处罚、法院判决或双段提单闭环。已证实逃避反倾销税：未发现。")

doc.add_heading("二、政策范围、税率与有效期", level=1)
add_para(doc, "商务部公告2022年第3号自2022年1月11日起对原产于美国的相关乙二醇和丙二醇的单烷基醚征收反倾销税，期限5年。税号29094400、29094990只是入口；同税号项下未列明产品不在范围。截至审计日，商务部2026年3月13日实施中案件总表仍列该措施，预计到期日为2027年1月10日；公开检索未见已启动期终复审公告。")
add_table(doc, ["生产商", "反倾销税率", "若完全漏征：AD+13%连带VAT增量"], [
    ["The Dow Chemical Company", "57.4%", "海关审定完税价格×64.862%"],
    ["其他美国公司", "65.3%", "海关审定完税价格×73.789%"],
], [3900, 1700, 3760], size=9.0)
add_callout(doc, "税种口径", "商务部2022年第4号反补贴终裁虽认定16.8%补贴率，但明确决定暂不实施反补贴措施。本报告只核反倾销税及其导致的进口增值税差额，不计算反补贴税。", fill=LIGHT_BLUE)

doc.add_heading("2.1 公告列明20项产品", level=2)
add_table(doc, ["体系", "列明产品"], [
    ["乙二醇系（13项）", "乙二醇/二乙二醇/三乙二醇/四乙二醇甲醚；乙二醇/二乙二醇/三乙二醇乙醚；乙二醇/二乙二醇丙醚；乙二醇/二乙二醇己醚；乙二醇/二乙二醇异辛醚"],
    ["丙二醇系（7项）", "丙二醇/二丙二醇乙醚；丙二醇/二丙二醇丙醚；丙二醇/二丙二醇/三丙二醇丁醚"],
    ["典型排除", "乙二醇或二乙二醇单丁醚、丙二醇甲醚、PTMEG聚合物、苯基醚、醋酸酯、氨基醚等"],
], [2200, 7160], size=8.6)

page_break(doc)
doc.add_heading("三、数据完整性、去重与范围筛查", level=1)
q = AUDIT["query"]
add_table(doc, ["层级", "记录数", "重量字段", "解释"], [
    ["原始读取", "84", "—", "5页全部读取；与all84顺序一致"],
    ["可见字段唯一", "81", "—", "3组完全相同记录额外出现3次"],
    ["保守商业组", "约78", "—", "再合并2组高概率镜像和1组中等概率美国Hexyl候选"],
    ["明确范围内", "18唯一", "373,748.02", "美国5+德国13"],
    ["CAS/结构待核", "4唯一", "95,177.00", "比利时generic/isopropyl glycol ether"],
    ["明确排除", "59唯一", "750,925.36", "不属于公告20项"],
], [2300, 1300, 1800, 3960], size=8.7)
add_callout(doc, "字段限制", "原始JSON未提供列头截图、重量单位、数量单位或金额币种。报告使用“重量字段/数量字段/金额字段”；外国贸易数据的重量还可能是毛重。不得把重量字段直接写成中国海关净重，也不得把金额字段写成人民币或美元。", fill=PALE_RED, title_color=RED)

doc.add_heading("3.1 三组完全重复与疑似同票", level=2)
add_table(doc, ["类型", "记录", "处理"], [
    ["完全重复", "2026-04-19 德国PnB 23,870；同日PnB Harmless 23,720；2026-02-21 比利时Isopropyl 25,555", "并列原始84与唯一81；不直接删票"],
    ["高概率镜像", "2025-09-20 德国PnB 23,820，两条仅数量空值/0不同；2025-09-19 德国DPnB 16,674，跨数据源同日同重同主体", "保守合并，调B/L确认"],
    ["中等概率", "2025-11-11 美国Hexyl两条各32,216.03；2025-10-09美国EPh两条各21,132（范围外）", "无柜号前均保留"],
], [1500, 5260, 2600], size=8.3)

doc.add_heading("四、范围纠偏：哪些记录不能计税", level=1)
add_table(doc, ["排除品类", "唯一记录", "字段量", "排除理由"], [
    ["PTMEG聚醚", "29", "越南数量字段721,520；墨西哥重量字段48,700", "HS39072910/CAS25190-06-1聚合物，不是小分子单烷基醚"],
    ["沙特EG/DEG单丁醚", "15", "重量字段489,457.68", "不在公告20项列明产品中"],
    ["美国EPh/PPh苯基醚", "8", "重量字段212,767.18", "芳基醚，不是单烷基醚"],
    ["印度醋酸酯", "4", "数量字段128,000", "酯类，不是醇醚本体"],
    ["印度氨基醚", "2", "数量字段8", "结构不符"],
    ["印尼DPM样品", "1", "重量字段0.5", "二丙二醇甲醚未列入"],
], [1900, 1050, 2850, 3560], size=8.4)
add_callout(doc, "本轮最重要修正", "此前若把沙特15条或越南28条PTMEG当作第三国绕道，会把至少489,457.68的重量字段和721,520的数量字段错误纳入风险量。当前已从计税候选中全部剔除。", fill=PALE_GREEN)

page_break(doc)
doc.add_heading("五、美国直达中国：5条优先核税记录", level=1)
add_table(doc, ["日期", "产品", "卖方", "重量字段", "数量字段", "研判"], [
    ["2025-12-02", "DOWANOL PnP", "The Dow Chemical Company", "15,753.98", "80", "明确范围；直达核税"],
    ["2025-11-11", "HEXYL CELLOSOLVE", "The Dow Chemical Company", "32,216.03", "160", "两条同值，疑似镜像之一"],
    ["2025-11-11", "HEXYL CELLOSOLVE", "The Dow Chemical Company", "32,216.03", "160", "无柜号前保留"],
    ["2025-10-25", "DOWANOL DPnP", "The Dow Chemical Company", "18,426.00", "80", "明确范围；直达核税"],
    ["2025-10-13", "DOWANOL PnP", "The Dow Chemical Company", "15,753.98", "80", "明确范围；直达核税"],
], [1200, 1900, 2300, 1350, 1100, 1510], size=8.4)
add_para(doc, "5条重量字段合计114,366.02；若两条Hexyl最终证实同一票，保守重量池为82,149.99。数量字段从重量/包装关系看更像桶数，不能当千克。5条买方字段均为“---”、金额为空，因此既不能落中国进口人，也不能测算实际税额。")
add_callout(doc, "核税而非定罪", "美国直达记录本身不异常。只有调取中国报关单和税款缴款书，确认产品范围、美国原产、生产商税档及未缴/少缴反倾销税后，才可计算税差。若为Dow且完全漏征，合计增量系数为完税价格×64.862%。", fill=PALE_GOLD)

doc.add_heading("六、德国→中国B腿：13个唯一记录", level=1)
add_table(doc, ["口径", "记录", "重量字段", "说明"], [
    ["可见字段唯一", "13", "259,382", "PnB、DPnB、N-Hexyl Glycol"],
    ["保守货运组", "约11", "218,888", "合并两组高概率镜像"],
    ["明确主体", "BASF SE、Dow Europe、Dow Chemical Shanghai、Nanjing Golden、Polystar", "—", "Stolt/BDP/Newport多为货代、罐箱或运输映射主体"],
], [2200, 1350, 1800, 4010], size=8.5)
add_para(doc, "没有检出美国→德国同批A腿、#&US标记、共同柜号/PO/批号，也没有中国端原产国字段。平台“Germany”字段与德国装运事实只能建立B腿，不足以证明美国原产或洗产地。")

doc.add_heading("6.1 最强可落地调单键值", level=2)
add_table(doc, ["日期/主体", "产品与数量", "键值", "核查目的"], [
    ["2026-02-06｜BASF SE→Polystar Shanghai", "N-Hexyl Glycol；平台重量16,770；货描净重14,800kg", "柜MSKU510989-6；封志005610820；BILL2512040；6013458836/000010；3020907988/000010；黄埔新港", "以柜号/订单调CO、批次、报关单、税款书"],
    ["2026-04-14｜BDP→Nanjing Golden", "DOWANOL DPnB；平台16,674；净重15,200kg", "LC NB023IL001813300；SO117659063；PO250703G-0544-SH1/HKES2026NZ003-2；ref0118631017", "核德国生产批次、供应合同和原产地证"],
    ["2026-04-14｜BDP→Dow Shanghai", "DOWANOL DPnB；平台16,674；净重15,200kg", "PO/ref4010344886；USCI91310000760867002K；contract299593739", "集团关联链成立，但仍须生产地证据"],
    ["2025-11-12｜BDP→Nanjing Golden两票", "DOWANOL DPnB；平台33,356；净重30,400kg", "两套LC/PO/ref：1494800/JH007/0117981895；1440900/JH004-2/0117901237", "不同单号应分别调单，不得合并"],
], [2100, 2300, 3200, 1760], size=7.8)

doc.add_heading("七、合法德国生产的公开反证", level=1)
add_para(doc, "第三国B腿的判断必须同时检验合法本地产能。Dow官方Stade工厂资料将DOWANOL列为德国Stade生产产品；BASF官方N-Hexyl Glycol安全数据表列明BASF SE及德国Ludwigshafen联系信息。它们不能替代具体批次CO，但说明“德国真实生产”是强而合理的替代解释。")
add_table(doc, ["线索", "公开一手事实", "对本案意义"], [
    ["Dow Stade", "Dow德国工厂资料明确列有DOWANOL产品；Stade为集团综合生产基地", "Dow品牌从德国发华不能自动回推为美国原产"],
    ["BASF N-Hexyl Glycol", "官方SDS载CAS 112-25-4、BASF SE和Ludwigshafen信息", "2026-02-06 BASF票更可能具德国合法来源；仍需批次CO"],
    ["Dow Europe GmbH", "德国B腿卖方/集团商业节点", "关联公司或开票路径不等于物理转运"],
], [2200, 3500, 3660], size=8.6)
add_callout(doc, "不能使用的推理", "不能因商品为DOWANOL、集团总部在美国、某条“数据源”为美国，便认定货物美国原产。原产地判断必须落到生产工厂、批号、CO/制造商声明和中国进口申报。", fill=PALE_RED, title_color=RED)

doc.add_heading("八、比利时4条：CAS/结构待核", level=1)
add_para(doc, "比利时方向有5次原始出现、4个唯一记录，重量字段95,177，货描仅为generic monoalkylether或isopropyl glycol ether。由于未给具体化学结构、CAS和品牌，不能判断是否命中20项产品，暂不计入明确风险量。")
add_table(doc, ["必须补充", "判定作用"], [
    ["CAS、化学全名、SDS/TDS", "确认是乙二醇/丙二醇的哪一种单烷基醚"],
    ["生产商与工厂地址、批号、CO", "判断美国或比利时原产"],
    ["中国进口10位税号、商品规范申报", "判断归类和产品描述是否一致"],
    ["完税价格、反倾销税/增值税缴款书", "若确认美国原产且漏税，再测算税差"],
], [3300, 6060], size=8.8)

doc.add_heading("九、证据分级与税差门槛", level=1)
add_table(doc, ["等级", "当前证据", "可支持", "不能支持"], [
    ["A-核税", "美国直达5条明确范围内记录", "逐票核税；Dow税档候选", "未缴税或违法"],
    ["B+", "德国13个唯一B腿；部分有柜号、LC/PO", "德国发华路径及具体调单对象", "美国A腿、美国原产、规避"],
    ["B-待范围", "比利时4个唯一记录", "调CAS/CO", "落入产品范围"],
    ["排除", "59个唯一记录", "不纳入本案风险量", "反倾销税差"],
    ["未闭合", "无A/B同批、无中国原产申报和税单", "第三国风险仍可核查", "逃税/伪报原产已发生"],
], [1100, 3000, 2450, 2810], size=8.5)
add_callout(doc, "税差公式", "只有确认货物属于20项产品、实际美国原产、生产商税档和中国端未缴税后，才计算：反倾销税=海关审定完税价格×57.4%或65.3%；连带进口VAT差额=反倾销税×13%。现有记录没有可用完税价格，实际税额暂为“无法测算”，不是零。", fill=PALE_GOLD)

doc.add_heading("十、进一步核查清单", level=1)
for text in [
    "补做HS29094400、29094990→中国近两年全页查询，并用公告20项中英文名、CAS及DOWANOL/CELLOSOLVE/CARBITOL牌号补漏；每个查询记录总数、页数和日期边界。",
    "对美国直达5条向承运人/货代反查B/L、柜号和中国通知方；向中国海关底单核对进口人、生产商、原产国、税率、完税价格和税款缴款书。",
    "以德国四组最强键值调取合同、LC、PO、装箱单、CO、制造商声明、生产批号和工厂出库记录；核Dow Stade/BASF Ludwigshafen是否为实际生产厂。",
    "补查美国→德国/比利时A腿，窗口建议发华日前30—180天；重量容差±3%—5%，必须叠加同牌号、批号、柜号或PO，不能仅按同名商品匹配。",
    "对Dow Europe、Dow Shanghai、Nanjing Golden、Polystar区分生产商、销售方、进口人和货代角色；BDP、Stolt、Newport不得直接写成生产商。",
    "对比利时4条取得CAS/SDS后重新归类；不命中公告20项的记录继续排除。",
    "在2027年1月10日前持续核查商务部是否发布期终复审立案；若立案，措施通常在复审期间继续实施，应动态更新系统。",
]:
    add_number(doc, text)

doc.add_heading("十一、主要公开来源", level=1)
sources = [
    ("商务部公告2022年第3号", "https://dcj.mofcom.gov.cn/article/zcfb/zcdwmy/202201/20220103235940.shtml", "20项产品范围、HS、税率、公式和实施期"),
    ("商务部公告2022年第4号", "https://www.mofcom.gov.cn/zcfb/dwmygl/art/2022/art_e9fce28544324ab4bf016a64bed193c2.html", "补贴率16.8%，但明确暂不实施反补贴措施"),
    ("商务部实施中案件总表（2026-03-13）", "https://admin.cacs.mofcom.gov.cn/cacscms/article/sszaj?articleId=160267&type=", "预计措施到期日2027-01-10"),
    ("Dow Stade工厂资料", "https://de.dow.com/content/dam/corp/documents/location/903-076-03-dow-stade-fact-sheet.pdf", "德国Stade生产网络及DOWANOL产品"),
    ("Dow Stade厂区手册", "https://corporate.dow.com/documents/location/903-367-03-werk-stade-broschuere.pdf", "德国工厂补充资料"),
    ("BASF N-Hexyl Glycol SDS", "https://download.basf.com/p1/000000000030796740_SDS_GEN_DE/en/n-Hexylglycol_BMB_I_30796740_SDS_GEN_DE_en_6-0.pdf", "CAS112-25-4及BASF德国主体"),
    ("Dow DOWANOL PnP产品页", "https://www.dow.com/en-us/pdp.dowanol-pnp-glycol-ether.40614z.html", "丙二醇单丙醚结构"),
    ("Dow DOWANOL DPnP产品页", "https://www.dow.com/en-us/pdp.dowanol-dpnp-glycol-ether.41173z.html", "二丙二醇单丙醚结构"),
    ("Dow DOWANOL PnB产品页", "https://www.dow.com/en-us/pdp.dowanol-pnb-glycol-ether.5790z.html", "丙二醇单丁醚结构"),
    ("Dow溶剂选择指南", "https://www.dow.com/documents/327/327-00001-01-dow-oxygenated-solvents-selection-guide.pdf?iframe=true", "品牌、化学名称与用途对照"),
]
for label, url, note in sources:
    add_source(doc, label, url, note)

doc.add_heading("十二、公开检索结论的边界", level=1)
add_para(doc, "截至2026年8月13日，公开官方和企业一手来源中未找到本措施项下已公开的第三国反规避调查、海关行政处罚、刑事/行政判决或可核验双段提单，能够证明美国原产单烷基醚经德国、比利时或其他第三国更换原产地后进入中国。公开缺失不等于绝无风险；它决定了当前报告的措辞必须是“待核线索”，而不是“已发生逃税”。")

page_break(doc)
doc.add_heading("附录一：18个明确范围内唯一记录", level=1)
appendix_rows = []
for row in LEADS:
    appendix_rows.append([
        row["query_row"],
        row["date"],
        row["platform_origin"],
        row["product_class"],
        fmt(row["weight_num"], 2),
        row["buyer_or_consignee"][:30],
        row["seller_or_shipper"][:28],
    ])
add_table(doc, ["行", "日期", "平台原产字段", "品类", "重量字段", "第一方", "第二方"], appendix_rows, [550, 1050, 1200, 1550, 1200, 2020, 1790], size=7.3)

doc.add_heading("附录二：交付文件与状态", level=1)
for text in [
    "单烷基醚_易迅逐票判定台账_阶段审计.xlsx：含概览、18个明确范围唯一记录、4个范围待CAS记录、全部84条、重复审计、调单键值和政策范围。",
    "单烷基醚_易迅逐票标准化.csv/json：84条逐行底稿，含范围、路线、证据等级和可见单证键值。",
    "单烷基醚_全量阶段审计.json：统计、范围纠偏、税率、查询缺口和重点调单键值。",
    "单烷基醚_易迅_all84.json及p1—p5：保留原始页面读取结果。",
]:
    add_bullet(doc, text)
add_callout(doc, "报告状态", "本文件为阶段审计报告。完成HS、20项品名/CAS/品牌、美国A腿和中国端报关税单核验后，应覆盖更新为最终深度分析报告，并同步总报告、总清单和前端系统。", fill=PALE_RED, title_color=RED)

doc.core_properties.title = "乙二醇和丙二醇单烷基醚反倾销税与第三国转运风险阶段审计报告"
doc.core_properties.subject = "易迅数据逐条审计、美国—德国—中国供应链、产品范围与条件税差"
doc.core_properties.author = "Codex"
doc.core_properties.keywords = "单烷基醚, 乙二醇醚, 丙二醇醚, 反倾销税, 美国, 德国, 第三国转运"
doc.save(REPORT_PATH)
print(REPORT_PATH)
