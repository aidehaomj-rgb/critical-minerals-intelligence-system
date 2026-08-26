#!/usr/bin/env python3
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT=Path(r'D:\易迅数据\反倾销税深度分析报告\18_正丁醇')
DOCX=OUT/'正丁醇_反倾销与第三国转运风险_数据缺口阶段报告.docx'
NAVY=RGBColor(11,37,69); BLUE=RGBColor(46,116,181); GRAY=RGBColor(90,98,108); RED=RGBColor(155,28,28)

def ft(r,size=11,bold=False,color=None):
 r.font.name='Calibri'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'等线'); r.font.size=Pt(size); r.bold=bold
 if color:r.font.color.rgb=color
def sh(c,fill):
 x=OxmlElement('w:shd');x.set(qn('w:fill'),fill);c._tc.get_or_add_tcPr().append(x)
def mar(c):
 p=c._tc.get_or_add_tcPr();m=p.first_child_found_in('w:tcMar')
 if m is None:m=OxmlElement('w:tcMar');p.append(m)
 for n,v in [('top',80),('bottom',80),('start',120),('end',120)]:
  x=OxmlElement('w:'+n);x.set(qn('w:w'),str(v));x.set(qn('w:type'),'dxa');m.append(x)
def tbl(doc,headers,rows,widths):
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 pr=t.rows[0]._tr.get_or_add_trPr();x=OxmlElement('w:tblHeader');x.set(qn('w:val'),'true');pr.append(x)
 for i,h in enumerate(headers):
  c=t.rows[0].cells[i];c.width=Inches(widths[i]);sh(c,'F2F4F7');mar(c);p=c.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;ft(p.add_run(h),9.5,True,NAVY)
 for row in rows:
  cs=t.add_row().cells
  for i,v in enumerate(row):
   cs[i].width=Inches(widths[i]);cs[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;mar(cs[i]);p=cs[i].paragraphs[0];p.paragraph_format.space_after=Pt(0);ft(p.add_run(str(v)),9)
 return t
def bul(doc,s):
 p=doc.add_paragraph(style='List Bullet');p.paragraph_format.space_after=Pt(4);ft(p.add_run(s),10.3)
def lnk(doc,label,url):
 p=doc.add_paragraph();p.paragraph_format.space_after=Pt(2);ft(p.add_run(label+'：'),8.7,True,GRAY);ft(p.add_run(url),8.7,False,BLUE)

def main():
 OUT.mkdir(parents=True,exist_ok=True);d=Document();sec=d.sections[0]
 sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1);sec.header_distance=sec.footer_distance=Inches(.492)
 n=d.styles['Normal'];n.font.name='Calibri';n._element.rPr.rFonts.set(qn('w:eastAsia'),'等线');n.font.size=Pt(11);n.paragraph_format.space_after=Pt(6);n.paragraph_format.line_spacing=1.10
 for name,size,bef,aft,col in [('Heading 1',16,16,8,BLUE),('Heading 2',13,12,6,BLUE),('Heading 3',12,8,4,NAVY)]:
  s=d.styles[name];s.font.name='Calibri';s._element.rPr.rFonts.set(qn('w:eastAsia'),'等线');s.font.size=Pt(size);s.font.bold=True;s.font.color.rgb=col;s.paragraph_format.space_before=Pt(bef);s.paragraph_format.space_after=Pt(aft)
 h=sec.header.paragraphs[0];ft(h.add_run('反倾销税深度分析｜第18项'),9,True,GRAY)
 f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT;ft(f.add_run('阶段审计报告｜2026-08-13'),9,False,GRAY)
 p=d.add_paragraph();ft(p.add_run('风险核查备忘录'),10,True,BLUE)
 p=d.add_paragraph();p.paragraph_format.space_after=Pt(4);ft(p.add_run('正丁醇（N-Butanol）'),23,True,NAVY)
 p=d.add_paragraph();p.paragraph_format.space_after=Pt(14);ft(p.add_run('中国反倾销措施、企业税率穿透与第三国转运风险阶段审计'),13,False,GRAY)
 for a,b in [('受税来源','台湾地区、马来西亚、美国'),('产品标识','HS 29051300；CAS 71-36-3；UN 1120'),('审计日期','2026-08-13'),('本地范围','D:\易迅数据内304个XLSX/XLS/CSV/JSON文件内容级预筛'),('阶段结论','本地未取得正丁醇逐票底表；公开源未闭合受税来源→第三国→中国链')]:
  p=d.add_paragraph();p.paragraph_format.space_after=Pt(2);ft(p.add_run(a+'：'),10.5,True);ft(p.add_run(b),10.5)

 d.add_heading('一、结论先行',1)
 p=d.add_paragraph();p.paragraph_format.space_after=Pt(7);ft(p.add_run('当前没有证据证明正丁醇经第三国进入中国并逃避反倾销税。'),12,True,RED)
 bul(d,'措施自2024年12月29日起续征5年，预计至2029年12月28日；企业税率从6.0%到139.3%，必须穿透生产商和特定销售链。')
 bul(d,'本地304个文件只有主清单和进度台账两处项目名称；无日期、商品、主体、数量和路线字段的真实贸易票据。')
 bul(d,'公开资料确认台湾塑化、马来西亚PETRONAS体系、美国OQ等真实受税产能；德国BASF等非受税地区也有真实供应能力，品牌或集团国籍不能替代原产地。')
 bul(d,'公开检索未发现中国正丁醇专项海关处罚、反规避裁定、判决或闭合双段提单；当前只能建立查询与调证路线。')

 d.add_heading('二、现行税率与条件税差',1)
 tbl(d,['来源/企业','AD税率','AD及AD引致VAT增量*'],[
  ('台湾塑胶工业股份有限公司','6.0%','6.780% × 完税价格'),('台湾地区其他公司','56.1%','63.393% × 完税价格'),
  ('PETRONAS特定生产/销售链','12.7%','14.351% × 完税价格'),('BASF PETRONAS、Optimal及马来西亚其他','26.7%','30.171% × 完税价格'),
  ('OQ/OXEA美国','52.2%','58.986% × 完税价格'),('Eastman、Dow、BASF及美国其他','139.3%','157.409% × 完税价格')],[2.8,1.15,2.55])
 p=d.add_paragraph('* 按进口增值税13%筛查：完税价格×AD税率×1.13。正式税差须使用中国海关完税价格和税款书；不含正常关税及正常进口增值税。');ft(p.runs[0],8.5,False,GRAY)
 p=d.add_paragraph('马来西亚12.7%仅适用于马石化营销（纳闽）有限公司向中国出口、且由国油石化衍生公司生产的涉案产品。生产商、销售人或链条不满足时，不能套用低税率。')

 d.add_heading('三、产品范围及高频误纳边界',1)
 p=d.add_paragraph('正丁醇，又称1-丁醇、丙原醇、酪醇；英文Butan-1-ol、1-Butanol、N-Butanol、N-Butyl Alcohol，分子式CH₃(CH₂)₃OH。用于丙烯酸丁酯、醋酸丁酯、增塑剂、丁胺、乙二醇丁醚等下游产品。')
 bul(d,'明确排除：异丁醇、叔丁醇、仲丁醇、丁二醇、丁酮。')
 bul(d,'下游产品不是正丁醇：醋酸丁酯、丙烯酸丁酯、丁醚、丁胺及含正丁醇配方制剂；须核成分和归类。')
 bul(d,'外国HS 290513或货描“BUTANOL”不能单独替代中国报关税号、CAS和异构体判定。')

 d.add_heading('四、本地易迅数据全盘盘点',1)
 tbl(d,['审计项','结果','说明'],[
  ('文件总数','304','XLSX/XLS/CSV/JSON；排除本项输出目录'),('预筛候选','2','商品清单及总进度台账'),('HS/CAS/UN命中','0','未发现29051300、71-36-3或UN1120逐票底表'),('可确认贸易记录','0','无日期—商品—主体—数量—路线闭合字段'),('中国B腿/A腿/闭合链','0 / 0 / 0','不能统计数量、企业、口岸或税差'),('解析错误','0','本地文件内容级预筛完成')],[1.55,1.25,3.7])
 p=d.add_paragraph('边界：0条只代表现有D盘数据缺口，不代表易迅全库无数据、中国无进口或无第三国风险。')

 d.add_heading('五、实体穿透与合法产能反证',1)
 tbl(d,['节点','公开核验','核查意义'],[
  ('台湾塑胶','官方披露2024年正丁醇产能25万吨、产量24.65万吨，并有8.84万吨外销','真实台湾原产供货能力；适用6.0%须核生产商'),
  ('PETRONAS马来西亚','官方产品组合列有N-Butanol；终裁对特定生产/销售链单列12.7%','须同时核生产商与Labuan销售实体'),
  ('OQ/OXEA美国','官方ISO证书确认Bay City, Texas生产n-Butanol','真实美国原产；OQ适用52.2%'),
  ('BASF德国','官方产品/安全资料显示Ludwigshafen的n-Butanol供应','重要反证：BASF品牌不必然是美国原产'),
  ('第三国贸易商/储罐商','公开资料尚未形成同批A/B腿','开票或仓储国不等于法定原产国')],[1.35,3.25,1.9])

 d.add_heading('六、第三国风险判定门槛',1)
 bul(d,'A级：受税来源A腿与第三国→中国B腿具有同船/同罐号、同批号或同提单，第三国无实质生产，且中国申报非受税原产或未缴AD。当前0条。')
 bul(d,'B+级：同实体、CAS/纯度一致、0—90日、净重±2%，第三国只有贸易/储罐功能；仍需中国报关和原产证闭环。')
 bul(d,'C级：仅集团公司、品牌、启运国或宏观流量相同。不能据此指认企业逃税。')
 bul(d,'合法替代解释：第三国真实羰基合成/加氢/精馏生产、正常库存分拨，或第三国发货但中国仍申报受税原产并正常缴税。')

 d.add_heading('七、下一轮易迅全页查询矩阵',1)
 tbl(d,['查询','条件','目标'],[
  ('Q1 中国B腿','目的国中国；HS29051300/290513；2023-12-29至最新并补历史；7组名称/CAS/UN分别查','每页200条读至末页；跨查询去重并排除异构体/衍生物'),
  ('Q2 受税A腿','台湾、马来西亚、美国→全球；同税号/关键词；逐个补查列名企业','识别第三国收货人、贸易商、储罐商和生产厂'),
  ('Q3 第三国B腿','Q2第三国→中国；同品名并按主体精确补查','按0—30/60/90/180日和运输指纹闭合')],[1.25,3.85,1.4])
 p=d.add_paragraph('重点先看新加坡、韩国、泰国、印度、越南、阿联酋、荷兰、德国、比利时等贸易/仓储或真实生产节点；最终以Q2实际收货国扩展，不预设某国必然绕道。')

 d.add_heading('八、调证字段及优先实体',1)
 bul(d,'货物指纹：CAS 71-36-3、纯度、水分、色度、醛含量、批号、COA、SDS、UN1120、储罐/ISO罐号、净重。')
 bul(d,'商业运输：两段提单、船名航次、中转港、罐区进出记录、合同、发票、付款受益人、贸易术语。')
 bul(d,'生产原产：羰基合成、加氢与精馏装置，投料、BOM、能耗、工单、库存、生产日期及中国非优惠原产地核定。')
 bul(d,'中国端：进口人、消费使用单位、报关企业、口岸、启运国、原产国、境外生产商、完税价格、AD税率和税款缴款书。')
 bul(d,'优先实体：Formosa Plastics；PETRONAS Chemicals Derivatives/Marketing (Labuan)；BASF PETRONAS；Optimal；OQ/OXEA；Eastman；Dow；BASF。列名仅用于核单，不代表涉嫌违法。')

 d.add_heading('九、阶段结论',1)
 bul(d,'已完成现行政策、完整企业税率、产品边界、本地304文件盘点及公开产能反证。')
 bul(d,'未找到可指向中国进口人、境外第三国收发货人、报关行或口岸的正丁醇风险票据。')
 bul(d,'当前评级为“数据不足/无法判断”，并非“低风险”或“无风险”；待Q1—Q3全页数据后再计算数量和条件税差。')

 d.add_heading('十、主要公开来源',1)
 lnk(d,'商务部2024年第57号期终复审裁定','https://www.mofcom.gov.cn/zcfb/blgg/gg/2024/art/2024/art_d7750f7273a94cb2930545aadaf5ff3d.html')
 lnk(d,'商务部2018年第100号终裁','https://dcj.mofcom.gov.cn/article/zcfb/zcblgg/201812/20181202820832.shtml')
 lnk(d,'台湾塑胶正丁醇生产量值','https://www.fpc.com.tw/fpcw/index.php?c=14&id=12&op=res')
 lnk(d,'台湾塑胶正丁醇销售量值','https://www.fpc.com.tw/fpcw/index.php?c=15&id=12&op=res')
 lnk(d,'PETRONAS Chemicals产品组合','https://www.petronas.com/pcg/our-business/olefins-glycols-derivatives')
 lnk(d,'OQ Chemicals美国Bay City生产证书','https://chemicals.oq.com/fileadmin/user_upload/OQ-Chemicals/Company/Service/Certificat/North_America/QM15_10005680_QM15_EN.pdf')
 lnk(d,'中国非优惠原产地实质性改变标准','https://policy.mofcom.gov.cn/claw/clawContent.shtml?id=101350')
 d.save(DOCX);print(DOCX)
if __name__=='__main__':main()
