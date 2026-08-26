#!/usr/bin/env python3
"""Normalize exported or scraped 外贸公社 shipment records into a standard Excel workbook."""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import re
from pathlib import Path
from typing import Any

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except Exception as exc:  # pragma: no cover - environment guard
    raise SystemExit("openpyxl is required to write standardized Excel output") from exc


STANDARD_FIELDS = [
    "查询条件",
    "数据来源页面",
    "查询时间",
    "结果序号",
    "企业名称",
    "企业角色",
    "产品名称",
    "HS编码",
    "进口或出口方向",
    "国家或地区",
    "日期",
    "数量",
    "重量",
    "金额",
    "提运单号",
    "报关单号",
    "集装箱号",
    "封志号",
    "原始记录链接",
    "风险标签",
    "备注",
]


HEADER_ALIASES = {
    "企业名称": ["企业", "公司", "供应商", "采购商", "进口商", "出口商", "shipper", "supplier", "buyer", "consignee", "importer", "exporter"],
    "产品名称": ["产品", "品名", "货描", "商品描述", "description", "product", "goods", "cargo"],
    "HS编码": ["hs", "hscode", "hs编码", "税号", "商品编码", "hts"],
    "国家或地区": ["国家", "地区", "目的国", "原产国", "进口国", "出口国", "country", "origin", "destination"],
    "日期": ["日期", "到港", "出运", "申报日期", "date", "arrival", "departure"],
    "数量": ["数量", "件数", "qty", "quantity", "packages"],
    "重量": ["重量", "净重", "毛重", "weight", "net weight", "gross weight", "kg"],
    "金额": ["金额", "价格", "价值", "value", "amount", "usd"],
    "提运单号": ["提单", "提运单", "bill", "bol", "b/l", "master", "house"],
    "报关单号": ["报关单", "申报单", "declaration", "entry"],
    "集装箱号": ["集装箱", "container"],
    "封志号": ["封志", "seal"],
    "原始记录链接": ["链接", "url", "link", "href"],
}


def norm(text: Any) -> str:
    return re.sub(r"\s+", "", str(text or "").lower())


def read_rows(path: Path) -> list[dict[str, Any]]:
    suffix = path.suffix.lower()
    if suffix in {".csv", ".txt"}:
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            return list(csv.DictReader(handle))
    if suffix in {".xlsx", ".xlsm"}:
        wb = load_workbook(path, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(x or "").strip() for x in rows[0]]
        return [
            {headers[i] if i < len(headers) else f"列{i+1}": value for i, value in enumerate(row)}
            for row in rows[1:]
            if any(cell not in (None, "") for cell in row)
        ]
    if suffix == ".json":
        data = json.loads(path.read_text(encoding="utf-8"))
        if isinstance(data, list):
            return [row if isinstance(row, dict) else {"原始记录": row} for row in data]
        if isinstance(data, dict) and isinstance(data.get("records"), list):
            return [row if isinstance(row, dict) else {"原始记录": row} for row in data["records"]]
    raise SystemExit(f"Unsupported input file type: {path.suffix}")


def find_value(row: dict[str, Any], field: str) -> Any:
    aliases = HEADER_ALIASES.get(field, [])
    normalized_keys = {norm(k): k for k in row.keys()}
    for alias in aliases:
        alias_norm = norm(alias)
        for key_norm, original_key in normalized_keys.items():
            if alias_norm and alias_norm in key_norm:
                value = row.get(original_key)
                if value not in (None, ""):
                    return value
    return ""


def infer_role(row: dict[str, Any]) -> str:
    keys = " ".join(row.keys()).lower()
    if any(x in keys for x in ["supplier", "shipper", "供应商", "出口商"]):
        return "供应商/出口商"
    if any(x in keys for x in ["buyer", "consignee", "importer", "采购商", "进口商"]):
        return "采购商/进口商"
    return ""


def normalize_rows(rows: list[dict[str, Any]], args: argparse.Namespace) -> list[dict[str, Any]]:
    now = args.query_time or dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    output = []
    for index, row in enumerate(rows, 1):
        item = {field: "" for field in STANDARD_FIELDS}
        item["查询条件"] = args.query or ""
        item["数据来源页面"] = args.source_url or ""
        item["查询时间"] = now
        item["结果序号"] = index
        for field in ["企业名称", "产品名称", "HS编码", "国家或地区", "日期", "数量", "重量", "金额", "提运单号", "报关单号", "集装箱号", "封志号", "原始记录链接"]:
            item[field] = find_value(row, field)
        item["企业角色"] = infer_role(row)
        item["进口或出口方向"] = args.direction or ""
        item["备注"] = row.get("备注", "") or row.get("note", "") or ""
        output.append(item)
    return output


def autosize(ws) -> None:
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        width = 12
        for cell in ws[letter]:
            width = max(width, min(len(str(cell.value or "")) + 2, 42))
        ws.column_dimensions[letter].width = width


def write_workbook(records: list[dict[str, Any]], args: argparse.Namespace) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "标准化结果"
    ws.append(STANDARD_FIELDS)
    for record in records:
        ws.append([record.get(field, "") for field in STANDARD_FIELDS])
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    autosize(ws)

    meta = wb.create_sheet("查询条件")
    meta_rows = [
        ("查询条件", args.query or ""),
        ("数据来源页面", args.source_url or ""),
        ("查询时间", args.query_time or dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("进口或出口方向", args.direction or ""),
        ("结果数量", len(records)),
    ]
    for row in meta_rows:
        meta.append(row)
    meta.column_dimensions["A"].width = 18
    meta.column_dimensions["B"].width = 80
    args.output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.output)


def main() -> None:
    parser = argparse.ArgumentParser(description="Normalize 外贸公社 trade records into standard Excel.")
    parser.add_argument("--input", type=Path, required=True, help="Raw CSV/XLSX/JSON export or scraped rows file.")
    parser.add_argument("--output", type=Path, required=True, help="Standardized XLSX output path.")
    parser.add_argument("--query", default="", help="Human-readable query condition summary.")
    parser.add_argument("--source-url", default="", help="Source search/result page URL.")
    parser.add_argument("--query-time", default="", help="Query timestamp. Defaults to now.")
    parser.add_argument("--direction", default="", help="进口 or 出口 direction.")
    args = parser.parse_args()
    rows = read_rows(args.input)
    records = normalize_rows(rows, args)
    write_workbook(records, args)
    print(f"Wrote {len(records)} standardized rows to {args.output}")


if __name__ == "__main__":
    main()
