import json, os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE=os.path.dirname(__file__)
data=json.load(open(os.path.join(BASE,'.pc_analysis.json'),encoding='utf-8'))
outdir=r'D:\易迅数据\反倾销税风险文件夹'
os.makedirs(outdir,exist_ok=True)
out=os.path.join(outdir,'聚碳酸酯_易迅数据反倾销税风险分析报告_2026-08-12.docx')

doc=Document(); sec=doc.sections[0]
sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
styles=doc.styles
for s in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
    styles[s].font.name='Calibri'; styles[s]._element.rPr.rFonts.set(qn('w:eastAsia'),'微软雅黑')
styles['Normal'].font.size=Pt(10.5); styles['Normal'].paragraph_format.space_after=Pt(6); styles['Normal'].paragraph_format.line_spacing=1.1
for s,size,color in [('Heading 1',16,'2E74B5'),('Heading 2',13,'2E74B5'),('Heading 3',12,'1F4D78')]:
    styles[s].font.size=Pt(size); styles[s].font.color.rgb=RGBColor.from_string(color)

p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('聚碳酸酯易迅数据反倾销税风险分析报告'); r.bold=True; r.font.size=Pt(22); r.font.color.rgb=RGBColor.from_string('17365D')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('HS 390740｜中国目的地｜2024-08-06至2026-08-06'); r.font.size=Pt(11); r.font.color.rgb=RGBColor.from_string('666666')

doc.add_heading('一、结论摘要',level=1)
for t in [
'易迅最终检索口径共2,629条，设置为200条/页，共14页；前13页各200条、末页29条，已逐次点击“下一页”读取并完成首行变化校验，分页合计与页面总数一致。',
'现有记录显示，对华流向高度集中在越南（1,138票）、印度尼西亚（1,022票）、印度（230票）和菲律宾（171票）。其中大量货物描述为PC RECYCLE(D) PELLET、再生聚碳酸酯或改性/复合料，不能直接等同于反倾销措施所涉台湾地区原生聚碳酸酯。',
'最具体的穿透线索出现在越南记录：部分对华出口虽平台原产国字段为Vietnam，但商品描述保留“#&KR”或“#&CN”。这证明平台原产国字段不能单独作为原产地证据，应以描述尾标、生产商、配方和原产地证交叉核验。',
'本批数据没有形成“台湾地区聚碳酸酯→第三国→中国”的闭环证据：缺少台湾地区向越南、印尼、印度或菲律宾的对应上游进口记录及同型号、数量、时间、企业匹配。因此结论为“存在需穿透核查的第三国供应链风险”，不能据此认定绕道逃避反倾销税。'
]: doc.add_paragraph(t,style='List Bullet')

doc.add_heading('二、查询口径与完整性审计',level=1)
table=doc.add_table(rows=1,cols=2); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.style='Table Grid'
for i,x in enumerate(['项目','内容']): table.rows[0].cells[i].text=x
items=[('平台','易迅数据—全球海关数据'),('HS编码','39074000（聚碳酸酯，平台匹配各国扩展税号）'),('目的国/地区','CHINA标准选项'),('时间','2024-08-06—2026-08-06'),('进出口','全部；结果实际主要为第三国出口至中国'),('分页','200条/页，14页'),('记录数','2,629条'),('下载','未下载平台数据；仅逐页读取并形成统计')]
for a,b in items:
    c=table.add_row().cells;c[0].text=a;c[1].text=b

doc.add_heading('三、来源国/地区结构',level=1)
t=doc.add_table(rows=1,cols=4);t.style='Table Grid';t.alignment=WD_TABLE_ALIGNMENT.CENTER
for i,x in enumerate(['来源标记','票数','占比','研判']):t.rows[0].cells[i].text=x
notes={'Vietnam':'越南出口记录为主，部分描述含#&KR/#&CN，需穿透原料来源','Indonesia':'大量再生PC粒子，主要风险是品类/原生与再生界定','India':'需核对是否为原生PC、改性料或转口贸易','Philippines':'数量较少，仍应核查具体型号与生产能力'}
for x in data['topOrigins'][:4]:
    c=t.add_row().cells;c[0].text=x['name'];c[1].text=f"{x['count']:,}";c[2].text=f"{x['count']/2629:.1%}";c[3].text=notes.get(x['name'],'')

doc.add_heading('四、高风险实体与链路线索',level=1)
doc.add_heading('4.1 越南高频供应商',level=2)
for x in data['topSuppliers'][:12]: doc.add_paragraph(f"{x['name']}：{x['count']}票；平台数量字段合计约{x['qty']:,.2f}。",style='List Bullet')
doc.add_heading('4.2 中国/香港高频采购商',level=2)
for x in data['topBuyers'][:12]: doc.add_paragraph(f"{x['name']}：{x['count']}票；平台数量字段合计约{x['qty']:,.2f}。",style='List Bullet')

doc.add_heading('4.3 具体异常记录',level=2)
for text in [
'2026-06-30，越南Nagase Vietnam向NWP International Trading (Shenzhen)出口LEXAN 3412ECR-739，数量50，商品描述尾标“#&KR”，但平台原产国字段为Vietnam。该票直接证明“平台原产国字段—商品描述产地标记”存在冲突，建议调取原产地证、商业发票、COA和批号。',
'同日同一供应链另有LNP THERMOCOMP compound DC0041PE-7M1D145W，描述尾标“#&CN”，数量50，平台原产国同样显示Vietnam。这更可能是中国原料/产品在越南加工或返运，而非越南原产，应核查加工增值和原产地规则。',
'越南、印尼大量记录描述为PC RECYCLE PELLET/PC RECYCLED PELLETS。再生粒子可能不属于被调查原生聚碳酸酯的同一产品范围，但如实际为原生料掺混、简单换包或品名伪报，则存在归类和原产地风险。'
]: doc.add_paragraph(text,style='List Bullet')

doc.add_heading('五、证据等级',level=1)
for a,b in [('已证实','2,629条结果已全页读取；对华流向集中于越南、印尼、印度、菲律宾；越南个别记录存在#&KR/#&CN与平台Vietnam字段冲突。'),('高度疑似','Nagase Vietnam—深圳买方的LEXAN韩国标记记录具有明确的跨国产地线索，但尚未闭合台湾地区上游来源。'),('线索','越南和印尼高频再生PC供应商及高集中采购商值得进行主体穿透。'),('无法判断','是否逃避台湾地区聚碳酸酯反倾销税；缺少台湾上游进口、生产BOM、原产地证及中国进口报关数据。')]:
    p=doc.add_paragraph();p.add_run(a+'：').bold=True;p.add_run(b)

doc.add_heading('六、建议进一步核查的数据',level=1)
for x in ['台湾地区出口至越南、印度尼西亚、印度、菲律宾的HS 390740记录，至少覆盖同期两年；','LEXAN 3412ECR-739及其他品牌牌号的生产商、COA、批号和原厂销售链；','中国进口报关单中的境内收货人、境外发货人、原产国、贸易国、征免性质和反倾销税缴款书；','第三国工厂的产能、设备、BOM、原料入库与成品出库、加工增值和原产地证签发资料；','同集装箱号、提单号、船名航次、装卸港与15—90天时间窗的数量匹配。']:doc.add_paragraph(x,style='List Number')

doc.add_heading('附录：逐页审计',level=1)
t=doc.add_table(rows=1,cols=3);t.style='Table Grid'
for i,x in enumerate(['页码','读取条数','校验']):t.rows[0].cells[i].text=x
for a in data['audit']:
    c=t.add_row().cells;c[0].text=str(a['page']);c[1].text=str(a['count']);c[2].text='通过：逐次下一页，首行变化'

for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(2)
                for r in p.runs:r.font.size=Pt(9)
    for cell in table.rows[0].cells:
        shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'E8EEF5');cell._tc.get_or_add_tcPr().append(shd)
        for r in cell.paragraphs[0].runs:r.bold=True

doc.sections[0].header.paragraphs[0].text='聚碳酸酯反倾销税风险分析｜易迅数据'
doc.sections[0].footer.paragraphs[0].text='形成日期：2026-08-12｜结论仅用于风险识别，不替代海关执法认定'
doc.save(out)
print(out)
