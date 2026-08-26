from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path

OUT = Path(r"C:\Users\59809\Documents\关键矿产\outputs\Terra_Drone集团供应链与产能核查报告_2026-08-21.docx")
OUT.parent.mkdir(parents=True, exist_ok=True)

NAVY = "203748"; BLUE = "2E74B5"; GOLD = "C59A3D"; GRAY = "5B6573"; LIGHT = "F2F5F8"; PALE = "E8EFF5"; RED = "9C2F2F"; GREEN="2F6B4F"

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn('w:shd'))
    if shd is None: shd=OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tc=cell._tc; tcPr=tc.get_or_add_tcPr(); tcMar=tcPr.first_child_found_in('w:tcMar')
    if tcMar is None: tcMar=OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node=tcMar.find(qn('w:'+m))
        if node is None: node=OxmlElement('w:'+m); tcMar.append(node)
        node.set(qn('w:w'),str(v)); node.set(qn('w:type'),'dxa')

def font(run, size=9.5, bold=False, color=None, italic=False):
    run.font.name='Arial'; run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    run.font.size=Pt(size); run.bold=bold; run.italic=italic
    if color: run.font.color.rgb=RGBColor.from_string(color)

def set_repeat_table_header(row):
    trPr=row._tr.get_or_add_trPr(); rep=OxmlElement('w:tblHeader'); rep.set(qn('w:val'),'true'); trPr.append(rep)

def set_cell_width(cell, dxa):
    tcPr=cell._tc.get_or_add_tcPr(); tcW=tcPr.find(qn('w:tcW'))
    if tcW is None: tcW=OxmlElement('w:tcW'); tcPr.append(tcW)
    tcW.set(qn('w:w'),str(dxa)); tcW.set(qn('w:type'),'dxa')

doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.5); sec.page_height=Inches(11); sec.top_margin=Inches(.78); sec.bottom_margin=Inches(.72); sec.left_margin=Inches(.78); sec.right_margin=Inches(.78)
styles=doc.styles
normal=styles['Normal']; normal.font.name='Arial'; normal._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); normal.font.size=Pt(9.5)
normal.paragraph_format.space_after=Pt(5); normal.paragraph_format.line_spacing=1.18
for name,size,color,before,after in [('Heading 1',16,NAVY,14,7),('Heading 2',12.5,BLUE,10,5),('Heading 3',10.5,NAVY,7,3)]:
    s=styles[name]; s.font.name='Arial'; s._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor.from_string(color); s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after); s.paragraph_format.keep_with_next=True

def footer(section):
    p=section.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run('Terra Drone集团供应链与产能核查 | 截至2026年8月21日  ·  '); font(r,8,color=GRAY)
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); p._p.append(fld)
footer(sec)
hp=sec.header.paragraphs[0]; hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT; r=hp.add_run('供应链核查报告  |  内部分析'); font(r,8,color=GRAY)

def ptext(text,boldlead=None,italic=False,color=None,after=5,align=None):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(after)
    if align is not None: p.alignment=align
    if boldlead and text.startswith(boldlead):
        r=p.add_run(boldlead); font(r,bold=True,color=color)
        r=p.add_run(text[len(boldlead):]); font(r,color=color,italic=italic)
    else: r=p.add_run(text); font(r,color=color,italic=italic)
    return p

def bullet(text,level=0):
    p=doc.add_paragraph(style='List Bullet' if level==0 else 'List Bullet 2'); p.paragraph_format.space_after=Pt(3); p.paragraph_format.left_indent=Inches(.25+.22*level); p.paragraph_format.first_line_indent=Inches(-.16); r=p.add_run(text); font(r); return p

def table(headers, rows, widths, font_size=8.1):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    tPr=t._tbl.tblPr; tblW=tPr.find(qn('w:tblW')); tblW.set(qn('w:w'),str(sum(widths))); tblW.set(qn('w:type'),'dxa')
    grid=t._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for w in widths:
        c=OxmlElement('w:gridCol'); c.set(qn('w:w'),str(w)); grid.append(c)
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; set_cell_width(c,widths[i]); shade(c,NAVY); set_cell_margins(c)
        p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(str(h)); font(r,font_size,bold=True,color='FFFFFF')
    set_repeat_table_header(t.rows[0])
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,v in enumerate(row):
            set_cell_width(cells[i],widths[i]); set_cell_margins(cells[i]); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri%2: shade(cells[i],LIGHT)
            p=cells[i].paragraphs[0]; p.paragraph_format.space_after=Pt(0); r=p.add_run(str(v)); font(r,font_size)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

def callout(title,text,color=BLUE):
    t=doc.add_table(rows=1,cols=1); t.autofit=False; t.alignment=WD_TABLE_ALIGNMENT.CENTER; c=t.cell(0,0); set_cell_width(c,9360); shade(c,PALE); set_cell_margins(c,140,160,140,160)
    p=c.paragraphs[0]; r=p.add_run(title+'  '); font(r,10,bold=True,color=color); r=p.add_run(text); font(r,9.2)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

# Cover
doc.add_paragraph().paragraph_format.space_after=Pt(85)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('供应链与产能核查报告'); font(r,11,bold=True,color=GOLD)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(8); r=p.add_run('TERRA DRONE集团'); font(r,28,bold=True,color=NAVY)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(20); r=p.add_run('中国大陆、台湾及香港发货人｜子公司生产能力｜最新集团动向'); font(r,13,color=BLUE)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(105); r=p.add_run('贸易数据与公开情报交叉分析'); font(r,10,italic=True,color=GRAY)
table(['报告时点','数据底稿','证据口径'],[['2026年8月21日','用户提供的372行贸易明细（2025-01-03至2026-07-24）及易迅页面补查摘录','海关/贸易平台事实、公司公告、监管文件分层判断']],[2100,3500,3760],8.3)
ptext('重要说明：贸易平台的“金额、数量、原产地、进出口方向”以页面可见字段为准；同一商业批次可能拆成多行商品项。本文不把平台记录等同于海关全量，也不以HS编码单独认定管制属性。',italic=True,color=GRAY,after=0)
doc.add_page_break()

doc.add_heading('执行摘要',level=1)
callout('核心判断','Terra Drone已形成“日本总部整机与防务集成—印尼农业机组装—乌克兰战场型无人机制造—爱沙尼亚欧洲供应维护—欧美UTM软件”的多层结构。中国大陆、台湾及香港来源的动力、飞控、导航、链路、结构和电池足以支撑农业无人机成套组装，但现有贸易证据尚不能闭合证明这些部件进入日本军品或乌克兰战场。')
table(['结论','核查结果','置信度'],[
['贸易规模','底稿共372个商品明细行，覆盖2025-01-03至2026-07-24；PT Terra Drone Indonesia占257行、约1,037.6万元人民币申报价值，是集团中国供应链最集中节点。','高'],
['关键来源地','大陆发货人覆盖动力、机架、喷洒、遥控、飞控、测绘和电池；台湾链条主要为CubePilot飞控/RTK；香港链条主要为Emlid导航及数据传输模块。','中高'],
['组装能力','按电机、飞控、GPS、遥控等瓶颈交叉估算，印尼可形成约140–165架农业无人机的独立交付能力；187架仅是放宽“遥控器共享”假设后的理论上限。','中'],
['军用转用','G20、G30、T25/T50、B100等20L及以上、具自主航线/地形跟随能力的成品存在落入中国9A501参数口径的高风险；零件本身多不因单件性能自动受控。','中高'],
['战场流向','未发现PT Terra Drone Indonesia向乌克兰、日本或爱沙尼亚防务实体的可闭合出口提单，也未发现A1/A2/ATLA型号BOM出现印尼或中国来源部件。当前只能评为“供应链交叉风险”，不能写成已证实流入战场。','高'],
['集团新变化','2026年3月进入防务市场；6月并表Amazing Drones、WinnyLab并设Terra Defense Europe；7月获ATLA迎击项目并宣布日本电池量产；8月更新Terra Xross 1。','高']],[1550,5960,1850],8)

doc.add_heading('一、范围、口径与证据等级',level=1)
ptext('本报告的“发货人”优先采用申报/平台显示的供应商或境内发货人名称；品牌制造商、贸易商、报关行和最终制造商不自动等同。数量估算只对可识别为台/套/个的物项计算，桨叶和机架若以千克申报，不强行换算为架数。')
table(['等级','定义','报告中的用法'],[['A 已证实','公司公告、政府合同、海关申报或同批提单闭合','可陈述为事实'],['B 高度支持','多项独立信息一致，仍缺BOM/序列号/同票提单','写“高度可能/支持”'],['C 线索','单腿贸易、品牌或时间接近','只作核查方向'],['D 未发现','在当前数据范围内无命中','不等同于不存在']],[1450,3900,4010],8.2)

doc.add_heading('二、集团结构与各子公司生产/交付能力',level=1)
table(['实体/地区','主要角色与产品','制造/组装能力判断','状态'],[
['Terra Drone Corp.（日本）','集团总部；模块化UAV、Terra Xross 1、防务集成、电池规格定义','已签300架国产模块化训练机制造合同，2026-09-30交付；建立日本电池包设计、组装、检测与质保体系。具现实批量交付能力。','已证实'],
['PT Terra Drone Indonesia','G20、E16农业机；喷洒、测绘、培训','G20为“自研”并获31.26% TKDN；累计部署超150架，单日最高4,000架次；2025年计划向Yanmar交付约120架。贸易数据支持本地成套组装。','已证实'],
['Terra Drone Agri（马来西亚）','农业喷洒服务，使用E16/G20/G30等平台','收到54行零件/设备，但公开资料未确认独立工厂；更像运营、维修和区域调拨节点。','组装待核'],
['Amazing Drones（乌克兰）','Terra A1近程快速反应迎击无人机','2026-06并表；A1已在乌克兰实战拦截。具本地开发、制造和快速迭代能力，具体年产能未披露。','已证实'],
['WinnyLab LLC（乌克兰）','Terra A2固定翼迎击机；312km/h、覆盖75km、续航40分钟以上','Terra Defense Europe持股50%并成为最大股东；A2已实战部署。具体产线/产量未公开。','已证实'],
['Besomar（乌克兰）/拟设JV','Terra C1侦察监视固定翼UAV','2026-06启动JV筹备；Besomar自称开发制造固定翼系统。Terra C1处于产品展开阶段，集团控制关系尚弱于并表子公司。','推进中'],
['Terra Defense Europe（爱沙尼亚）','A1/A2/C1欧洲销售、供应、维修、物流和伙伴协调','公开定位不含制造；可能承担备件仓储、维修和本地集成，但不能据此认定整机生产。','非制造节点'],
['Terra Drone Arabia（沙特）','油气巡检、测绘、本地化；接收Sky Whale Mini等','短期服务本地化，长期规划R&D和生产设施；截至报告日未见量产公告。','规划阶段'],
['Terra Drone USA / Aloft（美国）','Aloft UTM/LAANC与机队管理；Terra Xross 1渠道','Aloft已转全资子公司，软件能力强；未发现美国整机制造证据，硬件以渠道销售和应用集成为主。','软件/渠道'],
['Unifly NV（比利时）','国家级UTM、BVLOS、与MBDA合作C-UAS','软件与空域管理，不是无人机整机制造商；可成为防务系统指挥/协同层。','软件平台'],
['Terra Inspectioneering（荷兰）','工业罐体/室内检测；集团对乌投资执行平台','具检测方案与产品化能力；公开信息不足以量化其独立整机产能。','系统集成']],[1800,2580,3900,1080],7.35)

doc.add_heading('三、贸易数据总览',level=1)
table(['集团收货实体','商品明细行','申报价值（人民币）','主要特征'],[
['PT Terra Drone Indonesia',257,'10,375,580.34','动力、桨叶、飞控/GPS、H12链路、机架、泵/喷洒、播撒、电池'],
['Terra Drone Agri Sdn Bhd',54,'799,028.05','农业机零件、运营及备件'],
['Terra Drone Japan',49,'3,350,244.22','工业/测绘设备、无人机及零件；部分经印尼或直接到日'],
['Terra Drones S.A.（哥斯达黎加）',6,'未统一','区域服务设备'],
['Terra Drone Arabia',4,'未统一','固定翼整机/跟踪与控制附件'],
['Terra Drone Inc.',2,'未统一','少量设备/部件']],[2550,1300,2100,3410],8)
ptext('注：372行是商品明细行，不等于372票提单；金额为底稿中的人民币价值字段。最新底稿日期为2026-07-24。易迅补查中可确认的最新香港来源记录为2026-06-22数据传输模块。',italic=True,color=GRAY)

doc.add_heading('四、中国大陆发货人及所供物项',level=1)
table(['申报/平台发货人','地区','主要物项','主要集团收货方','判断'],[
['深圳市好盈科技股份有限公司（Hobbywing）','广东深圳','X8/X9 Plus/X13动力套装、电机/电调、桨叶','印尼、马来西亚','核心动力供应商'],
['重庆智睿之鸟贸易有限公司','重庆','E16/G20类机架、碳管、连接件、喷头、播撒套件、H12遥控','印尼','成套化程度高的贸易/集成节点'],
['杭州启飞创新科技有限公司','浙江杭州','B100农业无人机、X13配套及农业机部件','印尼','整机及重载平台供应'],
['上海极翼智能科技有限公司','上海','K++ V2飞控、GPS、农业机控制系统','印尼、马来西亚','飞控供应'],
['合肥翼飞特电子科技有限公司（EFT）','安徽合肥','E16/G20类机架、药箱及结构件','印尼','机体平台供应'],
['深圳嘉盈时代科技有限公司','广东深圳','H12遥控、图数传及附件','印尼','控制链路供应'],
['宁波高新区阶梯科技有限公司','浙江宁波','GNSS/RTK与Emlid相关产品','印尼、日本等','导航产品渠道/技术供应'],
['南昌睿极贸易有限公司','江西南昌','无人机零部件及配套','印尼','贸易节点，制造方待穿透'],
['北京卫固新能源科技有限公司','北京','高压锂电池组/模块','印尼','能源系统供应'],
['深圳市扬达科技有限公司','广东深圳','工业/固定翼无人机及载荷','沙特等','整机/系统供应'],
['上海华测导航技术股份有限公司（CHCNAV）','上海','GNSS、测绘定位设备','日本/集团实体','高精度导航供应'],
['深圳市道通智能航空技术股份有限公司（Autel）','广东深圳','整机、相机/载荷及附件','集团实体','商用整机供应'],
['昆山科迪威动力科技有限公司','江苏昆山','动力/电池相关产品','印尼','能源/动力供应'],
['深圳市飞盈佳乐电子有限公司','广东深圳','电调/电机等电子动力件','印尼','动力电子供应'],
['广州市苏纳米实业有限公司','广东广州','无人机安全箱线索','拟/传闻对应Terra Drone','在现有易迅供应商清单中未命中，关系未证实']],[2500,950,2650,1800,1460],7.1)

doc.add_heading('五、台湾与香港发货链条',level=1)
table(['发货人/品牌','平台显示来源','日期及可见记录','货物','判断'],[
['CubePilot Global / Hex Technology链条','台湾','2026-05-08：Cube Orange+ Standard Set 10套；Here4 RTK/GNSS 10套','飞控与多频RTK导航','直接补齐农业机/工业机航电，台湾来源明确；是否由香港贸易主体开票需调原始提单。'],
['Emlid链条','香港','2026-04-16：Reach RS4 8台、Reach M2 16台及天线/转接件','RTK/GNSS测量与机载定位','可能兼用于测绘和无人机导航；原产国字段“香港”不等于香港制造。'],
['SHENZHEN YOUSAN TECH CO., LTD.','香港（平台原产地）','2026-06-22：数据传输模块16件，HS 85366924，金额字段6,295.79','数据传输/连接模块','公司名含Shenzhen但平台原产地为香港，需核对注册地、发票方与实际制造地。'],
['香港转口/贸易链（汇总）','香港','底稿另有4行目的地/来源与香港相关记录','导航、链路及配套','应将“起运地、原产地、发货人注册地”三字段分开，不能把香港起运自动写成香港制造。']],[2300,1300,2750,1750,2260],7.45)

doc.add_heading('六、可识别零件数量与可组装能力',level=1)
table(['系统','可识别数量','折算/约束'],[
['动力系统', '2,526台/套电机或动力单元','X8 636、X9 Plus 676、X13 228、5900W 280及其他；按四旋翼折算的动力上限高于航电瓶颈。'],
['飞控', 'K++ V2 195套','若一机一套，理论上限195架。'],
['GNSS/GPS', '187套（另有RTK补查件）','若一机一套，理论上限约187架。'],
['遥控/链路', 'H12 140套，其他遥控25套','独立交付的最强瓶颈约165套；若地面端共享则可高于此数。'],
['作业系统', '泵355；播撒套件238；独立播撒器15','足以覆盖大量农业机，但泵、喷头、管路、阀件需按机型配套核对。'],
['电池', '完整电池包1,407；模块100；电芯600','数量充分，但电压、接口、BMS和尺寸不通用。'],
['桨叶', '约1,690.76千克','因法定单位多为重量，无法可靠折算叶片对数。'],
['完整整机', 'B100农业机4架；Sky Whale Mini固定翼1架','与零件组装估算分开统计。']],[1750,2500,5110],8)
callout('组装结论','现有物料在功能类别上已覆盖动力、控制、升力、作业载荷、结构、链路、导航和能源，足以支持完整农业无人机的组装；但逐架BOM尚不闭合。按“飞控—导航—独立遥控链路”瓶颈估算约140–165架，187架为放宽遥控器共享后的理论上限。')
doc.add_heading('可能对应的机型',level=2)
table(['平台','关键匹配','理论动力折算','判断'],[
['E16/16L级','X8动力、E16机架、16L作业系统','636÷4≈159架','与遥控/飞控瓶颈叠加后，是140–159架区间的主要候选。'],
['G20/20L级','X9 Plus、36×19桨、G20结构/20L作业系统','676÷4≈169架','与官方G20本地化及Yanmar 120架交付计划高度吻合。'],
['B100/重载级','X13动力、50L喷洒/70L播撒','228÷6≈38架，另有4架整机','仅动力支持约38架，完整BOM和实际组装量未闭合。'],
['G30/其他重载平台','5900W或其他大功率动力','280÷4≈70或÷6≈46架','机架、ESC、药箱与飞控匹配不充分，置信度较低。']],[1500,3000,1900,2960],8)

doc.add_heading('七、出口管制与潜在转用风险',level=1)
ptext('中国现行两用物项出口管制清单中，9A501重点关注非9A012无人机在自主控制/导航或超视距操控条件下，具备20升及以上气溶胶喷洒系统的情形，以及航程300千米以上的无人机。是否受控取决于技术参数和用途，不由HS编码单独决定。')
table(['产品/部件','参数触发分析','风险'],[
['G20','20L/20kg，自动航线、避障、地形跟随；与9A501“≥20L+自主能力”高度重合。','高'],
['T25/T50、G30','20L/40L/30L级且具有自动作业能力，参数重合度高。','高'],
['B100','50L喷洒/70L播撒，重载自主农业机；参数重合度高。','高'],
['E16','16L低于20L门槛；如不满足300km航程等其他条件，单凭载荷不触发该项。','中低'],
['X8/X9/X13电机、电调、桨叶','单件功率低于无人机发动机16kW相关阈值，通常不因该参数单独受控；但仍需核查专门设计、最终用途和拆分规避。','中'],
['CubePilot/Emlid/H12','通用飞控/导航/链路并非自动等同受控；需核对抗干扰、加密、军标、作用距离及是否专门设计。','中'],
['Terra A1/A2/C1','明确防务用途、实战部署/侦察迎击；应按军品及相关法域出口管制审查，不宜沿用民用农业机判断。','极高']],[1800,5700,2060],8)

doc.add_heading('八、流入日本防务或乌克兰战场的证据评估',level=1)
table(['待证链条','当前发现','结论'],[
['印尼→乌克兰/爱沙尼亚防务实体','未发现PT Terra Drone Indonesia向Amazing Drones、WinnyLab或Terra Defense Europe的出口提单。','未证实'],
['印尼→日本军用项目','存在部分PT Indonesia货物最终目的地日本，以及集团对日贸易；但未出现ATLA合同号、军用型号或序列号闭合。','仅线索'],
['A1/A2/C1 BOM','公开公告披露性能和开发方，未披露印尼、中国大陆、台湾或香港部件。','无法判断'],
['供应链能力交叉','印尼已掌握农业无人机成套组装；集团2026年防务并购和欧洲供应体系形成。','组织风险上升，但不是物流证据'],
['集团去中国化信号','2026-07宣布在日本量产非中国来源电池，并拟推进电机、控制器国产化。','说明既有海外依赖真实存在，也说明防务项目可能主动隔离中国供应链']],[2350,4700,2310],8.1)
callout('证据边界','目前最稳妥的表述是：存在中国/台港零件进入Terra Drone民用与农业供应链的确证，也存在集团乌克兰实战和日本防务采购的确证；但两条证据链尚未通过BOM、序列号、同批提单或内部采购单闭合。')

doc.add_heading('九、2026年集团最新动向时间线',level=1)
table(['日期','事件','战略含义'],[
['2026-03-23','宣布正式进入防务装备市场','集团战略从民用服务/UTM扩展到防务硬件。'],
['2026-04-17至05-19','Terra A1、A2在乌克兰投入实战；先后披露拦截成果','形成真实战场验证与快速迭代闭环。'],
['2026-05-08','获ATLA 300架日本国产模块化训练机合同，金额1.15434亿日元','日本总部具备批量制造交付责任。'],
['2026-06-15','并表Amazing Drones与WinnyLab；成立Terra Defense Europe；推进Besomar JV/Terra C1；Unifly与MBDA合作','构建A1近程迎击+A2广域迎击+C1侦察+UTM/C-UAS集成。'],
['2026-07-07','与UAE EDGE集团IGG签防务应用合作MOU','中东防务市场和系统集成渠道扩大。'],
['2026-07-15','入选日本ATLA迎击无人机快速采购项目','乌克兰验证产品向日本防务采购体系迁移。'],
['2026-07-27','启动日本国产高功率圆柱电芯电池包量产体系','明确采用非中国部件满足NDAA方向，并计划外销。'],
['2026-08-04','与Chevron、Shell参与的DeepStar合作开发Terra Xross 1信号恢复功能','工业室内检测产品继续迭代并拓展北美高风险场景。']],[1600,4650,3110],8)

doc.add_heading('十、风险排序与建议',level=1)
table(['优先级','风险点','建议调证'],[
['P1','印尼成套供应链与集团防务业务在组织层面已汇合，但物流链未闭合','调取PT Indonesia 2026年全部出口、集团内部调拨单、最终用户声明；重点筛Estonia/Ukraine/Japan及Terra Defense/Amazing/WinnyLab/Besomar。'],
['P1','G20/B100等成品存在9A501参数触发风险','取得逐型号规格书、控制模式、喷洒容量、最大航程、最终用途及许可证。'],
['P1','日本ATLA 300架及迎击项目供应链来源不透明','调取BOM、供应商名录、原产地声明、NDAA/非中国来源合规文件，并与大陆/台港零件型号比对。'],
['P2','CubePilot/Emlid/H12链条可能经香港/台湾贸易主体转口','核对发票方、制造地、起运地、最终收货仓、序列号与同批提单。'],
['P2','中国发货人中存在贸易公司，实际制造商被遮蔽','穿透重庆智睿之鸟、南昌睿极等的采购来源、关联企业、收款与报关代理。'],
['P3','苏纳米“无人机安全箱”线索尚无平台命中','以SUNAMI/SUNAMI INDUSTRIAL/GUANGZHOU SUNAMI、品名DRONE CASE/SAFETY CASE/FLIGHT CASE及HS4202/3926/7616补查。']],[900,3550,4910],8)

doc.add_heading('附录A：数据缺口与估算限制',level=1)
for x in [
'用户此前提供的Excel附件在本次工作目录中已不可重新打开，因此本报告沿用已完成的372行分析底稿汇总值；如需逐票附表，应重新附上原始Excel。',
'易迅网页本次可连接并确认登录，但连续交互出现超时；报告采用此前已逐页读取并保存的最新补查摘录，不声称覆盖2026-07-24之后全部新增记录。',
'动力、桨叶、机架和结构件存在重量单位申报，无法按一行一件换算；组装架数是BOM瓶颈估算，不是实际产量。',
'集团公告中的“国产”“自研”“本地含量”分别指合同、知识产权/产品责任和TKDN口径，不能互相替代。',
'未发现某条提单仅表示当前数据库和关键词范围未命中，不构成“不存在”的证明。']:
    bullet(x)

doc.add_heading('附录B：主要公开来源',level=1)
sources=[
('Terra Drone 2026-08-04：Terra Xross 1信号恢复功能','https://terra-drone.net/global/'),
('Terra Drone 2026-07-27：日本国产无人机电池业务','https://terra-drone.net/global/2026/07/27/27927/'),
('Terra Drone 2026-07-15：ATLA迎击无人机快速采购项目','https://terra-drone.net/global/2026/07/15/27784/'),
('Terra Drone 2026-06-15：WinnyLab并表及Terra A2','https://terra-drone.net/global/2026/06/15/terra-drone-winnylab-subsidiary/'),
('Terra Drone 2026-06-15：Terra Defense Europe','https://terra-drone.net/global/2026/06/15/terra-drone-establishes-terra-defense-europe/'),
('Terra Drone 2026-05-08：300架模块化UAV合同','https://terra-drone.net/global/2026/05/08/terra-drone-secures-jpy-115-4-million-order-from-japans-acquisition-technology-logistics-agency-for-300-domestically-produced-modular-uav-general-purpose-units/'),
('Terra Drone 2025-10-16：印尼G20 TKDN认证','https://terra-drone.net/global/2025/10/16/terra-drones-group-company-terra-drone-indonesia-obtains-tkdn-certification-for-its-in-house-developed-agricultural-drone-g20/'),
('Terra Drone 2025-08-28：与Yanmar Indonesia销售合作','https://terra-drone.net/global/2025/08/28/terra-drone-signs-sales-partnership-agreement-with-yanmar-diesel-indonesia-for-its-in-house-developed-agricultural-drones/'),
('Terra Drone 2025-09-23：Aloft全资子公司化','https://terra-drone.net/global/2025/09/23/12042/'),
('Terra Drone 2025-04-03：与Aramco合作及沙特本地化','https://terra-drone.net/global/2025/04/03/terra-drone-signs-mou-with-aramco-to-drive-innovation-and-localization-in-drone-technology/'),
('中国两用物项出口管制信息平台：常见问题（HS并非唯一判断依据）','https://exportcontrol.mofcom.gov.cn/article/cjwt/202504/1130.html'),
('国家出口管制协调机制办公室：两用物项出口管制清单PDF','https://www.nca.gov.cn/sca/xwdt/2024-11/25/1061218/files/de436ad19392485b9eb0c7020858b2e4.pdf')]
for i,(name,url) in enumerate(sources,1): ptext(f'{i}. {name}\n{url}',after=4)

doc.add_heading('附录C：建议持续监控关键词',level=1)
ptext('收货人：PT TERRA DRONE INDONESIA；TERRA DRONE AGRI SDN BHD；TERRA DRONE USA；ALOFT TECHNOLOGIES；TERRA DRONE ARABIA；TERRA DEFENSE EUROPE；AMAZING DRONES；WINNYLAB；BESOMAR。')
ptext('商品：DRONE PARTS、UAV PARTS、AGRICULTURAL DRONE、INTERCEPTOR、FLIGHT CONTROLLER、RTK/GNSS、DATA TRANSMISSION MODULE、MOTOR ESC、PROPELLER、BATTERY PACK、DRONE CASE/SAFETY CASE。')
ptext('路线：China/Taiwan/Hong Kong → Indonesia/Malaysia/Japan/Estonia/Ukraine/Saudi Arabia；并重点观察印尼→日本、爱沙尼亚、乌克兰的二次流向。')

doc.core_properties.title='Terra Drone集团供应链与产能核查报告'
doc.core_properties.subject='中国大陆、台湾和香港发货人；子公司无人机生产能力；最新公开动态'
doc.core_properties.author='Codex'
doc.save(OUT)
print(OUT)
