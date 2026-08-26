from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path

OUT = Path('outputs/韩国两家材料企业国内合作线索及黄埔口岸风险专题_核查稿.docx')
OUT.parent.mkdir(parents=True, exist_ok=True)

NAVY = '17365D'; BLUE = '2E74B5'; LIGHT = 'E8EEF5'; GRAY = 'F2F4F7'; RED = '9B1C1C'; GOLD = '7A5A00'; INK='222222'

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
sec.header_distance = sec.footer_distance = Inches(.492)

styles = doc.styles
normal = styles['Normal']; normal.font.name='Microsoft YaHei'; normal.font.size=Pt(10.5)
normal._element.rPr.rFonts.set(qn('w:eastAsia'),'微软雅黑')
normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.1
for name,size,before,after,color in [('Title',24,0,6,NAVY),('Heading 1',16,16,8,BLUE),('Heading 2',13,12,6,BLUE),('Heading 3',11.5,8,4,NAVY)]:
    st=styles[name]; st.font.name='Microsoft YaHei'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'微软雅黑'); st.font.size=Pt(size); st.font.color.rgb=RGBColor.from_string(color); st.font.bold=True
    st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after); st.paragraph_format.keep_with_next=True

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn('w:shd'))
    if shd is None: shd=OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'),fill)
def margins(cell,top=80,start=120,bottom=80,end=120):
    tc=cell._tc.get_or_add_tcPr(); mar=tc.first_child_found_in('w:tcMar')
    if mar is None: mar=OxmlElement('w:tcMar'); tc.append(mar)
    for tag,val in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        el=OxmlElement('w:'+tag); el.set(qn('w:w'),str(val)); el.set(qn('w:type'),'dxa'); mar.append(el)
def set_cell_text(cell,text,bold=False,color=INK,size=9):
    cell.text=''; p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.05
    r=p.add_run(text); r.bold=bold; r.font.name='Microsoft YaHei'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'微软雅黑'); r.font.size=Pt(size); r.font.color.rgb=RGBColor.from_string(color)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; margins(cell)
def table(headers, rows, widths):
    t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.LEFT; t.autofit=False; t.style='Table Grid'
    trPr=t.rows[0]._tr.get_or_add_trPr(); rep=OxmlElement('w:tblHeader'); rep.set(qn('w:val'),'true'); trPr.append(rep)
    for i,h in enumerate(headers): set_cell_text(t.rows[0].cells[i],h,True,NAVY,9); shade(t.rows[0].cells[i],LIGHT); t.rows[0].cells[i].width=Inches(widths[i])
    for row in rows:
        cells=t.add_row().cells
        trPr=t.rows[-1]._tr.get_or_add_trPr(); cant=OxmlElement('w:cantSplit'); cant.set(qn('w:val'),'true'); trPr.append(cant)
        for i,v in enumerate(row): set_cell_text(cells[i],str(v),False,INK,8.6); cells[i].width=Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t
def bullet(text,level=0):
    p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.left_indent=Inches(.5); p.paragraph_format.first_line_indent=Inches(-.25); p.paragraph_format.space_after=Pt(5); p.add_run(text); return p
def callout(label,text,color=RED):
    t=doc.add_table(rows=1,cols=1); t.alignment=WD_TABLE_ALIGNMENT.LEFT; t.autofit=False; c=t.cell(0,0); c.width=Inches(6.5); shade(c,'F4F6F9'); margins(c,120,160,120,160)
    c.text=''; p=c.paragraphs[0]; r=p.add_run(label+'  '); r.bold=True; r.font.color.rgb=RGBColor.from_string(color); r2=p.add_run(text); r2.font.color.rgb=RGBColor.from_string(INK)
    for x in (r,r2): x.font.name='Microsoft YaHei'; x._element.rPr.rFonts.set(qn('w:eastAsia'),'微软雅黑'); x.font.size=Pt(10)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

# header/footer
h=sec.header.paragraphs[0]; h.text='关键矿产出口风险专题｜线索核查稿'; h.alignment=WD_ALIGN_PARAGRAPH.RIGHT
for r in h.runs: r.font.size=Pt(8.5); r.font.color.rgb=RGBColor(100,100,100)
f=sec.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=f.add_run('内部研判材料｜截至 2026年7月13日'); r.font.size=Pt(8.5); r.font.color.rgb=RGBColor(100,100,100)

p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(20); p.paragraph_format.space_after=Pt(2)
r=p.add_run('专题核查报告'); r.bold=True; r.font.size=Pt(10); r.font.color.rgb=RGBColor.from_string(GOLD)
doc.add_paragraph('韩国两家材料企业国内合作线索及黄埔口岸风险分析',style='Title')
p=doc.add_paragraph('围绕 Seoul Metal Tech、Korea Advanced Material 与境内供货主体的证据分级、风险识别及黄埔口岸指向性核查')
p.style='Subtitle'
callout('结论先行','现有材料足以形成风险线索池，但不足以把全部主体列为“已立案查实企业”。对黄埔口岸的直接指向目前仅来自匿名走访线索；东莞主体与黄埔海关关区存在监管联系，但不能据此证明货物实际经黄埔口岸出境。',RED)

doc.add_heading('一、核查结论',1)
bullet('企业层面：共梳理13个境内主体条目，其中1个匿名、1个为描述性名称，另有2个企业名称疑似不完整。公开资料仅对杭州新川、南京达迈的部分海关违法事实形成交叉印证；其余与两家韩企的交易关系、拆票数量、最终用户和转运日本等描述，尚需报关单、提运单、合同、付款及海关文书全文支撑。')
bullet('韩国主体层面：公开资料可确认名为 Korea Advanced Material Co., Ltd. 的韩国企业存在，但其公开主营为铂族金属化合物和催化剂，不能自然推导为镍基高温合金采购商；Seoul Metal Tech 公开信息更偏汽车零部件/机械加工集团，未找到其采购中国军工级镍合金并转供日本的公开证据。')
bullet('黄埔层面：匿名“深圳XX金属材料贸易有限公司”被描述为轮换盐田、黄埔口岸，这是唯一直接口岸指向；东莞长安、沙田企业可作为黄埔海关关区协查对象，但必须用申报地海关代码、出境关别、运输工具及舱单确认是否实际走黄埔口岸。')
bullet('文书号层面：材料中的“南东关知字”与知识产权类文书编号样式高度相似，和镍合金涉证/申报不实案的叙述存在类型错配风险；应优先调取原文核验案由、当事人、货物和处罚依据。')

doc.add_heading('二、证据分级与口径调整',1)
table(['等级','定义','本专题采用口径'],[
('A 已公开核实','官方文书/权威平台可对应到企业与具体违法事实','仅写公开文书明确记载的行为，不外推境外客户或最终用途'),
('B 关联可确认','企业官网、展会名录、官方口岸资料可确认主体、产品或管辖关系','可作为背景，不作为违法证明'),
('C 待核线索','走访、行业调研、匿名描述或单一来源指控','使用“据材料反映/待核”，不得写成定论'),
('D 分析推断','由地域、产品、增幅或经营能力推导的风险判断','必须标明推断依据及替代解释')
],[1.05,1.8,3.65])

doc.add_heading('三、境内企业清单与核查状态',1)
rows=[
('澜笙科技（深圳）有限公司','Seoul Metal Tech','C','材料称有“南东关知字〔2026〕0054号”；公开检索未核到原文，且“关知字”通常提示知识产权案件编号，需防止文书与案情错配。','未指向黄埔'),
('深圳XX金属材料贸易有限公司（匿名）','Seoul Metal Tech','C','材料称1.8吨合金丝拆36票并轮换盐田、黄埔；无全称、统一信用代码、报关单号。','直接线索，证据不足'),
('东莞市扬驰精密五金有限公司','Seoul Metal Tech','C','材料称“深邮关缉缉违字〔2025〕73号”；编号存在重复“缉”，公开未核到原文。','关区协查，不等于口岸'),
('深圳跨境新材料进出口商行（描述性名称）','Korea Advanced Material','C','缺完整登记名称及识别码；出口增幅未给出统计口径、基期和数据表。','未指向黄埔'),
('东莞鑫源金属制品贸易有限公司','Korea Advanced Material','C','拆22票、韩方拒绝产能证明均为单一调研说法。','关区协查，不等于口岸'),
('深圳煌铭旺和五金有限公司','未明确/韩商','C','仅称2026年二季度开始接收小额拆分订单。','未指向黄埔'),
('深圳市铧宁金属材料有限公司','两家韩国贸易商','C','“分时段小额申报”缺报关底表。','未指向黄埔'),
('惠州金泽丰贸易有限公司','未明确/韩商','C','“搭配管制镍材混装”属严重指控，未附查验或检测依据。','未指向黄埔'),
('杭州新川新材料有限公司','Seoul Metal Tech','A/C','公开可对应沪外港关缉违字〔2026〕143号及“申报不实、影响许可证件管理”；韩国客户、拆票和日本最终用户仍为C级。','无黄埔指向'),
('南京达迈科技实业股份有限公司','Korea Advanced Material','A/B/C','公开案例支持阿拉山口无证出口含钇镍合金丝并被罚1.6万元；官网支持其具备高温合金等生产能力。材料所称韩企客户、拆分混装、日本军工链未获公开支持。','无黄埔指向'),
('北京友兴联有色金属有限公司','两家韩企','C','批量供货和拆票说法待交易数据验证。','无黄埔指向'),
('无锡天裕金属（名称待补全）','未明确','C','企业全称不完整，跨口岸记录未附单号。','无黄埔指向'),
('冶韩实业（上海，名称待补全）','未明确','C','企业全称不完整，跨口岸记录未附单号。','无黄埔指向')]
table(['境内主体','所称韩方','等级','可确认事实/主要缺口','黄埔关联'],rows,[1.35,1.05,.55,2.7,.85])

doc.add_heading('四、黄埔口岸指向性专题分析',1)
doc.add_heading('（一）目前能“直接指向”的只有匿名线索',2)
bullet('匿名深圳XX：材料明确使用“轮换盐田、黄埔口岸”表述，因此可列为黄埔口岸直接线索；但缺少企业真名、报关单号、出境关别、口岸代码、提运单和船名航次，无法落到具体企业或具体票次。')
bullet('该线索还声称经韩国仓分装后直供日本关东特气企业。此段需要韩国进口申报、仓储/换标记录、日本进口记录或收款链闭环；仅凭国内走访不足以认定最终用户和军事用途。')
doc.add_heading('（二）东莞企业属于“黄埔海关关区协查线索”，不是口岸证据',2)
bullet('公开海关资料显示黄埔海关在东莞沙田开展监管业务，说明东莞企业可能处于黄埔海关关区管理或业务覆盖范围。')
bullet('但企业注册地、主管海关、申报地海关和实际出境口岸是四个不同字段。东莞长安/沙田企业可以在本地申报后转关至其他口岸，也可能通过深圳盐田、广州南沙等口岸出境。')
bullet('因此，东莞市扬驰、东莞鑫源只能列为黄埔关区优先筛查对象，不宜直接写作“经黄埔口岸拆分出口”。')
doc.add_heading('（三）建议锁定的黄埔数据字段',2)
table(['核查对象','关键字段','判定用途'],[
('报关单','境内发货人、生产销售单位、申报地海关、出境关别、监管方式、商品编码、品名规格、数量/重量、总价、许可证号','确认企业、票次、商品及口岸'),
('物流舱单','提运单号、船名航次、装货港/卸货港、集装箱号、订舱人、货代','确认实际从黄埔哪个作业区出境'),
('合同与资金','合同号、订单拆分时间、单票货值、收款人、韩方付款账户、关联发票','识别人为拆单和同一交易整体性'),
('查验与检测','查验记录、取样送检、材质成分、牌号、用途说明','识别伪报废料/普通合金及涉证属性'),
('境外收货人','韩文/英文全称、地址、企业注册号、统一拼写变体','归并两家韩企及其仓库/代理')
],[1.1,3.8,1.6])

doc.add_heading('五、主要风险类型',1)
for title,txt in [
('1. 出口管制与涉证风险','高温合金、镍粉或含受控元素合金是否属于两用物项，取决于具体成分、性能参数、形态和用途，不能仅凭“镍制品”概括认定。若属于管制物项而未取得许可，可能触发出口管制法责任。'),
('2. 申报不实及伪报品名风险','将高温合金伪报为普通线材、再生不锈钢废料，或隐瞒材质成分、牌号、用途，可能影响许可证件管理及海关监管。'),
('3. 人为拆单规避监管风险','同一买卖合同在短周期内由同一或关联主体拆成多票、跨货代跨口岸申报，若目的在规避许可证件、查验或监管阈值，应合并交易链研判。单纯“小额多票”本身只是风险指标，不等同违法。'),
('4. 监管方式错用风险','材料提及1039市场采购与国际快件。需核查货物是否符合对应监管方式、出口主体和场所资格，是否以不匹配渠道掩盖真实贸易。'),
('5. 最终用户与转运风险','韩国收货人若仅仓储、分装、换标并转运第三国，需审查最终用户、最终用途和再出口安排；目前“日本军工/特气企业”均为待证实线索。'),
('6. 名誉与执法质量风险','将未经核实的企业合作关系、军事最终用途或走私方式公开定性，可能导致错误执法线索、企业名誉损害及报告可信度下降。')]:
    doc.add_heading(title,2); doc.add_paragraph(txt)

doc.add_heading('六、核查优先级与行动清单',1)
table(['优先级','对象','立即核查事项','完成标准'],[
('P0','匿名深圳XX','从走访底稿反查联系人、地址、电话、货代及合同；以1.8吨/36票/韩国收货人组合筛查盐田与黄埔数据','识别真实主体并锁定至少1票报关单'),
('P0','杭州新川、南京达迈','调取处罚决定书全文与报关随附单证；分离“已处罚事实”和材料追加叙述','形成逐项证据对照表'),
('P1','东莞扬驰、东莞鑫源','按统一信用代码归集黄埔关区及全国跨关区出口，重点比对同日/邻近日期、同收货人、同品名规格','确认是否存在同一订单拆票'),
('P1','两家韩国主体','统一英文/韩文名称、地址、注册号和关联仓库；核对业务范围及进口记录','消除同名主体误并风险'),
('P2','其余企业','补全企业全称、信用代码、生产能力、海关注册和历史出口','由模糊名单转为可执行主体表')
],[.55,1.25,3.25,1.45])

doc.add_heading('七、建议采用的正式表述',1)
callout('建议标题','“韩国材料企业涉镍制品境内供货线索及黄埔口岸拆分申报风险核查专题（线索核查稿）”',NAVY)
bullet('不建议使用：“已查实有拆分小额报关违规记录的核心客户”，除非每一家均有可核验的处罚/刑事文书且文书事实与拆票、货物和韩方客户逐项一致。')
bullet('建议使用：“公开文书已确认部分违法事实的企业”“产业调研待核线索企业”“其他供货关系待核企业”三类。')
bullet('对黄埔表述建议限定为：“发现1条直接提及黄埔口岸的匿名走访线索，并有2家东莞主体具备黄埔关区协查必要性；尚未取得具体报关单证证明其实际经黄埔口岸出境。”')

doc.add_heading('八、来源与核验说明',1)
sources=[
('用户提供的基础材料','本专题所有未公开印证的企业关系、拆票数量、口岸轮换及最终用户描述均源于该材料，按C级线索处理。'),
('上海国际经贸合规法律服务平台—南京达迈案例','https://www.tclegal.cn/compliace-case-detail?id=2013522130024599553&type=1'),
('南京达迈科技官网','https://www.njdamai.cn/'),
('Korea Advanced Material 展会资料（IMID）','https://www.imid.or.kr/2023/download/ex/34_IMID_2023_Exhibitor_Company_Instruction_Form_Koadvanced.pdf'),
('海关总署英文网—东莞沙田黄埔海关监管活动','https://english.customs.gov.cn/statics/34b551ae-49e3-4356-85fd-b90dda81a9ae.html'),
('广州市黄埔区政府—黄埔老港镍等工业原料通关背景','https://www.hp.gov.cn/xwzx/bmdt/content/post_10891376.html'),
('关务资讯索引—杭州新川处罚文书条目（非官方镜像，需以海关原文为准）','https://www.jitinfo.net/Service/Category?type=%E6%95%B0%E6%8D%AE%E6%9C%8D%E5%8A%A1')]
for name,url in sources:
    p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(4); p.add_run(name+'：').bold=True; p.add_run(url)

p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(12); p.paragraph_format.keep_together=True
r=p.add_run('使用限制：'); r.bold=True; r.font.color.rgb=RGBColor.from_string(RED)
p.add_run('本稿用于内部线索研判，不构成对任何企业违法行为、客户关系、最终用户或军事用途的事实认定。正式立案、通报或对外发布前，应以海关业务系统、处罚文书原件及依法调取的交易物流证据复核。')

doc.core_properties.title='韩国两家材料企业国内合作线索及黄埔口岸风险专题（核查稿）'
doc.core_properties.subject='关键矿产与出口管制风险线索核查'
doc.core_properties.author=''
doc.core_properties.keywords='黄埔口岸, 镍合金, 出口管制, 拆分申报, 线索核查'
doc.save(OUT)
print(OUT.resolve())
