from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
from datetime import date

OUT=Path(r'D:\易迅数据\反倾销税风险\氯氰菊酯贸易流向及反倾销税风险分析报告.docx')
OUT.parent.mkdir(parents=True,exist_ok=True)

NAVY='17365D'; BLUE='2F5597'; LIGHT='D9EAF7'; PALE='EEF4F8'; RED='C00000'; AMBER='BF7000'; GRAY='666666'; GRID='B8C4CE'

def font(run,size=10.5,bold=False,color='222222',name='Microsoft YaHei'):
    run.font.name=name; run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),name)
    run._element.rPr.rFonts.set(qn('w:ascii'),'Arial'); run._element.rPr.rFonts.set(qn('w:hAnsi'),'Arial')
    run.font.size=Pt(size); run.bold=bold; run.font.color.rgb=RGBColor.from_string(color)
    return run
def shade(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn('w:shd'))
    if shd is None: shd=OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'),fill)
def margins(cell,top=90,start=110,bottom=90,end=110):
    tcPr=cell._tc.get_or_add_tcPr(); mar=tcPr.first_child_found_in('w:tcMar')
    if mar is None: mar=OxmlElement('w:tcMar'); tcPr.append(mar)
    for tag,val in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        e=mar.find(qn('w:'+tag))
        if e is None: e=OxmlElement('w:'+tag); mar.append(e)
        e.set(qn('w:w'),str(val)); e.set(qn('w:type'),'dxa')
def set_cell(cell,text,bold=False,color='222222',size=8.8,fill=None,align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text=''; p=cell.paragraphs[0]; p.alignment=align; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.05
    font(p.add_run(str(text)),size,bold,color); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; margins(cell)
    if fill: shade(cell,fill)
def table(doc,headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    trPr=t.rows[0]._tr.get_or_add_trPr(); rep=OxmlElement('w:tblHeader'); rep.set(qn('w:val'),'true'); trPr.append(rep)
    for i,h in enumerate(headers): set_cell(t.rows[0].cells[i],h,True,'FFFFFF',8.8,NAVY,WD_ALIGN_PARAGRAPH.CENTER)
    for ridx,row in enumerate(rows):
        cells=t.add_row().cells
        for i,v in enumerate(row): set_cell(cells[i],v,False,'222222',8.3,PALE if ridx%2 else 'FFFFFF')
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    return t
def para(doc,text='',bold_lead=None,color='222222',size=10.5,after=5,align=None):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(after); p.paragraph_format.line_spacing=1.22
    if align is not None: p.alignment=align
    if bold_lead and text.startswith(bold_lead):
        font(p.add_run(bold_lead),size,True,color); font(p.add_run(text[len(bold_lead):]),size,False,color)
    else: font(p.add_run(text),size,False,color)
    return p
def bullet(doc,text):
    p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.left_indent=Inches(.28); p.paragraph_format.first_line_indent=Inches(-.18); p.paragraph_format.space_after=Pt(3); p.paragraph_format.line_spacing=1.15
    font(p.add_run(text),10,False,'222222'); return p
def heading(doc,text,level=1):
    p=doc.add_paragraph(style=f'Heading {level}'); p.paragraph_format.keep_with_next=True
    r=p.add_run(text); font(r,15 if level==1 else 12.3,True,NAVY if level==1 else BLUE); return p
def callout(doc,label,text,fill=LIGHT,color=NAVY):
    t=doc.add_table(rows=1,cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER; c=t.cell(0,0); shade(c,fill); margins(c,150,180,150,180)
    c.text=''; p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.15
    font(p.add_run(label+'  '),10.5,True,color); font(p.add_run(text),10.5,False,'222222'); return t

d=Document(); s=d.sections[0]; s.page_width=Inches(8.5); s.page_height=Inches(11); s.top_margin=s.bottom_margin=Inches(.82); s.left_margin=s.right_margin=Inches(.85)
styles=d.styles
for nm in ['Normal','Heading 1','Heading 2','Heading 3']:
    st=styles[nm]; st.font.name='Microsoft YaHei'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
styles['Normal'].font.size=Pt(10.5)
for nm,sz,before,after in [('Heading 1',15,12,6),('Heading 2',12.3,9,4),('Heading 3',10.8,7,3)]:
    st=styles[nm]; st.font.size=Pt(sz); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(NAVY if nm=='Heading 1' else BLUE); st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after)

# header/footer
hp=s.header.paragraphs[0]; hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT; font(hp.add_run('反倾销税风险核查｜氯氰菊酯'),8.5,False,GRAY)
fp=s.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(fp.add_run('基于易迅数据与公开政策资料｜内部研判'),8,False,GRAY)

p=d.add_paragraph(); p.paragraph_format.space_before=Pt(22); p.paragraph_format.space_after=Pt(4); font(p.add_run('氯氰菊酯贸易流向及反倾销税风险分析报告'),23,True,NAVY)
p=d.add_paragraph(); p.paragraph_format.space_after=Pt(15); font(p.add_run('印度原产直接进口、第三国供应链与绕道风险核查'),13,False,BLUE)
for lab,val in [('数据范围','易迅下载结果5个文件；2024-08-06至2026-08-06'),('分析口径','1,145行原始记录；跨关键词/CAS去重后995条'),('政策基准','商务部公告2025年第24号；海关商品编号2926909013'),('出具日期','2026年8月12日'),('结论等级','直接涉税风险：高；第三国绕道证据：未闭环（观察级）')]:
    p=d.add_paragraph(); p.paragraph_format.space_after=Pt(2); font(p.add_run(lab+'：'),10.5,True,NAVY); font(p.add_run(val),10.5)
d.add_paragraph()
callout(d,'核心结论','本批易迅完整下载数据检出印度原产直接对华3条、52,000千克、出口侧金额325,600美元；其中2026年Tagros一票16,000千克发生在终裁生效后，按48.4%税率测算反倾销税约46,076.80美元（仅为线索测算，须以中国海关计税价格核定）。印度对越南存在44条、约218,627.5千克的上游供货，但下载数据未检出越南或其他第三国原产对华对应记录，尚不能证明绕道。')

heading(d,'一、任务范围与方法')
para(d,'本报告对用户已下载的5个易迅结果文件逐行合并、标准化并跨表去重，检索词包括CYPERMETHRIN TECHNICAL、CIPERMETHRIN及三个受控CAS号（52315-07-8、67375-30-8、1315501-18-8）。去重键综合日期、品名、买卖双方、重量/数量、目的国和原产国，避免同一票因多个关键词重复命中。')
table(d,['文件','原始行数'],[
 ['CAS 52315-07-8（2年）','799'],['CAS 67375-30-8（2年）','203'],['CYPERMETHRIN TECHNICAL（2年）','135'],['CAS 1315501-18-8（2年）','4'],['CIPERMETHRIN（2年）','4'],['合计 / 去重后','1,145 / 995']], [5.2,1.3])
para(d,'易迅网页当前另行验证“CYPERMETHRIN × 目的国中国 × 近三年”为346条、18页。由于下载账户额度限制，本报告不把网页未逐条导出的三年总数冒充为完整明细；数量、企业和税款测算以5个已下载文件的全量记录为基础，网页结果仅用于补充发现近期品名和税号异常。',color=GRAY,size=9.2,after=8)

heading(d,'二、适用反倾销措施')
table(d,['要素','政策内容'],[
 ['原产地','印度'],['产品范围','氯氰菊酯（Cypermethrin / Cypermethrin technical / Cipermethrin）'],['CAS','52315-07-8、67375-30-8、1315501-18-8'],['中国税则/商品编号','税则号列29269090；申报商品编号2926909013'],['临时措施','2025-01-08起提供保证金'],['终裁税','2025-05-07起征，期限5年'],['税率','Gharda 75.7%；UPL 166.2%；Tagros 48.4%；Meghmani 62.0%；Bharat Rasayan 62.0%；Heranba 62.0%；其他印度公司166.2%']], [1.45,5.05])
para(d,'政策来源：商务部公告2025年第24号（终裁）及海关总署公告2025年第3号（商品编号）。反倾销税以海关确定的进口货物计税价格从价计征；进口环节增值税计税基础相应包含反倾销税。')

heading(d,'三、易迅全量结果概览')
table(d,['指标','结果','研判'],[
 ['原始/去重','1,145 / 995','关键词和CAS之间存在重复命中'],['日期范围','2024-08-06—2026-08-06','覆盖措施前后'],['含“TECHNICAL”','315条','需结合CAS和产品形态确认范围'],['疑似标准品/实验室','130条','小样、标准品不宜与商业大货混算'],['印度原产记录','781条；约7,201,989.8 kg','主要流向全球第三国'],['印度→中国','3条；52,000 kg；USD 325,600','直接涉税核查主线'],['第三国原产→中国','0条','下载数据未形成绕道闭环'],['印度→越南','44条；约218,627.5 kg','存在上游供货事实，需补下游对华数据']], [1.5,2.15,2.85])
para(d,'金额字段混合不同国家数据源及币种，尤其越南记录呈现本币大额，不宜直接跨国汇总或折算。重量优先取“重量”，缺失时使用“数量”作为千克口径；因此第三国数量属于贸易情报估算值，最终应回核原始提单/申报单位。',color=GRAY,size=9.2)

heading(d,'四、印度直接对华记录与税款风险')
table(d,['日期','供应商','品名摘要','数量','金额USD','措施状态'],[
 ['2024-11-19','Bayer Vapi Private Limited','CYPERMETHRIN BI TC；CAS 52315-07-8','18,000 kg','115,200','措施前'],
 ['2024-11-28','Bayer Vapi Private Limited','CYPERMETHRIN TECH 94%；CAS 52315-07-8','18,000 kg','115,200','措施前'],
 ['2026-02-28','Tagros Chemicals India Pvt. Ltd.','CYPERMETHRIN TECHNICAL 92%；CAS 52315-07-8','16,000 kg','95,200','终裁后']], [0.72,1.45,2.35,.75,.75,.75])
callout(d,'重点核查票','2026-02-28 Tagros→IPO LTD SHNGHAI，16,000千克。Tagros终裁税率48.4%，按出口侧95,200美元暂测反倾销税=46,076.80美元；若进口增值税税率分别按9%或13%情景测算，因反倾销税增加的增值税约4,146.91美元或5,989.98美元。实际税额必须以中国进口报关单的海关计税价格、成交方式、运保费、关税和适用增值税率重新核定。','FFF2CC',AMBER)
bullet(d,'税号不一致风险：易迅印度出口侧记录使用HS 38089135，政策规定中国进口申报应关注2926909013。出口国税号不同本身不等于违法，但进入中国时若仍按农药制剂类税号申报，可能造成产品范围规避或税款漏征。')
bullet(d,'原产地/生产商风险：Bayer Vapi未列入单独税率企业，若措施生效后仍以其为印度生产商，原则上落入“其他印度公司166.2%”；本批两票发生在措施前，不追溯征税，但可用于企业历史画像。')
bullet(d,'收货人缺失：两票Bayer记录采购商为空，应以提单号、到港日期、承运人和中国舱单反查实际收货人及报关代理。')

heading(d,'五、第三国绕道核查')
heading(d,'5.1 已发现的印度→越南上游链条',2)
para(d,'印度原产对越南共44条、约218,627.5千克，其中按品名识别为technical/TC/CAS明确的31条、约67,376.5千克。主要越南接货方及印度供应方如下：')
table(d,['方向','主体','条数','数量'],[
 ['越南买方','Công ty TNHH UPL Việt Nam','8','144,600 kg'],['越南买方','Công Ty Cổ Phần Kiên Nam','6','24,000 kg'],['越南买方','Công Ty TNHH FUMAKILLA Việt Nam','17','18,275 kg'],
 ['印度供应','Meghmani Organics Limited','5','90,000 kg'],['印度供应','Heranba Industries（名称合并前）','3','37,500 kg'],['印度供应','UPL Mauritius Limited','2','18,600 kg'],['印度供应','Tagros Chemicals India Pvt. Ltd.','5','18,000 kg'],['印度供应','Bharat Rasayan Limited','2','14,000 kg']], [1.05,3.4,.65,1.2])
para(d,'上述主体中，Meghmani、Heranba、Tagros、Bharat Rasayan和UPL均与终裁列名企业直接相关，具备受税产品生产/贸易背景。因此越南链条具有较强“能力与动机”核查价值，但仍缺少越南对华下游发运这一必要证据。')
heading(d,'5.2 闭环证据检验',2)
table(d,['证据环节','本次结果','结论'],[
 ['印度→第三国','存在：印度→越南44条、约218.6吨','成立'],['第三国→中国（同品名/CAS）','5个下载结果去重后0条','未成立'],['主体/时间/数量匹配','无可匹配的越南对华票据','未成立'],['原产地声明异常','未取得中国进口报关单、原产地证','未核实'],['综合结论','仅有上游供货，不具备完整转运闭环','不得认定已绕道']], [2.0,2.8,1.7])
callout(d,'证据等级','第三国绕道：C级观察线索。现有数据能证明印度向越南供应氯氰菊酯，但不能证明越南再出口中国，更不能证明原产地被改报。建议将UPL Vietnam、Kien Nam、Fumakilla Vietnam作为下阶段名单核查对象，而非违法结论。')

heading(d,'六、网页补充发现与分类风险')
para(d,'当前易迅网页三年查询的可见记录还出现BETA CYPERMETHRIN TECHNICAL及HS 38086900/38089199等记录，例如Superform Chemistries、Tagros在2025年末至2026年初对华发运。该类记录提示仅以292690或普通CYPERMETHRIN关键词检索会漏掉β-氯氰菊酯、制剂税号或品名变体。由于这些网页记录尚未完整下载，本报告只作风险提示，不与52,000千克的下载口径合并。')
bullet(d,'扩词：BETA CYPERMETHRIN、ALPHA CYPERMETHRIN、高效氯氰菊酯、CYANO、52315、67375、1315501及常见含量92%/94%。')
bullet(d,'扩税号：同时核查出口侧292690、380891、380869，以及中国进口侧2926909013；不要仅凭HS判断产品是否在措施范围。')
bullet(d,'范围判定核心：原产地印度 + 产品化学身份/CAS + 进口日期；外包装、贸易国或出口国改变不当然改变原产地。')

heading(d,'七、风险分级与核查清单')
table(d,['优先级','对象/票据','核查目标'],[
 ['A1','2026-02-28 Tagros→IPO LTD SHNGHAI 16吨','取得中国报关单、税款缴款书、商品编号、原产地证、合同发票；核定48.4%反倾销税'],
 ['A2','网页所见2025-12至2026-01 Beta Cypermethrin大票','确认CAS和实际产品形态；排查以380869/380891申报规避2926909013'],
 ['B1','UPL Vietnam、Kien Nam、Fumakilla Vietnam','查其对华出口、保税仓出入库、加工记录和原产地证签发'],
 ['B2','Meghmani、Heranba、Tagros、Bharat Rasayan→越南','按发运日期后30—180日匹配越南对华同重量/包装/批号'],
 ['B3','Bayer Vapi历史36吨','反查中国实际收货人及报关行，建立后续进口画像']], [0.6,2.75,3.15])
heading(d,'建议调取字段',2)
for x in ['中国进口报关单号、申报日期、商品编号、规格型号、CAS、原产国、贸易国、启运国、境外发货人、境内收货人、消费使用单位、申报价格及币制。','提单号/主分单号、船名航次、箱号、封志号、装卸港、转运港、承运人、报关企业及报关员。','原产地证签发机构与证书号、印度供应商发票、越南加工工艺/增值比例、仓储和换单记录。','对同一主体采用30/60/90/180日窗口，按数量±5%、包装规格、批号和描述相似度进行匹配。']:
    bullet(d,x)

heading(d,'八、结论')
para(d,'本次对5个易迅下载文件的全部记录分析后，最明确的风险不是已证实的第三国绕道，而是终裁生效后的印度原产直接进口及可能的税号/品名范围错配。2026年Tagros 16吨记录具备具体日期、企业、数量、金额和CAS，应列为第一核查对象。')
para(d,'印度→越南供应链规模较大且涉及多家终裁列名企业，说明越南具备接货和后续加工/分销条件；但现有下载数据没有任何越南或其他第三国原产对华对应票，证据链在第二段中断。报告因此将其定为观察级线索，并明确禁止用上游数量直接推断逃税数量。')

heading(d,'附录：政策与数据来源')
para(d,'1. 中华人民共和国商务部公告2025年第24号：公布对原产于印度的进口氯氰菊酯反倾销调查最终裁定，https://www.mofcom.gov.cn/zcfb/zgdwjjmywg/art/2025/art_f4dda218753844ad8244cd84442c7992.html',size=8.8,color=GRAY)
para(d,'2. 中华人民共和国海关总署公告2025年第3号：明确商品编号2926909013，https://www.mofcom.gov.cn/zcfb/zgdwjjmywg/art/2025/art_891c7244e10b4cbdbe1688e288a18e87.html',size=8.8,color=GRAY)
para(d,'3. 易迅数据：用户下载的5个氯氰菊酯两年结果文件；网页查询“三年、目的国中国、CYPERMETHRIN”用于补充验证。易迅数据为贸易情报来源，可能存在字段缺失、币种混合、出口税号与中国进口税号不一致等限制。',size=8.8,color=GRAY)
para(d,'免责声明：本报告用于风险研判，不构成行政执法结论或法律意见。逃税额和违法性质必须以中国海关原始申报、计税价格、原产地认定和税款缴纳资料为准。',size=8.8,color=RED)

d.save(OUT)
print(OUT)
