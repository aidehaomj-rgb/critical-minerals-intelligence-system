from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT_DIR = Path(r"C:\Users\59809\Documents\关键矿产\outputs\antidumping_evasion_20260809")
OUT = OUT_DIR / "反倾销税偷逃及第三国绕道专项补充报告_20260809.docx"

NAVY = "17365D"
BLUE = "D9EAF7"
PALE = "EEF5FB"
RED = "F4CCCC"
ORANGE = "FCE4D6"
AMBER = "FFF2CC"
GREEN = "D9EAD3"
GRAY = "E7E6E6"
DARK_GRAY = "666666"
WHITE = "FFFFFF"

doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.27)
sec.page_height = Inches(11.69)
sec.top_margin = Inches(0.78)
sec.bottom_margin = Inches(0.72)
sec.left_margin = Inches(0.82)
sec.right_margin = Inches(0.82)
sec.header_distance = Inches(0.35)
sec.footer_distance = Inches(0.35)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.18

for name, size, before, after in [
    ("Heading 1", 16, 14, 7),
    ("Heading 2", 13, 10, 5),
    ("Heading 3", 11.5, 8, 4),
]:
    st = styles[name]
    st.font.name = "Arial"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(NAVY)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True


def set_run_font(run, size=None, color=None, bold=None):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.font.bold = bold


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def margins(cell, top=70, start=85, bottom=70, end=85):
    tcpr = cell._tc.get_or_add_tcPr()
    tc_mar = tcpr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tcpr.append(tc_mar)
    for key, val in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn("w:" + key))
        if node is None:
            node = OxmlElement("w:" + key)
            tc_mar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def cant_split(row):
    trpr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:cantSplit")
    trpr.append(node)


def style_table(table, widths=None, font_size=8.5, header=True, repeat_header=True):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for r_idx, row in enumerate(table.rows):
        cant_split(row)
        if r_idx == 0 and repeat_header:
            trpr = row._tr.get_or_add_trPr()
            tbl_header = OxmlElement("w:tblHeader")
            tbl_header.set(qn("w:val"), "true")
            trpr.append(tbl_header)
        for c_idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins(cell)
            if widths:
                cell.width = Inches(widths[c_idx])
            if header and r_idx == 0:
                shade(cell, NAVY)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(1.5)
                p.paragraph_format.line_spacing = 1.02
                if header and r_idx == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    set_run_font(run, font_size, WHITE if header and r_idx == 0 else None, header and r_idx == 0)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(color)
    rpr.append(underline)
    run.append(rpr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_source(label, url):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("来源：")
    set_run_font(r, 8, DARK_GRAY, True)
    add_hyperlink(p, label, url)
    return p


def add_bullet(text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.42 + level * 0.2)
    p.paragraph_format.first_line_indent = Inches(-0.2)
    p.paragraph_format.space_after = Pt(4)
    p.add_run(text)
    return p


def add_callout(title, text, fill=ORANGE, color="9C5700"):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.cell(0, 0)
    shade(c, fill)
    margins(c, 110, 130, 110, 130)
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(title + "：")
    set_run_font(r, 10.5, color, True)
    r = p.add_run(text)
    set_run_font(r, 10.5, color, False)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_page_field(paragraph, field_text):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field_text
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)


# Header/footer
header = sec.header.paragraphs[0]
header.text = "反倾销税偷逃及第三国绕道专项补充报告｜内部研判材料"
for r in header.runs:
    set_run_font(r, 8.5, DARK_GRAY, False)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = footer.add_run("内部研判材料  ·  2026-08-09")
set_run_font(r, 8, DARK_GRAY, False)

# Cover
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(7)
r = p.add_run("专项风险分析报告")
set_run_font(r, 11, NAVY, True)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(8)
r = p.add_run("反倾销税偷逃及第三国绕道风险")
set_run_font(r, 25, "111111", True)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(18)
r = p.add_run("基于海关处罚、商务部措施与易迅贸易数据的数量—税种—实体—路线交叉核验")
set_run_font(r, 12.5, DARK_GRAY, False)

meta = [
    ("核验日期", "2026年8月9日"),
    ("重点范围", "2024年以来已公开处罚；POM、酞菁蓝、氯氰菊酯、EPDM、KAMAX及PPS排除项"),
    ("核心口径", "处罚决定确认的税额与平台贸易数据情景测算严格分列；平台记录不等于违法认定"),
    ("成果附件", "《反倾销税偷逃风险证据清单_20260809.xlsx》"),
]
for label, value in meta:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(label + "：")
    set_run_font(r, 10.5, NAVY, True)
    r = p.add_run(value)
    set_run_font(r, 10.5, "333333", False)

add_callout(
    "结论状态",
    "已找到3宗近年处罚的明确少缴税款证据；找到1条POM原产地字段直接冲突线索和1条酞菁蓝大体量路线级线索。除历史多晶硅判例外，当前易迅记录尚未形成同一批货‘原产国—第三国—中国’的完整定案闭环。",
    fill=BLUE,
    color=NAVY,
)

doc.add_page_break()

doc.add_heading("一、执行摘要", level=1)
add_callout(
    "硬结论",
    "2024年以来3宗已公开行政处罚合计漏缴税款772,257.60元、罚款338,800元。上海酞菁蓝案明确4,000千克并拆分为反倾销税149,879.24元、进口增值税19,484.30元；EPDM与KAMAX案公开决定没有披露重量，不应估填。",
    fill=ORANGE,
    color="9C5700",
)

p = doc.add_paragraph()
p.add_run("当前第三国绕道证据强弱排序如下：").bold = True
add_bullet("第一位：POM / HOSTAFORM LW15EWX。越南对华25kg记录的平台原产地字段为Vietnam，货物描述却保留“#&DE”；这是单票内部直接矛盾，足以优先调取中国进口报关单、原产地证及越南上游采购记录，但尚不能单独认定逃税。")
add_bullet("第二位：酞菁蓝印度—越南—中国路线。印度向越南检得490条相关记录；越南向中国35票、合计738,620kg同CAS/品类记录，均列“#&VN”。体量大、商品相近，但尚未将两程匹配到同一越南企业、同一型号或同一批次。")
add_bullet("直接税率执行风险：措施实施后，Tagros与Meghmani氯氰菊酯对华记录去重后合计68,000kg。该方向是印度直接出口中国，不是第三国绕道；风险在是否按中国商品编号和列名生产商税率缴纳反倾销税。")
add_bullet("实体与口岸：厦门翔邦/厦门诚运通—海沧海关、上海嘉助—吴淞海关、浙江锐泰—奉化海关均来自处罚文书，是确定性最高的既往风险点。上海恒久百传动、越南Hamakyu等仅为POM核查对象，不代表违法。")

doc.add_heading("二、‘逃什么税’：税种链条与计算口径", level=1)
p = doc.add_paragraph()
p.add_run("反倾销规避通常不只影响反倾销税本身。").bold = True
p.add_run(" 商务部公告明确，反倾销税以海关确定的进口货物计税价格从价计征；进口环节增值税的计税价格包括完税价格、关税和反倾销税。因此少缴反倾销税，还会同步少缴由该反倾销税抬高税基所对应的进口增值税。")
add_source("商务部公告2025年第25号（共聚POM，含征税公式）", "https://www.mofcom.gov.cn/zcfb/zgdwjjmywg/art/2025/art_09eb6be1f50f4cdaa6a36dfbb09bb529.html")

t = doc.add_table(rows=1, cols=4)
for i, h in enumerate(["风险手法", "主要少缴税种", "附带影响", "本报告表述"]):
    t.rows[0].cells[i].text = h
tax_logic = [
    ("伪报非措施原产国/第三国原产地", "反倾销税", "进口增值税差额；如错误享受协定税率，还可能影响关税", "必须核中国申报原产国和原产地证"),
    ("错报生产商，套用较低企业税率", "反倾销税税率差额", "进口增值税差额", "KAMAX案：6.1%错套，实际26%"),
    ("错报HS/商品范围", "反倾销税", "进口增值税差额；可能还影响普通关税", "上海酞菁蓝案已明确拆分"),
    ("低报成交价格、漏报运保杂费", "关税、反倾销税、进口增值税", "三者计税基础均可能被压低", "KAMAX案同时存在EXW/FOB及费用漏报"),
    ("漏填原生产商和适用税率", "反倾销税", "进口增值税差额", "EPDM案公开总漏税额但未拆分"),
]
for row in tax_logic:
    cells = t.add_row().cells
    for i, v in enumerate(row):
        cells[i].text = v
style_table(t, [1.55, 1.35, 1.45, 2.25], 8.5)

p = doc.add_paragraph()
r = p.add_run("情景测算公式：")
set_run_font(r, 10.5, NAVY, True)
p.add_run("反倾销税=平台金额×适用税率；因反倾销税产生的进口增值税增量=反倾销税×13%；合计差额=反倾销税×1.13。平台金额不是中国海关审定完税价格，币种保持原币，不换算为人民币认定税额。")

doc.add_page_break()
doc.add_heading("三、已认定案件：数量、税种、实体与口岸", level=1)
t = doc.add_table(rows=1, cols=6)
for i, h in enumerate(["案件/口岸", "主体", "商品与数量", "手法", "少缴税种及金额", "处罚"]):
    t.rows[0].cells[i].text = h
confirmed_rows = [
    (
        "海沧关缉违字〔2025〕100号\n厦门海沧",
        "进口人：厦门翔邦高分子科技有限公司\n报关行：厦门诚运通报关行有限公司",
        "EPDM\nHS 4002701000\n1票；重量未公开\n涉案货值407,077.89元",
        "未填原生产商中英文名称及反倾销税率",
        "总漏缴58,663.29元。确定涉及反倾销税；是否及多少为进口增值税差额，决定书未拆分",
        "罚款17,600元",
    ),
    (
        "沪吴淞关缉违字〔2024〕15号\n上海吴淞",
        "上海嘉助贸易有限公司",
        "印度原产酞菁蓝\n4,000kg\nCIF 68,000美元",
        "申报3204170090，实际3204170020；应适用30.7%",
        "反倾销税149,879.24元\n进口增值税19,484.30元\n合计169,363.54元",
        "罚款76,200元",
    ),
    (
        "甬奉关缉违字〔2026〕2号\n宁波奉化",
        "浙江锐泰悬挂系统科技有限公司",
        "KAMAX螺栓\nHS 7318151001\n2票；重量未公开\n计税货值2,389,451.03元",
        "把KAMAX S.L.U报为KAMAX GmbH，税率6.1%而实际26%；EXW错报FOB并漏报费用",
        "总漏缴544,230.77元。涉及反倾销税差额、进口增值税，并可能含运保杂费引起的关税等差额；公开决定未正式拆分",
        "罚款245,000元",
    ),
]
for r_idx, row in enumerate(confirmed_rows, start=1):
    cells = t.add_row().cells
    for i, v in enumerate(row):
        cells[i].text = v
    for c in cells:
        shade(c, PALE if r_idx < 3 else ORANGE)
style_table(t, [1.05, 1.25, 1.15, 1.45, 1.75, 0.75], 7.8)

add_source("EPDM处罚决定镜像（页面链接至厦门海关原文）", "https://www.jitinfo.net/Service/Detail?detail=6883ba77-2a27-a778-0091-40de37b92d64")
add_source("上海嘉助披露文件（含吴淞海关处罚明细）", "https://pdf.dfcfw.com/pdf/H2_AN202407191638089496_1.pdf?1721403630000.pdf=")
add_source("KAMAX处罚决定公开转载", "https://m.sohu.com/a/1029239603_121218495")

doc.add_heading("3.1 EPDM案：风险点在生产商字段，不在第三国", level=2)
p = doc.add_paragraph("该案申报日期为2025年3月12日，报关单号370820251000007917，贸易方式为一般贸易。报关企业未填原生产商中英文名称及对应反倾销税率，导致税款未征。中国对原产于美国、韩国和欧盟的EPDM按生产商实施不同税率；2025年期终复审已继续维持调查期间的措施。易迅以“XIAMEN XIANGBANG+EPDM”检索未命中，说明当前平台字段不足以补出重量或历史提单，不能反向推定没有进口。")
add_source("商务部2020年第60号公告（EPDM原措施）", "https://dcj.mofcom.gov.cn/article/zcfb/gpmy/202012/20201203024350.shtml")
add_source("商务部2025年第81号公告（EPDM期终复审）", "https://www.mofcom.gov.cn/zwgk/zcfb/art/2025/art_50c8adab6b72437b84b0f04a69b5ab0d.html")

doc.add_page_break()
doc.add_heading("3.2 上海酞菁蓝案：唯一同时公开重量和税种拆分的近年样本", level=2)
p = doc.add_paragraph("该案最适合用作核税模板：4,000kg、印度原产、申报CIF 68,000美元，税号由3204170090纠正为3204170020，适用“其他印度公司”30.7%反倾销税率。海关明确少缴反倾销税149,879.24元和进口增值税19,484.30元。它证明税号错报可以同时造成反倾销税和进口增值税少缴，但没有第三国环节。")
add_source("商务部公告2023年第8号（印度酞菁类颜料）", "https://dcj.mofcom.gov.cn/article/zcfb/gpmy/202302/20230203393455.shtml")

doc.add_heading("3.3 KAMAX案：生产商错报与完税价格低报叠加", level=2)
p = doc.add_paragraph("两票报关单号分别为223320251000431866（2025-05-15）和310120251019948113（2025-07-23）。处罚认定实际生产商为西班牙KAMAX S.L.U，税率26%，而申报为德国KAMAX GmbH & Co.KG并套用6.1%。同时，实际成交方式为EXW，却申报FOB，漏报1,900欧元和1,155美元费用。仅以2,389,451.03元计税货值和19.9个百分点税率差测算，反倾销税差约475,500.75元、对应进口增值税增量约61,815.10元，合计537,315.85元；与海关认定总漏税544,230.77元相差6,914.92元，差额可能来自漏报费用等，不能把模型拆分当作正式认定。")

doc.add_page_break()
doc.add_heading("四、第三国绕道线索一：POM / HOSTAFORM LW15EWX", level=1)
add_callout(
    "判定",
    "这是目前最具体的第三国核查线索，但数量仅25kg。风险不在‘从越南发货’本身，而在同一条对华记录中结构化原产地字段为Vietnam、商品描述却写有“#&DE”，两字段指向不同原产地。",
    fill=RED,
    color="9C0006",
)

t = doc.add_table(rows=1, cols=6)
for i, h in enumerate(["环节", "日期", "交易主体", "商品", "数量", "原产地信号"]):
    t.rows[0].cells[i].text = h
pom_rows = [
    ("第一程", "2025-10-07", "卖方：Celanese Performance Solutions Switzerland SARL\n买方：CÔNG TY TNHH BỒ CÔNG ANH SÀI GÒN", "HOSTAFORM LW15EWX NATURAL\n25kg/袋", "1,000kg", "Germany"),
    ("第二程", "2026-05-26", "卖方：CÔNG TY TNHH HAMAKYU\n买方：SHANGHAI HENGJIU HUNDRED TRANSMISSION CO., LTD", "HOSTAFORM LW15EWX，Celanese", "25kg", "平台字段Vietnam；货描#&DE"),
    ("同票配料", "2026-05-26", "同上", "用于按2%比例混合的浅蓝色母粒", "2kg", "#&VN"),
]
for idx, row in enumerate(pom_rows):
    cells = t.add_row().cells
    for i, v in enumerate(row): cells[i].text = v
    if idx == 1:
        for c in cells: shade(c, RED)
style_table(t, [0.65, 0.8, 2.15, 1.55, 0.55, 1.0], 7.7)

p = doc.add_paragraph()
p.add_run("涉及数量：").bold = True
p.add_run("第二程对华POM为25kg；同票越南色母粒2kg。第一程德国原产同型号为1,000kg、25kg/袋。25kg等于一袋包装，是可衔接信号，但并非批次闭环。")
p = doc.add_paragraph()
p.add_run("可能少缴税种：").bold = True
p.add_run("如中国进口申报采用越南原产，而经核查该货物实际仍为德国Celanese原产、属于涉案共聚POM且越南未发生足以改变原产地的实质性加工，则可能少缴34.5%反倾销税以及反倾销税对应的进口增值税增量。")
p = doc.add_paragraph()
p.add_run("理论金额：").bold = True
p.add_run("以平台金额字段VND 4,774,382.50为基数，理论反倾销税VND 1,647,161.96、进口增值税增量VND 214,131.06、合计VND 1,861,293.02。该金额是越南数据字段的情景测算，不是中国海关认定税额，也不能直接换算为人民币案值。")

p = doc.add_paragraph()
p.add_run("当前不能定案的原因：").bold = True
p.add_run("第一程和第二程的越南企业不同；没有提单号、箱号、批次号、原产地证及中国进口申报字段；25kg亦可能是样品、合法转售或经实际加工后的货物。必须确认平台的“Vietnam”是否等同中国进口报关原产国，而不是平台自身映射。")
add_source("商务部公告2025年第25号（共聚POM范围及德国Celanese 34.5%）", "https://www.mofcom.gov.cn/zcfb/zgdwjjmywg/art/2025/art_09eb6be1f50f4cdaa6a36dfbb09bb529.html")

doc.add_page_break()
doc.add_heading("五、第三国绕道线索二：印度酞菁蓝—越南—中国", level=1)
add_callout(
    "判定",
    "这是大体量路线级风险池，不是已认定逃税数量。两程品类和CAS相近，但尚未找到同一越南企业既接收印度来料又向中国发出同型号货物的闭环。",
    fill=ORANGE,
    color="9C5700",
)

t = doc.add_table(rows=1, cols=5)
for i, h in enumerate(["方向", "检索结果", "代表性主体", "商品/数量", "证据意义"]):
    t.rows[0].cells[i].text = h
pig_rows = [
    ("印度→越南", "490条", "KILBURN CHEMICALS BHARUCH→U.C.C越南；Indian Chemical Industries→Brenntag Vietnam", "MEGHAFAST BLUE、OCTAFINE BLUE；CAS147-14-8；单日多批720–2,520kg", "证明印度酞菁蓝持续进入越南，但代表主体与第二程未闭合"),
    ("越南→中国", "35票", "YICAI COLOR PLASTIC / Vĩnh Gia→HUNAN YIGAO、DONGGUAN CHENGHAN等", "PIGMENT BLUE 150030；25kg/包；35票合计738,620kg；常见24,000kg/票", "大体量对华流向，记录均标#&VN；需验证越南原产形成过程"),
]
for row in pig_rows:
    cells = t.add_row().cells
    for i, v in enumerate(row): cells[i].text = v
style_table(t, [0.85, 0.7, 2.0, 1.75, 1.6], 8.0)

p = doc.add_paragraph()
p.add_run("涉及数量：").bold = True
p.add_run("越南对华35票合计738,620kg（738.62吨）；平台金额字段合计VND 98,781,668,980。印度对越南的490条是检索结果数量，未全部纳入吨数汇总，避免把不同口径相加。")
p = doc.add_paragraph()
p.add_run("可能少缴税种：").bold = True
p.add_run("如果逐票原产地核查认定这些货物实际仍为印度原产、产品仍在商务部酞菁类颜料措施范围，且生产商无法适用列名低税率，则可能少缴30.7%反倾销税及相应进口增值税。若越南发生了足以改变原产地的实质性加工，则不能因印度来料而当然认定规避。")
p = doc.add_paragraph()
p.add_run("全量极端情景：").bold = True
p.add_run("以VND 98,781,668,980全部按30.7%假设，理论反倾销税VND 30,325,972,376.86、进口增值税增量VND 3,942,376,408.99、合计VND 34,268,348,785.85。该数只用于排序调证资源，不是认定税额。")
p = doc.add_paragraph()
p.add_run("产品范围提醒：").bold = True
p.add_run("对华货描含“Acid stearic Phthalocyanine Crude”等成分表述。需通过配方、实验室检测及税则归类确认其是否仍属于无论是否精制/颜料化的涉案酞菁类颜料，不能只凭CAS号下结论。")
add_source("商务部公告2023年第8号（印度酞菁类颜料）", "https://dcj.mofcom.gov.cn/article/zcfb/gpmy/202302/20230203393455.shtml")

doc.add_page_break()
doc.add_heading("六、直接税率执行风险：印度氯氰菊酯", level=1)
p = doc.add_paragraph("商务部自2025年5月7日起对原产于印度的氯氰菊酯实施最终反倾销措施，税率48.4%—166.2%；海关执行商品编号为2926909013，涉及CAS 52315-07-8、67375-30-8和1315501-18-8。易迅记录显示，措施后列名生产商仍直接向中国发货。直接发运本身合法，关键在进口申报是否如实列明CAS、生产商并缴纳相应反倾销税。")
add_source("商务部公告2025年第24号（氯氰菊酯终裁）", "https://www.mofcom.gov.cn/zcfb/zgdwjjmywg/art/2025/art_f4dda218753844ad8244cd84442c7992.html")
add_source("海关总署公告2025年第3号（氯氰菊酯商品编号）", "https://www.mofcom.gov.cn/zcfb/zgdwjjmywg/art/2025/art_891c7244e10b4cbdbe1688e288a18e87.html")

t = doc.add_table(rows=1, cols=6)
for i, h in enumerate(["生产商", "去重数量", "平台金额", "适用税率", "理论反倾销税+增值税增量", "核查对象"]):
    t.rows[0].cells[i].text = h
cyper_rows = [
    ("TAGROS CHEMICALS INDIA PRIVATE LIMITED", "32,000kg（16,000kg×2）", "USD191,200", "48.4%", "USD104,571.10", "IPO LTD SHNGHAI/未展示买方；中国报关单与缴款书"),
    ("MEGHMANI ORGANICS LIMITED", "36,000kg（同字段两日期按1票去重）", "USD201,600", "62.0%", "USD141,240.96", "买方字段---；先识别中国进口人"),
]
for row in cyper_rows:
    cells = t.add_row().cells
    for i, v in enumerate(row): cells[i].text = v
style_table(t, [1.65, 1.1, 0.85, 0.7, 1.05, 1.7], 8.0)

p = doc.add_paragraph()
p.add_run("结论边界：").bold = True
p.add_run("68,000kg是去重后的高匹配调单量，不是偷逃税数量。平台中的印度出口HS与中国执行商品编号不同，可能只是两国税则细分差异；只有中国进口底单和税款缴款书能证明税率是否执行。")

doc.add_heading("七、PPS排除项：印度发运不等于印度原产，也不当然涉税", level=1)
p = doc.add_paragraph("中国现行PPS反倾销措施针对原产于日本、美国、韩国和马来西亚的进口聚苯硫醚。印度不是该措施所列原产国。因此，“印度向中国发运PPS”本身没有因启运国印度而产生反倾销税风险；真正风险只在货物是否实际为美国等措施原产国、是否通过印度简单分装/转售后改报印度原产，或商品/改性料仍落入措施范围而申报为其他税号。")
add_source("商务部公告2025年第77号（PPS期终复审）", "https://trb.mofcom.gov.cn/myjjdc/art/2025/art_9fa4ea413043469c8794445b2a5ac31f.html")
add_bullet("需要比对美国出口印度与印度出口中国的批次号、牌号、包装、数量和时间窗。此前检得美国RYTON进入印度的具体批次，但未匹配到印度对华同批次，因此没有形成第三国证据。")
add_bullet("印度对华FORTRON/DURAFIDE小票存在跨HS、样品、零件或改性料可能，须先判定商品范围。不能因品牌来自美国/日本企业就推定原产国。")

doc.add_page_break()
doc.add_heading("八、已确认的第三国绕道判例：多晶硅", level=1)
add_callout(
    "证明标准参照",
    "（2018）津02刑初45号公开摘要显示，美国太阳能级多晶硅先运至台湾进行简单加工并取得台湾原产地证明，再进口中国，偷逃税款380余万元，公司被判罚金381万元，4名责任人员获缓刑。该案具备‘原始美国货物—第三地简单加工—原产地证—回流中国—税额’完整链条。",
    fill=GREEN,
    color="274E13",
)
add_source("反倾销税稽查与案例汇总（含多晶硅判例摘要）", "https://www.tradesichuan.com/jmzx/1987.html")
p = doc.add_paragraph("这也说明当前POM与酞菁蓝线索的决定性缺口：不仅要证明第三国节点存在，还要证明第三国加工不足以改变原产地、对华进口申报采用了不实原产地或生产商，并将少缴税额落实到具体报关单。")

doc.add_page_break()
doc.add_heading("九、实体和口岸风险清单", level=1)
t = doc.add_table(rows=1, cols=5)
for i, h in enumerate(["主体/口岸", "对应商品", "已知数量/票数", "风险性质", "建议动作"]):
    t.rows[0].cells[i].text = h
entity_rows = [
    ("厦门翔邦高分子科技有限公司 / 厦门诚运通报关行 / 海沧海关", "EPDM", "1票；重量未公开", "处罚已认定：漏报生产商/税率，漏税58,663.29元", "核同报关行对多税率商品的生产商字段完整性"),
    ("上海嘉助贸易有限公司 / 吴淞海关", "印度酞菁蓝", "4,000kg", "处罚已认定：税号错报，漏税169,363.54元", "扩查同进口人、同供应商、同税号历史申报"),
    ("浙江锐泰悬挂系统科技有限公司 / 奉化海关", "KAMAX螺栓", "2票；重量未公开", "处罚已认定：生产商错报+完税价格低报，漏税544,230.77元", "按KAMAX S.L.U、西班牙发货史扩查生产商证据"),
    ("上海恒久百传动 / 越南Hamakyu", "POM LW15EWX", "25kg+同票色母粒2kg", "A-线索：Vietnam与#&DE冲突", "调2026-05-26中国底单、原产地证、包装批号"),
    ("YICAI COLOR PLASTIC / Vĩnh Gia / HUNAN YIGAO / DONGGUAN CHENGHAN", "酞菁蓝", "35票、738,620kg", "B+路线池：越南原产声明下的持续大票对华", "先按同型号、日期、数量筛批次，再调CO与越南生产记录"),
    ("WUHAN ASIA EUROPE INTERNATIONAL SUPPLY CHAIN MANAGEMENT", "KAMAX S.L.U螺钉", "993.4kg", "供应链基准，不是处罚同批", "核其是否代理其他KAMAX S.L.U对华交易"),
]
for row in entity_rows:
    cells = t.add_row().cells
    for i, v in enumerate(row): cells[i].text = v
style_table(t, [1.85, 1.0, 0.85, 1.75, 1.75], 7.8)

p = doc.add_paragraph()
p.add_run("口岸表述边界：").bold = True
p.add_run("海沧、吴淞、奉化是处罚文书中已发生案件的口岸，并不代表这些口岸整体风险率最高。POM对华买方在上海，但易迅第二程记录未公开中国申报口岸；酞菁蓝路线的买方位于湖南、东莞，也不能据企业所在地推定进口口岸。")

doc.add_page_break()
doc.add_heading("十、下一步需要的数据：按可形成证据闭环排序", level=1)
t = doc.add_table(rows=1, cols=4)
for i, h in enumerate(["优先级", "数据", "核心字段", "可回答的问题"]):
    t.rows[0].cells[i].text = h
data_rows = [
    ("1", "中国进口申报底单+税款缴款书", "报关单号、口岸、收货人、消费使用单位、境外发货人、生产商、原产国、启运国、HS、规格型号、品牌、数量、单价、总价、成交方式、运保杂费、随附单证", "是否错报原产地/生产商/税号；到底少缴哪种税、多少税"),
    ("2", "两程运输单证", "主/分提单号、集装箱号、封志号、船名航次、到离港时间、毛净重、件数、包装、唛头", "能否把原产国→第三国和第三国→中国匹配为同一批货"),
    ("3", "原产地与第三国加工材料", "CO、生产商声明、BOM/配方、领料单、生产批记录、能源/人工/设备、出入库、增值比例、税号变化", "第三国是否发生足以改变原产地的实质性加工"),
    ("4", "合同、发票与资金链", "采购/转售合同、信用证、付款、保险、运费、关联关系、贸易商毛利", "是否仅为转售/低附加值加工；价格是否异常"),
    ("5", "企业历史申报批量数据", "同进口人、同报关行、同境外发货人、同牌号、同生产商、同税号按月聚合", "发现重复错报模式和扩线范围"),
]
for row in data_rows:
    cells = t.add_row().cells
    for i, v in enumerate(row): cells[i].text = v
style_table(t, [0.55, 1.35, 3.15, 2.2], 8.0)

doc.add_page_break()
doc.add_heading("十一、建议立即执行的具体查询", level=1)
for text in [
    "POM：调取2026-05-26上海恒久百传动对应25kg HOSTAFORM LW15EWX的进口报关单、原产地证和税款缴款书；核中国申报原产国是否Vietnam、货物批号及包装是否保留德国Celanese原标；追查越南Hamakyu的上游采购。",
    "酞菁蓝：将越南对华35票按买方、卖方、型号150030、24吨常见票重、日期聚类；优先调2025年5月以后大票的中国底单与CO；再以KILBURN/U.C.C等印度—越南记录按10—90天时间窗、数量和包装反向匹配。",
    "氯氰菊酯：调取Tagros 32吨和Meghmani 36吨对应中国进口底单，核商品编号2926909013、CAS、列名生产商和反倾销税缴款。若税款已足额缴纳，应把该线索降为正常贸易。",
    "KAMAX：以KAMAX S.L.U为生产商关键词扩查宁波、上海、武汉等进口记录，重点识别被申报为德国KAMAX GmbH而实际发货/生产来自西班牙的票。",
    "EPDM：以报关行厦门诚运通为维度，筛查同类多生产商差别税率商品中生产商字段为空、中文/英文名称不一致、税率为0或异常低的申报。",
]:
    add_bullet(text)

doc.add_page_break()
doc.add_heading("十二、最终结论", level=1)
p = doc.add_paragraph()
p.add_run("可以确认的：").bold = True
p.add_run("近年已公开处罚中，三宗案件明确涉及反倾销税执行错误，合计漏税772,257.60元；上海酞菁蓝案同时明确4,000kg、反倾销税和进口增值税。历史多晶硅案则明确证明了经台湾简单加工改变原产地证明、规避对美反倾销措施的第三国路径。")
p = doc.add_paragraph()
p.add_run("可以进一步核查但尚不能定案的：").bold = True
p.add_run("POM 25kg票存在“Vietnam/#&DE”单票冲突，是最具体线索；酞菁蓝35票738.62吨构成路线级高风险池；两者均缺中国申报、原产地证、加工记录和同批物流闭环。")
p = doc.add_paragraph()
p.add_run("不应误判的：").bold = True
p.add_run("印度直接向中国发PPS或氯氰菊酯并不等于绕道。PPS要看实际原产国是否为美国等措施国家；氯氰菊酯要看中国进口申报是否按印度列名生产商税率足额缴税。任何平台金额测算都只能用于风险排序，不能写成逃税认定额。")

doc.add_heading("附录：主要来源", level=1)
sources = [
    ("商务部2025年第25号：共聚POM", "https://www.mofcom.gov.cn/zcfb/zgdwjjmywg/art/2025/art_09eb6be1f50f4cdaa6a36dfbb09bb529.html"),
    ("商务部2025年第24号：氯氰菊酯", "https://www.mofcom.gov.cn/zcfb/zgdwjjmywg/art/2025/art_f4dda218753844ad8244cd84442c7992.html"),
    ("海关总署2025年第3号：氯氰菊酯商品编号", "https://www.mofcom.gov.cn/zcfb/zgdwjjmywg/art/2025/art_891c7244e10b4cbdbe1688e288a18e87.html"),
    ("商务部2023年第8号：印度酞菁类颜料", "https://dcj.mofcom.gov.cn/article/zcfb/gpmy/202302/20230203393455.shtml"),
    ("商务部2020年第60号：EPDM", "https://dcj.mofcom.gov.cn/article/zcfb/gpmy/202012/20201203024350.shtml"),
    ("商务部2025年第81号：EPDM期终复审", "https://www.mofcom.gov.cn/zwgk/zcfb/art/2025/art_50c8adab6b72437b84b0f04a69b5ab0d.html"),
    ("商务部2025年第77号：PPS期终复审", "https://trb.mofcom.gov.cn/myjjdc/art/2025/art_9fa4ea413043469c8794445b2a5ac31f.html"),
    ("上海嘉助披露文件", "https://pdf.dfcfw.com/pdf/H2_AN202407191638089496_1.pdf?1721403630000.pdf="),
    ("易迅数据", "https://dd.data1688.com/custom/search"),
]
for label, url in sources:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    add_hyperlink(p, label, url)

doc.core_properties.title = "反倾销税偷逃及第三国绕道专项补充报告"
doc.core_properties.subject = "数量、税种、实体、口岸与供应链证据核验"
doc.core_properties.author = "Codex"
doc.core_properties.keywords = "反倾销税, 第三国绕道, 原产地, 易迅数据, 海关处罚"

OUT_DIR.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
