import json, os, re, zipfile
from pathlib import Path
from collections import Counter, defaultdict
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION

BASE=Path(__file__).parent
REPORT_DIR=BASE/'矿种专项报告'
REPORT_DIR.mkdir(exist_ok=True)
rows=json.loads((BASE/'analysis_rows.json').read_text(encoding='utf-8'))
summaries=json.loads((BASE/'analysis_summary.json').read_text(encoding='utf-8'))
by=defaultdict(list)
for r in rows: by[r['矿种']].append(r)

BLUE='17365D'; ACCENT='2E75B6'; PALE='D9EAF7'; LIGHT='F2F4F7'; RED='9C0006'; AMBER='9C6500'; GREEN='006100'; GRAY='666666'

def set_cell_shading(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn('w:shd'))
    if shd is None: shd=OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'),fill)

def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tc=cell._tc; tcPr=tc.get_or_add_tcPr(); tcMar=tcPr.first_child_found_in('w:tcMar')
    if tcMar is None: tcMar=OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node=tcMar.find(qn('w:'+m))
        if node is None: node=OxmlElement('w:'+m); tcMar.append(node)
        node.set(qn('w:w'),str(v)); node.set(qn('w:type'),'dxa')

def set_col_widths(table,widths):
    table.autofit=False
    tblPr=table._tbl.tblPr
    tblW=tblPr.find(qn('w:tblW'))
    if tblW is None: tblW=OxmlElement('w:tblW'); tblPr.append(tblW)
    total=int(sum(widths)*1440); tblW.set(qn('w:w'),str(total)); tblW.set(qn('w:type'),'dxa')
    tblInd=tblPr.find(qn('w:tblInd'))
    if tblInd is None: tblInd=OxmlElement('w:tblInd'); tblPr.append(tblInd)
    tblInd.set(qn('w:w'),'120'); tblInd.set(qn('w:type'),'dxa')
    grid=table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for w in widths:
        gc=OxmlElement('w:gridCol'); gc.set(qn('w:w'),str(int(w*1440))); grid.append(gc)
    for row in table.rows:
        for i,(cell,w) in enumerate(zip(row.cells,widths)):
            cell.width=Inches(w); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcW=cell._tc.get_or_add_tcPr().find(qn('w:tcW'))
            if tcW is None: tcW=OxmlElement('w:tcW'); cell._tc.get_or_add_tcPr().append(tcW)
            tcW.set(qn('w:w'),str(int(w*1440))); tcW.set(qn('w:type'),'dxa')
            set_cell_margins(cell)

def style_doc(doc, title, subtitle):
    sec=doc.sections[0]; sec.page_width=Inches(8.5); sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
    sec.header_distance=sec.footer_distance=Inches(.492)
    st=doc.styles['Normal']; st.font.name='Microsoft YaHei'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); st.font.size=Pt(10.5)
    st.paragraph_format.space_after=Pt(6); st.paragraph_format.line_spacing=1.1
    for name,size,color,bef,aft in [('Heading 1',16,ACCENT,16,8),('Heading 2',13,ACCENT,12,6),('Heading 3',12,BLUE,8,4)]:
        s=doc.styles[name]; s.font.name='Microsoft YaHei'; s._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); s.font.size=Pt(size); s.font.color.rgb=RGBColor.from_string(color); s.font.bold=True
        s.paragraph_format.space_before=Pt(bef); s.paragraph_format.space_after=Pt(aft); s.paragraph_format.keep_with_next=True
    header=sec.header.paragraphs[0]; header.text='关键矿产贸易数据穿透分析  |  内部研判材料'; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs: run.font.name='Microsoft YaHei'; run.font.size=Pt(8); run.font.color.rgb=RGBColor.from_string(GRAY)
    foot=sec.footer.paragraphs[0]; foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=foot.add_run('数据来源：易迅数据页面查询 ｜ 生成日期：2026年8月13日'); run.font.name='Microsoft YaHei'; run.font.size=Pt(8); run.font.color.rgb=RGBColor.from_string(GRAY)
    doc.add_paragraph().paragraph_format.space_after=Pt(26)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.LEFT; p.paragraph_format.space_after=Pt(8)
    r=p.add_run(title); r.font.name='Microsoft YaHei'; r.font.size=Pt(24); r.bold=True; r.font.color.rgb=RGBColor.from_string(BLUE)
    p2=doc.add_paragraph(); p2.paragraph_format.space_after=Pt(20); rr=p2.add_run(subtitle); rr.font.name='Microsoft YaHei'; rr.font.size=Pt(12); rr.font.color.rgb=RGBColor.from_string(GRAY)
    return doc

def callout(doc,label,text,color=PALE):
    t=doc.add_table(rows=1,cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER; set_col_widths(t,[6.5]); set_cell_shading(t.cell(0,0),color)
    trPr=t.rows[0]._tr.get_or_add_trPr(); hdr=OxmlElement('w:tblHeader'); hdr.set(qn('w:val'),'true'); trPr.append(hdr)
    p=t.cell(0,0).paragraphs[0]; r=p.add_run(label+'：'); r.bold=True; r.font.color.rgb=RGBColor.from_string(BLUE); p.add_run(text)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

def add_bullets(doc,items):
    for x in items:
        p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.left_indent=Inches(.5); p.paragraph_format.first_line_indent=Inches(-.25); p.paragraph_format.space_after=Pt(5); p.add_run(x)

def add_table(doc,headers,data,widths,font=8):
    t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
    trPr=t.rows[0]._tr.get_or_add_trPr(); hdr=OxmlElement('w:tblHeader'); hdr.set(qn('w:val'),'true'); trPr.append(hdr)
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=str(h); set_cell_shading(c,LIGHT)
        for r in c.paragraphs[0].runs: r.bold=True; r.font.size=Pt(font); r.font.color.rgb=RGBColor.from_string(BLUE)
    for row in data:
        cells=t.add_row().cells
        for i,v in enumerate(row):
            cells[i].text=str(v or '')
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0
                for r in p.runs: r.font.size=Pt(font); r.font.name='Microsoft YaHei'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    set_col_widths(t,widths); return t

def clean(s,n=90):
    s=re.sub(r'\s+',' ',str(s or '')).strip()
    return s if len(s)<=n else s[:n-1]+'…'

def conclusion(s):
    if s['逐条审阅数']==0: return '本轮代表性英文检索词未命中页面记录，不能据此认定不存在贸易活动；建议补充同义词、化合物、合金、设备和HS编码检索。'
    if s['高']:
        return f"发现{s['高']}条高优先级、{s['中']}条中优先级核查线索。高优先级表示品名、目的地节点、主体类型和申报信息存在多项叠加特征，尚不等同于已证实的第三国绕道。"
    if s['中']:
        return f"发现{s['中']}条中优先级核查线索，主要适合做许可证、最终用户及后续流向核对；当前未形成可独立证明转口绕道的闭环。"
    return '当前记录以低风险或暂停期监测为主，未形成明确绕道指向；仍应关注政策恢复、最终用途和再出口变化。'

def top_counts(rs,key,n=8): return Counter(r[key] for r in rs if r.get(key)).most_common(n)

def build_one(s):
    m=s['矿种']; rs=by[m]
    doc=style_doc(Document(),f'{m}贸易数据第三国绕道风险专项分析',f"基于易迅数据页面逐条查看结果 ｜ 检索词：{s['检索词'] or '—'}")
    callout(doc,'核心结论',conclusion(s))
    doc.add_heading('一、查询范围与覆盖情况',level=1)
    add_table(doc,['指标','结果'],[
        ['页面命中',s['总命中数']],['逐条审阅',s['逐条审阅数']],['覆盖说明',s['覆盖范围']],['政策状态',s['政策状态']],['查询时间','2026年8月12—13日'],['数据获取方式','Chrome页面查看；未使用站点下载功能']], [1.5,5.0],9)
    doc.add_heading('二、风险分层结果',level=1)
    add_table(doc,['高优先','中优先','低风险','暂停期监测','排除误命中'],[[s['高'],s['中'],s['低'],s['监测'],s['排除']]],[1.3]*5,9)
    add_bullets(doc,[
        '高/中优先级仅用于确定核查顺序，不代表已经构成走私、规避许可或虚假申报。',
        '仅有“中国原产—第三国进口”记录时，只能识别为第一程候选链；确认绕道还需匹配第三国再出口、提单、集装箱、原产地变化或主体关联关系。',
        '明显与矿产品无关的关键词误命中已单独标记排除。'
    ])
    doc.add_heading('三、主要贸易节点与主体',level=1)
    dest=top_counts(rs,'目的国/地区'); buyers=top_counts(rs,'采购商'); suppliers=top_counts(rs,'供应商')
    if dest: add_table(doc,['主要目的地','记录数'],dest,[4.8,1.7],9)
    if buyers or suppliers:
        pairs=[]
        for i in range(max(len(buyers),len(suppliers))):
            b=buyers[i] if i<len(buyers) else ('',0); sp=suppliers[i] if i<len(suppliers) else ('',0)
            pairs.append([clean(b[0],38),b[1],clean(sp[0],38),sp[1]])
        add_table(doc,['采购方','记录数','供应方','记录数'],pairs,[2.5,.75,2.5,.75],8)
    doc.add_heading('四、优先核查记录',level=1)
    cand=sorted([r for r in rs if r['风险等级'] in ('高','中','监测')], key=lambda x:({'高':0,'中':1,'监测':2}.get(x['风险等级'],3),x['日期']))[:12]
    if cand:
        data=[[r['风险等级'],r['日期'],clean(r['商品描述'],58),clean(r['采购商'],25),r['目的国/地区'],clean(r['主要依据'],45)] for r in cand]
        add_table(doc,['等级','日期','商品/申报描述','采购商','目的地','核查依据'],data,[.55,.78,2.15,1.1,.72,1.2],7)
    else: doc.add_paragraph('本轮无可列示的高、中优先或监测记录。')
    doc.add_heading('五、第三国绕道风险研判',level=1)
    hub=sum(1 for r in rs if r['中转节点']=='是')
    trader=sum(1 for r in rs if r['主体特征']=='贸易/物流/分销')
    vague=sum(1 for r in rs if r['申报信息完整性']=='不足')
    add_bullets(doc,[
        f'节点特征：{hub}条记录进入区域加工或转运型目的地，需结合该地企业产能、进口规模和再出口方向复核。',
        f'主体特征：{trader}条记录涉及贸易、物流或分销型名称；该特征只能提升核查优先级，不能单独证明绕道。',
        f'申报质量：{vague}条记录存在HS缺失、位数不足或品名过宽情形，建议核对材质、纯度、粒径、牌号及最终用途。',
        '闭环条件：应至少再取得第三国后续出口记录，以及主体关联、时间数量匹配、提单/集装箱或原产地变更中的一项以上证据。'
    ])
    doc.add_heading('六、建议核查动作',level=1)
    add_bullets(doc,[
        '核对出口许可证、报关单、合同、发票、装箱单、技术参数及最终用户声明。',
        '对高频第三国采购商建立主体画像，核验注册资本、经营范围、产能、员工和仓储条件。',
        '检索同一商品在目的国向其他市场再出口的时间、重量和价格，进行近似匹配。',
        '对品名宽泛或HS异常样本开展实验室成分、纯度、粒径和用途核验。',
        '对重复记录先按日期、商品、双方主体、数量和金额去重，再进行票级调查。'
    ])
    doc.add_heading('七、局限性说明',level=1)
    doc.add_paragraph('本报告基于易迅数据页面当前可见记录及代表性英文关键词，页面数据库可能存在镜像记录、字段缺失、金额币种不一致、国别转译和数据时滞。报告用于线索排序，不替代海关原始单证、许可证系统、执法调查和法律定性。')
    out=REPORT_DIR/f'{m}_易迅数据第三国绕道风险专项分析.docx'; doc.save(out); return out

def build_total():
    doc=style_doc(Document(),'关键矿产贸易数据第三国绕道风险综合分析报告','26类关键矿产/相关物项 ｜ 易迅数据页面逐条分析 ｜ 2026年度查询范围')
    total=sum(s['逐条审阅数'] for s in summaries); highs=sum(s['高'] for s in summaries); meds=sum(s['中'] for s in summaries)
    callout(doc,'总体判断',f'共逐条审阅{total:,}条页面记录，形成{highs}条高优先级、{meds:,}条中优先级线索。现有数据主要提供第一程贸易特征和核查排序，尚不足以把这些记录直接认定为第三国绕道。')
    doc.add_heading('一、任务与方法',level=1)
    add_bullets(doc,[
        '查询对象：26类关键矿产及其代表性受控物项、化合物、材料或产业链产品。',
        '查询方式：使用Chrome浏览器逐页查看易迅数据结果，不调用站点导出或下载。',
        '逐条字段：日期、商品描述、HS编码、采购商、供应商、数量/重量/金额、目的地和原产地。',
        '风险框架：物项相关性、第三国节点特征、主体角色、申报完整性、政策状态和反向解释因素。',
        '结论等级：高/中优先仅为核查顺序；“确认绕道”需要后续再出口与物流、主体或原产地证据闭环。'
    ])
    doc.add_heading('二、总体数据覆盖',level=1)
    data=[[s['矿种'],s['检索词'],s['总命中数'],s['逐条审阅数'],s['政策状态'],s['高'],s['中'],s['监测']] for s in summaries]
    add_table(doc,['矿种','检索词','命中','审阅','状态','高','中','监测'],data,[.55,2.0,.65,.65,.75,.55,.55,.8],7)
    doc.add_heading('三、关键发现',level=1)
    add_bullets(doc,[
        '钨类记录量最大，供应链覆盖越南、印度、美国等加工与终端市场；重点不是将全部记录视为风险，而是从牌号、HS、主体和节点叠加中筛出许可核验链。',
        '锑、石墨、镓、锗等受控物项出现第三国节点和申报信息不完整样本，适合优先核对最终用户、加工能力和后续再出口。',
        '稀土样本总体数量较少，但部分记录进入贸易/分销主体或区域加工节点，适合做跨矿种主体关联分析。',
        '暂停执行或产业链监测物项不宜按照现行管制直接判定违规，应作为政策恢复后的历史供应链基线。',
        '关键词检索存在成品、工具及非相关描述误命中，报告已标记排除；后续宜转为HS、CAS、牌号和技术参数的组合检索。'
    ])
    doc.add_heading('四、重点目的地与共性风险模式',level=1)
    hubc=Counter(r['目的国/地区'] for r in rows if r['中转节点']=='是')
    add_table(doc,['第三国/地区节点','记录数','初步用途'],[[k,v,'核验加工产能、再出口流向及最终用户'] for k,v in hubc.most_common(12)],[2.3,1.0,3.2],9)
    add_bullets(doc,[
        '品名宽泛或HS缺失：可能掩盖材质、纯度、粒径或牌号，也可能仅是数据库字段不完整。',
        '贸易/物流/分销主体：缺少明显加工能力时，应核验仓储、转售及最终买方。',
        '第三国加工节点：越南、马来西亚、新加坡、韩国、阿联酋等需结合具体行业产能，不能仅凭国别判定风险。',
        '同日同品同主体重复：可能是拆分申报或数据库镜像，必须先去重后再计算票数和规模。',
        '小批量科研或样品：具有参数和许可核验价值，但金额/数量较小可作为反向因素，避免过度判定。'
    ])
    doc.add_heading('五、优先核查清单',level=1)
    cand=sorted([r for r in rows if r['风险等级']=='高'],key=lambda x:(x['矿种'],x['日期']))[:30]
    data=[[r['矿种'],r['日期'],clean(r['商品描述'],50),clean(r['采购商'],22),clean(r['供应商'],22),r['目的国/地区'],clean(r['主要依据'],36)] for r in cand]
    add_table(doc,['矿种','日期','商品','采购商','供应商','目的地','核查依据'],data,[.48,.72,1.55,1.0,1.0,.65,1.1],6.8)
    doc.add_heading('六、调查路径建议',level=1)
    add_bullets(doc,[
        '第一层：政策和物项核验——锁定生效日期、技术参数、管制编码与许可证要求。',
        '第二层：贸易单证核验——回溯报关单、合同、发票、装箱单、运输方式、提运单与集装箱号。',
        '第三层：主体穿透——识别境内供应商、报关企业、货代、第三国进口商及最终用户的股权和业务关联。',
        '第四层：路径闭环——匹配第三国再出口日期、数量、价格、品名、港口和物流单证，验证是否存在改换原产地或目的地。',
        '第五层：执法处置——只有在许可证、申报、实际货物、主观故意及后续流向证据完整后，才进入违法定性。'
    ])
    doc.add_heading('七、数据质量与法律边界',level=1)
    doc.add_paragraph('易迅数据属于商业贸易情报页面，记录可能不是中国出口报关单的完整镜像，也可能出现相同贸易的多国进口端镜像。报告中的记录数不等同于独立票数；风险标签不构成对企业或个人违法的认定。')
    doc.add_heading('附录：矿种报告索引',level=1)
    add_table(doc,['序号','矿种','专项报告文件'],[[i+1,s['矿种'],f"{s['矿种']}_易迅数据第三国绕道风险专项分析.docx"] for i,s in enumerate(summaries)],[.65,.85,5.0],9)
    out=BASE/'关键矿产易迅数据第三国绕道风险综合分析报告.docx'; doc.save(out); return out

paths=[build_one(s) for s in summaries]
total_path=build_total()
zip_path=BASE/'26种关键矿产专项分析报告.zip'
with zipfile.ZipFile(zip_path,'w',zipfile.ZIP_DEFLATED) as z:
    for p in paths: z.write(p,p.name)
print(json.dumps({'total_report':str(total_path),'reports':len(paths),'zip':str(zip_path)},ensure_ascii=False))
