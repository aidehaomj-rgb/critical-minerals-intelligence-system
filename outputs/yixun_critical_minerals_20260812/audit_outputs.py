import json, zipfile, re
from pathlib import Path
from docx import Document

BASE=Path(__file__).parent
files=[BASE/'关键矿产易迅数据第三国绕道风险综合分析报告.docx']+sorted((BASE/'矿种专项报告').glob('*.docx'))
issues=[]
required_total=['任务与方法','总体数据覆盖','关键发现','重点目的地与共性风险模式','调查路径建议','数据质量与法律边界']
required_one=['查询范围与覆盖情况','风险分层结果','主要贸易节点与主体','第三国绕道风险研判','建议核查动作','局限性说明']
for p in files:
    try:
        d=Document(p)
        texts=[x.text.strip() for x in d.paragraphs if x.text.strip()]
        joined='\n'.join(texts)
        req=required_total if '综合分析报告' in p.name else required_one
        missing=[x for x in req if x not in joined]
        if missing: issues.append({'file':p.name,'missing_sections':missing})
        if len(texts)<10: issues.append({'file':p.name,'issue':'paragraph_count_low','count':len(texts)})
        for ti,t in enumerate(d.tables):
            if not t.rows or not t.columns: issues.append({'file':p.name,'table':ti,'issue':'empty_table'})
            for ri,row in enumerate(t.rows):
                if len(row.cells)!=len(t.columns): issues.append({'file':p.name,'table':ti,'row':ri,'issue':'cell_mismatch'})
    except Exception as e: issues.append({'file':p.name,'error':str(e)})

xlsx=BASE/'关键矿产易迅数据第三国绕道风险逐条分析.xlsx'
try:
    with zipfile.ZipFile(xlsx) as z:
        names=set(z.namelist())
        if 'xl/workbook.xml' not in names: issues.append({'file':xlsx.name,'issue':'missing_workbook_xml'})
        for n in names:
            if n.endswith('.xml'):
                text=z.read(n).decode('utf-8','ignore')
                if re.search(r'#(REF!|DIV/0!|VALUE!|NAME\?|N/A)',text): issues.append({'file':xlsx.name,'part':n,'issue':'formula_error'})
except Exception as e: issues.append({'file':xlsx.name,'error':str(e)})

result={'docx_count':len(files),'xlsx_exists':xlsx.exists(),'issues':issues}
(BASE/'output_audit.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(result,ensure_ascii=False))
